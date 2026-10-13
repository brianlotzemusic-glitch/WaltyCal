#!/usr/bin/env python3
"""DMP1 Unit 3: Chords by Ear (TpT 011): build the teacher resource, the intro deck and the listing files.

    python3 gen.py            build everything into this folder
    python3 gen.py --lowres   same, but square images at scale 1 and small page shots (self-check only; --check rejects them)
    python3 gen.py --check    verify the music theory, data, files, page/slide counts, day slots, speaker notes, fonts, label fit,
                              page overflow, banned strings, tpt.json and UPLOAD.md

The owner's own curriculum (Brian Lotze, Digital Music Production 1), reformatted in the Hudson Beat style: a 5-day,
ear-first chords unit in Logic Pro. The DMP1 Restructure defines Unit 3 "Chords by Ear" (5 days, 2/2/1, October, mixing
concept: balance): students hear and play I-IV-V-I and describe what chords feel like before any label, then get the KCSRN
worksheet ("keys, chords, scales and Roman numerals") as names for what they already heard; C major or A minor, white keys only.
The owner's Drive folder "3 - The Building Blocks" supplies the material: Part 1 The Keyboard (teacher notes with the opening
listening activity, Day 1 post), the scale pages of Part 2, Part 3 Chords & Roman Numerals (project sheet and /15 rubric,
teacher notes, classroom post), the Scary Stuff packet (Parts 1-3) and reference sheet, the Key Reference Cards (C major,
A minor, blank) and the owner's KCSRN worksheet. Part 2's melody project and Part 4 Basslines belong to Unit 4 (Bass & Melody
Writing), a separate product. New here (disclosed in the teacher notes): the Day 1 Chord Feel Log and post (the Restructure
lists the ear-first activity as still to be written), the Logic file recipe (Logic sessions are not shipped), the worked
example and the answer keys. Songs are named by title and artist only (find them on YouTube or Spotify). School names and
logos removed; no student work, names or grades; no Logic sessions, audio or video.

Every chord spelling, scale, Roman numeral, key signature and answer in the kit is computed by the small theory engine
below and checked against the owner's cards (verify_theory()).

Outputs: <slug>.pdf (teacher + student resource), <slug>-intro-deck.pptx and its PDF copy, a ZIP of those three
(the TpT product file), a 4-page watermarked preview PDF, cover.png and preview-1..3.png (2000 x 2000).
Renders with Playwright + Chromium (../tools/render.js) and LibreOffice (deck PDF), fonts from ../fonts. $0 AI.
"""
import json
import os
import re
import shutil
import subprocess
import sys
import zipfile
from html import escape

HERE = os.path.dirname(os.path.abspath(__file__))
TPT = os.path.dirname(HERE)
FONTS = os.path.join(TPT, "fonts")
BUILD = os.path.join(HERE, "build")
RENDER_JS = os.path.join(TPT, "tools", "render.js")

SLUG = "chords-by-ear"
PDF = os.path.join(HERE, f"{SLUG}.pdf")
DECK = os.path.join(HERE, f"{SLUG}-intro-deck.pptx")
DECK_PDF = os.path.join(HERE, f"{SLUG}-intro-deck.pdf")
ZIP = os.path.join(HERE, f"{SLUG}.zip")
PREVIEW_PDF = os.path.join(HERE, f"{SLUG}-PREVIEW.pdf")
COVER_PNG = os.path.join(HERE, "cover.png")
PREVIEW_PNGS = [os.path.join(HERE, f"preview-{i}.png") for i in (1, 2, 3)]
UPLOAD = os.path.join(HERE, "UPLOAD.md")
TPT_JSON = os.path.join(HERE, "tpt.json")
OVERFLOW_JSON = os.path.join(BUILD, "overflow.json")
ZIP_DIR = "Chords by Ear - Hudson Beat"

STORE = "Brian Lotze"      # copyright holder and TpT seller
BRAND = "Hudson Beat"      # store line on covers and previews (owner, 4 Oct)
YEAR = 2026
ACCENT = "#0E8C7F"         # this product's single accent: teal (distinct from 008 red, 009 orange, 010 violet)
ACC_LIGHT = "#BFE6E1"      # light tint of the accent, for text on black in the deck
TITLE = "Chords by Ear"
PRICE = "12.00"
UNIT = 3                   # DMP1 Restructure numbering (the Building Blocks docs also say Unit 3)
DAYS_N = 5
MINUTES = 53
LOW_RES = "--lowres" in sys.argv
DISCLOSURE = "Designed with the help of digital and AI tools, and checked by hand."
LISTING_TITLE = "Chord Progressions in Logic Pro | Roman Numerals | High School Music Technology"
PROJECT_NAME = "Your Name — Building Blocks Part 3"   # owner's file-naming rule, kept as written (flagged to the owner)

# Strings that must never appear in any output or in any file in this folder (school branding, LMS). Split so this file never contains them.
BANNED = ["Blind" + " Brook", "BB" + "HS", "blind" + "brook", "Google" + " Classroom"]

# ================================================================ music theory engine (letter-based spelling)
LETTERS = "CDEFGAB"
NAT = dict(C=0, D=2, E=4, F=5, G=7, A=9, B=11)
STEP = {"H": 1, "W": 2, "WH": 3}


def pc(n):
    return (NAT[n[0]] + n.count("#") - n.count("b")) % 12


def name_for(letter, target):
    d = (target - NAT[letter]) % 12
    if d > 6:
        d -= 12
    if abs(d) > 2:
        raise ValueError(f"cannot spell pitch class {target} on {letter}")
    return letter + ("#" * d if d > 0 else "b" * -d)


def up(note, semis, letters):
    """The note `semis` half steps above `note`, spelled on the letter `letters` steps higher."""
    return name_for(LETTERS[(LETTERS.index(note[0]) + letters) % 7], (pc(note) + semis) % 12)


def steps_of(formula):
    return [s.strip() for s in formula.split("-")]


def scale(tonic, formula):
    """Notes of a scale from a W/H/WH formula (WH skips a letter, as in a pentatonic), last note = the octave."""
    out = [tonic]
    for s in steps_of(formula):
        out.append(up(out[-1], STEP[s], 2 if s == "WH" else 1))
    return out


def chord(root, formula):
    """Stack thirds from a root with the owner's half-step formula, e.g. 'R-4-3' -> major."""
    out = [root]
    for n in formula.split("-")[1:]:
        out.append(up(out[-1], int(n), 2))
    return out


QUALITY = {(4, 3): "Major", (3, 4): "Minor", (3, 3): "Dim.", (4, 4): "Aug."}
FORMULA = {"Major": "R-4-3", "Minor": "R-3-4", "Dim.": "R-3-3", "Aug.": "R-4-4"}
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII"]


def quality(notes):
    return QUALITY[tuple((pc(notes[i + 1]) - pc(notes[i])) % 12 for i in range(2))]


def roman(degree, q):
    r = ROMAN[degree]
    return (r if q in ("Major", "Aug.") else r.lower()) + ("°" if q == "Dim." else "")


def diatonic(sc):
    """The 7 triads of a 7-note scale: (Roman numeral, quality, notes), stacking every other scale note."""
    s7 = sc[:7]
    out = []
    for i in range(7):
        notes = [s7[i], s7[(i + 2) % 7], s7[(i + 4) % 7]]
        q = quality(notes)
        out.append((roman(i, q), q, notes))
    return out


def chord_name(notes):
    """The owner's card style: 'C Major', 'A Minor', 'B Dim.'."""
    return f"{notes[0]} {quality(notes)}"


def dash(notes):
    return " - ".join(notes)


def arrow(notes):   # compact spelling for tight spaces
    return "–".join(notes)


def accidentals(notes):
    return [n for n in notes if len(n) > 1]


def semis_between(a, b):
    return (pc(b) - pc(a)) % 12


# ---------------------------------------------------------------- the owner's formulas and keys (Unit Overview, packet, reference sheet)
MAJOR_SCALE = "W - W - H - W - W - W - H"
MINOR_SCALE = "W - H - W - W - H - W - W"
MAJOR_PENT = "W - W - WH - W - WH"
KEYS = {"C major": ("C", MAJOR_SCALE), "A minor": ("A", MINOR_SCALE)}


def key_scale(key):
    t, f = KEYS[key]
    return scale(t, f)


def key_chords(key):
    return diatonic(key_scale(key))


def by_rn(key, rn):
    return next(n for r, q, n in key_chords(key) if r == rn)


# owner's Reference Cards (Key reference Cards folder): RN, Type, Chord, Notes, Feels like
CARD = {
    "C major": [("I", "Major", "C Major", "C - E - G", "Home. Settled. Complete."),
                ("ii", "Minor", "D Minor", "D - F - A", "Gentle tension. Moving somewhere."),
                ("iii", "Minor", "E Minor", "E - G - B", "Wistful. In between."),
                ("IV", "Major", "F Major", "F - A - C", "Lifting. Hopeful. Away from home."),
                ("V", "Major", "G Major", "G - B - D", "Homesick. Wants to resolve to I."),
                ("vi", "Minor", "A Minor", "A - C - E", "Tender. Melancholy. Emotional."),
                ("vii°", "Dim.", "B Dim.", "B - D - F", "Tense. Unstable. Rarely used alone.")],
    "A minor": [("i", "Minor", "A Minor", "A - C - E", "Home. But darker and more unsettled than major."),
                ("ii°", "Dim.", "B Dim.", "B - D - F", "Tense. Rarely used alone."),
                ("III", "Major", "C Major", "C - E - G", "Relative major. Bright moment in a darker key."),
                ("iv", "Minor", "D Minor", "D - F - A", "Heavy. Emotional. Pulling back toward home."),
                ("v", "Minor", "E Minor", "E - G - B", "Less forceful tension than major V."),
                ("VI", "Major", "F Major", "F - A - C", "Warm and open. Common resolution point."),
                ("VII", "Major", "G Major", "G - B - D", "Lifting. Drives back toward i strongly.")],
}
CARD_SCALE_TEXT = {"C major": "C Major  —  W - W - H - W - W - W - H", "A minor": "A Minor  —  W - H - W - W - H - W - W"}
BLANK_CARD = [("I", "Major"), ("ii", "Minor"), ("iii", "Minor"), ("IV", "Major"), ("V", "Major"), ("vi", "Minor"), ("vii°", "Dim.")]
CARD_REMEMBER = ("Remember: Roman numerals describe chord function, not just chord names.",
                 "The I chord always feels like home.  ·  The V chord always feels unresolved.  ·  The vi chord always feels tender or melancholy.",
                 "These feelings are the same in every key — only the note names change.")

# owner's packet, Part 1: half step / whole step exercise (first two done as examples)
HW_EXAMPLES = [("C", "D"), ("C#", "D")]
HW_PAIRS = [("B", "C"), ("E", "F"), ("D", "E"), ("A", "B"), ("F", "G"), ("G#", "A")]
# owner's packet, Part 3: chord construction boxes (C major and A minor pre-filled as examples)
PACKET_CHORDS = [("Major", "R-4-3", "sounds bright and stable", ["C", "F", "G", "D"]),
                 ("Minor", "R-3-4", "sounds darker and more emotional", ["A", "D", "E", "B"]),
                 ("Diminished", "R-3-3", "sounds tense and unstable", ["B", "D"])]
PACKET_GIVEN = {("C", "Major"), ("A", "Minor")}   # the C major and A minor boxes are filled in on the owner's packet
# owner's packet, Roman numerals table (RN, Type, Chord, Feels like)
PACKET_RN = [("I / i", "Major/Minor", "Home chord", "Home. Settled. Complete. The resting point of the key."),
             ("ii / ii°", "Minor/Dim.", "Tension chord", "Gentle tension. Moving somewhere."),
             ("IV / iv", "Major/Minor", "Lifting chord", "Hopeful. Away from home but not lost."),
             ("V / v", "Major/Minor", "Dominant chord", "Homesick. Wants to resolve back to I. Very strong pull."),
             ("vi / VI", "Minor/Major", "Emotional chord", "Tender, melancholy. The emotional heart of the key.")]
# owner's reference sheet: scale formulas and chord types (example root C)
REF_SCALES = [("Major", MAJOR_SCALE, "Bright, resolved, familiar"), ("Natural Minor", MINOR_SCALE, "Darker, more emotional than major"),
              ("Major Pentatonic", MAJOR_PENT, "5-note major scale — bright and simple"),
              ("Minor Pentatonic", "WH - W - W - WH - W", "5-note minor scale — common in blues and rock"),
              ("Harmonic Minor", "W - H - W - W - H - WH - H", "Minor with a raised 7th — dramatic, Eastern quality"),
              ("Melodic Minor", "W - H - W - W - W - W - H", "Minor with raised 6th and 7th ascending"),
              ("Blues Scale", "WH - W - H - H - WH - W", "Minor pentatonic + flat 5 — essential for blues and jazz")]
REF_CHORDS = [("Major", "R-4-3", "C major", ["C", "E", "G"], "Bright, stable, resolved"),
              ("Minor", "R-3-4", "C minor", ["C", "Eb", "G"], "Darker, more emotional"),
              ("Diminished", "R-3-3", "C dim", ["C", "Eb", "Gb"], "Unstable, tense, unresolved"),
              ("Augmented", "R-4-4", "C aug", ["C", "E", "G#"], "Dreamy, unsettled"),
              ("Dominant 7th", "R-4-3-3", "C7", ["C", "E", "G", "Bb"], "Bluesy, wants to resolve"),
              ("Major 7th", "R-4-3-4", "Cmaj7", ["C", "E", "G", "B"], "Lush, jazzy, sophisticated"),
              ("Minor 7th", "R-3-4-3", "Cmin7", ["C", "Eb", "G", "Bb"], "Smooth, melancholy, jazzy")]
REF_RN = [("I", "Major", "Home. Complete. The resting point."), ("ii", "Minor", "Gentle tension. Moving somewhere."),
          ("iii", "Minor", "Wistful. Transitional."), ("IV", "Major", "Lifting. Hopeful. Away from home."),
          ("V", "Major", "Homesick. Wants to resolve back to I."), ("vi", "Minor", "Tender. Melancholy. Emotional center."),
          ("vii°", "Diminished", "Unstable. Tense. Rarely used alone.")]

# ---------------------------------------------------------------- ear-first material (Day 1) and the worked example
# Day 1 Chord Feel Log: the seven white-key chords, by notes only (no names, no numerals: feel it before you name it)
FEEL_ORDER = ["I", "IV", "V", "vi", "ii", "iii", "vii°"]           # C major degrees, played in this order
FEEL_CHORDS = [by_rn("C major", rn) for rn in FEEL_ORDER]
LOOP_MAJOR = ["I", "IV", "V", "I"]          # Restructure: "students play I-IV-V-I ... before they see a single Roman numeral"
LOOP_MINOR = ["i", "VI", "III", "VII"]      # owner's Part 3 teacher notes, Day 2: "Play i-VI-III-VII. Ask: how is that different?"
# Worked example (Hudson Beat, for the Part 3 Logic file recipe): one progression per key, the same four chords in both
EX_PROG = {"C major": ["I", "vi", "IV", "V"], "A minor": ["i", "VI", "III", "VII"]}
# Songs (title and artist only; facts checked against the engine): hook and listening examples
SONGS = [("Twist and Shout", "The Beatles", "D major", ["I", "IV", "V"], "D", ["D", "G", "A"],
          "Day 1 hook: three chords, I, IV and V, round and round"),
         ("Let It Be", "The Beatles", "C major", ["I", "V", "vi", "IV"], "C", ["C", "G", "Am", "F"],
          "Day 4: find the tender vi chord")]


def short_name(notes):
    q = quality(notes)
    return notes[0] + {"Major": "", "Minor": "m", "Dim.": "dim", "Aug.": "aug"}[q]


# ================================================================ the owner's words (lightly edited: LMS, unit labels)
OVERVIEW = ("You have written a melody using scale notes. Now you are going to build the chords that give that melody its harmonic context. "
            "You will use the R-4-3 and R-3-4 formulas to construct chords, learn how Roman numerals describe their function, and build a chord "
            "layer on top of the provided bassline in Logic.")
REQUIREMENTS = ["Complete the chord construction pages and Roman numeral section in the Scary Stuff packet",
                "Complete your Reference Card for your chosen key (C Major or A Minor)",
                "Download the provided Logic file — a bassline is already written",
                "Build the 4 chords from the progression in Logic using the piano roll — one chord per measure",
                "Each chord correctly constructed using R-4-3 (major) or R-3-4 (minor)",
                "Label each chord region in Logic with its Roman numeral",
                "Identify the Roman numeral and function of each chord in your progression in writing (see below)",
                f"Project named: {PROJECT_NAME}",
                "Exported as MP3 and .zip (compressed .logicx file) and submitted to your class page before presenting"]
TIPS = ["Build one chord at a time and play it before moving on. If it sounds wrong, check the formula — count half steps carefully from the root.",
        "In the piano roll, all three notes of a chord should start at exactly the same position and last for the same duration.",
        "Label each region with its Roman numeral as you go. You will need these labels for Part 4.",
        "Your Reference Card is your most important tool here. Keep it open."]
CHECKLIST = ["Chord construction pages and Roman numeral section of the Scary Stuff packet completed",
             "Reference Card completed for your chosen key",
             "4 chords built in Logic using correct formulas",
             "Each chord region labeled with its Roman numeral",
             "Chord function table filled in above",
             f"Project named: {PROJECT_NAME}",
             "Exported as MP3 and .zip and submitted to your class page before presenting"]
RUBRIC = [  # the owner's Part 3 rubric, word for word: levels 2-5, 3 categories = /15
    ("Chord Construction",
     ["Chords built with wrong notes or polyphonic errors. Formula not applied correctly.",
      "Most chords correct but at least one has a construction error. Some formula misapplication.",
      "All chords correctly constructed using R-4-3 or R-3-4. Notes stacked correctly in piano roll.",
      "All chords correct, well-voiced, and labeled with Roman numerals. Demonstrates understanding of the formula."]),
    ("Harmonic Understanding",
     ["Chord function table incomplete or descriptions show no understanding of how chords feel.",
      "Table complete but descriptions are vague or generic (e.g. 'sounds good').",
      "Table complete. Student can describe what each chord feels like in their own words.",
      "Table complete with specific, accurate descriptions. Student clearly understands why each chord was chosen."]),
    ("Documentation & Presentation",
     ["Not submitted or not presented.",
      "Submitted late or named incorrectly. Minimal presentation engagement.",
      "Correctly named and submitted. Presentation made. Reference Card completed.",
      "Correctly named, on time. Student explains chord choices using Roman numeral language."]),
]
LEVELS = [2, 3, 4, 5]
RUBRIC_TOTAL = 5 * len(RUBRIC)       # 15, as on the owner's sheet
RUBRIC_SHORT = ["All chords correct, well-voiced and labeled with Roman numerals.",
                "Specific, accurate descriptions; knows why each chord was chosen.",
                "Named correctly, on time; explains choices in Roman numeral language."]
FORMULAS_BOX = [("Major", "R - 4 - 3", "bright and stable"), ("Minor", "R - 3 - 4", "darker and more emotional"),
                ("Diminished", "R - 3 - 3", "tense and unstable")]
OPENING = [  # owner's Part 1 teacher notes: the opening activity, teacher-led, first 10 minutes
    "Open Logic. Load a piano sound. Make sure students can hear.",
    "Play a single C. Ask: what does that feel like? (Empty. Waiting.) Good.",
    "Play C major (C-E-G). Ask: what changed? Let them answer freely. Happy, full, finished.",
    "Play A minor (A-C-E). Ask: what is different? Dark, heavy, sad.",
    "Play I-IV-V-I slowly. Ask: what happened at the end? Most will say it resolved. Why?",
    "Tell them: the Scary Stuff packet gives names to everything you just felt. That is all music theory is."]
STRUGGLES = [
    ("Rushing to the worksheet", "The most common failure mode in this unit. Spend the full first day just playing chords and describing feelings "
     "in plain language. ‘That one sounds like something good is about to happen.’ ‘That one sounds unfinished.’ Those descriptions are music "
     "theory. The Roman numerals just give them shorter names."),
    ("Finding C on the MPK Mini", "The MPK Mini keys are small and not labeled. Show students the pattern: C always sits just to the left of the "
     "group of 2 black keys. Have every student find C before they do anything else."),
    ("Counting half steps from the root incorrectly", "The single most common chord construction error. Students count the root as step 1. "
     "Demonstrate explicitly: for C major (R-4-3), start from C# as step 1 — C#(1), D(2), D#(3), E(4). Then count 3 from E: F(1), F#(2), G(3). "
     "Walk around and ask every student to verbally count their first chord before they enter it in Logic."),
    ("Chords entered as melody notes", "Students sometimes enter chord notes one at a time horizontally instead of stacked vertically. Show the "
     "piano roll with a correct chord: all three notes start at the same position and have the same length. Any chord that looks like a melody "
     "line is wrong."),
    ("Roman numerals feeling abstract", "Students who built their chords mechanically and did not play them will struggle to connect Roman "
     "numerals to feelings. Return to the opening activity: play I and V on the class speaker and ask what the difference feels like. Then "
     "explain: I is labeled I because it is the home chord, built on the first note of the scale. That is all the Roman numeral tells you."),
    ("Rushing through the packet", "The half/whole step exercises are not optional. They build the mechanical understanding that makes chord "
     "construction possible. Check completion before students open Logic.")]
VOCAB = [("Half step", "moving one key in either direction, including black keys. The smallest possible interval."),
         ("Whole step", "two half steps: skip one key. Most white key pairs are whole steps."),
         ("Scale", "an ordered set of notes built with a formula of whole and half steps. The ingredient list for a key."),
         ("Key", "the home a piece of music settles on: its home note, its scale and the chords built from that scale (this unit: C major or A minor)."),
         ("Chord", "three or more notes played at the same time, built by stacking intervals on a root."),
         ("Root", "the note a chord is built on: the R in R-4-3."),
         ("Major chord", "R-4-3. Sounds bright and stable."),
         ("Minor chord", "R-3-4. Sounds darker and more emotional."),
         ("Diminished chord", "R-3-3. Sounds tense and unstable."),
         ("Roman numeral", "a chord’s name from its position in the scale. Uppercase = major, lowercase = minor, ° = diminished."),
         ("Chord function", "what a chord does in the key: the emotional role it plays (home, lifting, homesick, tender)."),
         ("Balance", "how loud each layer is compared with the others. The part the listener should follow sits on top; the rest steps back.")]
VOCAB_QUIZ = [("Three or more notes played at the same time are a ______.", "chord"),
              ("C to C# is one ______.", "half step"),
              ("The I, IV and V in I–IV–V–I are ______s.", "Roman numeral"),
              ("The note a chord is built on is its ______.", "root")]

# ================================================================ pacing: 5 days of about 53 minutes (DMP1 Restructure, Unit 3: 5 days, 2/2/1)
# block = (minutes, heading, text, deck slides (first, last) or None). Every deck slide belongs to exactly one block (checked).
DAYS = [
    {"n": 1, "kind": "Learning", "title": "Chords by ear: just listen",
     "goal": "Every student has played the seven white-key chords and described each one in plain words, and heard the four-chord loop come home. "
             "No Roman numerals yet.",
     "blocks": [("0–10", "Opening: one note, then chords", "Slides 1–4. Teacher-led, no worksheets yet (steps on page @TN1). Play a single C, then "
                 "C major, then A minor, then the four chords C–F–G–C slowly. Ask what each one feels like and what happened at the end.", (1, 4)),
                ("10–18", "Listen: a real song", "Slide 5: Twist and Shout by The Beatles, the first minute. How many different chords? Which one "
                 "feels like home? When does it want to go back?", (5, 5)),
                ("18–45", "Play and describe", "Slides 6–7. Students load a piano sound in Logic and play the seven chords on the Chord Feel Log "
                 "(page @FEEL), writing what each one feels like in their own words. Then they loop the two four-chord progressions and compare.",
                 (6, 7)),
                ("45–53", "Share-out", "Slide 8. Collect the class’s words on the board and keep them: on Day 4 the Roman numerals give them shorter "
                 "names.", (8, 8))]},
    {"n": 2, "kind": "Learning", "title": "The keyboard and scales",
     "goal": "Every student can name all 7 white keys, tell a half step from a whole step, and build the C major and A natural minor scales.",
     "blocks": [("0–5", "Bright or dark?", "Play the C major scale, then the A natural minor scale. Ask: which sounds brighter? Which sounds darker? "
                 "Do not explain yet.", None),
                ("5–20", "The keyboard", "Slide 9. Students complete Part 1 of the packet (page @PK1): label the white keys, then the half step / "
                 "whole step exercises. Find C first.", (9, 9)),
                ("20–30", "Logic: name every key", "Students load a piano sound and play every white key from C to C, naming each one aloud. Then C "
                 "and D, then C and C#: can you hear the difference between a whole step and a half step?", None),
                ("30–50", "Scales", "Slide 10. Students complete Part 2 of the packet (page @PK2): C major, A natural minor and C major pentatonic, "
                 "playing each scale as they fill it in.", (10, 10)),
                ("50–53", "Exit question", "Which two notes from C major are missing from C major pentatonic? Collect answers at the door.", None)]},
    {"n": 3, "kind": "Project day", "title": "Build chords from formulas",
     "goal": "Chord construction exercises complete. Students can build any major or minor chord from any root and enter it in the piano roll.",
     "blocks": [("0–5", "What notes are those?", "Play C major and A minor again. Ask: now that you know notes, can you figure out what notes those "
                 "are?", None),
                ("5–12", "The project", "Slides 11–13: the project, what to turn in and the rubric. Hand out the project sheet (pages @PROJ).", (11, 13)),
                ("12–35", "Chord construction", "Slide 14. Students work through Part 3 of the packet (page @PK3): major, minor and diminished "
                 "chords. Ask every student to count their first chord aloud before they write it.", (14, 14)),
                ("35–50", "Logic: chords in the piano roll", "Slide 15: demonstrate one chord in the piano roll (page @LOGIC). Students build the "
                 "chords from their packet work as MIDI regions on a new track and play each one to confirm it sounds correct.", (15, 15)),
                ("50–53", "Save + check", "Students save. Exit check: spell one major and one minor chord out loud.", None)]},
    {"n": 4, "kind": "Project day", "title": "Roman numerals + Reference Card",
     "goal": "Reference Card complete. Chords built over the bassline in Logic and labeled with their Roman numerals.",
     "blocks": [("0–15", "Name what you already heard", "Slides 16–17. Play I–IV–V–I: home, away, very away, home. Play i–VI–III–VII: how is that "
                 "different? Name the feelings before naming the numerals, then match the class’s Day 1 words to the numerals. Let It Be by "
                 "The Beatles for the vi chord.", (16, 17)),
                ("15–40", "Packet, KCSRN + Reference Card", "Slide 18. Students complete the Roman numeral section of the packet (page @PK4), "
                 "the KCSRN worksheet for their key (page @KCSRN) and their Reference Card (page @CARDB). Circulate and check that chord spellings "
                 "are correct.", (18, 18)),
                ("40–53", "Chords over the bassline", "Slide 19: the worked example. Students open the Logic file, listen to the bassline twice, "
                 "build the 4 chords (one per measure) and label each region with its Roman numeral.", (19, 19))]},
    {"n": 5, "kind": "Finish + present", "title": "Balance, finish, present",
     "goal": "Every project submitted and presented; every student can point to their I chord and say why it feels like home.",
     "blocks": [("0–8", "Mixing moment: balance", "Slide 20. Solo the chords, then the bass, then play both: which one should the listener "
                 "follow? Turn the chords down until the bass is clear.", (20, 20)),
                ("8–30", "Finish + submit", "Students complete the chord function table on the project sheet, finalize the Logic project, export "
                 "and submit. Circulate and ask every student: point to your I chord. Play it. Does it feel like home?", None),
                ("30–50", "Presentations", "Each student plays their four chords and explains them using Roman numeral language: which chord "
                 "is home, which one wants to go back, and what each one feels like.", None),
                ("50–53", "Close", "Slide 21: point to your I chord. Does it feel like home? That is the test.", (21, 21))]},
]
DECK_DAY_LINES = ["Play seven chords and describe them in your own words; hear I–IV–V–I come home",
                  "Name every key; half steps and whole steps; build the C major and A minor scales",
                  "Build major, minor and diminished chords from formulas; enter them in the piano roll",
                  "Roman numerals = names for feelings; KCSRN + Reference Card; chords over the bassline",
                  "Balance the chords under the bass; finish, submit and present"]

# ================================================================ intro deck (the owner's "Looping in Logic" pattern); (day, kind, title, body, notes)
TS, LIB = SONGS[0], SONGS[1]
SLIDES = [
    (1, "title", "Chords by Ear", f"Feat. {TS[0]}",
     f"Day 1, the first minute. Today is an ear day: no worksheets, no Roman numerals. Have {TS[0]} ({TS[1]}) and, for Day 4, {LIB[0]} "
     f"({LIB[1]}) cued up on YouTube or Spotify from your own account; no audio or video is included in this deck. Open Logic with a piano "
     "sound on the class speakers. Open with a question: why do some songs feel finished at the end and others feel like they are waiting?"),
    (1, "bullets", "What the heck is a chord?", "Three or more notes played at the same time|Notes from the scale, stacked together|"
     "Major sounds bright. Minor sounds dark.|Today: no names. Just your ears.",
     "Plain-English definition, kept short on purpose. The unit's principle is feel it before you name it: students play and describe chords "
     "before they see a single Roman numeral. Tell them the only thing you want today is their honest words. There are no wrong answers."),
    (1, "listen", "Listen: one note, then a chord", "C  ›  C–E–G  ›  A–C–E · teacher plays on a piano sound|Play a single C. Then C major. Then A minor.|"
     "What does one note feel like?|What changed when it became a chord?|What is different about the third one?",
     "The opening activity from the teacher notes, steps 2 to 4. Play a single C: what does that feel like? (Empty. Waiting.) Play C-E-G: what "
     "changed? Let them answer freely: happy, full, finished. Play A-C-E: what is different? Dark, heavy, sad. Write their words on the board "
     "and keep them for Day 4."),
    (1, "listen", "Listen: four chords in a row", "C–E–G  ›  F–A–C  ›  G–B–D  ›  C–E–G · slowly, then on a loop|Play the four chords slowly, then loop them.|"
     "What happened at the end?|Which chord wanted to go somewhere?|Why does the last one feel finished?",
     "Step 5 of the opening activity: play I-IV-V-I slowly, but do not say I-IV-V-I yet. Ask what happened at the end. Most will say it resolved. "
     "Ask why. Then say step 6: the packet gives names to everything you just felt. That is all music theory is."),
    (1, "listen", "Listen: a real song", f"{TS[0]} · {TS[1]}|Teacher: play the first minute from YouTube or Spotify.|"
     "How many different chords do you hear?|Which one feels like home?|When does it want to go back?",
     f"About eight minutes. {TS[0]} by {TS[1]} goes round three chords over and over: they are the same three as the loop they just heard "
     "(I, IV and V, in a different key), but do not name them yet. Students usually hear that one chord is home and another one pulls back to it. "
     "Play it twice if they need it."),
    (1, "feel", "Your turn: play and describe", "",
     "Students open Logic, load a piano sound and play the seven chords on the Chord Feel Log one at a time (MPK Mini or Musical Typing), "
     "writing what each one feels like in their own words and ticking bright or dark. The chords are given by their notes only, all white keys. "
     "Circulate and ask: say more. Sad how? Waiting for what?"),
    (1, "compare", "Same chords, a different home", "",
     "Second half of the log. Students loop the two four-chord progressions. Loop 1 lands on C-E-G; loop 2 starts and feels at home on A-C-E. "
     "Same white keys, nearly the same chords, a different home. Ask: where is home now, and how does it feel compared with loop 1? This is "
     "C major against A minor, but keep the names for later."),
    (1, "bullets", "Your words are music theory", "“That one sounds like something good is about to happen.”|“That one sounds unfinished.”|"
     "Write the class’s words on the board|On Day 4 they get shorter names",
     "The share-out. Collect the best words for each chord and keep them visible (a photo of the board works). These descriptions are music "
     "theory. On Day 4 the Roman numerals simply give them shorter names: home, lifting, homesick, tender."),
    (2, "keys", "The keyboard", "",
     "Day 2 opens with the C major and A natural minor scales played for the class: which sounds brighter? Then this slide. The keyboard repeats "
     "7 white keys and 5 black keys. Find C just to the left of every group of 2 black keys. E to F and B to C are the only white-key half steps. "
     "Students then do Part 1 of the packet."),
    (2, "scales", "Scales: the ingredient list", "",
     "A scale is the ingredient list for a key: it tells you which notes will always sound right together. Walk the C major formula on the "
     "keyboard, then A natural minor. Same seven white keys, a different starting note, a different feel. Students do Part 2 of the packet "
     "and play each scale as they go."),
    (3, "days", "Your project: Chords & Roman Numerals", "",
     "Day 3 starts with C major and A minor played again: now that you know notes, can you figure out what notes those are? Then hand out the "
     "project sheet. The project runs Days 3 to 5: build the chords, name them with Roman numerals, build them over the bassline in Logic, "
     "fill in the chord function table, present."),
    (3, "check", "What you turn in", "Chord pages + Roman numeral section of the packet|Reference Card for your key|"
     "4 chords in Logic, one per measure|Each region labeled with its Roman numeral|Chord function table filled in|"
     "MP3 + .zip, named Your Name — Building Blocks Part 3",
     "The submission checklist. Show where students submit (your class page). Export: File › Bounce › Project or Section, MP3. Compress: in "
     "Finder, Control-click the .logicx project and choose Compress. The file name follows the project sheet."),
    (3, "rubric", "How it’s scored", "",
     f"Three categories, each scored 2 to 5, for a total out of {RUBRIC_TOTAL}. Read the 5 column aloud: every chord correct, well-voiced and "
     "labeled; specific, accurate descriptions of how each chord feels and why it was chosen; on time, named correctly, and chord choices "
     "explained in Roman numeral language."),
    (3, "chordkeys", "Build a chord: count half steps", "",
     "Model counting on the keyboard. The root is not step 1: for C major, C# is 1, D 2, D# 3, E 4; then from E, F is 1, F# 2, G 3. C-E-G. "
     "Minor flips it: 3 then 4. Diminished is 3 and 3. Ask every student to count their first chord out loud before writing it. "
     "Then Part 3 of the packet."),
    (3, "steps", "Watch me: a chord in the piano roll", "New software instrument track, a piano sound|Draw an empty MIDI region, one measure long|"
     "Open the piano roll and add C, E and G|Same start, same length: stacked, not a melody|Play it. Name the region.",
     "The worked example in Logic, about 3 minutes, then students build every chord from their packet as a region. Point at the piano roll: all "
     "three notes start at the same position and last the same length. Anything that looks like a melody line is wrong. Steps and shortcuts "
     "are on the Logic quick reference page."),
    (4, "rnfeel", "Roman numerals = names for feelings", "",
     "Day 4. Play I-IV-V-I and say it: home, away, very away, home. Then play i-VI-III-VII: how is that different? Name the feelings before "
     "naming the numerals. Then match the class's Day 1 words to these tiles. A chord's letter name tells you its notes; its Roman numeral tells "
     "you how it functions. The I chord feels like home in every key."),
    (4, "listen", "Listen: find the tender chord", f"{LIB[0]} · {LIB[1]}|Teacher: play the first minute from YouTube or Spotify.|"
     "Four chords: I, V, vi, IV|Which one is home?|Which one feels tender?",
     f"{LIB[0]} by {LIB[1]} uses I, V, vi and IV in C major: C, G, A minor, F. Students now have names, so ask them to raise a hand on the "
     "tender one (vi, A minor). Point out that vi is the same chord as the A minor home from Day 1: a chord's job depends on the key."),
    (4, "card", "Your Reference Card", "",
     "This is the finished C major card. Students fill in their own on the blank card (and the KCSRN worksheet for their key) after the Roman "
     "numeral section of the packet. It is a living document: they will use it in the next unit, in the EDM project and beyond, so filling it "
     "in carefully now is an investment. Check spellings as you circulate."),
    (4, "roll", "Worked example: chords over the bassline", "",
     "One possible answer for a C major file: the bass plays C, A, F, G, one note per measure, and the chords above it are I, vi, IV and V, "
     "each labeled with its Roman numeral. Every chord's root is the bass note. Students listen to their bassline twice, work out the chords from "
     "the roots, then build and label them. The A minor version is in the kit."),
    (5, "bullets", "Mixing moment: balance", "Solo the chords, then the bass, then both|Which one should the listener follow?|"
     "Turn the chords down until the bass is clear|Turning down is a musical decision",
     "Day 5 opens with the unit's mixing concept. When you layer parts, the ones in the background have to step back. When chords and melody "
     "are the same volume, the melody disappears; today it is the chords and the bass. Turning the chords down is not a mixing trick, it is a "
     "musical decision about what the listener should focus on."),
    (5, "closer", "One last thing…", "Point to your I chord.|Play it. Does it feel like home?|That is the test.",
     "The last three minutes, after the presentations. Close the way the teacher notes do: point to your I chord, play it, does it feel like "
     "home? Name what students did this week: they felt the chords first and named them second. The bassline unit comes next."),
]
DECK_SLIDES = len(SLIDES)

# ================================================================ html scaffolding (Hudson Beat, from 006-010)


def font_css(base=BUILD):
    f = os.path.relpath(FONTS, base)
    faces = [("Oswald", 300, "normal", "oswald-latin-300-normal"), ("Oswald", 500, "normal", "oswald-latin-500-normal"),
             ("Oswald", 700, "normal", "oswald-latin-700-normal"),
             ("Source Sans 3", 400, "normal", "source-sans-3-latin-400-normal"),
             ("Source Sans 3", 400, "italic", "source-sans-3-latin-400-italic"),
             ("Source Sans 3", 600, "normal", "source-sans-3-latin-600-normal"),
             ("Source Sans 3", 700, "normal", "source-sans-3-latin-700-normal")]
    return "".join(f"@font-face{{font-family:'{n}';font-weight:{w};font-style:{s};src:url('{f}/{file}.woff2') format('woff2');}}"
                   for n, w, s, file in faces)


CSS = """
@page { size: 8.5in 11in; margin: 0; }
* { box-sizing: border-box; }
html, body { margin: 0; padding: 0; background: #fff; color: #000; }
body { font-family: 'Source Sans 3', sans-serif; -webkit-print-color-adjust: exact; print-color-adjust: exact; }
.page { width: 8.5in; height: 11in; position: relative; overflow: hidden; page-break-after: always; break-after: page;
        padding: 0.42in 0.5in 0.38in; display: flex; flex-direction: column; }
.page:last-child { page-break-after: auto; break-after: auto; }
.l { font-family: 'Oswald'; font-weight: 300; }
.top { display: flex; justify-content: space-between; align-items: flex-end; border-bottom: 2.2pt solid #000; padding-bottom: 6pt; }
.top .t { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; letter-spacing: .01em; line-height: 1; font-size: 25pt; }
.top .t small { display: block; font-family: 'Oswald'; font-weight: 300; text-transform: none; font-size: 12pt; letter-spacing: 0; margin-top: 4pt; }
.tag { text-align: right; }
.tag .lv { display: inline-block; background: #000; color: #fff; font-family: 'Oswald'; font-weight: 700; font-size: 14pt; padding: 3pt 9pt 4pt; border-radius: 4pt; letter-spacing: .04em; }
.tag .lv.ak { background: ACCENT; }
.tag .ln { font-family: 'Oswald'; font-weight: 300; font-size: 11pt; margin-top: 3pt; }
.name { display: flex; gap: 18pt; font-size: 12pt; margin: 9pt 0 4pt; }
.name span { flex: 1; border-bottom: 1pt solid #000; height: 15pt; }
.name span.d { flex: 0 0 1.6in; }
.main { flex: 1; display: flex; flex-direction: column; min-height: 0; padding-top: 6pt; overflow: hidden; }
.foot { display: flex; justify-content: space-between; font-size: 8.5pt; color: #333; border-top: .8pt solid #000; padding-top: 3pt; }
.wm { position: absolute; left: 50%; top: 50%; transform: translate(-50%, -50%) rotate(-38deg); font-family: 'Oswald'; font-weight: 700;
      font-size: 150pt; color: rgba(0,0,0,.13); letter-spacing: .06em; white-space: nowrap; pointer-events: none; z-index: 9; }
h2.s { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 13pt; margin: 9pt 0 3pt; letter-spacing: .02em; }
h2.s .acc, .acc { color: ACCENT; }
p, li { font-size: 10.5pt; line-height: 1.32; margin: 0 0 4pt; }
ul, ol { margin: 0 0 4pt; padding-left: 15pt; }
.tn p, .tn li { font-size: 11.4pt; line-height: 1.45; margin-bottom: 6pt; }
.tn.big p { font-size: 12.8pt; line-height: 1.5; margin-bottom: 11pt; }
.lines { flex: 1; min-height: 44pt; background: repeating-linear-gradient(to bottom, transparent 0, transparent 21pt, #000 21pt, #000 22pt); }
table.t { border-collapse: collapse; width: 100%; font-size: 9.8pt; margin: 2pt 0 4pt; }
table.t td, table.t th { border: 1pt solid #000; padding: 3pt 5pt; text-align: left; vertical-align: top; line-height: 1.25; }
table.t.tight td { padding: 1.4pt 5pt; font-size: 9.4pt; line-height: 1.2; }
table.t th { font-family: 'Oswald'; font-weight: 500; text-transform: uppercase; font-size: 9.5pt; background: #000; color: #fff; letter-spacing: .04em; vertical-align: middle; }
.cols { display: flex; gap: 20pt; }
.cols > div { flex: 1; min-width: 0; }
.small { font-size: 9pt !important; color: #333; }
.box { border: 1.4pt solid #000; border-radius: 8pt; padding: 7pt 10pt; }
.box.acc-b { border-color: ACCENT; border-width: 2pt; }
.box.inv { background: #000; color: #fff; border-color: #000; }
.k { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 10pt; letter-spacing: .05em; color: ACCENT; display: block; margin-bottom: 2pt; }
.box.inv .k { color: ACC_LIGHT; }
.ln2 { border-bottom: 1pt solid #000; height: 22pt; }
.fl { display: flex; align-items: flex-end; gap: 6pt; font-size: 10.5pt; margin: 0 0 8pt; }
.fl b { font-weight: 600; white-space: nowrap; }
.fl span { flex: 1; border-bottom: 1pt solid #000; height: 14pt; }
.num { list-style: none; padding: 0; margin: 2pt 0 0; }
.num li { display: flex; gap: 10pt; margin: 0 0 9pt; }
.num .n { flex: 0 0 24pt; height: 24pt; border-radius: 12pt; background: #000; color: #fff; font-family: 'Oswald'; font-weight: 700; font-size: 13pt;
          display: flex; align-items: center; justify-content: center; }
.num .hd { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 12pt; letter-spacing: .02em; }
.why { font-style: italic; }
.blk { display: flex; gap: 9pt; margin: 0 0 6pt; font-size: 10.6pt; line-height: 1.32; }
.blk .tm { flex: 0 0 40pt; text-align: center; background: #000; color: #fff; font-family: 'Oswald'; font-weight: 700; font-size: 10pt;
           border-radius: 4pt; padding: 2pt 0 3pt; align-self: flex-start; }
.blk b.h { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 10.5pt; letter-spacing: .02em; }
.day { border: 1.4pt solid #000; border-radius: 8pt; padding: 7pt 10pt 3pt; margin-bottom: 8pt; }
.day .dh { display: flex; justify-content: space-between; align-items: baseline; border-bottom: 1pt solid #000; padding-bottom: 3pt; margin-bottom: 6pt; }
.day .dh b { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 14pt; }
.day .dh span { font-family: 'Oswald'; font-weight: 300; font-size: 11pt; }
.goal { font-size: 10.6pt; margin: 0 0 7pt; }
.chk { display: flex; gap: 8pt; align-items: flex-start; font-size: 11pt; margin: 0 0 8pt; line-height: 1.3; }
table.rub { border-collapse: collapse; width: 100%; }
table.rub td, table.rub th { border: 1.1pt solid #000; padding: 9pt 6pt; vertical-align: top; font-size: 11pt; line-height: 1.38; }
table.rub th { background: #000; color: #fff; font-family: 'Oswald'; font-weight: 500; font-size: 12pt; text-align: center; vertical-align: middle; }
table.rub td.c { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 11pt; width: 15%; vertical-align: middle; }
table.rub td.sc { width: 8%; }
.cover { align-items: center; justify-content: space-between; padding: 0.55in 0.6in 0.5in;
         background: linear-gradient(180deg, #fff 0%, #fff 55%, #ededed 100%); }
.store { font-family: 'Oswald'; font-weight: 500; font-size: 13pt; letter-spacing: .32em; text-transform: uppercase; }
.ctitle { text-align: center; }
.ctitle .a, .ctitle .b, .ctitle .c { white-space: nowrap; }
.ctitle .a { font-family: 'Oswald'; font-weight: 700; line-height: 1; letter-spacing: .02em; text-transform: uppercase; }
.ctitle .b { font-family: 'Oswald'; font-weight: 700; line-height: .98; letter-spacing: .01em; text-transform: uppercase; }
.ctitle .c { font-family: 'Oswald'; font-weight: 300; margin-top: 8pt; }
.cband { width: 100%; background: #000; color: #fff; text-align: center; font-family: 'Oswald'; font-weight: 500; font-size: 14pt;
         letter-spacing: .08em; text-transform: uppercase; padding: 8pt 0 9pt; border-radius: 4pt; white-space: nowrap; }
.cband em { font-style: normal; color: ACCENT; }
.lab { font-family: 'Oswald'; font-weight: 500; font-size: 9pt; text-transform: uppercase; letter-spacing: .05em; text-align: center; }
table.nt { border-collapse: collapse; }
table.nt td { border: 1.2pt solid #000; width: 40pt; height: 30pt; text-align: center; vertical-align: middle; font-family: 'Oswald'; font-weight: 700; font-size: 15pt; }
table.nt td.g { color: #555; }
table.nt td.st { border: none; height: 16pt; font-family: 'Oswald'; font-weight: 500; font-size: 10.5pt; color: ACCENT; }
.cbox { border: 1.2pt solid #000; border-radius: 6pt; padding: 5pt 7pt 6pt; }
.cbox b.cn { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 12pt; display: block; margin-bottom: 3pt; }
table.ch { border-collapse: collapse; width: 100%; }
table.ch td { border: 1.1pt solid #000; height: 26pt; text-align: center; font-family: 'Oswald'; font-weight: 700; font-size: 15pt; width: 33.3%; }
table.ch td.lb { height: 13pt; font-family: 'Source Sans 3'; font-weight: 400; font-size: 8.5pt; border: none; color: #333; }
table.card { border-collapse: collapse; width: 100%; }
table.card td, table.card th { border: 1.2pt solid #000; padding: 6pt 7pt; vertical-align: middle; font-size: 11.5pt; }
table.card th { background: #000; color: #fff; font-family: 'Oswald'; font-weight: 500; text-transform: uppercase; font-size: 11pt; letter-spacing: .04em; }
table.card td.rn { font-family: 'Oswald'; font-weight: 700; font-size: 17pt; text-align: center; width: 11%; }
table.card td.nt { font-family: 'Oswald'; font-weight: 500; font-size: 13pt; white-space: nowrap; }
""".replace("ACCENT", ACCENT).replace("ACC_LIGHT", ACC_LIGHT)

COPY = f"© {YEAR} {STORE} · {BRAND} · For single-classroom use"


def html_doc(pages, title="doc"):
    return (f'<!doctype html><html><head><meta charset="utf-8"><title>{escape(title)}</title>'
            f'<style>{font_css()}{CSS}</style></head><body>{"".join(pages)}</body></html>')


def foot(right):
    return f'<div class="foot"><span>{COPY}</span><span>{right}</span></div>'


def wm(on):
    return '<div class="wm">PREVIEW</div>' if on else ""


def page(sub, tag, tagsub, inner, n, watermark=False, name=False, ak=False, cls="", key=""):
    top = (f'<div class="top"><div class="t">{escape(TITLE)}<small>{sub}</small></div>'
           f'<div class="tag"><div class="lv{" ak" if ak else ""}">{tag}</div><div class="ln">{tagsub}</div></div></div>')
    nm = '<div class="name">Name<span></span>Date<span class="d"></span></div>' if name else ""
    return (f'<section class="page {cls}" data-key="{key}">{wm(watermark)}{top}{nm}<div class="main">{inner}</div>{foot(f"Page {n}")}</section>')


def checkbox(size=12):
    return (f'<svg width="{size}pt" height="{size}pt" viewBox="0 0 {size} {size}" style="flex:0 0 auto;margin-top:1pt">'
            f'<rect x=".7" y=".7" width="{size - 1.4}" height="{size - 1.4}" rx="2" fill="#fff" stroke="#000" stroke-width="1.2"/></svg>')


def t_(x, y, s, size, weight=700, family="Oswald", anchor="middle", fill="#000", extra=""):
    fam = f"'{family}'" if " " in family else family
    return (f'<text x="{x:.1f}" y="{y:.1f}" font-family="{fam}" font-weight="{weight}" font-size="{size}" text-anchor="{anchor}" '
            f'fill="{fill}" {extra}>{escape(s)}</text>')


def svgwrap(vw, vh, body, width="100%"):
    return f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {vw} {vh}" width="{width}" style="display:block">{body}</svg>'


_OSWALD_ADV = {}


def text_w(s, size):
    """Advance width of s in Oswald 700 at size (viewBox units / px / pt), from the font's own metrics."""
    if not _OSWALD_ADV:
        from fontTools.ttLib import TTFont
        f = TTFont(os.path.join(FONTS, "oswald-latin-700-normal.woff2"))
        cmap, hm = f.getBestCmap(), f["hmtx"].metrics
        _OSWALD_ADV["upm"] = f["head"].unitsPerEm
        _OSWALD_ADV.update({chr(c): hm[g][0] for c, g in cmap.items() if c < 0x2200})
    return sum(_OSWALD_ADV[ch] for ch in s) * size / _OSWALD_ADV["upm"]


# ================================================================ drawings (original, code-drawn: black + accent only); every label is logged for the fit check
LABELS = []   # (diagram, text, size, room) : text_w(text, size) must be <= room


def lab(diagram, text, size, room):
    LABELS.append((diagram, text, size, room))


BLACK_AFTER = {"C": ("C#", "Db"), "D": ("D#", "Eb"), "F": ("F#", "Gb"), "G": ("G#", "Ab"), "A": ("A#", "Bb")}
WHITE = "CDEFGAB"


def kb_keys(n_white, ww, wh, bw, bh):
    keys = []
    for i in range(n_white):
        L = WHITE[i % 7]
        keys.append(dict(kind="w", x=i * ww, w=ww, h=wh, s=NAT[L] + 12 * (i // 7), names=(L,)))
    for i in range(n_white - 1):
        L = WHITE[i % 7]
        if L in BLACK_AFTER:
            keys.append(dict(kind="b", x=(i + 1) * ww - bw / 2, w=bw, h=bh, s=NAT[L] + 1 + 12 * (i // 7), names=BLACK_AFTER[L]))
    return keys


def keyboard_svg(n_white=8, hl=None, labels="white", black_labels=False, ww=48, wh=190, bw=28, bh=118, x0=0, y0=0, name="keyboard",
                 wrap=True, width="100%", label_size=None):
    """A keyboard from C: hl = {semitone offset: colour}. Labels sit at the bottom of white keys and (optionally) inside black keys."""
    hl = hl or {}
    keys = kb_keys(n_white, ww, wh, bw, bh)
    out = []
    ls = label_size or round(ww * 0.42)
    for k in keys:
        if k["kind"] != "w":
            continue
        col = hl.get(k["s"])
        out.append(f'<rect x="{x0 + k["x"] + 0.8:.1f}" y="{y0}" width="{k["w"] - 1.6:.1f}" height="{k["h"]}" rx="5" fill="{col or "#fff"}" '
                   f'stroke="#000" stroke-width="2"/>')
        if labels in ("white", "all") or (labels == "hl" and col):
            out.append(t_(x0 + k["x"] + k["w"] / 2, y0 + k["h"] - ls * 0.62, k["names"][0], ls, 700, fill="#fff" if col else "#000"))
            lab(name, k["names"][0], ls, k["w"] - 8)
    for k in keys:
        if k["kind"] != "b":
            continue
        col = hl.get(k["s"])
        out.append(f'<rect x="{x0 + k["x"]:.1f}" y="{y0}" width="{k["w"]}" height="{k["h"]}" rx="4" fill="{col or "#000"}" stroke="#000" stroke-width="1.5"/>')
        if black_labels:
            bs = round(bw * 0.42)
            for j, nm in enumerate(k["names"]):
                out.append(t_(x0 + k["x"] + k["w"] / 2, y0 + k["h"] - 10 - (1 - j) * (bs + 3), nm, bs, 700, fill="#fff"))
                lab(name, nm, bs, k["w"] - 4)
    W, H = n_white * ww, wh
    body = "".join(out)
    return svgwrap(W + 2, H + 2, body, width) if wrap else (body, W, H)


def hero_svg(width="100%"):
    """Cover art: a two-octave keyboard with the home chord C-E-G in the accent, and I-IV-V-I as four tiles with the owner's feelings."""
    ww, wh = 64, 230
    body, W, H = keyboard_svg(14, {0: ACCENT, 4: ACCENT, 7: ACCENT}, labels="hl", ww=ww, wh=wh, bw=38, bh=142, wrap=False, name="hero")
    tiles = [("I", "HOME"), ("IV", "LIFTING"), ("V", "HOMESICK"), ("I", "HOME")]
    gap, ty, th = 18, H + 26, 132
    tw = (W - gap * 3) / 4
    t = []
    for i, (rn, feel) in enumerate(tiles):
        x = i * (tw + gap)
        home = rn == "I"
        t.append(f'<rect x="{x + 1:.1f}" y="{ty}" width="{tw - 2:.1f}" height="{th}" rx="12" fill="{ACCENT if home else "#000"}"/>')
        t.append(t_(x + tw / 2, ty + 72, rn, 66, 700, fill="#fff"))
        t.append(t_(x + tw / 2, ty + 114, feel, 26, 500, fill="#fff", extra='letter-spacing="2"'))
        lab("hero tiles", rn, 66, tw - 20)
        lab("hero tiles", feel, 26, tw - 20 - 2 * len(feel))
    for i in range(3):   # arrows between tiles
        x = (i + 1) * (tw + gap) - gap / 2
        t.append(f'<path d="M{x - 6:.1f} {ty + th / 2 - 9} L{x + 6:.1f} {ty + th / 2} L{x - 6:.1f} {ty + th / 2 + 9}" fill="none" stroke="#000" '
                 f'stroke-width="3" stroke-linecap="round" stroke-linejoin="round"/>')
    return svgwrap(W + 2, ty + th + 2, body + "".join(t), width)


def keys_svg(width="100%"):
    """Deck/packet art: one octave and a bit, every key named, the two white-key half steps marked."""
    ww, wh = 70, 250
    body, W, H = keyboard_svg(8, {}, labels="white", black_labels=True, ww=ww, wh=wh, bw=42, bh=150, wrap=False, name="keys")
    marks = []
    for a in (2, 6):   # E-F and B-C
        x = (a + 1) * ww
        marks.append(f'<path d="M{x - ww / 2:.1f} {H + 22} Q{x:.1f} {H + 52} {x + ww / 2:.1f} {H + 22}" fill="none" stroke="{ACCENT}" stroke-width="4"/>')
        marks.append(t_(x, H + 82, "H", 26, 700, fill=ACCENT))
    return svgwrap(W + 2, H + 92, body + "".join(marks), width)


CHORDKEYS = [("C Major", "R-4-3", "C"), ("A Minor", "R-3-4", "A"), ("B Dim.", "R-3-3", "B")]


def chordkeys_svg(width="100%"):
    """Three keyboards (C to the next F), each with one chord built from the owner's formula, half steps counted above."""
    ww, wh, gap = 34, 150, 46
    one = 11 * ww
    out = []
    for i, (title, formula, root) in enumerate(CHORDKEYS):
        notes = chord(root, formula)
        offs = [NAT[root]]
        for a, b in zip(notes, notes[1:]):
            offs.append(offs[-1] + semis_between(a, b))
        x0 = i * (one + gap)
        body, W, H = keyboard_svg(11, {o: ACCENT for o in offs}, labels="hl", ww=ww, wh=wh, bw=20, bh=92, x0=x0, y0=74, wrap=False,
                                  name="chordkeys", label_size=15)
        out.append(body)
        out.append(t_(x0 + one / 2, 30, title.upper(), 32, 700))
        lab("chordkeys", title.upper(), 32, one)
        cap = f"{formula}  ·  {' - '.join(notes)}"
        out.append(t_(x0 + one / 2, 64, cap, 24, 500, fill=ACCENT))
        lab("chordkeys", cap, 24, one)
    H = 74 + wh + 4
    return svgwrap(3 * one + 2 * gap + 2, H, "".join(out), width)


def voicing(notes):
    """Close root-position voicing: roots C-E in octave 4, F-B in octave 3 (MIDI numbers); bass root in octave 2 (C-E) or 1 (F-B)."""
    r = pc(notes[0])
    base = 60 + r if r <= 4 else 48 + r
    out = [base]
    for nn in notes[1:]:
        out.append(out[-1] + semis_between(notes[len(out) - 1], nn))
    bass = 36 + r if r <= 4 else 24 + r
    return out, bass


ROLL_CH = (53, 67)   # chord rows F3..G4
ROLL_BS = (28, 37)   # bass rows E1..C#2
ROLL_LS = 15         # row label size (note labels inside bars are 2 larger)
ROLL_RH = 25         # px per semitone row; every row label and note label must fit inside one row


def roll_layout(key):
    chords = [by_rn(key, rn) for rn in EX_PROG[key]]
    return [(rn, ch) + voicing(ch) for rn, ch in zip(EX_PROG[key], chords)]


def midi_name(m):
    names = ["C", "C#", "D", "D#", "E", "F", "F#", "G", "G#", "A", "A#", "B"]
    return names[m % 12]


def roll_svg(key="C major", width="100%"):
    """Worked example as a piano roll: Roman numeral tiles, the chord lane (3 stacked notes per measure) and the bass lane."""
    lay = roll_layout(key)
    gut, mw, rh, top = 130, 172, ROLL_RH, 92
    W = gut + 4 * mw + 4
    out = []
    lanes = [("Chords", ROLL_CH, top), ("Bass", ROLL_BS, top + (ROLL_CH[1] - ROLL_CH[0] + 1) * rh + 26)]
    for m, (rn, ch, vo, bass) in enumerate(lay):
        x = gut + m * mw
        home = rn.upper() == "I"
        out.append(f'<rect x="{x + 4}" y="6" width="{mw - 8}" height="70" rx="9" fill="{ACCENT if home else "#000"}"/>')
        out.append(t_(x + mw / 2, 42, rn, 32, 700, fill="#fff"))
        nm = f"{ch[0]} {quality(ch).replace('Dim.', 'dim').lower()}"
        out.append(t_(x + mw / 2, 66, nm, 17, 500, fill="#fff"))
        lab(f"roll {key}", rn, 32, mw - 20)
        lab(f"roll {key}", nm, 17, mw - 20)
    for title, (lo, hi), y0 in lanes:
        n = hi - lo + 1
        for i, mnote in enumerate(range(hi, lo - 1, -1)):
            y = y0 + i * rh
            black = midi_name(mnote).endswith("#")
            out.append(f'<rect x="{gut}" y="{y}" width="{4 * mw}" height="{rh}" fill="{"#ececec" if black else "#fff"}"/>')
            if not black:
                lb = f"{midi_name(mnote)}{mnote // 12 - 1}"
                out.append(t_(gut - 8, y + rh * 0.74, lb, ROLL_LS, 500, anchor="end"))
                lab(f"roll {key} gutter", lb, ROLL_LS, 44)
        out.append(t_(10, y0 + n * rh / 2 + 8, title.upper(), 20, 700, anchor="start"))
        lab(f"roll {key} gutter", title.upper(), 20, gut - 60)
        for m in range(5):
            out.append(f'<line x1="{gut + m * mw}" y1="{y0}" x2="{gut + m * mw}" y2="{y0 + n * rh}" stroke="#000" stroke-width="{1.6 if m in (0, 4) else 1}"/>')
        out.append(f'<rect x="{gut}" y="{y0}" width="{4 * mw}" height="{n * rh}" fill="none" stroke="#000" stroke-width="1.8"/>')
        for m, (rn, ch, vo, bass) in enumerate(lay):
            x = gut + m * mw
            notes = vo if title == "Chords" else [bass]
            for mnote in notes:
                if not lo <= mnote <= hi:
                    raise ValueError(f"{key}: note {mnote} outside the {title} lane")
                y = y0 + (hi - mnote) * rh
                out.append(f'<rect x="{x + 6}" y="{y + 1.5}" width="{mw - 12}" height="{rh - 3}" rx="4" fill="{ACCENT if title == "Chords" else "#000"}"/>')
                out.append(t_(x + 18, y + rh * 0.74, midi_name(mnote), ROLL_LS + 2, 700, anchor="start", fill="#fff"))
                lab(f"roll {key}", midi_name(mnote), ROLL_LS + 2, mw - 30)
    H = lanes[1][2] + (ROLL_BS[1] - ROLL_BS[0] + 1) * rh + 4
    for m in range(4):
        out.append(t_(gut + m * mw + mw / 2, H + 18, f"Measure {m + 1}", 14, 500))
        lab(f"roll {key}", f"Measure {m + 1}", 14, mw - 10)
    return svgwrap(W, H + 26, "".join(out), width)


def verify_labels():
    """Every drawn label fits its key, tile or region (Oswald 700 metrics); row text fits its row height."""
    LABELS.clear()
    hero_svg(), keys_svg(), chordkeys_svg(), roll_svg("C major"), roll_svg("A minor")
    errs = [f"{d}: label '{t}' at {s} is {text_w(t, s):.0f} wide, room {r:.0f}" for d, t, s, r in LABELS if text_w(t, s) > r]
    if not LABELS:
        errs.append("label check found no labels")
    if (ROLL_LS + 2) * 0.72 > ROLL_RH - 6:   # cap height of the largest row label inside one row, with 3px clear above and below
        errs.append("piano-roll labels are too tall for their rows")
    return errs


# ================================================================ page plan
PLAN = ["cover", "howto", "overview", "setup", "logicref", "pacing", "day12", "day34", "day5", "tn1", "tn2", "tn3", "tn4",
        "post1", "feellog", "post2", "post3", "pk1", "pk2", "pk3", "pk4", "kcsrn", "cardC", "cardA", "cardB", "refsheet",
        "proj1", "proj2", "vocab", "checklist", "rubric", "key1", "key2", "example", "terms"]
IX = {k: i + 1 for i, k in enumerate(PLAN)}
PREVIEW_PICKS = ["pacing", "feellog", "cardC", "rubric"]


def rng(a, b):
    return str(a) if a == b else f"{a}–{b}"


def day_slides(d):
    sl = [s for blk in d["blocks"] if blk[3] for s in range(blk[3][0], blk[3][1] + 1)]
    return (min(sl), max(sl)) if sl else None


def slides_label(d):
    r = day_slides(d)
    return f"Intro deck slides {rng(*r)}" if r and r[0] != r[1] else (f"Intro deck slide {r[0]}" if r else "No slides")


def fill(s):
    """Page placeholders in the day plans are written @NAME so they never collide with ordinary words."""
    pages = {"PROJ": f"{IX['proj1']}–{IX['proj2']}", "FEEL": IX["feellog"], "TN1": IX["tn1"], "PK1": IX["pk1"], "PK2": IX["pk2"],
             "PK3": IX["pk3"], "PK4": IX["pk4"], "KCSRN": IX["kcsrn"], "CARDB": IX["cardB"], "LOGIC": IX["logicref"]}
    out = re.sub(r"@([A-Z0-9]+)", lambda m: str(pages[m.group(1)]), s)
    if "@" in out:
        raise ValueError(f"unfilled placeholder in: {s}")
    return out


def H2(s, first=False):
    st = ' style="margin-top:2pt"' if first else ""
    return f'<h2 class="s"{st}>{s}<span class="acc"> ·</span></h2>'


SONG_NOTE = "Songs are named by title and artist only. Play them from your own YouTube or Spotify access; no audio, video or lyrics are included."
COVER_A = "Chord Progressions in Logic Pro"
COVER_B = "Chords by Ear"
COVER_C = f"DMP1 · Unit {UNIT} · hear I–IV–V–I, then name it · C major + A minor"
COVER_BAND = f"Grades 9–12 · {DAYS_N} days · {len(PLAN)} pages + {DECK_SLIDES}-slide deck"
# cover sizes (pt on the PDF cover, px on the square); fit checked against the available width with Oswald 700 metrics
PDF_COVER = dict(a=30, b=76, c=15, band=14, room=7.3 * 72)
SQ_COVER = dict(a=52, b=136, c=29, band=26, room=940)


def cover_page():
    s = PDF_COVER
    band = COVER_BAND.replace(" · ", " <em>·</em> ")
    return (f'<section class="page cover" data-key="cover"><div class="store">{BRAND}</div>'
            f'<div style="width:7.2in">{hero_svg()}</div>'
            f'<div class="ctitle"><div class="a" style="font-size:{s["a"]}pt">{COVER_A}</div><div class="b" style="font-size:{s["b"]}pt">{COVER_B}</div>'
            f'<div class="c" style="font-size:{s["c"]}pt">{COVER_C}</div></div>'
            f'<div class="cband">{band}</div></section>')


def deck_day_line():
    parts = []
    for d in DAYS:
        r = day_slides(d)
        parts.append(f"Day {d['n']} (slide{'s' if r[0] != r[1] else ''} {rng(*r)})")
    return ", ".join(parts[:-1]) + " and " + parts[-1]


def howto_page(n):
    parts = [("Teacher pages", "howto", "tn4", "How to use, overview, setup + the Logic file recipe, Logic reference, pacing, day plans, teacher notes"),
             ("Student pages", "post1", "rubric", "Classroom posts, Chord Feel Log, Scary Stuff packet (Parts 1–3), KCSRN, Reference Cards, "
              "reference sheet, project sheet, vocabulary, checklist, rubric"),
             ("Answer keys + example", "key1", "example", "Packet and KCSRN answers, Feel Log sample words, a worked example in both keys"),
             ("Intro deck (separate file)", None, None, f"{DECK_SLIDES} slides, .pptx + PDF, speaker notes on every slide")]
    trs = "".join(f"<tr><td><b>{a}</b></td><td style='white-space:nowrap'>{('pages ' + str(IX[f]) + '–' + str(IX[l])) if f else 'deck'}</td><td>{u}</td></tr>"
                  for a, f, l, u in parts)
    inner = (f'<div class="cols"><div>'
             f'{H2("What this is", True)}'
             f'<p>A complete {DAYS_N}-day chords unit for high school music technology, taught ear-first. On Day 1 students only listen and play: '
             f'seven chords, two four-chord loops and a real song, described in their own words. Day 2 gives them the keyboard and the C major and '
             f'A minor scales. Days 3 and 4 build chords from formulas and name them with Roman numerals (the KCSRN worksheet and a Reference Card). '
             f'Then they build four chords over a bassline in Logic Pro and present on Day {DAYS_N}.</p>'
             f'<p>This is Unit {UNIT} of my Digital Music Production 1 course. All students work in C major or A minor, white keys only. No music '
             f'reading is needed. Unit 4, Bass &amp; Melody Writing, is the next unit and a separate resource.</p>'
             f'{H2("Quick start")}<ol>'
             f'<li>Read the overview and pacing (pages {IX["overview"]}–{IX["pacing"]}) and the teacher notes (pages {IX["tn1"]}–{IX["tn4"]}).</li>'
             f'<li>Make the Part 3 Logic file: a 4-measure bassline (10 minutes, recipe on page {IX["setup"]}).</li>'
             f'<li>Cue up Twist and Shout and Let It Be (The Beatles) on YouTube or Spotify.</li>'
             f'<li>Practice the opening activity once (page {IX["tn1"]}).</li>'
             f'<li>Print the student pages (pages {IX["post1"]}–{IX["rubric"]}): the packet, KCSRN and cards as one set per student.</li></ol>'
             f'{H2("The intro deck")}'
             f'<p>Every slide has a day: {deck_day_line()}. The day plans name the slide for each block. Every slide has speaker notes.</p>'
             f'<p class="small">Fonts: the deck uses Oswald and Source Sans 3 (free on Google Fonts, built into Google Slides). If PowerPoint or Keynote '
             f'shows a different font, install the two fonts or present the PDF copy, which shows the intended look.</p>'
             f'</div><div>'
             f'{H2("What’s inside", True)}<table class="t"><tr><th>Part</th><th>Pages</th><th>Use</th></tr>{trs}</table>'
             f'<div class="box acc-b" style="margin-top:6pt"><span class="k">About the songs</span><p style="margin:0">{SONG_NOTE}</p></div>'
             f'{H2("Make it fit your class")}'
             f'<ul><li><b>File names:</b> the project sheet keeps my naming rule, “{PROJECT_NAME}” (The Building Blocks is my name for this '
             f'theory sequence). Change it to suit your class.</li>'
             f'<li><b>More accessible / more advanced:</b> two variations on page {IX["tn4"]}.</li>'
             f'<li><b>No MIDI keyboard?</b> Musical Typing (Command-K) or the Pencil tool works for every step.</li>'
             f'<li><b>Short on time?</b> Run Days 1 and 2 as one day only if your students already know the keyboard; never skip the ear day.</li></ul>'
             f'</div></div>')
    return page("How to use this resource", "START HERE", "Teacher", inner, n, key="howto")


def overview_page(n):
    inner = (f'<div class="cols"><div>'
             f'{H2("What students make", True)}'
             f'<p>A Chord Feel Log in their own words, the Scary Stuff packet (keyboard, scales, chords, Roman numerals), the KCSRN worksheet and a '
             f'Reference Card for their key. Then a short Logic project: four chords built over a provided bassline, one per measure, each '
             f'labeled with its Roman numeral, plus a chord function table that says what each chord feels like.</p>'
             f'{H2("What this unit is really about")}'
             f'<p>What a chord sounds like and feels like. The I chord feels like home. The V chord feels like homesick. The vi chord feels like '
             f'something tender or melancholy. Major sounds bright; minor sounds dark. Chords are just notes from the scale stacked together. '
             f'Roman numerals are names for feelings, not abstract theory.</p>'
             f'<p>Feel it before you name it: students play I–IV–V–I and describe how it feels before they know what those symbols mean. The '
             f'KCSRN worksheet then arrives as “now name what you already heard,” not as a cold fill-in exercise.</p>'
             f'{H2("Mixing concept: balance")}'
             f'<p>When you layer parts, the background has to step back. When chords and melody are the same volume, the melody disappears. Turning '
             f'the chords down is not a mixing trick; it is a musical decision about what the listener should focus on.</p>'
             f'</div><div>'
             f'{H2("Goals", True)}<ul>'
             f'<li>Describe what a chord feels like in plain language before naming it</li>'
             f'<li>Name every white key and tell a half step from a whole step</li>'
             f'<li>Build the C major and A natural minor scales from their formulas</li>'
             f'<li>Build major, minor and diminished chords with R-4-3, R-3-4 and R-3-3</li>'
             f'<li>Use Roman numerals to describe chord function in C major or A minor</li>'
             f'<li>Build and label four chords over a bassline in Logic</li></ul>'
             f'{H2("At a glance")}<table class="t">'
             f'<tr><td><b>Grades</b></td><td>9–12 (no music reading needed)</td></tr>'
             f'<tr><td><b>Length</b></td><td>{DAYS_N} periods of about {MINUTES} minutes: listen, keyboard and scales, chords, Roman numerals, present</td></tr>'
             f'<tr><td><b>Keys</b></td><td>C major or A minor, white keys only</td></tr>'
             f'<tr><td><b>Software</b></td><td>Logic Pro: software instrument tracks, MIDI regions, the piano roll, region names, volume, export</td></tr>'
             f'<tr><td><b>Product</b></td><td>Packet, KCSRN and Reference Card; four labeled chords over a bassline; chord function table; MP3 + compressed project</td></tr>'
             f'<tr><td><b>Assessment</b></td><td>3-category rubric, /{RUBRIC_TOTAL} (page {IX["rubric"]})</td></tr></table>'
             f'{H2("Standards")}'
             f'<p>Supports the National Core Arts Standards for Music (Music Technology, high school): analyzing how musical elements are used '
             f'(MU:Re7.2), developing musical ideas (MU:Cr2.1) and presenting work (MU:Pr6.1).</p>'
             f'</div></div>')
    return page("Unit overview", "OVERVIEW", f"DMP1 · Unit {UNIT}", inner, n, key="overview")


def setup_page(n):
    c_bass = " – ".join(f"{midi_name(b)}" for _, _, _, b in roll_layout("C major"))
    a_bass = " – ".join(f"{midi_name(b)}" for _, _, _, b in roll_layout("A minor"))
    recipe = [
        "New project, one Software Instrument track with a bass sound (Library: Bass). Tempo between 80 and 100 BPM is comfortable.",
        f"Draw a 4-measure MIDI region. In the piano roll, add one whole note per measure, the root of each chord. C major example: "
        f"{c_bass} (low octave). A minor example: {a_bass}.",
        "Turn on Cycle over the 4 measures so it loops.",
        "Optional, as in my file: arrangement markers with chord labels visible.",
        "Leave the chord track empty: students build it. Save it as a template or share the project with your class.",
    ]
    owner_setup = ["Same key as the scales work (C major or A minor)", "Bassline clearly following chord roots so students can infer the "
                   "progression from the bass", "Arrangement markers with chord labels visible", "No chord track — leave that blank for students to build"]
    inner = (f'<div class="cols"><div>'
             f'{H2("You need", True)}<ul>'
             f'<li><b>A Mac with Logic Pro</b> and <b>headphones</b> for each student; a MIDI keyboard (an MPK Mini or similar) if you have one.</li>'
             f'<li><b>A piano sound on the class speakers</b> for the Day 1 opening and the Day 4 Roman numeral demo.</li>'
             f'<li><b>Projector</b> for the deck; <b>two songs</b> from your own YouTube or Spotify access: Twist and Shout and Let It Be (The Beatles).</li>'
             f'<li><b>Printed student pages</b>, one set each: Chord Feel Log, packet, KCSRN, Reference Cards (C major, A minor, blank), '
             f'reference sheet, project sheet.</li>'
             f'<li><b>The Part 3 Logic file</b> with a bassline (recipe on the right).</li>'
             f'<li><b>A place to submit</b> an MP3 and a .zip.</li></ul>'
             f'{H2("Before Day 1")}'
             + "".join(f'<div class="chk">{checkbox()}<span>{c}</span></div>' for c in [
                 "Open Logic with a piano sound and test the class speakers.",
                 f"Practice the opening activity (page {IX['tn1']}) and cue up the songs.",
                 "Check that students can play chords: MPK Mini, or Musical Typing (Command-K).",
                 "Make and test the Part 3 Logic file; decide C major, A minor or both.",
                 "Print the student pages; set up the submission spot."]) +
             f'</div><div>'
             f'<div class="box acc-b"><span class="k">Make the Part 3 Logic file (10 minutes)</span>'
             f'<p style="margin:0 0 3pt">I give students a Logic file with the bassline already written. Logic projects are not included here, so '
             f'build your own once. My recommended setup:</p><ul style="margin:0 0 4pt">' + "".join(f"<li>{x}</li>" for x in owner_setup) + '</ul>'
             f'<p style="margin:0 0 3pt"><b>Recipe:</b></p><ol style="margin:0">' + "".join(f"<li>{x}</li>" for x in recipe) + '</ol>'
             f'<p class="small" style="margin:4pt 0 0">The examples use the same four chords in both keys: I–vi–IV–V in C major and i–VI–III–VII in '
             f'A minor (worked example on page {IX["example"]}). Any four chords from the Reference Card work.</p></div>'
             f'<div class="box" style="margin-top:8pt"><span class="k">What comes next</span><p style="margin:0">Unit 4, Bass &amp; Melody '
             f'Writing, picks up here: the rule behind every bassline, and a short melody over the chords students just built.</p></div>'
             f'</div></div>')
    return page("Setup and requirements", "SETUP", "Teacher", inner, n, key="setup")


LOGIC_REF = [("New instrument track", "Track › New Software Instrument Track (Option-Command-S). Open the Library (Y) and pick a piano or "
              "electric piano for chords, a bass sound for the bassline."),
             ("Play without a keyboard", "Window › Show Musical Typing (Command-K) turns the computer keys into a piano. Z and X move the octave."),
             ("Empty MIDI region", "With the Pencil tool, click in the track’s lane, or Control-click the lane and choose Create MIDI Region. Drag "
              "its right edge to make it one measure long."),
             ("Piano roll", "Double-click the region, or press P. With the Pencil tool, click to add a note; drag its right edge to set the length."),
             ("Stack a chord", "All three notes start at the same position and last the same length. Select the notes and press Q to quantize them "
              "to the grid."),
             ("Copy a chord", "Option-drag a region to copy it; then drag its notes up or down in the piano roll to make the next chord."),
             ("Name a region", "Select it and type the Roman numeral in the Region inspector’s name field (or press Shift-N)."),
             ("Loop it", "Press C for Cycle and drag the cycle area over the four measures."),
             ("Balance", "Open the Mixer (X). Pull the chord track’s fader down until the bass is clear. Solo (S) a track to hear it alone."),
             ("Export", "File › Bounce › Project or Section (Command-B), MP3. In Finder, Control-click the .logicx project and choose Compress.")]


def logicref_page(n):
    rows = "".join(f"<tr><td style='width:23%'><b>{a}</b></td><td>{b}</td></tr>" for a, b in LOGIC_REF)
    inner = (f'<p>Everything students need in Logic for this unit, in one table. On Day 3 demonstrate one chord (slide 15), then let students '
             f'build the rest.</p>'
             f'<table class="t" style="font-size:11.4pt;line-height:1.42">{rows}</table>'
             f'<div class="cols" style="margin-top:8pt"><div class="box"><span class="k">Chord check in the piano roll</span>'
             f'<p style="margin:0">Three notes stacked in one column, same start, same length. Count the rows between them: 4 then 3 half steps is '
             f'major, 3 then 4 is minor, 3 then 3 is diminished. Every grey row (a black key) counts.</p></div>'
             f'<div class="box"><span class="k">Version note</span>'
             f'<p style="margin:0" class="small">Menu names are from current Logic Pro versions; older versions may differ slightly. Logic Pro is a '
             f'trademark of Apple Inc., named only to describe what the unit needs. This resource is not affiliated with Apple. MPK Mini is named '
             f'only as an example keyboard.</p></div></div>')
    return page("Logic quick reference", "SETUP", "Teacher", inner, n, key="logicref")


def pacing_page(n, watermark=False):
    rows = "".join(
        f"<tr><td style='white-space:nowrap'><b class='acc' style='font-family:Oswald;font-size:13pt'>DAY {d['n']}</b><br><span class='small'>{d['kind']}</span></td>"
        f"<td><b>{d['title']}</b><br><span class='small'>{d['goal']}</span></td>"
        f"<td style='font-size:9.4pt;line-height:1.38'>{'<br>'.join(f'{a} · {h}' for a, h, _, _ in d['blocks'])}</td>"
        f"<td style='font-size:9.4pt'>{slides_label(d)}</td></tr>" for d in DAYS)
    inner = (f'<p>Five class periods of about {MINUTES} minutes: two learning days (ear first, then the keyboard and scales), two project days '
             f'(chords, then Roman numerals) and a finishing and presenting day. Each day’s full plan is on pages {IX["day12"]}–{IX["day5"]}.</p>'
             f'<table class="t"><tr><th style="width:13%">Day</th><th style="width:37%">Focus and goal</th><th style="width:31%">Minute blocks</th>'
             f'<th>Deck</th></tr>{rows}</table>'
             f'{H2("Student pages by day")}<table class="t">'
             f'<tr><td style="width:13%"><b>Day 1</b></td><td>Day 1 post (p. {IX["post1"]}), Chord Feel Log (p. {IX["feellog"]})</td></tr>'
             f'<tr><td><b>Day 2</b></td><td>Day 2 post (p. {IX["post2"]}), packet Parts 1–2 (pp. {IX["pk1"]}–{IX["pk2"]}), reference sheet (p. {IX["refsheet"]})</td></tr>'
             f'<tr><td><b>Day 3</b></td><td>Project sheet (pp. {IX["proj1"]}–{IX["proj2"]}), chords post (p. {IX["post3"]}, step 1), packet Part 3 chords (p. {IX["pk3"]})</td></tr>'
             f'<tr><td><b>Day 4</b></td><td>Chords post steps 2–5, Roman numerals (p. {IX["pk4"]}), KCSRN (p. {IX["kcsrn"]}), Reference Cards (pp. {IX["cardC"]}–{IX["cardB"]})</td></tr>'
             f'<tr><td><b>Day 5</b></td><td>Chord function table (p. {IX["proj1"]}), checklist (p. {IX["checklist"]}), rubric (p. {IX["rubric"]})</td></tr></table>')
    return page(f"Pacing guide · {DAYS_N} days", "PACING", "Teacher", inner, n, watermark, key="pacing")


def day_box(d):
    def blk(a, h, t, s):
        tag = f' <span class="small">[{"slide" if s[0] == s[1] else "slides"} {rng(*s)}]</span>' if s else ""
        return f'<div class="blk"><div class="tm">{a}</div><div><b class="h">{h}</b>{tag} · {fill(t)}</div></div>'
    return (f'<div class="day"><div class="dh"><b>Day {d["n"]} · {d["title"]}</b><span>{d["kind"]} · {slides_label(d)}</span></div>'
            f'<p class="goal"><b>Goal:</b> {d["goal"]}</p>{"".join(blk(*b) for b in d["blocks"])}</div>')


def day12_page(n):
    return page("Day-by-day plans · Days 1 and 2", "DAYS 1–2", "Learning days", day_box(DAYS[0]) + day_box(DAYS[1]), n, key="day12")


def spell_strip():
    rows = "".join(f"<tr><td style='white-space:nowrap'><b>{key}</b></td>" + "".join(
        f"<td style='text-align:center;white-space:nowrap'><b style='font-family:Oswald;font-size:11pt'>{rn}</b><br>{arrow(ns)}</td>"
        for rn, q, ns in key_chords(key)) + "</tr>" for key in KEYS)
    return (f'<div class="box"><span class="k">Spell-check at a glance (for circulating)</span>'
            f'<table class="t" style="font-size:10pt;margin:2pt 0 0">{rows}</table>'
            f'<p class="small" style="margin:3pt 0 0">Same seven chords in both keys, white keys only; only the Roman numerals move. '
            f'Full answers on pages {IX["key1"]}–{IX["key2"]}.</p></div>')


def day34_page(n):
    return page("Day-by-day plans · Days 3 and 4", "DAYS 3–4", "Project days", day_box(DAYS[2]) + day_box(DAYS[3]) + spell_strip(), n, key="day34")


def day5_page(n):
    prompts = ('<div class="box"><span class="k">Circulating prompts (Days 3–5)</span><ul style="margin:0">'
               '<li>‘Count your first chord out loud for me, starting from the root.’</li>'
               '<li>‘Do all three notes start in the same place? Does it look like a chord or a melody?’</li>'
               '<li>‘Point to your I chord. Play it. Does it feel like home?’</li>'
               '<li>‘Which chord wants to go back home? What is its Roman numeral?’</li></ul></div>')
    pres = ('<div class="box acc-b" style="margin-top:8pt"><span class="k">Presentation questions (Day 5)</span><ul style="margin:0">'
            '<li>What are your four chords, as Roman numerals and as chord names?</li>'
            '<li>Which one feels like home? Which one wants to resolve?</li>'
            '<li>What would you change if you could swap one chord?</li></ul>'
            '<p class="small" style="margin:4pt 0 0">The rubric’s top level asks students to explain their chord choices using Roman numeral '
            'language: listen for “the V wants to go back to the I,” not just “it sounds good.”</p></div>')
    notes = f'{H2("Notes for next year")}<div class="lines"></div>'
    return page("Day-by-day plans · Day 5", "DAY 5", "Finish + present", day_box(DAYS[4]) + prompts + pres + notes, n, key="day5")


def tn1_page(n):
    steps = "".join(f'<li><div class="n">{i + 1}</div><div>{s}</div></li>' for i, s in enumerate(OPENING))
    inner = (f'<div class="tn">{H2("Feel it before you name it", True)}'
             f'<p>This is the single most important fix in my course sequence: theory used to arrive as a cold worksheet. Taught ear-first, every '
             f'later project has real harmonic understanding behind it. The most common failure mode is rushing to the worksheet. Spend the full first '
             f'day just playing chords and describing feelings in plain language. Those descriptions are music theory; the Roman numerals just give '
             f'them shorter names.</p>'
             f'{H2("Day 1: the opening activity (first 10 minutes)")}'
             f'<p>Ten minutes of playing notes and chords before any theory framing plants an emotional reference point that the packet then labels. '
             f'Do not skip it, even if it feels like lost time.</p><ol class="num" style="font-size:11.2pt">{steps}</ol>'
             f'<div class="box"><span class="k">How this kit differs from my classroom version</span><ul style="margin:0">'
             f'<li>My Building Blocks material runs Parts 1–3 in 6 days; this kit follows the 5 days of my course plan: the ear day, then the '
             f'keyboard and scales in one day, then the three Part 3 days (presentations on Day 5).</li>'
             f'<li>The Chord Feel Log, the Day 1 post and the Logic file recipe are new: my plan calls for an ear-first activity before the KCSRN '
             f'worksheet. Part 2’s melody project and Part 4 (basslines) move to Unit 4.</li>'
             f'<li>My LMS is written as “your class page”. The balance mini-lesson on Day 5 applies my mixing concept for this unit to chords and '
             f'bass, since melody comes in Unit 4.</li></ul></div></div>')
    return page("Teacher notes · the big idea and Day 1", "TEACHER NOTES", "1 of 4", inner, n, key="tn1")


def tn2_page(n):
    inner = (f'<div class="tn">{H2("Roman numerals: the key framing", True)}'
             f'<p>Part 3 has two distinct moments. The first is mechanical: students learn to build chords from formulas and enter them correctly in '
             f'the piano roll. The second is conceptual: Roman numerals are introduced as a way to describe chord function rather than just chord '
             f'names. The second moment is more important and should not be rushed.</p>'
             f'<p>A chord’s letter name tells you what notes it contains. Its Roman numeral tells you how it functions: what emotional role it plays in '
             f'the key. The I chord always feels like home regardless of whether the key is C major or G major or Bb major. That principle transfers '
             f'to every key students will ever work in.</p>'
             f'<p>On Day 4, play I–IV–V–I together: home, away, very away, home. Play i–VI–III–VII: how is that different? Name the feelings before '
             f'naming the numerals. Then bring back the class’s words from Day 1: the KCSRN worksheet is “now name what you already heard.”</p>'
             f'{H2("The Reference Card is a living document")}'
             f'<p>Students will use it again in the next unit, in the EDM project and beyond. Make clear that filling it in carefully now is an '
             f'investment. The completed C major and A minor cards (pages {IX["cardC"]}–{IX["cardA"]}) are for checking; the blank card '
             f'(page {IX["cardB"]}) is the one students fill in.</p>'
             f'{H2("About the Logic file")}'
             f'<p>The Part 3 file contains a bassline already written; students add only the chord layer. The bass follows the chord roots so students '
             f'can infer the progression from the bass. Recipe on page {IX["setup"]}.</p>'
             f'{H2("Mixing concept: balance (Day 5)")}'
             f'<p>Solo the chords, solo the bass, then play both. Ask which one the listener should follow. Students pull the chord fader down until '
             f'the bass is clear. It is not graded; it is the habit the next unit needs when the melody arrives.</p></div>')
    return page("Teacher notes · Roman numerals, the card, balance", "TEACHER NOTES", "2 of 4", inner, n, key="tn2")


def tn3_page(n):
    st = "".join(f'<p><b>{a}:</b> {b}</p>' for a, b in STRUGGLES)
    inner = f'<div class="tn big">{H2("Common student struggles", True)}{st}</div>'
    return page("Teacher notes · common struggles", "TEACHER NOTES", "3 of 4", inner, n, key="tn3")


def tn4_page(n):
    inner = (f'<div class="tn">{H2("Suggested variations", True)}'
             f'<p><b>More accessible:</b> Pre-fill the chord spellings in the Scary Stuff packet for the key students are working in. Ask them to '
             f'verify each chord in Logic rather than calculate it from scratch.</p>'
             f'<p><b>More advanced:</b> Ask students to add a second chord layer with a different instrument sound and slightly varied rhythm. '
             f'Introduce the concept of harmonic rhythm (how often the chords change) and ask students to experiment with half-measure chord changes.</p>'
             f'{H2("Grading tips")}<ul>'
             f'<li>Three categories, each scored 2–5, for a total out of {RUBRIC_TOTAL} (page {IX["rubric"]}).</li>'
             f'<li>Chord Construction: open the piano roll and count. The answer keys (pages {IX["key1"]}–{IX["example"]}) list every chord in both keys.</li>'
             f'<li>Harmonic Understanding: the chord function table. ‘Sounds good’ is a 3; ‘wants to go back home’ is the kind of specific answer the 5 '
             f'asks for.</li>'
             f'<li>The quickest evidence: point to your I chord. Play it. Does it feel like home? That is the test.</li></ul>'
             f'{H2("Resources")}<ul>'
             f'<li>The intro deck (.pptx and PDF), included.</li>'
             f'<li>musictheory.net, lesson 44 (free), as extra reading for the KCSRN worksheet.</li>'
             f'<li>The two songs, from your own YouTube or Spotify access.</li></ul>'
             f'{H2("Observation notes")}'
             f'<p class="small">What worked, what to adjust, and anything worth remembering for next year.</p></div><div class="lines"></div>')
    return page("Teacher notes · variations, grading, resources", "TEACHER NOTES", "4 of 4", inner, n, key="tn4")


def post(n, key, title, sub, intro, bold, steps, box_title, box_body, tag):
    lis = "".join(f'<li style="margin-bottom:15pt;font-size:12.6pt;line-height:1.4"><div class="n">{i + 1}</div><div><b class="hd" style="font-size:13.5pt">{a}</b> · {bb}<br>'
                  f'<span class="why">{cc}</span></div></li>' for i, (a, bb, cc) in enumerate(steps))
    inner = (f'<h2 class="s" style="font-size:20pt;margin:4pt 0 2pt">{title}</h2>'
             f'<p class="l" style="font-size:12pt">{sub}</p>'
             f'<p style="font-size:12.6pt;margin:8pt 0 6pt">{intro}</p>'
             f'<p style="font-size:12.6pt;margin:0 0 14pt"><b>{bold}</b></p>'
             f'<ol class="num">{lis}</ol>'
             f'<div class="box acc-b" style="margin-top:6pt"><span class="k" style="font-size:12pt">{box_title}</span>{box_body}</div>')
    return page("Classroom post", "STUDENT", tag, inner, n, key=key)


def post1_page(n):
    steps = [("Listen", "Your teacher will play one note, then chords, in Logic. Describe what you hear in plain language.",
              "No wrong answers. Use any words that come to mind."),
             ("Listen to a real song", "Your teacher will play Twist and Shout by The Beatles. How many different chords do you hear? Which one "
              "feels like home?", "You do not need any names yet. Just notice."),
             ("Play the seven chords", f"Open Logic, load a piano sound, and play each chord on the Chord Feel Log (page {IX['feellog']}). "
              "Write what it feels like in your own words.", "Play it more than once. Say more than ‘good’ or ‘bad’: sad how? Waiting for what?"),
             ("Loop the two progressions", "Play Loop 1 and Loop 2 from the log on repeat. Where does each one feel at home?",
              "Same white keys, a different home."),
             ("Share your words", "Bring your best words to the share-out. We will keep them on the board.",
              "On Day 4 these words get shorter names. You are already doing music theory.")]
    box = ('<p style="margin:0;font-size:11pt">Three or more notes played at the same time make a <b>chord</b>. Today you only need your ears.</p>')
    return post(n, "post1", "Chords by Ear · Day 1", f"Digital Music Production 1 · Unit {UNIT}",
                "Today is a listening day. You are going to play chords and describe what they feel like before anyone tells you what they are called.",
                "Complete the Chord Feel Log before you leave today.", steps, "Today’s one idea", box, "Day 1")


def post2_page(n):
    steps = [("Listen", "Your teacher will play two scales. Which sounds brighter? Which sounds darker?", "Do not worry about names yet."),
             ("Part 1 — Keyboard diagram", f"Label all 7 white keys on the keyboard diagram in the packet (page {IX['pk1']}).",
              "Find C first — it sits just to the left of every group of 2 black keys."),
             ("Part 1 — Half/whole steps", "Complete the half step / whole step exercises. The first two are done for you as examples.",
              "Play each interval on your MPK Mini as you label it."),
             ("Logic activity", "Open Logic, load a piano sound, and play every white key from C to C, naming each one as you press it.",
              "Then play C and D. Then C and C#. Can you hear the difference between a whole step and a half step?"),
             ("Part 2 — Scales", f"Build the C Major, A Natural Minor and C Major Pentatonic scales in the packet (page {IX['pk2']}).",
              "Play each scale on your MPK Mini as you complete it.")]
    box = ('<p style="margin:0;font-size:12pt;font-family:Oswald;font-weight:700">C  –  D  –  E  –  F  –  G  –  A  –  B  –  (back to C)</p>'
           '<p style="margin:2pt 0 0;font-size:10.5pt">E→F and B→C are the only half steps between white keys — no black key between them.</p>'
           .replace("→", " to "))
    return post(n, "post2", "The Keyboard and Scales · Day 2", f"Digital Music Production 1 · Unit {UNIT}",
                "Today we start at the beginning — the notes on a keyboard. Before you can write music, you need to know what keys you are pressing.",
                "Complete the Scary Stuff packet Part 1 before opening Logic.", steps, "The white keys in order", box, "Day 2")


def post3_page(n):
    steps = [("Day 3 · Part 3 — Chord construction", "Complete the major, minor, and diminished chord exercises in the packet. C Major and A Minor "
              "are pre-filled as examples.", "Play each chord in Logic as you complete it on paper. Does it sound like what the formula predicts?"),
             ("Day 4 · Part 3 — Roman numerals", "Read the Roman numeral section and complete the chord feel table.",
              "The I chord always feels like home. The V chord always wants to resolve. These feelings are the same in every key."),
             ("Day 4 · Complete your Reference Card", "Fill in your Reference Card (C Major or A Minor) using your packet work and the reference sheets.",
              "Keep this card. You will use it for the rest of the year."),
             ("Day 4 · Open the Logic file", "Download today’s Logic file. A bassline is already written. Listen to it twice before adding anything.",
              "What are the chords? What Roman numerals would you use to label them?"),
             ("Day 4 · Build the chords", "Add a new track and build the chords for the progression using the piano roll. One chord per measure. "
              "Each chord should be 3 notes stacked vertically.", "Use your Reference Card to check your chord spellings.")]
    box = "".join(f'<p style="margin:0 0 2pt;font-size:11.5pt"><b style="font-family:Oswald;font-size:13pt">{a}: {f}</b> — {d}</p>'
                  for a, f, d in FORMULAS_BOX)
    return post(n, "post3", "Chords &amp; Roman Numerals · Days 3–4", f"Digital Music Production 1 · Unit {UNIT}",
                "Today you are building chords in Logic — and learning the system that explains why certain chords feel the way they do.",
                "Complete Part 3 of the packet before building anything in Logic.", steps, "The chord formulas", box, "Days 3–4")


def feellog_page(n, watermark=False):
    rows = "".join(f"<tr style='height:34pt'><td style='text-align:center;font-family:Oswald;font-weight:700;font-size:13pt'>{i + 1}</td>"
                   f"<td style='font-family:Oswald;font-weight:700;font-size:15pt;white-space:nowrap'>{arrow(ch)}</td><td></td>"
                   f"<td style='text-align:center;font-size:10pt'>bright · dark</td></tr>" for i, ch in enumerate(FEEL_CHORDS))

    def loop(rns, key):
        return "  ›  ".join(arrow(by_rn(key, r)) for r in rns)
    inner = (f'<p style="font-size:11pt;margin:0 0 5pt">Play each chord (all three notes at once) on a piano sound. Play it more than once. Then write '
             f'what it feels like in your own words and circle bright or dark. No names needed today.</p>'
             f'<table class="t" style="font-size:10.5pt"><tr><th style="width:7%">#</th><th style="width:18%">Play these notes</th>'
             f'<th>What does it feel like? (your words)</th><th style="width:16%">Circle one</th></tr>{rows}</table>'
             f'<div class="cols" style="gap:14pt;margin-top:4pt"><div class="box"><span class="k">Loop 1 · play it on repeat</span>'
             f'<p style="font-family:Oswald;font-weight:700;font-size:13.5pt;margin:2pt 0 4pt;white-space:nowrap">{loop(LOOP_MAJOR, "C major")}</p>'
             f'<div class="fl"><b>What happens at the end?</b><span></span></div><div class="fl" style="margin:0"><b>Home is:</b><span></span></div></div>'
             f'<div class="box"><span class="k">Loop 2 · play it on repeat</span>'
             f'<p style="font-family:Oswald;font-weight:700;font-size:13.5pt;margin:2pt 0 4pt;white-space:nowrap">{loop(LOOP_MINOR, "A minor")}</p>'
             f'<div class="fl"><b>How is it different?</b><span></span></div><div class="fl" style="margin:0"><b>Home is:</b><span></span></div></div></div>'
             f'<div class="fl" style="margin-top:8pt"><b>Which chord felt most like home?</b><span></span></div>'
             f'<div class="fl"><b>Which one wanted to go somewhere?</b><span></span></div>')
    return page("Chord Feel Log · play it, then describe it", "STUDENT", "Day 1 · by ear", inner, n, watermark, name=True, key="feellog")


def keyboard_html(labels, width="5.6in"):
    return f'<div style="width:{width};margin:2pt 0 4pt">{keyboard_svg(8, {}, labels=labels, black_labels=True, ww=56, wh=170, bw=34, bh=104, name="packet keyboard")}</div>'


def pk1_page(n):
    ex = "".join(f'<td><b>{a} → {b}</b></td><td style="font-family:Oswald;font-weight:700">{"W" if semis_between(a, b) == 2 else "H"}</td>'
                 .replace("→", "to") for a, b in HW_EXAMPLES)
    rows = "".join(f'<tr style="height:26pt"><td><b>{a} to {b}</b></td><td></td><td><b>{c} to {d}</b></td><td></td></tr>'
                   for (a, b), (c, d) in zip(HW_PAIRS[0::2], HW_PAIRS[1::2]))
    inner = (f'<p style="margin:0 0 3pt">The keyboard repeats the same pattern of 7 white keys and 5 black keys over and over. Learn one octave and you '
             f'know all of them.</p>'
             f'<div class="cols" style="gap:14pt;align-items:flex-start"><div style="flex:0 0 3.6in">{keyboard_html("white", "3.6in")}</div>'
             f'<div><p style="margin:0 0 4pt;font-size:10.2pt">Find C on your MPK Mini. It sits just to the left of every group of 2 black keys. Play C, '
             f'then continue upward naming each white key: C D E F G A B. Then it starts again at C.</p>'
             f'<div class="box" style="padding:4pt 8pt;margin-bottom:5pt"><span class="k">Half step (H)</span><p style="margin:0;font-size:9.8pt">Moving one '
             f'key in either direction — including black keys. The smallest possible interval. <i>E to F is a half step. B to C is a half step. These '
             f'are the only two white-key pairs with no black key between them.</i></p></div>'
             f'<div class="box" style="padding:4pt 8pt"><span class="k">Whole step (W)</span><p style="margin:0;font-size:9.8pt">Two half steps — skip '
             f'one key. Most white key pairs are whole steps. <i>C to D is a whole step. D to E is a whole step. G to A is a whole step.</i></p></div></div></div>'
             f'{H2("Now try it. Label the white keys without looking above")}{keyboard_html("none", "3.6in")}'
             f'{H2("Half step (H) or whole step (W)? The first two are done for you")}'
             f'<table class="t" style="font-size:11.5pt;width:80%"><tr style="height:24pt">{ex}</tr>{rows}</table>'
             f'<div class="box inv" style="margin-top:6pt"><p style="margin:0;font-size:11pt"><b>When you finish:</b> open Logic, load any piano sound, '
             f'and play every white key from C to C naming each one out loud as you press it.</p></div>')
    return page("The Scary Stuff · Part 1: The Keyboard", "PACKET", "Part 1 · Day 2", inner, n, name=True, key="pk1")


def note_row(first, n_cells, formula, given=1, final=None):
    cells = "".join(f'<td class="{"" if i < given else ""}">{first[i] if i < given else ""}</td>' for i in range(n_cells))
    cells += f'<td class="g">({final})</td>' if final else ""
    steps = "".join(f'<td class="st">{s}</td>' for s in steps_of(formula))
    return (f'<table class="nt"><tr>{cells}</tr></table>'
            f'<table class="nt" style="margin-left:20pt"><tr>{steps}</tr></table>')


def pk2_page(n):
    cmaj = key_scale("C major")
    inner = (f'<p style="margin:0 0 4pt">A scale is an ordered set of notes built using a formula of whole steps (W) and half steps (H). The formula tells '
             f'you exactly which notes to use. Play each scale on your MPK Mini as you complete it.</p>'
             f'{H2("C Major Scale · " + MAJOR_SCALE)}<p style="margin:0 0 3pt">Starting on C, apply the formula. The first two notes are done for you.</p>'
             f'{note_row(cmaj, 7, MAJOR_SCALE, given=2, final="=C")}'
             f'<div class="fl" style="margin-top:5pt"><b>Play it. Describe what it sounds like in one word:</b><span></span></div>'
             f'{H2("A Natural Minor Scale · " + MINOR_SCALE)}<p style="margin:0 0 3pt">Starting on A, apply the formula. Fill in all 8 notes.</p>'
             f'{note_row(["A"], 7, MINOR_SCALE, given=1, final="=A")}'
             f'<div class="fl" style="margin-top:5pt"><b>Play it. How does it feel different from C Major?</b><span></span></div>'
             f'<div class="box" style="margin:6pt 0"><span class="k">About the pentatonic scale</span><p style="margin:0">A pentatonic scale removes 2 '
             f'notes from the major or minor scale, leaving only the ones that sound most agreeable together. WH = one whole step + one half step '
             f'combined (3 half steps).</p></div>'
             f'{H2("C Major Pentatonic · " + MAJOR_PENT + " (5 notes)")}'
             f'{note_row(["C"], 5, MAJOR_PENT, given=1, final="=C")}'
             f'<div class="fl" style="margin-top:6pt"><b>Which two notes from C Major are missing from C Major Pentatonic?</b><span style="flex:0 0 1in"></span>'
             f'<span style="flex:0 0 1in"></span></div>'
             f'<div class="box inv" style="margin-top:4pt"><p style="margin:0;font-size:11pt"><b>Keep going:</b> the notes of these scales are '
             f'the notes of your chords. Every chord in C major is built only from C D E F G A B.</p></div>')
    return page("The Scary Stuff · Part 2: Scales", "PACKET", "Part 2 · Day 2", inner, n, name=True, key="pk2")


def chord_box(root, kind, formula, filled=False):
    notes = chord(root, formula)
    show = filled or (root, kind) in PACKET_GIVEN
    title = f"{root} {'Diminished' if kind == 'Diminished' else kind}"
    cells = "".join(f'<td>{nn if show else ""}</td>' for nn in notes)
    return (f'<div class="cbox"><b class="cn">{title}</b><table class="ch"><tr>{cells}</tr>'
            f'<tr><td class="lb">Root</td><td class="lb">3rd</td><td class="lb">5th</td></tr></table></div>')


def chord_grid(filled=False):
    out = []
    for kind, formula, feel, roots in PACKET_CHORDS:
        boxes = "".join(f'<div style="flex:1;min-width:0">{chord_box(r, kind, formula, filled)}</div>' for r in roots)
        out.append(f'{H2(f"{kind} Chords · {formula} — {feel}")}<div style="display:flex;gap:10pt">{boxes}</div>')
    return "".join(out)


def pk3_page(n):
    inner = (f'<p style="margin:0 0 3pt">A chord is three or more notes played at the same time. You build chords by stacking intervals on top of a root '
             f'note. Play each chord in Logic as you complete it.</p>'
             f'<div class="box" style="padding:4pt 9pt"><p style="margin:0"><b>R = Root</b> &nbsp;·&nbsp; 4 = 4 half steps up from the root &nbsp;·&nbsp; '
             f'3 = 3 half steps up from the previous note. <i>Do not count the note you start on: for C, C# is 1.</i></p></div>'
             f'{chord_grid()}'
             f'<div style="margin-top:12pt">{chordkeys_svg()}</div>'
             f'<p class="small" style="margin-top:3pt">Teal keys: the examples, counted from the formula. Use the keyboard to count yours.</p>')
    return page("The Scary Stuff · Part 3: Chords", "PACKET", "Part 3 · Day 3", inner, n, name=True, key="pk3")


def pk4_page(n):
    rows = "".join(f"<tr><td style='font-family:Oswald;font-weight:700;font-size:14pt;text-align:center'>{a}</td><td>{b}</td><td><b>{c}</b></td>"
                   f"<td><i>{d}</i></td><td></td></tr>" for a, b, c, d in PACKET_RN)
    inner = (f'{H2("Roman Numerals — Names for Chord Function", True)}'
             f'<p style="font-size:11pt">Every chord in a key has a Roman numeral based on its position in the scale. The Roman numeral tells you something '
             f'the letter name cannot — it tells you how the chord functions. The I chord always feels like home. The V chord always feels like it '
             f'needs to resolve. These feelings are the same in every key.</p>'
             f'<p style="font-size:10.5pt">Uppercase = major chord. Lowercase = minor chord. ° = diminished. In the last column, write the word '
             f'the class used for that feeling on Day 1.</p>'
             f'<table class="t" style="font-size:11pt;line-height:1.35"><tr><th style="width:11%">RN</th><th style="width:14%">Type</th>'
             f'<th style="width:17%">Chord</th><th>Feels like…</th><th style="width:22%">Our Day 1 word</th></tr>{rows}</table>'
             f'{H2("Hear it")}'
             f'<div class="fl"><b>Play I–IV–V–I in C major. Write the chord names:</b><span></span></div>'
             f'<div class="fl"><b>Play i–VI–III–VII in A minor. How is it different?</b><span></span></div>'
             f'<div class="fl"><b>Which chord pulls hardest back to home?</b><span></span></div>'
             f'<div class="box inv" style="margin-top:6pt"><p style="margin:0;font-size:11pt">Now complete your Reference Card for your chosen key. Use the '
             f'Scary Stuff Reference Sheet if you need help.</p><p style="margin:4pt 0 0;font-size:11pt"><b>When you finish:</b> open the Logic file '
             f'your teacher provides. A bassline is already written. Build the chords on top of it using the formulas above.</p></div>'
             f'<div class="lines" style="margin-top:8pt"></div>')
    return page("The Scary Stuff · Part 3: Roman Numerals", "PACKET", "Part 3 · Day 4", inner, n, name=True, key="pk4")


def kcsrn_col(kind, filled=False):
    key = "C major" if kind == "major" else "A minor"
    sc, chs = key_scale(key), key_chords(key)
    blank = '<span style="display:inline-block;width:30pt;border-bottom:1pt solid #000;height:15pt;margin-right:4pt"></span>'

    def val(x, w=30):
        return (f'<span style="display:inline-block;min-width:{w}pt;border-bottom:1pt solid #000;margin-right:4pt;font-family:Oswald;font-weight:700;'
                f'text-align:center;color:{ACCENT}">{x}</span>') if filled else blank.replace("30pt", f"{w}pt")
    head = "Major Key" if kind == "major" else "Minor Key"
    scl = "Scale" if kind == "major" else "Scale (natural minor)"
    mg = "2pt" if filled else "9pt"
    names = "".join(f'<div style="margin:0 0 {mg}"><b style="display:inline-block;width:34pt">{rn}</b>{val(chord_name(ns), 74)}</div>'
                    for rn, q, ns in chs)
    spell = "".join(f'<div style="margin:0 0 {mg};white-space:nowrap"><b style="display:inline-block;width:34pt">{rn}:</b>' + "".join(val(x, 24) for x in ns)
                    + '</div>' for rn, q, ns in chs)
    rem = (MAJOR_SCALE if kind == "major" else MINOR_SCALE)
    return (f'<div><h2 class="s" style="margin-top:0;font-size:15pt">{head}</h2>'
            f'<div class="fl"><b>{head}:</b><span style="text-align:center;font-family:Oswald;font-weight:700;color:{ACCENT}">{key.split()[0] + (" major" if kind == "major" else " minor") if filled else ""}</span></div>'
            f'<div style="margin:0 0 8pt;white-space:nowrap"><b>{scl}:</b><br>' + "".join(val(x, 22) for x in sc[:8]) + '</div>'
            f'<span class="k">Chord names</span>{names}'
            f'<span class="k" style="margin-top:6pt">Chord spellings</span>{spell}'
            + ("" if filled else f'<div class="box" style="padding:4pt 8pt;margin-top:4pt"><span class="k">Reminders!</span><p style="margin:0;font-size:9.8pt">'
            f'{"Major" if kind == "major" else "Natural Minor"} Scale: {rem.replace(" - ", " ")}<br>Major chord: R 4 3 · Minor chord: R 3 4<br>'
            f'Diminished (°) chord: R 3 3</p></div>') + '</div>')


def kcsrn_page(n):
    inner = (f'<p style="margin:0 0 4pt;font-size:10.8pt"><b>Keys, chords, scales and Roman numerals.</b> Now name what you already heard. Fill in the '
             f'column for your key (major = C major, minor = A minor). Extra reading: musictheory.net, lesson 44.</p>'
             f'<div class="cols" style="gap:24pt">{kcsrn_col("major")}{kcsrn_col("minor")}</div>')
    return page("KCSRN worksheet", "KCSRN", "Day 4", inner, n, name=True, key="kcsrn")


def card_table(key):
    rows = "".join(f"<tr><td class='rn'>{rn}</td><td>{t}</td><td><b>{c}</b></td><td class='nt'>{nt}</td><td><i>{f}</i></td></tr>"
                   for rn, t, c, nt, f in CARD[key])
    return (f"<table class='card'><tr><th>RN</th><th style='width:11%'>Type</th><th style='width:17%'>Chord</th><th style='width:17%'>Notes</th>"
            f"<th>Feels like…</th></tr>{rows}</table>")


def scale_cells(notes, n_cells=8, extra=None):
    cells = "".join(f"<td>{x}</td>" for x in notes[:n_cells])
    if extra:
        cells += f"<td class='g'>{extra}</td>"
    return f"<table class='nt' style='margin:2pt 0 6pt'><tr>{cells}</tr></table>"


def remember_box():
    a, b, c = CARD_REMEMBER
    return (f'<div class="box acc-b" style="margin-top:8pt"><p style="margin:0;font-size:11.5pt"><b>{a}</b></p>'
            f'<p style="margin:3pt 0 0;font-size:10.5pt"><i>{b}</i></p><p style="margin:3pt 0 0;font-size:10.5pt">{c}</p></div>')


def card_page(n, key, watermark=False):
    sc = key_scale(key)
    extra = "(8)" if key == "C major" else None
    inner = (f'<h2 class="s" style="font-size:20pt;margin:2pt 0 0">{key.split()[0]} {key.split()[1].title()} — Reference Card</h2>'
             f'<p class="l" style="font-size:12pt;margin-bottom:4pt">The Building Blocks · Unit {UNIT} · Keep this open while you work</p>'
             f'{H2("Scale")}<p style="font-size:12pt;margin:0"><b>{CARD_SCALE_TEXT[key]}</b></p>{scale_cells(sc, 8, extra)}'
             f'{H2("Chords in this Key")}{card_table(key)}{remember_box()}'
             f'<p class="small" style="margin-top:6pt">No sharps or flats in this key: every chord is white keys only. '
             f'{"A minor uses the same seven notes: it is the relative minor." if key == "C major" else "C major uses the same seven notes: it is the relative major."}</p>')
    return page(f"Reference Card · {key}", "REFERENCE", "Keep it open", inner, n, watermark, key=f"card{key[0]}")


def cardB_page(n):
    rows = "".join(f"<tr style='height:38pt'><td class='rn'>{rn}</td><td>{t}</td><td></td><td></td><td></td></tr>" for rn, t in BLANK_CARD)
    cells = "<td style='font-size:9pt;font-family:Source Sans 3;font-weight:400;font-style:italic'>ROOT</td>" + "<td></td>" * 6 + \
            "<td style='font-size:9pt;font-family:Source Sans 3;font-weight:400;font-style:italic'>ROOT</td>"
    inner = (f'<h2 class="s" style="font-size:20pt;margin:2pt 0 0">My Key — Reference Card</h2>'
             f'<p class="l" style="font-size:12pt;margin-bottom:4pt">The Building Blocks · Unit {UNIT} · Keep this open while you work</p>'
             f'<div class="fl" style="font-size:12pt"><b>My key:</b><span></span></div>'
             f'{H2("Scale")}<div class="fl"><b>Key:</b><span style="flex:0 0 1.6in"></span><b>Formula:</b><span></span></div>'
             f'<table class="nt" style="margin:2pt 0 6pt"><tr>{cells}</tr></table>'
             f'{H2("Chords in this Key")}'
             f"<table class='card'><tr><th>RN</th><th style='width:11%'>Type</th><th style='width:17%'>Chord</th><th style='width:17%'>Notes</th>"
             f"<th>Feels like…</th></tr>{rows}</table>{remember_box()}"
             f'<p class="small" style="margin-top:6pt">The RN and Type columns are for a major key, as on my card. Working in A minor? Cross them out and '
             f'use i, ii°, III, iv, v, VI, VII (check the A minor card).</p>')
    return page("Reference Card · blank", "REFERENCE", "Fill it in", inner, n, name=True, key="cardB")


def refsheet_page(n):
    scales = "".join(f"<tr><td><b>{a}</b></td><td style='white-space:nowrap'>{b}</td><td><i>{c}</i></td></tr>" for a, b, c in REF_SCALES)
    chords = "".join(f"<tr><td><b>{a}</b></td><td>{len(ns)} notes</td><td style='white-space:nowrap'>{f.replace('-', ' - ')}</td>"
                     f"<td style='white-space:nowrap'>{ex} = {dash(ns)}</td><td><i>{feel}</i></td></tr>" for a, f, ex, ns, feel in REF_CHORDS)
    rn = "".join(f"<tr><td style='font-family:Oswald;font-weight:700;font-size:12pt;text-align:center'>{a}</td><td>{b}</td><td><i>{c}</i></td></tr>"
                 for a, b, c in REF_RN)
    inner = (f'<p class="why" style="margin:0 0 2pt">Keep this open while you work. You do not need to memorize it.</p>'
             f'<div class="cols" style="gap:14pt;align-items:flex-start"><div style="flex:0 0 2.7in">'
             f'{keyboard_html("white", "2.7in")}</div><div><p style="margin:0;font-size:10pt"><b>The Keyboard.</b> Every octave has the same 7 white keys '
             f'and 5 black keys. The pattern repeats. Find C — it sits just to the left of every group of 2 black keys.</p>'
             f'<p style="margin:3pt 0 0;font-size:10pt"><b>Half step (H):</b> moving one key in either direction — the smallest interval. E to F and B to '
             f'C are half steps on white keys. <b>Whole step (W):</b> two half steps — skip one key. C to D, D to E, F to G, G to A, A to B are whole '
             f'steps.</p></div></div>'
             f'{H2("Scale Formulas")}<table class="t tight">{scales}</table>'
             f'{H2("Chord Types")}<table class="t tight"><tr><th>Type</th><th>Notes</th><th>Formula</th><th>Example (root C)</th>'
             f'<th>Sounds like</th></tr>{chords}</table>'
             f'{H2("Chords in a Key — Roman Numerals")}<p style="margin:0 0 2pt;font-size:9.8pt">Every key has 7 chords built from its scale. The Roman '
             f'numeral tells you which scale degree the chord is built on. Uppercase = major chord. Lowercase = minor chord.</p>'
             f'<table class="t tight">{rn}</table>'
             f'<p class="why" style="font-size:9.6pt">In a minor key, the i chord is home — but it feels more unsettled and emotional than the major I. '
             f'The iv, v, VI, and VII chords take on the same functions relative to their new home.</p>')
    return page("The Scary Stuff · reference sheet", "REFERENCE", "Notes, scales, chords", inner, n, key="refsheet")


def proj1_page(n):
    rows = "".join(f"<tr style='height:40pt'><td style='text-align:center;font-family:Oswald;font-weight:700;font-size:14pt'>{i}</td><td></td><td></td><td></td></tr>"
                   for i in range(1, 5))
    inner = (f'<h2 class="s" style="font-size:20pt;margin:2pt 0 0">Chords &amp; Roman Numerals</h2>'
             f'<p class="l" style="font-size:12pt;margin-bottom:4pt">The Building Blocks · Part 3 · Digital Music Production 1 · Unit {UNIT} · '
             f'3 project days (Days 3–5)</p>'
             f'{H2("Overview")}<p style="font-size:11.2pt">{OVERVIEW}</p>'
             f'{H2("Requirements")}'
             + "".join(f'<div class="chk" style="font-size:10.8pt;margin-bottom:5pt">{checkbox(12)}<span>{r}</span></div>' for r in REQUIREMENTS) +
             f'{H2("Chord Function — Fill This In")}<p style="font-size:10.8pt;margin:0 0 3pt">After building your chords, complete this table:</p>'
             f'<table class="t" style="font-size:10.5pt"><tr><th style="width:11%">Chord #</th><th style="width:17%">Roman Numeral</th>'
             f'<th style="width:20%">Chord Name</th><th>What it feels like (in your own words)</th></tr>{rows}</table>')
    return page("Project sheet · overview, requirements, chord function", "STUDENT", "Project · 1 of 2", inner, n, name=True, key="proj1")


def proj2_page(n):
    inner = (f'{H2("Tips for Success", True)}<ul>' + "".join(f"<li style='font-size:11.5pt;margin-bottom:6pt'>{t}</li>" for t in TIPS) + '</ul>'
             f'{H2("The chord formulas")}'
             f'<div style="display:flex;gap:10pt">' + "".join(
                 f'<div class="box" style="flex:1"><span class="k">{a}</span><p style="margin:0;font-family:Oswald;font-weight:700;font-size:18pt">{f}</p>'
                 f'<p style="margin:0;font-size:10.5pt">{d}</p></div>' for a, f, d in FORMULAS_BOX) + '</div>'
             f'{H2("Counting half steps")}<p style="font-size:11pt">Do not count the root. For C major (R-4-3): C#(1), D(2), D#(3), <b>E</b>(4). Then '
             f'count 3 from E: F(1), F#(2), <b>G</b>(3). C major = C - E - G.</p>'
             f'<div class="box inv" style="margin-top:6pt"><span class="k">The test</span>'
             f'<p style="margin:0;font-size:13pt">Point to your I chord. Play it. Does it feel like home?</p></div>'
             f'{H2("Notes")}<div class="lines"></div>')
    return page("Project sheet · tips and formulas", "STUDENT", "Project · 2 of 2", inner, n, name=True, key="proj2")


def vocab_page(n):
    rows = "".join(f"<tr><td style='width:19%'><b>{a}</b></td><td>{b}</td></tr>" for a, b in VOCAB)
    qz = "".join(f'<div class="fl" style="margin-bottom:6pt"><b>{i + 1}. {q}</b><span style="flex:0 0 1.2in"></span></div>' for i, (q, _) in enumerate(VOCAB_QUIZ))
    inner = (f'<p style="font-size:11pt">Read each word and its meaning. These are the words you should own by the end of the unit.</p>'
             f'<table class="t" style="font-size:10.8pt;line-height:1.38"><tr><th>Word</th><th>Meaning</th></tr>{rows}</table>'
             f'<div class="box" style="margin-top:6pt"><span class="k">Quick check · answers on page {IX["key1"]}</span>{qz}</div>')
    return page("Key vocabulary", "STUDENT", f"{len(VOCAB)} words", inner, n, name=True, key="vocab")


def checklist_page(n):
    inner = (f'<p style="font-size:11.5pt">Tick each box before you submit and present. If a box isn’t ticked, you aren’t done yet.</p>'
             + "".join(f'<div class="chk" style="font-size:12pt;margin-bottom:9pt">{checkbox(14)}<span>{c}</span></div>' for c in CHECKLIST) +
             f'{H2("Get ready to present")}'
             f'<div class="fl"><b>My four chords (Roman numerals):</b><span></span></div>'
             f'<div class="fl"><b>My home chord is</b><span style="flex:0 0 1.2in"></span><b>because</b><span></span></div>'
             f'<div class="fl"><b>The chord that wants to go home is</b><span></span></div>'
             f'{H2("Notes")}<div class="lines"></div>')
    return page("Submission checklist", "STUDENT", "Before you present", inner, n, name=True, key="checklist")


def rubric_page(n, watermark=False):
    head = "<tr><th>Category</th>" + "".join(f"<th>{lv}</th>" for lv in LEVELS) + "<th>Score</th></tr>"
    rows = "".join(f"<tr><td class='c'>{c}</td>" + "".join(f"<td>{d}</td>" for d in ds) + "<td class='sc'></td></tr>" for c, ds in RUBRIC)
    total = (f"<tr><td class='c' colspan='5' style='text-align:right;font-size:14pt'>Total</td>"
             f"<td style='font-family:Oswald;font-weight:700;font-size:14pt;text-align:right;vertical-align:bottom'>/ {RUBRIC_TOTAL}</td></tr>")
    inner = (f'<h2 class="s" style="font-size:18pt;margin:2pt 0 0">Part 3: Chords &amp; Roman Numerals — Rubric</h2>'
             f'<p class="l" style="font-size:12pt;margin-bottom:6pt">The Building Blocks · Unit {UNIT} · /{RUBRIC_TOTAL}</p>'
             f'<p style="font-size:11pt">Each category is scored from 2 to 5.</p>'
             f'<table class="rub">{head}{rows}{total}</table>'
             f'<div class="cols" style="margin-top:14pt"><div><h2 class="s" style="margin-top:0">My I chord feels like home because…</h2>'
             f'<div class="ln2"></div><div class="ln2"></div><div class="ln2"></div></div>'
             f'<div><h2 class="s" style="margin-top:0">Teacher comments</h2><div class="ln2"></div><div class="ln2"></div><div class="ln2"></div></div></div>')
    return page("Rubric", "RUBRIC", f"/ {RUBRIC_TOTAL}", inner, n, watermark, name=True, key="rubric")


def key1_page(n):
    hw = " · ".join(f"{a} to {b} <b>{'W' if semis_between(a, b) == 2 else 'H'}</b>" for a, b in HW_PAIRS)
    cs, am, pent = key_scale("C major"), key_scale("A minor"), scale("C", MAJOR_PENT)
    missing = [x for x in cs[:7] if x not in pent]
    grid = chord_grid(filled=True)
    vq = " &nbsp;·&nbsp; ".join(f"<b>{i + 1}.</b> {a}" for i, (_, a) in enumerate(VOCAB_QUIZ))
    inner = (f'<div class="cols" style="gap:14pt"><div class="box" style="padding:5pt 9pt"><span class="k">Part 1 · keyboard and steps</span>'
             f'<p style="margin:0;font-size:10.2pt">White keys: C D E F G A B. {hw}.</p></div>'
             f'<div class="box" style="padding:5pt 9pt"><span class="k">Part 2 · scales</span><p style="margin:0;font-size:10.2pt">'
             f'C major: {" ".join(cs)}<br>A natural minor: {" ".join(am)}<br>C major pentatonic: {" ".join(pent)}<br>Missing from the pentatonic: '
             f'<b>{" and ".join(missing)}</b></p></div></div>'
             f'<p style="margin:6pt 0 0;font-size:10.2pt"><b>Part 3 · chords</b> (every spelling counted from the formula; D major and B minor use a '
             f'black key, D diminished uses A-flat, written Ab):</p>{grid}'
             f'<div class="cols" style="gap:14pt;margin-top:8pt"><div class="box" style="padding:5pt 9pt"><span class="k">Vocabulary · quick check</span>'
             f'<p style="margin:0;font-size:10.2pt">{vq}</p></div>'
             f'<div class="box" style="padding:5pt 9pt"><span class="k">Roman numerals · hear it</span><p style="margin:0;font-size:10.2pt">'
             f'I–IV–V–I in C: {", ".join(short_name(by_rn("C major", r)) for r in LOOP_MAJOR)}. i–VI–III–VII in A minor: '
             f'{", ".join(short_name(by_rn("A minor", r)) for r in LOOP_MINOR)}. The pull home: V (G major) in C; in A minor, VII (G major) '
             f'drives back toward i.</p></div></div>')
    return page("Answer key · packet Parts 1–3", "ANSWER KEY", "Teacher", inner, n, ak=True, key="key1")


def key2_page(n):
    feel = "".join(f"<tr><td style='white-space:nowrap;font-family:Oswald;font-weight:700'>{i + 1} · {arrow(ch)}</td>"
                   f"<td style='white-space:nowrap'>{rn} in C · {chord_name(ch)}</td><td><i>{next(r[4] for r in CARD['C major'] if r[0] == rn)}</i></td>"
                   f"<td>{dict(Major='bright', Minor='dark').get(quality(ch), 'tense')}</td></tr>"
                   for i, (rn, ch) in enumerate(zip(FEEL_ORDER, FEEL_CHORDS)))
    inner = (f'<div class="cols" style="gap:18pt">{kcsrn_col("major", filled=True)}{kcsrn_col("minor", filled=True)}</div>'
             f'{H2("Chord Feel Log · words from my cards (any honest description counts)")}'
             f'<table class="t" style="font-size:9.8pt"><tr><th>Chord</th><th>What it is</th><th>Feels like (my card)</th><th>Usually</th></tr>{feel}</table>'
             f'<p class="small">Loop 1 is I–IV–V–I in C major and comes home on C–E–G. Loop 2 is i–VI–III–VII in A minor: home is A–C–E. '
             f'The diminished chord (B–D–F) is neither bright nor dark: accept either circle, or “tense”.</p>')
    return page("Answer key · KCSRN and the Feel Log", "ANSWER KEY", "Teacher", inner, n, ak=True, key="key2")


def ex_table(key):
    rows = []
    for i, (rn, ch, vo, bass) in enumerate(roll_layout(key), start=1):
        feel = next(r[4] for r in CARD[key] if r[0] == rn)
        rows.append(f"<tr><td style='text-align:center'><b>{i}</b></td><td style='font-family:Oswald;font-weight:700;font-size:12pt;text-align:center'>"
                    f"{rn}</td><td>{chord_name(ch)} ({dash(ch)})</td><td>{midi_name(bass)}</td><td><i>{feel}</i></td></tr>")
    return (f"<table class='t' style='font-size:9.8pt'><tr><th style='width:8%'>Chord #</th><th style='width:11%'>Roman numeral</th>"
            f"<th style='width:27%'>Chord name</th><th style='width:9%'>Bass</th><th>What it feels like (from the card)</th></tr>{''.join(rows)}</table>")


def example_page(n):
    inner = (f'<p style="font-size:10.2pt;margin:0 0 4pt">One possible answer for the Part 3 Logic file: four chords, one per measure, root position, '
             f'over a bassline that plays each chord’s root. C major file: <b>{"–".join(EX_PROG["C major"])}</b>; A minor file: '
             f'<b>{"–".join(EX_PROG["A minor"])}</b>. Teal regions: the chords students build; black: the provided bassline.</p>'
             f'<div class="cols" style="gap:16pt;align-items:center"><div style="flex:0 0 4.2in">{roll_svg("C major")}</div>'
             f'<div><span class="k">How to read it</span><ul style="margin:0">'
             f'<li>Each column is one measure: one chord, one bass note.</li>'
             f'<li>The three teal notes are stacked: same start, same length. Root on the bottom.</li>'
             f'<li>The black bass note is the chord’s root, an octave or two lower.</li>'
             f'<li>Label each region with the Roman numeral on top: I, vi, IV, V.</li>'
             f'<li>Grey rows are black keys. Count rows to check a chord: 4 + 3 is major, 3 + 4 is minor.</li></ul></div></div>'
             f'{H2("Chord function table · C major file")}{ex_table("C major")}'
             f'{H2("Chord function table · A minor file")}{ex_table("A minor")}'
             f'<p class="small">Same four chords, two homes: A minor is the relative minor of C major. A student who writes only “sounds good” scores a 3 '
             f'for Harmonic Understanding; “wants to go back to the I” is the kind of answer the 5 describes.</p>')
    return page("Example notes · a worked example in both keys", "EXAMPLE", "Teacher", inner, n, ak=True, key="example")


def terms_page(n):
    return (f'<section class="page" data-key="terms"><div class="top"><div class="t">Terms of Use<small>Thank you for your purchase!</small></div></div>'
            f'<div class="main"><p>© {YEAR} {STORE}. {BRAND} by {STORE}. All rights reserved. This purchase gives <b>one teacher</b> a license to use this '
            f'resource (the PDF and the intro deck) with their own students.</p>'
            f'<h2 class="s">You may</h2><ul><li>Print and copy pages for your own classroom and students, year after year.</li>'
            f'<li>Present the intro deck to your own classes and edit it for your own teaching.</li>'
            f'<li>Share pages with your own students through a password-protected class site or learning platform.</li>'
            f'<li>Leave the pages for a substitute teacher who teaches your own students.</li></ul>'
            f'<h2 class="s">You may not</h2><ul><li>Share, email or copy this resource for other teachers, a whole school or a district. '
            f'Additional licenses are available on Teachers Pay Teachers.</li>'
            f'<li>Post any part of it on a public website, shared drive or social media.</li>'
            f'<li>Sell, give away or claim any part of it as your own, or edit it into a new product.</li></ul>'
            f'<h2 class="s">Questions?</h2><p>Please use the Q&amp;A tab on the {BRAND} by {STORE} store on Teachers Pay Teachers. '
            f'If something looks wrong, let me know there and I will fix it.</p>'
            f'<h2 class="s">Credits</h2><ul>'
            f'<li>Curriculum, packet, reference cards, KCSRN worksheet, project sheet, rubric, teacher notes and intro deck: original work by {STORE}, {BRAND}.</li>'
            f'<li>Keyboard and piano-roll drawings, worked example and layouts: original work, {BRAND}.</li>'
            f'<li>Text fonts: Oswald (The Oswald Project Authors) and Source Sans 3 (Adobe), SIL Open Font License 1.1.</li>'
            f'<li>No audio, video, lyrics or album art are included. Songs are named for reference only (title and artist); play them from your own '
            f'YouTube or Spotify access. Song titles belong to their owners; this resource is not affiliated with or endorsed by the artists or their '
            f'labels. musictheory.net is named only as free extra reading.</li>'
            f'<li>Logic Pro is a trademark of Apple Inc., named only to describe compatibility; this resource is not affiliated with or endorsed by '
            f'Apple. MPK Mini, YouTube and Spotify are named only as examples of tools you may already use.</li></ul>'
            f'<p class="small">{DISCLOSURE}</p><div style="flex:1"></div></div>{foot(f"Page {n}")}</section>')


def render_page(key, watermark=False):
    n = IX[key]
    f = {"cover": cover_page, "howto": lambda: howto_page(n), "overview": lambda: overview_page(n), "setup": lambda: setup_page(n),
         "logicref": lambda: logicref_page(n), "pacing": lambda: pacing_page(n, watermark), "day12": lambda: day12_page(n),
         "day34": lambda: day34_page(n), "day5": lambda: day5_page(n), "tn1": lambda: tn1_page(n), "tn2": lambda: tn2_page(n),
         "tn3": lambda: tn3_page(n), "tn4": lambda: tn4_page(n), "post1": lambda: post1_page(n), "feellog": lambda: feellog_page(n, watermark),
         "post2": lambda: post2_page(n), "post3": lambda: post3_page(n), "pk1": lambda: pk1_page(n), "pk2": lambda: pk2_page(n),
         "pk3": lambda: pk3_page(n), "pk4": lambda: pk4_page(n), "kcsrn": lambda: kcsrn_page(n),
         "cardC": lambda: card_page(n, "C major", watermark), "cardA": lambda: card_page(n, "A minor", watermark), "cardB": lambda: cardB_page(n),
         "refsheet": lambda: refsheet_page(n), "proj1": lambda: proj1_page(n), "proj2": lambda: proj2_page(n), "vocab": lambda: vocab_page(n),
         "checklist": lambda: checklist_page(n), "rubric": lambda: rubric_page(n, watermark), "key1": lambda: key1_page(n),
         "key2": lambda: key2_page(n), "example": lambda: example_page(n), "terms": lambda: terms_page(n)}
    return f[key]()


# ================================================================ intro deck (.pptx, python-pptx) + PDF copy (LibreOffice)
def slide_day(i):
    return SLIDES[i - 1][0]


def build_deck(img):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

    BLACK, WHITE, ACC = RGBColor(0, 0, 0), RGBColor(255, 255, 255), RGBColor.from_string(ACCENT[1:])
    LIGHT = RGBColor.from_string(ACC_LIGHT[1:])
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.author = f"{STORE} · {BRAND}"
    prs.core_properties.title = f"{TITLE} · Intro deck"
    prs.core_properties.comments = DISCLOSURE
    layout = prs.slide_layouts[6]

    def box(sl, x, y, w, h, fillc=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=1.5):
        s = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        if fillc is None:
            s.fill.background()
        else:
            s.fill.solid()
            s.fill.fore_color.rgb = fillc
        if line is None:
            s.line.fill.background()
        else:
            s.line.color.rgb = line
            s.line.width = Pt(lw)
        s.shadow.inherit = False
        return s

    def text(sl, x, y, w, h, runs, size=24, font="Source Sans 3", bold=False, color=BLACK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             spacing=None, caps=False):
        tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.04)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        paras = runs if isinstance(runs, list) else [runs]
        for k, ptxt in enumerate(paras):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = ptxt.upper() if caps else ptxt
            r.font.size = Pt(size)
            r.font.name = font
            r.font.bold = bold
            r.font.color.rgb = color
            if spacing:
                r._r.get_or_add_rPr().set("spc", str(spacing))
        return tb

    def chrome(sl, i, title, day):
        text(sl, 0.6, 0.32, 6, 0.4, BRAND, 13, "Oswald", False, BLACK, spacing=600, caps=True)
        box(sl, 11.33, 0.3, 1.4, 0.42, ACC, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        text(sl, 11.33, 0.3, 1.4, 0.42, f"Day {day}", 15, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, caps=True)
        text(sl, 0.6, 0.72, 12.1, 0.95, title, 40, "Oswald", True, BLACK, caps=True, anchor=MSO_ANCHOR.BOTTOM)
        box(sl, 0.6, 1.72, 12.13, 0.04, BLACK)
        box(sl, 0.6, 7.0, 12.13, 0.012, BLACK)
        text(sl, 0.6, 7.04, 9, 0.3, f"{TITLE} · Intro deck · © {YEAR} {STORE} · {BRAND}", 10, "Source Sans 3", False, RGBColor(0x33, 0x33, 0x33))
        text(sl, 11.2, 7.04, 1.53, 0.3, f"{i}", 10, "Oswald", True, BLACK, align=PP_ALIGN.RIGHT)

    def bullets(sl, items, x=0.6, y=2.05, w=12.1, size=28, gap=1.75):
        for k, it in enumerate(items):
            yy = y + k * (size / 72 * gap)
            box(sl, x, yy + size / 72 * 0.36, 0.16, 0.16, ACC, shape=MSO_SHAPE.OVAL)
            text(sl, x + 0.38, yy, w - 0.4, size / 72 * 1.6, it, size, "Source Sans 3", False, BLACK)

    def pic(sl, key, x, y, w):
        path, (pw, ph) = img[key]
        sl.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))
        return y + w * ph / pw

    def tile(sl, x, y, w, h, head, sub, home=False, hs=40, ss=20):
        box(sl, x, y, w, h, ACC if home else BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        text(sl, x, y + 0.08, w, h * 0.55, head, hs, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(sl, x + 0.1, y + h * 0.55, w - 0.2, h * 0.42, sub, ss, "Source Sans 3", False, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)

    for i, (day, kind, title, body, notes) in enumerate(SLIDES, start=1):
        sl = prs.slides.add_slide(layout)
        items = [b for b in body.split("|") if b]
        if kind == "title":
            text(sl, 0, 0.42, 13.333, 0.5, BRAND, 16, "Oswald", False, BLACK, align=PP_ALIGN.CENTER, spacing=800, caps=True)
            pic(sl, "hero", 3.67, 0.98, 6.0)
            text(sl, 0, 4.2, 13.333, 1.2, title, 80, "Oswald", True, BLACK, align=PP_ALIGN.CENTER, caps=True)
            text(sl, 0, 5.5, 13.333, 0.6, body, 28, "Oswald", False, BLACK, align=PP_ALIGN.CENTER)
            box(sl, 0, 6.55, 13.333, 0.95, BLACK)
            text(sl, 0, 6.68, 13.333, 0.6, f"Digital Music Production · Unit {UNIT} · Day {day} · Chords and Roman numerals in Logic Pro", 20, "Oswald",
                 False, WHITE, align=PP_ALIGN.CENTER, spacing=200, caps=True)
        else:
            chrome(sl, i, title, day)
        if kind == "listen":
            head, _, sub = items[0].partition(" · ")
            box(sl, 0.6, 2.0, 12.13, 2.0, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            text(sl, 0.9, 2.1, 11.5, 0.7, head, 36, "Oswald", True, WHITE, caps=True)
            text(sl, 0.9, 2.82, 11.5, 0.5, sub, 24, "Oswald", False, LIGHT)
            text(sl, 0.9, 3.38, 11.5, 0.5, items[1], 18, "Source Sans 3", False, WHITE)
            bullets(sl, items[2:], y=4.3, size=30, gap=1.6)
        elif kind == "bullets":
            long_ = max(len(it) for it in items) > 46
            bullets(sl, items, size=(33 if long_ else 38) if len(items) <= 4 else 33, y=2.3, gap=2.35 if long_ else (2.2 if len(items) <= 4 else 1.95))
        elif kind == "closer":
            bullets(sl, items, size=38, y=2.1, gap=1.5)
            pic(sl, "hero", 4.12, 4.45, 5.1)
        elif kind == "feel":
            text(sl, 0.6, 1.95, 12.1, 0.5, "Play each one on a piano sound. Write what it feels like in your own words. Bright or dark?", 20,
                 "Source Sans 3", False, BLACK)
            for k, ch in enumerate(FEEL_CHORDS):
                row, col = k // 4, k % 4
                n_in_row = 4 if row == 0 else 3
                w = 2.85
                x0 = 0.6 + (12.13 - n_in_row * w - (n_in_row - 1) * 0.24) / 2
                x = x0 + col * (w + 0.24)
                y = 2.65 + row * 2.1
                tile(sl, x, y, w, 1.85, arrow(ch), f"Chord {k + 1} · feels like…", hs=34, ss=18)
        elif kind == "compare":
            for r, (lbl, rns, key, q) in enumerate([("Loop 1", LOOP_MAJOR, "C major", "Home: the first and last chord"),
                                                    ("Loop 2", LOOP_MINOR, "A minor", "Where is home now?")]):
                y = 2.05 + r * 2.35
                text(sl, 0.6, y, 2.0, 0.6, lbl, 30, "Oswald", True, ACC, caps=True)
                text(sl, 0.6, y + 0.62, 2.1, 1.2, q, 17, "Source Sans 3", False, BLACK)
                for k, rn in enumerate(rns):
                    ch = by_rn(key, rn)
                    tile(sl, 2.85 + k * 2.5, y, 2.3, 1.9, arrow(ch), "on a loop", home=(k == 0), hs=34, ss=16)
            text(sl, 0.6, 6.45, 12.1, 0.45, "Same white keys, nearly the same chords. A different home.", 20, "Source Sans 3", False, BLACK)
        elif kind == "keys":
            yb = pic(sl, "keys", 0.6, 2.0, 6.6)
            for k, (h, b) in enumerate([("Half step (H)", "One key in either direction, black keys included. E to F and B to C are the only "
                                         "white-key half steps."), ("Whole step (W)", "Two half steps: skip one key. C to D, D to E, G to A.")]):
                y = 2.0 + k * 2.4
                box(sl, 7.6, y, 5.13, 2.15, WHITE, BLACK if k else ACC, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=2.5)
                text(sl, 7.85, y + 0.15, 4.7, 0.6, h, 30, "Oswald", True, ACC if k == 0 else BLACK, caps=True)
                text(sl, 7.85, y + 0.8, 4.7, 1.3, b, 20, "Source Sans 3", False, BLACK)
            text(sl, 0.6, min(yb + 0.1, 6.45), 6.6, 0.45, "Find C: just left of every group of 2 black keys.", 18, "Source Sans 3", False, BLACK)
        elif kind == "scales":
            for r, (key, f) in enumerate([("C major", MAJOR_SCALE), ("A minor", MINOR_SCALE)]):
                sc, y = key_scale(key), 2.05 + r * 2.3
                text(sl, 0.6, y, 12.1, 0.55, f"{key.split()[0]} {'major' if 'major' in key else 'natural minor'}  ·  {f}", 26, "Oswald", True,
                     ACC if r == 0 else BLACK, caps=True)
                for k, nt in enumerate(sc):
                    x = 0.6 + k * 1.5
                    box(sl, x, y + 0.75, 1.0, 1.0, BLACK if k not in (0, 7) else ACC, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                    text(sl, x, y + 0.75, 1.0, 1.0, nt, 34, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                    if k < 7:
                        text(sl, x + 1.0, y + 0.95, 0.5, 0.6, steps_of(f)[k], 20, "Oswald", True, ACC, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            text(sl, 0.6, 6.45, 12.1, 0.45, "Same seven white keys, a different starting note, a different feel.", 20, "Source Sans 3", False, BLACK)
        elif kind == "days":
            for k, d in enumerate(DAYS):
                y = 1.98 + k * 0.98
                proj = d["n"] >= 3
                box(sl, 0.6, y, 1.55, 0.8, ACC if proj else BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                text(sl, 0.6, y, 1.55, 0.8, f"Day {d['n']}", 24, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, caps=True)
                text(sl, 2.4, y - 0.04, 10.3, 0.45, d["title"] + ("  ·  project" if proj else ""), 22, "Oswald", True, BLACK, caps=True)
                text(sl, 2.4, y + 0.4, 10.3, 0.4, DECK_DAY_LINES[k], 17, "Source Sans 3", False, BLACK)
        elif kind == "check":
            for k, it in enumerate(items):
                col, row = k % 2, k // 2
                x, y = 0.6 + col * 6.15, 2.15 + row * 1.5
                box(sl, x, y + 0.08, 0.42, 0.42, WHITE, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=2.25)
                text(sl, x + 0.6, y, 5.35, 1.4, it, 24, "Source Sans 3", False, BLACK)
        elif kind == "rubric":
            for k, ((c, _), short) in enumerate(zip(RUBRIC, RUBRIC_SHORT)):
                y = 2.1 + k * 1.35
                box(sl, 0.6, y, 4.4, 1.12, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                text(sl, 0.8, y, 4.0, 1.12, c, 22, "Oswald", True, WHITE, anchor=MSO_ANCHOR.MIDDLE, caps=True)
                text(sl, 5.25, y, 6.4, 1.12, "5 = " + short, 22, "Source Sans 3", False, BLACK, anchor=MSO_ANCHOR.MIDDLE)
                text(sl, 11.7, y, 1.03, 1.12, "/ 5", 26, "Oswald", True, ACC, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
            text(sl, 7.0, 6.3, 5.73, 0.5, f"Scored 2 to 5 · Total / {RUBRIC_TOTAL}", 24, "Oswald", True, BLACK, align=PP_ALIGN.RIGHT, caps=True)
        elif kind == "chordkeys":
            yb = pic(sl, "chordkeys", 0.6, 2.0, 12.13)
            for k, (a, f, d) in enumerate(FORMULAS_BOX):
                x = 0.6 + k * 4.14
                box(sl, x, yb + 0.25, 3.85, 1.55, WHITE, ACC if k == 0 else BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=2.25)
                text(sl, x, yb + 0.32, 3.85, 0.65, f"{a}: {f}", 28, "Oswald", True, ACC if k == 0 else BLACK, align=PP_ALIGN.CENTER, caps=True)
                text(sl, x, yb + 1.0, 3.85, 0.6, d, 20, "Source Sans 3", False, BLACK, align=PP_ALIGN.CENTER)
            text(sl, 0.6, 6.5, 12.1, 0.42, "Do not count the root: for C, C# is 1. C-E-G = 4 up, then 3 up.", 19, "Source Sans 3", False, BLACK)
        elif kind == "steps":
            for k, it in enumerate(items):
                y = 2.05 + k * 0.95
                box(sl, 0.6, y + 0.05, 0.62, 0.62, BLACK, shape=MSO_SHAPE.OVAL)
                text(sl, 0.6, y + 0.05, 0.62, 0.62, str(k + 1), 24, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                text(sl, 1.45, y - 0.05, 11.2, 0.9, it, 28, "Source Sans 3", False, BLACK, anchor=MSO_ANCHOR.MIDDLE)
        elif kind == "rnfeel":
            text(sl, 0.6, 1.95, 12.1, 0.6, "I – IV – V – I   =   home, away, very away, home", 28, "Oswald", True, ACC, caps=True)
            for k, rn in enumerate(["I", "IV", "V", "vi"]):
                row = next(r for r in CARD["C major"] if r[0] == rn)
                tile(sl, 0.6 + k * 3.08, 2.85, 2.85, 3.0, rn, row[4], home=(rn == "I"), hs=66, ss=21)
            text(sl, 0.6, 6.15, 12.1, 0.7, "Uppercase = major · lowercase = minor · ° = diminished. A letter name tells you the notes; a Roman numeral "
                 "tells you the job.", 18, "Source Sans 3", False, BLACK)
        elif kind == "card":
            cols = [("RN", 1.15), ("Chord", 2.3), ("Notes", 2.3), ("Feels like…", 6.38)]
            rows = [tuple(c for c in (r[0], r[2], r[3], r[4])) for r in CARD["C major"]]
            y = 1.95
            text(sl, 0.6, y, 12.1, 0.45, CARD_SCALE_TEXT["C major"], 20, "Oswald", True, ACC, caps=True)
            y += 0.55
            x = 0.6
            for (h, w) in cols:
                box(sl, x, y, w - 0.04, 0.48, BLACK)
                text(sl, x + 0.1, y, w - 0.2, 0.48, h, 17, "Oswald", True, WHITE, anchor=MSO_ANCHOR.MIDDLE, caps=True)
                x += w
            for k, r in enumerate(rows):
                yy = y + 0.52 + k * 0.54
                x = 0.6
                home = r[0] == "I"
                for c, (val, (h, w)) in enumerate(zip(r, cols)):
                    box(sl, x, yy, w - 0.04, 0.5, ACC if home and c == 0 else WHITE, BLACK, lw=1.25)
                    text(sl, x + 0.1, yy, w - 0.2, 0.5, val, 24 if c == 0 else 19, "Oswald" if c < 3 else "Source Sans 3", c == 0,
                         WHITE if home and c == 0 else BLACK, align=PP_ALIGN.CENTER if c == 0 else PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
                    x += w
        elif kind == "roll":   # the worked example as native blocks (projector-sized): RN tile, chord tones (5th on top), bass root
            lay = roll_layout("C major")
            cw, gap, x0 = 2.45, 0.2, 2.1
            text(sl, 0.6, 3.35, 1.4, 0.6, "Chords", 24, "Oswald", True, ACC, caps=True)
            text(sl, 0.6, 5.2, 1.4, 0.6, "Bass", 24, "Oswald", True, BLACK, caps=True)
            for m, (rn, ch, vo, bass) in enumerate(lay):
                x = x0 + m * (cw + gap)
                home = rn == "I"
                box(sl, x, 1.95, cw, 0.92, ACC if home else BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                text(sl, x, 1.95, cw, 0.58, rn, 34, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                text(sl, x, 2.5, cw, 0.35, chord_name(ch), 16, "Oswald", False, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                for k, nt in enumerate(reversed(ch)):
                    y = 3.0 + k * 0.66
                    box(sl, x, y, cw, 0.58, ACC, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                    text(sl, x, y, cw, 0.58, nt, 26, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                box(sl, x, 5.18, cw, 0.58, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                text(sl, x, 5.18, cw, 0.58, midi_name(bass), 26, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                text(sl, x, 5.85, cw, 0.4, f"Measure {m + 1}", 16, "Source Sans 3", False, BLACK, align=PP_ALIGN.CENTER)
            text(sl, 0.6, 6.4, 12.1, 0.45, "Teal: the chords you build (3 notes stacked). Black: the bassline in the file. Every chord’s root = the bass note.",
                 18, "Source Sans 3", False, BLACK, align=PP_ALIGN.CENTER)
        sl.notes_slide.notes_text_frame.text = notes
    prs.save(DECK)


def ttf_fonts():
    """woff2 -> ttf for LibreOffice (deck PDF), plus a fontconfig file that adds them."""
    from fontTools.ttLib import TTFont
    d = os.path.join(BUILD, "ttf")
    os.makedirs(d, exist_ok=True)
    styles = {"oswald-latin-500-normal": ("Oswald", "Regular", 400), "oswald-latin-700-normal": ("Oswald", "Bold", 700),
              "source-sans-3-latin-400-normal": ("Source Sans 3", "Regular", 400), "source-sans-3-latin-700-normal": ("Source Sans 3", "Bold", 700),
              "source-sans-3-latin-400-italic": ("Source Sans 3", "Italic", 400)}
    for stem, (fam, sub, wt) in styles.items():
        out = os.path.join(d, stem + ".ttf")
        if os.path.exists(out):
            continue
        t = TTFont(os.path.join(FONTS, stem + ".woff2"))
        t.flavor = None
        name = t["name"]
        ps = (fam.replace(" ", "") + "-" + sub)
        for rec in list(name.names):
            if rec.nameID in (16, 17, 21, 22):
                name.removeNames(nameID=rec.nameID)
        for nid, val in ((1, fam), (2, sub), (4, f"{fam} {sub}" if sub != "Regular" else fam), (6, ps)):
            name.setName(val, nid, 3, 1, 0x409)
            name.setName(val, nid, 1, 0, 0)
        t["OS/2"].usWeightClass = wt
        sel = 0
        if sub == "Bold":
            sel |= 1 << 5
        if sub == "Italic":
            sel |= 1
        if sub == "Regular":
            sel |= 1 << 6
        t["OS/2"].fsSelection = (t["OS/2"].fsSelection & ~0b1100001) | sel
        t["head"].macStyle = (1 if sub == "Bold" else 0) | (2 if sub == "Italic" else 0)
        t.save(out)
    conf = os.path.join(BUILD, "fonts.conf")
    with open(conf, "w") as fh:
        fh.write(f'<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig><include ignore_missing="yes">/etc/fonts/fonts.conf</include>'
                 f'<dir>{d}</dir><cachedir>{os.path.join(BUILD, "fccache")}</cachedir></fontconfig>')
    return conf


def deck_pdf(src=DECK, out=DECK_PDF, fonts=True):
    env = dict(os.environ, FONTCONFIG_FILE=ttf_fonts()) if fonts else dict(os.environ)
    prof = os.path.join(BUILD, "lo-profile" if fonts else "lo-profile-fallback")
    tmp = os.path.join(BUILD, "lo-out")
    os.makedirs(tmp, exist_ok=True)
    subprocess.run(["soffice", f"-env:UserInstallation=file://{prof}", "--headless", "--convert-to", "pdf", "--outdir", tmp, src],
                   check=True, env=env, capture_output=True)
    shutil.move(os.path.join(tmp, os.path.basename(src).replace(".pptx", ".pdf")), out)


# ================================================================ square images (2000 x 2000)
SQ_CSS = """
.sq { width: 1000px; height: 1000px; position: relative; overflow: hidden; background: linear-gradient(180deg,#fff 0%,#fff 60%,#eeeeee 100%); }
.sq .store { position: absolute; left: 0; right: 0; top: 30px; text-align: center; font-family: 'Oswald'; font-weight: 500; letter-spacing: .32em;
             text-transform: uppercase; font-size: 22px; }
.sq .fillw { position: absolute; left: 0; right: 0; top: 66px; bottom: 70px; display: flex; flex-direction: column; align-items: center; justify-content: space-evenly; }
.sq .hd { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; text-align: center; line-height: .95; white-space: nowrap; }
.sq .sub { font-family: 'Oswald'; font-weight: 300; text-align: center; white-space: nowrap; }
.sq .band { position: absolute; left: 0; right: 0; bottom: 0; background: #000; color: #fff; text-align: center; font-family: 'Oswald';
            font-weight: 500; text-transform: uppercase; letter-spacing: .1em; font-size: 26px; padding: 16px 0 18px; white-space: nowrap; }
.sq .band em { font-style: normal; color: ACCENT; }
.shot { background: #fff; border: 2px solid #000; box-shadow: 8px 8px 0 #000; position: relative; overflow: hidden; }
.shot img { display: block; width: 100%; }
.lbl { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 22px; text-align: center; margin-top: 14px; }
.lbl span { font-weight: 300; text-transform: none; display: block; font-size: 18px; }
.swm { position: absolute; left: 50%; top: 50%; transform: translate(-50%,-50%) rotate(-38deg); font-family: 'Oswald'; font-weight: 700;
       color: rgba(0,0,0,.14); white-space: nowrap; letter-spacing: .06em; }
""".replace("ACCENT", ACCENT)
FONT_PRIME = ('<div style="position:absolute;left:-9999px;font-family:\'Source Sans 3\'">'
              '<span style="font-weight:400">a</span><span style="font-weight:600">a</span><i style="font-weight:400">a</i></div>')


def sq_doc(body):
    return (f'<!doctype html><html><head><meta charset="utf-8"><style>{font_css()}'
            f'html,body{{margin:0;background:#fff}}{SQ_CSS}</style></head><body>{FONT_PRIME}{body}</body></html>')


def sq(inner, band):
    return sq_doc(f'<div class="sq"><div class="store">{BRAND}</div><div class="fillw">{inner}</div><div class="band">{band}</div></div>')


SQ_SUB = f"{DAYS_N}-day unit · I–IV–V–I by ear · Roman numerals · intro slides"
SQ_BAND = f"Grades 9–12 · {DAYS_N} days · {len(PLAN)} pages + {DECK_SLIDES} slides"


def cover_square():
    s = SQ_COVER
    return sq(f'<div style="width:880px">{hero_svg()}</div>'
              f'<div><div class="hd" style="font-size:{s["a"]}px">{COVER_A}</div>'
              f'<div class="hd" style="font-size:{s["b"]}px;margin-top:4px">{COVER_B}</div>'
              f'<div class="sub" style="font-size:{s["c"]}px;margin-top:10px">{SQ_SUB}</div></div>',
              SQ_BAND.replace(" · ", " <em>·</em> "))


SQ_HEADS = {1: ("Ready to teach Monday", 76, "Teacher notes, five day plans, student pages and the rubric in one PDF", 28),
            2: ("Intro deck included", 76, f"{DECK_SLIDES} editable slides (.pptx + PDF) with speaker notes on every slide", 28),
            3: ("Everything inside", 72, "Feel Log, packet, KCSRN, reference cards, answer keys", 28)}
SQ_BANDS = {1: "Hear it <em>·</em> name it <em>·</em> build it <em>·</em> present", 2: "Hook <em>·</em> listen <em>·</em> worked example "
            "<em>·</em> steps <em>·</em> rubric", 3: f"{len(PLAN)} pages <em>·</em> US Letter PDF <em>·</em> prints in black &amp; white"}


def sq_head(i):
    h, hs, s, ss = SQ_HEADS[i]
    return f'<div><div class="hd" style="font-size:{hs}px">{h}</div><div class="sub" style="font-size:{ss}px;margin-top:8px">{s}</div></div>'


def preview_square_1(img):
    rot = [-4, 0, 4]
    shots = "".join(f'<div style="position:absolute;left:{x}px;top:{y}px;width:372px;transform:rotate({r}deg);z-index:{z}">'
                    f'<div class="shot"><img src="{s}"><div class="swm" style="font-size:80px">PREVIEW</div></div></div>'
                    for (s, _, _), x, y, r, z in zip(img, (40, 314, 588), (24, 0, 24), rot, (1, 3, 2)))
    labels = "".join(f'<div class="lbl" style="width:300px">{a}<span>{b}</span></div>' for _, a, b in img)
    return sq(f'{sq_head(1)}<div style="position:relative;width:1000px;height:520px">{shots}</div>'
              f'<div style="display:flex;gap:30px">{labels}</div>', SQ_BANDS[1])


def preview_square_2(slides):
    cells = "".join(f'<div class="shot" style="box-shadow:6px 6px 0 #000"><img src="{s}"><div class="swm" style="font-size:52px">PREVIEW</div></div>'
                    for s in slides)
    return sq(f'{sq_head(2)}<div style="display:grid;grid-template-columns:388px 388px;gap:20px 34px">{cells}</div>', SQ_BANDS[2])


def preview_square_3(imgs):
    cells = "".join(f'<div style="width:214px"><div class="shot" style="box-shadow:5px 5px 0 #000"><img src="{s}">'
                    f'<div class="swm" style="font-size:46px">PREVIEW</div></div><div class="lbl" style="font-size:18px;margin-top:9px">{lab_}</div></div>'
                    for s, lab_ in imgs)
    return sq(f'{sq_head(3)}<div style="display:grid;grid-template-columns:repeat(4,214px);gap:22px 24px">{cells}</div>', SQ_BANDS[3])


SHOTS1 = [("pacing", "Pacing + plans", "5 days, minute by minute"), ("feellog", "Chord Feel Log", "the ear-first Day 1"),
          ("cardC", "Reference cards", "C major and A minor")]
SHOTS2 = [1, 4, 6, 14, 16, 19]   # deck slides
SHOTS3 = [("post1", "Day 1 post"), ("pk1", "Keyboard"), ("pk3", "Chords"), ("kcsrn", "KCSRN"),
          ("cardA", "A minor card"), ("proj1", "Project sheet"), ("rubric", "Rubric / 15"), ("example", "Worked example")]


# ================================================================ render
def node_env():
    env = dict(os.environ)
    if not env.get("NODE_PATH"):
        env["NODE_PATH"] = subprocess.run(["npm", "root", "-g"], capture_output=True, text=True).stdout.strip()
    return env


def render(jobs):
    path = os.path.join(BUILD, "jobs.json")
    with open(path, "w") as f:
        json.dump(jobs, f)
    subprocess.run(["node", RENDER_JS, path], check=True, env=node_env())


OVERFLOW_JS = r"""
const { chromium } = require('playwright');
const url = require('url');
(async () => {
  const browser = await chromium.launch();
  const page = await browser.newPage();
  await page.goto(url.pathToFileURL(process.argv[2]).href, { waitUntil: 'load' });
  await page.evaluate(() => document.fonts.ready);
  await page.emulateMedia({ media: 'print' });
  const res = await page.evaluate(() => [...document.querySelectorAll('section.page')].map((s, i) => {
    const m = s.querySelector('.main');
    const out = { page: i + 1, key: s.dataset.key, main_over: m ? Math.max(0, m.scrollHeight - m.clientHeight) : 0, cells: [] };
    // anything that sticks out of the page's main area (right or bottom) is clipped or overprinted
    if (m) {
      const mb = m.getBoundingClientRect();
      m.querySelectorAll('*').forEach(el => {
        const r = el.getBoundingClientRect();
        if (r.width && (r.right > mb.right + 1 || r.bottom > mb.bottom + 1 || r.left < mb.left - 1))
          out.cells.push((el.tagName + ': ' + (el.textContent || '').trim()).slice(0, 50));
      });
    }
    return out;
  }));
  require('fs').writeFileSync(process.argv[3], JSON.stringify(res, null, 1));
  await browser.close();
})().catch(e => { console.error(e); process.exit(1); });
"""


def overflow_scan(html):
    js = os.path.join(BUILD, "overflow.js")
    with open(js, "w") as f:
        f.write(OVERFLOW_JS)
    subprocess.run(["node", js, html, OVERFLOW_JSON], check=True, env=node_env())


def write(name, txt):
    p = os.path.join(BUILD, name)
    with open(p, "w") as f:
        f.write(txt)
    return p


def pdf_png(src, page_no, out, res=110):
    subprocess.run(["pdftoppm", "-png", "-r", str(res), "-f", str(page_no), "-l", str(page_no), "-singlefile", src,
                    os.path.join(BUILD, out)], check=True)


def art_doc(w, h, inner):
    return (f'<!doctype html><html><head><meta charset="utf-8"><style>{font_css()}{CSS}'
            f'html,body{{margin:0;background:#fff}}</style></head>'
            f'<body><div style="width:{w}px;height:{h}px;overflow:hidden;background:#fff">{inner}</div></body></html>')


def deck_art():
    """PNG art for the deck (original drawings and the C major card table, rendered with the brand fonts)."""
    svgs = {"hero": (1800, hero_svg()), "keys": (1500, keys_svg()), "chordkeys": (2000, chordkeys_svg())}
    arts = {}
    for k, (w, svg) in svgs.items():   # exact height from the drawing's own viewBox, so nothing is cropped and nothing is padded
        vw, vh = map(float, re.search(r'viewBox="0 0 ([\d.]+) ([\d.]+)"', svg).groups())
        arts[k] = (w, int(round(w * vh / vw)) + 2, f'<div style="width:{w}px">{svg}</div>')
    jobs, out = [], {}
    for k, (w, h, inner) in arts.items():
        p = os.path.join(BUILD, f"deck-{k}.png")
        jobs.append({"html": write(f"deck-{k}.html", art_doc(w, h, inner)), "png": p, "width": w, "height": h, "scale": 1})
        out[k] = (p, (w, h))
    return jobs, out


def tpt_json():
    return {
        "price": PRICE, "license_price": PRICE,
        "tax_code": "Other Digital Goods - No Physical Media",
        "grades": ["9th-grade", "10th-grade", "11th-grade", "12th-grade"],
        "subjects": ["Music"],
        "tags": ["Autumn"],
        "formats": ["PDF"],
        "pages": len(PLAN) + DECK_SLIDES,
        "answer_key": "Included",
        "zip_contents": [os.path.basename(PDF), os.path.basename(DECK), os.path.basename(DECK_PDF)],
        "files": {"product": os.path.basename(ZIP), "preview": os.path.basename(PREVIEW_PDF), "thumb1": "cover.png",
                  "thumb2": "preview-1.png", "thumb3": "preview-2.png", "thumb4": "preview-3.png"},
    }


def make_zip():
    if os.path.exists(ZIP):
        os.remove(ZIP)
    with zipfile.ZipFile(ZIP, "w", zipfile.ZIP_DEFLATED) as z:
        for p in (PDF, DECK, DECK_PDF):
            zi = zipfile.ZipInfo(f"{ZIP_DIR}/{os.path.basename(p)}", date_time=(2026, 10, 13, 12, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            with open(p, "rb") as fh:
                z.writestr(zi, fh.read())


def build():
    os.makedirs(BUILD, exist_ok=True)
    problems = verify_theory() + verify_data() + verify_labels()
    if problems:
        print("\n".join(problems))
        sys.exit("data check failed; not rendering")
    full = write("resource.html", html_doc([render_page(p) for p in PLAN], TITLE))
    prev = write("preview.html", html_doc([render_page(p, watermark=True) for p in PREVIEW_PICKS], "Preview"))
    art_jobs, art = deck_art()
    render([{"html": full, "pdf": PDF}, {"html": prev, "pdf": PREVIEW_PDF}] + art_jobs)
    overflow_scan(full)
    build_deck(art)
    deck_pdf()
    make_zip()
    res = 45 if LOW_RES else 120
    for key, _, _ in SHOTS1:
        pdf_png(PDF, IX[key], key, res)
    for key, _ in SHOTS3:
        pdf_png(PDF, IX[key], key, 40 if LOW_RES else 80)
    for s in SHOTS2:
        pdf_png(DECK_PDF, s, f"slide{s}", 30 if LOW_RES else 70)
    sqs = [write("cover.html", cover_square()), write("p1.html", preview_square_1([(f"{k}.png", a, b) for k, a, b in SHOTS1])),
           write("p2.html", preview_square_2([f"slide{s}.png" for s in SHOTS2])),
           write("p3.html", preview_square_3([(f"{k}.png", lab_) for k, lab_ in SHOTS3]))]
    render([{"html": h, "png": o, "width": 1000, "height": 1000, "scale": 1 if LOW_RES else 2} for h, o in zip(sqs, [COVER_PNG] + PREVIEW_PNGS)])
    with open(TPT_JSON, "w") as f:
        json.dump(tpt_json(), f, indent=2, ensure_ascii=False)
        f.write("\n")
    print(f"built {PDF} ({len(PLAN)} pages), {DECK} + PDF ({DECK_SLIDES} slides), {ZIP}, preview, cover + 3 previews, tpt.json"
          + (" [LOW RES]" if LOW_RES else ""))


# ================================================================ checks
def verify_theory():
    """Music theory, computed: scales, key signatures, chord spellings, Roman numerals, the owner's cards and every answer printed."""
    errs = []
    for f in [MAJOR_SCALE, MINOR_SCALE, MAJOR_PENT] + [b for _, b, _ in REF_SCALES]:
        if sum(STEP[s] for s in steps_of(f)) != 12:
            errs.append(f"scale formula {f} does not add up to an octave")
    exp = {"C major": "C D E F G A B C", "A minor": "A B C D E F G A"}
    for key, want in exp.items():
        sc = key_scale(key)
        if " ".join(sc) != want:
            errs.append(f"{key} scale is {' '.join(sc)}, expected {want}")
        if accidentals(sc):
            errs.append(f"{key}: key signature should have no sharps or flats, found {accidentals(sc)}")
        if any(pc(n) not in NAT.values() for n in sc):
            errs.append(f"{key}: a scale note is not a white key")
    if set(key_scale("C major")[:7]) != set(key_scale("A minor")[:7]):
        errs.append("C major and A minor must share the same seven notes (relative keys)")
    if " ".join(scale("C", MAJOR_PENT)) != "C D E G A C":
        errs.append(f"C major pentatonic is {' '.join(scale('C', MAJOR_PENT))}")
    # the owner's cards against the computed diatonic triads
    for key, rows in CARD.items():
        comp = key_chords(key)
        if len(rows) != 7:
            errs.append(f"{key} card: {len(rows)} rows")
        for (rn, typ, name, notes, feel), (crn, cq, cn) in zip(rows, comp):
            if rn != crn:
                errs.append(f"{key} card: Roman numeral {rn}, computed {crn}")
            if typ != cq:
                errs.append(f"{key} card {rn}: type {typ}, computed {cq}")
            if notes != dash(cn):
                errs.append(f"{key} card {rn}: notes {notes}, computed {dash(cn)}")
            if name != chord_name(cn):
                errs.append(f"{key} card {rn}: name {name}, computed {chord_name(cn)}")
            if chord(cn[0], FORMULA[cq]) != cn:
                errs.append(f"{key} {rn}: formula {FORMULA[cq]} from {cn[0]} gives {chord(cn[0], FORMULA[cq])}, not {cn}")
            if accidentals(cn):
                errs.append(f"{key} {rn}: not white keys only")
            if not feel:
                errs.append(f"{key} card {rn}: no feeling")
    if [(r, q) for r, q, _ in key_chords("C major")] != BLANK_CARD:
        errs.append("blank card RN/Type column must be the major-key pattern")
    # packet RN table: each 'X / y' pair matches C major / A minor (major / minor) on the same degree
    for pair, typ, _, _ in PACKET_RN:
        a, b = [x.strip() for x in pair.split("/")]
        ta, tb = [x.strip() for x in typ.split("/")]
        ra = next(((r, q) for r, q, _ in key_chords("C major") if r == a), None)
        rb = next(((r, q) for r, q, _ in key_chords("A minor") if r == b), None)
        qmap = {"Major": "Major", "Minor": "Minor", "Dim.": "Dim."}
        if not ra or not rb or qmap[ra[1]] != ta or qmap[rb[1]] != tb or ROMAN.index(a.upper().rstrip("°")) != ROMAN.index(b.upper().rstrip("°")):
            errs.append(f"packet RN row {pair} ({typ}) does not match C major / A minor")
    for rn, typ, _ in REF_RN:
        q = next((q for r, q, _ in key_chords("C major") if r == rn), None)
        if q is None or {"Major": "Major", "Minor": "Minor", "Dim.": "Diminished"}[q] != typ:
            errs.append(f"reference sheet RN {rn} {typ} does not match C major")
    # packet chord exercises and the reference sheet examples
    want = {("C", "R-4-3"): "C E G", ("F", "R-4-3"): "F A C", ("G", "R-4-3"): "G B D", ("D", "R-4-3"): "D F# A",
            ("A", "R-3-4"): "A C E", ("D", "R-3-4"): "D F A", ("E", "R-3-4"): "E G B", ("B", "R-3-4"): "B D F#",
            ("B", "R-3-3"): "B D F", ("D", "R-3-3"): "D F Ab"}
    for kind, formula, _, roots in PACKET_CHORDS:
        for r in roots:
            got = " ".join(chord(r, formula))
            if want.get((r, formula)) != got:
                errs.append(f"packet {r} {kind}: {got}, expected {want.get((r, formula))}")
    for typ, formula, ex, notes, _ in REF_CHORDS:
        if chord("C", formula) != notes:
            errs.append(f"reference sheet {ex}: printed {notes}, formula gives {chord('C', formula)}")
    # half step / whole step answers
    hw = {("C", "D"): 2, ("C#", "D"): 1, ("B", "C"): 1, ("E", "F"): 1, ("D", "E"): 2, ("A", "B"): 2, ("F", "G"): 2, ("G#", "A"): 1}
    for a, b in HW_EXAMPLES + HW_PAIRS:
        if semis_between(a, b) != hw[(a, b)]:
            errs.append(f"half/whole step {a}-{b} wrong")
    # Day 1: the seven white-key chords, both loops; the worked example (bass = chord root, chords from the key)
    if sorted(arrow(c) for c in FEEL_CHORDS) != sorted(arrow(n) for _, _, n in key_chords("C major")):
        errs.append("Feel Log must hold all seven white-key triads")
    if [arrow(by_rn("C major", r)) for r in LOOP_MAJOR] != ["C–E–G", "F–A–C", "G–B–D", "C–E–G"]:
        errs.append("Loop 1 must be I-IV-V-I in C: C-E-G F-A-C G-B-D C-E-G")
    if [arrow(by_rn("A minor", r)) for r in LOOP_MINOR] != ["A–C–E", "F–A–C", "C–E–G", "G–B–D"]:
        errs.append("Loop 2 must be i-VI-III-VII in A minor: A-C-E F-A-C C-E-G G-B-D")
    for key in EX_PROG:
        sc = set(key_scale(key))
        for rn, ch, vo, bass in roll_layout(key):
            if not set(ch) <= sc:
                errs.append(f"example {key} {rn}: {ch} not in the key")
            if midi_name(bass) != ch[0] or midi_name(vo[0]) != ch[0]:
                errs.append(f"example {key} {rn}: bass/voicing root is not {ch[0]}")
            if [midi_name(v) for v in vo] != ch or not (ROLL_CH[0] <= min(vo) and max(vo) <= ROLL_CH[1]):
                errs.append(f"example {key} {rn}: voicing {vo} wrong or off the roll")
            if [vo[1] - vo[0], vo[2] - vo[1]] != [int(x) for x in FORMULA[quality(ch)].split("-")[1:]]:
                errs.append(f"example {key} {rn}: voicing is not root position by formula")
            if not ROLL_BS[0] <= bass <= ROLL_BS[1]:
                errs.append(f"example {key} {rn}: bass off the roll")
    if {arrow(by_rn("C major", r)) for r in EX_PROG["C major"]} != {arrow(by_rn("A minor", r)) for r in EX_PROG["A minor"]}:
        errs.append("the two worked examples should use the same four chords")
    # songs: the progression facts named in the deck match the engine
    for title, artist, key, rns, tonic, names, _ in SONGS:
        sc = scale(tonic, MAJOR_SCALE)
        got = [short_name(next(n for r, q, n in diatonic(sc) if r == rn)) for rn in rns]
        if got != names:
            errs.append(f"{title}: {rns} in {key} gives {got}, not {names}")
    return errs


def verify_data():
    errs = []
    counts = dict(req=(REQUIREMENTS, 9), check=(CHECKLIST, 7), tips=(TIPS, 4), struggles=(STRUGGLES, 6), opening=(OPENING, 6),
                  vocab=(VOCAB, 12), quiz=(VOCAB_QUIZ, 4), formulas=(FORMULAS_BOX, 3))
    for nm, (lst, nexp) in counts.items():
        if len(lst) != nexp:
            errs.append(f"{nm}: {len(lst)} items, expected {nexp}")
    vw = {w.lower() for w, _ in VOCAB}
    for _, a in VOCAB_QUIZ:
        if a.lower() not in vw:
            errs.append(f"vocab quiz answer '{a}' is not a vocabulary word")
    if len(RUBRIC) != 3 or any(len(dd) != len(LEVELS) for _, dd in RUBRIC) or RUBRIC_TOTAL != 15 or len(RUBRIC_SHORT) != 3:
        errs.append("rubric must be the owner's 3 categories x levels 2-5 = /15")
    if RUBRIC[2][1][3] != "Correctly named, on time. Student explains chord choices using Roman numeral language.":
        errs.append("rubric: Documentation level 5 must keep the owner's wording")
    # pacing: 5 days, blocks 0-53, every deck slide in exactly one block of the day the slide says
    if len(DAYS) != DAYS_N or [d["n"] for d in DAYS] != list(range(1, DAYS_N + 1)) or len(DECK_DAY_LINES) != DAYS_N:
        errs.append(f"pacing must be {DAYS_N} days (and {DAYS_N} on the deck)")
    seen = {}
    for d in DAYS:
        blocks = [tuple(map(int, x.split("–"))) for x, _, _, _ in d["blocks"]]
        if blocks[0][0] != 0 or blocks[-1][1] != MINUTES or any(blocks[i][1] != blocks[i + 1][0] for i in range(len(blocks) - 1)):
            errs.append(f"Day {d['n']} minute blocks do not run 0-{MINUTES} without gaps")
        if not day_slides(d):
            errs.append(f"Day {d['n']} has no deck slide")
        for _, h, txt, sl in d["blocks"]:
            nameds = set()
            for a_, b_ in re.findall(r"[Ss]lides? (\d+)(?:–(\d+))?", txt):
                nameds |= set(range(int(a_), int(b_ or a_) + 1))
            if not sl:
                if nameds:
                    errs.append(f"Day {d['n']} '{h}' names a slide but has no slide slot")
                continue
            if nameds != set(range(sl[0], sl[1] + 1)):
                errs.append(f"Day {d['n']} '{h}': text names slides {sorted(nameds)} but the slot is {sl}")
            for s in range(sl[0], sl[1] + 1):
                if s in seen:
                    errs.append(f"slide {s} is slotted twice (Day {seen[s]} and Day {d['n']})")
                seen[s] = d["n"]
                if not 1 <= s <= DECK_SLIDES:
                    errs.append(f"Day {d['n']} names slide {s}, which does not exist")
                elif slide_day(s) != d["n"]:
                    errs.append(f"slide {s} says Day {slide_day(s)} but the Day {d['n']} plan uses it")
    for i in range(1, DECK_SLIDES + 1):
        if i not in seen:
            errs.append(f"deck slide {i} has no day slot in the day plans")
    if [s[0] for s in SLIDES] != sorted(s[0] for s in SLIDES):
        errs.append("deck slides must run in day order")
    if len(SLIDES) < 12 or any(len(s[4].strip()) < 80 for s in SLIDES):
        errs.append("deck needs 12+ slides, each with real speaker notes")
    kinds = [k for _, k, *_ in SLIDES]
    for need in ("title", "listen", "steps", "check", "rubric", "roll"):
        if need not in kinds:
            errs.append(f"deck pattern: no '{need}' slide")
    # Day 1 is ear-only: no Roman numerals on Day 1 slides or the Day 1 student pages' text
    rn_pat = re.compile(r"\b(?:I|II|III|IV|V|VI|VII|i|ii|iii|iv|v|vi|vii)°?(?:[–-](?:I|II|III|IV|V|VI|VII|i|ii|iii|iv|v|vi|vii)°?)+\b")
    for day, kind, title, body, notes in SLIDES:
        if day == 1 and rn_pat.search(title + " " + body):
            errs.append(f"Day 1 slide '{title}' shows Roman numerals (feel it before you name it)")
    if len(PLAN) != len(set(PLAN)):
        errs.append("plan has a repeated page")
    if len(LISTING_TITLE) > 80:
        errs.append(f"listing title is {len(LISTING_TITLE)} chars")
    # cover and square lines fit their width on one line (Oswald 700 metrics; light weights are narrower)
    for label, sz, room in (("cover A", PDF_COVER["a"], PDF_COVER["room"]), ("cover B", PDF_COVER["b"], PDF_COVER["room"])):
        txt = COVER_A if label.endswith("A") else COVER_B
        if text_w(txt.upper(), sz) * 1.01 > room:
            errs.append(f"PDF {label} '{txt}' is {text_w(txt.upper(), sz):.0f}pt wide, room {room:.0f}")
    if text_w(COVER_C, PDF_COVER["c"]) > PDF_COVER["room"]:
        errs.append("PDF cover subtitle too wide")
    if text_w(COVER_BAND.upper(), PDF_COVER["band"]) * 1.12 > PDF_COVER["room"]:
        errs.append("PDF cover band too wide")
    for txt, sz in ((COVER_A.upper(), SQ_COVER["a"]), (COVER_B.upper(), SQ_COVER["b"]), (SQ_SUB, SQ_COVER["c"])):
        if text_w(txt, sz) * 1.01 > SQ_COVER["room"]:
            errs.append(f"square cover line '{txt}' is {text_w(txt, sz):.0f}px wide, room {SQ_COVER['room']}")
    if text_w(SQ_BAND.upper(), 26) * 1.12 > 960:
        errs.append("square cover band too wide")
    if text_w(COVER_B.upper(), SQ_COVER["b"]) / SQ_COVER["b"] < 0 or SQ_COVER["b"] * 0.75 < 100:
        errs.append("square cover big line under 10% of the image height")
    for i, (h, hs, s, ss) in SQ_HEADS.items():
        if text_w(h.upper(), hs) * 1.01 > 960 or text_w(s, ss) > 960:
            errs.append(f"preview {i} headline/subtitle too wide")
    return errs


def pdf_fonts(path):
    out = subprocess.run(["pdffonts", path], capture_output=True, text=True, check=True).stdout.splitlines()[2:]
    return [(ln.split()[0], ln.split()[-5]) for ln in out]


def pdf_text(path):
    return subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True, check=True).stdout


def deck_texts():
    from pptx import Presentation
    prs = Presentation(DECK)
    slides, notes = [], []
    for s in prs.slides:
        slides.append(" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame))
        notes.append(s.notes_slide.notes_text_frame.text if s.has_notes_slide else "")
    media = [p for p in zipfile.ZipFile(DECK).namelist() if p.startswith("ppt/media/")]
    return prs, slides, notes, media


def deck_overflow(prs):
    """Every text box on every slide stays inside the slide; Oswald titles fit their box on one line."""
    errs = []
    W, H = prs.slide_width, prs.slide_height
    for i, s in enumerate(prs.slides, start=1):
        for sh in s.shapes:
            if sh.left < 0 or sh.top < 0 or sh.left + sh.width > W + 10 or sh.top + sh.height > H + 10:
                errs.append(f"deck slide {i}: a shape runs off the slide")
            if sh.has_text_frame and sh.text_frame.text:
                r = sh.text_frame.paragraphs[0].runs
                if r and r[0].font.name == "Oswald" and r[0].font.size:
                    widest = max((text_w(p.text, r[0].font.size.pt) for p in sh.text_frame.paragraphs), default=0)
                    room = sh.width / 12700 - 6
                    if widest > room * 1.0 and len(sh.text_frame.text) > 0:
                        lines_avail = int((sh.height / 12700) // (r[0].font.size.pt * 1.15)) or 1
                        if widest > room * lines_avail:
                            errs.append(f"deck slide {i}: '{sh.text_frame.text[:30]}' ({widest:.0f}pt) does not fit {room:.0f}pt x {lines_avail} line(s)")
    return errs


def verify_files():
    errs = []
    from pypdf import PdfReader
    from PIL import Image
    for p in [PDF, DECK, DECK_PDF, ZIP, PREVIEW_PDF, COVER_PNG, *PREVIEW_PNGS, UPLOAD, TPT_JSON, OVERFLOW_JSON]:
        if not os.path.exists(p):
            errs.append(f"missing {os.path.basename(p)}")
    if errs:
        return errs
    for path, n_exp in ((PDF, len(PLAN)), (PREVIEW_PDF, len(PREVIEW_PICKS))):
        rd = PdfReader(path)
        if len(rd.pages) != n_exp:
            errs.append(f"{os.path.basename(path)}: {len(rd.pages)} pages, expected {n_exp}")
        for i, pg_ in enumerate(rd.pages):
            w, h = float(pg_.mediabox.width), float(pg_.mediabox.height)
            if abs(w - 612) > 1 or abs(h - 792) > 1:
                errs.append(f"{os.path.basename(path)} p{i + 1}: {w:.0f}x{h:.0f}pt, not US Letter")
    if len(PdfReader(PREVIEW_PDF).pages) != 4:
        errs.append("preview PDF must be 4 pages")
    # rendered-layout scan: nothing clipped at the bottom of a page, no nowrap cell or line wider than its box
    for o in json.load(open(OVERFLOW_JSON)):
        if o["main_over"] > 1:
            errs.append(f"page {o['page']} ({o['key']}): content overflows by {o['main_over']}px (clipped)")
        for c in o["cells"]:
            errs.append(f"page {o['page']} ({o['key']}): text wider than its box: '{c}'")
    drd = PdfReader(DECK_PDF)
    if len(drd.pages) != DECK_SLIDES:
        errs.append(f"deck PDF: {len(drd.pages)} pages, expected {DECK_SLIDES}")
    for i, pg_ in enumerate(drd.pages):
        if abs(float(pg_.mediabox.width) / float(pg_.mediabox.height) - 16 / 9) > 0.01:
            errs.append(f"deck PDF p{i + 1}: not 16:9")
    for path in (PDF, PREVIEW_PDF, DECK_PDF):
        fonts = pdf_fonts(path)
        for name, emb in fonts:
            if emb != "yes":
                errs.append(f"{os.path.basename(path)}: font {name} not embedded")
            if not re.search(r"oswald|sourcesans", name.lower().replace(" ", "")):
                errs.append(f"{os.path.basename(path)}: unexpected font {name} (fallback?)")
        names = " ".join(n for n, _ in fonts).lower()
        for fam in ("oswald", "sourcesans3"):
            if fam not in names:
                errs.append(f"{os.path.basename(path)}: {fam} missing")
    prs, stext, notes, media = deck_texts()
    if len(prs.slides) != DECK_SLIDES:
        errs.append(f"deck has {len(prs.slides)} slides, expected {DECK_SLIDES}")
    for i, nt in enumerate(notes):
        if len(nt.strip()) < 80:
            errs.append(f"deck slide {i + 1}: speaker notes missing or too short")
        if not re.search(rf"\bDay {slide_day(i + 1)}\b", stext[i], re.I):
            errs.append(f"deck slide {i + 1}: no 'Day {slide_day(i + 1)}' slot shown on the slide")
    errs += deck_overflow(prs)
    if any(not m.lower().endswith(".png") for m in media) or len(media) != 3:
        errs.append(f"deck media should be exactly the 3 drawn PNGs (hero, keys, chordkeys), found {media}")
    if abs(prs.slide_width / prs.slide_height - 16 / 9) > 0.01:
        errs.append("deck is not 16:9")
    alldeck = " ".join(stext + notes)
    for t, a, *_ in SONGS:
        if t.lower() not in alldeck.lower() or a.lower() not in alldeck.lower():
            errs.append(f"deck must name {t} by {a}")
    for i, (_, k, _, body, _) in enumerate(SLIDES):
        if k == "listen" and any(t in body for t, *_ in SONGS) and "YouTube or Spotify" not in stext[i]:
            errs.append(f"deck slide {i + 1}: a song slide must say YouTube or Spotify")
    if "YouTube or Spotify" not in notes[0]:
        errs.append("deck slide 1 notes must say YouTube or Spotify")
    names = zipfile.ZipFile(ZIP).namelist()
    for p in (PDF, DECK, DECK_PDF):
        n = f"{ZIP_DIR}/{os.path.basename(p)}"
        if n not in names:
            errs.append(f"zip is missing {os.path.basename(p)}")
        elif zipfile.ZipFile(ZIP).read(n) != open(p, "rb").read():
            errs.append(f"zip copy of {os.path.basename(p)} is stale; rebuild")
    if len(names) != 3:
        errs.append(f"zip should hold exactly 3 files, has {names}")
    rd = PdfReader(PDF)
    pages = [subprocess.run(["pdftotext", "-layout", "-f", str(i + 1), "-l", str(i + 1), PDF, "-"], capture_output=True, text=True,
                            check=True).stdout for i in range(len(rd.pages))]
    flat = [re.sub(r"\s+", " ", t) for t in pages]
    raw0 = subprocess.run(["pdftotext", "-f", "1", "-l", "1", PDF, "-"], capture_output=True, text=True, check=True).stdout
    lines0 = [re.sub(r"\s+", " ", ln.strip()) for ln in raw0.splitlines() if ln.strip()]
    for need in (COVER_A.upper(), COVER_B.upper(), COVER_C):
        if need not in lines0:
            errs.append(f"PDF cover: '{need}' is not on a line of its own (orphaned word?)")
    for i, t in enumerate(flat):
        if i > 0 and "Brian Lotze · Hudson Beat" not in t:
            errs.append(f"page {i + 1}: footer missing")
        if i > 0 and not re.search(rf"Page {i + 1}\s*$", t.strip()):
            errs.append(f"page {i + 1}: footer page number wrong (overflow?)")
    for m in re.finditer(r"\b(?:pages?|pp?\.) (\d+)(?:–(\d+))?", " ".join(flat[1:])):
        for g in m.groups():
            if g and not 1 <= int(g) <= len(PLAN):
                errs.append(f"cross-reference to page {g} does not exist")
    for m in re.finditer(r"\bslides? (\d+)(?:–(\d+))?", " ".join(flat[1:])):
        for g in m.groups():
            if g and not 1 <= int(g) <= DECK_SLIDES:
                errs.append(f"PDF names deck slide {g}, which does not exist")
    rub = flat[IX["rubric"] - 1]
    if f"/ {RUBRIC_TOTAL}" not in rub or "/ 25" in rub:
        errs.append(f"rubric page: /{RUBRIC_TOTAL} total missing")
    for c, ds in RUBRIC:
        for d in ds:
            if not all(w in rub for w in re.findall(r"[A-Za-z]{6,}", d)):
                errs.append(f"rubric text missing: {d[:50]}")
    for w, _ in VOCAB:
        if w not in flat[IX["vocab"] - 1]:
            errs.append(f"vocab page missing {w}")
    for key, pk in (("C major", "cardC"), ("A minor", "cardA")):
        t = flat[IX[pk] - 1]
        for rn, typ, name, notes_, feel in CARD[key]:
            if notes_.replace(" ", "") not in t.replace(" ", "") or feel.split(".")[0] not in t:
                errs.append(f"{key} card page missing {rn} {notes_}")
    k2 = flat[IX["key2"] - 1]
    for key in ("C major", "A minor"):
        for rn, q, ns in key_chords(key):
            if not all(x in k2 for x in ns):
                errs.append(f"KCSRN key missing {key} {rn}")
    k1 = flat[IX["key1"] - 1]
    for need in ("F A C", "G B D", "D F# A", "D F A", "E G B", "B D F#", "B D F", "D F Ab"):
        if not re.search(r"\s".join(re.escape(x) for x in need.split()), k1):
            errs.append(f"answer key 1 missing chord {need}")
    for d in DAYS:
        if f"DAY {d['n']}" not in flat[IX["pacing"] - 1]:
            errs.append(f"pacing page missing Day {d['n']}")
    # Day 1 pages stay ear-only (no Roman numerals): the Day 1 post and the Feel Log
    for k in ("post1", "feellog"):
        if re.search(r"\b(?:IV|vi|ii°|vii°|I–IV)\b", flat[IX[k] - 1]):
            errs.append(f"{k}: Roman numerals on a Day 1 page")
    # Roman numerals are case-sensitive (vi is minor, VI is major): no progression may come out upper-cased by a CSS/caps heading
    progs = list(EX_PROG.values()) + [LOOP_MAJOR, LOOP_MINOR] + [x[3] for x in SONGS]
    pdf_all, deck_all = " ".join(flat), " ".join(stext)
    for p in progs:
        if [r.upper() for r in p] == p:
            continue
        for sep in ("–", "-", ", ", " – "):
            bad = sep.join(r.upper() for r in p)
            if bad in pdf_all or bad in deck_all:
                errs.append(f"progression {sep.join(p)} printed in capitals as {bad} (case changes the chord quality)")
    ptxt = [pg.extract_text() for pg in PdfReader(PREVIEW_PDF).pages]
    if sum("PREVIEW" in t for t in ptxt) != len(PREVIEW_PICKS):
        errs.append("preview PDF: watermark missing on some pages")
    blobs = {"resource PDF": pdf_text(PDF), "preview PDF": pdf_text(PREVIEW_PDF), "deck PDF": pdf_text(DECK_PDF), "deck pptx": alldeck,
             "zip names": " ".join(names)}
    with zipfile.ZipFile(DECK) as z:
        blobs["deck XML"] = " ".join(z.read(n).decode("utf8", "ignore") for n in z.namelist() if n.endswith(".xml") or n.endswith(".rels"))
    for path in (PDF, PREVIEW_PDF, DECK_PDF):
        meta = PdfReader(path).metadata or {}
        blobs[f"{os.path.basename(path)} metadata"] = " ".join(str(v) for v in meta.values())
    for fn in os.listdir(HERE):
        if fn == "QA.md":
            continue
        if fn.endswith((".md", ".json", ".py", ".txt")) or fn in ("release", "uploaded"):
            blobs[fn] = open(os.path.join(HERE, fn), encoding="utf8", errors="ignore").read()
    for label, txt in blobs.items():
        low = txt.lower()
        for b in BANNED:
            if b.lower() in low:
                errs.append(f"{label}: contains banned string '{b}'")
    for p in [COVER_PNG, *PREVIEW_PNGS]:
        if Image.open(p).size != (2000, 2000):
            errs.append(f"{os.path.basename(p)}: not 2000x2000 (low-res build?)")
    tj = json.load(open(TPT_JSON))
    if tj != tpt_json():
        errs.append("tpt.json is stale; rebuild")
    if tj["price"] != PRICE or tj["license_price"] != PRICE or tj.get("tax_code") != "Other Digital Goods - No Physical Media":
        errs.append("tpt.json price/license_price/tax_code")
    if tj["grades"] != ["9th-grade", "10th-grade", "11th-grade", "12th-grade"] or tj["subjects"] != ["Music"]:
        errs.append("tpt.json grades 9-12 / subject Music")
    proven_tags = {"Halloween", "Autumn", "Winter", "Thanksgiving"}
    if not (1 <= len(tj["tags"]) <= 6 and len(tj["formats"]) <= 3) or not set(tj["tags"]) <= proven_tags or tj["formats"] != ["PDF"]:
        errs.append("tpt.json tags/formats: only proven values (tags from 001-007, format PDF)")
    for k, f in tj["files"].items():
        if not os.path.exists(os.path.join(HERE, f)):
            errs.append(f"tpt.json file {k} missing: {f}")
    if tj["pages"] != len(rd.pages) + len(prs.slides):
        errs.append("tpt.json pages != PDF pages + slides")
    up_ = open(UPLOAD).read()
    m = re.search(r"^## 1\. Title\s*\n+```\n(.+?)\n```\n\((\d+) characters", up_, re.M | re.S)
    if not m:
        errs.append("UPLOAD.md: title block or its character count not found")
    else:
        n = len(m.group(1).strip())
        if n > 80:
            errs.append(f"UPLOAD.md: title is {n} chars (max 80)")
        if int(m.group(2)) != n:
            errs.append(f"UPLOAD.md: says {m.group(2)} characters, title has {n}")
        if m.group(1).strip() != LISTING_TITLE:
            errs.append("UPLOAD.md title differs from LISTING_TITLE in gen.py")
    m2 = re.search(r"^## 2\. Description\s*\n+```\n(.+?)\n```", up_, re.M | re.S)
    if not m2 or not m2.group(1).rstrip().endswith(DISCLOSURE):
        errs.append("UPLOAD.md: description must end with the disclosure line")
    else:
        desc = m2.group(1)
        if not desc.startswith("Chord progressions for high school music technology:"):
            errs.append("UPLOAD.md description: first line must lead with the search phrase (METADATA-GUIDE)")
        for need in (f"{len(PLAN)}-page", f"{DECK_SLIDES}-slide", "Logic Pro", f"/{RUBRIC_TOTAL}", "YouTube or Spotify", "GarageBand",
                     f"{DAYS_N}-day", "I-IV-V-I", "Roman numerals"):
            if need not in desc:
                errs.append(f"UPLOAD.md description: missing '{need}'")
        if re.search(r"  +\S|\s[,.]\s|\bteh\b|\s{2,}$", desc, re.M):
            errs.append("UPLOAD.md description: double space or stray punctuation")
        if "/25" in desc:
            errs.append("UPLOAD.md description: the owner's rubric is /15, not /25")
    for need in (f"${PRICE}", f"{len(PLAN)} pages", os.path.basename(ZIP), os.path.basename(DECK), os.path.basename(DECK_PDF)):
        if need not in up_:
            errs.append(f"UPLOAD.md: missing '{need}'")
    for f in ("release", "uploaded"):
        if os.path.exists(os.path.join(HERE, f)):
            errs.append(f"{f} exists: only the Manager / uploader writes it")
    return errs


def check():
    errs = verify_theory() + verify_data() + verify_labels() + verify_files()
    if errs:
        print("CHECK FAILED")
        print("\n".join(" - " + e for e in errs))
        sys.exit(1)
    print(f"CHECK OK: theory engine confirms C major / A minor scales (no sharps or flats, white keys only, relative keys), all 14 diatonic "
          f"triads, Roman numerals and spellings on the owner's cards, packet RN and reference-sheet tables, 10 packet chord answers, 8 "
          f"half/whole-step answers, both Day 1 loops, both worked examples (bass = chord root, root-position voicings) and the song facts; "
          f"{len(LABELS)} diagram labels fit; owner's lists and rubric 3 x (2-5) = /{RUBRIC_TOTAL} printed; {DAYS_N} days run 0-{MINUTES}; every "
          f"one of {DECK_SLIDES} slides has exactly one day slot matching its Day badge; Day 1 ear-only; {len(PLAN)} Letter pages, no overflow, "
          f"footers and valid page/slide references; cover lines whole; preview 4 watermarked pages; deck 16:9, notes on every slide, text in "
          f"bounds, 3 drawn PNGs, PDF copy {DECK_SLIDES} pages; only Oswald/Source Sans 3, embedded; no banned strings; zip current; images "
          f"2000x2000; tpt.json proven values; UPLOAD.md title {len(LISTING_TITLE)} chars, SEO opening and disclosure.")


if __name__ == "__main__":
    if "--fallback" in sys.argv:
        deck_pdf(DECK, os.path.join(BUILD, "deck-fallback.pdf"), fonts=False)
        print("wrote build/deck-fallback.pdf")
    elif "--check" in sys.argv:
        check()
    else:
        build()
