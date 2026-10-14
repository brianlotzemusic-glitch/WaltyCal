#!/usr/bin/env python3
"""DMP1 Unit 4: Bass & Melody Writing (TpT 012): build the teacher resource, the intro deck and the listing files.

    python3 gen.py            build everything into this folder
    python3 gen.py --lowres   same, but square images at scale 1 and small page shots (self-check only; --check rejects them)
    python3 gen.py --check    verify the music theory, data, files, page/slide counts, day slots, speaker notes, fonts, label fit,
                              page overflow, banned strings, tpt.json and UPLOAD.md

The owner's own curriculum (Brian Lotze, Digital Music Production 1), reformatted in the Hudson Beat style: a 4-day bass and
melody unit in Logic Pro. The DMP1 Restructure defines Unit 4 "Bass & Melody Writing" (4 days, 1/2/1, October-November, mixing
concept: frequency separation): "a bassline and short melody written over a provided chord progression"; basslines are
monophonic and sit on the root; melody is short, repetitive and singable; four ways to write a melody.
The owner's Drive folder "3 - The Building Blocks" supplies the material: Part 2 Scales & Melody (project sheet and /15 rubric,
teacher notes, Day 1 post), Part 4 Basslines (project sheet and /15 rubric, teacher notes, Day 1 post), the Scary Stuff packet
(Part 2 scales, Part 4 bassline analysis) and the Key Reference Cards (C major, A minor). Unit 3 (Chords by Ear, TpT 011) used
Parts 1 and 3 and left Part 2's melody project and Part 4 to this unit. Students continue from their Unit 3 project (a bassline
plus four chords); students without one use the teacher's chords-and-bassline file (recipe on the setup page).
Owner text is kept word for word except: the LMS name ("your class page"), unit and day labels (Unit 4, day numbers), the
Day 1 post's "Parts 1 and 2 of the packet" (Part 1 is in Unit 3), the Part 4 post's "Download today's Logic file" (the same
file as the melody), and the Padlet lines (a class feedback board). All are disclosed on teacher notes page 1. Questions for
the owner (kept as written, not changed): the Restructure's "no rubric, just a check-in" against the two /15 rubrics; the
Part 4 post's "I to V feels resolved"; Billie Jean as an example of the bass rule; chord length in the Part 2 file.
New here (disclosed): the four-methods page, the frequency separation mini-lesson, the Logic file recipe, worked examples,
answer keys and the intro deck. Songs are named by title and artist only (find them on YouTube or Spotify). School names and
logos removed; no student work, names or grades; no Logic sessions, audio or video.

Every scale, chord, Roman numeral, key signature, melody note and bass note in the kit is computed or checked by the small
theory engine below (verify_theory()).

Outputs: <slug>.pdf (teacher + student resource), <slug>-intro-deck.pptx and its PDF copy, a ZIP of those three
(the TpT product file), a 4-page watermarked preview PDF, cover.png and preview-1..3.png (2000 x 2000).
Renders with Playwright + Chromium (../tools/render.js) and LibreOffice (deck PDF), fonts from ../fonts. $0 AI.
"""
import json
import os
import re
import shutil
import subprocess
import sys
import zipfile
from html import escape

HERE = os.path.dirname(os.path.abspath(__file__))
TPT = os.path.dirname(HERE)
FONTS = os.path.join(TPT, "fonts")
BUILD = os.path.join(HERE, "build")
RENDER_JS = os.path.join(TPT, "tools", "render.js")

SLUG = "bass-melody-writing"
PDF = os.path.join(HERE, f"{SLUG}.pdf")
DECK = os.path.join(HERE, f"{SLUG}-intro-deck.pptx")
DECK_PDF = os.path.join(HERE, f"{SLUG}-intro-deck.pdf")
ZIP = os.path.join(HERE, f"{SLUG}.zip")
PREVIEW_PDF = os.path.join(HERE, f"{SLUG}-PREVIEW.pdf")
COVER_PNG = os.path.join(HERE, "cover.png")
PREVIEW_PNGS = [os.path.join(HERE, f"preview-{i}.png") for i in (1, 2, 3)]
UPLOAD = os.path.join(HERE, "UPLOAD.md")
TPT_JSON = os.path.join(HERE, "tpt.json")
OVERFLOW_JSON = os.path.join(BUILD, "overflow.json")
ZIP_DIR = "Bass and Melody Writing - Hudson Beat"

STORE = "Brian Lotze"      # copyright holder and TpT seller
BRAND = "Hudson Beat"      # store line on covers and previews (owner, 4 Oct)
YEAR = 2026
ACCENT = "#2159C9"         # this product's single accent: blue (distinct from 008 red, 009 orange, 010 violet, 011 teal)
ACC_LIGHT = "#C9D8F5"      # light tint of the accent, for text on black in the deck
TITLE = "Bass & Melody Writing"
TITLE_H = escape(TITLE)
PRICE = "10.00"
UNIT = 4                   # DMP1 Restructure numbering (the Building Blocks docs say Unit 3 for Parts 1-4)
DAYS_N = 4
MINUTES = 53
LOW_RES = "--lowres" in sys.argv
DISCLOSURE = "Designed with the help of digital and AI tools, and checked by hand."
LISTING_TITLE = "Melody Writing in Logic Pro | Basslines & Scales | High School Music Technology"
PROJECT2 = "Your Name — Building Blocks Part 2"   # owner's file-naming rules, kept as written (flagged to the owner in 011)
PROJECT4 = "Your Name — Building Blocks Part 4"

# Strings that must never appear in any output or in any file in this folder (school branding, LMS). Split so this file never contains them.
BANNED = ["Blind" + " Brook", "BB" + "HS", "blind" + "brook", "Google" + " Classroom"]

# ================================================================ music theory engine (letter-based spelling)
LETTERS = "CDEFGAB"
NAT = dict(C=0, D=2, E=4, F=5, G=7, A=9, B=11)
STEP = {"H": 1, "W": 2, "WH": 3}


def pc(n):
    return (NAT[n[0]] + n.count("#") - n.count("b")) % 12


def name_for(letter, target):
    d = (target - NAT[letter]) % 12
    if d > 6:
        d -= 12
    if abs(d) > 2:
        raise ValueError(f"cannot spell pitch class {target} on {letter}")
    return letter + ("#" * d if d > 0 else "b" * -d)


def up(note, semis, letters):
    """The note `semis` half steps above `note`, spelled on the letter `letters` steps higher."""
    return name_for(LETTERS[(LETTERS.index(note[0]) + letters) % 7], (pc(note) + semis) % 12)


def steps_of(formula):
    return [s.strip() for s in formula.split("-")]


def scale(tonic, formula):
    """Notes of a scale from a W/H/WH formula (WH skips a letter, as in a pentatonic), last note = the octave."""
    out = [tonic]
    for s in steps_of(formula):
        out.append(up(out[-1], STEP[s], 2 if s == "WH" else 1))
    return out


def chord(root, formula):
    """Stack thirds from a root with the owner's half-step formula, e.g. 'R-4-3' -> major."""
    out = [root]
    for n in formula.split("-")[1:]:
        out.append(up(out[-1], int(n), 2))
    return out


QUALITY = {(4, 3): "Major", (3, 4): "Minor", (3, 3): "Dim.", (4, 4): "Aug."}
FORMULA = {"Major": "R-4-3", "Minor": "R-3-4", "Dim.": "R-3-3", "Aug.": "R-4-4"}
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII"]


def quality(notes):
    return QUALITY[tuple((pc(notes[i + 1]) - pc(notes[i])) % 12 for i in range(2))]


def roman(degree, q):
    r = ROMAN[degree]
    return (r if q in ("Major", "Aug.") else r.lower()) + ("°" if q == "Dim." else "")


def diatonic(sc):
    """The 7 triads of a 7-note scale: (Roman numeral, quality, notes), stacking every other scale note."""
    s7 = sc[:7]
    out = []
    for i in range(7):
        notes = [s7[i], s7[(i + 2) % 7], s7[(i + 4) % 7]]
        q = quality(notes)
        out.append((roman(i, q), q, notes))
    return out


def chord_name(notes):
    """The owner's card style: 'C Major', 'A Minor', 'B Dim.'."""
    return f"{notes[0]} {quality(notes)}"


def short_name(notes):
    return notes[0] + {"Major": "", "Minor": "m", "Dim.": "dim", "Aug.": "aug"}[quality(notes)]


def dash(notes):
    return " - ".join(notes)


def accidentals(notes):
    return [n for n in notes if len(n) > 1]


def semis_between(a, b):
    return (pc(b) - pc(a)) % 12


def midi(n):
    """'C5' -> 72 (MIDI numbers, middle C = C4 = 60 here; the printed kit shows letter names only, never octave numbers)."""
    m = re.fullmatch(r"([A-G][#b]*)(-?\d)", n)
    return 12 * (int(m.group(2)) + 1) + pc(m.group(1))


def letter(n):
    return re.sub(r"-?\d$", "", n)


def midi_name(m):
    return ["C", "C#", "D", "D#", "E", "F", "F#", "G", "G#", "A", "A#", "B"][m % 12]


# ---------------------------------------------------------------- the owner's formulas and keys (Unit Overview, packet, reference cards)
MAJOR_SCALE = "W - W - H - W - W - W - H"
MINOR_SCALE = "W - H - W - W - H - W - W"
MAJOR_PENT = "W - W - WH - W - WH"
KEYS = {"C major": ("C", MAJOR_SCALE), "A minor": ("A", MINOR_SCALE)}


def key_scale(key):
    t, f = KEYS[key]
    return scale(t, f)


def key_chords(key):
    return diatonic(key_scale(key))


def by_rn(key, rn):
    return next(n for r, q, n in key_chords(key) if r == rn)


def tonic(key):
    return KEYS[key][0]


# owner's Reference Cards (Key reference Cards folder): RN, Type, Chord, Notes, Feels like (as printed in Unit 3, TpT 011)
CARD = {
    "C major": [("I", "Major", "C Major", "C - E - G", "Home. Settled. Complete."),
                ("ii", "Minor", "D Minor", "D - F - A", "Gentle tension. Moving somewhere."),
                ("iii", "Minor", "E Minor", "E - G - B", "Wistful. In between."),
                ("IV", "Major", "F Major", "F - A - C", "Lifting. Hopeful. Away from home."),
                ("V", "Major", "G Major", "G - B - D", "Homesick. Wants to resolve to I."),
                ("vi", "Minor", "A Minor", "A - C - E", "Tender. Melancholy. Emotional."),
                ("vii°", "Dim.", "B Dim.", "B - D - F", "Tense. Unstable. Rarely used alone.")],
    "A minor": [("i", "Minor", "A Minor", "A - C - E", "Home. But darker and more unsettled than major."),
                ("ii°", "Dim.", "B Dim.", "B - D - F", "Tense. Rarely used alone."),
                ("III", "Major", "C Major", "C - E - G", "Relative major. Bright moment in a darker key."),
                ("iv", "Minor", "D Minor", "D - F - A", "Heavy. Emotional. Pulling back toward home."),
                ("v", "Minor", "E Minor", "E - G - B", "Less forceful tension than major V."),
                ("VI", "Major", "F Major", "F - A - C", "Warm and open. Common resolution point."),
                ("VII", "Major", "G Major", "G - B - D", "Lifting. Drives back toward i strongly.")],
}
CARD_SCALE_TEXT = {"C major": "C Major  —  W - W - H - W - W - W - H", "A minor": "A Minor  —  W - H - W - W - H - W - W"}
CARD_REMEMBER = ("Remember: Roman numerals describe chord function, not just chord names.",
                 "The I chord always feels like home.  ·  The V chord always feels unresolved.  ·  The vi chord always feels tender or melancholy.",
                 "These feelings are the same in every key — only the note names change.")
# As in Unit 3 (QA round 1 there): the A minor card's reminder uses its own numerals (i, v / VII, VI), in the words of its own rows.
CARD_REMEMBER_MINOR = "The i chord always feels like home.  ·  The v and VII chords pull back toward i.  ·  The VI chord feels warm and open."

# ---------------------------------------------------------------- the Logic file (Unit 3 project or the teacher's file) and the worked examples
# The chords-and-bassline file: Unit 3's worked example progressions (TpT 011), one chord per measure, bass = whole-note roots.
FILE_PROG = {"C major": ["I", "vi", "IV", "V"], "A minor": ["i", "VI", "III", "VII"]}


def bass_root_midi(note):
    """Bass register of the Unit 3 file: roots C-E in octave 2, F-B in octave 1 (MIDI numbers)."""
    r = pc(note)
    return 36 + r if r <= 4 else 24 + r


# Worked melody (Hudson Beat, original): 4 measures of 4/4, one per chord of the file, the same rhythm in every measure
# (quarter, two eighths, half); starts on the root; up, down, up, down. (note, beats)
MELODY = {
    "C major": [[("C5", 1), ("D5", .5), ("E5", .5), ("G5", 2)], [("E5", 1), ("D5", .5), ("C5", .5), ("A4", 2)],
                [("C5", 1), ("D5", .5), ("E5", .5), ("F5", 2)], [("D5", 1), ("E5", .5), ("D5", .5), ("B4", 2)]],
    "A minor": [[("A4", 1), ("B4", .5), ("C5", .5), ("E5", 2)], [("C5", 1), ("B4", .5), ("A4", .5), ("F4", 2)],
                [("G4", 1), ("A4", .5), ("B4", .5), ("C5", 2)], [("B4", 1), ("C5", .5), ("B4", .5), ("G4", 2)]],
}
# Worked bassline (Part 4, Hudson Beat, original): a student-chosen progression from the card, one chord per measure.
# Rhythm idea A = dotted quarter, eighth, half; idea B = four quarters; the last quarter of measure 4 walks back to the loop's first root.
BASS_PROG = {"C major": ["I", "V", "vi", "IV"], "A minor": ["i", "VI", "VII", "i"]}
BASS = {
    "C major": [[("C2", 1.5), ("C2", .5), ("C2", 2)], [("G1", 1.5), ("G1", .5), ("G1", 2)],
                [("A1", 1), ("A1", 1), ("A1", 1), ("A1", 1)], [("F1", 1), ("F1", 1), ("F1", 1), ("G1", 1)]],
    "A minor": [[("A1", 1.5), ("A1", .5), ("A1", 2)], [("F1", 1.5), ("F1", .5), ("F1", 2)],
                [("G1", 1), ("G1", 1), ("G1", 1), ("G1", 1)], [("A1", 1), ("A1", 1), ("A1", 1), ("G1", 1)]],
}
BASS_IDEAS = ["Idea A", "Idea A", "Idea B", "Idea B + walk"]
BASS_WHY = {"C major": "I start at home, the V pulls away, the vi turns tender, and the IV lifts it so the loop wants to start again.",
            "A minor": "I start at a darker home, the VI opens it up, and the VII drives back to the i, so it always comes home."}
# Cover art: the file progression I-vi-IV-V with the worked melody and a rhythmic bassline (roots on every downbeat).
HERO_BASS = [[("C2", 1.5), ("C2", .5), ("C2", 2)], [("A1", 1.5), ("A1", .5), ("A1", 2)],
             [("F1", 1), ("F1", 1), ("F1", 1), ("F1", 1)], [("G1", 1), ("G1", 1), ("G1", 1), ("B1", 1)]]
# owner's Part 2 variation: "Restrict students to just 3 notes from the scale. C, E, G for C Major or A, C, E for A Minor."
THREE_NOTES = {"C major": ["C", "E", "G"], "A minor": ["A", "C", "E"]}


def flat(measures):
    """[(midi, start beat, beats)] across measures of 4/4."""
    out, t = [], 0.0
    for m in measures:
        for n, d in m:
            out.append((midi(n), t, d))
            t += d
    return out


# Songs (title and artist only; facts checked against the engine)
SNA = ("Seven Nation Army", "The White Stripes")
SNA_RIFF = ["E", "E", "G", "E", "D", "C", "B"]      # the riff a stadium sings: seven notes, five different ones, all in E natural minor
SBM = ("Stand by Me", "Ben E. King")
SBM_FACT = ("A major", "A", ["I", "vi", "IV", "V"], ["A", "F#m", "D", "E"])   # the bass riff moves to each chord's root
SONG_NOTE = "Songs are named by title and artist only. Play them from your own YouTube or Spotify access; no audio, video or lyrics are included."

# ================================================================ the owner's words (lightly edited: LMS, unit and day labels)
# ---- Part 2: Scales & Melody (project sheet)
OVERVIEW2 = ("You have learned the notes on the keyboard and how whole and half steps work. Now you are going to use a scale — the ingredient "
             "list for a key — to write a short melody in Logic. A chord progression and bassline are already provided in the Logic file. Your job "
             "is to add a melody on top that stays in key and is short enough to repeat.")
SUBMIT2 = ("A short melody written over the provided chord progression and bassline. The melody must be 2 or 4 measures long, stay within the "
           "notes of your chosen scale, and be entered into Logic using either the play-in or click-in method.")
REQ2 = ["Complete the scale pages in the Scary Stuff packet before opening Logic",
        "Choose C Major or A Minor as your scale — white keys only on your MPK Mini",
        "Download the provided Logic file — it contains a chord progression and bassline already written",
        "Write a melody that is 2 or 4 measures long",
        "All melody notes must come from the notes of your chosen scale",
        "Melody should start or end on the root note of the key (C for C Major, A for A Minor)",
        "Entered using play-in (record + quantize) or click-in (pencil tool)",
        "Melody on a separate track with a different sound from the bass",
        f"Project named: {PROJECT2}",
        "Exported as MP3 and .zip (compressed .logicx file) and submitted to your class page before presenting"]
HOWTO2 = [("Keep it short", "2 or 4 measures only. Once you have something you like, copy and paste it across the progression. Short melodies are "
           "easier to remember and easier to build on."),
          ("White keys only", "In C Major or A Minor, every white key is a safe note. You cannot play a wrong note if you stay on the white keys."),
          ("Give it shape", "A good melody goes up for a bit, then comes down — or starts high and descends. Avoid playing the same note over and "
           "over, or jumping randomly in both directions."),
          ("Start or end on the root", "Anchor your melody to C (for C Major) or A (for A Minor). This creates a sense of resolution."),
          ("Simple and repetitive wins", "A two-note melody with a strong rhythm is more powerful than eight random notes. Write something you "
           "could hum after hearing it once.")]
TIPS2 = ["If you cannot find a melody, set a 2-measure loop over the biggest part of the progression and just play around on the white keys. "
         "Record anything. You can delete it. The goal is to start.",
         "Quantize (Q) is not cheating. It fixes timing, not notes. Use it.",
         "If your melody feels aimless, try this: pick one note and play it in a rhythm you like. Then move one note up or down. Build from there."]
CHECK2 = ["Scale pages in the Scary Stuff packet completed", "Melody is 2 or 4 measures long",
          "All notes from C Major or A Minor scale (white keys only)", "Melody starts or ends on the root note of the key",
          "Entered using play-in or click-in — quantized if played in", "Melody on separate track with different sound from bass",
          f"Project named: {PROJECT2}", "Exported as MP3 and .zip and submitted to your class page before presenting"]
RUBRIC2 = [  # the owner's Part 2 rubric, word for word: levels 2-5, 3 categories = /15
    ("Melody Writing",
     ["Melody does not stay in key or has no recognizable shape. Cannot be sung back.",
      "Melody stays mostly in key but lacks a clear phrase or contour.",
      "Melody is in key, 2 or 4 measures, and has a recognizable shape. Could be sung back.",
      "Melody is in key, short, starts or ends on the root, and has clear contour. Sounds like it belongs to the progression."]),
    ("Technical Skill",
     ["Notes entered incorrectly or piano roll not used. Play-in material not quantized.",
      "Notes mostly correct but some placement errors or timing issues. Inconsistent use of quantize.",
      "Piano roll used correctly. Notes placed with intention. Quantize applied after play-in.",
      "Piano roll used confidently. Notes are well-placed and well-timed. Both entry methods understood."]),
    ("Documentation & Presentation",
     ["Not submitted or not presented.",
      "Submitted late or named incorrectly. Minimal presentation engagement.",
      "Correctly named and submitted. Presentation made.",
      "Correctly named, on time. Student explains their melody choices clearly."]),
]
# ---- Part 4: Basslines (project sheet)
OVERVIEW4 = ("You have written a melody and built chords. Now you are going to discover the rule behind every bassline you have ever heard — and "
             "write one yourself. Using your Reference Card to choose a chord progression, you will write a bassline that follows the root of each "
             "chord and gives the progression its rhythmic foundation.")
REQ4 = ["Complete the bassline analysis section of the Scary Stuff packet before writing anything in Logic",
        "Choose a 4-chord progression using Roman numerals from your Reference Card",
        "Write the root note for each chord before opening Logic",
        "Write a bassline that is monophonic throughout — one note at a time, no overlapping notes",
        "Root note of each chord used when that chord is playing",
        "At least 2 distinct rhythmic ideas in the bassline — not the same note value throughout",
        "Progression clearly reflects a harmonic intention — be ready to explain why you chose these chords",
        f"Project named: {PROJECT4}",
        "Exported as MP3 and .zip (compressed .logicx file) and submitted to your class page before presenting"]
WHY4 = "Why did you choose this progression? What did you want it to feel like?"
TIPS4 = ["Get the root notes right first. Place the root note of each chord for the full duration of that chord before adding anything else.",
         "A bassline that is too busy is worse than a bassline that is too simple. If in doubt, simplify.",
         "The rhythm of your bassline is what makes it interesting. Try long notes for slow sections and shorter notes for busier sections.",
         "Check for overlapping notes in the piano roll. Any overlap in the bass track is an error."]
CHECK4 = ["Bassline analysis section of Scary Stuff packet completed", "Chord progression planned with Roman numerals and root notes",
          "Bassline monophonic throughout — no overlapping notes", "Root note used at each chord change",
          "At least 2 distinct rhythmic ideas in the bassline", "Written explanation of chord progression choices completed above",
          f"Project named: {PROJECT4}", "Exported as MP3 and .zip and submitted to your class page before presenting"]
RUBRIC4 = [  # the owner's Part 4 rubric, word for word: levels 2-5, 3 categories = /15
    ("Bassline Construction",
     ["Bassline not monophonic, or root notes absent or wrong. Does not follow chord progression.",
      "Bassline mostly monophonic with some correct root notes, but does not consistently follow chord changes.",
      "Bassline is monophonic. Root note used at each chord change. Some rhythmic variation present.",
      "Bassline is monophonic, follows every chord change, and has a clear rhythmic identity. Feels like it belongs."]),
    ("Harmonic Intention",
     ["Chord choices appear random. No written explanation or explanation shows no harmonic thinking.",
      "Progression uses correct chords but written explanation is vague. No clear emotional intention.",
      "Progression has a discernible feel. Written explanation describes what student intended.",
      "Progression has a clear emotional character. Student can explain their choices using Roman numeral language."]),
    ("Documentation & Presentation",
     ["Not submitted or not presented.",
      "Submitted late or named incorrectly. Minimal presentation engagement.",
      "Correctly named and submitted. Presentation made. Planning table completed.",
      "Correctly named, on time. Student presents chord choices using Roman numeral language confidently."]),
]
LEVELS = [2, 3, 4, 5]
RUBRIC_TOTAL = 15                    # each of the owner's two rubrics: 3 categories x 5 = /15
RUBRIC_SHORT = {2: ["In key, short, starts or ends on the root, clear contour.", "Confident piano roll; well placed, well timed; both entry methods.",
                    "Named correctly, on time; explains the melody choices clearly."],
                4: ["Monophonic, follows every chord change, clear rhythmic identity.", "Clear emotional character, explained in Roman numeral language.",
                    "Named correctly, on time; presents chord choices with Roman numerals."]}
RULE = "When the chord changes, the bass note changes to the root of the new chord."
RULE_SUB = "Everything else — rhythm, extra notes, movement between roots — is decoration on top of this one rule."
# ---- Part 4 packet: analysis questions (owner's wording)
ANALYSIS = ["What key is this bassline in?", "What chord is playing in the first section?", "What note does the bassline play during that chord?",
            "Is that note the root of the chord?", "What happens to the bass note when the chord changes?"]

# ---- the owner's teacher notes (Part 2, Part 4) and course plan (Restructure, Unit 4)
WHAT2 = ("Part 2 makes the first musical demand: write something. The scale gives students a safe set of notes, the Logic file removes the "
         "blank-page problem, and the constraint of 2 or 4 measures prevents the endless-tinkering failure mode. By the end students should have a "
         "short, singable melody and a physical understanding of what a scale is for.")
KEY2 = ("The key insight to deliver: a scale is not an exercise. It is the ingredient list for a key. Every note in a melody should come from "
        "the scale. If it does, it will sound like it belongs. That is the only rule.")
FILE2 = ["C major or A minor key — white keys only for accessibility", "4 chords, each lasting 4 measures (or 2 measures for a shorter feel)",
         "Bassline using root notes of each chord", "Arrangement markers labeling each section",
         "No melody track — leave that blank for students to fill in"]
WHAT4 = ["Part 4 has a reveal at its center. Students have been given basslines in every Logic file so far. Now they are asked to look carefully "
         "at one and figure out the rule themselves before it is stated. The analysis section of the packet is the discovery activity. The rule box "
         "in the packet confirms what they found. This sequence — discover, confirm, apply — is more memorable than being told the rule directly.",
         "The second half of Part 4 connects everything. Students choose a chord progression using Roman numerals from their Reference Card and "
         "write a bassline that follows those roots. For the first time, they are making a compositional decision — not just executing an "
         "exercise. The written explanation of why they chose their progression is the assessment moment that shows you whether the Roman "
         "numeral instruction actually worked."]
REVEAL = ["After students have completed the analysis questions in the packet and identified that the bass note always follows the chord root, "
          "bring the class back together and ask: what did you find? Let two or three students answer. Then write the rule on the board:",
          "Then say: every bassline you have ever heard follows this rule. Another One Bites the Dust. Billie Jean. Your favorite song. The rule "
          "is always the same. Everything else is rhythm and decoration on top of this one idea."]
PRESQ = ["Every student should be asked this during their presentation: why did you choose this chord progression? What did you want it to "
         "feel like?",
         "A student who answers with Roman numeral language — ‘I wanted it to feel unresolved so I ended on the V’ or ‘I used the vi because I "
         "wanted it to feel more emotional’ — has internalized the unit. A student who says ‘it sounded good’ has not. Both are valid at this "
         "stage, but the difference tells you what to reinforce in Unit 5 and beyond."]
STRUGGLES = [
    ("Blank page paralysis", "Enforce the loop-and-play rule: set a 2-measure loop over the biggest part of the progression and just play. Record "
     "anything. Students who resist starting are almost always overthinking. The prompt: play one note in a rhythm you like. Then move one note "
     "up or down."),
    ("Melodies that go nowhere", "Students who try to use all 7 scale notes end up with aimless lines. Redirect them: use only 3 notes. Go up, "
     "then come back down. That is a melody."),
    ("Not quantizing", "Play-in melodies with bad timing are discouraging. Show the Q shortcut immediately after the first recording attempt. "
     "Make it automatic."),
    ("Wrong octave", "Students who do not shift the MPK Mini up an octave will record a melody that sits in the bass register. Show the OCT "
     "button on the keyboard before recording begins."),
    ("Not changing the bass note at chord changes", "Walk around during working time and ask every student: point to where your chord changes. "
     "Does your bass note change there? This is the most common error and the most important one to catch early."),
    ("Polyphonic basslines", "Students leave previous notes sustaining when new ones start. Zoom into the piano roll and look for any notes that "
     "overlap. Any overlap in the bass is an error."),
    ("Chord progressions chosen randomly", "Students who pick chords without thinking about function will have nothing to say during the "
     "presentation. Push back before they start building: tell me what you want this to feel like. Then ask: which Roman numerals would create "
     "that feeling?"),
    ("Rhythmic basslines that are too busy", "A bassline that plays sixteenth notes throughout is exhausting. Redirect with a simple rule: your "
     "bassline should breathe. Give some notes room. Less is almost always more.")]
VAR_EASY = ("Restrict students to just 3 notes from the scale. C, E, G for C Major or A, C, E for A Minor. A melody built from only 3 notes is "
            "still a real melody and removes the decision paralysis entirely.")
VAR_HARD = ("Ask students to write a second melody that acts as a response to the first — call and response. Or ask them to change the last "
            "measure of the melody so it ends differently on the final repetition.")
PLAN_LEARN = ("Basslines are monophonic — one note at a time. The root of each chord is the safest and most powerful bass note. Melody is short, "
              "repetitive, and singable. Quantize exists and it is not cheating. The scale gives you a guaranteed set of notes that will sound "
              "right. Four methods of writing melody: play in, draw a rhythm first, click one note and move it, use the step sequencer.")
PLAN_WHY = ("Unit 3 gave students the harmonic landscape. This unit teaches them how to move through it — one line at a time. Bass anchors the "
            "bottom, melody defines the top. These are the two skills the EDM project requires but currently cannot assume. Teaching them here, "
            "with no other pressures, means students can focus entirely on the musical challenge.")
PLAN_CHANGE = ("This is a new standalone unit carved from content currently buried inside EDM project videos. Students practice bassline and "
               "melody writing before they need it in a high-stakes project. The 'start on the root of each chord' rule is introduced explicitly "
               "and repeated.")
PLAN_MIX = ("Frequency separation is introduced here because students immediately encounter the problem — their bassline and melody sound "
            "muddy when stacked. A brief introduction to EQ: high-pass the melody to remove low frequencies it does not need, roll off the "
            "extreme highs of the bass. Students do not need to understand parametric EQ yet. They need to understand that instruments have "
            "frequency homes and mixing means giving each one space to live.")
SIMPLE = "Simple and repetitive will always beat complicated and forgettable."
BASS_LINE = "Your bassline loves to sit on the root of the chord. That is the safest and most powerful note you can play."
METHODS = [  # the four methods named in the owner's course plan; the Logic steps are written for this kit
    ("Play it in", "Loop 2 measures with Cycle (C). Press R, play on the white keys, stop. Select the notes and press Q to quantize.",
     "you can already hear an idea in your head"),
    ("Draw a rhythm first", "With the Pencil, click notes in a rhythm you like, all on the root. Play it back. Then drag some notes up or down to "
     "other scale notes.", "you have a groove but no tune yet"),
    ("Click one note, move it", "Click in one note. Listen. Drag it up or down until it fits the chord. Then click the next one.",
     "you want to build slowly, one decision at a time"),
    ("Use the step sequencer", "Control-click the track lane and choose Create Pattern Region. Turn steps on to place notes; each row is one pitch.",
     "you want a short, repeating, rhythmic hook")]
VOCAB = [("Scale", "an ordered set of notes built with a formula of whole and half steps. The ingredient list for a key."),
         ("Key", "the home a piece of music settles on: its home note, its scale and its chords (this unit: C major or A minor)."),
         ("Root", "the note a chord is built on, and the home note of a key (C in C major, A in A minor)."),
         ("Melody", "the tune: one note at a time, on top. Short, repetitive and singable."),
         ("Contour", "the shape of a melody: up for a bit, then down, or high and then descending."),
         ("Play-in", "recording a part by playing it on a MIDI keyboard (R), then quantizing it (Q)."),
         ("Click-in", "entering notes one at a time with the Pencil tool in the piano roll."),
         ("Quantize", "snapping notes to the grid. It fixes timing, not notes. It is not cheating."),
         ("Bassline", "the lowest part. It follows the roots of the chords and gives the music its rhythmic foundation."),
         ("Monophonic", "one note at a time, with no overlapping notes. Every bassline in this unit is monophonic."),
         ("Chord progression", "the order the chords come in, written as Roman numerals (I–V–vi–IV) or chord names."),
         ("Frequency separation", "giving each part its own frequency home so bass and melody do not turn to mud. A high-pass filter on the "
          "melody is the first tool.")]
VOCAB_QUIZ = [("One note at a time, no overlaps: the bassline is ______.", "monophonic"),
              ("______ fixes timing, not notes.", "quantize"),
              ("The note a chord is built on is its ______.", "root"),
              ("The shape of a melody, up and then down, is its ______.", "contour")]

# ================================================================ pacing: 4 days of about 53 minutes (DMP1 Restructure, Unit 4: 4 days, 1/2/1)
# block = (minutes, heading, text, deck slides (first, last) or None). Every deck slide belongs to exactly one block (checked).
DAYS = [
    {"n": 1, "kind": "Intro day", "title": "Scales + melody start",
     "goal": "Every student has checked or built the C major and A natural minor scales, chosen a key, and started a 2- or 4-measure melody "
             "over the chords and bassline.",
     "blocks": [("0–8", "The hook", f"Slides 1–3. What is a melody? Play the first 30 seconds of {SNA[0]} ({SNA[1]}): how many different "
                 "notes, how many repeats, could you hum it after one listen?", (1, 3)),
                ("8–12", "Bright or dark?", "Slide 4. Play the C major scale, then the A natural minor scale, for the class. Ask: which sounds "
                 "brighter? Which sounds darker? Do not explain yet.", (4, 4)),
                ("12–25", "Scales", "Slide 5. Students work through Part 2 of the packet (page @PK2): C major, A natural minor and C major "
                 "pentatonic. Students who did it in Unit 3 check it against the Reference Card (pages @CARDS) and play each scale.", (5, 5)),
                ("25–32", "Two projects, four ways", "Slides 6–7. The four days, the melody project sheet (pages @PROJ2) and the four ways to "
                 "write a melody (page @METH).", (6, 7)),
                ("32–38", "Watch me", "Slides 8–9: the worked example and a 2-minute demo. Shift the keyboard up an octave, loop 2 measures, "
                 "record, quantize, copy across the progression.", (8, 9)),
                ("38–50", "Write", "Slide 10 stays up. Students open the chords-and-bassline file (their Unit 3 project saved as Part 2, or "
                 "your file), listen twice, and begin the melody. Enforce the 2-or-4-measure rule immediately.", (10, 10)),
                ("50–53", "Save + hum it", "Students save. Exit check: hum your melody to a neighbor.", None)]},
    {"n": 2, "kind": "Work day", "title": "Finish the melody, meet the bass rule",
     "goal": "Melody projects submitted. Every student has looked hard at a bassline and found the rule before you say it.",
     "blocks": [("0–5", "Does it belong?", "Ask two students to play their melody alongside the Logic file. Class listens. Does it sound like "
                 "it belongs?", None),
                ("5–10", "What you turn in", "Slides 11–12: the melody checklist and rubric (pages @PROJ2, @RUB2).", (11, 12)),
                ("10–17", "Mixing moment: frequency separation", "Slide 13. Solo the bass, solo the melody, then both: muddy? Students "
                 "put Channel EQ on the melody and turn on the low cut (high-pass); on the bass, the high cut.", (13, 13)),
                ("17–35", "Finish + submit the melody", "Students complete the melody, quantize, export and submit. Early finishers add a "
                 "second melody variation (page @TN4).", None),
                ("35–47", "Basslines: look before you write", f"Slide 14: {SBM[0]} ({SBM[1]}), follow the bass. Then students open the "
                 "piano roll of the bassline in the same file and answer the Part 4 analysis questions (page @PK4). Do not reveal the rule "
                 "yet; circulate and ask what they are finding.", (14, 14)),
                ("47–53", "The reveal", "Slide 15. What did you find? Let two or three students answer. Confirm the rule and write it on the "
                 "board (page @TN2).", (15, 15))]},
    {"n": 3, "kind": "Work day", "title": "Plan + build the bassline",
     "goal": "Progression planned with Roman numerals and root notes; chords and a monophonic bassline built in Logic and submitted.",
     "blocks": [("0–12", "Plan your progression", "Slide 16. Hand out the bassline project sheet (pages @PROJ4). Students fill in the "
                 "planning table and the why question before opening Logic. Circulate and ask: why these chords? What feeling are you going "
                 "for?", (16, 16)),
                ("12–18", "Watch me: roots first", "Slides 17–18: the worked example, then a demo: whole-note roots first, then rhythm.",
                 (17, 18)),
                ("18–45", "Build", "Full working time. New Logic file: build the four chords (Reference Card spellings), then the bassline. "
                 "Key check: play the bassline alone. Does it clearly follow the chord roots? Then play it with the chords. Does everything "
                 "agree?", None),
                ("45–50", "What you turn in", "Slides 19–20: the bassline checklist and rubric (pages @PROJ4, @RUB4).", (19, 20)),
                ("50–53", "Export + submit", "Students export and submit before the presentation day.", None)]},
    {"n": 4, "kind": "Present", "title": "Presentations",
     "goal": "Every student presents; every student can say what key, what Roman numerals, and what they wanted it to feel like.",
     "blocks": [("0–5", "Last checks", "Stragglers export and submit. Students post one question about their project on the class feedback "
                 "board.", None),
                ("5–48", "Presentations", "Slide 21. Each student plays their melody, then their bassline over their chord progression, "
                 "and answers: what key, what Roman numerals, what did you want it to feel like? Classmates respond to at least 2 projects on "
                 "the feedback board.", (21, 21)),
                ("48–53", "Close", "Slide 22: the full picture. In this unit you wrote a melody, built chords, and wrote a bassline.", (22, 22))]},
]
DECK_DAY_LINES = ["A melody is short, repetitive and singable; scales; four ways to write it; start your melody",
                  "Finish the melody; frequency separation; look hard at a bassline and find the rule",
                  "Plan four chords with Roman numerals; roots first, then rhythm; build the bassline",
                  "Play it, explain it: what key, what Roman numerals, what did you want it to feel like?"]

# ================================================================ intro deck (the owner's "Looping in Logic" pattern); (day, kind, title, body, notes)
SLIDES = [
    (1, "title", "Bass & Melody Writing", f"Feat. {SNA[0]}",
     f"Day 1, the first minute. Have {SNA[0]} ({SNA[1]}) and, for Day 2, {SBM[0]} ({SBM[1]}) cued up on YouTube or Spotify from your own "
     "account; no audio or video is included in this deck. Open Logic with the chords-and-bassline file on the class speakers. Open with a "
     "question: what is the part of a song you hum in the shower?"),
    (1, "bullets", "What the heck is a melody?", "The tune you hum after the song ends|One note at a time, on top of the chords|"
     "Short, repetitive and singable|Simple and repetitive wins",
     f"Plain-English definition, kept short. Say the line from my course plan now and keep saying it all week: {SIMPLE} The bass is the other "
     "half of this unit: it anchors the bottom while the melody defines the top."),
    (1, "listen", "Listen: a melody a stadium can sing", f"{SNA[0]} · {SNA[1]}|Teacher: play the first 30 seconds from YouTube or Spotify.|"
     "How many different notes do you hear?|How many times does it repeat?|Could you hum it after one listen?",
     f"About five minutes. The famous riff of {SNA[0]} is seven notes long and uses only five different notes, all from one scale, and whole "
     "stadiums sing it back. Students usually guess far more notes. Point out: short, repetitive, singable. That is the target for their own melody."),
    (1, "listen", "Bright or dark?", "C major scale  ›  A natural minor scale · teacher plays on a piano sound|Same seven white keys, two starting notes.|"
     "Which one sounds brighter?|Which one sounds darker?|What changed?",
     "From my Part 2 teacher notes: play the C major scale and the A natural minor scale for the class. Ask which sounds brighter and which "
     "darker. Do not explain yet. Most students hear C major as bright and A minor as darker, even though both use only white keys: the home "
     "note changed."),
    (1, "scales", "A scale is your ingredient list", "",
     "The key insight from my teacher notes: a scale is not an exercise. It is the ingredient list for a key. Every note in a melody should come "
     "from the scale; if it does, it will sound like it belongs. Students who built these scales in Unit 3 check them against the Reference Card "
     "and play them; others complete Part 2 of the packet now."),
    (1, "days", "Four days, two short projects", "",
     "The shape of the unit. Days 1 and 2: a short melody over the chords and bassline you already have. End of Day 2: look hard at a bassline "
     "and discover its rule. Day 3: choose your own four chords and write a bassline that follows them. Day 4: present both. Hand out the "
     "melody project sheet now."),
    (1, "methods", "Four ways to write a melody", "",
     "The four methods from my course plan. None is more correct than another; students should try at least two today. Play it in suits "
     "students who can already hear an idea; drawing a rhythm first and clicking one note at a time suit students who freeze at the keyboard; "
     "the step sequencer suits short rhythmic hooks. Quantize is not cheating: it fixes timing, not notes."),
    (1, "pic", "Worked example: a melody over the chords", "melroll",
     "One possible answer in a C major file, over the chords and bassline from Unit 3: one measure per chord, the same rhythm in every measure, "
     "starting on the root C, going up then down. Every note is a white key from the C major scale, and each measure starts on a note of its "
     "chord, so it sounds like it belongs. The A minor version is in the kit."),
    (1, "steps", "Watch me: record, quantize, copy", "Open the chords-and-bass file; save it as Part 2|Shift the keyboard up an octave (OCT +)|"
     "New track, a lead or synth sound|Loop 2 measures, press R, play; then Q|Copy and paste it across the progression",
     f"A two-minute demo, then students work. Unit 3 students open their own chords project and Save As {PROJECT2}; others open your file. "
     "Show the OCT button before anyone records, or the melody lands in the bass register. Record, select the notes, press Q. Then copy the "
     "region across the progression."),
    (1, "bullets", "How to write a melody", "Keep it short: 2 or 4 measures|White keys only: no wrong notes|Give it shape: up, then down|"
     "Start or end on the root|Simple and repetitive wins",
     "The five rules from my project sheet, on screen while students write. If a student cannot start: set a 2-measure loop and just play "
     "around on the white keys; record anything. If a melody feels aimless: pick one note, play it in a rhythm you like, then move one note up "
     "or down."),
    (2, "check", "Melody: what you turn in", "Scale pages of the packet done|2 or 4 measures, in your scale|Starts or ends on the root|"
     "Played in and quantized, or clicked in|Own track, a different sound from the bass|MP3 + .zip, named Your Name — Building Blocks Part 2",
     "Day 2 opens with two students playing their melody alongside the file: does it sound like it belongs? Then the submission checklist. "
     "Show where students submit (your class page). Export: File › Bounce › Project or Section, MP3. Compress: in Finder, Control-click the "
     ".logicx project and choose Compress."),
    (2, "rubric", "Melody: how it’s scored", "2",
     f"Three categories, each scored 2 to 5, for a total out of {RUBRIC_TOTAL}. Read the 5 column aloud: in key, short, starts or ends on the "
     "root, with a clear contour that sounds like it belongs; the piano roll used confidently with both entry methods understood; on time, "
     "named correctly, and the melody choices explained clearly."),
    (2, "freq", "Mixing moment: frequency separation", "",
     "The mixing concept for this unit, about seven minutes. Solo the bass, solo the melody, then play both: does it sound muddy? Instruments "
     "have frequency homes, and mixing means giving each one space to live. Students open Channel EQ on the melody and turn on the low cut "
     "(high-pass); on the bass they turn on the high cut. They do not need parametric EQ yet."),
    (2, "listen", "Listen: follow the bass", f"{SBM[0]} · {SBM[1]}|Teacher: play the first minute from YouTube or Spotify.|"
     "Hum the bass, not the singer.|Where does the bass pattern start again?|What happens to it when the chords change?",
     f"About three minutes, then the packet. The bass riff in {SBM[0]} plays the same pattern four times, starting each time on the root of "
     "the new chord: I, vi, IV, V in A major. Do not say that yet. Let students notice that the riff moves when the harmony moves, then send "
     "them to the bassline in their own file and the Part 4 analysis questions."),
    (2, "rule", "The one rule for basslines", "",
     "The reveal, from my Part 4 teacher notes. After students have answered the analysis questions, ask: what did you find? Let two or three "
     "students answer. Then write the rule on the board and say: every bassline you have ever heard follows this rule. Everything else is "
     "rhythm and decoration on top of this one idea."),
    (3, "plan", "Plan your progression", "",
     "Day 3. Students choose four chords from their Reference Card and write the Roman numeral and root note for each before opening Logic, "
     "then answer the why question. Push back before they build: tell me what you want this to feel like. Which Roman numerals would create "
     "that feeling? The example shown is one possible plan."),
    (3, "pic", "Worked example: a bassline", "bassroll",
     "One possible answer for the plan on the last slide, I–V–vi–IV in C major. Every measure starts on the chord's root. Measures 1 and 2 use "
     "one rhythm idea, measures 3 and 4 a second one, and the last note walks back toward C. One note at a time: no overlaps. It breathes: "
     "nothing busier than eighth notes."),
    (3, "steps", "Watch me: roots first, then rhythm", "New file: build your four chords from the card|Bass track: one whole-note root per chord|"
     "Play the bass alone, then with the chords|Add rhythm: at least two ideas|Zoom in: no overlapping notes",
     "A short demo. Get the roots right first: place the root of each chord for the full duration of that chord before adding anything else. "
     "Then add rhythm. A bassline that is too busy is worse than one that is too simple. Any overlap in the bass track is an error."),
    (3, "check", "Bassline: what you turn in", "Part 4 analysis in the packet done|Planning table: Roman numerals + roots|"
     "Monophonic: no overlapping notes|Root at every chord change|At least 2 rhythmic ideas + your why|MP3 + .zip, named Your Name — Building Blocks Part 4",
     "The Part 4 submission checklist, condensed. The written explanation of why students chose their progression is on the project sheet "
     "under the planning table. Students export and submit before the presentation day, the same way as the melody."),
    (3, "rubric", "Bassline: how it’s scored", "4",
     f"Three categories, each scored 2 to 5, for a total out of {RUBRIC_TOTAL}. Read the 5 column aloud: monophonic, following every chord "
     "change with a clear rhythmic identity; a progression with a clear emotional character explained in Roman numeral language; on time, "
     "named correctly, and chord choices presented with Roman numerals."),
    (4, "bullets", "Presentation day", "Play your melody, then your bassline|What key? What Roman numerals?|What did you want it to feel like?|"
     "Respond to at least 2 classmates",
     "Day 4. Each student plays their melody, then their bassline over their chords, and answers the question that matters: why did you choose "
     "this progression, what did you want it to feel like? A student who answers in Roman numeral language has internalized the unit; one who "
     "says it sounded good has not yet."),
    (4, "closer", "One last thing…", "You wrote a melody.|You built chords.|You wrote a bassline.|The foundation of every song.",
     "Close with the full picture from my teacher notes: in this unit you wrote a melody, built chords, and wrote a bassline. Those three things "
     "together are the foundation of every song ever written. Next unit, Song Form, arranges them into sections."),
]
DECK_SLIDES = len(SLIDES)
DECK_PICS = ["hero", "melroll", "bassroll"]

# ================================================================ html scaffolding (Hudson Beat, from 006-011)


def font_css(base=BUILD):
    f = os.path.relpath(FONTS, base)
    faces = [("Oswald", 300, "normal", "oswald-latin-300-normal"), ("Oswald", 500, "normal", "oswald-latin-500-normal"),
             ("Oswald", 700, "normal", "oswald-latin-700-normal"),
             ("Source Sans 3", 400, "normal", "source-sans-3-latin-400-normal"),
             ("Source Sans 3", 400, "italic", "source-sans-3-latin-400-italic"),
             ("Source Sans 3", 600, "normal", "source-sans-3-latin-600-normal"),
             ("Source Sans 3", 700, "normal", "source-sans-3-latin-700-normal")]
    return "".join(f"@font-face{{font-family:'{n}';font-weight:{w};font-style:{s};src:url('{f}/{file}.woff2') format('woff2');}}"
                   for n, w, s, file in faces)


CSS = """
@page { size: 8.5in 11in; margin: 0; }
* { box-sizing: border-box; }
html, body { margin: 0; padding: 0; background: #fff; color: #000; }
body { font-family: 'Source Sans 3', sans-serif; -webkit-print-color-adjust: exact; print-color-adjust: exact; }
.page { width: 8.5in; height: 11in; position: relative; overflow: hidden; page-break-after: always; break-after: page;
        padding: 0.42in 0.5in 0.38in; display: flex; flex-direction: column; }
.page:last-child { page-break-after: auto; break-after: auto; }
.l { font-family: 'Oswald'; font-weight: 300; }
.top { display: flex; justify-content: space-between; align-items: flex-end; border-bottom: 2.2pt solid #000; padding-bottom: 6pt; }
.top .t { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; letter-spacing: .01em; line-height: 1; font-size: 25pt; }
.top .t small { display: block; font-family: 'Oswald'; font-weight: 300; text-transform: none; font-size: 12pt; letter-spacing: 0; margin-top: 4pt; }
.tag { text-align: right; }
.tag .lv { display: inline-block; background: #000; color: #fff; font-family: 'Oswald'; font-weight: 700; font-size: 14pt; padding: 3pt 9pt 4pt; border-radius: 4pt; letter-spacing: .04em; }
.tag .lv.ak { background: ACCENT; }
.tag .ln { font-family: 'Oswald'; font-weight: 300; font-size: 11pt; margin-top: 3pt; }
.name { display: flex; gap: 18pt; font-size: 12pt; margin: 9pt 0 4pt; }
.name span { flex: 1; border-bottom: 1pt solid #000; height: 15pt; }
.name span.d { flex: 0 0 1.6in; }
.main { flex: 1; display: flex; flex-direction: column; min-height: 0; padding-top: 6pt; overflow: hidden; }
.foot { display: flex; justify-content: space-between; font-size: 8.5pt; color: #333; border-top: .8pt solid #000; padding-top: 3pt; }
.wm { position: absolute; left: 50%; top: 50%; transform: translate(-50%, -50%) rotate(-38deg); font-family: 'Oswald'; font-weight: 700;
      font-size: 150pt; color: rgba(0,0,0,.13); letter-spacing: .06em; white-space: nowrap; pointer-events: none; z-index: 9; }
h2.s { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 13pt; margin: 9pt 0 3pt; letter-spacing: .02em; }
h2.s .acc, .acc { color: ACCENT; }
p, li { font-size: 10.5pt; line-height: 1.32; margin: 0 0 4pt; }
ul, ol { margin: 0 0 4pt; padding-left: 15pt; }
.tn p, .tn li { font-size: 11.4pt; line-height: 1.45; margin-bottom: 6pt; }
.tn.big p, .tn.big li { font-size: 12.4pt; line-height: 1.48; margin-bottom: 9pt; }
.tn.huge p { font-size: 13.4pt; line-height: 1.5; margin-bottom: 12pt; }
.lines { flex: 1; min-height: 44pt; background: repeating-linear-gradient(to bottom, transparent 0, transparent 21pt, #000 21pt, #000 22pt); }
table.t { border-collapse: collapse; width: 100%; font-size: 9.8pt; margin: 2pt 0 4pt; }
table.t td, table.t th { border: 1pt solid #000; padding: 3pt 5pt; text-align: left; vertical-align: top; line-height: 1.25; }
table.t.tight td { padding: 1.6pt 5pt; font-size: 9.4pt; line-height: 1.2; }
table.t th { font-family: 'Oswald'; font-weight: 500; text-transform: uppercase; font-size: 9.5pt; background: #000; color: #fff; letter-spacing: .04em; vertical-align: middle; }
.cols { display: flex; gap: 20pt; }
.cols > div { flex: 1; min-width: 0; }
.small { font-size: 9pt !important; color: #333; }
.box { border: 1.4pt solid #000; border-radius: 8pt; padding: 7pt 10pt; }
.box.acc-b { border-color: ACCENT; border-width: 2pt; }
.box.inv { background: #000; color: #fff; border-color: #000; }
.k { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 10pt; letter-spacing: .05em; color: ACCENT; display: block; margin-bottom: 2pt; }
.box.inv .k { color: ACC_LIGHT; }
.ln2 { border-bottom: 1pt solid #000; height: 22pt; }
.fl { display: flex; align-items: flex-end; gap: 6pt; font-size: 10.5pt; margin: 0 0 8pt; }
.fl b { font-weight: 600; white-space: nowrap; }
.fl span { flex: 1; border-bottom: 1pt solid #000; height: 14pt; }
.num { list-style: none; padding: 0; margin: 2pt 0 0; }
.num li { display: flex; gap: 10pt; margin: 0 0 9pt; }
.num .n { flex: 0 0 24pt; height: 24pt; border-radius: 12pt; background: #000; color: #fff; font-family: 'Oswald'; font-weight: 700; font-size: 13pt;
          display: flex; align-items: center; justify-content: center; }
.num .hd { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 12pt; letter-spacing: .02em; }
.why { font-style: italic; }
.blk { display: flex; gap: 9pt; margin: 0 0 5pt; font-size: 10.4pt; line-height: 1.3; }
.blk .tm { flex: 0 0 40pt; text-align: center; background: #000; color: #fff; font-family: 'Oswald'; font-weight: 700; font-size: 10pt;
           border-radius: 4pt; padding: 2pt 0 3pt; align-self: flex-start; }
.blk b.h { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 10.5pt; letter-spacing: .02em; }
.day { border: 1.4pt solid #000; border-radius: 8pt; padding: 7pt 10pt 3pt; margin-bottom: 8pt; }
.day .dh { display: flex; justify-content: space-between; align-items: baseline; border-bottom: 1pt solid #000; padding-bottom: 3pt; margin-bottom: 6pt; }
.day .dh b { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 14pt; }
.day .dh span { font-family: 'Oswald'; font-weight: 300; font-size: 11pt; }
.goal { font-size: 10.6pt; margin: 0 0 7pt; }
.chk { display: flex; gap: 8pt; align-items: flex-start; font-size: 11pt; margin: 0 0 8pt; line-height: 1.3; }
table.rub { border-collapse: collapse; width: 100%; }
table.rub td, table.rub th { border: 1.1pt solid #000; padding: 9pt 6pt; vertical-align: top; font-size: 11pt; line-height: 1.38; }
table.rub th { background: #000; color: #fff; font-family: 'Oswald'; font-weight: 500; font-size: 12pt; text-align: center; vertical-align: middle; }
table.rub td.c { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 11pt; width: 15%; vertical-align: middle; }
table.rub td.sc { width: 8%; }
.cover { align-items: center; justify-content: space-between; padding: 0.55in 0.6in 0.5in;
         background: linear-gradient(180deg, #fff 0%, #fff 55%, #ededed 100%); }
.store { font-family: 'Oswald'; font-weight: 500; font-size: 13pt; letter-spacing: .32em; text-transform: uppercase; }
.ctitle { text-align: center; }
.ctitle .a, .ctitle .b, .ctitle .c { white-space: nowrap; }
.ctitle .a { font-family: 'Oswald'; font-weight: 700; line-height: 1; letter-spacing: .02em; text-transform: uppercase; }
.ctitle .b { font-family: 'Oswald'; font-weight: 700; line-height: .98; letter-spacing: .01em; text-transform: uppercase; }
.ctitle .c { font-family: 'Oswald'; font-weight: 300; margin-top: 8pt; }
.cband { width: 100%; background: #000; color: #fff; text-align: center; font-family: 'Oswald'; font-weight: 500; font-size: 14pt;
         letter-spacing: .08em; text-transform: uppercase; padding: 8pt 0 9pt; border-radius: 4pt; white-space: nowrap; }
.cband em { font-style: normal; color: ACCENT; }
table.nt { border-collapse: collapse; }
table.nt td { border: 1.2pt solid #000; width: 40pt; height: 30pt; text-align: center; vertical-align: middle; font-family: 'Oswald'; font-weight: 700; font-size: 15pt; }
table.nt td.g { color: #555; }
table.nt td.st { border: none; height: 16pt; font-family: 'Oswald'; font-weight: 500; font-size: 10.5pt; color: ACCENT; }
table.card { border-collapse: collapse; width: 100%; }
table.card td, table.card th { border: 1.2pt solid #000; padding: 6pt 7pt; vertical-align: middle; font-size: 11.5pt; }
table.card th { background: #000; color: #fff; font-family: 'Oswald'; font-weight: 500; text-transform: uppercase; font-size: 11pt; letter-spacing: .04em; }
table.card td.rn { font-family: 'Oswald'; font-weight: 700; font-size: 17pt; text-align: center; width: 11%; }
table.card td.nt { font-family: 'Oswald'; font-weight: 500; font-size: 13pt; white-space: nowrap; }
table.plan { border-collapse: collapse; width: 100%; }
table.plan td, table.plan th { border: 1.2pt solid #000; padding: 6pt 8pt; font-size: 11pt; }
table.plan th { background: #000; color: #fff; font-family: 'Oswald'; font-weight: 500; text-transform: uppercase; font-size: 11pt; letter-spacing: .04em; text-align: center; }
.rule { background: #000; color: #fff; border-radius: 8pt; padding: 9pt 12pt; }
.rule .k { color: ACC_LIGHT; }
.rule b { font-size: 14pt; line-height: 1.3; display: block; }
""".replace("ACCENT", ACCENT).replace("ACC_LIGHT", ACC_LIGHT)

COPY = f"© {YEAR} {STORE} · {BRAND} · For single-classroom use"


def html_doc(pages, title="doc"):
    return (f'<!doctype html><html><head><meta charset="utf-8"><title>{escape(title)}</title>'
            f'<style>{font_css()}{CSS}</style></head><body>{"".join(pages)}</body></html>')


def foot(right):
    return f'<div class="foot"><span>{COPY}</span><span>{right}</span></div>'


def wm(on):
    return '<div class="wm">PREVIEW</div>' if on else ""


def page(sub, tag, tagsub, inner, n, watermark=False, name=False, ak=False, cls="", key=""):
    top = (f'<div class="top"><div class="t">{TITLE_H}<small>{sub}</small></div>'
           f'<div class="tag"><div class="lv{" ak" if ak else ""}">{tag}</div><div class="ln">{tagsub}</div></div></div>')
    nm = '<div class="name">Name<span></span>Date<span class="d"></span></div>' if name else ""
    return (f'<section class="page {cls}" data-key="{key}">{wm(watermark)}{top}{nm}<div class="main">{inner}</div>{foot(f"Page {n}")}</section>')


def checkbox(size=12):
    return (f'<svg width="{size}pt" height="{size}pt" viewBox="0 0 {size} {size}" style="flex:0 0 auto;margin-top:1pt">'
            f'<rect x=".7" y=".7" width="{size - 1.4}" height="{size - 1.4}" rx="2" fill="#fff" stroke="#000" stroke-width="1.2"/></svg>')


def t_(x, y, s, size, weight=700, family="Oswald", anchor="middle", fill="#000", extra=""):
    fam = f"'{family}'" if " " in family else family
    return (f'<text x="{x:.1f}" y="{y:.1f}" font-family="{fam}" font-weight="{weight}" font-size="{size}" text-anchor="{anchor}" '
            f'fill="{fill}" {extra}>{escape(s)}</text>')


def svgwrap(vw, vh, body, width="100%"):
    return f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {vw} {vh}" width="{width}" style="display:block">{body}</svg>'


_OSWALD_ADV = {}


def text_w(s, size):
    """Advance width of s in Oswald 700 at size (viewBox units / px / pt), from the font's own metrics."""
    if not _OSWALD_ADV:
        from fontTools.ttLib import TTFont
        f = TTFont(os.path.join(FONTS, "oswald-latin-700-normal.woff2"))
        cmap, hm = f.getBestCmap(), f["hmtx"].metrics
        _OSWALD_ADV["upm"] = f["head"].unitsPerEm
        _OSWALD_ADV.update({chr(c): hm[g][0] for c, g in cmap.items() if c < 0x2200})
    return sum(_OSWALD_ADV[ch] for ch in s) * size / _OSWALD_ADV["upm"]


# ================================================================ drawings (original, code-drawn: black + accent only); every label is logged for the fit check
LABELS = []   # (diagram, text, size, room) : text_w(text, size) must be <= room


def lab(diagram, text, size, room):
    LABELS.append((diagram, text, size, room))


def tiles_for(key, prog):
    return [(rn, short_name(by_rn(key, rn)), rn.upper() == "I") for rn in prog]


ROLL_LS = 16          # gutter row label size (letters only, never octave numbers: Logic and other DAWs number octaves differently)


def roll_svg(name, tiles, lanes, foot=None, mw=200, rh=24, gut=150, width="100%"):
    """A piano roll: Roman numeral tiles per measure, then lanes (title, lowest MIDI, highest MIDI, notes, fill); notes = [(midi, start beat, beats)].
    Grey rows are black keys. Every note carries its letter name inside the bar; every label is checked for fit."""
    out, top = [], 92
    W = gut + 4 * mw + 4
    for m, (rn, nm, home) in enumerate(tiles):
        x = gut + m * mw
        out.append(f'<rect x="{x + 4}" y="6" width="{mw - 8}" height="70" rx="9" fill="{ACCENT if home else "#000"}"/>')
        out.append(t_(x + mw / 2, 42, rn, 32, 700, fill="#fff"))
        out.append(t_(x + mw / 2, 66, nm, 17, 500, fill="#fff"))
        lab(name, rn, 32, mw - 20)
        lab(name, nm, 17, mw - 20)
    y0 = top
    for title, lo, hi, notes, fillc in lanes:
        n = hi - lo + 1
        for i, mn in enumerate(range(hi, lo - 1, -1)):
            y = y0 + i * rh
            black = midi_name(mn).endswith("#")
            out.append(f'<rect x="{gut}" y="{y}" width="{4 * mw}" height="{rh}" fill="{"#e6e6e6" if black else "#fff"}"/>')
            if not black:
                out.append(t_(gut - 8, y + rh * 0.74, midi_name(mn), ROLL_LS, 500, anchor="end"))
                lab(f"{name} gutter", midi_name(mn), ROLL_LS, 16)
        out.append(t_(10, y0 + n * rh / 2 + 8, title.upper(), 20, 700, anchor="start"))
        lab(f"{name} gutter", title.upper(), 20, gut - 40)
        for b in range(17):   # beat lines (light), measure lines (dark)
            xx = gut + b * mw / 4
            meas = b % 4 == 0
            out.append(f'<line x1="{xx:.1f}" y1="{y0}" x2="{xx:.1f}" y2="{y0 + n * rh}" stroke="{"#000" if meas else "#bdbdbd"}" '
                       f'stroke-width="{1.6 if b in (0, 16) else (1.1 if meas else 0.7)}"/>')
        out.append(f'<rect x="{gut}" y="{y0}" width="{4 * mw}" height="{n * rh}" fill="none" stroke="#000" stroke-width="1.8"/>')
        for mn, st, d in notes:
            if not lo <= mn <= hi:
                raise ValueError(f"{name}: note {mn} outside the {title} lane")
            x, w = gut + st * mw / 4 + 2, d * mw / 4 - 4
            y = y0 + (hi - mn) * rh
            out.append(f'<rect x="{x:.1f}" y="{y + 1.5}" width="{w:.1f}" height="{rh - 3}" rx="4" fill="{fillc}"/>')
            out.append(t_(x + 5, y + rh * 0.76, midi_name(mn), ROLL_LS + 1, 700, anchor="start", fill="#fff"))
            lab(name, midi_name(mn), ROLL_LS + 1, w - 7)
        y0 += n * rh + 26
    H = y0 - 26
    for m in range(4):
        s = foot[m] if foot else f"Measure {m + 1}"
        out.append(t_(gut + m * mw + mw / 2, H + 20, s, 15, 500))
        lab(name, s, 15, mw - 10)
    return svgwrap(W, H + 28, "".join(out), width)


def mel_lane(key, title="Melody"):
    notes = flat(MELODY[key])
    lo, hi = min(n for n, _, _ in notes), max(n for n, _, _ in notes)
    return (title, lo, hi, notes, ACCENT)


def file_bass(key):
    return [(bass_root_midi(by_rn(key, rn)[0]), 4 * i, 4) for i, rn in enumerate(FILE_PROG[key])]


def bass_lane(notes, title="Bass"):
    lo, hi = min(n for n, _, _ in notes) - 1, max(n for n, _, _ in notes) + 1
    return (title, lo, hi, notes, "#000")


def melroll_svg(key="C major", width="100%", bass=True):
    lanes = [mel_lane(key)] + ([bass_lane(file_bass(key), "Bass (file)")] if bass else [])
    return roll_svg(f"melroll {key}{'' if bass else ' page'}", tiles_for(key, FILE_PROG[key]), lanes, width=width)


def bassroll_svg(key="C major", width="100%"):
    return roll_svg(f"bassroll {key}", tiles_for(key, BASS_PROG[key]), [bass_lane(flat(BASS[key]))], foot=BASS_IDEAS, rh=30, width=width)


def hero_svg(width="100%"):
    """Cover art: I-vi-IV-V tiles, the worked melody in the accent and a rhythmic bassline in black, as a piano roll."""
    return roll_svg("hero", tiles_for("C major", FILE_PROG["C major"]), [mel_lane("C major"), bass_lane(flat(HERO_BASS))],
                    foot=["Home", "Tender", "Lifting", "Homesick"], width=width)


def verify_labels():
    """Every drawn label fits its tile, row or note bar (Oswald 700 metrics); row text fits its row height."""
    LABELS.clear()
    hero_svg(), melroll_svg("C major"), melroll_svg("A minor"), bassroll_svg("C major"), bassroll_svg("A minor")
    melroll_svg("C major", bass=False), melroll_svg("A minor", bass=False)
    errs = [f"{d}: label '{t}' at {s} is {text_w(t, s):.0f} wide, room {r:.0f}" for d, t, s, r in LABELS if text_w(t, s) > r]
    if not LABELS:
        errs.append("label check found no labels")
    if (ROLL_LS + 1) * 0.72 > 24 - 6:   # cap height of the largest row label inside one 24px row, with 3px clear above and below
        errs.append("piano-roll labels are too tall for their rows")
    return errs


# ================================================================ page plan
PLAN = ["cover", "howto", "overview", "setup", "logicref", "pacing", "day12", "day34", "tn1", "tn2", "tn3", "tn4",
        "post1", "post2", "pk2", "pk4", "cardC", "cardA", "methods", "proj2a", "proj2b", "rub2", "proj4a", "proj4b", "rub4", "vocab",
        "key1", "exmel", "exbass", "terms"]
IX = {k: i + 1 for i, k in enumerate(PLAN)}
PREVIEW_PICKS = ["pacing", "methods", "pk4", "rub4"]


def rng(a, b):
    return str(a) if a == b else f"{a}–{b}"


def day_slides(d):
    sl = [s for blk in d["blocks"] if blk[3] for s in range(blk[3][0], blk[3][1] + 1)]
    return (min(sl), max(sl)) if sl else None


def slides_label(d):
    r = day_slides(d)
    return f"Intro deck slides {rng(*r)}" if r and r[0] != r[1] else (f"Intro deck slide {r[0]}" if r else "No slides")


def fill(s):
    """Page placeholders in the day plans are written @NAME so they never collide with ordinary words."""
    pages = {"PROJ2": f"{IX['proj2a']}–{IX['proj2b']}", "PROJ4": f"{IX['proj4a']}–{IX['proj4b']}", "PK2": IX["pk2"], "PK4": IX["pk4"],
             "CARDS": f"{IX['cardC']}–{IX['cardA']}", "METH": IX["methods"], "RUB2": IX["rub2"], "RUB4": IX["rub4"], "TN2": IX["tn2"],
             "TN4": IX["tn4"]}
    out = re.sub(r"@([A-Z0-9]+)", lambda m: str(pages[m.group(1)]), s)
    if "@" in out:
        raise ValueError(f"unfilled placeholder in: {s}")
    return out


def H2(s, first=False):
    st = ' style="margin-top:2pt"' if first else ""
    return f'<h2 class="s"{st}>{s}<span class="acc"> ·</span></h2>'


COVER_A = "Melody Writing in Logic Pro"
COVER_B = "Bass & Melody"
COVER_C = f"DMP1 · Unit {UNIT} · basslines that follow the roots · melodies you can hum"
COVER_BAND = f"Grades 9–12 · {DAYS_N} days · {len(PLAN)} pages + {DECK_SLIDES}-slide deck"
# cover sizes (pt on the PDF cover, px on the square); fit checked against the available width with Oswald 700 metrics
PDF_COVER = dict(a=30, b=76, c=15, band=14, room=7.3 * 72)
SQ_COVER = dict(a=52, b=136, c=29, band=26, room=940)


def cover_page():
    s = PDF_COVER
    band = COVER_BAND.replace(" · ", " <em>·</em> ")
    return (f'<section class="page cover" data-key="cover"><div class="store">{BRAND}</div>'
            f'<div style="width:7.2in">{hero_svg()}</div>'
            f'<div class="ctitle"><div class="a" style="font-size:{s["a"]}pt">{COVER_A}</div><div class="b" style="font-size:{s["b"]}pt">{escape(COVER_B)}</div>'
            f'<div class="c" style="font-size:{s["c"]}pt">{COVER_C}</div></div>'
            f'<div class="cband">{band}</div></section>')


def deck_day_line():
    parts = []
    for d in DAYS:
        r = day_slides(d)
        parts.append(f"Day {d['n']} (slide{'s' if r[0] != r[1] else ''} {rng(*r)})")
    return ", ".join(parts[:-1]) + " and " + parts[-1]


def howto_page(n):
    parts = [("Teacher pages", "howto", "tn4", "How to use, overview, setup + the Logic file, Logic reference, pacing, day plans, teacher notes"),
             ("Student pages", "post1", "vocab", "Two classroom posts, packet Part 2 (scales) and Part 4 (basslines), Reference Cards, four ways to "
              "write a melody, two project sheets with their rubrics, vocabulary"),
             ("Answer keys + examples", "key1", "exbass", "Packet answers, a worked melody and a worked bassline in both keys"),
             ("Intro deck (separate file)", None, None, f"{DECK_SLIDES} slides, .pptx + PDF, speaker notes on every slide")]
    trs = "".join(f"<tr><td><b>{a}</b></td><td style='white-space:nowrap'>{('pages ' + str(IX[f]) + '–' + str(IX[l])) if f else 'deck'}</td><td>{u}</td></tr>"
                  for a, f, l, u in parts)
    inner = (f'<div class="cols"><div>'
             f'{H2("What this is", True)}'
             f'<p>A complete {DAYS_N}-day bass and melody unit for high school music technology. Days 1 and 2: students use a scale to write a short, '
             f'singable melody over a chord progression and bassline they already have. At the end of Day 2 they look hard at that bassline and '
             f'discover the rule behind every bassline. Day 3: they choose four chords with Roman numerals and write their own bassline. Day 4: they '
             f'present both.</p>'
             f'<p>This is Unit {UNIT} of my Digital Music Production 1 course. It picks up where Unit 3, Chords by Ear (a separate resource), '
             f'ends: students continue in their Unit 3 project. Students who did not do Unit 3 use your chords-and-bassline file (recipe on page '
             f'{IX["setup"]}) and the Reference Cards in this kit. C major or A minor, white keys only. No music reading is needed.</p>'
             f'{H2("Quick start")}<ol>'
             f'<li>Read the overview and pacing (pages {IX["overview"]}–{IX["pacing"]}) and the teacher notes (pages {IX["tn1"]}–{IX["tn4"]}).</li>'
             f'<li>Check who has a Unit 3 project; for everyone else make the chords-and-bassline file once (10 minutes, page {IX["setup"]}).</li>'
             f'<li>Cue up {SNA[0]} ({SNA[1]}) and {SBM[0]} ({SBM[1]}) on YouTube or Spotify.</li>'
             f'<li>Print the student pages (pages {IX["post1"]}–{IX["vocab"]}), one set per student.</li></ol>'
             f'{H2("The intro deck")}'
             f'<p>Every slide has a day: {deck_day_line()}. The day plans name the slide for each block. Every slide has speaker notes.</p>'
             f'<p class="small">Fonts: the deck uses Oswald and Source Sans 3 (free on Google Fonts, built into Google Slides). If PowerPoint or Keynote '
             f'shows a different font, install the two fonts or present the PDF copy, which shows the intended look.</p>'
             f'</div><div>'
             f'{H2("What’s inside", True)}<table class="t"><tr><th>Part</th><th>Pages</th><th>Use</th></tr>{trs}</table>'
             f'<div class="box acc-b" style="margin-top:6pt"><span class="k">About the songs</span><p style="margin:0">{SONG_NOTE}</p></div>'
             f'{H2("Make it fit your class")}'
             f'<ul><li><b>File names:</b> the project sheets keep my naming rule, “{PROJECT2}” and “Part 4” (The Building Blocks is my name '
             f'for this theory sequence). Change it to suit your class.</li>'
             f'<li><b>Grading:</b> each project has my /{RUBRIC_TOTAL} rubric. My course plan treats this unit as a check-in; if you prefer '
             f'that, use the submission checklists alone.</li>'
             f'<li><b>More accessible / more advanced:</b> variations on page {IX["tn4"]}.</li>'
             f'<li><b>No MIDI keyboard?</b> Musical Typing (Command-K), click-in or the step sequencer works for every step.</li></ul>'
             f'</div></div>')
    return page("How to use this resource", "START HERE", "Teacher", inner, n, key="howto")


def overview_page(n):
    inner = (f'<div class="cols"><div>'
             f'{H2("What students make", True)}'
             f'<p>A bassline and short melody written over a chord progression in Logic: the foundational skills a full-song or EDM project will '
             f'require. Two short projects: a 2- or 4-measure melody over the chords and bassline from Unit 3 (or your file), then an original '
             f'bassline over four chords each student chooses from the Reference Card.</p>'
             f'{H2("What students learn")}<p>{PLAN_LEARN}</p>'
             f'{H2("Why this unit lands here")}<p>{PLAN_WHY}</p>'
             f'{H2("Mixing concept: frequency separation")}'
             f'<p>Bass and melody stacked together sound muddy. A brief EQ introduction gives students the tool to solve the problem they just '
             f'created: high-pass the melody, roll off the extreme highs of the bass.</p>'
             f'</div><div>'
             f'{H2("Goals", True)}<ul>'
             f'<li>Use a scale as the ingredient list for a melody</li>'
             f'<li>Write a short, singable melody that starts or ends on the root</li>'
             f'<li>Enter notes by playing in (and quantizing) and by clicking in</li>'
             f'<li>Discover and state the bassline rule: the bass changes to the root of the new chord</li>'
             f'<li>Choose four chords with Roman numerals and explain the feeling they create</li>'
             f'<li>Write a monophonic bassline with at least two rhythmic ideas</li>'
             f'<li>Give bass and melody their own frequency homes with EQ</li></ul>'
             f'{H2("At a glance")}<table class="t">'
             f'<tr><td><b>Grades</b></td><td>9–12 (no music reading needed)</td></tr>'
             f'<tr><td><b>Length</b></td><td>{DAYS_N} periods of about {MINUTES} minutes: intro, two work days, present</td></tr>'
             f'<tr><td><b>Keys</b></td><td>C major or A minor, white keys only</td></tr>'
             f'<tr><td><b>Software</b></td><td>Logic Pro: piano roll, record + quantize, Pencil, step sequencer, copy and paste, Channel EQ, export</td></tr>'
             f'<tr><td><b>Product</b></td><td>A melody project and a bassline project, each an MP3 + compressed project; packet Parts 2 and 4; a planning table</td></tr>'
             f'<tr><td><b>Assessment</b></td><td>Two 3-category rubrics, /{RUBRIC_TOTAL} each (pages {IX["rub2"]} and {IX["rub4"]})</td></tr></table>'
             f'{H2("Standards")}'
             f'<p>Supports the National Core Arts Standards for Music (Music Technology, high school): developing musical ideas (MU:Cr2.1), '
             f'refining work (MU:Cr3.1) and presenting work (MU:Pr6.1).</p>'
             f'</div></div>')
    return page("Unit overview", "OVERVIEW", f"DMP1 · Unit {UNIT}", inner, n, key="overview")


def setup_page(n):
    cb = " – ".join(midi_name(b) for b, _, _ in file_bass("C major"))
    ab = " – ".join(midi_name(b) for b, _, _ in file_bass("A minor"))
    cch = ", ".join(f"{rn} {dash(by_rn('C major', rn))}" for rn in FILE_PROG["C major"])
    recipe = [
        "New project, a Software Instrument track with a bass sound (Library: Bass). A tempo between 80 and 100 BPM is comfortable.",
        f"Draw a 4-measure MIDI region. In the piano roll add one whole note per measure, the root of each chord, low on the keyboard. "
        f"C major: {cb}. A minor: {ab}.",
        f"A second track with a piano or pad sound: the four chords, three notes stacked per measure. C major: {cch}. A minor: "
        f"{', '.join(FILE_PROG['A minor'])} (spellings on the A minor card).",
        "Turn on Cycle over the 4 measures; add arrangement markers labeling the chords.",
        "Leave the melody track out: students add it. Share the project with your class.",
    ]
    inner = (f'<div class="cols"><div>'
             f'{H2("You need", True)}<ul>'
             f'<li><b>A Mac with Logic Pro</b> and <b>headphones</b> for each student; a MIDI keyboard (an MPK Mini or similar) if you have one.</li>'
             f'<li><b>A piano sound on the class speakers</b> for the Day 1 scales and demos.</li>'
             f'<li><b>Projector</b> for the deck; <b>two songs</b> from your own YouTube or Spotify access: {SNA[0]} ({SNA[1]}) and {SBM[0]} '
             f'({SBM[1]}).</li>'
             f'<li><b>Printed student pages</b>, one set each: two posts, packet Parts 2 and 4, the Reference Cards, four ways to write a melody, '
             f'both project sheets and rubrics, vocabulary.</li>'
             f'<li><b>A chords-and-bassline Logic file</b> for every student (right).</li>'
             f'<li><b>A place to submit</b> an MP3 and a .zip, and a class feedback board for Day 4 (I use Padlet).</li></ul>'
             f'{H2("Before Day 1")}'
             + "".join(f'<div class="chk">{checkbox()}<span>{c}</span></div>' for c in [
                 "List who has a Unit 3 project; make and share your file for everyone else.",
                 "Test the class speakers; cue up the two songs.",
                 "Check that students can play notes: MPK Mini (find OCT +), or Musical Typing (Command-K).",
                 "Print the student pages; set up the submission spot and the feedback board."]) +
             f'</div><div>'
             f'<div class="box acc-b"><span class="k">The Logic file: continue from Unit 3</span>'
             f'<p style="margin:0">Students who did Unit 3 already have it: their Part 3 project is a bassline plus four chords. On Day 1 they '
             f'open it and save a copy as “{PROJECT2}”, then add the melody on a new track. On Day 2 the bassline in the same file is the one '
             f'they analyze.</p></div>'
             f'<div class="box" style="margin-top:8pt"><span class="k">No Unit 3 project? Make the file once (10 minutes)</span>'
             f'<p style="margin:0 0 3pt">My recommended setup for the Part 2 file:</p><ul style="margin:0 0 4pt">'
             + "".join(f"<li>{x}</li>" for x in FILE2) + '</ul>'
             f'<p style="margin:0 0 3pt"><b>Recipe</b> (one chord per measure, as in the Unit 3 project):</p><ol style="margin:0">'
             + "".join(f"<li>{x}</li>" for x in recipe) + '</ol>'
             f'<p class="small" style="margin:4pt 0 0">Worked examples on pages {IX["exmel"]}–{IX["exbass"]}. Logic sessions are not included in '
             f'this kit.</p></div>'
             f'</div></div>')
    return page("Setup and requirements", "SETUP", "Teacher", inner, n, key="setup")


LOGIC_REF = [("New instrument track", "Track › New Software Instrument Track (Option-Command-S). Open the Library (Y): a lead or synth for "
              "the melody, a bass sound for the bassline."),
             ("Octave up for melody", "Press OCT + on the MPK Mini once before recording. With Musical Typing (Command-K), X moves up an "
              "octave and Z down."),
             ("Play it in", "Turn on Cycle (C) over 2 measures. Press R to record, play, press the Space bar to stop."),
             ("Quantize", "Select the notes in the piano roll (Command-A) and press Q. It fixes timing, not notes."),
             ("Click it in", "Open the piano roll (P). Choose the Pencil tool (T menu) and click to add a note; drag its right edge to set the "
              "length; drag it up or down to change the pitch."),
             ("Step sequencer", "Control-click the track lane and choose Create Pattern Region. Turn steps on to place notes; each row is one pitch."),
             ("Copy across", "Option-drag a region to copy it, or select it and press Command-R to repeat it."),
             ("Stack a chord", "Three notes in one column, same start, same length. Spellings are on the Reference Card."),
             ("No overlaps", "Zoom in on the bass in the piano roll: every note must end before the next one starts."),
             ("EQ", "Click the EQ area at the top of a channel strip to open Channel EQ. Melody: turn on the low cut (high-pass) and drag it up. "
              "Bass: turn on the high cut."),
             ("Export", "File › Bounce › Project or Section (Command-B), MP3. In Finder, Control-click the .logicx project and choose Compress.")]


def logicref_page(n):
    rows = "".join(f"<tr><td style='width:23%;padding:5pt 6pt'><b>{a}</b></td><td style='padding:5pt 6pt'>{b}</td></tr>" for a, b in LOGIC_REF)
    inner = (f'<p>Everything students need in Logic for this unit, in one table. Demonstrate once (slides 9 and 18), then let students work.</p>'
             f'<table class="t" style="font-size:12.2pt;line-height:1.45">{rows}</table>'
             f'<div class="cols" style="margin-top:8pt"><div class="box"><span class="k">Melody check in the piano roll</span>'
             f'<p style="margin:0">One note at a time, high on the keyboard, on its own track. White keys only. The first or last note is C '
             f'(C major) or A (A minor).</p></div>'
             f'<div class="box"><span class="k">Version note</span>'
             f'<p style="margin:0" class="small">Menu names are from current Logic Pro versions; older versions may differ slightly. Logic Pro is a '
             f'trademark of Apple Inc., named only to describe what the unit needs. This resource is not affiliated with Apple. MPK Mini is named '
             f'only as an example keyboard.</p></div></div>')
    return page("Logic quick reference", "SETUP", "Teacher", inner, n, key="logicref")


def pacing_page(n, watermark=False):
    rows = "".join(
        f"<tr><td style='white-space:nowrap'><b class='acc' style='font-family:Oswald;font-size:13pt'>DAY {d['n']}</b><br><span class='small'>{d['kind']}</span></td>"
        f"<td><b>{d['title']}</b><br><span class='small'>{d['goal']}</span></td>"
        f"<td style='font-size:9.4pt;line-height:1.38'>{'<br>'.join(f'{a} · {h}' for a, h, _, _ in d['blocks'])}</td>"
        f"<td style='font-size:9.4pt'>{slides_label(d)}</td></tr>" for d in DAYS)
    inner = (f'<p>Four class periods of about {MINUTES} minutes, as in my course plan: an intro day, two work days and a presenting day. The '
             f'melody comes first, then the bassline. Each day’s full plan is on pages {IX["day12"]}–{IX["day34"]}.</p>'
             f'<table class="t"><tr><th style="width:13%">Day</th><th style="width:37%">Focus and goal</th><th style="width:31%">Minute blocks</th>'
             f'<th>Deck</th></tr>{rows}</table>'
             f'{H2("Student pages by day")}<table class="t">'
             f'<tr><td style="width:13%"><b>Day 1</b></td><td>Melody post (p. {IX["post1"]}), packet Part 2 (p. {IX["pk2"]}), Reference Cards '
             f'(pp. {IX["cardC"]}–{IX["cardA"]}), four ways (p. {IX["methods"]}), melody project sheet (pp. {IX["proj2a"]}–{IX["proj2b"]})</td></tr>'
             f'<tr><td><b>Day 2</b></td><td>Melody rubric (p. {IX["rub2"]}), basslines post (p. {IX["post2"]}, steps 1–3), packet Part 4 '
             f'analysis (p. {IX["pk4"]})</td></tr>'
             f'<tr><td><b>Day 3</b></td><td>Basslines post steps 4–5, bassline project sheet with the planning table (pp. {IX["proj4a"]}–'
             f'{IX["proj4b"]}), rubric (p. {IX["rub4"]})</td></tr>'
             f'<tr><td><b>Day 4</b></td><td>Vocabulary (p. {IX["vocab"]}) and the two rubrics for scoring</td></tr></table>'
             f'<div class="box" style="margin-top:6pt"><span class="k">My line for the whole unit</span><p style="margin:0;font-size:12pt">'
             f'<b>“{SIMPLE}”</b> Say it on Day 1. Say it again when a student plays you a melody that goes nowhere because they tried to use every '
             f'note in the scale.</p></div>')
    return page(f"Pacing guide · {DAYS_N} days", "PACING", "Teacher", inner, n, watermark, key="pacing")


def day_box(d):
    def blk(a, h, t, s):
        tag = f' <span class="small">[{"slide" if s[0] == s[1] else "slides"} {rng(*s)}]</span>' if s else ""
        return f'<div class="blk"><div class="tm">{a}</div><div><b class="h">{h}</b>{tag} · {fill(t)}</div></div>'
    return (f'<div class="day"><div class="dh"><b>Day {d["n"]} · {d["title"]}</b><span>{d["kind"]} · {slides_label(d)}</span></div>'
            f'<p class="goal"><b>Goal:</b> {d["goal"]}</p>{"".join(blk(*b) for b in d["blocks"])}</div>')


def day12_page(n):
    return page("Day-by-day plans · Days 1 and 2", "DAYS 1–2", "Melody", day_box(DAYS[0]) + day_box(DAYS[1]), n, key="day12")


def day34_page(n):
    prompts = ('<div class="box"><span class="k">Circulating prompts (Days 1–3)</span><ul style="margin:0">'
               '<li>‘Hum it for me. Could you hum it after hearing it once?’</li>'
               '<li>‘Point to where your chord changes. Does your bass note change there?’</li>'
               '<li>‘Tell me what you want this to feel like. Which Roman numerals would create that feeling?’</li></ul></div>')
    pres = ('<div class="box acc-b" style="margin-top:8pt"><span class="k">The presentation question that matters (Day 4)</span>'
            f'<p style="margin:0">Why did you choose this chord progression? What did you want it to feel like? Listen for Roman numeral language: '
            f'“I wanted it to feel unresolved so I ended on the V.” More on page {IX["tn2"]}.</p></div>')
    return page("Day-by-day plans · Days 3 and 4", "DAYS 3–4", "Bassline + present", day_box(DAYS[2]) + day_box(DAYS[3]) + prompts + pres, n,
                key="day34")


def tn1_page(n):
    inner = (f'<div class="tn big">{H2("Part one: a melody (Days 1–2)", True)}'
             f'<p>{WHAT2}</p><p>{KEY2}</p>'
             f'<p>My course plan adds four ways to write a melody (page {IX["methods"]}): play it in, draw a rhythm first, click one note and move '
             f'it, use the step sequencer. And the most important thing you can say in this unit: “{SIMPLE}”</p>'
             f'{H2("Why a standalone unit")}<p>{PLAN_CHANGE}</p>'
             f'<div class="box" style="margin-top:4pt"><span class="k">How this kit differs from my classroom version</span><ul style="margin:0">'
             f'<li>My Building Blocks material gives Part 2 (melody) 2 days and Part 4 (basslines) 2 days plus presentations, labeled Unit 3. This kit '
             f'follows the 4 days of my course plan (intro, two work days, present) and is labeled Unit {UNIT}. Part 4’s analysis and reveal move '
             f'to the end of Day 2 (12 + 6 minutes instead of 15 + 10). Project sheets say Unit {UNIT} and the days.</li>'
             f'<li>Students write the melody in their Unit 3 project (it holds a bassline and four chords) or in your file, and analyze the bassline '
             f'in that same file, so the basslines post says “Open the Logic file from your melody project” instead of “Download today’s Logic file”. '
             f'The melody post asks for Part 2 of the packet only (Part 1 is in Unit 3).</p>'
             f'<li>My LMS is written as “your class page” and my Padlet steps as “the class feedback board”.</li>'
             f'<li>New: the four-ways page, the frequency separation mini-lesson, the Logic file recipe, the worked examples, the answer keys and '
             f'the deck.</li></ul></div></div>')
    return page("Teacher notes · the melody", "TEACHER NOTES", "1 of 4", inner, n, key="tn1")


def tn2_page(n):
    inner = (f'<div class="tn big">{H2("Part two: basslines (Days 2–4)", True)}'
             f'<p>{WHAT4[0]}</p><p>{WHAT4[1]}</p>'
             f'{H2("The reveal moment: how to frame it")}<p>{REVEAL[0]}</p>'
             f'<div class="rule" style="margin:2pt 0 8pt"><b>{RULE}</b></div>'
             f'<p>{REVEAL[1]}</p>'
             f'{H2("The presentation question that matters")}<p>{PRESQ[0]}</p><p>{PRESQ[1]}</p>'
             f'<p class="small">In class I also quote my tutorial line: “{BASS_LINE}”</p></div>')
    return page("Teacher notes · basslines", "TEACHER NOTES", "2 of 4", inner, n, key="tn2")


def tn3_page(n):
    st = "".join(f'<p><b>{a}:</b> {b}</p>' for a, b in STRUGGLES)
    inner = (f'<div class="tn huge">{H2("Common student struggles · melody", True)}{st.split("<p><b>Not changing")[0]}'
             f'{H2("Common student struggles · basslines")}<p><b>Not changing{st.split("<p><b>Not changing")[1]}</div>')
    return page("Teacher notes · common struggles", "TEACHER NOTES", "3 of 4", inner, n, key="tn3")


def tn4_page(n):
    inner = (f'<div class="tn">{H2("Suggested variations (melody)", True)}'
             f'<p><b>More accessible:</b> {VAR_EASY}</p><p><b>More advanced:</b> {VAR_HARD}</p>'
             f'{H2("Mixing concept: frequency separation (Day 2)")}<p>{PLAN_MIX}</p>'
             f'{H2("Grading tips")}<ul>'
             f'<li>Two rubrics, each with three categories scored 2–5, for /{RUBRIC_TOTAL} each (pages {IX["rub2"]} and {IX["rub4"]}).</li>'
             f'<li>Melody: open the piano roll. In key? 2 or 4 measures? First or last note C or A? Quantized?</li>'
             f'<li>Bassline: point to each chord change. Is the root there? Any overlapping notes? Two rhythm ideas? Then read the why answer.</li>'
             f'<li>Worked examples and answer keys: pages {IX["key1"]}–{IX["exbass"]}.</li></ul>'
             f'{H2("Resources")}<ul>'
             f'<li>The intro deck (.pptx and PDF), included. The two songs, from your own YouTube or Spotify access.</li></ul>'
             f'{H2("Observation notes")}'
             f'<p class="small">What worked, what to adjust, and anything worth remembering for next year.</p></div><div class="lines"></div>')
    return page("Teacher notes · variations, mixing, grading", "TEACHER NOTES", "4 of 4", inner, n, key="tn4")


def post(n, key, title, sub, intro, bold, steps, box_title, box_body, tag):
    lis = "".join(f'<li style="margin-bottom:13pt;font-size:12.4pt;line-height:1.38"><div class="n">{i + 1}</div><div><b class="hd" style="font-size:13.5pt">{a}</b> · {bb}<br>'
                  f'<span class="why">{cc}</span></div></li>' for i, (a, bb, cc) in enumerate(steps))
    inner = (f'<h2 class="s" style="font-size:20pt;margin:4pt 0 2pt">{title}</h2>'
             f'<p class="l" style="font-size:12pt">{sub}</p>'
             f'<p style="font-size:12.4pt;margin:8pt 0 6pt">{intro}</p>'
             f'<p style="font-size:12.4pt;margin:0 0 12pt"><b>{bold}</b></p>'
             f'<ol class="num">{lis}</ol>'
             f'<div class="box acc-b" style="margin-top:4pt"><span class="k" style="font-size:12pt">{box_title}</span>{box_body}</div>')
    return page("Classroom post", "STUDENT", tag, inner, n, key=key)


POST1_STEPS = [
    ("Part 2 — Scales", "Complete the C Major, A Natural Minor, and C Major Pentatonic scale exercises in the packet.",
     "Apply the formula one step at a time. Play each scale on your MPK Mini when you finish it. (Did them in Unit 3? Check them against your Reference Card.)"),
    ("Open the Logic file", f"Open your Unit 3 chords project and save a copy as {PROJECT2}. No Unit 3 project? Download today’s Logic file from "
     "your class page.", "A chord progression and bassline are already written. Do not change them."),
    ("Listen twice", "Play the full progression from beginning to end twice before you write anything.",
     "What key is this in? What do the chords feel like? Where do they change?"),
    ("Choose your scale", "Decide whether you are writing in C Major or A Minor. Both use the same white keys — the difference is which note feels "
     "like home.", "C Major: start and end on C. A Minor: start and end on A."),
    ("Write your melody", "Add a new track. Choose a lead or synth sound. Use the white keys on your MPK Mini to find a short phrase — 2 or 4 "
     "measures — that you like. Record it (R) and quantize it (Q).", "Simple and repetitive will always beat complicated and forgettable.")]
POST1_BOX = ("<p style='margin:0;font-size:11.5pt'><b>Play-in:</b> press R to record, play on the MPK Mini, press Q to quantize after.</p>"
             "<p style='margin:2pt 0 0;font-size:11.5pt'><b>Click-in:</b> select the pencil tool, hold Command, and click notes into the piano roll.</p>")
POST2_STEPS = [
    ("Open the Logic file", "Open the Logic file from your melody project. A bassline is already written. Open the piano roll and look at it "
     "carefully.", "Do not change anything yet. Just look."),
    ("Part 4 — Analysis", "Complete the analysis questions in the packet. What note does the bassline play when each chord is happening?",
     "You are looking for the pattern. What does the bass note have to do with the chord?"),
    ("The reveal", "Read the rule in the packet. Does it match what you found in your analysis?", "Every bassline you have ever heard follows this rule."),
    ("Plan your progression", "Fill in the chord planning table in the packet. Choose 4 chords from your Reference Card using Roman numerals. Write "
     "the root note for each.", "Think about what you want it to feel like. I to V feels resolved. Ending on V feels unfinished."),
    ("Write your bassline", "Open a new Logic file. Build your chord progression, then write a bassline that follows the roots. Start with just the "
     "root note for the full duration of each chord. Then add rhythm and interest.", "Get the roots right first. Add everything else second.")]


def post1_page(n):
    return post(n, "post1", "Scales &amp; Melody · Days 1–2", f"The Building Blocks · Part 2 · Digital Music Production 1 · Unit {UNIT}",
                "Today you are writing your first melody in Logic. Before you touch the DAW, complete the scale pages in the packet.",
                f"Do not open Logic until you have finished Part 2 of the packet (page {IX['pk2']}).", POST1_STEPS,
                "Two ways to enter your melody", POST1_BOX, "Day 1")


def post2_page(n):
    box = (f'<p style="margin:0;font-size:12.5pt"><b>{RULE}</b></p>'
           f'<p style="margin:3pt 0 0;font-size:11pt" class="why">Everything else is rhythm and decoration on top of this rule.</p>')
    return post(n, "post2", "Basslines · Days 2–3", f"The Building Blocks · Part 4 · Digital Music Production 1 · Unit {UNIT}",
                "Today you discover the rule behind every bassline you have ever heard — and write one yourself.",
                f"Complete Part 4 of the packet (page {IX['pk4']}) before writing your bassline in Logic.", POST2_STEPS,
                "The one rule for basslines", box, "Days 2–3")


def note_row(first, n_cells, formula, given=1, final=None):
    cells = "".join(f'<td>{first[i] if i < given else ""}</td>' for i in range(n_cells))
    cells += f'<td class="g">({final})</td>' if final else ""
    steps = "".join(f'<td class="st">{s}</td>' for s in steps_of(formula))
    return (f'<table class="nt"><tr>{cells}</tr></table>'
            f'<table class="nt" style="margin-left:20pt"><tr>{steps}</tr></table>')


def pk2_page(n):
    cmaj = key_scale("C major")
    inner = (f'<p style="margin:0 0 4pt">A scale is an ordered set of notes built using a formula of whole steps (W) and half steps (H). The formula tells '
             f'you exactly which notes to use. Play each scale on your MPK Mini as you complete it.</p>'
             f'{H2("C Major Scale · " + MAJOR_SCALE)}<p style="margin:0 0 3pt">Starting on C, apply the formula. The first two notes are done for you.</p>'
             f'{note_row(cmaj, 7, MAJOR_SCALE, given=2, final="=C")}'
             f'<div class="fl" style="margin-top:5pt"><b>Play it. Describe what it sounds like in one word:</b><span></span></div>'
             f'{H2("A Natural Minor Scale · " + MINOR_SCALE)}<p style="margin:0 0 3pt">Starting on A, apply the formula. Fill in all 8 notes.</p>'
             f'{note_row(["A"], 7, MINOR_SCALE, given=1, final="=A")}'
             f'<div class="fl" style="margin-top:5pt"><b>Play it. How does it feel different from C Major?</b><span></span></div>'
             f'<div class="box" style="margin:6pt 0"><span class="k">About the pentatonic scale</span><p style="margin:0">A pentatonic scale removes 2 '
             f'notes from the major or minor scale, leaving only the ones that sound most agreeable together. WH = one whole step + one half step '
             f'combined (3 half steps).</p></div>'
             f'{H2("C Major Pentatonic · " + MAJOR_PENT + " (5 notes)")}'
             f'{note_row(["C"], 5, MAJOR_PENT, given=1, final="=C")}'
             f'<div class="fl" style="margin-top:6pt"><b>Which two notes from C Major are missing from C Major Pentatonic?</b><span style="flex:0 0 1in"></span>'
             f'<span style="flex:0 0 1in"></span></div>'
             f'<div class="box inv" style="margin-top:6pt"><p style="margin:0;font-size:11.5pt"><b>When you finish:</b> your teacher will open the '
             f'Logic file. Listen to the chords and bassline, then use the notes from C Major or A Minor to write a short melody on top.</p></div>')
    return page("The Scary Stuff · Part 2: Scales", "PACKET", "Part 2 · Day 1", inner, n, name=True, key="pk2")


def plan_table(filled=None):
    """The owner's planning table (Chord 1-4: Roman numeral, root note); filled = (key, progression) for the worked example."""
    if filled:
        key, prog = filled
        cells = "".join(f'<td style="text-align:center"><span class="small">Roman numeral</span><br><b style="font-family:Oswald;font-size:17pt">{rn}</b>'
                        f'<br><span class="small">Root note</span><br><b style="font-family:Oswald;font-size:15pt;color:{ACCENT}">{by_rn(key, rn)[0]}</b></td>'
                        for rn in prog)
    else:
        cells = "".join('<td style="height:64pt;vertical-align:top"><span>Roman numeral:</span><div class="ln2" style="height:16pt"></div>'
                        '<span style="display:block;margin-top:8pt">Root note:</span><div class="ln2" style="height:16pt"></div></td>' for _ in range(4))
    return f'<table class="plan"><tr>{"".join(f"<th>Chord {i}</th>" for i in range(1, 5))}</tr><tr>{cells}</tr></table>'


def pk4_page(n, watermark=False):
    qs = "".join(f'<div class="fl" style="margin-bottom:7pt"><b>{q}</b><span></span></div>' for q in ANALYSIS)
    rn_row = "".join(f'<b>Chord {i}:</b><span style="flex:0 0 0.8in"></span>' for i in range(1, 5))
    inner = (f'<p style="margin:0 0 4pt">Open the Logic file your teacher provides. A bassline is already written. Before you do anything else, '
             f'answer the questions below.</p>'
             f'{H2("Look at the bassline in the piano roll. Answer these questions")}{qs}'
             f'<p style="margin:2pt 0 6pt">Write the Roman numeral for each chord in the progression:</p>'
             f'<div class="fl" style="gap:6pt">{rn_row}</div>'
             f'<div class="rule" style="margin:6pt 0 4pt"><span class="k">The Rule</span><b>{RULE}</b>'
             f'<p style="margin:3pt 0 0;font-size:10.5pt" class="why">{RULE_SUB}</p></div>'
             f'{H2("Now write your own bassline")}'
             f'<p style="margin:0 0 4pt">Choose a chord progression from your Reference Card. Write the Roman numerals and root notes below, then '
             f'build it in Logic.</p>{plan_table()}'
             f'<p style="margin:7pt 0 2pt"><b>{WHY4}</b></p><div class="lines" style="min-height:52pt"></div>')
    return page("The Scary Stuff · Part 4: Basslines", "PACKET", "Part 4 · Days 2–3", inner, n, watermark, name=True, key="pk4")


def card_table(key):
    rows = "".join(f"<tr><td class='rn'>{rn}</td><td>{t}</td><td><b>{c}</b></td><td class='nt'>{nt}</td><td><i>{f}</i></td></tr>"
                   for rn, t, c, nt, f in CARD[key])
    return (f"<table class='card'><tr><th>RN</th><th style='width:11%'>Type</th><th style='width:17%'>Chord</th><th style='width:17%'>Notes</th>"
            f"<th>Feels like…</th></tr>{rows}</table>")


def scale_cells(notes, n_cells=8, extra=None):
    cells = "".join(f"<td>{x}</td>" for x in notes[:n_cells])
    if extra:
        cells += f"<td class='g'>{extra}</td>"
    return f"<table class='nt' style='margin:2pt 0 6pt'><tr>{cells}</tr></table>"


def remember_box(key="C major"):
    a, b, c = CARD_REMEMBER
    if key == "A minor":
        b = CARD_REMEMBER_MINOR
    return (f'<div class="box acc-b" style="margin-top:8pt"><p style="margin:0;font-size:11.5pt"><b>{a}</b></p>'
            f'<p style="margin:3pt 0 0;font-size:10.5pt"><i>{b}</i></p><p style="margin:3pt 0 0;font-size:10.5pt">{c}</p></div>')


def card_page(n, key, watermark=False):
    sc = key_scale(key)
    extra = "(8)" if key == "C major" else None
    t = tonic(key)
    inner = (f'<h2 class="s" style="font-size:20pt;margin:2pt 0 0">{key.split()[0]} {key.split()[1].title()} — Reference Card</h2>'
             f'<p class="l" style="font-size:12pt;margin-bottom:4pt">The Building Blocks · Unit {UNIT} · Keep this open while you work</p>'
             f'{H2("Scale")}<p style="font-size:12pt;margin:0"><b>{CARD_SCALE_TEXT[key]}</b></p>{scale_cells(sc, 8, extra)}'
             f'{H2("Chords in this Key")}{card_table(key)}{remember_box(key)}'
             f'<div class="cols" style="gap:12pt;margin-top:8pt"><div class="box"><span class="k">For your melody</span><p style="margin:0">Every note '
             f'of the scale above is a safe note. Start or end on <b>{t}</b>, the root of the key.</p></div>'
             f'<div class="box"><span class="k">For your bassline</span><p style="margin:0">The root of each chord is the first note in its Notes '
             f'column: play it when that chord is playing.</p></div></div>'
             f'<p class="small" style="margin-top:6pt">No sharps or flats in this key: every chord is white keys only. '
             f'{"A minor uses the same seven notes: it is the relative minor." if key == "C major" else "C major uses the same seven notes: it is the relative major."}</p>')
    return page(f"Reference Card · {key}", "REFERENCE", "Keep it open", inner, n, watermark, key=f"card{key[0]}")


def methods_page(n, watermark=False):
    boxes = "".join(f'<div class="box" style="margin-bottom:8pt"><div style="display:flex;gap:10pt;align-items:baseline">'
                    f'<b style="font-family:Oswald;font-size:24pt;color:{ACCENT}">{i + 1}</b><div><b style="font-family:Oswald;font-size:15pt;'
                    f'text-transform:uppercase">{a}</b><p style="margin:2pt 0 0;font-size:11.5pt">{b}</p>'
                    f'<p style="margin:2pt 0 0;font-size:10.5pt" class="why">Good when {c}.</p></div></div></div>'
                    for i, (a, b, c) in enumerate(METHODS))
    three = " · ".join(f"{k.split()[0]} {k.split()[1]}: {', '.join(v)}" for k, v in THREE_NOTES.items())
    inner = (f'<p style="font-size:11.5pt;margin:0 0 6pt">There is more than one way to write a melody. Try at least two of these today. Whichever '
             f'you use, the rules are the same: 2 or 4 measures, notes from your scale, start or end on the root.</p>{boxes}'
             f'<div class="cols" style="gap:12pt"><div class="box inv"><span class="k">Stuck?</span><p style="margin:0">Use only 3 notes. {three}. '
             f'Go up, then come back down. That is a melody.</p></div>'
             f'<div class="box acc-b"><span class="k">Remember</span><p style="margin:0;font-size:12pt"><b>{SIMPLE}</b></p></div></div>'
             f'<div class="fl" style="margin-top:10pt"><b>The two methods I tried:</b><span></span></div>'
             f'<div class="fl"><b>The one that worked best for me, and why:</b><span></span></div>')
    return page("Four ways to write a melody", "STUDENT", "Day 1", inner, n, watermark, name=True, key="methods")


def proj2a_page(n):
    rows = "".join(f"<tr><td style='width:24%'><b>{a}</b></td><td>{b}</td></tr>" for a, b in HOWTO2)
    inner = (f'<h2 class="s" style="font-size:20pt;margin:2pt 0 0">Scales &amp; Melody</h2>'
             f'<p class="l" style="font-size:12pt;margin-bottom:4pt">The Building Blocks · Part 2 · Digital Music Production 1 · Unit {UNIT} · Days 1–2</p>'
             f'{H2("Overview")}<p style="font-size:11pt">{OVERVIEW2}</p>'
             f'{H2("What You Are Submitting")}<p style="font-size:11pt">{SUBMIT2}</p>'
             f'{H2("Requirements")}'
             + "".join(f'<div class="chk" style="font-size:10.6pt;margin-bottom:4pt">{checkbox(12)}<span>{r}</span></div>' for r in REQ2) +
             f'{H2("How to Write a Melody")}<table class="t" style="font-size:10.4pt">{rows}</table>')
    return page("Project sheet · melody", "STUDENT", "Melody · 1 of 2", inner, n, name=True, key="proj2a")


def proj2b_page(n):
    inner = (f'{H2("Tips for Success", True)}<ul>' + "".join(f"<li style='font-size:11.5pt;margin-bottom:6pt'>{t}</li>" for t in TIPS2) + '</ul>'
             f'{H2("Submission Checklist")}'
             + "".join(f'<div class="chk" style="font-size:11.5pt;margin-bottom:7pt">{checkbox(13)}<span>{c}</span></div>' for c in CHECK2) +
             f'{H2("Get ready to present")}'
             f'<div class="fl"><b>My key and my root note:</b><span></span></div>'
             f'<div class="fl"><b>My melody goes up, then down? Describe its shape:</b><span></span></div>'
             f'{H2("Notes")}<div class="lines"></div>')
    return page("Project sheet · tips and checklist", "STUDENT", "Melody · 2 of 2", inner, n, name=True, key="proj2b")


def proj4a_page(n):
    inner = (f'<h2 class="s" style="font-size:20pt;margin:2pt 0 0">Basslines</h2>'
             f'<p class="l" style="font-size:12pt;margin-bottom:4pt">The Building Blocks · Part 4 · Digital Music Production 1 · Unit {UNIT} · Days 2–4</p>'
             f'{H2("Overview")}<p style="font-size:11pt">{OVERVIEW4}</p>'
             f'{H2("Requirements")}'
             + "".join(f'<div class="chk" style="font-size:10.8pt;margin-bottom:5pt">{checkbox(12)}<span>{r}</span></div>' for r in REQ4) +
             f'{H2("Your Chord Progression — Plan It Here")}<p style="font-size:10.8pt;margin:0 0 3pt">Before you open Logic, complete this planning table:</p>'
             f'{plan_table()}<p style="margin:7pt 0 2pt;font-size:11pt"><b>{WHY4}</b></p><div class="lines" style="min-height:66pt"></div>')
    return page("Project sheet · basslines", "STUDENT", "Bassline · 1 of 2", inner, n, name=True, key="proj4a")


def proj4b_page(n):
    inner = (f'{H2("Tips for Success", True)}<ul>' + "".join(f"<li style='font-size:11.5pt;margin-bottom:6pt'>{t}</li>" for t in TIPS4) + '</ul>'
             f'{H2("Submission Checklist")}'
             + "".join(f'<div class="chk" style="font-size:11.5pt;margin-bottom:7pt">{checkbox(13)}<span>{c}</span></div>' for c in CHECK4) +
             f'<div class="rule" style="margin-top:6pt"><span class="k">The one rule</span><b>{RULE}</b></div>'
             f'{H2("Notes")}<div class="lines"></div>')
    return page("Project sheet · tips and checklist", "STUDENT", "Bassline · 2 of 2", inner, n, name=True, key="proj4b")


def rubric_page(n, part, watermark=False):
    rub = RUBRIC2 if part == 2 else RUBRIC4
    head = "<tr><th>Category</th>" + "".join(f"<th>{lv}</th>" for lv in LEVELS) + "<th>Score</th></tr>"
    rows = "".join(f"<tr><td class='c'>{escape(c)}</td>" + "".join(f"<td>{d}</td>" for d in ds) + "<td class='sc'></td></tr>" for c, ds in rub)
    total = (f"<tr><td class='c' colspan='5' style='text-align:right;font-size:14pt'>Total</td>"
             f"<td style='font-family:Oswald;font-weight:700;font-size:14pt;text-align:right;vertical-align:bottom'>/ {RUBRIC_TOTAL}</td></tr>")
    name = "Part 2: Scales &amp; Melody" if part == 2 else "Part 4: Basslines"
    prompt = "My melody choices, in one sentence…" if part == 2 else "I chose my progression because…"
    inner = (f'<h2 class="s" style="font-size:18pt;margin:2pt 0 0">{name} — Rubric</h2>'
             f'<p class="l" style="font-size:12pt;margin-bottom:6pt">The Building Blocks · Unit {UNIT} · /{RUBRIC_TOTAL}</p>'
             f'<p style="font-size:11pt">Each category is scored from 2 to 5.</p>'
             f'<table class="rub">{head}{rows}{total}</table>'
             f'<div class="cols" style="margin-top:14pt"><div><h2 class="s" style="margin-top:0">{prompt}</h2>'
             f'<div class="ln2"></div><div class="ln2"></div><div class="ln2"></div></div>'
             f'<div><h2 class="s" style="margin-top:0">Teacher comments</h2><div class="ln2"></div><div class="ln2"></div><div class="ln2"></div></div></div>')
    return page("Rubric · melody" if part == 2 else "Rubric · bassline", "RUBRIC", f"/ {RUBRIC_TOTAL}", inner, n, watermark, name=True,
                key=f"rub{part}")


def vocab_page(n):
    rows = "".join(f"<tr><td style='width:21%;padding:5pt 6pt'><b>{a}</b></td><td style='padding:5pt 6pt'>{b}</td></tr>" for a, b in VOCAB)
    qz = "".join(f'<div class="fl" style="margin-bottom:6pt"><b>{i + 1}. {q}</b><span style="flex:0 0 1.2in"></span></div>' for i, (q, _) in enumerate(VOCAB_QUIZ))
    inner = (f'<p style="font-size:11pt">Read each word and its meaning. These are the words you should own by the end of the unit.</p>'
             f'<table class="t" style="font-size:12pt;line-height:1.45"><tr><th>Word</th><th>Meaning</th></tr>{rows}</table>'
             f'<div class="box" style="margin-top:6pt"><span class="k">Quick check · answers on page {IX["key1"]}</span>{qz}</div>')
    return page("Key vocabulary", "STUDENT", f"{len(VOCAB)} words", inner, n, name=True, key="vocab")


def analysis_answers(key):
    prog = FILE_PROG[key]
    first = by_rn(key, prog[0])
    bass_first = midi_name(file_bass(key)[0][0])
    return [key.split()[0] + " " + key.split()[1], f"{chord_name(first)} ({prog[0]})", bass_first,
            "Yes" if bass_first == first[0] else "No", "It changes to the root of the new chord."], prog


def key1_page(n):
    cs, am, pent = key_scale("C major"), key_scale("A minor"), scale("C", MAJOR_PENT)
    missing = [x for x in cs[:7] if x not in pent]
    vq = " &nbsp;·&nbsp; ".join(f"<b>{i + 1}.</b> {a}" for i, (_, a) in enumerate(VOCAB_QUIZ))

    def an(key):
        ans, prog = analysis_answers(key)
        rows = "".join(f"<tr><td>{q}</td><td><b>{a}</b></td></tr>" for q, a in zip(ANALYSIS, ans))
        rows += f"<tr><td>Roman numerals, chords 1–4</td><td><b>{' – '.join(prog)}</b> ({', '.join(short_name(by_rn(key, r)) for r in prog)})</td></tr>"
        return (f'<div><span class="k">{key} file · bass {" – ".join(midi_name(b) for b, _, _ in file_bass(key))}</span>'
                f'<table class="t" style="font-size:10pt">{rows}</table></div>')
    inner = (f'<div class="box" style="padding:5pt 9pt"><span class="k">Packet Part 2 · scales</span><p style="margin:0;font-size:10.8pt">'
             f'C major: {" ".join(cs)}<br>A natural minor: {" ".join(am)}<br>C major pentatonic: {" ".join(pent)}<br>Missing from the pentatonic: '
             f'<b>{" and ".join(missing)}</b></p></div>'
             f'{H2("Packet Part 4 · bassline analysis")}'
             f'<p style="font-size:10.4pt;margin:0 0 4pt">Answers for the chords-and-bassline file from the setup page (and the Unit 3 worked example). '
             f'If your file uses other chords, the pattern is the same: the bass plays the root of each chord.</p>'
             f'<div class="cols" style="gap:14pt">{an("C major")}{an("A minor")}</div>'
             f'<div class="cols" style="gap:14pt;margin-top:8pt"><div class="box" style="padding:5pt 9pt"><span class="k">Vocabulary · quick check</span>'
             f'<p style="margin:0;font-size:10.4pt">{vq}</p></div>'
             f'<div class="box" style="padding:5pt 9pt"><span class="k">The 3-note melody set</span><p style="margin:0;font-size:10.4pt">'
             + "<br>".join(f"{k}: {', '.join(v)} (the notes of the home chord)" for k, v in THREE_NOTES.items()) +
             '</p></div></div>'
             f'{H2("Listening answers")}<table class="t" style="font-size:10.4pt">'
             f'<tr><td style="width:24%"><b>{SNA[0]}</b> ({SNA[1]})</td><td>The riff is seven notes long with only five different notes '
             f'({", ".join(dict.fromkeys(SNA_RIFF))}), all from E natural minor, and it repeats over and over. Short, repetitive, singable.</td></tr>'
             f'<tr><td><b>{SBM[0]}</b> ({SBM[1]})</td><td>The bass riff starts again on each new chord’s root: {" – ".join(SBM_FACT[2])} in '
             f'{SBM_FACT[0]} = {", ".join(SBM_FACT[3])}. The rule, in a song they know.</td></tr>'
             f'<tr><td><b>Bright or dark?</b></td><td>Most students hear C major as brighter and A natural minor as darker: same seven white keys, a '
             f'different home note.</td></tr></table>'
             f'<p class="small" style="margin-top:8pt">The planning table and the why answer have no single right answer. Check that each root is the '
             f'first note of that chord on the Reference Card, and listen for Roman numeral language in the why.</p>')
    return page("Answer key · packet Parts 2 and 4", "ANSWER KEY", "Teacher", inner, n, ak=True, key="key1")


def mel_table(key):
    rows = []
    for m, (rn, meas) in enumerate(zip(FILE_PROG[key], MELODY[key]), start=1):
        ch = by_rn(key, rn)
        rows.append(f"<tr><td style='text-align:center'><b>{m}</b></td><td style='font-family:Oswald;font-weight:700;text-align:center'>{rn}</td>"
                    f"<td>{chord_name(ch)} ({dash(ch)})</td><td style='font-family:Oswald;font-weight:500;white-space:nowrap'>"
                    f"{' '.join(letter(x) for x, _ in meas)}</td><td>{letter(meas[0][0])}, {letter(meas[-1][0])}</td></tr>")
    return (f"<table class='t tight' style='font-size:9.4pt'><tr><th style='width:10%'>Measure</th><th style='width:9%'>Chord</th><th style='width:28%'>"
            f"Chord notes</th><th>Melody</th><th style='width:24%'>Beats 1 and 3 (chord notes)</th></tr>{''.join(rows)}</table>")


def exmel_page(n):
    inner = (f'<p style="font-size:10.2pt;margin:0 0 4pt">One possible answer for the melody project, over the chords-and-bassline file (one chord '
             f'per measure); the bass plays each chord’s root as a whole note (slide 8 shows both). Same rhythm in every measure: quarter, two '
             f'eighths, half.</p>'
             f'{H2("C major file · starts on C")}<div style="width:5.5in;margin:0 auto">{melroll_svg("C major", bass=False)}</div>{mel_table("C major")}'
             f'{H2("A minor file · starts on A")}<div style="width:5.5in;margin:0 auto">{melroll_svg("A minor", bass=False)}</div>{mel_table("A minor")}'
             f'<p class="small">Every note is a white key from the scale. Beats 1 and 3 land on notes of the chord underneath, which is why it sounds '
             f'like it belongs; the notes between are passing notes from the scale. Shape: up, down, up, down. Copy it, or let students vary it.</p>')
    return page("Example notes · a worked melody", "EXAMPLE", "Teacher", inner, n, ak=True, key="exmel")


def bass_table(key):
    rows = []
    for m, (rn, meas, idea) in enumerate(zip(BASS_PROG[key], BASS[key], BASS_IDEAS), start=1):
        rows.append(f"<tr><td style='text-align:center'><b>{m}</b></td><td style='font-family:Oswald;font-weight:700;text-align:center'>{rn}</td>"
                    f"<td>{by_rn(key, rn)[0]}</td><td style='font-family:Oswald;font-weight:500'>{' '.join(letter(x) for x, _ in meas)}</td>"
                    f"<td>{idea}</td></tr>")
    return (f"<table class='t tight' style='font-size:9.4pt'><tr><th style='width:10%'>Measure</th><th style='width:9%'>Chord</th><th style='width:10%'>"
            f"Root</th><th>Bass notes</th><th style='width:24%'>Rhythm</th></tr>{''.join(rows)}</table>")


def exbass_page(n):
    inner = (f'<p style="font-size:10.2pt;margin:0 0 4pt">One possible answer for the bassline project. Idea A: dotted quarter, eighth, half. '
             f'Idea B: four quarters. The last note of measure 4 walks back toward the first root of the loop.</p>'
             f'{H2("C major · the plan")}{plan_table(("C major", BASS_PROG["C major"]))}'
             f'<p style="margin:5pt 0 2pt;font-size:10.4pt"><b>Why:</b> <i>{BASS_WHY["C major"]}</i></p>'
             f'<div style="width:5.7in;margin:4pt auto 0">{bassroll_svg("C major")}</div>{bass_table("C major")}'
             f'{H2("A minor · the plan")}<p style="margin:0 0 3pt;font-size:10.4pt"><b>{" – ".join(BASS_PROG["A minor"])}</b> '
             f'({", ".join(short_name(by_rn("A minor", r)) for r in BASS_PROG["A minor"])}). <b>Why:</b> <i>{BASS_WHY["A minor"]}</i></p>'
             f'{bass_table("A minor")}'
             f'<p class="small">Checks: one note at a time (no overlaps); the root on the first beat of every chord; two rhythm ideas; every note '
             f'from the scale. A why answer like this one, in Roman numeral language, is what the top level of Harmonic Intention describes.</p>')
    return page("Example notes · a worked bassline", "EXAMPLE", "Teacher", inner, n, ak=True, key="exbass")


def terms_page(n):
    return (f'<section class="page" data-key="terms"><div class="top"><div class="t">Terms of Use<small>Thank you for your purchase!</small></div></div>'
            f'<div class="main"><p>© {YEAR} {STORE}. {BRAND} by {STORE}. All rights reserved. This purchase gives <b>one teacher</b> a license to use this '
            f'resource (the PDF and the intro deck) with their own students.</p>'
            f'<h2 class="s">You may</h2><ul><li>Print and copy pages for your own classroom and students, year after year.</li>'
            f'<li>Present the intro deck to your own classes and edit it for your own teaching.</li>'
            f'<li>Share pages with your own students through a password-protected class site or learning platform.</li>'
            f'<li>Leave the pages for a substitute teacher who teaches your own students.</li></ul>'
            f'<h2 class="s">You may not</h2><ul><li>Share, email or copy this resource for other teachers, a whole school or a district. '
            f'Additional licenses are available on Teachers Pay Teachers.</li>'
            f'<li>Post any part of it on a public website, shared drive or social media.</li>'
            f'<li>Sell, give away or claim any part of it as your own, or edit it into a new product.</li></ul>'
            f'<h2 class="s">Questions?</h2><p>Please use the Q&amp;A tab on the {BRAND} by {STORE} store on Teachers Pay Teachers. '
            f'If something looks wrong, let me know there and I will fix it.</p>'
            f'<h2 class="s">Credits</h2><ul>'
            f'<li>Curriculum, packet, reference cards, project sheets, rubrics, teacher notes and intro deck: original work by {STORE}, {BRAND}.</li>'
            f'<li>Piano-roll drawings, worked examples and layouts: original work, {BRAND}.</li>'
            f'<li>Text fonts: Oswald (The Oswald Project Authors) and Source Sans 3 (Adobe), SIL Open Font License 1.1.</li>'
            f'<li>No audio, video, lyrics or album art are included. Songs are named for reference only (title and artist); play them from your own '
            f'YouTube or Spotify access. Song titles belong to their owners; this resource is not affiliated with or endorsed by the artists or their '
            f'labels.</li>'
            f'<li>Logic Pro is a trademark of Apple Inc., named only to describe compatibility; this resource is not affiliated with or endorsed by '
            f'Apple. MPK Mini, Padlet, YouTube and Spotify are named only as examples of tools you may already use.</li></ul>'
            f'<p class="small">{DISCLOSURE}</p><div style="flex:1"></div></div>{foot(f"Page {n}")}</section>')


def render_page(key, watermark=False):
    n = IX[key]
    f = {"cover": cover_page, "howto": lambda: howto_page(n), "overview": lambda: overview_page(n), "setup": lambda: setup_page(n),
         "logicref": lambda: logicref_page(n), "pacing": lambda: pacing_page(n, watermark), "day12": lambda: day12_page(n),
         "day34": lambda: day34_page(n), "tn1": lambda: tn1_page(n), "tn2": lambda: tn2_page(n), "tn3": lambda: tn3_page(n),
         "tn4": lambda: tn4_page(n), "post1": lambda: post1_page(n), "post2": lambda: post2_page(n), "pk2": lambda: pk2_page(n),
         "pk4": lambda: pk4_page(n, watermark), "cardC": lambda: card_page(n, "C major", watermark), "cardA": lambda: card_page(n, "A minor", watermark),
         "methods": lambda: methods_page(n, watermark), "proj2a": lambda: proj2a_page(n), "proj2b": lambda: proj2b_page(n),
         "rub2": lambda: rubric_page(n, 2, watermark), "proj4a": lambda: proj4a_page(n), "proj4b": lambda: proj4b_page(n),
         "rub4": lambda: rubric_page(n, 4, watermark), "vocab": lambda: vocab_page(n), "key1": lambda: key1_page(n),
         "exmel": lambda: exmel_page(n), "exbass": lambda: exbass_page(n), "terms": lambda: terms_page(n)}
    return f[key]()


# ================================================================ intro deck (.pptx, python-pptx) + PDF copy (LibreOffice)
def slide_day(i):
    return SLIDES[i - 1][0]


def build_deck(img):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.dml import MSO_LINE

    BLACK, WHITE, ACC = RGBColor(0, 0, 0), RGBColor(255, 255, 255), RGBColor.from_string(ACCENT[1:])
    LIGHT = RGBColor.from_string(ACC_LIGHT[1:])
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.author = f"{STORE} · {BRAND}"
    prs.core_properties.title = f"{TITLE} · Intro deck"
    prs.core_properties.comments = DISCLOSURE
    layout = prs.slide_layouts[6]

    def box(sl, x, y, w, h, fillc=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=1.5):
        s = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        if fillc is None:
            s.fill.background()
        else:
            s.fill.solid()
            s.fill.fore_color.rgb = fillc
        if line is None:
            s.line.fill.background()
        else:
            s.line.color.rgb = line
            s.line.width = Pt(lw)
        s.shadow.inherit = False
        return s

    def text(sl, x, y, w, h, runs, size=24, font="Source Sans 3", bold=False, color=BLACK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             spacing=None, caps=False):
        tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.04)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        paras = runs if isinstance(runs, list) else [runs]
        for k, ptxt in enumerate(paras):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = ptxt.upper() if caps else ptxt
            r.font.size = Pt(size)
            r.font.name = font
            r.font.bold = bold
            r.font.color.rgb = color
            if spacing:
                r._r.get_or_add_rPr().set("spc", str(spacing))
        return tb

    def chrome(sl, i, title, day):
        text(sl, 0.6, 0.32, 6, 0.4, BRAND, 13, "Oswald", False, BLACK, spacing=600, caps=True)
        box(sl, 11.33, 0.3, 1.4, 0.42, ACC, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        text(sl, 11.33, 0.3, 1.4, 0.42, f"Day {day}", 15, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, caps=True)
        text(sl, 0.6, 0.72, 12.1, 0.95, title, 40, "Oswald", True, BLACK, caps=True, anchor=MSO_ANCHOR.BOTTOM)
        box(sl, 0.6, 1.72, 12.13, 0.04, BLACK)
        box(sl, 0.6, 7.0, 12.13, 0.012, BLACK)
        text(sl, 0.6, 7.04, 9, 0.3, f"{TITLE} · Intro deck · © {YEAR} {STORE} · {BRAND}", 10, "Source Sans 3", False, RGBColor(0x33, 0x33, 0x33))
        text(sl, 11.2, 7.04, 1.53, 0.3, f"{i}", 10, "Oswald", True, BLACK, align=PP_ALIGN.RIGHT)

    def bullets(sl, items, x=0.6, y=2.05, w=12.1, size=28, gap=1.75):
        for k, it in enumerate(items):
            yy = y + k * (size / 72 * gap)
            box(sl, x, yy + size / 72 * 0.36, 0.16, 0.16, ACC, shape=MSO_SHAPE.OVAL)
            text(sl, x + 0.38, yy, w - 0.4, size / 72 * 1.6, it, size, "Source Sans 3", False, BLACK)

    def pic(sl, key, x, y, w):
        path, (pw, ph) = img[key]
        sl.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))
        return y + w * ph / pw

    def tile(sl, x, y, w, h, head, sub, home=False, hs=40, ss=20):
        box(sl, x, y, w, h, ACC if home else BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        text(sl, x, y + 0.08, w, h * 0.42, head, hs, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(sl, x + 0.15, y + h * 0.45, w - 0.3, h * 0.52, sub, ss, "Source Sans 3", False, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)

    for i, (day, kind, title, body, notes) in enumerate(SLIDES, start=1):
        sl = prs.slides.add_slide(layout)
        items = [b for b in body.split("|") if b]
        if kind == "title":
            text(sl, 0, 0.42, 13.333, 0.5, BRAND, 16, "Oswald", False, BLACK, align=PP_ALIGN.CENTER, spacing=800, caps=True)
            pic(sl, "hero", 4.42, 0.92, 4.5)
            text(sl, 0, 4.3, 13.333, 1.1, title, 72, "Oswald", True, BLACK, align=PP_ALIGN.CENTER, caps=True)
            text(sl, 0, 5.5, 13.333, 0.6, body, 28, "Oswald", False, BLACK, align=PP_ALIGN.CENTER)
            box(sl, 0, 6.55, 13.333, 0.95, BLACK)
            text(sl, 0, 6.68, 13.333, 0.6, f"Digital Music Production · Unit {UNIT} · Day {day} · Basslines and melodies in Logic Pro", 20, "Oswald",
                 False, WHITE, align=PP_ALIGN.CENTER, spacing=200, caps=True)
        else:
            chrome(sl, i, title, day)
        if kind == "listen":
            head, _, sub = items[0].partition(" · ")
            box(sl, 0.6, 2.0, 12.13, 2.0, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            text(sl, 0.9, 2.1, 11.5, 0.7, head, 36, "Oswald", True, WHITE, caps=True)
            text(sl, 0.9, 2.82, 11.5, 0.5, sub, 24, "Oswald", False, LIGHT)
            text(sl, 0.9, 3.38, 11.5, 0.5, items[1], 18, "Source Sans 3", False, WHITE)
            bullets(sl, items[2:], y=4.3, size=30, gap=1.6)
        elif kind == "bullets":
            long_ = max(len(it) for it in items) > 40
            n_ = len(items)
            size = 34 if n_ <= 4 else 31
            bullets(sl, items, size=size, y=2.2, gap=2.2 if n_ <= 4 else 1.95)
        elif kind == "closer":
            bullets(sl, items, w=6.9, size=32, y=2.1, gap=1.75)
            pic(sl, "hero", 7.9, 2.05, 4.83)
        elif kind == "scales":
            for r, (key, f) in enumerate([("C major", MAJOR_SCALE), ("A minor", MINOR_SCALE)]):
                sc, y = key_scale(key), 2.05 + r * 2.15
                text(sl, 0.6, y, 12.1, 0.55, f"{key.split()[0]} {'major' if 'major' in key else 'natural minor'}  ·  {f}", 26, "Oswald", True,
                     ACC if r == 0 else BLACK, caps=True)
                for k, nt in enumerate(sc):
                    x = 0.6 + k * 1.5
                    box(sl, x, y + 0.72, 1.0, 1.0, BLACK if k not in (0, 7) else ACC, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                    text(sl, x, y + 0.72, 1.0, 1.0, nt, 34, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                    if k < 7:
                        text(sl, x + 1.0, y + 0.92, 0.5, 0.6, steps_of(f)[k], 20, "Oswald", True, ACC, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            box(sl, 0.6, 6.25, 12.13, 0.62, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            text(sl, 0.6, 6.25, 12.13, 0.62, "Every white key is a safe note. Start or end on the root: C or A.", 22, "Source Sans 3", True, WHITE,
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        elif kind == "days":
            tags = {1: "melody", 2: "melody + the bass rule", 3: "bassline", 4: "present"}
            for k, d in enumerate(DAYS):
                y = 2.0 + k * 1.2
                box(sl, 0.6, y, 1.7, 0.95, ACC if d["n"] in (1, 2) else BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                text(sl, 0.6, y, 1.7, 0.95, f"Day {d['n']}", 26, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, caps=True)
                text(sl, 2.55, y - 0.02, 10.2, 0.5, d["title"] + "  ·  " + tags[d["n"]], 24, "Oswald", True, BLACK, caps=True)
                text(sl, 2.55, y + 0.48, 10.2, 0.45, DECK_DAY_LINES[k], 18, "Source Sans 3", False, BLACK)
        elif kind == "methods":
            for k, (a, b, c) in enumerate(METHODS):
                col, row = k % 2, k // 2
                x, y = 0.6 + col * 6.17, 1.98 + row * 2.42
                tile(sl, x, y, 5.96, 2.22, f"{k + 1}  ·  {a}", f"Good when {c}.", home=(k == 0), hs=30, ss=20)
            text(sl, 0.6, 6.55, 12.1, 0.4, "Whichever way: 2 or 4 measures, notes from your scale, start or end on the root.", 18, "Source Sans 3",
                 False, BLACK, align=PP_ALIGN.CENTER)
        elif kind == "pic":
            pw_, ph_ = img[body][1]
            yb = pic(sl, body, 0.6, 1.98, min(7.6 if body == "melroll" else 8.2, (6.85 - 1.98) * pw_ / ph_))
            side = (["Blue: your melody, one note at a time", "Black: the bassline already in the file", "Each measure starts on a note of "
                     "its chord", "Same rhythm every measure: easy to remember"] if body == "melroll" else
                    ["Root on beat 1 of every chord", "Two rhythm ideas: A and B", "One note at a time: no overlaps",
                     "The last note walks back home"])
            xs = 8.55 if body == "melroll" else 9.1
            for k, s in enumerate(side):
                yy = 2.1 + k * 1.12
                box(sl, xs, yy + 0.12, 0.16, 0.16, ACC, shape=MSO_SHAPE.OVAL)
                text(sl, xs + 0.32, yy, 12.73 - xs - 0.32, 1.05, s, 20, "Source Sans 3", False, BLACK)
            if body == "bassroll":
                prog = BASS_PROG["C major"]
                text(sl, 0.6, min(yb + 0.25, 6.2), 8.2, 0.6, f"The plan: {' – '.join(prog)} in C major = "
                     f"{', '.join(short_name(by_rn('C major', r)) for r in prog)}", 22, "Oswald", True, ACC, caps=False)
        elif kind == "freq":
            box(sl, 0.6, 2.0, 12.13, 0.06, BLACK)
            text(sl, 0.6, 2.1, 2.0, 0.4, "Low", 18, "Oswald", True, BLACK, caps=True)
            text(sl, 10.73, 2.1, 2.0, 0.4, "High", 18, "Oswald", True, BLACK, caps=True, align=PP_ALIGN.RIGHT)
            text(sl, 3.0, 2.1, 7.33, 0.4, "Frequency  ›", 18, "Oswald", False, BLACK, align=PP_ALIGN.CENTER, caps=True)
            box(sl, 0.6, 2.7, 5.0, 0.9, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            text(sl, 0.6, 2.7, 5.0, 0.9, "Bass lives down here", 24, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, caps=True)
            box(sl, 4.1, 3.8, 8.63, 0.9, ACC, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            text(sl, 4.1, 3.8, 8.63, 0.9, "Melody lives up here", 24, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, caps=True)
            d = box(sl, 4.0, 2.6, 1.7, 2.2, None, BLACK, lw=2.25)
            d.line.dash_style = MSO_LINE.DASH
            text(sl, 1.9, 3.9, 2.0, 0.6, "Overlap = mud", 18, "Oswald", True, BLACK, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE, caps=True)
            bullets(sl, ["Melody: Channel EQ, turn on the low cut (high-pass)", "Bass: roll off the extreme highs (high cut)",
                         "Every instrument gets a frequency home"], y=5.0, size=24, gap=1.45)
        elif kind == "rule":
            box(sl, 0.6, 2.0, 12.13, 2.1, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            text(sl, 0.95, 2.05, 11.43, 2.0, RULE, 36, "Oswald", True, WHITE, anchor=MSO_ANCHOR.MIDDLE)
            text(sl, 0.6, 4.25, 12.1, 0.5, "Everything else is rhythm and decoration on top of this rule.", 22, "Source Sans 3", False, BLACK)
            key = "C major"
            for k, rn in enumerate(FILE_PROG[key]):
                ch = by_rn(key, rn)
                x = 0.6 + k * 3.1
                tile(sl, x, 4.95, 2.85, 1.85, f"{rn} · {short_name(ch)}", f"bass plays {ch[0]}", home=(rn == "I"), hs=30, ss=22)
        elif kind == "plan":
            key, prog = "C major", BASS_PROG["C major"]
            for k, rn in enumerate(prog):
                x = 0.6 + k * 3.1
                box(sl, x, 1.98, 2.85, 0.55, BLACK)
                text(sl, x, 1.98, 2.85, 0.55, f"Chord {k + 1}", 22, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, caps=True)
                box(sl, x, 2.53, 2.85, 2.05, WHITE, BLACK, lw=2)
                text(sl, x, 2.6, 2.85, 0.4, "Roman numeral", 16, "Source Sans 3", False, BLACK, align=PP_ALIGN.CENTER)
                text(sl, x, 2.95, 2.85, 0.75, rn, 40, "Oswald", True, BLACK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                text(sl, x, 3.68, 2.85, 0.35, "Root note", 16, "Source Sans 3", False, BLACK, align=PP_ALIGN.CENTER)
                text(sl, x, 3.98, 2.85, 0.55, by_rn(key, rn)[0], 30, "Oswald", True, ACC, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            box(sl, 0.6, 4.85, 12.13, 1.95, WHITE, ACC, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=2.5)
            text(sl, 0.9, 4.95, 11.5, 0.55, WHY4, 24, "Oswald", True, ACC)
            text(sl, 0.9, 5.55, 11.5, 1.2, f"Example: “{BASS_WHY['C major']}”", 20, "Source Sans 3", False, BLACK)
        elif kind == "check":
            for k, it in enumerate(items):
                col, row = k % 2, k // 2
                x, y = 0.6 + col * 6.15, 2.15 + row * 1.5
                box(sl, x, y + 0.08, 0.42, 0.42, WHITE, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=2.25)
                text(sl, x + 0.6, y, 5.35, 1.4, it, 24, "Source Sans 3", False, BLACK)
        elif kind == "rubric":
            part = int(body)
            rub = RUBRIC2 if part == 2 else RUBRIC4
            for k, ((c, _), short) in enumerate(zip(rub, RUBRIC_SHORT[part])):
                y = 2.1 + k * 1.35
                box(sl, 0.6, y, 4.4, 1.12, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                text(sl, 0.8, y, 4.0, 1.12, c, 22, "Oswald", True, WHITE, anchor=MSO_ANCHOR.MIDDLE, caps=True)
                text(sl, 5.25, y, 6.4, 1.12, "5 = " + short, 22, "Source Sans 3", False, BLACK, anchor=MSO_ANCHOR.MIDDLE)
                text(sl, 11.7, y, 1.03, 1.12, "/ 5", 26, "Oswald", True, ACC, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
            text(sl, 7.0, 6.3, 5.73, 0.5, f"Scored 2 to 5 · Total / {RUBRIC_TOTAL}", 24, "Oswald", True, BLACK, align=PP_ALIGN.RIGHT, caps=True)
        elif kind == "steps":
            for k, it in enumerate(items):
                y = 2.05 + k * 0.95
                box(sl, 0.6, y + 0.05, 0.62, 0.62, BLACK, shape=MSO_SHAPE.OVAL)
                text(sl, 0.6, y + 0.05, 0.62, 0.62, str(k + 1), 24, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                text(sl, 1.45, y - 0.05, 11.2, 0.9, it, 28, "Source Sans 3", False, BLACK, anchor=MSO_ANCHOR.MIDDLE)
        sl.notes_slide.notes_text_frame.text = notes
    prs.save(DECK)


def ttf_fonts():
    """woff2 -> ttf for LibreOffice (deck PDF), plus a fontconfig file that adds them."""
    from fontTools.ttLib import TTFont
    d = os.path.join(BUILD, "ttf")
    os.makedirs(d, exist_ok=True)
    styles = {"oswald-latin-500-normal": ("Oswald", "Regular", 400), "oswald-latin-700-normal": ("Oswald", "Bold", 700),
              "source-sans-3-latin-400-normal": ("Source Sans 3", "Regular", 400), "source-sans-3-latin-700-normal": ("Source Sans 3", "Bold", 700),
              "source-sans-3-latin-400-italic": ("Source Sans 3", "Italic", 400)}
    for stem, (fam, sub, wt) in styles.items():
        out = os.path.join(d, stem + ".ttf")
        if os.path.exists(out):
            continue
        t = TTFont(os.path.join(FONTS, stem + ".woff2"))
        t.flavor = None
        name = t["name"]
        ps = (fam.replace(" ", "") + "-" + sub)
        for rec in list(name.names):
            if rec.nameID in (16, 17, 21, 22):
                name.removeNames(nameID=rec.nameID)
        for nid, val in ((1, fam), (2, sub), (4, f"{fam} {sub}" if sub != "Regular" else fam), (6, ps)):
            name.setName(val, nid, 3, 1, 0x409)
            name.setName(val, nid, 1, 0, 0)
        t["OS/2"].usWeightClass = wt
        sel = 0
        if sub == "Bold":
            sel |= 1 << 5
        if sub == "Italic":
            sel |= 1
        if sub == "Regular":
            sel |= 1 << 6
        t["OS/2"].fsSelection = (t["OS/2"].fsSelection & ~0b1100001) | sel
        t["head"].macStyle = (1 if sub == "Bold" else 0) | (2 if sub == "Italic" else 0)
        t.save(out)
    conf = os.path.join(BUILD, "fonts.conf")
    with open(conf, "w") as fh:
        fh.write(f'<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig><include ignore_missing="yes">/etc/fonts/fonts.conf</include>'
                 f'<dir>{d}</dir><cachedir>{os.path.join(BUILD, "fccache")}</cachedir></fontconfig>')
    return conf


def deck_pdf(src=DECK, out=DECK_PDF, fonts=True):
    env = dict(os.environ, FONTCONFIG_FILE=ttf_fonts()) if fonts else dict(os.environ)
    prof = os.path.join(BUILD, "lo-profile" if fonts else "lo-profile-fallback")
    tmp = os.path.join(BUILD, "lo-out")
    os.makedirs(tmp, exist_ok=True)
    subprocess.run(["soffice", f"-env:UserInstallation=file://{prof}", "--headless", "--convert-to", "pdf", "--outdir", tmp, src],
                   check=True, env=env, capture_output=True)
    shutil.move(os.path.join(tmp, os.path.basename(src).replace(".pptx", ".pdf")), out)


# ================================================================ square images (2000 x 2000)
SQ_CSS = """
.sq { width: 1000px; height: 1000px; position: relative; overflow: hidden; background: linear-gradient(180deg,#fff 0%,#fff 60%,#eeeeee 100%); }
.sq .store { position: absolute; left: 0; right: 0; top: 30px; text-align: center; font-family: 'Oswald'; font-weight: 500; letter-spacing: .32em;
             text-transform: uppercase; font-size: 22px; }
.sq .fillw { position: absolute; left: 0; right: 0; top: 66px; bottom: 70px; display: flex; flex-direction: column; align-items: center; justify-content: space-evenly; }
.sq .hd { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; text-align: center; line-height: .95; white-space: nowrap; }
.sq .sub { font-family: 'Oswald'; font-weight: 300; text-align: center; white-space: nowrap; }
.sq .band { position: absolute; left: 0; right: 0; bottom: 0; background: #000; color: #fff; text-align: center; font-family: 'Oswald';
            font-weight: 500; text-transform: uppercase; letter-spacing: .1em; font-size: 26px; padding: 16px 0 18px; white-space: nowrap; }
.sq .band em { font-style: normal; color: ACCENT; }
.shot { background: #fff; border: 2px solid #000; box-shadow: 8px 8px 0 #000; position: relative; overflow: hidden; }
.shot img { display: block; width: 100%; }
.lbl { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 22px; text-align: center; margin-top: 14px; }
.lbl span { font-weight: 300; text-transform: none; display: block; font-size: 18px; }
.swm { position: absolute; left: 50%; top: 50%; transform: translate(-50%,-50%) rotate(-38deg); font-family: 'Oswald'; font-weight: 700;
       color: rgba(0,0,0,.14); white-space: nowrap; letter-spacing: .06em; }
""".replace("ACCENT", ACCENT)
FONT_PRIME = ('<div style="position:absolute;left:-9999px;font-family:\'Source Sans 3\'">'
              '<span style="font-weight:400">a</span><span style="font-weight:600">a</span><i style="font-weight:400">a</i></div>')


def sq_doc(body):
    return (f'<!doctype html><html><head><meta charset="utf-8"><style>{font_css()}'
            f'html,body{{margin:0;background:#fff}}{SQ_CSS}</style></head><body>{FONT_PRIME}{body}</body></html>')


def sq(inner, band):
    return sq_doc(f'<div class="sq"><div class="store">{BRAND}</div><div class="fillw">{inner}</div><div class="band">{band}</div></div>')


SQ_SUB = f"{DAYS_N}-day unit · scales · melody · the bassline rule · intro slides"
SQ_BAND = f"Grades 9–12 · {DAYS_N} days · {len(PLAN)} pages + {DECK_SLIDES} slides"


def cover_square():
    s = SQ_COVER
    return sq(f'<div style="width:860px">{hero_svg()}</div>'
              f'<div><div class="hd" style="font-size:{s["a"]}px">{COVER_A}</div>'
              f'<div class="hd" style="font-size:{s["b"]}px;margin-top:4px">{escape(COVER_B)}</div>'
              f'<div class="sub" style="font-size:{s["c"]}px;margin-top:10px">{SQ_SUB}</div></div>',
              SQ_BAND.replace(" · ", " <em>·</em> "))


SQ_HEADS = {1: ("Ready to teach Monday", 76, "Teacher notes, four day plans, student pages and two rubrics in one PDF", 28),
            2: ("Intro deck included", 76, f"{DECK_SLIDES} editable slides (.pptx + PDF) with speaker notes on every slide", 28),
            3: ("Everything inside", 72, "Posts, packet, reference cards, project sheets, worked examples", 28)}
SQ_BANDS = {1: "Melody <em>·</em> the bass rule <em>·</em> bassline <em>·</em> present", 2: "Hook <em>·</em> listen <em>·</em> worked example "
            "<em>·</em> steps <em>·</em> rubric", 3: f"{len(PLAN)} pages <em>·</em> US Letter PDF <em>·</em> prints in black &amp; white"}


def sq_head(i):
    h, hs, s, ss = SQ_HEADS[i]
    return f'<div><div class="hd" style="font-size:{hs}px">{h}</div><div class="sub" style="font-size:{ss}px;margin-top:8px">{s}</div></div>'


def preview_square_1(img):
    rot = [-4, 0, 4]
    shots = "".join(f'<div style="position:absolute;left:{x}px;top:{y}px;width:372px;transform:rotate({r}deg);z-index:{z}">'
                    f'<div class="shot"><img src="{s}"><div class="swm" style="font-size:80px">PREVIEW</div></div></div>'
                    for (s, _, _), x, y, r, z in zip(img, (40, 314, 588), (24, 0, 24), rot, (1, 3, 2)))
    labels = "".join(f'<div class="lbl" style="width:300px">{a}<span>{b}</span></div>' for _, a, b in img)
    return sq(f'{sq_head(1)}<div style="position:relative;width:1000px;height:520px">{shots}</div>'
              f'<div style="display:flex;gap:30px">{labels}</div>', SQ_BANDS[1])


def preview_square_2(slides):
    cells = "".join(f'<div class="shot" style="box-shadow:6px 6px 0 #000"><img src="{s}"><div class="swm" style="font-size:52px">PREVIEW</div></div>'
                    for s in slides)
    return sq(f'{sq_head(2)}<div style="display:grid;grid-template-columns:388px 388px;gap:20px 34px">{cells}</div>', SQ_BANDS[2])


def preview_square_3(imgs):
    cells = "".join(f'<div style="width:214px"><div class="shot" style="box-shadow:5px 5px 0 #000"><img src="{s}">'
                    f'<div class="swm" style="font-size:46px">PREVIEW</div></div><div class="lbl" style="font-size:18px;margin-top:9px">{lab_}</div></div>'
                    for s, lab_ in imgs)
    return sq(f'{sq_head(3)}<div style="display:grid;grid-template-columns:repeat(4,214px);gap:22px 24px">{cells}</div>', SQ_BANDS[3])


SHOTS1 = [("pacing", "Pacing + plans", "4 days, minute by minute"), ("methods", "Four ways", "to write a melody"),
          ("exmel", "Worked examples", "melody and bassline")]
SHOTS2 = [1, 7, 8, 13, 15, 17]   # deck slides
SHOTS3 = [("post1", "Melody post"), ("pk2", "Scales"), ("pk4", "Bass rule"), ("cardC", "C major card"),
          ("proj2a", "Melody project"), ("rub2", "Rubric / 15"), ("proj4a", "Bass project"), ("exbass", "Worked bassline")]


# ================================================================ render
def node_env():
    env = dict(os.environ)
    if not env.get("NODE_PATH"):
        env["NODE_PATH"] = subprocess.run(["npm", "root", "-g"], capture_output=True, text=True).stdout.strip()
    return env


def render(jobs):
    path = os.path.join(BUILD, "jobs.json")
    with open(path, "w") as f:
        json.dump(jobs, f)
    subprocess.run(["node", RENDER_JS, path], check=True, env=node_env())


OVERFLOW_JS = r"""
const { chromium } = require('playwright');
const url = require('url');
(async () => {
  const browser = await chromium.launch();
  const page = await browser.newPage();
  await page.goto(url.pathToFileURL(process.argv[2]).href, { waitUntil: 'load' });
  await page.evaluate(() => document.fonts.ready);
  await page.emulateMedia({ media: 'print' });
  const res = await page.evaluate(() => [...document.querySelectorAll('section.page')].map((s, i) => {
    const m = s.querySelector('.main');
    const out = { page: i + 1, key: s.dataset.key, main_over: m ? Math.max(0, m.scrollHeight - m.clientHeight) : 0, cells: [] };
    // anything that sticks out of the page's main area (right or bottom) is clipped or overprinted
    if (m) {
      const mb = m.getBoundingClientRect();
      m.querySelectorAll('*').forEach(el => {
        const r = el.getBoundingClientRect();
        if (r.width && (r.right > mb.right + 1 || r.bottom > mb.bottom + 1 || r.left < mb.left - 1))
          out.cells.push((el.tagName + ': ' + (el.textContent || '').trim()).slice(0, 50));
      });
    }
    // free space at the bottom of the main area (dead space check)
    if (m) {
      let maxb = m.getBoundingClientRect().top;
      m.querySelectorAll('*').forEach(el => { const r = el.getBoundingClientRect(); if (r.height) maxb = Math.max(maxb, r.bottom); });
      out.free = Math.round(m.getBoundingClientRect().bottom - maxb);
    }
    return out;
  }));
  require('fs').writeFileSync(process.argv[3], JSON.stringify(res, null, 1));
  await browser.close();
})().catch(e => { console.error(e); process.exit(1); });
"""


def overflow_scan(html):
    js = os.path.join(BUILD, "overflow.js")
    with open(js, "w") as f:
        f.write(OVERFLOW_JS)
    subprocess.run(["node", js, html, OVERFLOW_JSON], check=True, env=node_env())


def write(name, txt):
    p = os.path.join(BUILD, name)
    with open(p, "w") as f:
        f.write(txt)
    return p


def pdf_png(src, page_no, out, res=110):
    subprocess.run(["pdftoppm", "-png", "-r", str(res), "-f", str(page_no), "-l", str(page_no), "-singlefile", src,
                    os.path.join(BUILD, out)], check=True)


def art_doc(w, h, inner):
    return (f'<!doctype html><html><head><meta charset="utf-8"><style>{font_css()}{CSS}'
            f'html,body{{margin:0;background:#fff}}</style></head>'
            f'<body><div style="width:{w}px;height:{h}px;overflow:hidden;background:#fff">{inner}</div></body></html>')


def deck_art():
    """PNG art for the deck (original drawings rendered with the brand fonts)."""
    svgs = {"hero": (1800, hero_svg()), "melroll": (1900, melroll_svg("C major")), "bassroll": (1900, bassroll_svg("C major"))}
    arts = {}
    for k, (w, svg) in svgs.items():   # exact height from the drawing's own viewBox, so nothing is cropped and nothing is padded
        vw, vh = map(float, re.search(r'viewBox="0 0 ([\d.]+) ([\d.]+)"', svg).groups())
        arts[k] = (w, int(round(w * vh / vw)) + 2, f'<div style="width:{w}px">{svg}</div>')
    jobs, out = [], {}
    for k, (w, h, inner) in arts.items():
        p = os.path.join(BUILD, f"deck-{k}.png")
        jobs.append({"html": write(f"deck-{k}.html", art_doc(w, h, inner)), "png": p, "width": w, "height": h, "scale": 1})
        out[k] = (p, (w, h))
    return jobs, out


def tpt_json():
    return {
        "price": PRICE, "license_price": PRICE,
        "tax_code": "Other Digital Goods - No Physical Media",
        "grades": ["9th-grade", "10th-grade", "11th-grade", "12th-grade"],
        "subjects": ["Music"],
        "tags": ["Autumn"],
        "formats": ["PDF"],
        "pages": len(PLAN) + DECK_SLIDES,
        "answer_key": "Included",
        "zip_contents": [os.path.basename(PDF), os.path.basename(DECK), os.path.basename(DECK_PDF)],
        "files": {"product": os.path.basename(ZIP), "preview": os.path.basename(PREVIEW_PDF), "thumb1": "cover.png",
                  "thumb2": "preview-1.png", "thumb3": "preview-2.png", "thumb4": "preview-3.png"},
    }


def make_zip():
    if os.path.exists(ZIP):
        os.remove(ZIP)
    with zipfile.ZipFile(ZIP, "w", zipfile.ZIP_DEFLATED) as z:
        for p in (PDF, DECK, DECK_PDF):
            zi = zipfile.ZipInfo(f"{ZIP_DIR}/{os.path.basename(p)}", date_time=(2026, 10, 14, 12, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            with open(p, "rb") as fh:
                z.writestr(zi, fh.read())


def build():
    os.makedirs(BUILD, exist_ok=True)
    problems = verify_theory() + verify_data() + verify_labels()
    if problems:
        print("\n".join(problems))
        sys.exit("data check failed; not rendering")
    full = write("resource.html", html_doc([render_page(p) for p in PLAN], TITLE))
    prev = write("preview.html", html_doc([render_page(p, watermark=True) for p in PREVIEW_PICKS], "Preview"))
    art_jobs, art = deck_art()
    render([{"html": full, "pdf": PDF}, {"html": prev, "pdf": PREVIEW_PDF}] + art_jobs)
    overflow_scan(full)
    build_deck(art)
    deck_pdf()
    make_zip()
    res = 45 if LOW_RES else 120
    for key, _, _ in SHOTS1:
        pdf_png(PDF, IX[key], key, res)
    for key, _ in SHOTS3:
        pdf_png(PDF, IX[key], key, 40 if LOW_RES else 80)
    for s in SHOTS2:
        pdf_png(DECK_PDF, s, f"slide{s}", 30 if LOW_RES else 70)
    sqs = [write("cover.html", cover_square()), write("p1.html", preview_square_1([(f"{k}.png", a, b) for k, a, b in SHOTS1])),
           write("p2.html", preview_square_2([f"slide{s}.png" for s in SHOTS2])),
           write("p3.html", preview_square_3([(f"{k}.png", lab_) for k, lab_ in SHOTS3]))]
    render([{"html": h, "png": o, "width": 1000, "height": 1000, "scale": 1 if LOW_RES else 2} for h, o in zip(sqs, [COVER_PNG] + PREVIEW_PNGS)])
    with open(TPT_JSON, "w") as f:
        json.dump(tpt_json(), f, indent=2, ensure_ascii=False)
        f.write("\n")
    print(f"built {PDF} ({len(PLAN)} pages), {DECK} + PDF ({DECK_SLIDES} slides), {ZIP}, preview, cover + 3 previews, tpt.json"
          + (" [LOW RES]" if LOW_RES else ""))


# ================================================================ checks
def check_line(key, prog, measures, what):
    """Shared checks for a melody or bassline over one chord per measure: 4/4 measures, monophonic, every note in the key's scale."""
    errs, sc = [], set(key_scale(key))
    if len(measures) != len(prog):
        errs.append(f"{what} {key}: {len(measures)} measures for {len(prog)} chords")
    for m, meas in enumerate(measures, start=1):
        if abs(sum(d for _, d in meas) - 4) > 1e-9:
            errs.append(f"{what} {key} m{m}: beats add to {sum(d for _, d in meas)}, not 4")
        for nm, _ in meas:
            if letter(nm) not in sc:
                errs.append(f"{what} {key} m{m}: {nm} is not in the {key} scale")
            if len(letter(nm)) > 1:
                errs.append(f"{what} {key} m{m}: {nm} is not a white key")
    notes = flat(measures)
    for (a, sa, da), (b, sb, db) in zip(notes, notes[1:]):
        if sa + da > sb + 1e-9:
            errs.append(f"{what} {key}: notes overlap at beat {sb} (not monophonic)")
    return errs


def verify_theory():
    """Music theory, computed: scales, key signatures, the owner's cards, melodies and basslines against their chords, song facts, answers."""
    errs = []
    for f in [MAJOR_SCALE, MINOR_SCALE, MAJOR_PENT]:
        if sum(STEP[s] for s in steps_of(f)) != 12:
            errs.append(f"scale formula {f} does not add up to an octave")
    exp = {"C major": "C D E F G A B C", "A minor": "A B C D E F G A"}
    for key, want in exp.items():
        sc = key_scale(key)
        if " ".join(sc) != want:
            errs.append(f"{key} scale is {' '.join(sc)}, expected {want}")
        if accidentals(sc):
            errs.append(f"{key}: key signature should have no sharps or flats, found {accidentals(sc)}")
    if set(key_scale("C major")[:7]) != set(key_scale("A minor")[:7]):
        errs.append("C major and A minor must share the same seven notes (relative keys)")
    if " ".join(scale("C", MAJOR_PENT)) != "C D E G A C":
        errs.append(f"C major pentatonic is {' '.join(scale('C', MAJOR_PENT))}")
    if [x for x in key_scale("C major")[:7] if x not in scale("C", MAJOR_PENT)] != ["F", "B"]:
        errs.append("pentatonic: the missing notes must be F and B")
    for key, rows in CARD.items():
        comp = key_chords(key)
        for (rn, typ, name, notes, feel), (crn, cq, cn) in zip(rows, comp):
            if rn != crn or typ != cq or notes != dash(cn) or name != chord_name(cn):
                errs.append(f"{key} card row {rn} {typ} {name} {notes} does not match computed {crn} {cq} {dash(cn)}")
            if chord(cn[0], FORMULA[cq]) != cn:
                errs.append(f"{key} {rn}: formula {FORMULA[cq]} does not give {cn}")
    # the chords-and-bassline file: bass = whole-note root of every chord, in the bass register
    for key, prog in FILE_PROG.items():
        for (b, st, d), rn in zip(file_bass(key), prog):
            if midi_name(b) != by_rn(key, rn)[0] or d != 4 or not 28 <= b <= 40:
                errs.append(f"file bass {key} {rn}: {midi_name(b)} is not the root {by_rn(key, rn)[0]} as a low whole note")
    # worked melodies: in key, white keys, 4/4, monophonic, 2 or 4 measures, starts or ends on the root, strong beats on chord notes, a shape
    for key, meas in MELODY.items():
        prog = FILE_PROG[key]
        errs += check_line(key, prog, meas, "melody")
        if len(meas) not in (2, 4):
            errs.append(f"melody {key}: must be 2 or 4 measures")
        t = tonic(key)
        if letter(meas[0][0][0]) != t and letter(meas[-1][-1][0]) != t:
            errs.append(f"melody {key}: neither starts nor ends on the root {t}")
        for m, (rn, ms) in enumerate(zip(prog, meas), start=1):
            ch, beat = by_rn(key, rn), 0.0
            for nm, d in ms:
                if beat in (0.0, 2.0) and letter(nm) not in ch:
                    errs.append(f"melody {key} m{m}: {nm} on beat {beat + 1:g} is not in the {rn} chord {ch}")
                beat += d
        ns = [n for n, _, _ in flat(meas)]
        if not any(b > a for a, b in zip(ns, ns[1:])) or not any(b < a for a, b in zip(ns, ns[1:])):
            errs.append(f"melody {key}: no shape (needs to go up and come down)")
        if min(ns) < 60:
            errs.append(f"melody {key}: dips into the bass register (shift up an octave)")
        if len({tuple(d for _, d in m) for m in meas}) != 1:
            errs.append(f"melody {key}: the worked example keeps one rhythm in every measure")
    # worked basslines: in key, monophonic, root on beat 1 of every chord, at least 2 distinct rhythmic ideas, in the bass register
    for key, meas in BASS.items():
        prog = BASS_PROG[key]
        errs += check_line(key, prog, meas, "bass")
        for m, (rn, ms) in enumerate(zip(prog, meas), start=1):
            if letter(ms[0][0]) != by_rn(key, rn)[0]:
                errs.append(f"bass {key} m{m}: first note {ms[0][0]} is not the root of {rn} ({by_rn(key, rn)[0]})")
        if len({tuple(d for _, d in m) for m in meas}) < 2:
            errs.append(f"bass {key}: needs at least 2 distinct rhythmic ideas")
        if any(not 24 <= midi(n) <= 43 for m in meas for n, _ in m):
            errs.append(f"bass {key}: a note is outside the bass register")
        if min(d for m in meas for _, d in m) < 0.5:
            errs.append(f"bass {key}: busier than eighth notes")
        ideas = {"Idea A": (1.5, .5, 2), "Idea B": (1, 1, 1, 1)}
        for ms, idea in zip(meas, BASS_IDEAS):
            if tuple(d for _, d in ms) != ideas[idea.split(" +")[0]]:
                errs.append(f"bass {key}: rhythm label {idea} does not match the notes")
    # cover art bass under I-vi-IV-V: roots on every downbeat, in key, monophonic
    errs += check_line("C major", FILE_PROG["C major"], HERO_BASS, "cover bass")
    for ms, rn in zip(HERO_BASS, FILE_PROG["C major"]):
        if letter(ms[0][0]) != by_rn("C major", rn)[0]:
            errs.append(f"cover bass: {ms[0][0]} is not the root of {rn}")
    # the 3-note variation = the notes of the home chord, from the scale
    for key, notes in THREE_NOTES.items():
        if notes != by_rn(key, "I" if "major" in key else "i"):
            errs.append(f"3-note set {key} {notes} is not the home chord")
    # plan + why: every Roman numeral named in a why sentence is in that progression
    for key, why in BASS_WHY.items():
        for rn in re.findall(r"\b(?:I|II|III|IV|V|VI|VII|i|ii|iii|iv|v|vi|vii)\b", why.replace("I start", "_ start")):
            if rn not in BASS_PROG[key]:
                errs.append(f"why {key}: names {rn}, which is not in the progression {BASS_PROG[key]}")
    # songs
    sna_scale = scale("E", MINOR_SCALE)
    if not set(SNA_RIFF) <= set(sna_scale) or len(SNA_RIFF) != 7 or len(set(SNA_RIFF)) != 5:
        errs.append("Seven Nation Army riff: must be 7 notes, 5 different, all in E natural minor")
    key_, ton, rns, names = SBM_FACT
    got = [short_name(next(n for r, q, n in diatonic(scale(ton, MAJOR_SCALE)) if r == rn)) for rn in rns]
    if got != names:
        errs.append(f"{SBM[0]}: {rns} in {key_} gives {got}, not {names}")
    # answer key: analysis answers computed from the file
    for key in FILE_PROG:
        ans, prog = analysis_answers(key)
        if ans[3] != "Yes" or ans[2] != tonic(key):
            errs.append(f"analysis key {key}: the file's first bass note must be the root {tonic(key)}")
    return errs


def verify_data():
    errs = []
    counts = dict(req2=(REQ2, 10), check2=(CHECK2, 8), tips2=(TIPS2, 3), howto2=(HOWTO2, 5), req4=(REQ4, 9), check4=(CHECK4, 8),
                  tips4=(TIPS4, 4), struggles=(STRUGGLES, 8), vocab=(VOCAB, 12), quiz=(VOCAB_QUIZ, 4), methods=(METHODS, 4),
                  analysis=(ANALYSIS, 5), file2=(FILE2, 5))
    for nm, (lst, nexp) in counts.items():
        if len(lst) != nexp:
            errs.append(f"{nm}: {len(lst)} items, expected {nexp}")
    vw = {w.lower() for w, _ in VOCAB}
    for _, a in VOCAB_QUIZ:
        if a.lower() not in vw:
            errs.append(f"vocab quiz answer '{a}' is not a vocabulary word")
    for part, rub in ((2, RUBRIC2), (4, RUBRIC4)):
        if len(rub) != 3 or any(len(dd) != len(LEVELS) for _, dd in rub) or 5 * len(rub) != RUBRIC_TOTAL or len(RUBRIC_SHORT[part]) != 3:
            errs.append(f"rubric part {part}: must be the owner's 3 categories x levels 2-5 = /15")
    # no new rubric requirements: the requirement and checklist lists keep the owner's items (only the LMS name changes)
    for lst in (REQ2, CHECK2, REQ4, CHECK4):
        for it in lst:
            if "Google" in it:
                errs.append("an LMS name was left in a list")
    if not any("Download the provided Logic file" in r for r in REQ2) or not any("I to V feels resolved" in s[2] for s in POST2_STEPS):
        errs.append("owner wording changed: keep 'Download the provided Logic file' and 'I to V feels resolved' (flagged to the owner)")
    # pacing: 4 days, blocks 0-53, every deck slide in exactly one block of the day the slide says
    if len(DAYS) != DAYS_N or [d["n"] for d in DAYS] != list(range(1, DAYS_N + 1)) or len(DECK_DAY_LINES) != DAYS_N:
        errs.append(f"pacing must be {DAYS_N} days (and {DAYS_N} on the deck)")
    seen = {}
    for d in DAYS:
        blocks = [tuple(map(int, x.split("–"))) for x, _, _, _ in d["blocks"]]
        if blocks[0][0] != 0 or blocks[-1][1] != MINUTES or any(blocks[i][1] != blocks[i + 1][0] for i in range(len(blocks) - 1)):
            errs.append(f"Day {d['n']} minute blocks do not run 0-{MINUTES} without gaps")
        if not day_slides(d):
            errs.append(f"Day {d['n']} has no deck slide")
        for _, h, txt, sl in d["blocks"]:
            nameds = set()
            for a_, b_ in re.findall(r"[Ss]lides? (\d+)(?:–(\d+))?", txt):
                nameds |= set(range(int(a_), int(b_ or a_) + 1))
            if not sl:
                if nameds:
                    errs.append(f"Day {d['n']} '{h}' names a slide but has no slide slot")
                continue
            if nameds != set(range(sl[0], sl[1] + 1)):
                errs.append(f"Day {d['n']} '{h}': text names slides {sorted(nameds)} but the slot is {sl}")
            for s in range(sl[0], sl[1] + 1):
                if s in seen:
                    errs.append(f"slide {s} is slotted twice (Day {seen[s]} and Day {d['n']})")
                seen[s] = d["n"]
                if not 1 <= s <= DECK_SLIDES:
                    errs.append(f"Day {d['n']} names slide {s}, which does not exist")
                elif slide_day(s) != d["n"]:
                    errs.append(f"slide {s} says Day {slide_day(s)} but the Day {d['n']} plan uses it")
    for i in range(1, DECK_SLIDES + 1):
        if i not in seen:
            errs.append(f"deck slide {i} has no day slot in the day plans")
    if [s[0] for s in SLIDES] != sorted(s[0] for s in SLIDES):
        errs.append("deck slides must run in day order")
    if len(SLIDES) < 12 or any(len(s[4].strip()) < 80 for s in SLIDES):
        errs.append("deck needs 12+ slides, each with real speaker notes")
    kinds = [k for _, k, *_ in SLIDES]
    for need in ("title", "bullets", "listen", "pic", "steps", "check", "rubric"):
        if need not in kinds:
            errs.append(f"deck pattern: no '{need}' slide")
    if sorted([b for _, k, _, b, _ in SLIDES if k == "pic"] + ["hero"]) != sorted(DECK_PICS):
        errs.append("deck pictures must be hero + melroll + bassroll")
    # the Day 1 post and the deck's Day 1 demo agree with the day plan on the file (Unit 3 project or the teacher's file)
    if "Unit 3" not in POST1_STEPS[1][1] or "No Unit 3 project" not in POST1_STEPS[1][1]:
        errs.append("Day 1 post: open the Unit 3 project, with a fallback")
    if "Unit 3" not in DAYS[0]["blocks"][5][2] or "your file" not in DAYS[0]["blocks"][5][2]:
        errs.append("Day 1 plan: the Write block must name the Unit 3 project and your file")
    if len(PLAN) != len(set(PLAN)):
        errs.append("plan has a repeated page")
    if len(LISTING_TITLE) > 80 or "Logic Pro" not in LISTING_TITLE[:30]:
        errs.append(f"listing title is {len(LISTING_TITLE)} chars, or 'Logic Pro' is not in the first 30")
    # cover and square lines fit their width on one line (Oswald 700 metrics; light weights are narrower)
    for label, txt, sz in (("cover A", COVER_A, PDF_COVER["a"]), ("cover B", COVER_B, PDF_COVER["b"])):
        if text_w(txt.upper(), sz) * 1.01 > PDF_COVER["room"]:
            errs.append(f"PDF {label} '{txt}' is {text_w(txt.upper(), sz):.0f}pt wide, room {PDF_COVER['room']:.0f}")
    if text_w(COVER_C, PDF_COVER["c"]) > PDF_COVER["room"]:
        errs.append("PDF cover subtitle too wide")
    if text_w(COVER_BAND.upper(), PDF_COVER["band"]) * 1.12 > PDF_COVER["room"]:
        errs.append("PDF cover band too wide")
    for txt, sz in ((COVER_A.upper(), SQ_COVER["a"]), (COVER_B.upper(), SQ_COVER["b"]), (SQ_SUB, SQ_COVER["c"])):
        if text_w(txt, sz) * 1.01 > SQ_COVER["room"]:
            errs.append(f"square cover line '{txt}' is {text_w(txt, sz):.0f}px wide, room {SQ_COVER['room']}")
    if text_w(SQ_BAND.upper(), 26) * 1.12 > 960:
        errs.append("square cover band too wide")
    if SQ_COVER["b"] * 0.75 < 100:
        errs.append("square cover big line under 10% of the image height")
    for i, (h, hs, s, ss) in SQ_HEADS.items():
        if text_w(h.upper(), hs) * 1.01 > 960 or text_w(s, ss) > 960:
            errs.append(f"preview {i} headline/subtitle too wide")
    return errs


def pdf_fonts(path):
    out = subprocess.run(["pdffonts", path], capture_output=True, text=True, check=True).stdout.splitlines()[2:]
    return [(ln.split()[0], ln.split()[-5]) for ln in out]


def pdf_text(path):
    return subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True, check=True).stdout


def deck_texts():
    from pptx import Presentation
    prs = Presentation(DECK)
    slides, notes = [], []
    for s in prs.slides:
        slides.append(" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame))
        notes.append(s.notes_slide.notes_text_frame.text if s.has_notes_slide else "")
    media = [p for p in zipfile.ZipFile(DECK).namelist() if p.startswith("ppt/media/")]
    return prs, slides, notes, media


def deck_overflow(prs):
    """Every shape stays inside the slide; Oswald text fits its box (lines available x width); no two text boxes overlap."""
    errs = []
    W, H = prs.slide_width, prs.slide_height
    for i, s in enumerate(prs.slides, start=1):
        boxes = []
        for sh in s.shapes:
            if sh.left < 0 or sh.top < 0 or sh.left + sh.width > W + 10 or sh.top + sh.height > H + 10:
                errs.append(f"deck slide {i}: a shape runs off the slide")
            if sh.has_text_frame and sh.text_frame.text:
                r = sh.text_frame.paragraphs[0].runs
                if r and r[0].font.name == "Oswald" and r[0].font.size:
                    widest = max((text_w(p.text.upper() if p.text.upper() == p.text else p.text, r[0].font.size.pt)
                                  for p in sh.text_frame.paragraphs), default=0)
                    room = sh.width / 12700 - 6
                    lines_avail = int((sh.height / 12700) // (r[0].font.size.pt * 1.15)) or 1
                    if widest > room * lines_avail:
                        errs.append(f"deck slide {i}: '{sh.text_frame.text[:30]}' ({widest:.0f}pt) does not fit {room:.0f}pt x {lines_avail} line(s)")
                if sh.shape_type == 17:   # text boxes (not text on filled shapes)
                    boxes.append((sh.left, sh.top, sh.left + sh.width, sh.top + sh.height, sh.text_frame.text[:20]))
        pics = [(sh.left, sh.top, sh.left + sh.width, sh.top + sh.height) for sh in s.shapes if sh.shape_type == 13]
        for P in pics:
            for A in boxes:
                if min(A[2], P[2]) - max(A[0], P[0]) > 12700 * 4 and min(A[3], P[3]) - max(A[1], P[1]) > 12700 * 4:
                    errs.append(f"deck slide {i}: text '{A[4]}' overlaps a picture")
        for a in range(len(boxes)):
            for b in range(a + 1, len(boxes)):
                A, B = boxes[a], boxes[b]
                ox = min(A[2], B[2]) - max(A[0], B[0])
                oy = min(A[3], B[3]) - max(A[1], B[1])
                if ox > 12700 * 4 and oy > 12700 * 4:
                    inside = (A[0] >= B[0] and A[2] <= B[2] and A[1] >= B[1] and A[3] <= B[3]) or \
                             (B[0] >= A[0] and B[2] <= A[2] and B[1] >= A[1] and B[3] <= A[3])
                    if not inside:
                        errs.append(f"deck slide {i}: text boxes '{A[4]}' and '{B[4]}' overlap")
    return errs


def verify_files():
    errs = []
    from pypdf import PdfReader
    from PIL import Image
    for p in [PDF, DECK, DECK_PDF, ZIP, PREVIEW_PDF, COVER_PNG, *PREVIEW_PNGS, UPLOAD, TPT_JSON, OVERFLOW_JSON]:
        if not os.path.exists(p):
            errs.append(f"missing {os.path.basename(p)}")
    if errs:
        return errs
    for path, n_exp in ((PDF, len(PLAN)), (PREVIEW_PDF, len(PREVIEW_PICKS))):
        rd = PdfReader(path)
        if len(rd.pages) != n_exp:
            errs.append(f"{os.path.basename(path)}: {len(rd.pages)} pages, expected {n_exp}")
        for i, pg_ in enumerate(rd.pages):
            w, h = float(pg_.mediabox.width), float(pg_.mediabox.height)
            if abs(w - 612) > 1 or abs(h - 792) > 1:
                errs.append(f"{os.path.basename(path)} p{i + 1}: {w:.0f}x{h:.0f}pt, not US Letter")
    if len(PdfReader(PREVIEW_PDF).pages) != 4:
        errs.append("preview PDF must be 4 pages")
    # rendered-layout scan: nothing clipped at the bottom of a page, no nowrap cell or line wider than its box, no big dead space
    for o in json.load(open(OVERFLOW_JSON)):
        if o["main_over"] > 1:
            errs.append(f"page {o['page']} ({o['key']}): content overflows by {o['main_over']}px (clipped)")
        for c in o["cells"]:
            errs.append(f"page {o['page']} ({o['key']}): text wider than its box: '{c}'")
        if o["key"] not in ("cover", "terms") and o.get("free", 0) > 200:
            errs.append(f"page {o['page']} ({o['key']}): {o['free']}px empty at the bottom (dead space)")
    drd = PdfReader(DECK_PDF)
    if len(drd.pages) != DECK_SLIDES:
        errs.append(f"deck PDF: {len(drd.pages)} pages, expected {DECK_SLIDES}")
    for i, pg_ in enumerate(drd.pages):
        if abs(float(pg_.mediabox.width) / float(pg_.mediabox.height) - 16 / 9) > 0.01:
            errs.append(f"deck PDF p{i + 1}: not 16:9")
    for path in (PDF, PREVIEW_PDF, DECK_PDF):
        fonts = pdf_fonts(path)
        for name, emb in fonts:
            if emb != "yes":
                errs.append(f"{os.path.basename(path)}: font {name} not embedded")
            if not re.search(r"oswald|sourcesans", name.lower().replace(" ", "")):
                errs.append(f"{os.path.basename(path)}: unexpected font {name} (fallback?)")
        names = " ".join(n for n, _ in fonts).lower()
        for fam in ("oswald", "sourcesans3"):
            if fam not in names:
                errs.append(f"{os.path.basename(path)}: {fam} missing")
    prs, stext, notes, media = deck_texts()
    if len(prs.slides) != DECK_SLIDES:
        errs.append(f"deck has {len(prs.slides)} slides, expected {DECK_SLIDES}")
    for i, nt in enumerate(notes):
        if len(nt.strip()) < 80:
            errs.append(f"deck slide {i + 1}: speaker notes missing or too short")
        if not re.search(rf"\bDay {slide_day(i + 1)}\b", stext[i], re.I):
            errs.append(f"deck slide {i + 1}: no 'Day {slide_day(i + 1)}' slot shown on the slide")
    errs += deck_overflow(prs)
    if any(not m.lower().endswith(".png") for m in media) or len(media) != len(DECK_PICS):
        errs.append(f"deck media should be exactly the {len(DECK_PICS)} drawn PNGs, found {media}")
    if abs(prs.slide_width / prs.slide_height - 16 / 9) > 0.01:
        errs.append("deck is not 16:9")
    alldeck = " ".join(stext + notes)
    for t, a in (SNA, SBM):
        if t.lower() not in alldeck.lower() or a.lower() not in alldeck.lower():
            errs.append(f"deck must name {t} by {a}")
    for i, (_, k, _, body, _) in enumerate(SLIDES):
        if k == "listen" and any(t in body for t, _ in (SNA, SBM)) and "YouTube or Spotify" not in stext[i]:
            errs.append(f"deck slide {i + 1}: a song slide must say YouTube or Spotify")
    if "YouTube or Spotify" not in notes[0]:
        errs.append("deck slide 1 notes must say YouTube or Spotify")
    names = zipfile.ZipFile(ZIP).namelist()
    for p in (PDF, DECK, DECK_PDF):
        n = f"{ZIP_DIR}/{os.path.basename(p)}"
        if n not in names:
            errs.append(f"zip is missing {os.path.basename(p)}")
        elif zipfile.ZipFile(ZIP).read(n) != open(p, "rb").read():
            errs.append(f"zip copy of {os.path.basename(p)} is stale; rebuild")
    if len(names) != 3:
        errs.append(f"zip should hold exactly 3 files, has {names}")
    rd = PdfReader(PDF)
    pages = [subprocess.run(["pdftotext", "-layout", "-f", str(i + 1), "-l", str(i + 1), PDF, "-"], capture_output=True, text=True,
                            check=True).stdout for i in range(len(rd.pages))]
    flat_ = [re.sub(r"\s+", " ", t) for t in pages]
    raw0 = subprocess.run(["pdftotext", "-f", "1", "-l", "1", PDF, "-"], capture_output=True, text=True, check=True).stdout
    lines0 = [re.sub(r"\s+", " ", ln.strip()) for ln in raw0.splitlines() if ln.strip()]
    for need in (COVER_A.upper(), COVER_B.upper(), COVER_C):
        if need not in lines0:
            errs.append(f"PDF cover: '{need}' is not on a line of its own (orphaned word?)")
    for i, t in enumerate(flat_):
        if i > 0 and "Brian Lotze · Hudson Beat" not in t:
            errs.append(f"page {i + 1}: footer missing")
        if i > 0 and not re.search(rf"Page {i + 1}\s*$", t.strip()):
            errs.append(f"page {i + 1}: footer page number wrong (overflow?)")
    for m in re.finditer(r"\b(?:pages?|pp?\.) (\d+)(?:–(\d+))?", " ".join(flat_[1:])):
        for g in m.groups():
            if g and not 1 <= int(g) <= len(PLAN):
                errs.append(f"cross-reference to page {g} does not exist")
    for m in re.finditer(r"\bslides? (\d+)(?:–(\d+))?", " ".join(flat_[1:])):
        for g in m.groups():
            if g and not 1 <= int(g) <= DECK_SLIDES:
                errs.append(f"PDF names deck slide {g}, which does not exist")
    for pk, rub in (("rub2", RUBRIC2), ("rub4", RUBRIC4)):
        t = flat_[IX[pk] - 1]
        if f"/ {RUBRIC_TOTAL}" not in t:
            errs.append(f"{pk}: /{RUBRIC_TOTAL} total missing")
        for c, ds in rub:
            for d in ds:
                if not all(w in t for w in re.findall(r"[A-Za-z]{6,}", d)):
                    errs.append(f"{pk}: rubric text missing: {d[:50]}")
    for pk, lst in (("proj2a", REQ2), ("proj2b", CHECK2), ("proj4a", REQ4), ("proj4b", CHECK4)):
        t = flat_[IX[pk] - 1]
        for it in lst:
            if not all(w in t for w in re.findall(r"[A-Za-z]{6,}", it)):
                errs.append(f"{pk}: owner item missing: {it[:50]}")
    for w, _ in VOCAB:
        if not all(x in flat_[IX["vocab"] - 1] for x in w.split()):
            errs.append(f"vocab page missing {w}")
    for key, pk in (("C major", "cardC"), ("A minor", "cardA")):
        t = flat_[IX[pk] - 1]
        for rn, typ, name, notes_, feel in CARD[key]:
            if notes_.replace(" ", "") not in t.replace(" ", "") or feel.split(".")[0] not in t:
                errs.append(f"{key} card page missing {rn} {notes_}")
    if re.search(r"\b(?:I|V|vi) chord", flat_[IX["cardA"] - 1]):
        errs.append("A minor card: reminder must use i / v / VII / VI, not I / V / vi")
    if RULE not in flat_[IX["pk4"] - 1] or RULE not in flat_[IX["post2"] - 1] or RULE not in flat_[IX["tn2"] - 1]:
        errs.append("the bassline rule must be printed on the packet, the post and the teacher notes")
    k1 = flat_[IX["key1"] - 1]
    for key in FILE_PROG:
        ans, prog = analysis_answers(key)
        if " – ".join(prog) not in k1:
            errs.append(f"answer key: {key} Roman numerals {prog} missing")
    for key in MELODY:
        t = flat_[IX["exmel"] - 1]
        for meas in MELODY[key]:
            if "".join(letter(x) for x, _ in meas) not in t.replace(" ", ""):
                errs.append(f"worked melody {key}: a measure is missing from the page")
    for d in DAYS:
        if f"DAY {d['n']}" not in flat_[IX["pacing"] - 1]:
            errs.append(f"pacing page missing Day {d['n']}")
    # pacing, Day 1 post, demo steps and deck all name the same file (Unit 3 project or your file)
    for pk in ("post1", "setup", "day12"):
        if "Unit 3" not in flat_[IX[pk] - 1]:
            errs.append(f"{pk}: must name the Unit 3 project as the file to continue in")
    # Roman numerals are case-sensitive (vi is minor, VI is major): no progression may come out upper-cased by a caps heading
    progs = list(FILE_PROG.values()) + list(BASS_PROG.values()) + [SBM_FACT[2]]
    pdf_all, deck_all = " ".join(flat_), " ".join(stext)
    for p in progs:
        if [r.upper() for r in p] == p:
            continue
        for sep in ("–", "-", ", ", " – "):
            bad = sep.join(r.upper() for r in p)
            if bad in pdf_all or bad in deck_all:
                errs.append(f"progression {sep.join(p)} printed in capitals as {bad} (case changes the chord quality)")
    ptxt = [pg.extract_text() for pg in PdfReader(PREVIEW_PDF).pages]
    if sum("PREVIEW" in t for t in ptxt) != len(PREVIEW_PICKS):
        errs.append("preview PDF: watermark missing on some pages")
    blobs = {"resource PDF": pdf_text(PDF), "preview PDF": pdf_text(PREVIEW_PDF), "deck PDF": pdf_text(DECK_PDF), "deck pptx": alldeck,
             "zip names": " ".join(names)}
    with zipfile.ZipFile(DECK) as z:
        blobs["deck XML"] = " ".join(z.read(n).decode("utf8", "ignore") for n in z.namelist() if n.endswith(".xml") or n.endswith(".rels"))
    for path in (PDF, PREVIEW_PDF, DECK_PDF):
        meta = PdfReader(path).metadata or {}
        blobs[f"{os.path.basename(path)} metadata"] = " ".join(str(v) for v in meta.values())
    for fn in os.listdir(HERE):
        if fn == "QA.md":
            continue
        if fn.endswith((".md", ".json", ".py", ".txt")) or fn in ("release", "uploaded"):
            blobs[fn] = open(os.path.join(HERE, fn), encoding="utf8", errors="ignore").read()
    for label, txt in blobs.items():
        low = txt.lower()
        for b in BANNED:
            if b.lower() in low:
                errs.append(f"{label}: contains banned string '{b}'")
    for p in [COVER_PNG, *PREVIEW_PNGS]:
        if Image.open(p).size != (2000, 2000):
            errs.append(f"{os.path.basename(p)}: not 2000x2000 (low-res build?)")
    tj = json.load(open(TPT_JSON))
    if tj != tpt_json():
        errs.append("tpt.json is stale; rebuild")
    if tj["price"] != PRICE or tj["license_price"] != PRICE or tj.get("tax_code") != "Other Digital Goods - No Physical Media":
        errs.append("tpt.json price/license_price/tax_code")
    if tj["grades"] != ["9th-grade", "10th-grade", "11th-grade", "12th-grade"] or tj["subjects"] != ["Music"]:
        errs.append("tpt.json grades 9-12 / subject Music")
    proven_tags = {"Halloween", "Autumn", "Winter", "Thanksgiving"}
    if not (1 <= len(tj["tags"]) <= 6 and len(tj["formats"]) <= 3) or not set(tj["tags"]) <= proven_tags or tj["formats"] != ["PDF"]:
        errs.append("tpt.json tags/formats: only proven values (tags from 001-007, format PDF)")
    for k, f in tj["files"].items():
        if not os.path.exists(os.path.join(HERE, f)):
            errs.append(f"tpt.json file {k} missing: {f}")
    if tj["pages"] != len(rd.pages) + len(prs.slides):
        errs.append("tpt.json pages != PDF pages + slides")
    up_ = open(UPLOAD).read()
    m = re.search(r"^## 1\. Title\s*\n+```\n(.+?)\n```\n\((\d+) characters", up_, re.M | re.S)
    if not m:
        errs.append("UPLOAD.md: title block or its character count not found")
    else:
        n = len(m.group(1).strip())
        if n > 80:
            errs.append(f"UPLOAD.md: title is {n} chars (max 80)")
        if int(m.group(2)) != n:
            errs.append(f"UPLOAD.md: says {m.group(2)} characters, title has {n}")
        if m.group(1).strip() != LISTING_TITLE:
            errs.append("UPLOAD.md title differs from LISTING_TITLE in gen.py")
    m2 = re.search(r"^## 2\. Description\s*\n+```\n(.+?)\n```", up_, re.M | re.S)
    if not m2 or not m2.group(1).rstrip().endswith(DISCLOSURE):
        errs.append("UPLOAD.md: description must end with the disclosure line")
    else:
        desc = m2.group(1)
        if not desc.startswith("Melody writing for high school music technology:"):
            errs.append("UPLOAD.md description: first line must lead with the search phrase (METADATA-GUIDE)")
        for need in (f"{len(PLAN)}-page", f"{DECK_SLIDES}-slide", "Logic Pro", f"/{RUBRIC_TOTAL}", "YouTube or Spotify", "GarageBand",
                     f"{DAYS_N}-day", "bassline", "frequency separation", SNA[0], SBM[0]):
            if need not in desc:
                errs.append(f"UPLOAD.md description: missing '{need}'")
        if re.search(r"  +\S|\s[,.]\s|\bteh\b|\s{2,}$", desc, re.M):
            errs.append("UPLOAD.md description: double space or stray punctuation")
        if "/25" in desc or "checked by computer" in desc:
            errs.append("UPLOAD.md description: wrong rubric total or an unverifiable claim")
    for need in (f"${PRICE}", f"{len(PLAN)} pages", os.path.basename(ZIP), os.path.basename(DECK), os.path.basename(DECK_PDF)):
        if need not in up_:
            errs.append(f"UPLOAD.md: missing '{need}'")
    for f in ("release", "uploaded"):
        if os.path.exists(os.path.join(HERE, f)):
            errs.append(f"{f} exists: only the Manager / uploader writes it")
    return errs


def check():
    errs = verify_theory() + verify_data() + verify_labels() + verify_files()
    if errs:
        print("CHECK FAILED")
        print("\n".join(" - " + e for e in errs))
        sys.exit(1)
    print(f"CHECK OK: theory engine confirms C major / A minor scales (no sharps or flats, relative keys, pentatonic missing F and B), the "
          f"owner's two Reference Cards, the file bassline (whole-note roots), both worked melodies (in key, white keys, 4/4, monophonic, "
          f"start on the root, beats 1 and 3 on chord notes, a shape, above the bass), both worked basslines and the cover bass (root on every "
          f"chord, monophonic, 2 rhythm ideas, in key and register), the 3-note sets, the why sentences and the song facts; {len(LABELS)} "
          f"diagram labels fit; owner's lists and two rubrics 3 x (2-5) = /{RUBRIC_TOTAL} printed; {DAYS_N} days run 0-{MINUTES}; every one "
          f"of {DECK_SLIDES} slides has exactly one day slot matching its Day badge; {len(PLAN)} Letter pages, no overflow or dead space, "
          f"footers and valid page/slide references; cover lines whole; preview 4 watermarked pages; deck 16:9, notes on every slide, text in "
          f"bounds, no overlapping text boxes, {len(DECK_PICS)} drawn PNGs, PDF copy {DECK_SLIDES} pages; only Oswald/Source Sans 3, embedded; "
          f"no banned strings; zip current; images 2000x2000; tpt.json proven values; UPLOAD.md title {len(LISTING_TITLE)} chars, SEO "
          f"opening and disclosure.")


if __name__ == "__main__":
    if "--fallback" in sys.argv:
        deck_pdf(DECK, os.path.join(BUILD, "deck-fallback.pdf"), fonts=False)
        print("wrote build/deck-fallback.pdf")
    elif "--check" in sys.argv:
        check()
    else:
        build()
