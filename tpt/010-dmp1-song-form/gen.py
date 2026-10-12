#!/usr/bin/env python3
"""DMP1 Unit 5: Song Form (TpT 010): build the teacher resource, the intro deck and the listing files.

    python3 gen.py            build everything into this folder
    python3 gen.py --lowres   same, but square images at scale 1 and small page shots (self-check only; --check rejects them)
    python3 gen.py --check    verify data, files, page/slide counts, day slots, speaker notes, fonts, banned strings, tpt.json and UPLOAD.md

The owner's own curriculum (Brian Lotze, Digital Music Production 1, "Song Form"), reformatted in the Hudson Beat style:
a 4-day song-structure unit in Logic Pro. Day 1 is a listening day: students map the form of real pop songs on the owner's
Song Form sheet (two songs as a class, two on their own) and answer the pattern question. Days 2-4: they plan a form on
paper, build it in Logic with arrangement markers (at least 4 sections, each at least 4 measures; the worked example is a
32-measure skeleton), differentiate verse and chorus with a chord layer and drum transitions, then present.
Source documents (owner's Drive, read-only): Project Sheet & Rubric, Teacher Notes, Day 1 Classroom Post, Song Form - Pop
Music (Complete) and (Blank) sheets, and the DMP1 Restructure (Unit 5: 4 days, "32-measure skeleton, then form-spreadsheet
comparison"). The owner's Teacher Notes pace the unit in 3 days; the 4-day plan keeps every block of those 3 days and gives
building its own day (see TN page "How this kit differs").
Songs are named by title and artist only (find them on YouTube or Spotify); the answer key keeps only facts from the owner's
Complete sheet: section names and measure counts. No audio, video, lyrics or album art. School names and logos removed;
no student work, names or grades.

Outputs: <slug>.pdf (teacher + student resource), <slug>-intro-deck.pptx and its PDF copy, a ZIP of those three
(the TpT product file), a 4-page watermarked preview PDF, cover.png and preview-1..3.png (2000 x 2000).
Renders with Playwright + Chromium (../tools/render.js) and LibreOffice (deck PDF), fonts from ../fonts. $0 AI.
"""
import json
import os
import re
import shutil
import subprocess
import sys
import zipfile
from html import escape

HERE = os.path.dirname(os.path.abspath(__file__))
TPT = os.path.dirname(HERE)
FONTS = os.path.join(TPT, "fonts")
BUILD = os.path.join(HERE, "build")
RENDER_JS = os.path.join(TPT, "tools", "render.js")

SLUG = "song-form"
PDF = os.path.join(HERE, f"{SLUG}.pdf")
DECK = os.path.join(HERE, f"{SLUG}-intro-deck.pptx")
DECK_PDF = os.path.join(HERE, f"{SLUG}-intro-deck.pdf")
ZIP = os.path.join(HERE, f"{SLUG}.zip")
PREVIEW_PDF = os.path.join(HERE, f"{SLUG}-PREVIEW.pdf")
COVER_PNG = os.path.join(HERE, "cover.png")
PREVIEW_PNGS = [os.path.join(HERE, f"preview-{i}.png") for i in (1, 2, 3)]
UPLOAD = os.path.join(HERE, "UPLOAD.md")
TPT_JSON = os.path.join(HERE, "tpt.json")
ZIP_DIR = "Song Form - Hudson Beat"

STORE = "Brian Lotze"      # copyright holder and TpT seller
BRAND = "Hudson Beat"      # store line on covers and previews (owner, 4 Oct)
YEAR = 2026
ACCENT = "#6E44D6"         # this product's single accent: violet (year-round, distinct from 008 red and 009 orange)
TITLE = "Song Form"
PRICE = "10.00"
UNIT = 5                   # DMP1 Restructure numbering (the owner's Drive docs still say "Unit 4"; flagged to the owner)
DAYS_N = 4
LOW_RES = "--lowres" in sys.argv
DISCLOSURE = "Designed with the help of digital and AI tools, and checked by hand."
LISTING_TITLE = "Song Structure in Logic Pro | Pop Song Form Unit | High School Music Technology"

# Strings that must never appear in any output or in any file in this folder (school branding, LMS). Split so this file never contains them.
BANNED = ["Blind" + " Brook", "BB" + "HS", "blind" + "brook", "Google" + " Classroom"]

# ---------------------------------------------------------------- the owner's Song Form sheet (facts only: title, artist, section names, measure counts)
# (title, artist, use, [(section, measures, shown-as or None)], labels pre-filled on the owner's Blank sheet)
SONGS = [
    ("Crazy", "Gnarls Barkley", "class",
     [("Intro", 1, None), ("Verse", 16, None), ("Chorus", 8, None), ("Verse", 16, None), ("Chorus", 8, None), ("Verse", 16, None),
      ("Chorus", 8, None), ("Outro", 8, None)], 8),
    ("Blinding Lights", "The Weeknd", "class",
     [("Intro (Pop Drop)", 12, "12 (4+8)"), ("Verse", 16, None), ("Pre-chorus", 8, None), ("Chorus", 16, None), ("Pop Drop", 8, None),
      ("Verse", 8, None), ("Pre-chorus", 8, None), ("Chorus", 16, None), ("Bridge", 8, None), ("Chorus", 8, None), ("Pop Drop (2x)", 16, None),
      ("Outro", 8, None)], 5),
    ("Anti-Hero", "Taylor Swift", "choice",
     [("Intro", 2, None), ("Verse", 16, None), ("Chorus", 8, None), ("Post chorus", 2, None), ("Verse", 16, None), ("Chorus", 8, None),
      ("Post chorus", 2, None), ("Verse", 8, None), ("Bridge/Chorus?", 8, None), ("Chorus", 8, None), ("Outro", 2, None)], 4),
    ("good 4 u", "Olivia Rodrigo", "choice",
     [("Intro", 4, None), ("Verse", 16, None), ("Chorus", 16, None), ("Post chorus", 4, None), ("Verse", 16, None), ("Chorus", 16, None),
      ("Post chorus", 4, None), ("Bridge A", 8, None), ("Bridge B", 16, None), ("Chorus", 16, None), ("Outro", 6, None)], 4),
    ("Levitating", "Dua Lipa", "choice",
     [("Intro", 4, None), ("Verse", 8, None), ("Pre-chorus", 4, None), ("Chorus", 8, None), ("Post chorus", 2, None), ("Verse", 8, None),
      ("Pre-chorus", 4, None), ("Chorus", 8, None), ("Bridge A", 8, None), ("Bridge B", 8, None), ("Pre-chorus", 4, None), ("Chorus (2x)", 16, None)], 0),
    ("Uptown Funk", "Bruno Mars", "choice",
     [("Intro", 8, None), ("Verse", 8, None), ("Pre-chorus", 8, None), ("Chorus", 8, None), ("Pop Drop (Horns)", 8, None), ("Post Chorus", 4, None),
      ("Verse", 8, None), ("Pre-chorus", 8, None), ("Chorus", 8, None), ("Pop Drop (Horns)", 8, None), ("Post Chorus", 4, None), ("Bridge", 20, None),
      ("Pop Drop (Horns)", 8, None), ("Post Chorus", 4, None), ("Outro", 16, None)], 0),
]
SHEET_ROWS = 15            # rows per song table on the student sheet (the longest analysis, Uptown Funk, has 15 sections)
# owner's "Suggested Songs" table (Teacher Notes): why each one works
SONG_WHY = {
    "Blinding Lights": "Verse is sparse and unsettled. Chorus erupts with full synth pad and drums. Section change is physically felt. Textbook structure.",
    "Crazy": "Very simple form.",
    "Anti-Hero": "Clear verse/pre-chorus/chorus/bridge. Good for introducing what a pre-chorus does: builds tension before the drop.",
    "good 4 u": "Strong bridge that genuinely contrasts with the verse-chorus cycle. Good for finding a bridge in the wild.",
    "Uptown Funk": "Near-universal recognition. Has a clear break section. Good for showing not all songs follow the same template.",
    "Levitating": "Clean pop structure with a pre-chorus students will not have consciously noticed before. Good for more advanced students.",
}
# owner's Day 1 post: the one-line hint for each independent choice
SONG_HINT = {
    "Anti-Hero": "Clear verse / pre-chorus / chorus / bridge structure",
    "good 4 u": "Strong bridge that genuinely contrasts with the verse-chorus cycle",
    "Uptown Funk": "Classic structure with a break section that functions differently from a bridge",
    "Levitating": "Clean pop form with a pre-chorus you may not have consciously noticed before",
}

# ---------------------------------------------------------------- the owner's words (Project Sheet & Rubric, Day 1 post, Teacher Notes; lightly edited: LMS, unit numbers)
OVERVIEW = ("Every song you have ever heard is built on a structure. Verses, choruses, bridges, intros, outros: these sections are not accidents. "
            "They are deliberate decisions about energy, tension, and release. In this unit you will train your ear to hear those decisions in real "
            "songs, then make your own. You will build a structured arrangement of your chords, bass and melody with clearly differentiated sections: "
            "ones that feel different from each other because of what is playing, not just because a label says so.")
GOALS = ["Identify and label sections in real songs using the Song Form sheet",
         "Understand what makes each section feel different from the others",
         "Recognize common structural patterns across different songs and genres",
         "Plan a song form before building it: intro, verse, chorus at minimum",
         "Build a structured arrangement of your chords, bass and melody in Logic with at least 4 distinct sections",
         "Use the presence and absence of elements to differentiate sections",
         "Signal every section change with a deliberate transition"]
VOCAB = [("Song form", "the overall structure of a song: how its sections are organized from beginning to end"),
         ("Intro", "opens the song: establishes the sound and draws the listener in before the verse begins. Usually sparse."),
         ("Verse", "tells the story. Sparser than the chorus. Repeats with varied content. Sets up the emotional release of the chorus."),
         ("Pre-chorus", "a short section between verse and chorus that builds tension before the drop. Not all songs have one."),
         ("Chorus", "the emotional peak. Fuller, louder, the most memorable part. Usually returns multiple times unchanged."),
         ("Bridge", "a contrasting section that breaks the verse-chorus cycle. Creates tension or surprise before the final chorus."),
         ("Outro", "brings the song to a close. Often strips elements back one by one. Mirrors the intro in energy."),
         ("Texture", "how thick or thin the sound is: determined by how many elements are playing at once"),
         ("Arrangement", "the decisions about which elements play in which sections: what comes in, what drops out, and when"),
         ("Transition", "the moment between sections that signals a change: drum fill, breakdown, cut, or build")]
REQUIREMENTS = ["Song Form sheet completed for all four songs: two as a class, two independently",
                "Written response: what did you notice across all four songs? (see the Day 1 question)",
                "Melody added to your chords-and-bassline Logic file so all three elements are in one place",
                "Song form planned on paper before building: at least 4 sections with labels and lengths",
                "Arrangement built in Logic with arrangement markers labeling every section",
                "At least 4 distinct sections: intro, verse, chorus required; bridge and outro encouraged",
                "Each section at least 4 measures long",
                "Verse and chorus must feel clearly different from each other",
                "At least one chord layer added on a second instrument for texture contrast",
                "Drum track with a transition at every section change",
                "Project named: Your Name – Song Form",
                "Exported as MP3 and .zip (compressed .logicx file) and submitted before presenting"]
SECTION_JOBS = [("Intro", "Earns the listener’s attention. Usually starts with one or two elements and builds. Does not need to be complicated: just needs "
                          "to feel like a beginning."),
                ("Verse", "Gives the listener space to settle in. Sparser than the chorus. If the chorus is the payoff, the verse is the setup. A weak "
                          "verse makes a chorus feel less powerful."),
                ("Pre-Chorus", "Optional but effective. Builds tension in the last 2–4 measures before the chorus drops. Often raises energy slightly so "
                               "the chorus feels like a release."),
                ("Chorus", "The emotional peak. Should feel fuller, louder, and more energetic than the verse. This is where the melody is most "
                           "prominent. Copy and paste: it is the same every time it returns."),
                ("Bridge", "Breaks the pattern. Usually appears once, after the second chorus. Different chord, different energy, or stripped back "
                           "completely. Makes the final chorus feel earned."),
                ("Outro", "Brings it home. Often mirrors the intro, stripping elements back one by one. Ends cleanly. Does not just fade out in the "
                          "middle of a loop.")]
PLAN_ROWS = ["Intro", "Verse 1", "Chorus 1", "Verse 2", "Chorus 2", "Bridge (optional)", "Outro"]
STEPS = [  # (days, heading, text): the owner's six steps
    ("Day 1", "Analyze real songs", "Complete the Song Form sheet for all four songs. Two as a class, two independently. For each song: label every "
                                    "section, count the measures, and note what elements are playing. Ask yourself after each one: why does the chorus "
                                    "feel the way it does?"),
    ("Day 2", "Add your melody", "Open your chords-and-bassline Logic file. Add your melody so all three elements (chords, bass, and melody) are in one "
                                 "place. Call your teacher over to check before moving on."),
    ("Day 2", "Plan before you build", "Fill in the planning table on this sheet. Every section needs a name, a length, a list of what is playing, and a "
                                       "transition. Do this on paper. Do not open the arrangement view until your plan is complete."),
    ("Days 2–3", "Build your structure", "Add arrangement markers in Logic for every section in your plan. Copy your existing elements into the correct "
                                         "sections. Build the skeleton first (correct elements in the right places) before refining anything."),
    ("Day 3", "Differentiate your sections", "Go back through and make sure the chorus feels different from the verse. Add a chord layer on a second "
                                             "instrument for the chorus. Add the melody only where it belongs. Make sure every section change has a drum "
                                             "transition."),
    ("Day 4", "Mix and submit", "Play the full arrangement from start to finish. Can you hear which section you are in without looking at the screen? "
                                "If yes, your arrangement is working. Adjust volumes, export, and submit."),
]
TIPS = ["The chorus should feel bigger than the verse. If they sound the same, something is missing. Add the melody, layer the chords on a second "
        "instrument, or thicken the drums.",
        "Sections are almost always 4 or 8 bars long. If you analyzed the four songs carefully, you already know this from experience.",
        "Copy and paste is not lazy. Every professional arranger does it. Write once, arrange many times.",
        "Your intro does not need to be complicated. One element playing alone for 4 bars is a stronger intro than everything at full volume from bar one.",
        "Every section change needs a transition. If you can hear where one section ends and another begins without looking at the screen, your "
        "transitions are working."]
CHECKLIST = ["Song Form sheet completed for all four songs",
             "Written response to the pattern question completed",
             "Melody added to my chords-and-bassline Logic file",
             "Song form planning table filled in before building",
             "At least 4 distinct sections with arrangement markers",
             "Each section at least 4 measures long",
             "Verse and chorus clearly different from each other",
             "Chord layer added on a second instrument for texture contrast",
             "Drum track with a transition at every section change",
             "Project named: Your Name – Song Form",
             "Exported as MP3 and .zip (compressed .logicx file) and submitted before presenting",
             "Three feedback questions posted to the class feedback board before presenting"]
RUBRIC = [  # the owner's rubric, levels 2-5, 5 categories = /25 (Padlet written as "feedback questions"; nothing added)
    ("Song Form Analysis",
     ["Fewer than two spreadsheets completed. Sections not labeled or measures not counted. Written response missing.",
      "All four spreadsheets complete but with labeling errors or missing measure counts. Written response is vague.",
      "All four spreadsheets complete and mostly accurate. Written response identifies at least one pattern across songs.",
      "All four spreadsheets complete and accurate. Written response shows genuine insight about how song structure works across different songs."]),
    ("Song Form Structure",
     ["Fewer than 3 sections built. No arrangement markers. No clear structural plan evident.",
      "At least 3 sections but some are unlabeled or too short. Planning table not completed before building.",
      "At least 4 labeled sections, each at least 4 measures. Planning table completed. Intro, verse, and chorus all present.",
      "Form is well-constructed and clearly planned. All required sections present. Structure feels like a complete song with a beginning, middle, and end."]),
    ("Section Differentiation",
     ["All sections sound the same. No meaningful difference in texture or elements between sections.",
      "Some difference between sections but verse and chorus are not clearly distinct. Layering absent or ineffective.",
      "Verse and chorus clearly different. At least one chord layer added on a second instrument. Melody placement is intentional.",
      "Multiple sections have distinct textures. Addition and removal of elements is deliberate. Listener can identify the section by ear without "
      "looking at the screen."]),
    ("Arrangement as Mix",
     ["No evidence of mix decisions tied to arrangement. All elements at equal volume throughout with no change between sections.",
      "Some volume awareness but mix does not meaningfully reflect arrangement. Dense and sparse sections feel similar in the mix.",
      "Mix reflects the arrangement plan. Denser sections are fuller and louder. Sparse sections have room and clarity.",
      "Arrangement and mix work as one. Every addition and removal of an element is audible and intentional. Nothing is clipping."]),
    ("Documentation & Presentation",
     ["Project not submitted or not presented.",
      "Project submitted late or named incorrectly. Minimal presentation engagement.",
      "Project named correctly and submitted. Presentation made. Feedback questions posted.",
      "Project correctly named, on time. Student explains form and arrangement choices clearly, referencing specific sections. Three thoughtful "
      "feedback questions posted."]),
]
LEVELS = [2, 3, 4, 5]
RUN_ANALYSIS = [  # Teacher Notes: how to run each class example (10-12 minutes each; first 2-3 minutes of the song only)
    ("First listen (90 seconds)", "Play without stopping. Ask students to raise their hand when they feel the section change. Do not define: just listen."),
    ("Second listen", "Pause at each section change. Ask: what changed? What came in or dropped out? Students call it out. You label it on the board."),
    ("Measure count", "Play one more time and count together out loud. Students will usually discover on their own that sections land in multiples of 4."),
    ("Key question before moving on", "Why does the chorus feel bigger? What specifically is different? Push for musical answers, not just ‘it sounds louder’."),
]
PLAN_CHECK = ["The chorus has more elements than the verse", "The intro starts with fewer elements than the verse",
              "Every section change has a transition type written down", "Sections are labeled with real names, not just ‘Part 1’ and ‘Part 2’"]
LAYER_DEMO = ["Option-drag a chord region to duplicate it onto a new track below",
              "Change the instrument to something with a different character: a pad, a Rhodes, a slow-attack synth",
              "Play both chord tracks together: the texture thickens without the notes changing",
              "Mute the second track in the verse, unmute it in the chorus",
              "Play the full arrangement and ask: can you hear the chorus get bigger? That is arrangement as mix."]
CALLBACKS = [("Drum beats (Beat From Nothing)", "Drum beats return here to anchor sections. Remind students: a drum pattern that does not change across "
                                                 "section changes sounds like a loop, not a song."),
             ("Transitions (Stems & Grooves)", "The four transition types (fill, breakdown, cut, build) are the vocabulary for every section change "
                                               "in this unit. Ask students to name which type they used and why."),
             ("Chords, bass and melody (Building Blocks)", "These are the raw materials. This unit is purely about where those materials go and what "
                                                           "they do in each section.")]
STRUGGLES = [("All sections sound the same", "The most common failure. The test: mute the melody track. Does the chorus still feel bigger? If not, the "
                                             "arrangement is not doing any work. Ask the student to listen to their analyzed song again and compare."),
             ("Rushing through the spreadsheet", "Students who treat the analysis as a checkbox exercise will not have the vocabulary or intuition to plan "
                                                 "their own form. Enforce quality: every section labeled, every measure counted, the written response "
                                                 "answered thoughtfully."),
             ("Intros at full energy", "Students start with everything playing because they are not thinking about the listener. Ask: if someone walked "
                                       "into the room mid-sentence, would they know what was coming? The intro should earn attention, not assume it."),
             ("No transitions", "Students paste their drum groove everywhere and call it done. The transition is what tells the listener something "
                                "changed. A section change with no acknowledgment sounds like a loop.")]
TRANSITIONS = [("Fill", "a short drum pattern that leads into the next section"),
               ("Breakdown", "elements drop out so the next section hits harder"),
               ("Cut", "everything stops for a beat, then the new section starts"),
               ("Build", "energy rises (snare roll, rising notes) into the next section")]
PATTERNS = ["Sections are almost always 4 or 8 bars", "The chorus is always (texturally) fuller than the verse", "Intros build in and outros strip back"]
VOCAB_QUIZ = [("The ______ is the emotional peak of a song: fuller, louder, the most memorable part.", "chorus"),
              ("How thick or thin the sound is (how many elements play at once) is its ______.", "texture"),
              ("A drum fill, breakdown, cut or build between sections is a ______.", "transition"),
              ("The decisions about which elements play in which sections are the ______.", "arrangement")]

# ---------------------------------------------------------------- worked example: a 32-measure skeleton (Hudson Beat example, C major, I-V-vi-IV)
LANES = ["Chords (piano)", "Chord layer (pad)", "Bass", "Melody", "Drums"]
EX_KEY, EX_PROG, EX_CHORDS = "C major", ["I", "V", "vi", "IV"], ["C", "G", "Am", "F"]
# (section, measures, {lane: True | (first, last) measure within the section, 1-based}, transition into the next section, how)
EXAMPLE = [
    ("Intro", 4, {"Chords (piano)": True}, "Fill", "a one-measure drum fill in m. 4 brings the drums in"),
    ("Verse 1", 4, {"Chords (piano)": True, "Bass": True, "Drums": True}, "Build", "snare on every eighth note in m. 8, rising into the chorus"),
    ("Chorus 1", 8, {"Chords (piano)": True, "Chord layer (pad)": True, "Bass": True, "Melody": True, "Drums": True}, "Cut",
     "everything stops on beat 4 of m. 16"),
    ("Verse 2", 4, {"Chords (piano)": True, "Bass": True, "Drums": True}, "Fill", "a drum fill in m. 20"),
    ("Chorus 2", 8, {"Chords (piano)": True, "Chord layer (pad)": True, "Bass": True, "Melody": True, "Drums": True}, "Breakdown",
     "the drums and the pad drop out after m. 28"),
    ("Outro", 4, {"Chords (piano)": True, "Bass": (1, 2)}, "End", "the bass drops out after m. 30; the last chord rings on the I chord"),
]
EX_MEASURES = sum(s[1] for s in EXAMPLE)   # 32


def ex_regions(lane):
    """[(first, last)] absolute measures where a lane plays in the worked example."""
    out, start = [], 1
    for _, bars, lanes, _, _ in EXAMPLE:
        v = lanes.get(lane)
        if v is True:
            out.append((start, start + bars - 1))
        elif v:
            out.append((start + v[0] - 1, start + v[1] - 1))
        start += bars
    return out   # one region per section, as if copied in section by section


def ex_starts():
    s, out = 1, []
    for name, bars, *_ in EXAMPLE:
        out.append((name, s, s + bars - 1))
        s += bars
    return out


# ---------------------------------------------------------------- pacing: 4 days of about 53 minutes (DMP1 Restructure, Unit 5); blocks from the owner's 3-day Teacher Notes plan
# block = (minutes, heading, text, deck slides (first, last) or None). Every deck slide belongs to exactly one block (checked).
DAYS = [
    {"n": 1, "kind": "Learning", "title": "Song form analysis: a listening day",
     "goal": "All four Song Form sheets complete. Written response answered. Students have the vocabulary and structural intuition to plan their own form.",
     "blocks": [("0–12", "Hook + class example 1", "Slides 1–2 (two minutes): every song has a blueprint, and what song form means. Then slide 3: "
                 "Crazy by Gnarls Barkley, first 2–3 minutes only. Three listens: feel the change, label the sections, count the measures. "
                 "Key question: why does the chorus feel bigger?", (1, 3)),
                ("12–25", "Class example 2", "Slide 4 (one minute): what each section does. Slide 5: Blinding Lights by The Weeknd. Students lead more of "
                 "the labeling. Faster pace. Compare to example 1: same structure? Different? What is the same across both?", (4, 5)),
                ("25–45", "Independent analysis", "Slide 6. Students independently analyze two songs from the list on the Song Form sheet (pages SHEETS). "
                 "Circulate and check that sections are labeled and measures counted.", (6, 6)),
                ("45–53", "Share-out", "Slide 7. What patterns showed up in every song? Surface the two key observations: sections are almost always "
                 "4 or 8 bars, chorus is always texturally fuller than verse. Students write the response (page COMPARE).", (7, 7))]},
    {"n": 2, "kind": "Work day", "title": "Plan + build the skeleton",
     "goal": "Planning table complete, melody added, arrangement skeleton built with a marker on every section.",
     "blocks": [("0–5", "Chord layer demo", "Slide 8: arrangement is what plays where. Then demo the chord layer technique live in Logic "
                 "(slide 9; 5 minutes only, steps on page TN2).", (8, 9)),
                ("5–12", "The project", "Slides 10–16: the worked 32-measure skeleton, plan before you build, the four transitions, the project, "
                 "the steps, what to turn in and the rubric.", (10, 16)),
                ("12–25", "Plan + melody", "Students complete the planning table (page PLAN) and add their melody to their chords-and-bassline Logic "
                 "file. No Logic arrangement work until the plan is done: check each student.", None),
                ("25–50", "Build the skeleton", "Students add arrangement markers and copy their elements into the sections. Skeleton first: the "
                 "correct elements in the right places, before refining anything.", None),
                ("50–53", "Save + note", "Students save and write down what is left: chord layer, transitions, mix.", None)]},
    {"n": 3, "kind": "Work day", "title": "Differentiate the sections",
     "goal": "Chord layer added, melody placed on purpose, a drum transition at every section change; verse and chorus clearly different.",
     "blocks": [("0–5", "The mute test", "Mini-lesson: mute the melody track. Does the chorus still feel bigger? If not, the arrangement is not doing "
                 "any work yet.", None),
                ("5–45", "Differentiate", "Students add the chord layer, place the melody only where it belongs, and add a drum transition at every "
                 "section change (fill, breakdown, cut or build). Circulate with the prompts below.", None),
                ("45–53", "Play-through", "Each student plays the arrangement start to finish: can you hear which section you are in without looking "
                 "at the screen? Note one fix for Day 4. Students who are close to done export and submit.", None)]},
    {"n": 4, "kind": "Finish + present", "title": "Finish, submit, present",
     "goal": "Every project submitted and presented; every student can point to the chorus and say why it feels bigger.",
     "blocks": [("0–20", "Finish + submit", "Students finalize transitions, mix volumes, export, and submit. Three feedback questions go on the class "
                 "feedback board.", None),
                ("20–50", "Presentations", "Each student plays their full arrangement. Ask: point to your chorus. Why does it feel different from your "
                 "verse? What did you do to make that happen?", None),
                ("50–53", "Close", "Slide 17. ‘What you just heard (intro, verse, chorus, verse, chorus) is the structure of almost every pop song ever "
                 "released. You did not just copy that structure. You understood it well enough to build it yourself.’", (17, 17))]},
]

# ---------------------------------------------------------------- intro deck content (the owner's "Looping in Logic" pattern); (day, kind, title, body, notes)
L1, L2 = SONGS[0], SONGS[1]
SECTION_SHORT = [("Intro", "Earns attention. Starts sparse and builds."), ("Verse", "The setup. Sparser than the chorus."),
                 ("Pre-chorus", "Optional. Builds tension before the chorus."), ("Chorus", "The peak. Fuller, louder, same every time."),
                 ("Bridge", "Breaks the pattern once, before the last chorus."), ("Outro", "Brings it home. Strips back. Ends cleanly.")]
SLIDES = [
    (1, "title", "Song Form", "Feat. Blinding Lights",
     "Day 1, the first two minutes. Today is a listening day: nobody builds anything in Logic. Have Crazy (Gnarls Barkley) and Blinding Lights "
     "(The Weeknd) cued up on YouTube or Spotify from your own account; no audio or video is included in this deck. Open with a question: you always "
     "know when the chorus is coming, even in a song you have never heard. How?"),
    (1, "bullets", "What the heck is song form?", "The overall structure of a song|How its sections are organized, beginning to end|"
     "Intro · verse · chorus · bridge · outro|Sections feel different because of what is playing",
     "Plain-English definition, then the central insight of the unit, the line to land by the end of the week: sections feel different because of "
     "what is playing, not just because of a label. A chorus is not called a chorus because someone decided to call it that. It is called a chorus "
     "because it is texturally fuller, emotionally bigger and more memorable than the verse. That is a compositional decision, not a naming decision."),
    (1, "listen", "Listen: feel the change", f"{L1[0]} · {L1[1]}|Teacher: play the first 2–3 minutes from YouTube or Spotify.|"
     "Raise your hand when the section changes|What came in? What dropped out?|Count the measures: what number keeps showing up?",
     "Class example 1, about 10 to 12 minutes. First listen, about 90 seconds: play without stopping; students raise a hand when they feel the "
     "section change. Do not define, just listen. Second listen: pause at each change and ask what came in or dropped out; label it on the board. "
     "Third listen: count the measures together out loud. Students usually discover on their own that sections land in multiples of 4."),
    (1, "sections", "What each section does", "",
     "One minute, between the two class examples. A section is not just a label: it has a job, and the job tells you how to build it. Point back to "
     "Crazy: where was the setup, where was the payoff? Mention that not every song has a pre-chorus or a bridge; students will hunt for them in "
     "their independent songs. The full descriptions are on the project sheet."),
    (1, "listen", "Listen: song 2", f"{L2[0]} · {L2[1]}|Teacher: play the first 2–3 minutes from YouTube or Spotify.|"
     "You lead the labeling this time|Same structure as Crazy? What is different?|Why does the chorus feel bigger?",
     "Class example 2, same process at a faster pace, with students leading more of the labeling. The verse is sparse and unsettled; the chorus "
     "erupts with full synth pad and drums, so the section change is physically felt. Push for musical answers to the last question (what came in, "
     "what got thicker), not just ‘it sounds louder’."),
    (1, "sheet", "Your turn: two songs on your own", "|".join(f"{s[0]} · {s[1]}" for s in SONGS if s[2] == "choice"),
     "Independent analysis, about 20 minutes. Each student picks two of the four songs and completes a Song Form table for each: label every section, "
     "count the measures, note what is playing. Students play the songs from YouTube or Spotify on their own headphones. The image shows how a finished "
     "row looks, using the class example. Circulate and check that sections are labeled and measures counted."),
    (1, "bullets", "What showed up every time?", "Sections are almost always 4 or 8 bars|The chorus is fuller than the verse|"
     "Intros build in; outros strip back|Write it down: your written response",
     "The share-out at the end of Day 1 is important. Ask: what patterns did you notice across all four songs? Let students get there before you show "
     "the bullets. These are the principles they will apply the next day. Students finish with the written response, two or three sentences, on the "
     "Compare the Forms page."),
    (2, "arr", "Arrangement = what plays where", "",
     "Day 2 opens here. Arrangement is the decisions about which elements play in which sections: what comes in, what drops out, and when. Walk the "
     "picture left to right: the intro is one element, the verse adds bass and drums, the chorus adds the chord layer and the melody, the outro strips "
     "back. Count the black bars in each column: that count is the texture. The unit’s mixing concept is arrangement as mix."),
    (2, "steps", "Watch me: the chord layer", "Option-drag a chord region onto a new track below|Change the instrument: a pad, a Rhodes, a slow-attack synth|"
     "Play both together: thicker, same notes|Mute the second track in the verse, unmute it in the chorus|Play it all: can you hear the chorus get bigger?",
     "The worked example in Logic, a 5-minute whole-class demo. Do it live with one student’s chords or your own. After step 3 ask what changed; the notes "
     "did not, the texture did. After step 5 ask: can you hear the chorus get bigger? That is arrangement as mix. Students will repeat this on Day 3."),
    (2, "plan", "Worked example: a 32-measure skeleton", "",
     f"A filled-in planning table for a short song in {EX_KEY}, {'–'.join(EX_PROG)} ({' '.join(EX_CHORDS)}). Six sections, 32 measures, every section "
     "4 or 8 bars. Read down the elements column: the intro has one element, the verses three, the choruses five, the outro two and then one. Every "
     "row names its transition. Compare it to Crazy: same 4s and 8s, chorus fullest, but shorter verses because this is a skeleton."),
    (2, "bullets", "Plan before you build", "Every section needs a name and a length|…a list of what is playing|…and a transition into the next one|"
     "Fill in the planning table on paper first|No arrangement work in Logic until your plan is checked",
     "Enforce the planning table: do not let anyone open the arrangement view until the plan is filled in. Walk around and check: the chorus has more "
     "elements than the verse, the intro starts with fewer than the verse, every section change has a transition written down, and sections have real "
     "names, not Part 1 and Part 2."),
    (2, "columns", "Four ways to change sections", "|".join(f"{a}: {b}" for a, b in TRANSITIONS),
     "The four transition types from earlier units are the vocabulary for every section change: fill, breakdown, cut and build. Ask students to name "
     "which type they will use at each change and why. A section change with no acknowledgment sounds like a loop. If you can, tap each one on a drum "
     "track so students hear the difference."),
    (2, "bullets", "Your project: Song Form", "Four songs on the Song Form sheet + a written response|Plan 4+ sections, each 4+ measures|"
     "Build it in Logic with arrangement markers|Make the chorus feel bigger than the verse",
     "Hand out the project sheet and rubric. Students arrange their own chords, bass and melody into a song. Intro, verse and chorus are required; "
     "bridge and outro are encouraged. Sections are at least 4 measures each. The test for the whole project: can you hear which section you are in "
     "without looking at the screen?"),
    (2, "days", "The steps", "",
     "Walk the four days: Day 1 the listening and the Song Form sheet, Day 2 the plan, the melody and the skeleton, Day 3 the chord layer, melody "
     "placement and drum transitions, Day 4 finish, submit and present. The six steps on the project sheet follow the same order."),
    (2, "check", "What you turn in", "Song Form sheet: 4 songs + written response|Planning table, filled in before building|"
     "4+ sections with arrangement markers, 4+ measures each|Chord layer on a second instrument + a transition at every change|"
     "MP3 + compressed .logicx, named Your Name – Song Form|Three feedback questions before presenting",
     "These are the deliverables from the submission checklist. Show where students submit (your class page) and where the feedback questions go "
     "(any class feedback board). Export: File › Bounce › Project or Section, MP3. Compress: in Finder, Control-click the .logicx project › Compress."),
    (2, "rubric", "How it’s scored", "",
     "Five categories, each scored 2 to 5, for a total out of 25. Read the 5 column aloud so students know what great looks like: accurate sheets and "
     "real insight, a complete song with a beginning, middle and end, sections you can identify by ear, a mix that follows the arrangement, and a "
     "presentation that explains the choices section by section."),
    (4, "closer", "One last thing…", "Sections feel different because of what is playing.|Write once, arrange many times.|"
     "You built the form yourself.",
     "Day 4, the last three minutes, after the presentations. Close with this line: what you just heard (intro, verse, chorus, verse, chorus) "
     "is the structure of almost every pop song ever released. You did not just copy that structure. You understood it well enough to build it yourself."),
]
DECK_SLIDES = len(SLIDES)
RUBRIC_SHORT = ["All four sheets accurate; the written response shows real insight into song structure.",
                "A clearly planned, complete song with a beginning, middle and end.",
                "Distinct textures: a listener can name the section by ear.",
                "Arrangement and mix work as one; nothing is clipping.",
                "On time, named correctly; explains form choices section by section; three feedback questions."]


# ---------------------------------------------------------------- html scaffolding (Hudson Beat, from 006-009)
def font_css(base=BUILD):
    f = os.path.relpath(FONTS, base)
    faces = [("Oswald", 300, "normal", "oswald-latin-300-normal"), ("Oswald", 500, "normal", "oswald-latin-500-normal"),
             ("Oswald", 700, "normal", "oswald-latin-700-normal"),
             ("Source Sans 3", 400, "normal", "source-sans-3-latin-400-normal"),
             ("Source Sans 3", 400, "italic", "source-sans-3-latin-400-italic"),
             ("Source Sans 3", 600, "normal", "source-sans-3-latin-600-normal"),
             ("Source Sans 3", 700, "normal", "source-sans-3-latin-700-normal")]
    return "".join(f"@font-face{{font-family:'{n}';font-weight:{w};font-style:{s};src:url('{f}/{file}.woff2') format('woff2');}}"
                   for n, w, s, file in faces)


CSS = """
@page { size: 8.5in 11in; margin: 0; }
* { box-sizing: border-box; }
html, body { margin: 0; padding: 0; background: #fff; color: #000; }
body { font-family: 'Source Sans 3', sans-serif; -webkit-print-color-adjust: exact; print-color-adjust: exact; }
.page { width: 8.5in; height: 11in; position: relative; overflow: hidden; page-break-after: always; break-after: page;
        padding: 0.42in 0.5in 0.38in; display: flex; flex-direction: column; }
.page:last-child { page-break-after: auto; break-after: auto; }
.l { font-family: 'Oswald'; font-weight: 300; }
.top { display: flex; justify-content: space-between; align-items: flex-end; border-bottom: 2.2pt solid #000; padding-bottom: 6pt; }
.top .t { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; letter-spacing: .01em; line-height: 1; font-size: 25pt; }
.top .t small { display: block; font-family: 'Oswald'; font-weight: 300; text-transform: none; font-size: 12pt; letter-spacing: 0; margin-top: 4pt; }
.tag { text-align: right; }
.tag .lv { display: inline-block; background: #000; color: #fff; font-family: 'Oswald'; font-weight: 700; font-size: 14pt; padding: 3pt 9pt 4pt; border-radius: 4pt; letter-spacing: .04em; }
.tag .lv.ak { background: ACCENT; }
.tag .ln { font-family: 'Oswald'; font-weight: 300; font-size: 11pt; margin-top: 3pt; }
.name { display: flex; gap: 18pt; font-size: 12pt; margin: 9pt 0 4pt; }
.name span { flex: 1; border-bottom: 1pt solid #000; height: 15pt; }
.name span.d { flex: 0 0 1.6in; }
.main { flex: 1; display: flex; flex-direction: column; min-height: 0; padding-top: 6pt; }
.foot { display: flex; justify-content: space-between; font-size: 8.5pt; color: #333; border-top: .8pt solid #000; padding-top: 3pt; }
.wm { position: absolute; left: 50%; top: 50%; transform: translate(-50%, -50%) rotate(-38deg); font-family: 'Oswald'; font-weight: 700;
      font-size: 150pt; color: rgba(0,0,0,.13); letter-spacing: .06em; white-space: nowrap; pointer-events: none; z-index: 9; }
h2.s { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 13pt; margin: 9pt 0 3pt; letter-spacing: .02em; }
h2.s .acc, .acc { color: ACCENT; }
p, li { font-size: 10.5pt; line-height: 1.32; margin: 0 0 4pt; }
ul, ol { margin: 0 0 4pt; padding-left: 15pt; }
.tn p, .tn li { font-size: 11pt; line-height: 1.42; margin-bottom: 6pt; }
.tn.big p, .tn.big li { font-size: 12.2pt; line-height: 1.5; margin-bottom: 8pt; }
.tn.big h2.s { font-size: 14pt; margin-top: 12pt; }
.lines { flex: 1; min-height: 44pt; background: repeating-linear-gradient(to bottom, transparent 0, transparent 21pt, #000 21pt, #000 22pt); }
table.t { border-collapse: collapse; width: 100%; font-size: 9.8pt; margin: 2pt 0 4pt; }
table.t td, table.t th { border: 1pt solid #000; padding: 3pt 5pt; text-align: left; vertical-align: top; line-height: 1.25; }
table.t th { font-family: 'Oswald'; font-weight: 500; text-transform: uppercase; font-size: 9.5pt; background: #000; color: #fff; letter-spacing: .04em; vertical-align: middle; }
.cols { display: flex; gap: 20pt; }
.cols > div { flex: 1; min-width: 0; }
.small { font-size: 9pt !important; color: #333; }
.box { border: 1.4pt solid #000; border-radius: 8pt; padding: 7pt 10pt; }
.box.acc-b { border-color: ACCENT; border-width: 2pt; }
.box.inv { background: #000; color: #fff; border-color: #000; }
.k { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 10pt; letter-spacing: .05em; color: ACCENT; display: block; margin-bottom: 2pt; }
.ln2 { border-bottom: 1pt solid #000; height: 22pt; }
.fl { display: flex; align-items: flex-end; gap: 6pt; font-size: 10.5pt; margin: 0 0 8pt; }
.fl b { font-weight: 600; white-space: nowrap; }
.fl span { flex: 1; border-bottom: 1pt solid #000; height: 14pt; }
.num { list-style: none; padding: 0; margin: 2pt 0 0; }
.num li { display: flex; gap: 10pt; margin: 0 0 9pt; }
.num .n { flex: 0 0 24pt; height: 24pt; border-radius: 12pt; background: #000; color: #fff; font-family: 'Oswald'; font-weight: 700; font-size: 13pt;
          display: flex; align-items: center; justify-content: center; }
.num .hd { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 12pt; letter-spacing: .02em; display: block; }
.why { font-style: italic; }
.blk { display: flex; gap: 9pt; margin: 0 0 6pt; font-size: 10.6pt; line-height: 1.32; }
.blk .tm { flex: 0 0 40pt; text-align: center; background: #000; color: #fff; font-family: 'Oswald'; font-weight: 700; font-size: 10pt;
           border-radius: 4pt; padding: 2pt 0 3pt; align-self: flex-start; }
.blk b.h { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 10.5pt; letter-spacing: .02em; }
.day { border: 1.4pt solid #000; border-radius: 8pt; padding: 7pt 10pt 3pt; margin-bottom: 8pt; }
.day .dh { display: flex; justify-content: space-between; align-items: baseline; border-bottom: 1pt solid #000; padding-bottom: 3pt; margin-bottom: 6pt; }
.day .dh b { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 14pt; }
.day .dh span { font-family: 'Oswald'; font-weight: 300; font-size: 11pt; }
.goal { font-size: 10.6pt; margin: 0 0 7pt; }
.chk { display: flex; gap: 8pt; align-items: flex-start; font-size: 11pt; margin: 0 0 8pt; line-height: 1.3; }
table.rub { border-collapse: collapse; width: 100%; }
table.rub td, table.rub th { border: 1.1pt solid #000; padding: 3pt 4pt; vertical-align: top; font-size: 8.6pt; line-height: 1.22; }
table.rub th { background: #000; color: #fff; font-family: 'Oswald'; font-weight: 500; font-size: 11pt; text-align: center; vertical-align: middle; }
table.rub td.c { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 10pt; width: 12%; vertical-align: middle; }
table.rub td.sc { width: 6.5%; }
.cover { align-items: center; justify-content: space-between; padding: 0.55in 0.6in 0.5in;
         background: linear-gradient(180deg, #fff 0%, #fff 55%, #ededed 100%); }
.store { font-family: 'Oswald'; font-weight: 500; font-size: 13pt; letter-spacing: .32em; text-transform: uppercase; }
.ctitle { text-align: center; }
.ctitle .a { white-space: nowrap; font-family: 'Oswald'; font-weight: 700; font-size: 30pt; line-height: 1; letter-spacing: .02em; text-transform: uppercase; }
.ctitle .b { white-space: nowrap; font-family: 'Oswald'; font-weight: 700; font-size: 84pt; line-height: .98; letter-spacing: .01em; text-transform: uppercase; }
.ctitle .c { font-family: 'Oswald'; font-weight: 300; font-size: 15pt; margin-top: 8pt; }
.cband { width: 100%; background: #000; color: #fff; text-align: center; font-family: 'Oswald'; font-weight: 500; font-size: 14pt;
         letter-spacing: .08em; text-transform: uppercase; padding: 8pt 0 9pt; border-radius: 4pt; }
.cband em { font-style: normal; color: ACCENT; }
.slot { display: inline-block; border: 1.2pt solid #000; border-radius: 4pt; height: 26pt; min-width: 44pt; margin-right: 5pt; vertical-align: middle; }
.lab { font-family: 'Oswald'; font-weight: 500; font-size: 9pt; text-transform: uppercase; letter-spacing: .05em; text-align: center; }
""".replace("ACCENT", ACCENT)

COPY = f"© {YEAR} {STORE} · {BRAND} · For single-classroom use"


def html_doc(pages, title="doc"):
    return (f'<!doctype html><html><head><meta charset="utf-8"><title>{escape(title)}</title>'
            f'<style>{font_css()}{CSS}</style></head><body>{"".join(pages)}</body></html>')


def foot(right):
    return f'<div class="foot"><span>{COPY}</span><span>{right}</span></div>'


def wm(on):
    return '<div class="wm">PREVIEW</div>' if on else ""


def page(sub, tag, tagsub, inner, n, watermark=False, name=False, ak=False, cls=""):
    top = (f'<div class="top"><div class="t">{escape(TITLE)}<small>{sub}</small></div>'
           f'<div class="tag"><div class="lv{" ak" if ak else ""}">{tag}</div><div class="ln">{tagsub}</div></div></div>')
    nm = '<div class="name">Name<span></span>Date<span class="d"></span></div>' if name else ""
    return (f'<section class="page {cls}">{wm(watermark)}{top}{nm}<div class="main">{inner}</div>{foot(f"Page {n}")}</section>')


def checkbox(size=12):
    return (f'<svg width="{size}pt" height="{size}pt" viewBox="0 0 {size} {size}" style="flex:0 0 auto;margin-top:1pt">'
            f'<rect x=".7" y=".7" width="{size - 1.4}" height="{size - 1.4}" rx="2" fill="#fff" stroke="#000" stroke-width="1.2"/></svg>')


def t_(x, y, s, size, weight=700, family="Oswald", anchor="middle", fill="#000", extra=""):
    fam = f"'{family}'" if " " in family else family
    return (f'<text x="{x:.1f}" y="{y:.1f}" font-family="{fam}" font-weight="{weight}" font-size="{size}" text-anchor="{anchor}" '
            f'fill="{fill}" {extra}>{escape(s)}</text>')


def svgwrap(vw, vh, body, width="100%"):
    return f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {vw} {vh}" width="{width}" style="display:block">{body}</svg>'


CSS += """
table.sh { border-collapse: collapse; width: 100%; font-size: 9.6pt; }
table.sh td, table.sh th { border: 1pt solid #000; padding: 0 4pt; vertical-align: middle; }
table.sh th { background: #000; color: #fff; font-family: 'Oswald'; font-weight: 500; text-transform: uppercase; font-size: 8.6pt; letter-spacing: .04em; height: 17pt; }
table.sh td { height: 25.5pt; }
table.sh td.no { width: 8%; text-align: center; font-family: 'Oswald'; font-weight: 500; color: #555; font-size: 9pt; }
table.sh td.pre { color: #555; font-style: italic; }
.songh { display: block; border-bottom: 2pt solid #000; margin-bottom: 4pt; padding-bottom: 2pt; }
.songh b { display: block; white-space: nowrap; }
.songh > span { display: block; }
.songh b { font-family: 'Oswald'; font-weight: 700; font-size: 14pt; text-transform: uppercase; }
.songh span { font-family: 'Oswald'; font-weight: 300; font-size: 12pt; }
.pill { display: inline-block; font-family: 'Oswald'; font-weight: 500; font-size: 8.5pt; text-transform: uppercase; letter-spacing: .06em;
        border: 1pt solid #000; border-radius: 8pt; padding: 0 6pt; margin-left: 6pt; }
.pill.a { background: ACCENT; border-color: ACCENT; color: #fff; }
table.pl { border-collapse: collapse; width: 100%; }
table.pl td, table.pl th { border: 1.1pt solid #000; padding: 3pt 6pt; vertical-align: middle; font-size: 10.5pt; }
table.pl th { background: #000; color: #fff; font-family: 'Oswald'; font-weight: 500; text-transform: uppercase; font-size: 10pt; letter-spacing: .03em; }
table.pl td.s { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 11pt; width: 19%; }
""".replace("ACCENT", ACCENT)


# ---------------------------------------------------------------- drawings (original, code-drawn: black + accent only)
def arr_svg(width="100%", hero=False):
    """The worked example as a Logic-style arrangement: section markers on top, one lane per element, regions per section,
    and a texture count (elements playing) under each section. Chord-layer regions in the accent (the unit's new technique)."""
    lw = 0 if hero else 190
    W = 1000
    mw = (W - lw - 4) / EX_MEASURES
    X = lambda m: lw + (m - 1) * mw
    lane_h = 52 if hero else 44
    top = 8
    mk_h = 54 if hero else 44
    ruler = top + mk_h + 8
    y0 = ruler + (6 if hero else 26)
    H = y0 + len(LANES) * lane_h + (58 if not hero else 10)
    b = []
    for name, a, z in ex_starts():
        chorus = name.startswith("Chorus")
        b.append(f'<rect x="{X(a) + 2:.1f}" y="{top}" width="{(z - a + 1) * mw - 4:.1f}" height="{mk_h}" rx="7" fill="{ACCENT if chorus else "#000"}"/>')
        fs = (26 if hero else 21) if (z - a + 1) >= 8 else (20 if hero else 17)
        b.append(t_((X(a) + X(z + 1)) / 2, top + mk_h / 2 + fs * 0.36, name.upper(), fs, 700, fill="#fff"))
        if not hero:
            b.append(t_(X(a) + 4, ruler + 12, f"m. {a}", 13, 600, "Source Sans 3", anchor="start"))
    for name, a, z in ex_starts():
        b.append(f'<line x1="{X(a):.1f}" y1="{ruler if not hero else y0}" x2="{X(a):.1f}" y2="{y0 + len(LANES) * lane_h:.1f}" stroke="#000" '
                 f'stroke-width="{1.6 if hero else 1.2}"/>')
    b.append(f'<line x1="{X(EX_MEASURES + 1):.1f}" y1="{y0}" x2="{X(EX_MEASURES + 1):.1f}" y2="{y0 + len(LANES) * lane_h:.1f}" stroke="#000" stroke-width="1.2"/>')
    for r, lane in enumerate(LANES):
        y = y0 + r * lane_h
        b.append(f'<line x1="{lw}" y1="{y + lane_h:.1f}" x2="{X(EX_MEASURES + 1):.1f}" y2="{y + lane_h:.1f}" stroke="#000" stroke-width=".6"/>')
        if not hero:
            b.append(t_(lw - 10, y + lane_h / 2 + 6, lane, 16, 500, anchor="end"))
        for a, z in ex_regions(lane):
            col = ACCENT if lane.startswith("Chord layer") else "#000"
            b.append(f'<rect x="{X(a) + 2:.1f}" y="{y + 6}" width="{(z - a + 1) * mw - 4:.1f}" height="{lane_h - 12}" rx="6" fill="{col}"/>')
            if not hero:   # a few "notes" inside each region, in white, so it reads as a MIDI/audio region
                n = int((z - a + 1) * 2)
                for k in range(n):
                    xx = X(a) + 6 + (k + 0.5) * ((z - a + 1) * mw - 12) / n
                    yy = y + lane_h / 2 + ((k * 7 + r * 3) % 5 - 2) * 3.2
                    b.append(f'<rect x="{xx - 3:.1f}" y="{yy - 1.6:.1f}" width="6" height="3.2" rx="1.2" fill="#fff" opacity=".85"/>')
    b.append(f'<rect x="{lw}" y="{y0}" width="{X(EX_MEASURES + 1) - lw:.1f}" height="{len(LANES) * lane_h}" fill="none" stroke="#000" stroke-width="1.8"/>')
    if not hero:
        yt = y0 + len(LANES) * lane_h + 26
        b.append(t_(lw - 10, yt + 6, "Texture", 16, 700, anchor="end"))
        for name, a, z in ex_starts():
            cnt = sum(1 for lane in LANES for ra, rz in ex_regions(lane) if ra <= a <= rz)
            cx = (X(a) + X(z + 1)) / 2
            for k in range(cnt):
                b.append(f'<circle cx="{cx - (cnt - 1) * 7 + k * 14:.1f}" cy="{yt}" r="5.2" fill="{ACCENT if name.startswith("Chorus") else "#000"}"/>')
            b.append(t_(cx, yt + 26, f"{cnt} element{'s' if cnt != 1 else ''}" + (" › 1" if name == "Outro" else ""), 13, 600, "Source Sans 3"))
    return svgwrap(W, H, "".join(b), width)


def lanes_in(section_lanes):
    short = {"Chords (piano)": "chords", "Chord layer (pad)": "chord layer", "Bass": "bass", "Melody": "melody", "Drums": "drums"}
    parts = []
    for lane in LANES:
        v = section_lanes.get(lane)
        if v is True:
            parts.append(short[lane])
        elif v:
            parts.append(f"{short[lane]} (first {v[1] - v[0] + 1} measures)")
    return ", ".join(parts)


def plan_table(filled=False, font=10.5, row_h=None):
    """The owner's planning table (Section / Bars / Elements playing / Transition into next section); filled = the worked example."""
    ex = {name: (bars, lanes, tr, how) for name, bars, lanes, tr, how in EXAMPLE}
    rows = []
    for r in PLAN_ROWS:
        h = f" style='height:{row_h}pt'" if row_h else ""
        if filled:
            key = r if r in ex else None
            if key:
                bars, lanes, tr, how = ex[key]
                cells = f"<td style='text-align:center;font-family:Oswald;font-weight:700;font-size:{font + 2}pt'>{bars}</td><td>{lanes_in(lanes)}</td><td><b>{tr}:</b> {how}</td>"
            else:
                cells = "<td style='text-align:center'>–</td><td class='small'>not used in this skeleton</td><td></td>"
        else:
            cells = "<td></td><td></td><td></td>"
        rows.append(f"<tr{h}><td class='s' style='font-size:{font + 0.5}pt'>{r}</td>{cells}</tr>")
    return (f"<table class='pl' style='font-size:{font}pt'><tr><th>Section</th><th style='width:8%'>Bars</th>"
            f"<th>Elements playing<br><span style='font-weight:300;text-transform:none'>chords / bass / melody / drums / chord layer</span></th>"
            f"<th style='width:31%'>Transition into next section</th></tr>{''.join(rows)}</table>").replace("<td>", f"<td style='font-size:{font}pt'>")


def song_table(i, filled=False):
    """One song on the Song Form sheet: the owner's labels pre-filled where the Blank sheet has them; filled = answer view."""
    title, artist, use, secs, pre = SONGS[i]
    rows = []
    for r in range(SHEET_ROWS):
        if filled and r < len(secs):
            nm, ms, shown = secs[r]
            rows.append(f"<tr><td class='no'>{r + 1}</td><td><b>{escape(nm)}</b></td><td style='text-align:center'>{shown or ms}</td><td></td></tr>")
        elif r < pre:
            rows.append(f"<tr><td class='no'>{r + 1}</td><td class='pre'>{escape(secs[r][0])}</td><td></td><td></td></tr>")
        else:
            rows.append(f"<tr><td class='no'>{r + 1}</td><td></td><td></td><td></td></tr>")
    pill = '<span class="pill a">as a class</span>' if use == "class" else '<span class="pill">my choice?</span>'
    return (f'<div><div class="songh"><b>Song {i + 1} · {escape(title)}</b><span>{escape(artist)}{pill}</span></div>'
            f'<table class="sh"><tr><th style="width:8%">#</th><th style="width:31%">Section</th><th style="width:15%">Measures</th>'
            f'<th>What’s playing · what changed</th></tr>{"".join(rows)}</table></div>')


# answer-key map codes: every bar gets one; legend printed on the key page
MAP_CODES = [("bridge/chorus", "BR/CH"), ("bridge a", "BR A"), ("bridge b", "BR B"), ("bridge", "BR"), ("pre-chorus", "PRE"),
             ("post chorus", "POST"), ("pop drop", "DROP"), ("intro", "IN"), ("verse", "V"), ("chorus", "CH"), ("outro", "OUT")]
MAP_LEGEND = [("IN", "intro"), ("V", "verse"), ("PRE", "pre-chorus"), ("CH", "chorus"), ("POST", "post chorus"), ("DROP", "pop drop"),
              ("BR", "bridge"), ("BR A", "bridge A"), ("BR B", "bridge B"), ("BR/CH", "bridge/chorus?"), ("OUT", "outro")]
MAP_W, MAP_BAR, MAP_IN, MAP_OUT, MAP_ROW, MAP_PAD, MAP_GAP = 1000, 34, 15, 14, 17, 5, 6
_OSWALD_ADV = {}


def text_w(s, size):
    """Advance width of s in Oswald 700 at size (viewBox units), from the font's own metrics."""
    if not _OSWALD_ADV:
        from fontTools.ttLib import TTFont
        f = TTFont(os.path.join(FONTS, "oswald-latin-700-normal.woff2"))
        cmap, hm = f.getBestCmap(), f["hmtx"].metrics
        _OSWALD_ADV["upm"] = f["head"].unitsPerEm
        _OSWALD_ADV.update({chr(c): hm[g][0] for c, g in cmap.items() if c < 0x250})
    return sum(_OSWALD_ADV[ch] for ch in s) * size / _OSWALD_ADV["upm"]


def map_code(nm):
    low = nm.lower()
    return next(c for k, c in MAP_CODES if low.startswith(k))


def map_layout(i, scale_total):
    """Every bar gets a code: inside the bar when it fits with padding, otherwise in a call-out row above with a tick to the bar.
    Returns (bars, labels, rows); label = dict(code, bar, inside, x0, x1, row)."""
    secs = SONGS[i][3]
    bars, x = [], 0.0
    for nm, ms, _ in secs:
        w = ms / scale_total * MAP_W
        bars.append((nm, x, w))
        x += w
    labels, rows = [], []
    for k, (nm, bx, bw) in enumerate(bars):
        code = map_code(nm)
        inside = text_w(code, MAP_IN) + 2 * MAP_PAD <= bw - 2
        tw = text_w(code, MAP_IN if inside else MAP_OUT)
        c = bx + bw / 2
        if not inside:
            c = min(max(c, tw / 2), MAP_W - tw / 2)
        lab = dict(code=code, bar=k, inside=inside, x0=c - tw / 2, x1=c + tw / 2, row=None)
        if not inside:
            r = 0
            while r < len(rows) and any(lab["x0"] < o["x1"] + MAP_GAP and o["x0"] < lab["x1"] + MAP_GAP for o in rows[r]):
                r += 1
            if r == len(rows):
                rows.append([])
            rows[r].append(lab)
            lab["row"] = r
        labels.append(lab)
    return bars, labels, len(rows)


def form_map_svg(i, scale_total, width="100%"):
    """Answer key: one song's form as a proportional bar (common scale across songs). Chorus blocks in the accent.
    Every bar is labelled with a short code (legend on the page): inside when it fits, else in a call-out row above."""
    bars, labels, nrows = map_layout(i, scale_total)
    top = nrows * MAP_ROW + (4 if nrows else 0)
    H = top + MAP_BAR
    out = []
    for (nm, bx, bw), lab in zip(bars, labels):
        low = nm.lower()
        col = ACCENT if low.startswith("chorus") else ("#000" if any(low.startswith(k) for k in ("verse", "intro", "outro", "bridge")) else "#777")
        out.append(f'<rect x="{bx + 1:.1f}" y="{top + 2}" width="{max(bw - 2, 1.5):.1f}" height="{MAP_BAR - 4}" rx="{min(4, bw / 4):.1f}" fill="{col}"/>')
        if lab["inside"]:
            out.append(t_(bx + bw / 2, top + MAP_BAR / 2 + 5.5, lab["code"], MAP_IN, 700, fill="#fff"))
        else:
            ty = (nrows - lab["row"]) * MAP_ROW - 4      # row 0 sits just above the bar
            cx = bx + bw / 2
            out.append(f'<line x1="{cx:.1f}" y1="{ty + 2:.1f}" x2="{cx:.1f}" y2="{top + 2}" stroke="#000" stroke-width="1"/>')
            out.append(t_((lab["x0"] + lab["x1"]) / 2, ty, lab["code"], MAP_OUT, 700, fill="#000"))
    return svgwrap(MAP_W, H, "".join(out), width)


def verify_maps():
    """No map label is clipped or overprinted: in-bar codes fit their bar with padding, call-outs stay on the map and never touch."""
    errs = []
    total = max(sum(m for _, m, _ in s[3]) for s in SONGS)
    for i, s in enumerate(SONGS):
        bars, labels, nrows = map_layout(i, total)
        if len(labels) != len(s[3]):
            errs.append(f"answer key {s[0]}: {len(s[3])} sections but {len(labels)} labels")
        for lab in labels:
            nm, bx, bw = bars[lab["bar"]]
            if lab["inside"] and not (bx + MAP_PAD <= lab["x0"] and lab["x1"] <= bx + bw - MAP_PAD):
                errs.append(f"answer key {s[0]}: '{lab['code']}' clipped in its {nm} bar")
            if lab["x0"] < 0 or lab["x1"] > MAP_W:
                errs.append(f"answer key {s[0]}: '{lab['code']}' runs off the map")
        outs = [l for l in labels if not l["inside"]]
        for a in outs:
            for b in outs:
                if a is not b and a["row"] == b["row"] and a["x0"] < b["x1"] + MAP_GAP and b["x0"] < a["x1"] + MAP_GAP:
                    errs.append(f"answer key {s[0]}: call-outs '{a['code']}' and '{b['code']}' overlap")
        if nrows > 2:
            errs.append(f"answer key {s[0]}: {nrows} call-out rows (max 2)")
        legend = {c for c, _ in MAP_LEGEND}
        for lab in labels:
            if lab["code"] not in legend:
                errs.append(f"answer key {s[0]}: code '{lab['code']}' not in the legend")
    return errs


def hero_svg(width="100%"):
    return arr_svg(width, hero=True)


# ---------------------------------------------------------------- page plan
PLAN = ["cover", "howto", "overview", "setup", "logicref", "pacing", "day12", "day34", "tn1", "tn2", "tn3", "tn4",
        "post", "proj1", "proj2", "proj3", "proj4", "sheet1", "sheet2", "sheet3", "compare", "vocab", "checklist", "rubric",
        "key", "example", "terms"]
IX = {k: i + 1 for i, k in enumerate(PLAN)}
PREVIEW_PICKS = ["pacing", "proj3", "sheet1", "rubric"]


def rng(a, b):
    return str(a) if a == b else f"{a}–{b}"


def day_slides(d):
    """The deck slides a day uses, from its blocks (the only source of truth for slide-to-day slots)."""
    sl = [s for blk in d["blocks"] if blk[3] for s in range(blk[3][0], blk[3][1] + 1)]
    return (min(sl), max(sl)) if sl else None


def slides_label(d):
    r = day_slides(d)
    return f"Intro deck slides {rng(*r)}" if r and r[0] != r[1] else (f"Intro deck slide {r[0]}" if r else "No slides")


def fill(s):
    return (s.replace("SHEETS", f"{IX['sheet1']}–{IX['sheet3']}").replace("COMPARE", str(IX["compare"]))
             .replace("TN2", str(IX["tn2"])).replace("PLAN", str(IX["proj3"])))


def H2(s, first=False):
    return f'<h2 class="s"{" style=\"margin-top:2pt\"" if first else ""}>{s}<span class="acc"> ·</span></h2>'


def cover_page():
    return (f'<section class="page cover"><div class="store">{BRAND}</div>'
            f'<div style="width:7.4in">{hero_svg()}</div>'
            f'<div class="ctitle"><div class="a">Song Structure in Logic Pro</div><div class="b">Song Form</div>'
            f'<div class="c">Digital Music Production 1 · Unit {UNIT} · verse, chorus, bridge · arrangement · intro deck</div></div>'
            f'<div class="cband">Grades 9–12 <em>·</em> {DAYS_N} days <em>·</em> {len(PLAN)} pages + {DECK_SLIDES}-slide deck</div></section>')


SONG_NOTE = ("Songs are named by title and artist only. Play them from your own YouTube or Spotify access; no audio, video or lyrics are included.")


def deck_day_line():
    parts = []
    for d in DAYS:
        r = day_slides(d)
        if r:
            parts.append(f"Day {d['n']} (slide{'s' if r[0] != r[1] else ''} {rng(*r)})")
    return ", ".join(parts[:-1]) + " and " + parts[-1]


def howto_page(n):
    parts = [("Teacher pages", "howto", "tn4", "How to use, overview, setup, Logic reference, pacing, day plans, teacher notes"),
             ("Student pages", "post", "rubric", "Day 1 post, project sheet, planning table, Song Form sheets, compare page, vocabulary, checklist, rubric"),
             ("Answer key + example", "key", "example", "Form maps for all six songs, quick-check answers, a worked 32-measure skeleton"),
             ("Intro deck (separate file)", None, None, f"{DECK_SLIDES} slides, .pptx + PDF, speaker notes on every slide")]
    trs = "".join(f"<tr><td><b>{a}</b></td><td style='white-space:nowrap'>{('pages ' + str(IX[f]) + '–' + str(IX[l])) if f else 'deck'}</td><td>{u}</td></tr>"
                  for a, f, l, u in parts)
    inner = (f'<div class="cols"><div>'
             f'{H2("What this is", True)}'
             f'<p>A complete {DAYS_N}-day song structure unit for high school music technology. Day 1 is a listening day: students map the form of real '
             f'pop songs on a Song Form sheet, two as a class and two on their own. Then they plan a form on paper and build it in Logic Pro: their own '
             f'chords, bass and melody arranged into at least 4 sections (intro, verse and chorus required), with a chord layer and a drum transition at '
             f'every section change. They present on Day {DAYS_N}.</p>'
             f'<p>This is Unit {UNIT} of my Digital Music Production 1 course. It works on its own too, as long as students have a short chord '
             f'progression and bassline to arrange (page {IX["setup"]}).</p>'
             f'{H2("Quick start")}<ol>'
             f'<li>Read the overview and pacing (pages {IX["overview"]}–{IX["pacing"]}) and the teacher notes (pages {IX["tn1"]}–{IX["tn4"]}).</li>'
             f'<li>Cue up the two class songs on YouTube or Spotify: Crazy (Gnarls Barkley) and Blinding Lights (The Weeknd). Check that students can '
             f'play the four choice songs on headphones.</li>'
             f'<li>Check that every student has a Logic file with chords and a bassline (page {IX["setup"]}).</li>'
             f'<li>Practice the 5-minute chord layer demo once (page {IX["tn2"]}).</li>'
             f'<li>Print or post the student pages (pages {IX["post"]}–{IX["rubric"]}).</li></ol>'
             f'{H2("The intro deck")}'
             f'<p>Every slide has a day: {deck_day_line()}. The day plans name the slide for each block. Every slide has speaker notes.</p>'
             f'<p class="small">Fonts: the deck uses Oswald and Source Sans 3 (free on Google Fonts, built into Google Slides). If PowerPoint or Keynote '
             f'shows a different font, install the two fonts or present the PDF copy, which shows the intended look.</p>'
             f'</div><div>'
             f'{H2("What’s inside", True)}<table class="t"><tr><th>Part</th><th>Pages</th><th>Use</th></tr>{trs}</table>'
             f'<div class="box acc-b" style="margin-top:6pt"><span class="k">About the songs</span><p style="margin:0">{SONG_NOTE} The answer key '
             f'lists only section names and measure counts.</p></div>'
             f'{H2("Make it fit your class")}'
             f'<ul><li><b>Less time?</b> Run it in 3 days as I first wrote it: fold Day 3 into Day 2’s build time and finish on the last day '
             f'with 20 minutes of work, then presentations (page {IX["tn1"]}).</li>'
             f'<li><b>More accessible / more advanced:</b> two variations on page {IX["tn4"]}.</li>'
             f'<li><b>Feedback board:</b> I use Padlet; any discussion board, LMS forum or paper gallery walk works.</li></ul>'
             f'</div></div>')
    return page("How to use this resource", "START HERE", "Teacher", inner, n)


def overview_page(n):
    inner = (f'<div class="cols"><div>'
             f'{H2("What students make", True)}'
             f'<p>Four completed Song Form tables and a written response about the patterns they found. Then a structured arrangement in Logic of their '
             f'own chords, bass and melody: at least 4 distinct sections (intro, verse and chorus required; bridge and outro encouraged), each at least 4 '
             f'measures, labeled with arrangement markers, with a chord layer for texture and a drum transition at every section change. The worked '
             f'example is a 32-measure skeleton (page {IX["example"]}).</p>'
             f'{H2("What this unit is really about")}'
             f'<p>This unit is a listening unit first and a building unit second. The Song Form sheet is not a worksheet to fill in: it is a tool for '
             f'training students to hear structure the way a producer hears it. Students who genuinely engage with the analysis will arrive at the '
             f'building phase already knowing what they want to do.</p>'
             f'{H2("Mixing concept: arrangement as mix")}'
             f'<p>Denser sections are fuller and louder; sparse sections have room and clarity. Every addition and removal of an element is audible and '
             f'intentional. The chord layer demo makes it concrete: same notes, thicker texture, a bigger chorus.</p>'
             f'</div><div>'
             f'{H2("Goals", True)}<ul>' + "".join(f"<li>{g}</li>" for g in GOALS) + '</ul>'
             f'{H2("At a glance")}<table class="t">'
             f'<tr><td><b>Grades</b></td><td>9–12 (no music reading needed)</td></tr>'
             f'<tr><td><b>Length</b></td><td>{DAYS_N} periods of about 53 minutes: listen and analyze, plan and build, differentiate, present</td></tr>'
             f'<tr><td><b>Type</b></td><td>Structured project: listening analysis, then arranging</td></tr>'
             f'<tr><td><b>Software</b></td><td>Logic Pro: arrangement markers, copy and repeat, a second instrument track, volume, export</td></tr>'
             f'<tr><td><b>Product</b></td><td>4 song analyses + a written response; an arrangement with 4+ sections; MP3 + compressed project</td></tr>'
             f'<tr><td><b>Assessment</b></td><td>5-category rubric, /25 (page {IX["rubric"]})</td></tr></table>'
             f'{H2("Standards")}'
             f'<p>Supports the National Core Arts Standards for Music (Music Technology, high school): analyzing the structure of music (MU:Re7.2), '
             f'organizing and developing musical ideas (MU:Cr2.1), refining work (MU:Cr3.1) and presenting it (MU:Pr6.1).</p>'
             f'</div></div>')
    return page("Unit overview", "OVERVIEW", f"DMP1 · Unit {UNIT}", inner, n)


def setup_page(n):
    srows = "".join(f"<tr><td><b>{escape(t)}</b></td><td>{escape(a)}</td><td>{'Class example, Day 1' if u == 'class' else 'Student choice (pick 2 of 4)'}</td></tr>"
                    for t, a, u, _, _ in SONGS)
    inner = (f'<div class="cols"><div>'
             f'{H2("You need", True)}<ul>'
             f'<li><b>A Mac with Logic Pro</b> and <b>headphones</b> for each student.</li>'
             f'<li><b>Each student’s Logic file with a chord progression and a bassline.</b> In my course these come from the previous unit. Students '
             f'add their melody on Day 2.</li>'
             f'<li><b>Projector and speakers</b> for the deck and the two class songs.</li>'
             f'<li><b>The six songs</b> from your own YouTube or Spotify access (table below); students play their two choices on headphones.</li>'
             f'<li><b>A place to submit</b> an MP3 and a .zip, and a <b>feedback board</b> for the three feedback questions.</li></ul>'
             f'{H2("Optional")}<ul>'
             f'<li>A small MIDI keyboard (an MPK Mini or similar) for adding the melody.</li>'
             f'<li>A Logic template with the arrangement markers already built (the “more accessible” variation, page {IX["tn4"]}).</li></ul>'
             f'{H2("The songs")}<table class="t"><tr><th>Song</th><th>Artist</th><th>Use</th></tr>{srows}</table>'
             f'<p class="small">{SONG_NOTE}</p>'
             f'</div><div>'
             f'<div class="box acc-b"><span class="k">No chords-and-bassline file yet?</span>'
             f'<p style="margin:0 0 3pt">This unit arranges material students already have. If your students are new to Logic, give them one extra period '
             f'(or homework) before Day 2 to make it:</p><ul style="margin:0">'
             f'<li>A 4-chord progression on a software instrument, one chord per measure. The worked example uses '
             f'{"–".join(EX_PROG)} in {EX_KEY} ({", ".join(EX_CHORDS)}).</li>'
             f'<li>A bassline that plays the root of each chord.</li>'
             f'<li>A simple drum beat (any kick, snare and hi-hat pattern).</li></ul>'
             f'<p class="small" style="margin:3pt 0 0">The melody is added on Day 2 as part of the project.</p></div>'
             f'{H2("Before Day 1")}'
             + "".join(f'<div class="chk">{checkbox()}<span>{c}</span></div>' for c in [
                 "Cue up Crazy and Blinding Lights; test the projector sound. Know where the first two section changes are.",
                 "Check that students can reach the four choice songs on headphones (or plan to play them for small groups).",
                 "Check that every student can open a Logic file with chords and a bassline.",
                 f"Practice the chord layer demo once (page {IX['tn2']}) and look at the worked skeleton (page {IX['example']}).",
                 "Print or post the student pages; set up the submission spot and the feedback board.",
                 f"Skim the answer key (page {IX['key']}) so you can check measure counts quickly on Day 1."]) +
             f'</div></div>')
    return page("Setup and requirements", "SETUP", "Teacher", inner, n)


LOGIC_REF = [("Arrangement markers", "Show the global tracks (G), then the Arrangement track. Click + to add a marker; click its name to pick Intro, "
              "Verse, Chorus, Bridge or Outro, or rename it. Drag a marker’s right edge to set the section length."),
             ("Move or copy a section", "Drag an arrangement marker to move the whole section with its regions. Option-drag the marker to copy the section."),
             ("Copy a region", "Option-drag a region to copy it, or select it and press Command-R to repeat it."),
             ("New track", "Track › New Tracks (Option-Command-N), Software Instrument."),
             ("Change the sound", "Open the Library (Y) and pick a patch for the selected track: a pad, an electric piano, a slow synth."),
             ("Chord layer", "Option-drag a chord region onto a new track below, then change that track’s instrument. Same notes, new texture."),
             ("Take an element out", "Delete its region in that section (or never copy it in). To silence a whole track while you listen, click its M button."),
             ("Loop a section", "Press C for Cycle and drag the cycle area over the section you are working on."),
             ("Volume by section", "Track faders are in the Mixer (X). Show automation (A) and draw volume points to make a section louder or softer."),
             ("Drum transitions", "Edit the last measure of a drum region (Piano Roll or Step Sequencer) for a fill, a build or a cut. "
              "With a Drummer track, raise the Fills knob for that region."),
             ("Export", "File › Bounce › Project or Section (Command-B), MP3. In Finder, Control-click the .logicx project and choose Compress.")]


def logicref_page(n):
    rows = "".join(f"<tr><td style='width:22%'><b>{a}</b></td><td>{b}</td></tr>" for a, b in LOGIC_REF)
    inner = (f'<p>Everything students need in Logic for this unit, in one table. The planning table (page {IX["proj3"]}) comes before any of it.</p>'
             f'<table class="t" style="font-size:11.2pt;line-height:1.38">{rows}</table>'
             f'<div class="cols" style="margin-top:6pt"><div class="box"><span class="k">Measure math</span>'
             f'<p style="margin:0">Sections are almost always 4 or 8 measures. The worked skeleton: '
             + " + ".join(str(s[1]) for s in EXAMPLE) + f' = {EX_MEASURES} measures. At 4 beats per measure that is {EX_MEASURES * 4} beats: '
             f'about a minute at 120–128 BPM.</p></div>'
             f'<div class="box"><span class="k">Version note</span>'
             f'<p style="margin:0" class="small">Menu names are from current Logic Pro versions; older versions may differ slightly. Logic Pro and Drummer '
             f'are trademarks of Apple Inc., named only to describe what the unit needs. This resource is not affiliated with Apple.</p></div></div>')
    return page("Logic quick reference", "SETUP", "Teacher", inner, n)


def pacing_page(n, watermark=False):
    rows = "".join(
        f"<tr><td style='white-space:nowrap'><b class='acc' style='font-family:Oswald;font-size:13pt'>DAY {d['n']}</b><br><span class='small'>{d['kind']}</span></td>"
        f"<td><b>{d['title']}</b><br><span class='small'>{d['goal']}</span></td>"
        f"<td style='font-size:9.6pt;line-height:1.4'>{'<br>'.join(f'{a} · {h}' for a, h, _, _ in d['blocks'])}</td>"
        f"<td style='font-size:9.6pt'>{slides_label(d)}</td></tr>" for d in DAYS)
    inner = (f'<p>Four class periods of about 53 minutes: a listening day, a planning and building day, a day to differentiate the sections, then '
             f'finishing and presenting. Each day’s full plan is on pages {IX["day12"]}–{IX["day34"]}.</p>'
             f'<table class="t"><tr><th style="width:14%">Day</th><th style="width:38%">Focus and goal</th><th style="width:30%">Minute blocks</th>'
             f'<th>Deck</th></tr>{rows}</table>'
             f'{H2("Student pages by day")}<table class="t">'
             f'<tr><td><b>Day 1</b></td><td>Day 1 post (p. {IX["post"]}), Song Form sheets (pp. {IX["sheet1"]}–{IX["sheet3"]}), Compare the forms and '
             f'written response (p. {IX["compare"]})</td></tr>'
             f'<tr><td><b>Day 2</b></td><td>Project sheet (pp. {IX["proj1"]}–{IX["proj2"]}), planning table (p. {IX["proj3"]}), steps (p. {IX["proj4"]})</td></tr>'
             f'<tr><td><b>Day 3</b></td><td>Steps 4–5 (p. {IX["proj4"]}), vocabulary (p. {IX["vocab"]})</td></tr>'
             f'<tr><td><b>Day 4</b></td><td>Submission checklist (p. {IX["checklist"]}), rubric for self-score and teacher score (p. {IX["rubric"]})</td></tr></table>'
             f'<div class="box" style="margin-top:8pt"><span class="k">Three days instead of four</span>'
             f'<p style="margin:0">My Teacher Notes run this unit in 3 days: Day 1 as here; Day 2 plans and builds, adds the chord layer and starts the '
             f'transitions; Day 3 gives 20 minutes to finish and submit, then presentations. The four-day plan keeps the same blocks and gives building '
             f'its own day.</p></div>')
    return page(f"Pacing guide · {DAYS_N} days", "PACING", "Teacher", inner, n, watermark)


def day_box(d):
    def blk(a, h, t, s):
        tag = f' <span class="small">[{"slide" if s[0] == s[1] else "slides"} {rng(*s)}]</span>' if s else ""
        return f'<div class="blk"><div class="tm">{a}</div><div><b class="h">{h}</b>{tag} · {fill(t)}</div></div>'
    return (f'<div class="day"><div class="dh"><b>Day {d["n"]} · {d["title"]}</b><span>{d["kind"]} · {slides_label(d)}</span></div>'
            f'<p class="goal"><b>Goal:</b> {d["goal"]}</p>{"".join(blk(*b) for b in d["blocks"])}</div>')


def day12_page(n):
    return page("Day-by-day plans · Days 1 and 2", "DAYS 1–2", "Listen + plan", day_box(DAYS[0]) + day_box(DAYS[1]), n)


def day34_page(n):
    extra = ('<div class="box"><span class="k">Circulating prompts (Day 3)</span><ul style="margin:0">'
             '<li>‘Play me your verse, then your chorus. What is in the chorus that is not in the verse?’</li>'
             '<li>‘Mute your melody. Does the chorus still feel bigger?’</li>'
             '<li>‘Which transition type is this: fill, breakdown, cut or build? Why that one here?’</li></ul></div>')
    return page("Day-by-day plans · Days 3 and 4", "DAYS 3–4", "Build + present", day_box(DAYS[2]) + day_box(DAYS[3]) + extra, n)


def tn1_page(n):
    runs = "".join(f'<li><b>{a}:</b> {b}</li>' for a, b in RUN_ANALYSIS)
    inner = (f'<div class="tn">{H2("What this unit is really about", True)}'
             f'<p>This unit is a listening unit first and a building unit second. The Song Form sheet is not a worksheet to fill in: it is a tool for '
             f'training students to hear structure the way a producer hears it. Students who genuinely engage with the analysis will arrive at the building '
             f'phase already knowing what they want to do. Students who rush through it will spend Day 2 guessing.</p>'
             f'<p>The central insight you are trying to land: sections feel different because of what is playing, not just because of a label. A chorus is '
             f'not called a chorus because someone decided to call it that. It is called a chorus because it is texturally fuller, emotionally bigger, and '
             f'more memorable than the verse. That is a compositional decision, not a naming decision.</p>'
             f'{H2("Day 1: how to run the class analysis")}'
             f'<p>The two class examples should be tight: 10 to 12 minutes each. Play the first 2–3 minutes of each song only, enough to get through at least '
             f'two verse-chorus cycles. Here is how to structure each one:</p><ul>{runs}</ul>'
             f'<p>The share-out at the end of Day 1 is important. Ask: what patterns did you notice across all four songs? Students should arrive at: '
             f'sections are almost always 4 or 8 bars, chorus is always fuller than verse, intros build in and outros strip back. These are the principles '
             f'they will apply the next day.</p>'
             f'<div class="box"><span class="k">How this kit differs from my classroom version</span><ul style="margin:0">'
             f'<li>My Teacher Notes pace the unit in 3 days; this kit uses the 4 days of my course plan and gives building its own day (page {IX["pacing"]}).</li>'
             f'<li>Students arrange chords and a bassline they made earlier; if they have none, see page {IX["setup"]}. The worked 32-measure skeleton '
             f'(page {IX["example"]}) is new.</li>'
             f'<li>The Song Form spreadsheet is printed as tables, with the written response and an optional comparison grid on one page '
             f'(page {IX["compare"]}). My LMS and Padlet are written as “your class page” and “the class feedback board”.</li></ul></div></div>')
    return page("Teacher notes · the big idea and Day 1", "TEACHER NOTES", "1 of 4", inner, n)


def tn2_page(n):
    rows = "".join(f"<tr><td><b>{escape(t)}</b></td><td>{escape(a)}</td><td><i>{'Class example' if u == 'class' else 'Independent option'}</i></td>"
                   f"<td>{SONG_WHY[t]}</td></tr>" for t, a, u, _, _ in SONGS)
    demo = "".join(f'<li><div class="n">{i + 1}</div><div>{s}</div></li>' for i, s in enumerate(LAYER_DEMO))
    checks = "".join(f'<li>{c}</li>' for c in PLAN_CHECK)
    inner = (f'<div class="tn">{H2("Suggested songs", True)}'
             f'<table class="t" style="font-size:9.8pt"><tr><th>Song</th><th>Artist</th><th>Use</th><th>Why it works</th></tr>{rows}</table>'
             f'<p class="small">{SONG_NOTE}</p>'
             f'{H2("Day 2: enforce the planning table")}'
             f'<p>Do not let any student open the arrangement view until their planning table is filled in. Walk around at the start of the work time and '
             f'check each one. What to look for:</p><ul>{checks}</ul>'
             f'<p>If a student has the same elements in every section, send them back to the sheet: look at the song you analyzed. What dropped out in the '
             f'verse? That is what you need to do here.</p>'
             f'{H2("The chord layer demo (5 minutes, start of Day 2)")}<ol class="num" style="font-size:11pt">{demo}</ol></div>')
    return page("Teacher notes · songs, planning, the chord layer", "TEACHER NOTES", "2 of 4", inner, n)


def tn3_page(n):
    cb = "".join(f'<p><b>{a}:</b> {b}</p>' for a, b in CALLBACKS)
    st = "".join(f'<p><b>{a}:</b> {b}</p>' for a, b in STRUGGLES)
    inner = (f'<div class="tn big">{H2("Callbacks to earlier units", True)}'
             f'<p class="small" style="font-size:10pt !important">If your students did these projects (all part of my DMP1 course), name them out loud.</p>{cb}'
             f'{H2("Common student struggles")}{st}</div>')
    return page("Teacher notes · callbacks and struggles", "TEACHER NOTES", "3 of 4", inner, n)


def tn4_page(n):
    inner = (f'<div class="tn">{H2("Suggested variations", True)}'
             f'<p><b>More accessible:</b> Pre-build the arrangement markers in the Logic file and ask students to drag elements in rather than build from '
             f'scratch. Reduce the minimum to 3 sections: intro, verse, chorus.</p>'
             f'<p><b>More advanced:</b> Require a bridge that uses a chord not in the main progression (specifically the vi or iv) held for 8 bars before '
             f'the final chorus. Require all four transition types to appear at least once.</p>'
             f'{H2("Grading tips")}<ul>'
             f'<li>Each category is scored 2–5, for a total out of 25. Students self-score first; you score in the second column (page {IX["rubric"]}).</li>'
             f'<li>Song Form Analysis: check the measure counts against the answer key (page {IX["key"]}). Small differences (a pickup counted, a '
             f'section split in two) are fine when the student clearly counted.</li>'
             f'<li>Listen once eyes-closed for Section Differentiation: can you name each section by ear? Then open the project to check the markers, '
             f'section lengths and the chord layer track.</li>'
             f'<li>The presentation question is the quickest evidence of intention: point to your chorus; why does it feel different from your verse?</li></ul>'
             f'{H2("Resources")}<ul>'
             f'<li>The intro deck (.pptx and PDF), included.</li>'
             f'<li>The six songs, from your own YouTube or Spotify access (page {IX["setup"]}).</li>'
             f'<li>The worked skeleton (page {IX["example"]}) to model on Day 2 or to play from while you circulate.</li></ul>'
             f'{H2("Observation notes")}'
             f'<p class="small">What worked, what to adjust, and anything worth remembering for next year.</p></div><div class="lines"></div>')
    return page("Teacher notes · variations, grading, resources", "TEACHER NOTES", "4 of 4", inner, n)


def post_page(n):
    c, b = SONGS[0], SONGS[1]
    steps = [("Listen as a class: song 1", f"Your teacher will play {c[0]} by {c[1]}. Listen first without doing anything. Raise your hand when you feel "
              f"the section change.", "Do not try to name it yet. Just notice when the energy or texture shifts."),
             ("Label song 1 together", "Play it again and label every section as a class. Count the measures. What elements are playing in each section? "
              "What drops out? What comes in?", "Key question: why does the chorus feel bigger than the verse? What specifically is different?"),
             ("Listen as a class: song 2", f"Your teacher will play {b[0]} by {b[1]}. Same process, but this time you are leading more of the labeling.",
              "Compare to song 1: same structure? What is the same across both songs?"),
             ("Independent analysis: your two songs", f"Turn to the Song Form sheets (pages {IX['sheet2']}–{IX['sheet3']}). Choose two songs from the "
              "list and complete a table for each. Label every section, count the measures, note what elements are playing.",
              "The more carefully you do this, the easier Day 2 will be. These observations become your blueprint."),
             ("Written response", f"On the Compare the Forms page (page {IX['compare']}), write two or three sentences: what patterns did you notice "
              "across all four songs? What showed up every time?", "There is no single right answer, but there are patterns. Look for them.")]
    lis = "".join(f'<li style="margin-bottom:9pt"><div class="n">{i + 1}</div><div><b class="hd" style="display:inline;font-size:12.5pt">{a}</b> · {bb}<br>'
                  f'<span class="why">{cc}</span></div></li>' for i, (a, bb, cc) in enumerate(steps))
    choices = "".join(f'<p style="margin:0 0 3pt;font-size:11pt"><b>{escape(s[0])} · {escape(s[1])}</b> &nbsp;<span class="why">{SONG_HINT[s[0]]}</span></p>'
                      for s in SONGS if s[2] == "choice")
    inner = (f'<h2 class="s" style="font-size:20pt;margin:4pt 0 2pt">Song Form · Day 1</h2>'
             f'<p class="l" style="font-size:12pt">Digital Music Production 1 · Unit {UNIT}</p>'
             f'<p style="font-size:11.5pt;margin:6pt 0 4pt">Today is a listening day. You are not building anything in Logic. You are training your ear '
             f'to hear how songs are put together.</p>'
             f'<p style="font-size:11.5pt;margin:0 0 10pt"><b>All four Song Form tables must be complete before you leave today.</b></p>'
             f'<ol class="num" style="font-size:11.5pt">{lis}</ol>'
             f'<div class="box acc-b" style="margin-top:2pt"><span class="k" style="font-size:12pt">Choose two of these for your independent analysis</span>'
             f'{choices}</div>')
    return page("Day 1 classroom post", "STUDENT", "Day 1", inner, n)


def proj1_page(n):
    jobs = "".join(f"<tr><td style='width:17%'><b>{a}</b></td><td>{b}</td></tr>" for a, b in SECTION_JOBS)
    inner = (f'<h2 class="s" style="font-size:20pt;margin:2pt 0 0">Project sheet</h2>'
             f'<p class="l" style="font-size:12pt;margin-bottom:6pt">Digital Music Production 1 · Unit {UNIT} · Structured · {DAYS_N} Days · Arrangement as Mix</p>'
             f'{H2("Overview")}<p style="font-size:11.5pt">{OVERVIEW}</p>'
             f'{H2("Goals")}<ul>' + "".join(f"<li style='font-size:11pt'>{g}</li>" for g in GOALS) + '</ul>'
             f'{H2("What each section does")}<p style="font-size:11pt">A section is not just a label. It has a job. Understanding the job tells you how to '
             f'build it.</p><table class="t" style="font-size:10.2pt">{jobs}</table>')
    return page("Project sheet · overview, goals, sections", "STUDENT", "Project · 1 of 4", inner, n)


def proj2_page(n):
    inner = (f'{H2("Requirements", True)}'
             + "".join(f'<div class="chk" style="font-size:11.2pt;margin-bottom:6pt">{checkbox(12)}<span>{r}</span></div>' for r in REQUIREMENTS) +
             f'{H2("Tips for success")}<ul>' + "".join(f"<li style='font-size:11.2pt;margin-bottom:5pt'>{t}</li>" for t in TIPS) + '</ul>'
             f'<div class="box inv" style="margin-top:6pt"><span class="k">The test</span>'
             f'<p style="margin:0;font-size:13pt">Play it from start to finish. Can you hear which section you are in without looking at the screen?</p></div>')
    return page("Project sheet · requirements and tips", "STUDENT", "Project · 2 of 4", inner, n)


def proj3_page(n, watermark=False):
    slots = "".join('<span class="slot" style="min-width:58pt"></span>' + ('<b style="font-size:14pt;margin-right:5pt">–</b>' if i < 3 else '')
                    for i in range(4))
    inner = (f'<p style="font-size:11.5pt"><b>Fill this in before you open Logic.</b> Every section needs a job, a length, and a plan for what is playing.</p>'
             f'<div class="cols" style="margin:4pt 0 8pt"><div class="box" style="flex:0 0 2.6in"><span class="k">My key</span>'
             f'<div class="chk" style="margin:0 0 3pt">{checkbox()}<span>C major</span></div><div class="chk" style="margin:0 0 3pt">{checkbox()}<span>A minor</span></div>'
             f'<div class="fl" style="margin:0">{checkbox()}<b>Other:</b><span></span></div></div>'
             f'<div class="box"><span class="k">Chord progression (Roman numerals)</span><div style="margin-top:6pt">{slots}</div></div></div>'
             f'{plan_table(row_h=62)}'
             f'<p class="small" style="margin-top:5pt">Each section at least 4 measures. Transitions: fill, breakdown, cut or build. Worked example: page {IX["example"]}.</p>')
    return page("Project sheet · plan your song form", "STUDENT", "Project · 3 of 4", inner, n, watermark, name=True)


def proj4_page(n):
    lis = "".join(f'<li style="margin-bottom:9pt"><div class="n">{i + 1}</div><div style="font-size:11.4pt"><b class="hd">{d} · {h}</b>{t}</div></li>'
                  for i, (d, h, t) in enumerate(STEPS))
    inner = (f'{H2("Steps / process", True)}<ol class="num">{lis}</ol>'
             f'{H2("Notes")}<div class="lines"></div>')
    return page("Project sheet · steps", "STUDENT", "Project · 4 of 4", inner, n, name=True)


def sheet_page(n, idx, sub, intro, watermark=False):
    inner = (f'<p style="font-size:10.8pt;margin:0 0 6pt">{intro}</p>'
             f'<div class="cols" style="gap:16pt">{song_table(idx[0])}{song_table(idx[1])}</div>'
             f'<p class="small" style="margin-top:6pt">Grey labels were filled in for you. For each section: name it, count the measures, and write what is '
             f'playing or what changed (what came in, what dropped out).</p>')
    return page("Song Form in pop music", "SONG FORM", sub, inner, n, watermark, name=True)


def compare_page(n):
    rows = "".join(f"<tr style='height:40pt'><td><b>{lab}</b></td><td></td><td></td><td></td><td></td><td></td></tr>"
                   for lab in ["1 · Crazy", "2 · Blinding Lights", "3 · ______________", "4 · ______________", "My song (after Day 3)"])
    inner = (f'<p style="font-size:11pt">Line your four songs up side by side. The grid is optional: use it to find the patterns, then write your response.</p>'
             f'<table class="t" style="font-size:10pt"><tr><th style="width:22%">Song</th><th>How many sections?</th><th>Most common length</th>'
             f'<th>Pre-chorus?</th><th>Bridge?</th><th style="width:26%">What makes the chorus bigger?</th></tr>{rows}</table>'
             f'{H2("Written response")}'
             f'<p style="font-size:11.5pt">What patterns did you notice across all four songs? What showed up every time? Write two or three sentences.</p>'
             f'<p class="why" style="font-size:10.5pt">There is no single right answer, but there are patterns. Look for them.</p>'
             f'<div class="lines" style="flex:1"></div>')
    return page("Compare the forms · written response", "SONG FORM", "Day 1 · response", inner, n, name=True)


def vocab_page(n):
    rows = "".join(f"<tr><td style='width:18%'><b>{a}</b></td><td>{b}</td><td style='width:24%'></td></tr>" for a, b in VOCAB)
    qz = "".join(f'<div class="fl" style="margin-bottom:6pt"><b>{i + 1}. {q}</b><span style="flex:0 0 1.2in"></span></div>' for i, (q, _) in enumerate(VOCAB_QUIZ))
    inner = (f'<p style="font-size:11pt">Read each word and its meaning. In the last column, write where it shows up in your own song.</p>'
             f'<table class="t" style="font-size:10.4pt;line-height:1.35"><tr><th>Word</th><th>Meaning</th><th>In my song</th></tr>{rows}</table>'
             f'<div class="box" style="margin-top:6pt"><span class="k">Quick check · answers on page {IX["key"]}</span>{qz}</div>')
    return page("Key vocabulary", "STUDENT", f"{len(VOCAB)} words", inner, n, name=True)


def checklist_page(n):
    inner = (f'<p style="font-size:11.5pt">Tick each box before you submit and present. If a box isn’t ticked, you aren’t done yet.</p>'
             + "".join(f'<div class="chk" style="font-size:11.5pt;margin-bottom:7pt">{checkbox(14)}<span>{c}</span></div>' for c in CHECKLIST) +
             f'{H2("My three feedback questions")}'
             f'<p class="small">Ask about something you really want to know. Example: “Can you tell when my chorus starts without looking?”</p>'
             f'<div class="fl"><b>1.</b><span></span></div><div class="fl"><b>2.</b><span></span></div><div class="fl"><b>3.</b><span></span></div>'
             f'{H2("Notes")}<div class="lines"></div>')
    return page("Submission checklist", "STUDENT", "Before you present", inner, n, name=True)


def rubric_page(n, watermark=False):
    head = "<tr><th>Category</th>" + "".join(f"<th>{lv}</th>" for lv in LEVELS) + "<th>Self</th><th>Teacher</th></tr>"
    rows = "".join(f"<tr><td class='c'>{c}</td>" + "".join(f"<td>{d}</td>" for d in ds) + "<td class='sc'></td><td class='sc'></td></tr>"
                   for c, ds in RUBRIC)
    total = (f"<tr><td class='c' colspan='5' style='text-align:right;font-size:13pt'>Total</td>"
             f"<td style='font-family:Oswald;font-weight:700;font-size:13pt;text-align:right;vertical-align:bottom'>/ {5 * len(RUBRIC)}</td>"
             f"<td style='font-family:Oswald;font-weight:700;font-size:13pt;text-align:right;vertical-align:bottom'>/ {5 * len(RUBRIC)}</td></tr>")
    inner = (f'<p style="font-size:11pt">Score each category from 2 to 5. Score yourself first, then your teacher scores.</p>'
             f'<table class="rub">{head}{rows}{total}</table>'
             f'<div class="cols" style="margin-top:10pt"><div><h2 class="s" style="margin-top:0">My chorus is bigger because…</h2><div class="ln2"></div><div class="ln2"></div></div>'
             f'<div><h2 class="s" style="margin-top:0">Teacher comments</h2><div class="ln2"></div><div class="ln2"></div></div></div>')
    return page("Rubric", "RUBRIC", f"/ {5 * len(RUBRIC)}", inner, n, watermark, name=True)


def key_page(n):
    total = max(sum(m for _, m, _ in s[3]) for s in SONGS)
    maps = "".join(
        f'<div style="margin:0 0 7pt"><div style="display:flex;justify-content:space-between;align-items:baseline">'
        f'<b style="font-family:Oswald;font-size:12.5pt;text-transform:uppercase">{i + 1} · {escape(t)} <span style="font-weight:300;text-transform:none">· '
        f'{escape(a)}</span></b><span class="small">{len(secs)} sections</span></div>{form_map_svg(i, total)}'
        f'<p style="margin:2pt 0 0;font-size:9.4pt;line-height:1.3">' + " · ".join(f"{escape(nm)} <b>{shown or ms}</b>" for nm, ms, shown in secs) + '</p></div>'
        for i, (t, a, u, secs, _) in enumerate(SONGS))
    vq = " &nbsp;·&nbsp; ".join(f"<b>{i + 1}.</b> {a}" for i, (_, a) in enumerate(VOCAB_QUIZ))
    inner = (f'<p style="font-size:10pt;margin:0 0 5pt">Section names and measure counts from my classroom analysis (bars to scale; '
             f'<span class="acc"><b>violet</b></span> = chorus, black = intro, verse, bridge and outro, grey = other sections). Accept small differences '
             f'when a student clearly counted, for example a pickup measure or a section split in two. The label “Bridge/Chorus?” is my own: that section '
             f'works either way, so accept both. The teacher notes suggest Anti-Hero for introducing the pre-chorus, but my analysis does not list one '
             f'separately; accept a pre-chorus if a student splits one out of the end of a verse.</p>'
             f'<p style="font-size:9.4pt;margin:0 0 6pt"><b>Bar codes:</b> ' + " · ".join(f"<b>{c}</b> {d}" for c, d in MAP_LEGEND) + '. A code above the map with a short line belongs to the narrow bar the line points to.</p>'
             f'{maps}'
             f'<div class="cols" style="gap:14pt;margin-top:2pt"><div class="box" style="padding:5pt 9pt"><span class="k">Vocabulary · quick check</span>'
             f'<p style="margin:0;font-size:10pt">{vq}</p></div>'
             f'<div class="box" style="padding:5pt 9pt"><span class="k">Share-out patterns</span><p style="margin:0;font-size:10pt">'
             + "; ".join(PATTERNS) + '.</p></div></div>')
    return page("Answer key · form maps", "ANSWER KEY", "Teacher", inner, n, ak=True)


def example_page(n):
    inner = (f'<p style="font-size:10.2pt">A 32-measure skeleton to model on Day 2 or to play from while you circulate: chords {"–".join(EX_PROG)} in '
             f'{EX_KEY} ({" ".join(EX_CHORDS)}), one chord per measure. It is one possible answer, not the answer.</p>'
             f'{arr_svg()}'
             f'<div style="margin-top:6pt">{plan_table(filled=True, font=9.6)}</div>'
             f'<div class="cols" style="gap:14pt;margin-top:8pt"><div><span class="k">Compare it to the songs</span><p style="font-size:9.8pt;margin:0">'
             f'Like Crazy and Blinding Lights, every section is 4 or 8 measures and the chorus is the fullest part. There is no pre-chorus: the build in m. 8 '
             f'does that job. The verses are 4 measures because this is a skeleton; double them for a full song. The intro starts with one element and the '
             f'outro strips back to one.</p></div>'
             f'<div><span class="k">What a 5 sounds like</span><p style="font-size:9.8pt;margin:0">A listener can name every section by ear; the chorus '
             f'still feels bigger with the melody muted (the pad and full drums carry it); every change has a transition you can name; the denser sections '
             f'sit a little louder in the mix, and nothing clips.</p></div></div>')
    return page("Example notes · a worked 32-measure skeleton", "EXAMPLE", "Teacher", inner, n, ak=True)


def terms_page(n):
    return (f'<section class="page"><div class="top"><div class="t">Terms of Use<small>Thank you for your purchase!</small></div></div>'
            f'<div class="main"><p>© {YEAR} {STORE}. {BRAND} by {STORE}. All rights reserved. This purchase gives <b>one teacher</b> a license to use this '
            f'resource (the PDF and the intro deck) with their own students.</p>'
            f'<h2 class="s">You may</h2><ul><li>Print and copy pages for your own classroom and students, year after year.</li>'
            f'<li>Present the intro deck to your own classes and edit it for your own teaching.</li>'
            f'<li>Share pages with your own students through a password-protected class site or learning platform.</li>'
            f'<li>Leave the pages for a substitute teacher who teaches your own students.</li></ul>'
            f'<h2 class="s">You may not</h2><ul><li>Share, email or copy this resource for other teachers, a whole school or a district. '
            f'Additional licenses are available on Teachers Pay Teachers.</li>'
            f'<li>Post any part of it on a public website, shared drive or social media.</li>'
            f'<li>Sell, give away or claim any part of it as your own, or edit it into a new product.</li></ul>'
            f'<h2 class="s">Questions?</h2><p>Please use the Q&amp;A tab on the {BRAND} by {STORE} store on Teachers Pay Teachers. '
            f'If something looks wrong, let me know there and I will fix it.</p>'
            f'<h2 class="s">Credits</h2><ul>'
            f'<li>Curriculum, project sheet, rubric, teacher notes, Song Form sheets and intro deck: original work by {STORE}, {BRAND}.</li>'
            f'<li>Arrangement and form-map drawings, worked example and layouts: original work, {BRAND}.</li>'
            f'<li>Text fonts: Oswald (The Oswald Project Authors) and Source Sans 3 (Adobe), SIL Open Font License 1.1.</li>'
            f'<li>No audio, video, lyrics or album art are included. Songs are named for reference only (title and artist), and the answer key lists only '
            f'section names and measure counts; play the songs from your own YouTube or Spotify access. Song titles belong to their owners; this resource '
            f'is not affiliated with or endorsed by the artists or their labels.</li>'
            f'<li>Logic Pro and Drummer are trademarks of Apple Inc., named only to describe compatibility; this resource is not affiliated with or '
            f'endorsed by Apple. Padlet, YouTube and Spotify are named only as examples of tools you may already use.</li></ul>'
            f'<p class="small">{DISCLOSURE}</p><div style="flex:1"></div></div>{foot(f"Page {n}")}</section>')


SHEET_INTROS = {
    "sheet1": ([0, 1], "Class songs · Day 1", "Songs 1 and 2: we do these together. Listen, raise your hand at each section change, then label the sections and "
                                               "count the measures as a class."),
    "sheet2": ([2, 3], "Your choice · 1 of 2", "Songs 3–6: choose two of these four and complete a table for each. Play them on headphones from YouTube or "
                                              "Spotify."),
    "sheet3": ([4, 5], "Your choice · 2 of 2", "Choose two of songs 3–6 in total. No labels are filled in on these two: you name every section yourself."),
}


def render_page(key, watermark=False):
    n = IX[key]
    if key in SHEET_INTROS:
        idx, sub, intro = SHEET_INTROS[key]
        return sheet_page(n, idx, sub, intro, watermark)
    f = {"cover": cover_page, "howto": lambda: howto_page(n), "overview": lambda: overview_page(n), "setup": lambda: setup_page(n),
         "logicref": lambda: logicref_page(n), "pacing": lambda: pacing_page(n, watermark), "day12": lambda: day12_page(n),
         "day34": lambda: day34_page(n), "tn1": lambda: tn1_page(n), "tn2": lambda: tn2_page(n), "tn3": lambda: tn3_page(n),
         "tn4": lambda: tn4_page(n), "post": lambda: post_page(n), "proj1": lambda: proj1_page(n), "proj2": lambda: proj2_page(n),
         "proj3": lambda: proj3_page(n, watermark), "proj4": lambda: proj4_page(n), "compare": lambda: compare_page(n),
         "vocab": lambda: vocab_page(n), "checklist": lambda: checklist_page(n), "rubric": lambda: rubric_page(n, watermark),
         "key": lambda: key_page(n), "example": lambda: example_page(n), "terms": lambda: terms_page(n)}
    return f[key]()


# ---------------------------------------------------------------- intro deck (.pptx, python-pptx) + PDF copy (LibreOffice)
def slide_day(i):
    return SLIDES[i - 1][0]


def build_deck(img):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

    BLACK, WHITE, ACC = RGBColor(0, 0, 0), RGBColor(255, 255, 255), RGBColor.from_string(ACCENT[1:])
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.author = f"{STORE} · {BRAND}"
    prs.core_properties.title = f"{TITLE} · Intro deck"
    prs.core_properties.comments = DISCLOSURE
    layout = prs.slide_layouts[6]

    def box(sl, x, y, w, h, fillc=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=1.5):
        s = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        if fillc is None:
            s.fill.background()
        else:
            s.fill.solid()
            s.fill.fore_color.rgb = fillc
        if line is None:
            s.line.fill.background()
        else:
            s.line.color.rgb = line
            s.line.width = Pt(lw)
        s.shadow.inherit = False
        return s

    def text(sl, x, y, w, h, runs, size=24, font="Source Sans 3", bold=False, color=BLACK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             spacing=None, caps=False):
        tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.04)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        paras = runs if isinstance(runs, list) else [runs]
        for k, ptxt in enumerate(paras):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = ptxt.upper() if caps else ptxt
            r.font.size = Pt(size)
            r.font.name = font
            r.font.bold = bold
            r.font.color.rgb = color
            if spacing:
                r._r.get_or_add_rPr().set("spc", str(spacing))
        return tb

    def chrome(sl, i, title, day):
        text(sl, 0.6, 0.32, 6, 0.4, BRAND, 13, "Oswald", False, BLACK, spacing=600, caps=True)
        box(sl, 11.33, 0.3, 1.4, 0.42, ACC, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        text(sl, 11.33, 0.3, 1.4, 0.42, f"Day {day}", 15, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, caps=True)
        text(sl, 0.6, 0.72, 12.1, 0.95, title, 40, "Oswald", True, BLACK, caps=True, anchor=MSO_ANCHOR.BOTTOM)
        box(sl, 0.6, 1.72, 12.13, 0.04, BLACK)
        box(sl, 0.6, 7.0, 12.13, 0.012, BLACK)
        text(sl, 0.6, 7.04, 9, 0.3, f"{TITLE} · Intro deck · © {YEAR} {STORE} · {BRAND}", 10, "Source Sans 3", False, RGBColor(0x33, 0x33, 0x33))
        text(sl, 11.2, 7.04, 1.53, 0.3, f"{i}", 10, "Oswald", True, BLACK, align=PP_ALIGN.RIGHT)

    def bullets(sl, items, x=0.6, y=2.05, w=12.1, size=28, gap=1.75):
        for k, it in enumerate(items):
            yy = y + k * (size / 72 * gap)
            box(sl, x, yy + size / 72 * 0.36, 0.16, 0.16, ACC, shape=MSO_SHAPE.OVAL)
            text(sl, x + 0.38, yy, w - 0.4, size / 72 * 1.6, it, size, "Source Sans 3", False, BLACK)

    def pic(sl, key, x, y, w):
        path, (pw, ph) = img[key]
        sl.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))
        return y + w * ph / pw

    for i, (day, kind, title, body, notes) in enumerate(SLIDES, start=1):
        sl = prs.slides.add_slide(layout)
        items = [b for b in body.split("|") if b]
        if kind == "title":
            text(sl, 0, 0.42, 13.333, 0.5, BRAND, 16, "Oswald", False, BLACK, align=PP_ALIGN.CENTER, spacing=800, caps=True)
            pic(sl, "hero", 2.37, 1.0, 8.6)
            text(sl, 0, 4.05, 13.333, 1.3, title, 80, "Oswald", True, BLACK, align=PP_ALIGN.CENTER, caps=True)
            text(sl, 0, 5.45, 13.333, 0.6, body, 28, "Oswald", False, BLACK, align=PP_ALIGN.CENTER)
            box(sl, 0, 6.55, 13.333, 0.95, BLACK)
            text(sl, 0, 6.68, 13.333, 0.6, f"Digital Music Production · Unit {UNIT} · Day {day} · Song structure in Logic Pro", 20, "Oswald", False,
                 WHITE, align=PP_ALIGN.CENTER, spacing=200, caps=True)
        else:
            chrome(sl, i, title, day)
        if kind == "listen":
            song, _, artist = items[0].partition(" · ")
            box(sl, 0.6, 2.0, 12.13, 2.0, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            text(sl, 0.9, 2.1, 11.5, 0.7, song, 36, "Oswald", True, WHITE, caps=True)
            text(sl, 0.9, 2.82, 11.5, 0.5, artist, 24, "Oswald", False, RGBColor(0xC9, 0xB8, 0xF3))
            text(sl, 0.9, 3.38, 11.5, 0.5, items[1], 18, "Source Sans 3", False, WHITE)
            bullets(sl, items[2:], y=4.3, size=30, gap=1.6)
        elif kind == "bullets":
            bullets(sl, items, size=34 if len(items) <= 4 else 31, y=2.3, gap=2.0 if len(items) <= 4 else 1.85)
        elif kind == "closer":
            bullets(sl, items, size=36, y=2.15, gap=1.55)
            pic(sl, "hero", 3.42, 4.75, 6.5)
        elif kind == "sections":
            for k, (head, rest) in enumerate(SECTION_SHORT):
                col, row = k % 3, k // 3
                x, y = 0.6 + col * 4.14, 2.05 + row * 2.45
                chorus = head == "Chorus"
                box(sl, x, y, 3.85, 2.2, ACC if chorus else WHITE, None if chorus else BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=2.25)
                text(sl, x + 0.2, y + 0.18, 3.45, 0.7, head, 32, "Oswald", True, WHITE if chorus else ACC, align=PP_ALIGN.CENTER, caps=True)
                text(sl, x + 0.2, y + 0.95, 3.45, 1.15, rest, 22, "Source Sans 3", False, WHITE if chorus else BLACK, align=PP_ALIGN.CENTER)
        elif kind == "columns":
            for k, it in enumerate(items):
                x = 0.6 + k * 3.08
                head, _, rest = it.partition(":")
                box(sl, x, 2.15, 2.85, 4.55, WHITE, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=2.25)
                box(sl, x + 1.12, 2.45, 0.62, 0.62, ACC if k % 2 == 0 else BLACK, shape=MSO_SHAPE.OVAL)
                text(sl, x + 1.12, 2.45, 0.62, 0.62, str(k + 1), 24, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                text(sl, x + 0.15, 3.3, 2.55, 0.8, head, 34, "Oswald", True, BLACK, align=PP_ALIGN.CENTER, caps=True)
                text(sl, x + 0.2, 4.25, 2.45, 2.3, rest.strip(), 24, "Source Sans 3", False, BLACK, align=PP_ALIGN.CENTER)
        elif kind == "arr":
            yb = pic(sl, "arr", 0.6, 1.98, 12.13)
            text(sl, 0.6, yb + 0.08, 12.1, 0.45, "The worked example: one lane per element, one column per section. Violet = the chord layer.",
                 17, "Source Sans 3", False, BLACK)
        elif kind == "plan":
            pic(sl, "plan", 0.6, 1.98, 12.13)
        elif kind == "sheet":
            yb = pic(sl, "sheet", 0.6, 1.98, 12.13)
            text(sl, 0.6, yb + 0.06, 12.1, 0.4, "A finished table looks like this (class example). Now pick two of these. Find them on YouTube or Spotify:",
                 17, "Source Sans 3", False, BLACK)
            for k, it in enumerate(items):
                col, row = k % 2, k // 2
                x, y = 0.6 + col * 6.15, yb + 0.6 + row * 1.12
                song, _, artist = it.partition(" · ")
                box(sl, x, y, 5.98, 0.95, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                text(sl, x + 0.25, y, 3.6, 0.95, song, 28, "Oswald", True, WHITE, anchor=MSO_ANCHOR.MIDDLE, caps=True)
                text(sl, x + 3.7, y, 2.1, 0.95, artist, 20, "Oswald", False, RGBColor(0xC9, 0xB8, 0xF3), align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
        elif kind == "steps":
            for k, it in enumerate(items):
                y = 2.05 + k * 0.95
                box(sl, 0.6, y + 0.05, 0.62, 0.62, BLACK, shape=MSO_SHAPE.OVAL)
                text(sl, 0.6, y + 0.05, 0.62, 0.62, str(k + 1), 24, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                text(sl, 1.45, y - 0.05, 11.2, 0.9, it, 26, "Source Sans 3", False, BLACK, anchor=MSO_ANCHOR.MIDDLE)
        elif kind == "days":
            for k, d in enumerate(DAYS):
                y = 2.1 + k * 1.2
                box(sl, 0.6, y, 1.7, 0.95, ACC if k == 0 else BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                text(sl, 0.6, y, 1.7, 0.95, f"Day {d['n']}", 26, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, caps=True)
                text(sl, 2.55, y - 0.02, 10.1, 0.55, d["title"], 28, "Oswald", True, BLACK, caps=True)
                text(sl, 2.55, y + 0.5, 10.1, 0.5, DECK_DAY_LINES[k], 20, "Source Sans 3", False, BLACK)
        elif kind == "check":
            for k, it in enumerate(items):
                col, row = k % 2, k // 2
                x, y = 0.6 + col * 6.15, 2.15 + row * 1.55
                box(sl, x, y + 0.08, 0.42, 0.42, WHITE, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=2.25)
                text(sl, x + 0.6, y, 5.35, 1.4, it, 24, "Source Sans 3", False, BLACK)
        elif kind == "rubric":
            for k, ((c, _), short) in enumerate(zip(RUBRIC, RUBRIC_SHORT)):
                y = 2.05 + k * 0.92
                box(sl, 0.6, y, 4.3, 0.78, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                text(sl, 0.75, y, 4.0, 0.78, c, 18, "Oswald", True, WHITE, anchor=MSO_ANCHOR.MIDDLE, caps=True)
                text(sl, 5.1, y, 6.6, 0.78, "5 = " + short, 17, "Source Sans 3", False, BLACK, anchor=MSO_ANCHOR.MIDDLE)
                text(sl, 11.7, y, 1.03, 0.78, "/ 5", 22, "Oswald", True, ACC, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
            text(sl, 9.0, 6.6, 3.73, 0.4, f"Total / {5 * len(RUBRIC)}", 20, "Oswald", True, BLACK, align=PP_ALIGN.RIGHT, caps=True)
        sl.notes_slide.notes_text_frame.text = notes
    prs.save(DECK)


DECK_DAY_LINES = ["Listen to 4 songs and map their form on the Song Form sheet; write what you noticed",
                  "Add your melody, fill in the planning table, build the skeleton with arrangement markers",
                  "Chord layer, melody only where it belongs, a drum transition at every change",
                  "Finish the mix, export and submit, present: point to your chorus"]


def ttf_fonts():
    """woff2 -> ttf for LibreOffice (deck PDF), plus a fontconfig file that adds them."""
    from fontTools.ttLib import TTFont
    d = os.path.join(BUILD, "ttf")
    os.makedirs(d, exist_ok=True)
    styles = {"oswald-latin-500-normal": ("Oswald", "Regular", 400), "oswald-latin-700-normal": ("Oswald", "Bold", 700),
              "source-sans-3-latin-400-normal": ("Source Sans 3", "Regular", 400), "source-sans-3-latin-700-normal": ("Source Sans 3", "Bold", 700),
              "source-sans-3-latin-400-italic": ("Source Sans 3", "Italic", 400)}
    for stem, (fam, sub, wt) in styles.items():
        out = os.path.join(d, stem + ".ttf")
        if os.path.exists(out):
            continue
        t = TTFont(os.path.join(FONTS, stem + ".woff2"))
        t.flavor = None
        name = t["name"]
        ps = (fam.replace(" ", "") + "-" + sub)
        for rec in list(name.names):
            if rec.nameID in (16, 17, 21, 22):
                name.removeNames(nameID=rec.nameID)
        for nid, val in ((1, fam), (2, sub), (4, f"{fam} {sub}" if sub != "Regular" else fam), (6, ps)):
            name.setName(val, nid, 3, 1, 0x409)
            name.setName(val, nid, 1, 0, 0)
        t["OS/2"].usWeightClass = wt
        sel = 0
        if sub == "Bold":
            sel |= 1 << 5
        if sub == "Italic":
            sel |= 1
        if sub == "Regular":
            sel |= 1 << 6
        t["OS/2"].fsSelection = (t["OS/2"].fsSelection & ~0b1100001) | sel
        t["head"].macStyle = (1 if sub == "Bold" else 0) | (2 if sub == "Italic" else 0)
        t.save(out)
    conf = os.path.join(BUILD, "fonts.conf")
    with open(conf, "w") as fh:
        fh.write(f'<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig><include ignore_missing="yes">/etc/fonts/fonts.conf</include>'
                 f'<dir>{d}</dir><cachedir>{os.path.join(BUILD, "fccache")}</cachedir></fontconfig>')
    return conf


def deck_pdf(src=DECK, out=DECK_PDF, fonts=True):
    """fonts=False renders with system fonts only: the font-fallback view a teacher without Oswald would see (self-check)."""
    env = dict(os.environ, FONTCONFIG_FILE=ttf_fonts()) if fonts else dict(os.environ)
    prof = os.path.join(BUILD, "lo-profile" if fonts else "lo-profile-fallback")
    tmp = os.path.join(BUILD, "lo-out")
    os.makedirs(tmp, exist_ok=True)
    subprocess.run(["soffice", f"-env:UserInstallation=file://{prof}", "--headless", "--convert-to", "pdf", "--outdir", tmp, src],
                   check=True, env=env, capture_output=True)
    shutil.move(os.path.join(tmp, os.path.basename(src).replace(".pptx", ".pdf")), out)


# ---------------------------------------------------------------- square images (2000 x 2000)
SQ_CSS = """
.sq { width: 1000px; height: 1000px; position: relative; overflow: hidden; background: linear-gradient(180deg,#fff 0%,#fff 60%,#eeeeee 100%); }
.sq .store { position: absolute; left: 0; right: 0; top: 30px; text-align: center; font-family: 'Oswald'; font-weight: 500; letter-spacing: .32em;
             text-transform: uppercase; font-size: 22px; }
.sq .fillw { position: absolute; left: 0; right: 0; top: 66px; bottom: 66px; display: flex; flex-direction: column; align-items: center; justify-content: space-evenly; }
.sq .hd { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; text-align: center; line-height: .95; white-space: nowrap; }
.sq .sub { font-family: 'Oswald'; font-weight: 300; text-align: center; }
.sq .band { position: absolute; left: 0; right: 0; bottom: 0; background: #000; color: #fff; text-align: center; font-family: 'Oswald';
            font-weight: 500; text-transform: uppercase; letter-spacing: .1em; font-size: 26px; padding: 16px 0 18px; }
.sq .band em { font-style: normal; color: ACCENT; }
.shot { background: #fff; border: 2px solid #000; box-shadow: 8px 8px 0 #000; position: relative; overflow: hidden; }
.shot img { display: block; width: 100%; }
.lbl { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 22px; text-align: center; margin-top: 14px; }
.lbl span { font-weight: 300; text-transform: none; display: block; font-size: 18px; }
.swm { position: absolute; left: 50%; top: 50%; transform: translate(-50%,-50%) rotate(-38deg); font-family: 'Oswald'; font-weight: 700;
       color: rgba(0,0,0,.14); white-space: nowrap; letter-spacing: .06em; }
""".replace("ACCENT", ACCENT)
FONT_PRIME = ('<div style="position:absolute;left:-9999px;font-family:\'Source Sans 3\'">'
              '<span style="font-weight:400">a</span><span style="font-weight:600">a</span><i style="font-weight:400">a</i></div>')


def sq_doc(body):
    return (f'<!doctype html><html><head><meta charset="utf-8"><style>{font_css()}'
            f'html,body{{margin:0;background:#fff}}{SQ_CSS}</style></head><body>{FONT_PRIME}{body}</body></html>')


def sq(inner, band):
    return sq_doc(f'<div class="sq"><div class="store">{BRAND}</div><div class="fillw">{inner}</div><div class="band">{band}</div></div>')



def cover_square():
    return sq(f'<div style="width:960px">{hero_svg()}</div>'
              f'<div><div class="hd" style="font-size:58px">Song Structure in Logic Pro</div>'
              f'<div class="hd" style="font-size:168px;margin-top:4px">Song Form</div>'
              f'<div class="sub" style="font-size:31px;margin-top:10px">{DAYS_N}-day unit · verse, chorus, bridge · arrangement · intro slides</div></div>',
              f'Grades 9–12 <em>·</em> {DAYS_N} days <em>·</em> {len(PLAN)} pages + {DECK_SLIDES} slides')


def preview_square_1(img):
    rot = [-4, 0, 4]
    shots = "".join(f'<div style="position:absolute;left:{x}px;top:{y}px;width:372px;transform:rotate({r}deg);z-index:{z}">'
                    f'<div class="shot"><img src="{s}"><div class="swm" style="font-size:80px">PREVIEW</div></div></div>'
                    for (s, _, _), x, y, r, z in zip(img, (40, 314, 588), (24, 0, 24), rot, (1, 3, 2)))
    labels = "".join(f'<div class="lbl" style="width:300px">{a}<span>{b}</span></div>' for _, a, b in img)
    return sq(f'<div><div class="hd" style="font-size:76px">Ready to teach Monday</div>'
              f'<div class="sub" style="font-size:28px;margin-top:8px">Teacher notes, day plans, student pages and the rubric in one PDF</div></div>'
              f'<div style="position:relative;width:1000px;height:520px">{shots}</div>'
              f'<div style="display:flex;gap:30px">{labels}</div>',
              'Listen <em>·</em> plan <em>·</em> build <em>·</em> present')


def preview_square_2(slides):
    cells = "".join(f'<div class="shot" style="box-shadow:6px 6px 0 #000"><img src="{s}"><div class="swm" style="font-size:52px">PREVIEW</div></div>'
                    for s in slides)
    return sq(f'<div><div class="hd" style="font-size:76px">Intro deck included</div>'
              f'<div class="sub" style="font-size:28px;margin-top:8px">{DECK_SLIDES} editable slides (.pptx + PDF) with speaker notes on every slide</div></div>'
              f'<div style="display:grid;grid-template-columns:388px 388px;gap:20px 34px">{cells}</div>',
              'Hook <em>·</em> listen <em>·</em> worked example <em>·</em> steps <em>·</em> rubric')


def preview_square_3(imgs):
    cells = "".join(f'<div style="width:214px"><div class="shot" style="box-shadow:5px 5px 0 #000"><img src="{s}">'
                    f'<div class="swm" style="font-size:46px">PREVIEW</div></div><div class="lbl" style="font-size:18px;margin-top:9px">{lab}</div></div>'
                    for s, lab in imgs)
    return sq(f'<div><div class="hd" style="font-size:72px">Everything inside</div>'
              f'<div class="sub" style="font-size:28px;margin-top:8px">Song Form sheets, planning table, answer key and a worked 32-measure skeleton</div></div>'
              f'<div style="display:grid;grid-template-columns:repeat(4,214px);gap:22px 24px">{cells}</div>',
              f'{len(PLAN)} pages <em>·</em> US Letter PDF <em>·</em> prints in black &amp; white')


SHOTS1 = [("pacing", "Pacing + plans", "4 days, minute by minute"), ("sheet1", "Song Form sheet", "map real pop songs"),
          ("rubric", "Rubric", "5 categories, / 25")]
SHOTS2 = [1, 3, 4, 8, 10, 16]   # deck slides
SHOTS3 = [("post", "Day 1 post"), ("proj3", "Planning table"), ("sheet2", "Song Form sheet"), ("compare", "Compare the forms"),
          ("vocab", "Vocabulary"), ("checklist", "Checklist"), ("key", "Answer key"), ("example", "Worked skeleton")]


# ---------------------------------------------------------------- render
def node_env():
    env = dict(os.environ)
    if not env.get("NODE_PATH"):
        env["NODE_PATH"] = subprocess.run(["npm", "root", "-g"], capture_output=True, text=True).stdout.strip()
    return env


def render(jobs):
    path = os.path.join(BUILD, "jobs.json")
    with open(path, "w") as f:
        json.dump(jobs, f)
    subprocess.run(["node", RENDER_JS, path], check=True, env=node_env())


def write(name, txt):
    p = os.path.join(BUILD, name)
    with open(p, "w") as f:
        f.write(txt)
    return p


def pdf_png(src, page_no, out, res=110):
    subprocess.run(["pdftoppm", "-png", "-r", str(res), "-f", str(page_no), "-l", str(page_no), "-singlefile", src,
                    os.path.join(BUILD, out)], check=True)


def art_doc(w, h, inner, zoom=1.0):
    return (f'<!doctype html><html><head><meta charset="utf-8"><style>{font_css()}{CSS}'
            f'html,body{{margin:0;background:#fff}}</style></head>'
            f'<body><div style="width:{w}px;height:{h}px;overflow:hidden;background:#fff"><div style="zoom:{zoom}">{inner}</div></div></body></html>')


def sheet_strip_html():
    """Slide 6 art: the class example as one row of the owner's spreadsheet (Section / Measures), filled in."""
    t, a, _, secs, _ = SONGS[0]
    head = "".join(f"<td style='border:2px solid #000;padding:8px 6px;text-align:center;font-weight:700'>{escape(nm)}</td>" for nm, _, _ in secs)
    meas = "".join(f"<td style='border:2px solid #000;padding:8px 6px;text-align:center;font-family:Oswald;font-weight:700;font-size:30px'>"
                   f"{shown or ms}</td>" for _, ms, shown in secs)
    lab = "style='border:2px solid #000;padding:8px 10px;background:#000;color:#fff;font-family:Oswald;font-weight:500;text-transform:uppercase;font-size:22px'"
    return (f"<div style='font-family:\"Source Sans 3\";font-size:24px;padding:4px'>"
            f"<div style='font-family:Oswald;font-weight:700;font-size:30px;text-transform:uppercase;margin-bottom:8px'>{escape(t)} · "
            f"<span style='font-weight:300;text-transform:none'>{escape(a)}</span></div>"
            f"<table style='border-collapse:collapse;width:100%'><tr><td {lab}>Section</td>{head}</tr><tr><td {lab}>Measures</td>{meas}</tr></table></div>")


def deck_art():
    """PNG art for the deck (original drawings and tables, rendered with the brand fonts)."""
    arts = {"hero": (1720, 595, hero_svg()),
            "arr": (2000, 728, arr_svg()),
            "plan": (1600, 640, f'<div style="padding:4px">{plan_table(filled=True, font=21, row_h=50)}</div>'),
            "sheet": (1600, 210, sheet_strip_html())}
    jobs, out = [], {}
    for k, (w, h, inner) in arts.items():
        p = os.path.join(BUILD, f"deck-{k}.png")
        jobs.append({"html": write(f"deck-{k}.html", art_doc(w, h, inner)), "png": p, "width": w, "height": h, "scale": 1})
        out[k] = (p, (w, h))
    return jobs, out


def tpt_json():
    return {
        "price": PRICE, "license_price": PRICE,
        "tax_code": "Other Digital Goods - No Physical Media",
        "grades": ["9th-grade", "10th-grade", "11th-grade", "12th-grade"],
        "subjects": ["Music"],
        "tags": ["Autumn"],
        "formats": ["PDF"],
        "pages": len(PLAN) + DECK_SLIDES,
        "answer_key": "Included",
        "zip_contents": [os.path.basename(PDF), os.path.basename(DECK), os.path.basename(DECK_PDF)],
        "files": {"product": os.path.basename(ZIP), "preview": os.path.basename(PREVIEW_PDF), "thumb1": "cover.png",
                  "thumb2": "preview-1.png", "thumb3": "preview-2.png", "thumb4": "preview-3.png"},
    }


def make_zip():
    if os.path.exists(ZIP):
        os.remove(ZIP)
    with zipfile.ZipFile(ZIP, "w", zipfile.ZIP_DEFLATED) as z:
        for p in (PDF, DECK, DECK_PDF):
            zi = zipfile.ZipInfo(f"{ZIP_DIR}/{os.path.basename(p)}", date_time=(2026, 10, 12, 12, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            with open(p, "rb") as fh:
                z.writestr(zi, fh.read())


def build():
    os.makedirs(BUILD, exist_ok=True)
    problems = verify_data()
    if problems:
        print("\n".join(problems))
        sys.exit("data check failed; not rendering")
    full = write("resource.html", html_doc([render_page(p) for p in PLAN], TITLE))
    prev = write("preview.html", html_doc([render_page(p, watermark=True) for p in PREVIEW_PICKS], "Preview"))
    art_jobs, art = deck_art()
    render([{"html": full, "pdf": PDF}, {"html": prev, "pdf": PREVIEW_PDF}] + art_jobs)
    build_deck(art)
    deck_pdf()
    make_zip()
    res = 45 if LOW_RES else 120
    for key, _, _ in SHOTS1:
        pdf_png(PDF, IX[key], key, res)
    for key, _ in SHOTS3:
        pdf_png(PDF, IX[key], key, 40 if LOW_RES else 80)
    for s in SHOTS2:
        pdf_png(DECK_PDF, s, f"slide{s}", 30 if LOW_RES else 70)
    sqs = [write("cover.html", cover_square()), write("p1.html", preview_square_1([(f"{k}.png", a, b) for k, a, b in SHOTS1])),
           write("p2.html", preview_square_2([f"slide{s}.png" for s in SHOTS2])),
           write("p3.html", preview_square_3([(f"{k}.png", lab) for k, lab in SHOTS3]))]
    render([{"html": h, "png": o, "width": 1000, "height": 1000, "scale": 1 if LOW_RES else 2} for h, o in zip(sqs, [COVER_PNG] + PREVIEW_PNGS)])
    with open(TPT_JSON, "w") as f:
        json.dump(tpt_json(), f, indent=2, ensure_ascii=False)
        f.write("\n")
    print(f"built {PDF} ({len(PLAN)} pages), {DECK} + PDF ({DECK_SLIDES} slides), {ZIP}, preview, cover + 3 previews, tpt.json"
          + (" [LOW RES]" if LOW_RES else ""))


# ---------------------------------------------------------------- checks
def lane_count(section):
    return sum(1 for lane in LANES if section[2].get(lane))


def verify_data():
    errs = []
    # songs: facts from the owner's Complete sheet; the Blank sheet's pre-filled labels
    for t, a, use, secs, pre in SONGS:
        if not secs or len(secs) > SHEET_ROWS:
            errs.append(f"{t}: {len(secs)} sections do not fit the {SHEET_ROWS}-row table")
        if not 0 <= pre <= len(secs):
            errs.append(f"{t}: pre-filled labels out of range")
        for nm, ms, shown in secs:
            if ms <= 0:
                errs.append(f"{t}: {nm} has {ms} measures")
            if shown and sum(int(x) for x in re.findall(r"\d+", shown.split("(")[1])) != ms if "(" in (shown or "") else False:
                errs.append(f"{t}: '{shown}' does not add up to {ms}")
        if use not in ("class", "choice"):
            errs.append(f"{t}: use must be class or choice")
        if t not in SONG_WHY:
            errs.append(f"{t}: no 'why it works' line")
        if use == "choice" and t not in SONG_HINT:
            errs.append(f"{t}: no Day 1 hint")
    if [s[2] for s in SONGS].count("class") != 2 or [s[2] for s in SONGS].count("choice") != 4:
        errs.append("songs: 2 class examples and 4 choices expected (owner's Day 1 post)")
    if (SONGS[0][0], SONGS[1][0]) != ("Crazy", "Blinding Lights"):
        errs.append("class example order must follow the Day 1 post: Crazy first, then Blinding Lights")
    # worked example: a 32-measure skeleton that obeys the owner's rules
    if EX_MEASURES != 32:
        errs.append(f"worked example is {EX_MEASURES} measures, expected 32")
    names = [s[0] for s in EXAMPLE]
    for need in ("Intro", "Verse 1", "Chorus 1"):
        if need not in names:
            errs.append(f"worked example: missing {need} (intro, verse, chorus required)")
    if len(EXAMPLE) < 4:
        errs.append("worked example: fewer than 4 sections")
    starts = {nm: (a, z) for nm, a, z in ex_starts()}
    for sec in EXAMPLE:
        nm, bars, lanes, tr, how = sec
        if bars < 4 or bars % 4:
            errs.append(f"worked example: {nm} is {bars} measures (each section 4+ and in 4s)")
        if nm in starts and nm != "Outro":
            if tr not in ("Fill", "Breakdown", "Cut", "Build"):
                errs.append(f"worked example: {nm} has no named transition")
        a, z = starts[nm]
        for m in re.findall(r"m\. (\d+)", how):
            if not a <= int(m) <= z:
                errs.append(f"worked example: {nm} transition text names m. {m}, outside m. {a}-{z}")
        for lane in lanes:
            if lane not in LANES:
                errs.append(f"worked example: unknown lane {lane}")
    by = {s[0]: s for s in EXAMPLE}
    if by["Chorus 1"][2] != by["Chorus 2"][2] or by["Chorus 1"][1] != by["Chorus 2"][1]:
        errs.append("worked example: the chorus must be the same every time it returns (owner: copy and paste)")
    if not lane_count(by["Intro"]) < lane_count(by["Verse 1"]) < lane_count(by["Chorus 1"]):
        errs.append("worked example: intro < verse < chorus in elements (owner's planning-table checks)")
    if not ({"Chord layer (pad)", "Melody"} <= set(by["Chorus 1"][2]) and "Chord layer (pad)" not in by["Verse 1"][2]):
        errs.append("worked example: chord layer only in the chorus; melody in the chorus")
    no_mel = lambda s: sum(1 for lane in LANES if s[2].get(lane) and lane != "Melody")
    if no_mel(by["Chorus 1"]) <= no_mel(by["Verse 1"]):
        errs.append("worked example fails the owner's mute test (chorus not bigger with the melody muted)")
    if lane_count(by["Outro"]) >= lane_count(by["Verse 1"]):
        errs.append("worked example: the outro should strip back")
    if [r for r in PLAN_ROWS if r != "Bridge (optional)"] != [n for n in names]:
        errs.append("worked example sections must match the owner's planning-table rows (bridge left out)")
    if len(EX_PROG) != len(EX_CHORDS) or EX_CHORDS != ["C", "G", "Am", "F"] or EX_PROG != ["I", "V", "vi", "IV"]:
        errs.append("worked example: I-V-vi-IV in C is C G Am F")
    # owner's lists, counted
    counts = dict(goals=(GOALS, 7), vocab=(VOCAB, 10), req=(REQUIREMENTS, 12), check=(CHECKLIST, 12), steps=(STEPS, 6), tips=(TIPS, 5),
                  jobs=(SECTION_JOBS, 6), struggles=(STRUGGLES, 4), demo=(LAYER_DEMO, 5), plancheck=(PLAN_CHECK, 4), run=(RUN_ANALYSIS, 4),
                  transitions=(TRANSITIONS, 4), callbacks=(CALLBACKS, 3), planrows=(PLAN_ROWS, 7), quiz=(VOCAB_QUIZ, 4))
    for nm, (lst, nexp) in counts.items():
        if len(lst) != nexp:
            errs.append(f"{nm}: {len(lst)} items, expected {nexp}")
    vw = {w.lower() for w, _ in VOCAB}
    for _, a in VOCAB_QUIZ:
        if a not in vw:
            errs.append(f"vocab quiz answer '{a}' is not a vocabulary word")
    if len(RUBRIC) != 5 or any(len(dd) != len(LEVELS) for _, dd in RUBRIC) or 5 * len(RUBRIC) != 25 or len(RUBRIC_SHORT) != 5:
        errs.append("rubric must be 5 categories x levels 2-5 = /25")
    if RUBRIC[2][1][3] != ("Multiple sections have distinct textures. Addition and removal of elements is deliberate. Listener can identify the section "
                           "by ear without looking at the screen."):
        errs.append("rubric: Section Differentiation level 5 must keep the owner's wording")
    # pacing: 4 days, blocks 0-53, and every deck slide in exactly one block of the day the slide says
    if len(DAYS) != DAYS_N or [d["n"] for d in DAYS] != list(range(1, DAYS_N + 1)) or len(DECK_DAY_LINES) != DAYS_N:
        errs.append(f"pacing must be {DAYS_N} days (and {DAYS_N} on the deck)")
    seen = {}
    for d in DAYS:
        blocks = [tuple(map(int, x.split("–"))) for x, _, _, _ in d["blocks"]]
        if blocks[0][0] != 0 or blocks[-1][1] != 53 or any(blocks[i][1] != blocks[i + 1][0] for i in range(len(blocks) - 1)):
            errs.append(f"Day {d['n']} minute blocks do not run 0-53 without gaps")
        for _, h, txt, sl in d["blocks"]:
            if not sl:
                if re.search(r"\bslides? \d", txt, re.I):
                    errs.append(f"Day {d['n']} '{h}' names a slide but has no slide slot")
                continue
            for s in range(sl[0], sl[1] + 1):
                if s in seen:
                    errs.append(f"slide {s} is slotted twice (Day {seen[s]} and Day {d['n']})")
                seen[s] = d["n"]
                if not 1 <= s <= DECK_SLIDES:
                    errs.append(f"Day {d['n']} names slide {s}, which does not exist")
                elif slide_day(s) != d["n"]:
                    errs.append(f"slide {s} says Day {slide_day(s)} but the Day {d['n']} plan uses it")
            named = {int(x) for x in re.findall(r"\b(\d+)\b", " ".join(re.findall(r"[Ss]lides? [\d–]+", txt)).replace("–", " "))}
            rngset = set()
            for a_, b_ in re.findall(r"[Ss]lides? (\d+)(?:–(\d+))?", txt):
                rngset |= set(range(int(a_), int(b_ or a_) + 1))
            if rngset and rngset != set(range(sl[0], sl[1] + 1)):
                errs.append(f"Day {d['n']} '{h}': text names slides {sorted(rngset)} but the slot is {sl}")
            del named
    for i in range(1, DECK_SLIDES + 1):
        if i not in seen:
            errs.append(f"deck slide {i} has no day slot in the day plans")
    if [s[0] for s in SLIDES] != sorted(s[0] for s in SLIDES):
        errs.append("deck slides must run in day order")
    if len(SLIDES) < 12 or any(len(s[4].strip()) < 80 for s in SLIDES):
        errs.append("deck needs 12+ slides, each with real speaker notes")
    if [k for _, k, *_ in SLIDES].count("listen") != 2:
        errs.append("deck: two guided-listening slides expected")
    if len(PLAN) != len(set(PLAN)):
        errs.append("plan has a repeated page")
    if len(LISTING_TITLE) > 80:
        errs.append(f"listing title is {len(LISTING_TITLE)} chars")
    return errs


def pdf_fonts(path):
    out = subprocess.run(["pdffonts", path], capture_output=True, text=True, check=True).stdout.splitlines()[2:]
    return [(ln.split()[0], ln.split()[-5]) for ln in out]


def pdf_text(path):
    return subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True, check=True).stdout


def deck_texts():
    from pptx import Presentation
    prs = Presentation(DECK)
    slides, notes = [], []
    for s in prs.slides:
        slides.append(" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame))
        notes.append(s.notes_slide.notes_text_frame.text if s.has_notes_slide else "")
    media = [p for p in zipfile.ZipFile(DECK).namelist() if p.startswith("ppt/media/")]
    return prs, slides, notes, media


def verify_files():
    errs = []
    from pypdf import PdfReader
    from PIL import Image
    for p in [PDF, DECK, DECK_PDF, ZIP, PREVIEW_PDF, COVER_PNG, *PREVIEW_PNGS, UPLOAD, TPT_JSON]:
        if not os.path.exists(p):
            errs.append(f"missing {os.path.basename(p)}")
    if errs:
        return errs
    for path, n_exp in ((PDF, len(PLAN)), (PREVIEW_PDF, len(PREVIEW_PICKS))):
        rd = PdfReader(path)
        if len(rd.pages) != n_exp:
            errs.append(f"{os.path.basename(path)}: {len(rd.pages)} pages, expected {n_exp}")
        for i, pg_ in enumerate(rd.pages):
            w, h = float(pg_.mediabox.width), float(pg_.mediabox.height)
            if abs(w - 612) > 1 or abs(h - 792) > 1:
                errs.append(f"{os.path.basename(path)} p{i + 1}: {w:.0f}x{h:.0f}pt, not US Letter")
    if len(PdfReader(PREVIEW_PDF).pages) != 4:
        errs.append("preview PDF must be 4 pages")
    drd = PdfReader(DECK_PDF)
    if len(drd.pages) != DECK_SLIDES:
        errs.append(f"deck PDF: {len(drd.pages)} pages, expected {DECK_SLIDES}")
    for i, pg_ in enumerate(drd.pages):
        if abs(float(pg_.mediabox.width) / float(pg_.mediabox.height) - 16 / 9) > 0.01:
            errs.append(f"deck PDF p{i + 1}: not 16:9")
    for path in (PDF, PREVIEW_PDF, DECK_PDF):
        fonts = pdf_fonts(path)
        for name, emb in fonts:
            if emb != "yes":
                errs.append(f"{os.path.basename(path)}: font {name} not embedded")
            if not re.search(r"oswald|sourcesans", name.lower().replace(" ", "")):
                errs.append(f"{os.path.basename(path)}: unexpected font {name} (fallback?)")
        names = " ".join(n for n, _ in fonts).lower()
        for fam in ("oswald", "sourcesans3"):
            if fam not in names:
                errs.append(f"{os.path.basename(path)}: {fam} missing")
    # deck: slide count, speaker notes on every slide, day badge on every slide, only our 4 drawn PNGs as media
    prs, stext, notes, media = deck_texts()
    if len(prs.slides) != DECK_SLIDES:
        errs.append(f"deck has {len(prs.slides)} slides, expected {DECK_SLIDES}")
    for i, nt in enumerate(notes):
        if len(nt.strip()) < 80:
            errs.append(f"deck slide {i + 1}: speaker notes missing or too short")
        if not re.search(rf"\bDay {slide_day(i + 1)}\b", stext[i], re.I):
            errs.append(f"deck slide {i + 1}: no 'Day {slide_day(i + 1)}' slot shown on the slide")
    if any(not m.lower().endswith(".png") for m in media) or len(media) != 4:
        errs.append(f"deck media should be exactly the 4 drawn PNGs (hero, arrangement, plan, sheet), found {media}")
    if abs(prs.slide_width / prs.slide_height - 16 / 9) > 0.01:
        errs.append("deck is not 16:9")
    alldeck = " ".join(stext + notes)
    for t, a, *_ in SONGS:
        if t.lower() not in alldeck.lower() or a.lower() not in alldeck.lower():
            errs.append(f"deck must name {t} by {a}")
    for i, (_, k, *_) in enumerate(SLIDES):
        if k in ("listen", "sheet") and "YouTube or Spotify" not in stext[i]:
            errs.append(f"deck slide {i + 1}: must say YouTube or Spotify")
    if "YouTube or Spotify" not in notes[0]:
        errs.append("deck slide 1 notes must say YouTube or Spotify")
    names = zipfile.ZipFile(ZIP).namelist()
    for p in (PDF, DECK, DECK_PDF):
        n = f"{ZIP_DIR}/{os.path.basename(p)}"
        if n not in names:
            errs.append(f"zip is missing {os.path.basename(p)}")
        elif zipfile.ZipFile(ZIP).read(n) != open(p, "rb").read():
            errs.append(f"zip copy of {os.path.basename(p)} is stale; rebuild")
    if len(names) != 3:
        errs.append(f"zip should hold exactly 3 files, has {names}")
    rd = PdfReader(PDF)
    pages = [subprocess.run(["pdftotext", "-layout", "-f", str(i + 1), "-l", str(i + 1), PDF, "-"], capture_output=True, text=True,
                            check=True).stdout for i in range(len(rd.pages))]
    flat = [re.sub(r"\s+", " ", t) for t in pages]
    lines0 = [ln.strip() for ln in pages[0].splitlines() if ln.strip()]
    for need in ("SONG STRUCTURE IN LOGIC PRO", "SONG FORM"):
        if need not in lines0:
            errs.append(f"PDF cover: '{need}' is not on a line of its own (orphaned word?)")
    for i, t in enumerate(flat):
        if i > 0 and "Brian Lotze · Hudson Beat" not in t:
            errs.append(f"page {i + 1}: footer missing")
        if i > 0 and not re.search(rf"Page {i + 1}\s*$", t.strip()):
            errs.append(f"page {i + 1}: footer page number wrong (overflow?)")
    for m in re.finditer(r"\b(?:pages?|pp?\.) (\d+)(?:–(\d+))?", " ".join(flat[1:])):
        for g in m.groups():
            if g and not 1 <= int(g) <= len(PLAN):
                errs.append(f"cross-reference to page {g} does not exist")
    for m in re.finditer(r"\bslides? (\d+)(?:–(\d+))?", " ".join(flat[1:])):
        for g in m.groups():
            if g and not 1 <= int(g) <= DECK_SLIDES:
                errs.append(f"PDF names deck slide {g}, which does not exist")
    rub = flat[IX["rubric"] - 1]
    if f"/ {5 * len(RUBRIC)}" not in rub:
        errs.append("rubric page: /25 total missing")
    for c, ds in RUBRIC:
        for d in ds:
            if not all(w in rub for w in re.findall(r"[A-Za-z]{6,}", d)):
                errs.append(f"rubric text missing: {d[:50]}")
    for w, _ in VOCAB:
        if w not in flat[IX["vocab"] - 1]:
            errs.append(f"vocab page missing {w}")
    key = flat[IX["key"] - 1]
    for t, a, _, secs, _ in SONGS:
        if t.lower() not in key.lower():
            errs.append(f"answer key missing {t}")
    for d in DAYS:
        if f"DAY {d['n']}" not in flat[IX["pacing"] - 1]:
            errs.append(f"pacing page missing Day {d['n']}")
    for k in ("sheet1", "sheet2", "sheet3"):
        idx = SHEET_INTROS[k][0]
        for i in idx:
            if SONGS[i][0].lower() not in flat[IX[k] - 1].lower():
                errs.append(f"{k}: missing {SONGS[i][0]}")
    ptxt = [pg.extract_text() for pg in PdfReader(PREVIEW_PDF).pages]
    if sum("PREVIEW" in t for t in ptxt) != len(PREVIEW_PICKS):
        errs.append("preview PDF: watermark missing on some pages")
    # banned strings: every output, the deck XML, PDF metadata, and every text file in this folder (gen.py included)
    blobs = {"resource PDF": pdf_text(PDF), "preview PDF": pdf_text(PREVIEW_PDF), "deck PDF": pdf_text(DECK_PDF), "deck pptx": alldeck,
             "zip names": " ".join(names)}
    with zipfile.ZipFile(DECK) as z:
        blobs["deck XML"] = " ".join(z.read(n).decode("utf8", "ignore") for n in z.namelist() if n.endswith(".xml") or n.endswith(".rels"))
    for path in (PDF, PREVIEW_PDF, DECK_PDF):
        meta = PdfReader(path).metadata or {}
        blobs[f"{os.path.basename(path)} metadata"] = " ".join(str(v) for v in meta.values())
    for fn in os.listdir(HERE):
        if fn == "QA.md":   # QA's own review record (not shipped) names the strings it searched for
            continue
        if fn.endswith((".md", ".json", ".py", ".txt")) or fn in ("release", "uploaded"):
            blobs[fn] = open(os.path.join(HERE, fn), encoding="utf8", errors="ignore").read()
    for label, txt in blobs.items():
        low = txt.lower()
        for b in BANNED:
            if b.lower() in low:
                errs.append(f"{label}: contains banned string '{b}'")
    for p in [COVER_PNG, *PREVIEW_PNGS]:
        if Image.open(p).size != (2000, 2000):
            errs.append(f"{os.path.basename(p)}: not 2000x2000 (low-res build?)")
    tj = json.load(open(TPT_JSON))
    if tj != tpt_json():
        errs.append("tpt.json is stale; rebuild")
    if tj["price"] != PRICE or tj["license_price"] != PRICE or tj.get("tax_code") != "Other Digital Goods - No Physical Media":
        errs.append("tpt.json price/license_price/tax_code")
    if tj["grades"] != ["9th-grade", "10th-grade", "11th-grade", "12th-grade"] or tj["subjects"] != ["Music"]:
        errs.append("tpt.json grades 9-12 / subject Music")
    proven_tags = {"Halloween", "Autumn", "Winter", "Thanksgiving"}   # accepted on live items 001-007
    if not (1 <= len(tj["tags"]) <= 6 and len(tj["formats"]) <= 3) or not set(tj["tags"]) <= proven_tags or tj["formats"] != ["PDF"]:
        errs.append("tpt.json tags/formats: only proven values (tags from 001-007, format PDF)")
    for k, f in tj["files"].items():
        if not os.path.exists(os.path.join(HERE, f)):
            errs.append(f"tpt.json file {k} missing: {f}")
    if tj["pages"] != len(rd.pages) + len(prs.slides):
        errs.append("tpt.json pages != PDF pages + slides")
    up = open(UPLOAD).read()
    m = re.search(r"^## 1\. Title\s*\n+```\n(.+?)\n```\n\((\d+) characters", up, re.M | re.S)
    if not m:
        errs.append("UPLOAD.md: title block or its character count not found")
    else:
        n = len(m.group(1).strip())
        if n > 80:
            errs.append(f"UPLOAD.md: title is {n} chars (max 80)")
        if int(m.group(2)) != n:
            errs.append(f"UPLOAD.md: says {m.group(2)} characters, title has {n}")
        if m.group(1).strip() != LISTING_TITLE:
            errs.append("UPLOAD.md title differs from LISTING_TITLE in gen.py")
    m2 = re.search(r"^## 2\. Description\s*\n+```\n(.+?)\n```", up, re.M | re.S)
    if not m2 or not m2.group(1).rstrip().endswith(DISCLOSURE):
        errs.append("UPLOAD.md: description must end with the disclosure line")
    else:
        desc = m2.group(1)
        if not desc.startswith("Song structure for high school music technology:"):
            errs.append("UPLOAD.md description: first line must lead with the search phrase (METADATA-GUIDE)")
        for need in (f"{len(PLAN)}-page", f"{DECK_SLIDES}-slide", "Logic Pro", "/25", "YouTube or Spotify", "GarageBand", f"{DAYS_N}-day"):
            if need not in desc:
                errs.append(f"UPLOAD.md description: missing '{need}'")
        if re.search(r"  +\S|\s[,.]\s|\bteh\b|\s{2,}$", desc, re.M):
            errs.append("UPLOAD.md description: double space or stray punctuation")
    for need in (f"${PRICE}", f"{len(PLAN)} pages", os.path.basename(ZIP), os.path.basename(DECK), os.path.basename(DECK_PDF)):
        if need not in up:
            errs.append(f"UPLOAD.md: missing '{need}'")
    for f in ("release", "uploaded"):
        if os.path.exists(os.path.join(HERE, f)):
            errs.append(f"{f} exists: only the Manager / uploader writes it")
    return errs


def check():
    errs = verify_data() + verify_maps() + verify_files()
    if errs:
        print("CHECK FAILED")
        print("\n".join(" - " + e for e in errs))
        sys.exit(1)
    print(f"CHECK OK: 6 song analyses fit the sheet, owner's pre-filled labels in range; worked skeleton {EX_MEASURES} measures, sections in 4s, "
          f"intro < verse < chorus, chorus repeats unchanged, passes the mute test, transitions inside their sections; owner's lists and rubric "
          f"5 x (2-5) = /25 printed; {DAYS_N} days run 0-53; every one of {DECK_SLIDES} slides has exactly one day slot that matches its Day badge, and "
          f"every slide named in a plan exists; {len(PLAN)} Letter pages with footers and valid page/slide references; cover title lines whole; "
          f"preview 4 watermarked pages; deck 16:9 with speaker notes on every slide, only 4 drawn PNGs, PDF copy {DECK_SLIDES} pages; only "
          f"Oswald/Source Sans 3, all embedded; no banned strings in any output or file here; zip current; images 2000x2000; tpt.json uses only proven "
          f"values; UPLOAD.md title {len(LISTING_TITLE)} chars, SEO opening and disclosure.")


if __name__ == "__main__":
    if "--fallback" in sys.argv:   # self-check: the deck rendered without Oswald / Source Sans 3
        deck_pdf(DECK, os.path.join(BUILD, "deck-fallback.pdf"), fonts=False)
        print("wrote build/deck-fallback.pdf")
    elif "--check" in sys.argv:
        check()
    else:
        build()
