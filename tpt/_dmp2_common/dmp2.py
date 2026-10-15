"""Shared builder for the DMP2 project kits (Hudson Beat, TpT). See README.md in this folder.

A kit's gen.py holds only the content (the owner's words, the pacing, the slides, any custom pages and drawings) as
module-level names, then calls:

    import dmp2
    dmp2.main(sys.modules[__name__])

main() reads sys.argv: no flag = build everything, --lowres = build with small images (self-check only), --check = verify.
Everything renders with Playwright + Chromium (../tools/render.js) and LibreOffice (deck PDF), fonts from ../fonts. $0 AI.
"""
import json
import math
import os
import re
import shutil
import subprocess
import sys
import zipfile
from html import escape

COMMON = os.path.dirname(os.path.abspath(__file__))
TPT = os.path.dirname(COMMON)
FONTS = os.path.join(TPT, "fonts")
RENDER_JS = os.path.join(TPT, "tools", "render.js")

STORE = "Brian Lotze"      # copyright holder and TpT seller
BRAND = "Hudson Beat"      # store line (owner, 4 Oct 2026)
YEAR = 2026
DISCLOSURE = "Designed with the help of digital and AI tools, and checked by hand."
COPY = f"© {YEAR} {STORE} · {BRAND} · For single-classroom use"
TAX = "Other Digital Goods - No Physical Media"
PROVEN_TAGS = {"Halloween", "Autumn", "Winter", "Thanksgiving"}          # tag texts TpT has accepted (items 001-013)
PROVEN_GRADES = {f"{g}-grade" for g in ("9th", "10th", "11th", "12th")}   # slugs used on 008-013
# Strings that must never appear in any output or kit file (split so this file never contains them).
BANNED = ["Blind" + " Brook", "BB" + "HS", "blind" + "brook", "Google" + " Classroom", "Thunder" + "struck", "London" + "Bounce",
          "Falling" + " For U", "Brello", "nsd" + ".org"]

K = None   # the kit module, set by main()


# ================================================================ small helpers
def P(name):
    return os.path.join(K.HERE, name)


def build_dir():
    return os.path.join(K.HERE, "build")


def lowres():
    return "--lowres" in sys.argv


def rng(a, b):
    return str(a) if a == b else f"{a}–{b}"


def fill(s):
    """Page placeholders are written @KEY (a page key in PLAN, upper case) so they never collide with ordinary words."""
    ix = {k.upper(): i + 1 for i, k in enumerate(K.PLAN)}
    out = re.sub(r"@([A-Z0-9]+)", lambda m: str(ix[m.group(1)]), s)
    if "@" in out:
        raise ValueError(f"unfilled placeholder in: {s}")
    return out


def IX(key):
    return K.PLAN.index(key) + 1


_ADV = {}


def text_w(s, size, font="oswald"):
    """Advance width of s at size from the font's own metrics: Oswald 700 ("oswald") or Source Sans 3 400 ("ss")."""
    if font not in _ADV:
        from fontTools.ttLib import TTFont
        f = TTFont(os.path.join(FONTS, "oswald-latin-700-normal.woff2" if font == "oswald" else "source-sans-3-latin-400-normal.woff2"))
        cmap, hm = f.getBestCmap(), f["hmtx"].metrics
        d = {chr(c): hm[g][0] for c, g in cmap.items() if c < 0x2200}
        d["upm"] = f["head"].unitsPerEm
        _ADV[font] = d
    a = _ADV[font]
    return sum(a.get(ch, a["upm"] * 0.6) for ch in s) * size / a["upm"]


# ================================================================ html scaffolding (Hudson Beat, as 008-013)
def font_css(base=None):
    f = os.path.relpath(FONTS, base or build_dir())
    faces = [("Oswald", 300, "normal", "oswald-latin-300-normal"), ("Oswald", 500, "normal", "oswald-latin-500-normal"),
             ("Oswald", 700, "normal", "oswald-latin-700-normal"),
             ("Source Sans 3", 400, "normal", "source-sans-3-latin-400-normal"),
             ("Source Sans 3", 400, "italic", "source-sans-3-latin-400-italic"),
             ("Source Sans 3", 600, "normal", "source-sans-3-latin-600-normal"),
             ("Source Sans 3", 700, "normal", "source-sans-3-latin-700-normal")]
    return "".join(f"@font-face{{font-family:'{n}';font-weight:{w};font-style:{s};src:url('{f}/{file}.woff2') format('woff2');}}"
                   for n, w, s, file in faces)


CSS = """
@page { size: 8.5in 11in; margin: 0; }
* { box-sizing: border-box; }
html, body { margin: 0; padding: 0; background: #fff; color: #000; }
body { font-family: 'Source Sans 3', sans-serif; -webkit-print-color-adjust: exact; print-color-adjust: exact; }
.page { width: 8.5in; height: 11in; position: relative; overflow: hidden; page-break-after: always; break-after: page;
        padding: 0.42in 0.5in 0.38in; display: flex; flex-direction: column; }
.page:last-child { page-break-after: auto; break-after: auto; }
.l { font-family: 'Oswald'; font-weight: 300; }
.top { display: flex; justify-content: space-between; align-items: flex-end; border-bottom: 2.2pt solid #000; padding-bottom: 6pt; }
.top .t { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; letter-spacing: .01em; line-height: 1; font-size: 25pt; }
.top .t small { display: block; font-family: 'Oswald'; font-weight: 300; text-transform: none; font-size: 12pt; letter-spacing: 0; margin-top: 4pt; }
.tag { text-align: right; }
.tag .lv { display: inline-block; background: #000; color: #fff; font-family: 'Oswald'; font-weight: 700; font-size: 14pt; padding: 3pt 9pt 4pt; border-radius: 4pt; letter-spacing: .04em; }
.tag .lv.ak { background: ACCENT; }
.tag .ln { font-family: 'Oswald'; font-weight: 300; font-size: 11pt; margin-top: 3pt; }
.name { display: flex; gap: 18pt; font-size: 12pt; margin: 9pt 0 4pt; }
.name span { flex: 1; border-bottom: 1pt solid #000; height: 15pt; }
.name span.d { flex: 0 0 1.6in; }
.main { flex: 1; display: flex; flex-direction: column; min-height: 0; padding-top: 6pt; overflow: hidden; }
.foot { display: flex; justify-content: space-between; font-size: 8.5pt; color: #333; border-top: .8pt solid #000; padding-top: 3pt; }
.wm { position: absolute; left: 50%; top: 50%; transform: translate(-50%, -50%) rotate(-38deg); font-family: 'Oswald'; font-weight: 700;
      font-size: 150pt; color: rgba(0,0,0,.13); letter-spacing: .06em; white-space: nowrap; pointer-events: none; z-index: 9; }
h2.s { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 13pt; margin: 9pt 0 3pt; letter-spacing: .02em; }
h2.s .acc, .acc { color: ACCENT; }
p, li { font-size: 10.5pt; line-height: 1.32; margin: 0 0 4pt; }
ul, ol { margin: 0 0 4pt; padding-left: 15pt; }
.tn p, .tn li { font-size: 11.6pt; line-height: 1.45; margin-bottom: 6pt; }
.big p, .big li { font-size: 12.2pt; line-height: 1.46; margin-bottom: 7pt; }
.lines { flex: 1; min-height: 44pt; background: repeating-linear-gradient(to bottom, transparent 0, transparent 23pt, #000 23pt, #000 24pt); }
table.t { border-collapse: collapse; width: 100%; font-size: 9.8pt; margin: 2pt 0 4pt; }
table.t td, table.t th { border: 1pt solid #000; padding: 3pt 5pt; text-align: left; vertical-align: top; line-height: 1.25; }
table.t th { font-family: 'Oswald'; font-weight: 500; text-transform: uppercase; font-size: 9.5pt; background: #000; color: #fff; letter-spacing: .04em; vertical-align: middle; }
.cols { display: flex; gap: 20pt; }
.cols > div { flex: 1; min-width: 0; }
.small { font-size: 9pt !important; color: #333; }
.box { border: 1.4pt solid #000; border-radius: 8pt; padding: 7pt 10pt; }
.box.acc-b { border-color: ACCENT; border-width: 2pt; }
.box.inv { background: #000; color: #fff; border-color: #000; }
.k { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 10pt; letter-spacing: .05em; color: ACCENT; display: block; margin-bottom: 2pt; }
.box.inv .k { color: ACC_LIGHT; }
.ln2 { border-bottom: 1pt solid #000; height: 24pt; }
.fl { display: flex; align-items: flex-end; gap: 6pt; font-size: 11pt; margin: 0 0 9pt; }
.fl b { font-weight: 600; white-space: nowrap; }
.fl span { flex: 1; border-bottom: 1pt solid #000; height: 15pt; }
.num { list-style: none; padding: 0; margin: 2pt 0 0; }
.num li { display: flex; gap: 10pt; margin: 0 0 8pt; }
.num .n { flex: 0 0 24pt; height: 24pt; border-radius: 12pt; background: #000; color: #fff; font-family: 'Oswald'; font-weight: 700; font-size: 13pt;
          display: flex; align-items: center; justify-content: center; }
.chk { display: flex; gap: 8pt; align-items: flex-start; font-size: 11pt; margin: 0 0 8pt; line-height: 1.3; }
table.rub { border-collapse: collapse; width: 100%; }
table.rub td, table.rub th { border: 1.1pt solid #000; padding: 6pt 5pt; vertical-align: top; font-size: 10.4pt; line-height: 1.32; }
table.rub th { background: #000; color: #fff; font-family: 'Oswald'; font-weight: 500; font-size: 12pt; text-align: center; vertical-align: middle; }
table.rub td.c { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 10.5pt; width: 15%; vertical-align: middle; }
.cover { align-items: center; justify-content: space-between; padding: 0.55in 0.6in 0.5in;
         background: linear-gradient(180deg, #fff 0%, #fff 55%, #ededed 100%); }
.store { font-family: 'Oswald'; font-weight: 500; font-size: 13pt; letter-spacing: .32em; text-transform: uppercase; }
.ctitle { text-align: center; }
.ctitle .a, .ctitle .b, .ctitle .c { white-space: nowrap; }
.ctitle .a { font-family: 'Oswald'; font-weight: 700; line-height: 1; letter-spacing: .02em; text-transform: uppercase; }
.ctitle .b { font-family: 'Oswald'; font-weight: 700; line-height: .98; letter-spacing: .01em; text-transform: uppercase; }
.ctitle .c { font-family: 'Oswald'; font-weight: 300; margin-top: 8pt; }
.cband { width: 100%; background: #000; color: #fff; text-align: center; font-family: 'Oswald'; font-weight: 500; font-size: 14pt;
         letter-spacing: .08em; text-transform: uppercase; padding: 8pt 0 9pt; border-radius: 4pt; white-space: nowrap; }
.cband em { font-style: normal; color: ACCENT; }
table.map { border-collapse: collapse; width: 100%; }
table.map td, table.map th { border: 1.2pt solid #000; padding: 4pt 6pt; font-size: 10.5pt; height: 27pt; }
table.map th { background: #000; color: #fff; font-family: 'Oswald'; font-weight: 500; text-transform: uppercase; font-size: 9.6pt; letter-spacing: .04em; height: auto; }
.stp { font-family: 'Oswald'; font-weight: 500; letter-spacing: .12em; }
"""


def css():
    return CSS.replace("ACC_LIGHT", K.ACC_LIGHT).replace("ACCENT", K.ACCENT)


def html_doc(pages, title="doc"):
    return (f'<!doctype html><html><head><meta charset="utf-8"><title>{escape(title)}</title>'
            f'<style>{font_css()}{css()}</style></head><body>{"".join(pages)}</body></html>')


def foot(right):
    return f'<div class="foot"><span>{COPY}</span><span>{right}</span></div>'


def page(sub, tag, tagsub, inner, n, watermark=False, name=False, ak=False, key=""):
    top = (f'<div class="top"><div class="t">{escape(K.TITLE)}<small>{sub}</small></div>'
           f'<div class="tag"><div class="lv{" ak" if ak else ""}">{tag}</div><div class="ln">{tagsub}</div></div></div>')
    nm = '<div class="name">Name<span></span>Date<span class="d"></span></div>' if name else ""
    wm = '<div class="wm">PREVIEW</div>' if watermark else ""
    return f'<section class="page" data-key="{key}">{wm}{top}{nm}<div class="main">{inner}</div>{foot(f"Page {n}")}</section>'


def H2(s, first=False):
    st = ' style="margin-top:2pt"' if first else ""
    return f'<h2 class="s"{st}>{s}<span class="acc"> ·</span></h2>'


def checkbox(size=12):
    return (f'<svg width="{size}pt" height="{size}pt" viewBox="0 0 {size} {size}" style="flex:0 0 auto;margin-top:1pt">'
            f'<rect x=".7" y=".7" width="{size - 1.4}" height="{size - 1.4}" rx="2" fill="#fff" stroke="#000" stroke-width="1.2"/></svg>')


def chk(text, size=12, style=""):
    return f'<div class="chk" style="{style}">{checkbox(size)}<span>{text}</span></div>'


def lines(n, h=24):
    return "".join(f'<div class="ln2" style="height:{h}pt"></div>' for _ in range(n))


# ================================================================ SVG drawing helpers (labels logged for the fit check)
LABELS = []     # (diagram, text, size, room): text_w(text, size) must be <= room
SPANS = []      # (diagram, row, x0, x1, lo, hi): labels on one row must not overlap and must stay inside lo..hi


def lab(diagram, text, size, room):
    LABELS.append((diagram, text, size, room))


def span(diagram, row, text, size, cx, anchor="middle", lo=0.0, hi=1e9):
    w = text_w(text, size)
    x0 = cx - w / 2 if anchor == "middle" else (cx if anchor == "start" else cx - w)
    SPANS.append((diagram, row, x0, x0 + w, lo, hi))


def t_(x, y, s, size, weight=700, family="Oswald", anchor="middle", fill_="#000", extra=""):
    fam = f"'{family}'" if " " in family else family
    return (f'<text x="{x:.1f}" y="{y:.1f}" font-family="{fam}" font-weight="{weight}" font-size="{size}" text-anchor="{anchor}" '
            f'fill="{fill_}" {extra}>{escape(s)}</text>')


def svgwrap(vw, vh, body, width="100%"):
    return f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {vw:.0f} {vh:.0f}" width="{width}" style="display:block">{body}</svg>'


def body_svg(bwh, width="100%"):
    b, W, H = bwh
    return svgwrap(W, H, b, width)


def stack(parts, gap=26):
    W = max(w for _, w, _ in parts)
    out, y = [], 0
    for b, w, h in parts:
        out.append(f'<g transform="translate({(W - w) / 2:.1f},{y:.1f})">{b}</g>')
        y += h + gap
    return "".join(out), W, y - gap


def side(parts, gap=40):
    H = max(h for _, _, h in parts)
    out, x = [], 0
    for b, w, h in parts:
        out.append(f'<g transform="translate({x:.1f},{(H - h) / 2:.1f})">{b}</g>')
        x += w + gap
    return "".join(out), x - gap, H


STEPS16 = 16


def grid_body(name, measures, captions, lanes, cw=20, ch=22, gut=150, ls=15, accent_from=1, blank=False, vel=None):
    """A drum grid, one 16-step block per measure (dict lane -> 'x...' string). Grey columns are the beats. Hits in measures
    >= accent_from use the accent. blank=True draws an empty grid for students. vel: optional {measure index: [16 velocities or
    None]} drawn as bars under the grid (velocity lane). Returns (svg body, width, height)."""
    out, head = [], 22
    W = gut + STEPS16 * len(measures) * cw + 4
    for mi in range(len(measures)):
        for b in range(4):
            x = gut + (mi * STEPS16 + b * 4) * cw
            out.append(t_(x + cw / 2, 16, str(b + 1), 14, 700))
            lab(f"{name} beats", str(b + 1), 14, cw * 2)
    for r, ln in enumerate(lanes):
        y = head + r * ch
        out.append(t_(gut - 8, y + ch * 0.7, ln, ls, 500, anchor="end"))
        lab(f"{name} lanes", ln, ls, gut - 14)
        for mi, m in enumerate(measures):
            s = m.get(ln, "." * STEPS16)
            for i in range(STEPS16):
                x = gut + (mi * STEPS16 + i) * cw
                bg = "#e4e4e4" if i % 4 == 0 else "#fff"
                out.append(f'<rect x="{x:.1f}" y="{y:.1f}" width="{cw:.1f}" height="{ch}" fill="{bg}" stroke="#000" '
                           f'stroke-width="{1.3 if i % 4 == 0 else .55}"/>')
                if s[i] == "x" and not blank:
                    out.append(f'<rect x="{x + 3:.1f}" y="{y + 3:.1f}" width="{cw - 6:.1f}" height="{ch - 6}" rx="2.5" '
                               f'fill="{K.ACCENT if mi >= accent_from else "#000"}"/>')
    gh = len(lanes) * ch
    yv = head + gh
    if vel:
        vh = ch * 1.6
        out.append(t_(gut - 8, yv + vh * 0.62, "Velocity", ls, 500, anchor="end"))
        lab(f"{name} lanes", "Velocity", ls, gut - 14)
        for mi in range(len(measures)):
            out.append(f'<rect x="{gut + mi * STEPS16 * cw:.1f}" y="{yv:.1f}" width="{STEPS16 * cw:.1f}" height="{vh:.1f}" fill="#fff" '
                       f'stroke="#000" stroke-width="1"/>')
            for i, v in enumerate(vel.get(mi, [None] * STEPS16)):
                if v:
                    bh = (vh - 6) * v / 127
                    out.append(f'<rect x="{gut + (mi * STEPS16 + i) * cw + 4:.1f}" y="{yv + vh - 3 - bh:.1f}" width="{cw - 8:.1f}" '
                               f'height="{bh:.1f}" fill="{K.ACCENT if mi >= accent_from else "#555"}"/>')
        gh += vh
    for mi in range(len(measures) + 1):
        x = gut + mi * STEPS16 * cw
        out.append(f'<line x1="{x:.1f}" y1="{head - 4}" x2="{x:.1f}" y2="{head + gh + 4}" stroke="#000" stroke-width="3"/>')
    out.append(f'<rect x="{gut}" y="{head}" width="{STEPS16 * len(measures) * cw:.1f}" height="{gh}" fill="none" stroke="#000" stroke-width="1.8"/>')
    for mi, cap in enumerate(captions):
        cx = gut + (mi + .5) * STEPS16 * cw
        out.append(t_(cx, head + gh + 21, cap, 16, 700, fill_=K.ACCENT if mi >= accent_from else "#000"))
        lab(f"{name} captions", cap, 16, STEPS16 * cw - 12)
    return "".join(out), W, head + gh + (28 if captions else 6)


def form_body(name, sections, W=1600, h=110, fills=None, times=None, energy=None, label_size=26):
    """A song-form strip: one block per section (width = bars), the bar it starts at under it (and the start time if times is a
    function bar -> 'm:ss'). energy: optional {section name: 1-4} drawn as bars inside the block. Returns (body, W, H)."""
    total = sum(n for _, n in sections)
    mw = (W - 4) / total
    out, x, bar = [], 2.0, 1
    fills = fills or {}
    for k, (sec, n) in enumerate(sections):
        w = n * mw
        out.append(f'<rect x="{x + 2:.1f}" y="2" width="{w - 4:.1f}" height="{h}" rx="8" fill="{fills.get(sec.split()[0], "#000")}"/>')
        head = sec.upper()
        sz = label_size
        while sz > 14 and text_w(head, sz) > w - 14:
            sz -= 1
        out.append(t_(x + w / 2, 38, head, sz, 700, fill_="#fff"))
        lab(name, head, sz, w - 14)
        if energy:
            e = energy[sec.split()[0]]
            for j in range(4):
                bh = 10 + j * 9
                bx = x + w / 2 - 34 + j * 18
                out.append(f'<rect x="{bx:.1f}" y="{h - 8 - bh:.1f}" width="12" height="{bh}" rx="2" fill="#fff" '
                           f'opacity="{1 if j < e else .25}"/>')
        out.append(t_(x + 4, h + 24, f"bar {bar}", 15, 500, anchor="start"))
        lab(f"{name} bars", f"bar {bar}", 15, w - 8)
        if times:
            out.append(t_(x + 4, h + 44, times(bar), 15, 300, anchor="start"))
            lab(f"{name} times", times(bar), 15, w - 8)
        x += w
        bar += n
    return "".join(out), W, h + (52 if times else 32)


def verify_labels():
    """Every drawn label fits its cell, block or column; labels on one row do not overlap; nothing leaves the drawing."""
    LABELS.clear()
    SPANS.clear()
    for fn in K.DRAWINGS.values():
        fn()
    if hasattr(K, "label_extras"):
        K.label_extras()
    errs = [f"{d}: label '{t}' at {s} is {text_w(t, s):.0f} wide, room {r:.0f}" for d, t, s, r in LABELS if text_w(t, s) > r]
    if not LABELS:
        errs.append("label check found no labels")
    rows = {}
    for d, row, x0, x1, lo, hi in SPANS:
        rows.setdefault((d, row), []).append((x0, x1, lo, hi))
    for (d, row), iv in rows.items():
        iv.sort()
        for a, b in zip(iv, iv[1:]):
            if b[0] < a[1] + 10:
                errs.append(f"{d} {row}: labels overlap ({a[0]:.0f}-{a[1]:.0f} and {b[0]:.0f}-{b[1]:.0f})")
        for x0, x1, lo, hi in iv:
            if x0 < lo or x1 > hi:
                errs.append(f"{d} {row}: a label runs off the drawing")
    return errs


# ================================================================ standard pages
def cover_page(n=1):
    s = K.PDF_COVER
    band = cover_band().replace(" · ", " <em>·</em> ")
    return (f'<section class="page cover" data-key="cover"><div class="store">{BRAND}</div>'
            f'<div style="width:{s.get("art", 7.3)}in">{body_svg(K.hero_body())}</div>'
            f'<div class="ctitle"><div class="a" style="font-size:{s["a"]}pt">{K.COVER_A}</div>'
            f'<div class="b" style="font-size:{s["b"]}pt">{escape(K.TITLE)}</div>'
            f'<div class="c" style="font-size:{s["c"]}pt">{K.COVER_C}</div></div>'
            f'<div class="cband">{band}</div></section>')


def cover_band():
    return f"Grades {K.GRADES_TEXT} · {K.CLASSES} class periods · {len(K.PLAN)} pages + {len(K.SLIDES)}-slide deck"


def deck_class_line():
    parts = []
    for d in K.PACING:
        if d["slides"]:
            a, b = d["slides"]
            parts.append(f"Class {d['n']} (slide{'s' if a != b else ''} {rng(a, b)})")
    return ", ".join(parts[:-1]) + " and " + parts[-1]


def pacing_page(n, watermark=False):
    rows = []
    for k, d in enumerate(K.PACING):
        sl = f"Slides {rng(*d['slides'])}" if d["slides"] and d["slides"][0] != d["slides"][1] else (
            f"Slide {d['slides'][0]}" if d["slides"] else "No new slides")
        pg = ", ".join(f"p. {IX(k_)}" for k_ in d["pages"]) or "—"
        own = ""
        if d["owner"] is not None:   # owner None = this class shares the row above's text (e.g. the owner's "Classes 3–5")
            span_ = 1
            while k + span_ < len(K.PACING) and K.PACING[k + span_]["owner"] is None:
                span_ += 1
            own = f"<td rowspan='{span_}' style='font-size:11.2pt;line-height:1.4'>{d['owner']}</td>"
        rows.append(f"<tr><td style='white-space:nowrap'><b class='acc' style='font-family:Oswald;font-size:13pt'>CLASS {d['n']}</b></td>"
                    f"{own}<td style='font-size:10.4pt;line-height:1.36'>{fill(d['do'])}</td>"
                    f"<td style='font-size:10.4pt'>{sl}</td><td style='font-size:10.4pt'>{pg}</td></tr>")
    inner = (f'<p>{K.CLASSES} class periods, as on my project menu. The “My plan” column is my day-by-day pacing guide as I wrote it; the '
             f'other columns show what to put on the screen and in students’ hands.</p>'
             f'<table class="t"><tr><th style="width:11%">Class</th><th style="width:33%">My plan</th><th style="width:28%">Teacher moves</th>'
             f'<th style="width:14%">Deck</th><th>Student pages</th></tr>{"".join(rows)}</table>'
             f'{fill(K.PACING_EXTRA)}')
    return page(f"Pacing guide · {K.CLASSES} class periods", "PACING", "Teacher", inner, n, watermark, key="pacing")


def tn1_page(n):
    st = "".join(f'<li>{fill(s)}</li>' for s in K.STRUGGLES)
    sz = getattr(K, "TN1_SIZE", 11.6)
    inner = (f'<style>.tn1 p, .tn1 li {{ font-size: {sz}pt; line-height: 1.5; margin-bottom: {sz * 0.6:.1f}pt; }}</style><div class="tn tn1">{H2("What this project reinforces", True)}<p>{K.TN_REINFORCES}</p>'
             f'{H2("Common student struggles")}<ul>{st}</ul>'
             f'{H2("Suggested variations")}<p><b>Easier:</b> {K.VAR_EASY}</p><p><b>More Advanced:</b> {K.VAR_HARD}</p>'
             f'{fill(K.TN1_EXTRA)}</div>')
    return page("Teacher notes · the big idea", "TEACHER NOTES", "1 of 2", inner, n, key="tn1")


def tn2_page(n):
    res = "".join(f"<tr><td>{a}</td><td>{fill(b)}</td></tr>" for a, b in K.RESOURCES)
    inner = (f'<div class="tn">{H2("Suggested resources", True)}<table class="t" style="font-size:10.6pt"><tr><th style="width:44%">My list</th>'
             f'<th>In this kit</th></tr>{res}</table>'
             f'{H2("How to grade")}<ul>' + "".join(f"<li>{fill(g)}</li>" for g in K.GRADING) + '</ul>'
             f'{H2("Notes / observations")}<p>{K.TN_NOTE}</p></div><div class="lines"></div>')
    return page("Teacher notes · resources, grading, notes", "TEACHER NOTES", "2 of 2", inner, n, key="tn2")


def changes_page(n):
    lis = "".join(f"<li>{fill(x)}</li>" for x in K.CHANGES)
    inner = (f'<div class="big"><p>This kit is built from three of my documents: the <b>Project Sheet &amp; Rubric</b>, the <b>Teacher Notes</b> '
             f'and the <b>{K.TITLE}</b> entry on my DMP2 project menu. This page lists every change to my wording, so you know exactly what '
             f'is mine and what was added.</p>'
             f'{H2("Wording changes")}<ul>{lis}</ul>'
             f'{H2("Kept as I wrote it")}<p>{fill(K.KEPT)}</p>'
             f'{H2("New in this kit")}<p>{fill(K.NEW_MATERIAL)}</p></div>')
    return page("What changed from the original", "CHANGES", "Teacher", inner, n, key="changes")


def proj_head():
    return (f'<h2 class="s" style="font-size:21pt;margin:2pt 0 0">{escape(K.TITLE)}</h2>'
            f'<p class="l" style="font-size:12.5pt;margin-bottom:4pt">{K.PROJ_LINE}</p>')


def proj1_page(n, watermark=False):
    z = getattr(K, "PROJ_SIZE", 12)
    rows = "".join(f"<tr><td style='width:25%;padding:5pt 6pt'><b>{a}</b></td><td style='padding:5pt 6pt'>{b}</td></tr>" for a, b in K.VOCAB)
    inner = (proj_head() + f'{H2("Overview")}<p style="font-size:{z}pt;line-height:1.5">{K.OVERVIEW}</p>'
             f'{H2("Goals")}<ul>' + "".join(f"<li style='font-size:{z}pt;line-height:1.45;margin-bottom:5pt'>{g}</li>" for g in K.GOALS) + '</ul>'
             f'{H2("Key Vocabulary")}<table class="t" style="font-size:{z - 0.4}pt;line-height:1.4">{rows}</table>{fill(K.PROJ1_EXTRA)}')
    return page("Project sheet", "STUDENT", "1 of 3", inner, n, watermark, name=True, key="proj1")


def proj2_page(n, watermark=False):
    z = getattr(K, "PROJ_SIZE", 12)
    req = "".join(chk(fill(r), 13, f"font-size:{z}pt;margin-bottom:{z * 0.75:.1f}pt;line-height:1.35") for r in K.REQS)
    st = "".join(f'<li style="margin-bottom:{z * 0.85:.1f}pt"><div class="n">{i + 1}</div><div style="font-size:{z - 0.4}pt;line-height:1.42;padding-top:2pt">{fill(t)}</div></li>'
                 for i, t in enumerate(K.STEPS))
    inner = f'{H2("Requirements", True)}{req}{H2("Steps / Process")}<ol class="num">{st}</ol>'
    return page("Project sheet · requirements and steps", "STUDENT", "2 of 3", inner, n, watermark, name=True, key="proj2")


def proj3_page(n, watermark=False):
    inner = (f'{H2("Tips for Success", True)}<ul>' + "".join(f"<li style='font-size:11.6pt;line-height:1.4;margin-bottom:5pt'>{t}</li>" for t in K.TIPS)
             + '</ul>' + H2("Submission Checklist")
             + "".join(chk(fill(c), 12, "font-size:11.6pt;margin-bottom:6pt") for c in K.CHECK)
             + f'{H2("Notes")}<div class="lines"></div>')
    return page("Project sheet · tips and checklist", "STUDENT", "3 of 3", inner, n, watermark, name=True, key="proj3")


def rubric_page(n, watermark=False):
    head = "<tr><th>Category</th>" + "".join(f"<th>{lv}</th>" for lv in K.LEVELS) + "</tr>"
    rows = "".join(f"<tr><td class='c'>{escape(c)}</td>" + "".join(f"<td>{d}</td>" for d in ds) + "</tr>" for c, ds in K.RUBRIC)
    t = K.RUBRIC_TOTAL
    score = (f'<table class="t" style="margin-top:8pt;font-size:13pt"><tr><td style="padding:9pt 8pt">Student Score: ______ / {t}</td>'
             f'<td style="padding:9pt 8pt">Teacher Score: ______ / {t}</td></tr>'
             f'<tr><td colspan="2" style="padding:9pt 8pt"><b>Final Grade = Average of Student + Teacher = ______ / {t}</b></td></tr></table>')
    inner = (f'<h2 class="s" style="font-size:19pt;margin:2pt 0 0">Grading Rubric /{t}</h2>'
             f'<p class="l" style="font-size:12.5pt;margin-bottom:4pt">{escape(K.TITLE)} | {K.TYPE}</p>'
             f'<p style="font-size:11pt;margin-bottom:6pt">{K.RUBRIC_NOTE}</p>'
             f'<table class="rub">{head}{rows}</table>{score}{fill(K.RUBRIC_EXTRA)}')
    return page("Rubric", "RUBRIC", f"/ {t}", inner, n, watermark, name=True, key="rubric")


def vocab_page(n, watermark=False):
    cells = "".join(
        f'<div class="box" style="margin-bottom:{getattr(K, "VOCAB_GAP", 9)}pt;padding:8pt 10pt 5pt"><p style="margin:0 0 2pt;font-size:11.6pt"><b style="font-family:Oswald;'
        f'font-size:13pt;text-transform:uppercase;letter-spacing:.02em">{a}</b> — {b}</p>'
        f'<div class="fl" style="margin:9pt 0 9pt"><b>In my own words:</b><span></span></div>'
        f'<div class="fl" style="margin:0 0 4pt"><b>{K.VOCAB_PROMPT}</b><span></span></div></div>' for a, b in K.VOCAB)
    inner = (f'<p style="font-size:11.4pt;margin:0 0 7pt">The five words from your project sheet, with my definitions. Rewrite each in your '
             f'own words, then show where it happens in your own project.</p>{cells}')
    return page("Key vocabulary", "STUDENT", "Vocabulary", inner, n, watermark, name=True, key="vocab")


def present_page(n, watermark=False):
    qs = "".join(f'<div class="fl" style="margin-bottom:6pt"><b>Question {i}:</b><span></span></div><div class="ln2" style="margin-bottom:6pt"></div>'
                 for i in (1, 2, 3))
    rows = "".join("<tr><td style='height:40pt'></td><td></td><td></td></tr>" for _ in range(4))
    inner = (f'<div class="box acc-b"><span class="k">From your project sheet, step {len(K.STEPS)}</span><p style="margin:0;font-size:11.4pt">'
             f'{K.STEPS[-1]}</p></div>'
             f'{H2("My three feedback questions")}<p style="margin:0 0 6pt">Ask about things you really want to know, not yes-or-no questions.</p>{qs}'
             f'{H2("Classmates’ questions I answered (at least 2)")}'
             f'<table class="map"><tr><th style="width:24%">Classmate</th><th style="width:38%">Their question</th><th>My answer</th></tr>{rows}</table>'
             f'{H2("Before I present")}'
             + "".join(chk(fill(c), 12, "font-size:11.4pt;margin-bottom:6pt") for c in K.PRESENT_CHECK))
    return page("Present and get feedback", "STUDENT", f"Classes {K.CLASSES - 1}–{K.CLASSES}", inner, n, watermark, name=True, key="present")


def logicref_page(n, watermark=False):
    rows = "".join(f"<tr><td style='width:22%;padding:6pt 7pt'><b>{a}</b></td><td style='padding:6pt 7pt'>{b}</td></tr>" for a, b in K.LOGIC_REF)
    inner = (f'<p>{fill(K.LOGIC_INTRO)}</p><table class="t" style="font-size:{getattr(K, "LOGIC_SIZE", 11.4)}pt;line-height:1.45">{rows}</table>'
             f'<div class="box" style="margin-top:8pt"><span class="k">Version note</span>'
             f'<p style="margin:0" class="small">Menu names are from current Logic Pro versions; older versions may differ slightly. Logic Pro is a '
             f'trademark of Apple Inc., named only to describe what the project needs. This resource is not affiliated with Apple.</p></div>')
    return page("Logic quick reference", "SETUP", "Teacher + student", inner, n, watermark, key="logicref")


def terms_page(n, watermark=False):
    credits = "".join(f"<li>{c}</li>" for c in K.CREDITS)
    return (f'<section class="page" data-key="terms"><div class="top"><div class="t">Terms of Use<small>Thank you for your purchase!</small></div></div>'
            f'<div class="main"><p>© {YEAR} {STORE}. {BRAND} by {STORE}. All rights reserved. This purchase gives <b>one teacher</b> a license to use this '
            f'resource (the PDF and the intro deck) with their own students.</p>'
            f'<h2 class="s">You may</h2><ul><li>Print and copy pages for your own classroom and students, year after year.</li>'
            f'<li>Present the intro deck to your own classes and edit it for your own teaching.</li>'
            f'<li>Share pages with your own students through a password-protected class site or learning platform.</li>'
            f'<li>Leave the pages for a substitute teacher who teaches your own students.</li></ul>'
            f'<h2 class="s">You may not</h2><ul><li>Share, email or copy this resource for other teachers, a whole school or a district. '
            f'Additional licenses are available on Teachers Pay Teachers.</li>'
            f'<li>Post any part of it on a public website, shared drive or social media.</li>'
            f'<li>Sell, give away or claim any part of it as your own, or edit it into a new product.</li></ul>'
            f'<h2 class="s">Questions?</h2><p>Please use the Q&amp;A tab on the {BRAND} by {STORE} store on Teachers Pay Teachers. '
            f'If something looks wrong, let me know there and I will fix it.</p>'
            f'<h2 class="s">Credits</h2><ul>'
            f'<li>Project sheet, rubric, teacher notes and intro deck: original work by {STORE}, {BRAND}.</li>{credits}'
            f'<li>Text fonts: Oswald (The Oswald Project Authors) and Source Sans 3 (Adobe), SIL Open Font License 1.1.</li>'
            f'<li>Logic Pro is a trademark of Apple Inc., named only to describe compatibility; this resource is not affiliated with or endorsed '
            f'by Apple.</li></ul>'
            f'<p class="small">{DISCLOSURE}</p><div style="flex:1"></div></div>{foot(f"Page {n}")}</section>')


STD_PAGES = {"cover": lambda n, wm=False: cover_page(n), "pacing": pacing_page, "tn1": lambda n, wm=False: tn1_page(n),
             "tn2": lambda n, wm=False: tn2_page(n), "changes": lambda n, wm=False: changes_page(n), "proj1": proj1_page,
             "proj2": proj2_page, "proj3": proj3_page, "rubric": rubric_page, "vocab": vocab_page, "present": present_page,
             "logicref": logicref_page, "terms": terms_page}


def render_page(key, watermark=False):
    n = IX(key)
    fn = K.PAGES.get(key) or STD_PAGES[key]
    return fn(n, watermark)


# ================================================================ intro deck (.pptx, python-pptx) + PDF copy (LibreOffice)
def slide_class(i):
    return K.SLIDES[i - 1][0]


def build_deck(img):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

    BLACK, WHITE, ACC = RGBColor(0, 0, 0), RGBColor(255, 255, 255), RGBColor.from_string(K.ACCENT[1:])
    LIGHT = RGBColor.from_string(K.ACC_LIGHT[1:])
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.author = f"{STORE} · {BRAND}"
    prs.core_properties.title = f"{K.TITLE} · Intro deck"
    prs.core_properties.comments = DISCLOSURE
    layout = prs.slide_layouts[6]

    def box(sl, x, y, w, h, fillc=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=1.5):
        s = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        if fillc is None:
            s.fill.background()
        else:
            s.fill.solid()
            s.fill.fore_color.rgb = fillc
        if line is None:
            s.line.fill.background()
        else:
            s.line.color.rgb = line
            s.line.width = Pt(lw)
        s.shadow.inherit = False
        return s

    def text(sl, x, y, w, h, runs, size=24, font="Source Sans 3", bold=False, color=BLACK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             spacing=None, caps=False):
        tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.04)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        paras = runs if isinstance(runs, list) else [runs]
        for k, ptxt in enumerate(paras):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = ptxt.upper() if caps else ptxt
            r.font.size = Pt(size)
            r.font.name = font
            r.font.bold = bold
            r.font.color.rgb = color
            if spacing:
                r._r.get_or_add_rPr().set("spc", str(spacing))
        return tb

    def chrome(sl, i, title, cls):
        text(sl, 0.6, 0.32, 6, 0.4, BRAND, 13, "Oswald", False, BLACK, spacing=600, caps=True)
        box(sl, 11.18, 0.3, 1.55, 0.42, ACC, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        text(sl, 11.18, 0.3, 1.55, 0.42, f"Class {cls}", 15, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, caps=True)
        text(sl, 0.6, 0.72, 12.1, 0.95, title, 40, "Oswald", True, BLACK, caps=True, anchor=MSO_ANCHOR.BOTTOM)
        box(sl, 0.6, 1.72, 12.13, 0.04, BLACK)
        box(sl, 0.6, 7.0, 12.13, 0.012, BLACK)
        text(sl, 0.6, 7.04, 9, 0.3, f"{K.TITLE} · Intro deck · © {YEAR} {STORE} · {BRAND}", 10, "Source Sans 3", False, RGBColor(0x33, 0x33, 0x33))
        text(sl, 11.2, 7.04, 1.53, 0.3, f"{i}", 10, "Oswald", True, BLACK, align=PP_ALIGN.RIGHT)

    def bullets(sl, items, x=0.6, y=2.05, w=12.1, size=28, gap=1.75):
        for k, it in enumerate(items):
            yy = y + k * (size / 72 * gap)
            box(sl, x, yy + size / 72 * 0.36, 0.16, 0.16, ACC, shape=MSO_SHAPE.OVAL)
            text(sl, x + 0.38, yy, w - 0.4, size / 72 * 1.6, it, size, "Source Sans 3", False, BLACK)

    def pic(sl, key, x, y, w):
        path, (pw, ph) = img[key]
        sl.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))
        return y + w * ph / pw

    def fit(key, maxw, maxh):
        pw, ph = img[key][1]
        return min(maxw, maxh * pw / ph)

    for i, (cls, kind, title, body, notes) in enumerate(K.SLIDES, start=1):
        sl = prs.slides.add_slide(layout)
        items = [b for b in body.split("|") if b]
        if kind == "title":
            text(sl, 0, 0.42, 13.333, 0.5, BRAND, 16, "Oswald", False, BLACK, align=PP_ALIGN.CENTER, spacing=800, caps=True)
            w = fit("hero", 10.4, 3.25)
            pic(sl, "hero", (13.333 - w) / 2, 0.98, w)
            text(sl, 0, 4.3, 13.333, 1.1, title, 72, "Oswald", True, BLACK, align=PP_ALIGN.CENTER, caps=True)
            text(sl, 0, 5.5, 13.333, 0.6, body, 28, "Oswald", False, BLACK, align=PP_ALIGN.CENTER)
            box(sl, 0, 6.55, 13.333, 0.95, BLACK)
            text(sl, 0, 6.68, 13.333, 0.6, f"Digital Music Production 2 · Class {cls} · {K.DECK_STRAP}", 20, "Oswald",
                 False, WHITE, align=PP_ALIGN.CENTER, spacing=200, caps=True)
        else:
            chrome(sl, i, title, cls)
        if kind == "listen":
            head, _, sub = items[0].partition(" · ")
            box(sl, 0.6, 2.0, 12.13, 2.0, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            text(sl, 0.9, 2.1, 11.5, 0.7, head, 36, "Oswald", True, WHITE, caps=True)
            text(sl, 0.9, 2.82, 11.5, 0.5, sub, 24, "Oswald", False, LIGHT)
            text(sl, 0.9, 3.38, 11.5, 0.5, items[1], 18, "Source Sans 3", False, WHITE)
            bullets(sl, items[2:], y=4.3, size=30, gap=1.6)
        elif kind == "bullets":
            n_ = len(items)
            size = 34 if n_ <= 4 else 30
            bullets(sl, items, size=size, y=2.2, gap=2.2 if n_ <= 4 else 1.9)
        elif kind == "closer":
            bullets(sl, items, w=6.4, size=30, y=2.1, gap=1.9)
            w = fit("hero", 5.6, 4.6)
            pic(sl, "hero", 12.73 - w, 2.3, w)
        elif kind == "days":
            rowh = min(0.83, 4.85 / len(K.PACING))
            for k, d in enumerate(K.PACING):
                y = 1.98 + k * rowh
                box(sl, 0.6, y, 1.7, rowh - 0.12, ACC if d["slides"] else BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                text(sl, 0.6, y, 1.7, rowh - 0.12, f"Class {d['n']}", 22, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE,
                     caps=True)
                text(sl, 2.5, y, 10.2, rowh - 0.12, d["deck"], 22, "Source Sans 3", False, BLACK, anchor=MSO_ANCHOR.MIDDLE)
        elif kind == "pic":
            side_ = K.SIDE.get(body)
            w = fit(body, 12.13, 3.3 if side_ else 4.85)
            yb = pic(sl, body, 0.6 + (12.13 - w) / 2, 1.98, w)
            for k, s in enumerate(side_ or []):
                col, row = k % 2, k // 2
                x, yy = 0.6 + col * 6.17, max(yb + 0.3, 5.3) + row * 0.75
                box(sl, x, yy + 0.12, 0.16, 0.16, ACC, shape=MSO_SHAPE.OVAL)
                text(sl, x + 0.32, yy, 5.7, 0.6, s, 22, "Source Sans 3", False, BLACK)
        elif kind == "check":
            rows_ = math.ceil(len(items) / 2)
            step = min(1.5, 4.7 / rows_)
            for k, it in enumerate(items):
                col, row = k % 2, k // 2
                x, y = 0.6 + col * 6.15, 2.1 + row * step
                box(sl, x, y + 0.06, 0.4, 0.4, WHITE, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=2.25)
                text(sl, x + 0.6, y, 5.35, step - 0.08, it, 22, "Source Sans 3", False, BLACK)
        elif kind == "rubric":
            for k, ((c, _), short) in enumerate(zip(K.RUBRIC, K.RUBRIC_SHORT)):
                y = 1.98 + k * 0.9
                box(sl, 0.6, y, 4.2, 0.78, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                text(sl, 0.75, y, 3.9, 0.78, c, 19, "Oswald", True, WHITE, anchor=MSO_ANCHOR.MIDDLE, caps=True)
                text(sl, 5.0, y, 6.7, 0.78, "5 = " + short, 18, "Source Sans 3", False, BLACK, anchor=MSO_ANCHOR.MIDDLE)
                text(sl, 11.75, y, 0.98, 0.78, "/ 5", 24, "Oswald", True, ACC, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
            text(sl, 0.6, 6.5, 12.13, 0.45, K.RUBRIC_SLIDE_LINE, 20, "Oswald", True, BLACK, align=PP_ALIGN.RIGHT, caps=True)
        elif kind == "steps":
            gap = 0.95 if len(items) <= 5 else 0.8
            for k, it in enumerate(items):
                y = 2.05 + k * gap
                box(sl, 0.6, y + 0.05, 0.62, 0.62, BLACK, shape=MSO_SHAPE.OVAL)
                text(sl, 0.6, y + 0.05, 0.62, 0.62, str(k + 1), 24, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                text(sl, 1.45, y - 0.05, 11.2, 0.82, it, 28 if len(items) <= 5 else 26, "Source Sans 3", False, BLACK, anchor=MSO_ANCHOR.MIDDLE)
        sl.notes_slide.notes_text_frame.text = notes
    prs.save(P(f"{K.SLUG}-intro-deck.pptx"))


def ttf_fonts():
    """woff2 -> ttf for LibreOffice (deck PDF), plus a fontconfig file that adds them."""
    from fontTools.ttLib import TTFont
    d = os.path.join(build_dir(), "ttf")
    os.makedirs(d, exist_ok=True)
    styles = {"oswald-latin-500-normal": ("Oswald", "Regular", 400), "oswald-latin-700-normal": ("Oswald", "Bold", 700),
              "source-sans-3-latin-400-normal": ("Source Sans 3", "Regular", 400), "source-sans-3-latin-700-normal": ("Source Sans 3", "Bold", 700),
              "source-sans-3-latin-400-italic": ("Source Sans 3", "Italic", 400)}
    for stem, (fam, sub, wt) in styles.items():
        out = os.path.join(d, stem + ".ttf")
        if os.path.exists(out):
            continue
        t = TTFont(os.path.join(FONTS, stem + ".woff2"))
        t.flavor = None
        name = t["name"]
        ps = (fam.replace(" ", "") + "-" + sub)
        for rec in list(name.names):
            if rec.nameID in (16, 17, 21, 22):
                name.removeNames(nameID=rec.nameID)
        for nid, val in ((1, fam), (2, sub), (4, f"{fam} {sub}" if sub != "Regular" else fam), (6, ps)):
            name.setName(val, nid, 3, 1, 0x409)
            name.setName(val, nid, 1, 0, 0)
        t["OS/2"].usWeightClass = wt
        sel = 0
        if sub == "Bold":
            sel |= 1 << 5
        if sub == "Italic":
            sel |= 1
        if sub == "Regular":
            sel |= 1 << 6
        t["OS/2"].fsSelection = (t["OS/2"].fsSelection & ~0b1100001) | sel
        t["head"].macStyle = (1 if sub == "Bold" else 0) | (2 if sub == "Italic" else 0)
        t.save(out)
    conf = os.path.join(build_dir(), "fonts.conf")
    with open(conf, "w") as fh:
        fh.write(f'<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig><include ignore_missing="yes">/etc/fonts/fonts.conf</include>'
                 f'<dir>{d}</dir><cachedir>{os.path.join(build_dir(), "fccache")}</cachedir></fontconfig>')
    return conf


def deck_pdf():
    src, out = P(f"{K.SLUG}-intro-deck.pptx"), P(f"{K.SLUG}-intro-deck.pdf")
    env = dict(os.environ, FONTCONFIG_FILE=ttf_fonts())
    prof = os.path.join(build_dir(), "lo-profile")
    tmp = os.path.join(build_dir(), "lo-out")
    os.makedirs(tmp, exist_ok=True)
    subprocess.run(["soffice", f"-env:UserInstallation=file://{prof}", "--headless", "--convert-to", "pdf", "--outdir", tmp, src],
                   check=True, env=env, capture_output=True)
    shutil.move(os.path.join(tmp, os.path.basename(src).replace(".pptx", ".pdf")), out)


# ================================================================ square images (2000 x 2000)
SQ_CSS = """
.sq { width: 1000px; height: 1000px; position: relative; overflow: hidden; background: linear-gradient(180deg,#fff 0%,#fff 60%,#eeeeee 100%); }
.sq .store { position: absolute; left: 0; right: 0; top: 30px; text-align: center; font-family: 'Oswald'; font-weight: 500; letter-spacing: .32em;
             text-transform: uppercase; font-size: 22px; }
.sq .fillw { position: absolute; left: 0; right: 0; top: 66px; bottom: 70px; display: flex; flex-direction: column; align-items: center; justify-content: space-evenly; }
.sq .hd { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; text-align: center; line-height: .95; white-space: nowrap; }
.sq .sub { font-family: 'Oswald'; font-weight: 300; text-align: center; white-space: nowrap; }
.sq .band { position: absolute; left: 0; right: 0; bottom: 0; background: #000; color: #fff; text-align: center; font-family: 'Oswald';
            font-weight: 500; text-transform: uppercase; letter-spacing: .1em; font-size: 26px; padding: 16px 0 18px; white-space: nowrap; }
.sq .band em { font-style: normal; color: ACCENT; }
.shot { background: #fff; border: 2px solid #000; box-shadow: 8px 8px 0 #000; position: relative; overflow: hidden; }
.shot img { display: block; width: 100%; }
.lbl { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 22px; text-align: center; margin-top: 14px; }
.lbl span { font-weight: 300; text-transform: none; display: block; font-size: 18px; }
.swm { position: absolute; left: 50%; top: 50%; transform: translate(-50%,-50%) rotate(-38deg); font-family: 'Oswald'; font-weight: 700;
       color: rgba(0,0,0,.14); white-space: nowrap; letter-spacing: .06em; }
"""


def sq(inner, band):
    return (f'<!doctype html><html><head><meta charset="utf-8"><style>{font_css()}html,body{{margin:0;background:#fff}}'
            f'{SQ_CSS.replace("ACCENT", K.ACCENT)}</style></head><body><div class="sq"><div class="store">{BRAND}</div>'
            f'<div class="fillw">{inner}</div><div class="band">{band}</div></div></body></html>')


def sq_band():
    return f"Grades {K.GRADES_TEXT} · {K.CLASSES} classes · {len(K.PLAN)} pages + {len(K.SLIDES)} slides"


def cover_square():
    s = K.SQ_COVER
    return sq(f'<div style="width:{s.get("art", 900)}px">{body_svg(K.hero_body())}</div>'
              f'<div><div class="hd" style="font-size:{s["a"]}px">{K.COVER_A}</div>'
              + "".join(f'<div class="hd" style="font-size:{s["b"]}px;margin-top:4px">{escape(x)}</div>' for x in K.SQ_B_LINES) +
              f'<div class="sub" style="font-size:{s["c"]}px;margin-top:10px">{K.SQ_SUB}</div></div>',
              sq_band().replace(" · ", " <em>·</em> "))


def sq_head(i):
    h, hs, s, ss = K.SQ_HEADS[i]
    return f'<div><div class="hd" style="font-size:{hs}px">{h}</div><div class="sub" style="font-size:{ss}px;margin-top:8px">{s}</div></div>'


def preview_square_1(img):
    rot = [-4, 0, 4]
    shots = "".join(f'<div style="position:absolute;left:{x}px;top:{y}px;width:372px;transform:rotate({r}deg);z-index:{z}">'
                    f'<div class="shot"><img src="{s}"><div class="swm" style="font-size:80px">PREVIEW</div></div></div>'
                    for (s, _, _), x, y, r, z in zip(img, (40, 314, 588), (24, 0, 24), rot, (1, 3, 2)))
    labels = "".join(f'<div class="lbl" style="width:300px">{a}<span>{b}</span></div>' for _, a, b in img)
    return sq(f'{sq_head(1)}<div style="position:relative;width:1000px;height:520px">{shots}</div>'
              f'<div style="display:flex;gap:30px">{labels}</div>', K.SQ_BANDS[1])


def preview_square_2(slides):
    cells = "".join(f'<div class="shot" style="box-shadow:6px 6px 0 #000"><img src="{s}"><div class="swm" style="font-size:52px">PREVIEW</div></div>'
                    for s in slides)
    return sq(f'{sq_head(2)}<div style="display:grid;grid-template-columns:388px 388px;gap:20px 34px">{cells}</div>', K.SQ_BANDS[2])


def preview_square_3(imgs):
    cells = "".join(f'<div style="width:214px"><div class="shot" style="box-shadow:5px 5px 0 #000"><img src="{s}">'
                    f'<div class="swm" style="font-size:46px">PREVIEW</div></div><div class="lbl" style="font-size:18px;margin-top:9px">{lab_}</div></div>'
                    for s, lab_ in imgs)
    return sq(f'{sq_head(3)}<div style="display:grid;grid-template-columns:repeat(4,214px);gap:22px 24px">{cells}</div>', K.SQ_BANDS[3])


# ================================================================ render
def node_env():
    env = dict(os.environ)
    if not env.get("NODE_PATH"):
        env["NODE_PATH"] = subprocess.run(["npm", "root", "-g"], capture_output=True, text=True).stdout.strip()
    return env


def render(jobs):
    path = os.path.join(build_dir(), "jobs.json")
    with open(path, "w") as f:
        json.dump(jobs, f)
    subprocess.run(["node", RENDER_JS, path], check=True, env=node_env())


OVERFLOW_JS = r"""
const { chromium } = require('playwright');
const url = require('url');
(async () => {
  const browser = await chromium.launch();
  const page = await browser.newPage();
  await page.goto(url.pathToFileURL(process.argv[2]).href, { waitUntil: 'load' });
  await page.evaluate(() => document.fonts.ready);
  await page.emulateMedia({ media: 'print' });
  const res = await page.evaluate(() => [...document.querySelectorAll('section.page')].map((s, i) => {
    const m = s.querySelector('.main');
    const out = { page: i + 1, key: s.dataset.key, main_over: m ? Math.max(0, m.scrollHeight - m.clientHeight) : 0, cells: [], clipped: [] };
    if (m) {
      const mb = m.getBoundingClientRect();
      m.querySelectorAll('*').forEach(el => {
        if (el.closest('svg') && el.tagName.toLowerCase() !== 'svg') return;
        const r = el.getBoundingClientRect();
        if (r.width && (r.right > mb.right + 1 || r.bottom > mb.bottom + 1 || r.left < mb.left - 1))
          out.cells.push((el.tagName + ': ' + (el.textContent || '').trim()).slice(0, 50));
        if (el.children.length === 0 && el.scrollWidth > el.clientWidth + 1 && getComputedStyle(el).overflow !== 'visible')
          out.clipped.push((el.tagName + ': ' + (el.textContent || '').trim()).slice(0, 50));
      });
      let maxb = mb.top;
      m.querySelectorAll('*').forEach(el => { const r = el.getBoundingClientRect(); if (r.height && !el.classList.contains('lines')) maxb = Math.max(maxb, r.bottom); });
      out.free = Math.round(mb.bottom - maxb);
      if (m.querySelector('.lines')) out.free = 0;
    }
    return out;
  }));
  require('fs').writeFileSync(process.argv[3], JSON.stringify(res, null, 1));
  await browser.close();
})().catch(e => { console.error(e); process.exit(1); });
"""


def write(name, txt):
    p = os.path.join(build_dir(), name)
    with open(p, "w") as f:
        f.write(txt)
    return p


def overflow_scan(html):
    js = write("overflow.js", OVERFLOW_JS)
    subprocess.run(["node", js, html, os.path.join(build_dir(), "overflow.json")], check=True, env=node_env())


def pdf_png(src, page_no, out, res=110):
    subprocess.run(["pdftoppm", "-png", "-r", str(res), "-f", str(page_no), "-l", str(page_no), "-singlefile", src,
                    os.path.join(build_dir(), out)], check=True)


def art_doc(w, h, inner):
    return (f'<!doctype html><html><head><meta charset="utf-8"><style>{font_css()}{css()}'
            f'html,body{{margin:0;background:#fff}}</style></head>'
            f'<body><div style="width:{w}px;height:{h}px;overflow:hidden;background:#fff">{inner}</div></body></html>')


def deck_art():
    """PNG art for the deck: every drawing in K.DRAWINGS, rendered with the brand fonts at its own aspect ratio."""
    jobs, out = [], {}
    for k, fn in K.DRAWINGS.items():
        w = K.DRAWING_PX.get(k, 2000)
        svg = body_svg(fn())
        vw, vh = map(float, re.search(r'viewBox="0 0 ([\d.]+) ([\d.]+)"', svg).groups())
        h = int(round(w * vh / vw)) + 2
        p = os.path.join(build_dir(), f"deck-{k}.png")
        jobs.append({"html": write(f"deck-{k}.html", art_doc(w, h, f'<div style="width:{w}px">{svg}</div>')), "png": p, "width": w,
                     "height": h, "scale": 1})
        out[k] = (p, (w, h))
    return jobs, out


def files():
    s = K.SLUG
    return dict(pdf=P(f"{s}.pdf"), deck=P(f"{s}-intro-deck.pptx"), deck_pdf=P(f"{s}-intro-deck.pdf"), zip=P(f"{s}.zip"),
                preview=P(f"{s}-PREVIEW.pdf"), cover=P("cover.png"), pv=[P(f"preview-{i}.png") for i in (1, 2, 3)],
                upload=P("UPLOAD.md"), tpt=P("tpt.json"), overflow=os.path.join(build_dir(), "overflow.json"))


def tpt_json():
    F = files()
    tj = {
        "price": K.PRICE, "license_price": K.PRICE, "tax_code": TAX, "grades": K.GRADE_SLUGS, "subjects": ["Music"], "tags": K.TAGS,
        "formats": ["PDF"], "pages": len(K.PLAN) + len(K.SLIDES),
        "zip_contents": [os.path.basename(F["pdf"]), os.path.basename(F["deck"]), os.path.basename(F["deck_pdf"])],
        "files": {"product": os.path.basename(F["zip"]), "preview": os.path.basename(F["preview"]), "thumb1": "cover.png",
                  "thumb2": "preview-1.png", "thumb3": "preview-2.png", "thumb4": "preview-3.png"},
    }
    if K.ANSWER_KEY:            # only the proven option text "Included"; leave None when the kit has no answer key
        tj["answer_key"] = K.ANSWER_KEY
    return tj


def make_zip():
    F = files()
    if os.path.exists(F["zip"]):
        os.remove(F["zip"])
    with zipfile.ZipFile(F["zip"], "w", zipfile.ZIP_DEFLATED) as z:
        for p in (F["pdf"], F["deck"], F["deck_pdf"]):
            zi = zipfile.ZipInfo(f"{K.ZIP_DIR}/{os.path.basename(p)}", date_time=K.ZIP_DATE)
            zi.compress_type = zipfile.ZIP_DEFLATED
            with open(p, "rb") as fh:
                z.writestr(zi, fh.read())


def build():
    os.makedirs(build_dir(), exist_ok=True)
    problems = K.verify_content() + verify_data() + verify_labels()
    if problems:
        print("\n".join(problems))
        sys.exit("data check failed; not rendering")
    F = files()
    full = write("resource.html", html_doc([render_page(p) for p in K.PLAN], K.TITLE))
    prev = write("preview.html", html_doc([render_page(p, watermark=True) for p in K.PREVIEW_PICKS], "Preview"))
    art_jobs, art = deck_art()
    render([{"html": full, "pdf": F["pdf"]}, {"html": prev, "pdf": F["preview"]}] + art_jobs)
    overflow_scan(full)
    build_deck(art)
    deck_pdf()
    make_zip()
    lr = lowres()
    for key, _, _ in K.SHOTS1:
        pdf_png(F["pdf"], IX(key), key, 45 if lr else 120)
    for key, _ in K.SHOTS3:
        pdf_png(F["pdf"], IX(key), key, 40 if lr else 80)
    for s in K.SHOTS2:
        pdf_png(F["deck_pdf"], s, f"slide{s}", 30 if lr else 70)
    sqs = [write("cover.html", cover_square()), write("p1.html", preview_square_1([(f"{k}.png", a, b) for k, a, b in K.SHOTS1])),
           write("p2.html", preview_square_2([f"slide{s}.png" for s in K.SHOTS2])),
           write("p3.html", preview_square_3([(f"{k}.png", lab_) for k, lab_ in K.SHOTS3]))]
    render([{"html": h, "png": o, "width": 1000, "height": 1000, "scale": 1 if lr else 2} for h, o in zip(sqs, [F["cover"]] + F["pv"])])
    with open(F["tpt"], "w") as f:
        json.dump(tpt_json(), f, indent=2, ensure_ascii=False)
        f.write("\n")
    print(f"built {F['pdf']} ({len(K.PLAN)} pages), deck .pptx + PDF ({len(K.SLIDES)} slides), zip, preview, cover + 3 previews, tpt.json"
          + (" [LOW RES]" if lr else ""))


# ================================================================ checks
def verify_data():
    """Owner lists present, rubric shape and total, pacing <-> deck class slots, cover/square text fit, listing title."""
    errs = []
    if any(len(dd) != len(K.LEVELS) for _, dd in K.RUBRIC) or max(K.LEVELS) * len(K.RUBRIC) != K.RUBRIC_TOTAL:
        errs.append(f"rubric must be {len(K.RUBRIC)} categories x levels {K.LEVELS} = /{K.RUBRIC_TOTAL}")
    if len(K.RUBRIC_SHORT) != len(K.RUBRIC):
        errs.append("one short rubric line per category (deck)")
    # pacing: classes 1..N; every slide in exactly one class slot, matching its Class badge; slides run in class order
    if [d["n"] for d in K.PACING] != list(range(1, K.CLASSES + 1)):
        errs.append(f"pacing must list classes 1-{K.CLASSES}")
    seen = {}
    for d in K.PACING:
        if d["slides"]:
            a, b = d["slides"]
            for s in range(a, b + 1):
                if s in seen:
                    errs.append(f"slide {s} is slotted twice (Class {seen[s]} and Class {d['n']})")
                seen[s] = d["n"]
                if not 1 <= s <= len(K.SLIDES):
                    errs.append(f"Class {d['n']} names slide {s}, which does not exist")
                elif slide_class(s) != d["n"]:
                    errs.append(f"slide {s} says Class {slide_class(s)} but the Class {d['n']} plan uses it")
        for a_, b_ in re.findall(r"[Ss]lides? (\d+)(?:–(\d+))?", d["do"]):
            named = set(range(int(a_), int(b_ or a_) + 1))
            if not d["slides"] or not named <= set(range(d["slides"][0], d["slides"][1] + 1)):
                errs.append(f"Class {d['n']} teacher moves name slides {sorted(named)} outside its slot {d['slides']}")
        for k in d["pages"]:
            if k not in K.PLAN:
                errs.append(f"Class {d['n']} lists page '{k}', not in the plan")
    for i in range(1, len(K.SLIDES) + 1):
        if i not in seen:
            errs.append(f"deck slide {i} has no class slot in the pacing guide")
    if [s[0] for s in K.SLIDES] != sorted(s[0] for s in K.SLIDES):
        errs.append("deck slides must run in class order")
    if any(len(s[4].strip()) < 120 for s in K.SLIDES):
        errs.append("every slide needs real speaker notes (120+ characters)")
    kinds = [s[1] for s in K.SLIDES]
    for need in ("title", "listen", "bullets", "pic", "steps", "check", "rubric", "closer"):
        if need not in kinds:
            errs.append(f"deck pattern: no '{need}' slide")
    for _, k, _, body, _ in K.SLIDES:
        if k == "pic" and body not in K.DRAWINGS:
            errs.append(f"deck picture '{body}' is not a drawing")
    if len(K.PLAN) != len(set(K.PLAN)) or K.PLAN[0] != "cover" or K.PLAN[-1] != "terms" or "changes" not in K.PLAN:
        errs.append("plan: cover first, terms last, a 'changes' page, no repeats")
    for k in K.PLAN:
        if k not in K.PAGES and k not in STD_PAGES:
            errs.append(f"page '{k}' has no builder")
    if len(K.LISTING_TITLE) > 80 or "Logic Pro" not in K.LISTING_TITLE[:30]:
        errs.append(f"listing title is {len(K.LISTING_TITLE)} chars, or 'Logic Pro' is not in the first 30")
    room = 7.3 * 72
    for label, txt, sz in (("cover A", K.COVER_A.upper(), K.PDF_COVER["a"]), ("cover B", K.TITLE.upper(), K.PDF_COVER["b"])):
        if text_w(txt, sz) * 1.01 > room:
            errs.append(f"PDF {label} '{txt}' is {text_w(txt, sz):.0f}pt wide, room {room:.0f}")
    if text_w(K.COVER_C, K.PDF_COVER["c"]) > room:
        errs.append("PDF cover subtitle too wide")
    if text_w(cover_band().upper(), 14) * 1.12 > room:
        errs.append("PDF cover band too wide")
    if " ".join(K.SQ_B_LINES) != K.TITLE:
        errs.append("square cover lines must spell the title")
    S = K.SQ_COVER
    for txt, sz in [(K.COVER_A.upper(), S["a"]), (K.SQ_SUB, S["c"])] + [(x.upper(), S["b"]) for x in K.SQ_B_LINES]:
        if text_w(txt, sz) * 1.01 > 940:
            errs.append(f"square cover line '{txt}' is {text_w(txt, sz):.0f}px wide, room 940")
    if text_w(sq_band().upper(), 26) * 1.12 > 960:
        errs.append("square cover band too wide")
    if S["b"] * 0.75 < 100:
        errs.append("square cover big line under 10% of the image height")
    if any(len(x.split()) == 1 and len(K.SQ_B_LINES) > 1 and len(x) < 4 for x in K.SQ_B_LINES):
        errs.append("square cover: a very short word alone on a title line (orphan)")
    for i, (h, hs, s, ss) in K.SQ_HEADS.items():
        if text_w(h.upper(), hs) * 1.01 > 960 or text_w(s, ss) > 960:
            errs.append(f"preview {i} headline/subtitle too wide")
    if set(K.GRADE_SLUGS) - PROVEN_GRADES or not 1 <= len(K.GRADE_SLUGS) <= 4:
        errs.append("grades: 1-4 proven slugs only")
    if K.ANSWER_KEY not in (None, "Included"):
        errs.append("answer_key: only the proven 'Included', or None")
    if not set(K.TAGS) <= PROVEN_TAGS or not 1 <= len(K.TAGS) <= 6:
        errs.append("tags: proven values only (UPLOAD.md lists candidate extras)")
    return errs


def pdf_fonts(path):
    out = subprocess.run(["pdffonts", path], capture_output=True, text=True, check=True).stdout.splitlines()[2:]
    return [(ln.split()[0], ln.split()[-5]) for ln in out]


def pdf_text(path):
    return subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True, check=True).stdout


def page_texts(path):
    from pypdf import PdfReader
    n = len(PdfReader(path).pages)
    return [re.sub(r"\s+", " ", subprocess.run(["pdftotext", "-layout", "-f", str(i + 1), "-l", str(i + 1), path, "-"],
                                               capture_output=True, text=True, check=True).stdout) for i in range(n)]


def words_in(item, text):
    """Every word of 6+ letters in item appears in text (hyphenation- and line-break-proof check that owner text is printed)."""
    t = text.replace("- ", "").lower()      # lower case: headings and table heads are set in caps by CSS
    return [w for w in re.findall(r"[A-Za-z]{6,}", re.sub(r"<[^>]+>", "", item)) if w.lower() not in t]


def deck_texts():
    from pptx import Presentation
    F = files()
    prs = Presentation(F["deck"])
    slides, notes = [], []
    for s in prs.slides:
        slides.append(" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame))
        notes.append(s.notes_slide.notes_text_frame.text if s.has_notes_slide else "")
    media = [p for p in zipfile.ZipFile(F["deck"]).namelist() if p.startswith("ppt/media/")]
    return prs, slides, notes, media


def deck_overflow(prs):
    """Every shape stays inside the slide; Oswald text fits its box; Source Sans text boxes hold their wrapped lines; no two text
    boxes overlap; free text keeps 0.1 in clear of filled boxes; no text over a picture; pictures stay above the footer rule."""
    errs = []
    W, H = prs.slide_width, prs.slide_height
    for i, s in enumerate(prs.slides, start=1):
        boxes = []
        for sh in s.shapes:
            if sh.left < 0 or sh.top < 0 or sh.left + sh.width > W + 10 or sh.top + sh.height > H + 10:
                errs.append(f"deck slide {i}: a shape runs off the slide")
            if sh.has_text_frame and sh.text_frame.text:
                r = sh.text_frame.paragraphs[0].runs
                if r and r[0].font.name == "Oswald" and r[0].font.size:
                    widest = max((text_w(p.text, r[0].font.size.pt) for p in sh.text_frame.paragraphs), default=0)
                    room = sh.width / 12700 - 6
                    lines_avail = int((sh.height / 12700) // (r[0].font.size.pt * 1.15)) or 1
                    if widest > room * lines_avail:
                        errs.append(f"deck slide {i}: '{sh.text_frame.text[:30]}' ({widest:.0f}pt) does not fit {room:.0f}pt x {lines_avail} line(s)")
                if r and r[0].font.name == "Source Sans 3" and r[0].font.size and sh.shape_type == 17:
                    sz = r[0].font.size.pt
                    room = sh.width / 12700 - 6
                    need = sum(max(1, math.ceil(text_w(p.text, sz, "ss") / room)) for p in sh.text_frame.paragraphs)
                    if need * sz * 1.2 > sh.height / 12700 + sz * 0.45:
                        errs.append(f"deck slide {i}: '{sh.text_frame.text[:30]}' wraps to {need} line(s), box holds less (overlap risk)")
                if sh.shape_type == 17:
                    boxes.append((sh.left, sh.top, sh.left + sh.width, sh.top + sh.height, sh.text_frame.text[:20]))
        pics = [(sh.left, sh.top, sh.left + sh.width, sh.top + sh.height) for sh in s.shapes if sh.shape_type == 13]
        for Pq in pics:
            if Pq[3] > 914400 * 6.95:
                errs.append(f"deck slide {i}: a picture runs into the footer")
        pad = 914400 // 10
        fills = [(sh.left, sh.top, sh.left + sh.width, sh.top + sh.height) for sh in s.shapes
                 if sh.shape_type == 1 and sh.width > 914400 // 2 and sh.height > 914400 // 2]

        def inside(A, Fb):
            return A[0] >= Fb[0] - 12700 and A[2] <= Fb[2] + 12700 and A[1] >= Fb[1] - 12700 and A[3] <= Fb[3] + 12700
        for A in boxes:
            if any(inside(A, Fb) for Fb in fills):
                continue
            for Fb in fills:
                ox = min(A[2], Fb[2]) - max(A[0], Fb[0])
                near = min(A[3], Fb[3] + pad) - max(A[1], Fb[1] - pad) > 0
                if ox > 12700 and near:
                    errs.append(f"deck slide {i}: text '{A[4]}' touches a box (needs 0.1 in clearance)")
        for Pq in pics:
            for A in boxes:
                if min(A[2], Pq[2]) - max(A[0], Pq[0]) > 12700 * 4 and min(A[3], Pq[3]) - max(A[1], Pq[1]) > 12700 * 4:
                    errs.append(f"deck slide {i}: text '{A[4]}' overlaps a picture")
        for a in range(len(boxes)):
            for b in range(a + 1, len(boxes)):
                A, B = boxes[a], boxes[b]
                ox = min(A[2], B[2]) - max(A[0], B[0])
                oy = min(A[3], B[3]) - max(A[1], B[1])
                if ox > 12700 * 4 and oy > 12700 * 4:
                    inside_ = (A[0] >= B[0] and A[2] <= B[2] and A[1] >= B[1] and A[3] <= B[3]) or \
                              (B[0] >= A[0] and B[2] <= A[2] and B[1] >= A[1] and B[3] <= A[3])
                    if not inside_:
                        errs.append(f"deck slide {i}: text boxes '{A[4]}' and '{B[4]}' overlap")
    return errs


def verify_files():
    from pypdf import PdfReader
    from PIL import Image
    F = files()
    errs = []
    for p in [F["pdf"], F["deck"], F["deck_pdf"], F["zip"], F["preview"], F["cover"], *F["pv"], F["upload"], F["tpt"], F["overflow"]]:
        if not os.path.exists(p):
            errs.append(f"missing {os.path.basename(p)}")
    if errs:
        return errs
    for path, n_exp in ((F["pdf"], len(K.PLAN)), (F["preview"], 4)):
        rd = PdfReader(path)
        if len(rd.pages) != n_exp:
            errs.append(f"{os.path.basename(path)}: {len(rd.pages)} pages, expected {n_exp}")
        for i, pg_ in enumerate(rd.pages):
            if abs(float(pg_.mediabox.width) - 612) > 1 or abs(float(pg_.mediabox.height) - 792) > 1:
                errs.append(f"{os.path.basename(path)} p{i + 1}: not US Letter")
    if len(K.PREVIEW_PICKS) != 4:
        errs.append("preview must be 4 pages")
    for o in json.load(open(F["overflow"])):
        if o["main_over"] > 1:
            errs.append(f"page {o['page']} ({o['key']}): content overflows by {o['main_over']}px (clipped)")
        for c in o["cells"] + o.get("clipped", []):
            errs.append(f"page {o['page']} ({o['key']}): text wider than its box: '{c}'")
        if o["key"] not in ("cover", "terms") and o.get("free", 0) > 200:
            errs.append(f"page {o['page']} ({o['key']}): {o['free']}px empty at the bottom (dead space)")
    drd = PdfReader(F["deck_pdf"])
    if len(drd.pages) != len(K.SLIDES):
        errs.append(f"deck PDF: {len(drd.pages)} pages, expected {len(K.SLIDES)}")
    for i, pg_ in enumerate(drd.pages):
        if abs(float(pg_.mediabox.width) / float(pg_.mediabox.height) - 16 / 9) > 0.01:
            errs.append(f"deck PDF p{i + 1}: not 16:9")
    for path in (F["pdf"], F["preview"], F["deck_pdf"]):
        fonts = pdf_fonts(path)
        for name, emb in fonts:
            if emb != "yes":
                errs.append(f"{os.path.basename(path)}: font {name} not embedded")
            if not re.search(r"oswald|sourcesans", name.lower().replace(" ", "")):
                errs.append(f"{os.path.basename(path)}: unexpected font {name} (fallback?)")
        names = " ".join(n for n, _ in fonts).lower()
        for fam in ("oswald", "sourcesans3"):
            if fam not in names:
                errs.append(f"{os.path.basename(path)}: {fam} missing")
    prs, stext, notes, media = deck_texts()
    if len(prs.slides) != len(K.SLIDES):
        errs.append(f"deck has {len(prs.slides)} slides, expected {len(K.SLIDES)}")
    for i, nt in enumerate(notes):
        if len(nt.strip()) < 120:
            errs.append(f"deck slide {i + 1}: speaker notes missing or too short")
        if not re.search(rf"\bClass {slide_class(i + 1)}\b", stext[i], re.I):
            errs.append(f"deck slide {i + 1}: no 'Class {slide_class(i + 1)}' slot shown on the slide")
    errs += deck_overflow(prs)
    if any(not m.lower().endswith(".png") for m in media) or len(set(media)) > len(K.DRAWINGS):
        errs.append(f"deck media should be only the drawn PNGs, found {media}")
    if abs(prs.slide_width / prs.slide_height - 16 / 9) > 0.01:
        errs.append("deck is not 16:9")
    alldeck = " ".join(stext + notes)
    names = zipfile.ZipFile(F["zip"]).namelist()
    for p in (F["pdf"], F["deck"], F["deck_pdf"]):
        n = f"{K.ZIP_DIR}/{os.path.basename(p)}"
        if n not in names:
            errs.append(f"zip is missing {os.path.basename(p)}")
        elif zipfile.ZipFile(F["zip"]).read(n) != open(p, "rb").read():
            errs.append(f"zip copy of {os.path.basename(p)} is stale; rebuild")
    if len(names) != 3:
        errs.append(f"zip should hold exactly 3 files, has {names}")
    if any(re.search(r"\.(wav|mp3|aif|aiff|m4a|logicx|mov|mp4)$", nm, re.I) for nm in names):
        errs.append("zip holds audio, video or a Logic session")
    flat_ = page_texts(F["pdf"])
    raw0 = subprocess.run(["pdftotext", "-f", "1", "-l", "1", F["pdf"], "-"], capture_output=True, text=True, check=True).stdout
    lines0 = [re.sub(r"\s+", " ", ln.strip()) for ln in raw0.splitlines() if ln.strip()]
    for need in (K.COVER_A.upper(), K.TITLE.upper(), K.COVER_C):
        if need not in lines0:
            errs.append(f"PDF cover: '{need}' is not on a line of its own (orphaned word?)")
    for i, t in enumerate(flat_):
        if i > 0 and "Brian Lotze · Hudson Beat" not in t:
            errs.append(f"page {i + 1}: footer missing")
        if i > 0 and not re.search(rf"Page {i + 1}\s*$", t.strip()):
            errs.append(f"page {i + 1}: footer page number wrong (overflow?)")
    alltext = " ".join(flat_[1:])
    for m in re.finditer(r"\b(?:pages?|pp?\.) (\d+)(?:–(\d+))?", alltext):
        for g in m.groups():
            if g and not 1 <= int(g) <= len(K.PLAN):
                errs.append(f"cross-reference to page {g} does not exist")
    for m in re.finditer(r"\b[Ss]lides? (\d+)(?:–(\d+))?", alltext):
        for g in m.groups():
            if g and not 1 <= int(g) <= len(K.SLIDES):
                errs.append(f"PDF names deck slide {g}, which does not exist")
    # the owner's words are printed where they belong
    owner = {"proj1": [K.OVERVIEW] + K.GOALS + [a + " " + b for a, b in K.VOCAB], "proj2": K.REQS + K.STEPS,
             "proj3": K.TIPS + K.CHECK, "rubric": [d for _, ds in K.RUBRIC for d in ds] + [c for c, _ in K.RUBRIC] + [K.RUBRIC_NOTE],
             "tn1": [K.TN_REINFORCES, K.VAR_EASY, K.VAR_HARD] + K.STRUGGLES, "tn2": [a for a, _ in K.RESOURCES] + [K.TN_NOTE],
             "pacing": [d["owner"] for d in K.PACING if d["owner"]]}
    for pk, lst in owner.items():
        if pk not in K.PLAN:
            continue
        for it in lst:
            miss = words_in(fill(it), flat_[IX(pk) - 1])
            if miss:
                errs.append(f"{pk}: owner text missing words {miss[:4]}: {it[:50]}")
    rt = flat_[IX("rubric") - 1]
    if rt.count(f"/ {K.RUBRIC_TOTAL}") < 3 or f"/{K.RUBRIC_TOTAL}" not in rt:
        errs.append(f"rubric page: the owner's /{K.RUBRIC_TOTAL} total (student, teacher, final) missing")
    ch = flat_[IX("changes") - 1]
    for need in K.CHANGES_MUST_MENTION:
        if need not in ch:
            errs.append(f"'What changed' page must mention '{need}'")
    if re.search(r"@[A-Z]{3,}", alltext):
        errs.append("an @PLACEHOLDER was printed unfilled")
    if "word for word" in alltext.lower():
        errs.append("never claim 'word for word'")
    ptxt = [pg.extract_text() for pg in PdfReader(F["preview"]).pages]
    if sum("PREVIEW" in t for t in ptxt) != 4:
        errs.append("preview PDF: watermark missing on some pages")
    blobs = {"resource PDF": pdf_text(F["pdf"]), "preview PDF": pdf_text(F["preview"]), "deck PDF": pdf_text(F["deck_pdf"]),
             "deck pptx": alldeck, "zip names": " ".join(names)}
    with zipfile.ZipFile(F["deck"]) as z:
        blobs["deck XML"] = " ".join(z.read(n).decode("utf8", "ignore") for n in z.namelist() if n.endswith(".xml") or n.endswith(".rels"))
    for path in (F["pdf"], F["preview"], F["deck_pdf"]):
        meta = PdfReader(path).metadata or {}
        blobs[f"{os.path.basename(path)} metadata"] = " ".join(str(v) for v in meta.values())
    for fn in os.listdir(K.HERE):
        if fn == "QA.md":
            continue
        if fn.endswith((".md", ".json", ".py", ".txt")) or fn in ("release", "uploaded"):
            blobs[fn] = open(os.path.join(K.HERE, fn), encoding="utf8", errors="ignore").read()
    for fn in os.listdir(build_dir()):
        if fn.endswith(".html"):
            blobs["build/" + fn] = open(os.path.join(build_dir(), fn), encoding="utf8", errors="ignore").read()
    for label, txt in blobs.items():
        low = txt.lower()
        for b in BANNED + getattr(K, "BANNED_EXTRA", []):
            if b.lower() in low:
                errs.append(f"{label}: contains banned string '{b}'")
    for p in [F["cover"], *F["pv"]]:
        if Image.open(p).size != (2000, 2000):
            errs.append(f"{os.path.basename(p)}: not 2000x2000 (low-res build?)")
    tj = json.load(open(F["tpt"]))
    if tj != tpt_json():
        errs.append("tpt.json is stale; rebuild")
    if tj["price"] != K.PRICE or tj["license_price"] != K.PRICE or tj["tax_code"] != TAX:
        errs.append("tpt.json price/license_price/tax_code")
    for k, f in tj["files"].items():
        if not os.path.exists(P(f)):
            errs.append(f"tpt.json file {k} missing: {f}")
    if tj["pages"] != len(flat_) + len(prs.slides):
        errs.append("tpt.json pages != PDF pages + slides")
    up_ = open(F["upload"]).read()
    m = re.search(r"^## 1\. Title\s*\n+```\n(.+?)\n```\n\((\d+) characters", up_, re.M | re.S)
    if not m:
        errs.append("UPLOAD.md: title block or its character count not found")
    else:
        n = len(m.group(1).strip())
        if n > 80 or int(m.group(2)) != n or m.group(1).strip() != K.LISTING_TITLE:
            errs.append(f"UPLOAD.md: title {n} chars (max 80), count stated {m.group(2)}, must equal LISTING_TITLE")
    m2 = re.search(r"^## 2\. Description\s*\n+```\n(.+?)\n```", up_, re.M | re.S)
    if not m2 or not m2.group(1).rstrip().endswith(DISCLOSURE):
        errs.append("UPLOAD.md: description must end with the disclosure line")
    else:
        desc = m2.group(1)
        if not desc.startswith(K.DESC_START):
            errs.append("UPLOAD.md description: first line must lead with the search phrase (METADATA-GUIDE)")
        for need in [f"{len(K.PLAN)}-page", f"{len(K.SLIDES)}-slide", "Logic Pro", f"/{K.RUBRIC_TOTAL}", "GarageBand",
                     f"{K.CLASSES} class periods"] + K.DESC_MUST:
            if need not in desc:
                errs.append(f"UPLOAD.md description: missing '{need}'")
        if re.search(r"  +\S|\s[,.]\s|\bteh\b|\s{2,}$", desc, re.M):
            errs.append("UPLOAD.md description: double space or stray punctuation")
        if "any DAW" in desc or "word for word" in desc.lower():
            errs.append("UPLOAD.md description: unverifiable claim")
    for need in (f"${K.PRICE}", f"{len(K.PLAN)} pages", os.path.basename(F["zip"]), os.path.basename(F["deck"]), DISCLOSURE):
        if need not in up_:
            errs.append(f"UPLOAD.md: missing '{need}'")
    for f in ("release", "uploaded"):
        if os.path.exists(P(f)):
            errs.append(f"{f} exists: only the Manager / uploader writes it")
    errs += K.verify_outputs(flat_, stext, notes)
    return errs


def check():
    errs = K.verify_content() + verify_data() + verify_labels() + verify_files()
    if errs:
        print("CHECK FAILED")
        print("\n".join(" - " + e for e in errs))
        sys.exit(1)
    print(f"CHECK OK: {len(K.PLAN)} Letter pages (no overflow, no clipped text, no dead space, footers, valid page/slide references, "
          f"cover lines whole); owner's text printed on its pages; rubric {len(K.RUBRIC)} x {K.LEVELS[0]}-{K.LEVELS[-1]} = /{K.RUBRIC_TOTAL}; "
          f"{K.CLASSES} classes, each of {len(K.SLIDES)} slides in exactly one class slot matching its badge; notes on every slide; deck text "
          f"in bounds, no overlaps; {len(LABELS)} drawing labels fit; preview 4 watermarked pages; only Oswald/Source Sans 3, embedded; no "
          f"banned strings; zip current, no audio; images 2000x2000; tpt.json proven values; UPLOAD.md title {len(K.LISTING_TITLE)} chars, "
          f"SEO opening and disclosure. {K.check_summary()}")


def main(kit):
    global K
    K = kit
    if "--check" in sys.argv:
        check()
    else:
        build()
