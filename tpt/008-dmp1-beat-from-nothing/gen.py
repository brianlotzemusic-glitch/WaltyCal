#!/usr/bin/env python3
"""DMP1 Unit 1: Beat From Nothing (TpT 008): build the teacher resource, the intro deck and the listing files.

    python3 gen.py            build everything into this folder
    python3 gen.py --lowres   same, but square images at scale 1 and small page shots (self-check only; --check rejects them)
    python3 gen.py --check    verify data, files, page/slide counts, speaker notes, fonts, banned strings, tpt.json and UPLOAD.md

The owner's own curriculum (Brian Lotze, Digital Music Production 1, Unit 1), reformatted in the Hudson Beat
style: a 4-day drum-programming unit in Logic Pro (Drum Machine Designer + Step Sequencer). Source documents
(owner's Drive, read-only): Project Sheet & Rubric, Teacher Notes, Day 1 Classroom Post, Famous Drum Patterns,
Drumbeat Listening Examples, Parts of the Drumset, the "Looping in Logic" deck and the DMP1 Restructure (4 days, 2/1/1).
School names and logos are removed; no student work, names or grades; no audio, video or album art. Songs are
named by title and artist only, and the teacher is told to play them from YouTube or Spotify.

Outputs: <slug>.pdf (teacher + student resource), <slug>-intro-deck.pptx and its PDF copy, a ZIP of those three
(the TpT product file), a 4-page watermarked preview PDF, cover.png and preview-1..3.png (2000 x 2000).
Renders with Playwright + Chromium (../tools/render.js) and LibreOffice (deck PDF), fonts from ../fonts.
"""
import json
import os
import re
import shutil
import subprocess
import sys
import zipfile
from html import escape

HERE = os.path.dirname(os.path.abspath(__file__))
TPT = os.path.dirname(HERE)
FONTS = os.path.join(TPT, "fonts")
BUILD = os.path.join(HERE, "build")
RENDER_JS = os.path.join(TPT, "tools", "render.js")

SLUG = "beat-from-nothing"
PDF = os.path.join(HERE, f"{SLUG}.pdf")
DECK = os.path.join(HERE, f"{SLUG}-intro-deck.pptx")
DECK_PDF = os.path.join(HERE, f"{SLUG}-intro-deck.pdf")
ZIP = os.path.join(HERE, f"{SLUG}.zip")
PREVIEW_PDF = os.path.join(HERE, f"{SLUG}-PREVIEW.pdf")
COVER_PNG = os.path.join(HERE, "cover.png")
PREVIEW_PNGS = [os.path.join(HERE, f"preview-{i}.png") for i in (1, 2, 3)]
UPLOAD = os.path.join(HERE, "UPLOAD.md")
TPT_JSON = os.path.join(HERE, "tpt.json")

STORE = "Brian Lotze"      # copyright holder and TpT seller
BRAND = "Hudson Beat"      # store line on covers and previews (owner, 4 Oct)
YEAR = 2026
ACCENT = "#E03A3E"         # this product's single accent: drum-machine red
TITLE = "Beat From Nothing"
PRICE = "9.00"
LOW_RES = "--lowres" in sys.argv
DISCLOSURE = "Designed with the help of digital and AI tools, and checked by hand."

# Strings that must never appear in any output (school branding, student-data places, unlicensed media words)
BANNED = ["Blind Brook", "BBHS", "blindbrook", "Google Classroom", "Thunderstruck", "lyrics:"]

# ---------------------------------------------------------------- drum patterns (owner's sheets, 16 steps = one 4/4 bar in sixteenths)
STEP_LABELS = [f"{b}" if s == 1 else f"{b}.{s}" for b in range(1, 5) for s in range(1, 5)]   # 1 1.2 1.3 1.4 2 ... (Logic's positions)
FAMOUS = [  # (song, artist, tempo, rows) from the owner's "Famous Drum Patterns" sheet
    ("Billie Jean", "Michael Jackson", 116, [("Hi-hat", "x.x.x.x.x.x.x.x."), ("Snare", "....x.......x..."), ("Kick", "x.......x.......")]),
    ("Closer", "The Chainsmokers", 96, [("Hi-hat", "x.x.x.x.x.x.x.x."), ("Snare", "....x.......x..."), ("Kick", "x.........x....."),
                                         ("Clap", "x..x..x.........")]),
    ("When the Levee Breaks", "Led Zeppelin", 72, [("Hi-hat", "x.x.x.x.x.x.x.x."), ("Snare", "....x.......x..."), ("Kick", "xx.....x..xx....")]),
]
LISTEN = [  # the owner's three "Drumbeat Listening Examples" (answer key); the teacher programs them in Logic
    [("Hi-hat", "x.x.x.x.x.x.x.x."), ("Snare", "....x.......x..."), ("Kick", "x.......x.......")],
    [("Hi-hat", "x.x.x.x.x.x.xxx."), ("Snare", "....x.......x..."), ("Kick", "x.....x.x.....x."), ("Clap", "............x...")],
    [("Hi-hat", "xxxxxxxxxxxxxxxx"), ("Snare", "....x.......x..."), ("Kick", "x..x.x....x...xx"), ("Clap", "....x..x....x..x"),
     ("Shaker", "x.x.xxx.x.x.x.x.")],
]
LISTEN_ROWS = ["Hi-hat", "Snare", "Kick", "Clap", "Shaker"]
START = [("Hi-hat", "x.x.x.x.x.x.x.x."), ("Snare", "....x.......x..."), ("Kick", "x.......x.......")]   # Day 2 starting point
EX_A = [("Closed hi-hat", "x.x.x.x.x.x.x..."), ("Open hi-hat", "..............x."), ("Snare", "....x.......x..."),
        ("Kick", "x.....x.x......."), ("Clap", "............x...")]
EX_B = [("Closed hi-hat", "xxxxxxxxxxxxxx.."), ("Open hi-hat", "..............x."), ("Snare", "....x.......x..."),
        ("Kick", "x.....x.x......."), ("Clap", "............x...")]
MIX = [("Kick", "0 dB (loudest)", "Center", "The anchor: everything else sits around it."),
       ("Snare", "about −3 dB", "Center", "Cuts through on 2 and 4."),
       ("Clap", "about −6 dB", "Slightly left (L10)", "Doubles the snare on beat 4; tucked under it."),
       ("Closed hi-hat", "about −9 dB", "Slightly right (R15)", "Sits above the kick; never louder than the snare."),
       ("Open hi-hat", "about −11 dB", "Slightly left (L15)", "A quick lift into the next bar.")]
AABA = [("A", 1), ("A", 5), ("B", 9), ("A", 13)]   # section, first measure (4 measures each)

# ---------------------------------------------------------------- the owner's words (lightly edited: school, LMS and day numbers)
GOALS = ["Identify the kick, snare, and hi-hat by name, sound, and function",
         "Build an original drum beat using Logic Pro’s Drum Machine Designer",
         "Understand how volume and panning affect the feel of a mix",
         "Experience the creative process of building something musical from scratch"]
VOCAB = [("Kick drum", "the lowest, heaviest drum hit, typically placed on beats 1 and 3"),
         ("Snare drum", "the sharp, cracking hit that drives the groove, typically on beats 2 and 4"),
         ("Hi-hat", "a pair of cymbals creating a repeating pulse; can be closed (tight) or open (ringing)"),
         ("Loop", "a repeated musical idea: the building block of almost every song you hear"),
         ("Beat", "one steady pulse: what you clap or nod your head to"),
         ("Measure (bar)", "a group of 4 beats: how music gets counted and organized"),
         ("Step sequencer", "a grid where you place drum hits across a measure; each square is a subdivision"),
         ("Subdivision", "how a beat is divided: eighth notes split each beat in 2, sixteenth notes split it in 4"),
         ("Tempo", "the speed of a song, measured in beats per minute (BPM)"),
         ("Pattern region", "a looping block of music you build in the step sequencer"),
         ("Cycle", "Logic’s loop playback: repeats a measure or section over and over while you work on it"),
         ("AABA form", "a structure where A is your main pattern, B is a variation, then A returns"),
         ("Volume", "how loud a track plays relative to the others in the mix"),
         ("Panning", "the left-right position of a sound in the stereo field")]
REQUIREMENTS = ["Built entirely in Drum Machine Designer: no Apple Loops",
                "Must include: kick drum, snare drum, closed hi-hat, and at least one additional element (open hi-hat, clap, rim, shaker, etc.)",
                "Minimum 16 measures, structured in AABA form (A = main beat, B = variation, A = main beat returns)",
                "Volume adjusted thoughtfully on all tracks: nothing at default",
                "Panning applied intentionally to at least 2 tracks",
                "Project named correctly: Your Name – Beat From Nothing",
                "Exported as MP3 and .zip (compressed .logicx file) and submitted to your class page before presenting"]
STEPS = [  # (day, heading, text)
    (1, "Watch & Label", "Watch a short parts-of-the-drum-set video. Label the Parts of the Drum Set worksheet using the word bank. "
                         "Do not move on until you can name each piece of the kit without looking."),
    (1, "Loops & Listening", "Learn what a loop is using Billie Jean, then follow along with the Drumbeat Listening Examples as a class. "
                             "Notice where the kick lands, where the snare hits, what the hi-hat is doing. Count. Listen. Discuss."),
    (2, "Copy a Famous Pattern", "Open the Famous Drum Patterns page. Choose ONE pattern and recreate it exactly in Drum Machine Designer; "
                                 "tempo must match. Call over your teacher when you’re done."),
    (3, "Build Your A Section", "Delete the famous pattern. Now build your own. Start with kick on 1 and 3, snare on 2 and 4, hi-hat in eighth notes. "
                                "Then make it yours: adjust, add, subtract until it feels right. This is your A section."),
    (3, "Build Your B Section", "Create a variation on your A section. You don’t need to start over: try removing one element, adding one new sound, "
                                "or changing the hi-hat rhythm. It should feel related to A but noticeably different."),
    (3, "Structure in AABA", "Arrange your pattern regions in AABA form: A section, A section, B section, A section. "
                             "Each section should be 4 measures long for a 16-measure total."),
    (3, "Mix", "Adjust the volume of each track so every element has its own space in the mix. Pan at least two elements slightly left or right. "
               "Ask yourself: can I clearly hear each drum sound?"),
    (4, "Present", "Post three feedback questions to the class feedback board before class begins. Share your beat with the class. "
                   "Respond to at least two of your classmates’ projects with genuine, specific feedback."),
]
TIPS = ["Start simple. A groove that locks in is more powerful than a complicated pattern that doesn’t feel right.",
        "Listen while you build. If you wouldn’t nod your head to it, it isn’t there yet.",
        "Your B section doesn’t need to be completely different: one small change is often enough to give the listener something new.",
        "Turn things down before you turn things up. Your mix will almost always sound better when every element has space.",
        "The hi-hat is what makes a beat breathe. Try eighth notes vs. sixteenth notes and hear how the energy changes.",
        "When you copy the famous pattern, count carefully. Accuracy here trains your ear to recognize what ‘correct’ feels like."]
CHECKLIST = ["Beat built entirely in Drum Machine Designer: no Apple Loops",
             "Kick drum, snare drum, and closed hi-hat are all present",
             "At least one additional drum element included",
             "AABA structure with a clear, noticeable B section",
             "Volume adjusted on all drum tracks: nothing at default",
             "Panning applied to at least 2 tracks",
             "Project named: Your Name – Beat From Nothing",
             "Exported as MP3 (and compressed .logicx) and submitted to your class page",
             "Three feedback questions posted to the class feedback board before presenting"]
RUBRIC = [  # owner's rubric, levels 2-5 per category, 5 categories = /25
    ("Creativity & Originality",
     ["Beat is basic with little creative decision-making. Sounds like the default example.",
      "Beat is functional but predictable. B section exists but contrast is minimal.",
      "Beat shows clear creative choices. B section is noticeably different from A.",
      "Beat has a distinct feel and personality. B section creates meaningful contrast. Goes beyond the obvious."]),
    ("Technical Skill / DAW Use",
     ["Basic use of the step sequencer with significant errors or workflow confusion.",
      "Step sequencer used but some elements seem randomly placed. Some workflow confusion.",
      "Step sequencer used correctly. All required elements present and placed with intention.",
      "Drum Machine Designer used confidently. Pattern regions, 64-step, or note repeat used with control."]),
    ("Composition & Structure",
     ["No clear structure. Beat does not repeat in a musical way.",
      "Structure attempted but A and B sections are not clearly differentiated.",
      "AABA structure is present and mostly clear. Beat feels mostly locked in rhythmically.",
      "AABA structure is clear and musically satisfying. Beat locks in and feels complete."]),
    ("Volume & Balance",
     ["No meaningful volume or panning adjustments. Everything at default levels.",
      "Some volume adjustments made but mix feels unbalanced. Little or no panning applied.",
      "Most elements are balanced. Volume and panning adjustments are evident and mostly appropriate.",
      "All elements are balanced with intention. Each drum sound has its own space. Panning adds width without distraction."]),
    ("Documentation & Presentation",
     ["Project not submitted or not presented.",
      "Project submitted late or named incorrectly. Presentation made with minimal engagement.",
      "Project named correctly and submitted. Presentation made. Feedback questions posted.",
      "Project correctly named, submitted on time. Confident presentation. Three thoughtful feedback questions posted."]),
]
LEVELS = [2, 3, 4, 5]
KIT = [  # (number, word-bank name) in the order of the callouts on the drawing
    (1, "Crash Cymbal"), (2, "Hi-Hat"), (3, "High Tom"), (4, "Middle Tom"), (5, "Ride Cymbal"),
    (6, "Snare Drum"), (7, "Bass/Kick Drum"), (8, "Floor/Low Tom")]
WORD_BANK = ["Hi-Hat", "Snare Drum", "Bass/Kick Drum", "High Tom", "Middle Tom", "Floor/Low Tom", "Crash Cymbal", "Ride Cymbal"]

DAYS = [  # 4 days, 2/1/1 (owner's DMP1 Restructure); minute blocks for ~53-minute periods, from the owner's Teacher Notes
    {"n": 1, "kind": "Learning", "title": "Name the kit, hear the loop", "slides": "Intro deck slides 1–8",
     "goal": "Students can name every part of the kit, define a loop, and count beats and measures by ear. Nothing is built in Logic yet.",
     "blocks": [("0–5", "Hook", "Play about 60 seconds of Billie Jean (Michael Jackson) from YouTube or Spotify and ask what students hear. "
                                "What’s repeating? What enters? Don’t give them answers yet."),
                ("5–17", "Parts of the drum set", "Show a short parts-of-the-drum-set video, then students label the Parts of the Drum Set "
                                                  "worksheet from the word bank. Review: every student names kick (bass drum), snare, hi-hat, toms, crash and ride without looking."),
                ("17–35", "Intro deck: loops", "Define loop: a repeated musical idea that builds on what came before. Replay Billie Jean and name each "
                                               "looping instrument. Clap the pulse: each clap is a beat, 4 beats = 1 measure. Count how many measures the drums "
                                               "play alone, then map the intro layer by layer and show it as regions in Logic."),
                ("35–50", "Drumbeat listening", "Play the three listening examples (programmed from the key on page KEYPAGE). Students mark where each "
                                                "sound lands on the grid. Fill it in together, then compare: why does one feel different from another?"),
                ("50–53", "Exit check", "Words you should own: kick, snare, hi-hat, toms, crash, ride, loop, beat, measure, bass.")]},
    {"n": 2, "kind": "Learning + first time in Logic", "title": "Copy a famous pattern", "slides": "Intro deck slides 9–15",
     "goal": "Every student can create a pattern region, cycle a measure, add and duplicate tracks, change an instrument and set the tempo, "
             "and has one famous pattern checked off. The project launches at the end of class.",
     "blocks": [("0–5", "Vocabulary review", "Cold-call the parts of the kit. What is a loop? How many beats are in a measure? "
                                             "What made the three listening examples feel different?"),
                ("5–15", "First pattern region", "New project, Drum Machine Designer, then Control-click and Create Pattern Region. "
                                                 "Students mess around: no requirements yet. Hands on the grid before anything is graded."),
                ("15–23", "Logic techniques", "Demonstrate to the whole class: Cycle, add a track, duplicate a track, change the instrument on a track."),
                ("23–38", "Copy a famous pattern", "Students delete what they made, choose ONE pattern from the Famous Drum Patterns page and recreate it "
                                                   "exactly. Every student shows you theirs; use that check-in to show them how to change the project tempo."),
                ("38–53", "Launch the project", "Hand out the project sheet and rubric and read the requirements together. Explain the grid: 16 squares = "
                                                "16 sixteenth notes; every other square = eighth notes. Play the starting point (kick 1 and 3, snare 2 and 4, "
                                                "hi-hat in eighths), change one thing, and let them hear how far the feel moves.")]},
    {"n": 3, "kind": "Work day: build and mix", "title": "A, B, AABA, mix, submit", "slides": "Project sheet + checklist",
     "goal": "Project submitted by the end of class: 16 measures in AABA, every track’s volume set, two tracks panned.",
     "blocks": [("0–18", "A section", "Students build a 4-measure A section that is theirs, not the famous pattern. Circulate. The prompt when a student "
                                      "is stuck: ‘Can you nod your head to it? If not, what one thing would you change?’"),
                ("18–30", "B section", "Demonstrate: duplicate the A pattern region and change a single element. Redirect students who start from "
                                       "scratch. ‘Play me your A and then your B. Can I hear the difference?’"),
                ("30–37", "Arrange AABA", "A, A, B, A, four measures each, sixteen measures total. Check region lengths before anyone moves on."),
                ("37–48", "Mixing focus", "Play a mix with everything at max, then pull elements down one at a time. Students set every track’s volume "
                                          "and pan at least two. Remind: turn things down, not up."),
                ("48–53", "Export and submit", "Bounce an MP3, name it correctly, compress the .logicx, and submit both.")]},
    {"n": 4, "kind": "Presenting", "title": "Share and give feedback", "slides": "Rubric",
     "goal": "Every student plays their beat, hears feedback, and gives specific feedback to two classmates.",
     "blocks": [("0–5", "Feedback questions", "Students post three feedback questions to the class board (Padlet or similar) before class begins."),
                ("5–50", "Presentations", "Each student plays their beat once or twice and shares their three questions. Classmates respond to at least "
                                          "2 projects. Allow about 3–4 minutes per student."),
                ("50–53", "Close", "‘Every single one of those beats was made from scratch. Nothing you heard was a pre-made loop.’ Let that land.")]},
]

STRUGGLES = [
    ("The blank page problem", "Some students will stare at the step sequencer and not know where to start. The Day 2 sequence solves this: let them "
     "mess around in a pattern region with no requirements first, then copy a famous pattern, then delete it and build something new. Having done it "
     "twice removes the paralysis. Do not skip the copying step even if students are impatient to build their own."),
    ("Over-complicating the B section", "Students tend to want the B section to be completely different; they start over and lose their groove. "
     "Redirect them: remove one element, add one new sound, or change the hi-hat rhythm. Less is almost always more. The best B sections are 80% the same as A."),
    ("Volume at maximum", "Students default to pushing every fader to the top. Address this on Day 3 with a direct demonstration: play a mix with "
     "everything at max, then pull individual elements down and let students hear the difference. The cardinal rule (turn other things down rather "
     "than turning one thing up) should be stated explicitly and repeated."),
    ("Random hi-hat placement", "Students who do not yet understand subdivision will scatter hi-hats without a pattern. Spend 5 minutes at the Day 2 "
     "launch explaining the grid: 16 squares = 16 sixteenth notes. Every other square = eighth notes (the default groove feel). Every square = "
     "sixteenth notes (busier, more energy)."),
    ("Opening the piano roll instead of creating a pattern region", "This is the most common technical error on Day 2. Watch for students who "
     "double-click on a blank space and open an empty piano roll. The correct workflow is right-click (or Control-click), then Create Pattern Region."),
]

# ---------------------------------------------------------------- intro deck content (owner's "Looping in Logic" pattern)
SLIDES = [  # (kind, title, body, notes)
    ("title", "Beat From Nothing", "Feat. Billie Jean",
     "Day 1, before anything else. Welcome students to their first project: by the end of the week every one of them will have built an original drum "
     "beat from scratch in Logic. No loops, no pre-made grooves. Tell them today is about listening and naming; nobody touches Logic until tomorrow. "
     "Have Billie Jean by Michael Jackson cued up on YouTube or Spotify (your own account; no audio is included in this deck)."),
    ("listen", "Listen first", "Billie Jean · Michael Jackson|Teacher: play about 60 seconds from YouTube or Spotify.|What do you hear?|What’s repeating?|What enters, and when?",
     "Play roughly the first minute. Ask the three questions and collect answers on the board, but do not confirm anything yet: this is "
     "‘feel it before you name it’. Students usually say ‘the drums start alone’, ‘the bass comes in’, ‘it keeps going round’. Keep their words; "
     "you will name them on the next slides."),
    ("bullets", "What the heck is a loop?", "A repeated musical idea|Simple or complex!|Builds on previous material",
     "Plain-English definition: a loop is a musical idea that repeats. It can be one drum hit or a whole groove. Each new loop builds on what came "
     "before. Point back to their answers: ‘it keeps going round’ is exactly this. Ask: where else have you heard loops? (Almost every song they know.)"),
    ("columns", "Which instruments loop?", "Drums: hi-hat, snare, kick|Percussion: shaker|Bass: the lowest sound|Synth: chords and color",
     "Replay the opening of Billie Jean and have students raise a hand each time they hear a new looping layer. Reveal the list: drums (hi-hat, snare, "
     "kick), shaker, bass, synth. Point out that the drum part alone is just three sounds: kick, snare and hi-hat. That is what they will build this week."),
    ("bullets", "How long is each loop?", "Clap along to the pulse: each clap is a BEAT|Most of the time there are 4 beats in a MEASURE (bar)|"
     "How many beats do the drums play by themselves?|How many measures is that?",
     "Play Billie Jean again and have the class clap the pulse. Each clap is a beat; four beats make a measure. Then ask how long the drums play alone "
     "before the bass enters, and have students count it in measures while it plays. Let them count it themselves rather than giving the number; "
     "replay as many times as needed. Counting in measures is the skill they need for AABA later in the week."),
    ("intro", "Mapping an intro, layer by layer", "",
     "This is the worked example of how a song intro is built from loops. Walk through it like a timeline: the drums start alone, then each new layer "
     "enters at the start of a 4-measure block and keeps looping. In Logic, each block would be a region on its own track. Ask: what would it sound like "
     "if everything entered at once? (No build, nothing to look forward to.) This chart is a model, not a transcription: it shows the idea of "
     "staggered entrances."),
    ("kit", "Meet the kit", "Kick: the lowest, heaviest hit; usually beats 1 and 3|Snare: the sharp crack; usually beats 2 and 4|"
     "Hi-hat: two cymbals that keep a steady pulse; closed (tight) or open (ringing)",
     "Pairs with the Parts of the Drum Set worksheet. Focus on the three building blocks and their jobs: the kick grounds the beat, the snare drives "
     "it, the hi-hat makes it breathe. Mention the toms, crash and ride exist too; students can use them as their extra element."),
    ("grid", "Read the grid", "16 squares = 1 measure of sixteenth notes|Every other square = eighth notes (the default groove)|Every square = sixteenth notes (busier, more energy)",
     "Day 1 close or Day 2 launch. This is the same grid students used on the listening sheet and the one they will see in Logic’s Step Sequencer. "
     "The shaded squares are the four beats. Have students clap the hi-hat row (eighths), then sixteenths, and feel the energy change."),
    ("pattern", "Worked example: Billie Jean’s drums", "",
     "Day 2. This is the pattern from the Famous Drum Patterns page: hi-hat in eighth notes, snare on 2 and 4, kick on 1 and 3, at 116 BPM. "
     "Copying it is not busywork: it trains the ear to recognize what ‘correct’ feels like with a specific kick-snare-hi-hat relationship. "
     "Build it live in Logic while students watch, then play the song again (YouTube or Spotify) and compare."),
    ("steps", "In Logic", "Create a Software Instrument track and choose any Electronic Drum Kit (this loads Drum Machine Designer)|"
     "Control-click (right-click) in the track at bar 1, then choose Create Pattern Region|Double-click the region to open the Step Sequencer|"
     "Click squares to place hits; press C to Cycle the measure|Set the tempo in the control bar (116 for Billie Jean)",
     "Demonstrate each step on the projector while students follow. Watch for the most common Day 2 error: double-clicking a blank space opens an "
     "empty piano roll. The correct move is Control-click, Create Pattern Region. After the demo, students mess around with no requirements, then copy "
     "ONE famous pattern exactly and call you over. Use that check-in to show them how to change the tempo."),
    ("bullets", "Your project: Beat From Nothing", "An original drum beat, built from nothing|No loops, no pre-made grooves|"
     "Kick, snare, closed hi-hat + at least one extra sound|16 measures in AABA form, mixed with volume and panning",
     "Launch the project at the end of Day 2. Hand out the project sheet and rubric. Hold the line on ‘no loops’: some students will be frustrated "
     "they can’t just drag a loop. By the end of the unit they will feel more pride in four bars they built than in four minutes they assembled."),
    ("aaba", "The steps", "",
     "Walk through the plan: Day 3 they build an A section (start from kick 1 and 3, snare 2 and 4, hi-hat in eighths, then make it theirs), "
     "a B section (one change: remove one element, add one sound, or change the hi-hat rhythm), arrange A A B A at 4 measures each, and mix. "
     "Day 4 they present. The best B sections are about 80% the same as A."),
    ("check", "What you turn in", "Beat built in Drum Machine Designer, no Apple Loops|Volume set on every track, nothing at default|"
     "Panning on at least 2 tracks|Named: Your Name – Beat From Nothing|MP3 + compressed .logicx, submitted before presenting|"
     "Three feedback questions posted before presenting",
     "These are the deliverables from the submission checklist. Show where students submit (your class page) and where the feedback questions go "
     "(Padlet or any class feedback board). Export: File > Bounce > Project or Section, MP3. Compress: in Finder, Control-click the .logicx file > Compress."),
    ("rubric", "How it’s scored", "",
     "Five categories, each scored 2 to 5, for a total out of 25. Read the 5 column aloud for each category so students know what great looks like: "
     "distinct feel, confident use of Drum Machine Designer, a clear AABA that locks in, every sound with its own space, and a confident presentation "
     "with three thoughtful questions."),
    ("closer", "One last thing…", "Can you nod your head to it?|If not, what one thing would you change?",
     "Close the launch with the one question you will ask all week while circulating. It works for every student at every level, and it is "
     "the heart of the rubric’s Composition & Structure row. Then: turn things down, not up."),
]

# ---------------------------------------------------------------- html scaffolding (Hudson Beat, from 006/007)
def font_css(base=BUILD):
    f = os.path.relpath(FONTS, base)
    faces = [("Oswald", 300, "normal", "oswald-latin-300-normal"), ("Oswald", 500, "normal", "oswald-latin-500-normal"),
             ("Oswald", 700, "normal", "oswald-latin-700-normal"),
             ("Source Sans 3", 400, "normal", "source-sans-3-latin-400-normal"),
             ("Source Sans 3", 400, "italic", "source-sans-3-latin-400-italic"),
             ("Source Sans 3", 600, "normal", "source-sans-3-latin-600-normal"),
             ("Source Sans 3", 700, "normal", "source-sans-3-latin-700-normal")]
    return "".join(f"@font-face{{font-family:'{n}';font-weight:{w};font-style:{s};src:url('{f}/{file}.woff2') format('woff2');}}"
                   for n, w, s, file in faces)


CSS = """
@page { size: 8.5in 11in; margin: 0; }
* { box-sizing: border-box; }
html, body { margin: 0; padding: 0; background: #fff; color: #000; }
body { font-family: 'Source Sans 3', sans-serif; -webkit-print-color-adjust: exact; print-color-adjust: exact; }
.page { width: 8.5in; height: 11in; position: relative; overflow: hidden; page-break-after: always; break-after: page;
        padding: 0.42in 0.5in 0.38in; display: flex; flex-direction: column; }
.page:last-child { page-break-after: auto; break-after: auto; }
.l { font-family: 'Oswald'; font-weight: 300; }
.top { display: flex; justify-content: space-between; align-items: flex-end; border-bottom: 2.2pt solid #000; padding-bottom: 6pt; }
.top .t { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; letter-spacing: .01em; line-height: 1; font-size: 25pt; }
.top .t small { display: block; font-family: 'Oswald'; font-weight: 300; text-transform: none; font-size: 12pt; letter-spacing: 0; margin-top: 4pt; }
.tag { text-align: right; }
.tag .lv { display: inline-block; background: #000; color: #fff; font-family: 'Oswald'; font-weight: 700; font-size: 14pt; padding: 3pt 9pt 4pt; border-radius: 4pt; letter-spacing: .04em; }
.tag .lv.ak { background: ACCENT; }
.tag .ln { font-family: 'Oswald'; font-weight: 300; font-size: 11pt; margin-top: 3pt; }
.name { display: flex; gap: 18pt; font-size: 12pt; margin: 9pt 0 4pt; }
.name span { flex: 1; border-bottom: 1pt solid #000; height: 15pt; }
.name span.d { flex: 0 0 1.6in; }
.main { flex: 1; display: flex; flex-direction: column; min-height: 0; padding-top: 6pt; }
.foot { display: flex; justify-content: space-between; font-size: 8.5pt; color: #333; border-top: .8pt solid #000; padding-top: 3pt; }
.wm { position: absolute; left: 50%; top: 50%; transform: translate(-50%, -50%) rotate(-38deg); font-family: 'Oswald'; font-weight: 700;
      font-size: 150pt; color: rgba(0,0,0,.13); letter-spacing: .06em; white-space: nowrap; pointer-events: none; z-index: 9; }
h2.s { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 13pt; margin: 9pt 0 3pt; letter-spacing: .02em; }
h2.s .acc, .acc { color: ACCENT; }
p, li { font-size: 10.5pt; line-height: 1.32; margin: 0 0 4pt; }
ul, ol { margin: 0 0 4pt; padding-left: 15pt; }
.tn p, .tn li { font-size: 11pt; line-height: 1.42; margin-bottom: 6pt; }
.lines { flex: 1; min-height: 44pt; background: repeating-linear-gradient(to bottom, transparent 0, transparent 21pt, #000 21pt, #000 22pt); }
table.t { border-collapse: collapse; width: 100%; font-size: 9.8pt; margin: 2pt 0 4pt; }
table.t td, table.t th { border: 1pt solid #000; padding: 3pt 5pt; text-align: left; vertical-align: top; line-height: 1.25; }
table.t th { font-family: 'Oswald'; font-weight: 500; text-transform: uppercase; font-size: 9.5pt; background: #000; color: #fff; letter-spacing: .04em; vertical-align: middle; }
.cols { display: flex; gap: 20pt; }
.cols > div { flex: 1; min-width: 0; }
.small { font-size: 9pt !important; color: #333; }
.box { border: 1.4pt solid #000; border-radius: 8pt; padding: 7pt 10pt; }
.box.acc-b { border-color: ACCENT; border-width: 2pt; }
.k { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 10pt; letter-spacing: .05em; color: ACCENT; display: block; margin-bottom: 2pt; }
.ln2 { border-bottom: 1pt solid #000; height: 22pt; }
.fl { display: flex; align-items: flex-end; gap: 6pt; font-size: 10.5pt; margin: 0 0 8pt; }
.fl b { font-weight: 600; white-space: nowrap; }
.fl span { flex: 1; border-bottom: 1pt solid #000; height: 14pt; }
.num { list-style: none; padding: 0; margin: 2pt 0 0; }
.num li { display: flex; gap: 10pt; margin: 0 0 9pt; }
.num .n { flex: 0 0 24pt; height: 24pt; border-radius: 12pt; background: #000; color: #fff; font-family: 'Oswald'; font-weight: 700; font-size: 13pt;
          display: flex; align-items: center; justify-content: center; }
.num .n.d { flex: 0 0 auto; height: auto; border-radius: 4pt; padding: 2pt 6pt 3pt; font-size: 9.5pt; letter-spacing: .04em; align-self: flex-start; }
.num .hd { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 12pt; letter-spacing: .02em; display: block; }
.why { font-style: italic; }
.blk { display: flex; gap: 9pt; margin: 0 0 7pt; font-size: 10.8pt; line-height: 1.34; }
.blk .tm { flex: 0 0 40pt; text-align: center; background: #000; color: #fff; font-family: 'Oswald'; font-weight: 700; font-size: 10pt;
           border-radius: 4pt; padding: 2pt 0 3pt; align-self: flex-start; }
.blk b.h { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 10.5pt; letter-spacing: .02em; }
.day { border: 1.4pt solid #000; border-radius: 8pt; padding: 7pt 10pt 3pt; margin-bottom: 8pt; }
.day .dh { display: flex; justify-content: space-between; align-items: baseline; border-bottom: 1pt solid #000; padding-bottom: 3pt; margin-bottom: 6pt; }
.day .dh b { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 14pt; }
.day .dh span { font-family: 'Oswald'; font-weight: 300; font-size: 11pt; }
.goal { font-size: 10.6pt; margin: 0 0 7pt; }
.chk { display: flex; gap: 8pt; align-items: flex-start; font-size: 11pt; margin: 0 0 8pt; line-height: 1.3; }
table.rub { border-collapse: collapse; width: 100%; }
table.rub td, table.rub th { border: 1.1pt solid #000; padding: 5pt 5pt; vertical-align: top; font-size: 9.3pt; line-height: 1.27; }
table.rub th { background: #000; color: #fff; font-family: 'Oswald'; font-weight: 500; font-size: 11pt; text-align: center; vertical-align: middle; }
table.rub td.c { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 10.5pt; width: 17%; vertical-align: middle; }
table.rub td.sc { width: 9%; }
.cover { align-items: center; justify-content: space-between; padding: 0.55in 0.6in 0.5in;
         background: linear-gradient(180deg, #fff 0%, #fff 55%, #ededed 100%); }
.store { font-family: 'Oswald'; font-weight: 500; font-size: 13pt; letter-spacing: .32em; text-transform: uppercase; }
.ctitle { text-align: center; }
.ctitle .a { font-family: 'Oswald'; font-weight: 700; font-size: 36pt; line-height: 1; letter-spacing: .02em; text-transform: uppercase; }
.ctitle .b { white-space: nowrap; font-family: 'Oswald'; font-weight: 700; font-size: 60pt; line-height: .98; letter-spacing: .01em; text-transform: uppercase; }
.ctitle .c { font-family: 'Oswald'; font-weight: 300; font-size: 16pt; margin-top: 8pt; }
.cband { width: 100%; background: #000; color: #fff; text-align: center; font-family: 'Oswald'; font-weight: 500; font-size: 14pt;
         letter-spacing: .08em; text-transform: uppercase; padding: 8pt 0 9pt; border-radius: 4pt; }
.cband em { font-style: normal; color: ACCENT; }
""".replace("ACCENT", ACCENT)

COPY = f"© {YEAR} {STORE} · {BRAND} · For single-classroom use"


def html_doc(pages, title="doc"):
    return (f'<!doctype html><html><head><meta charset="utf-8"><title>{escape(title)}</title>'
            f'<style>{font_css()}{CSS}</style></head><body>{"".join(pages)}</body></html>')


def foot(right):
    return f'<div class="foot"><span>{COPY}</span><span>{right}</span></div>'


def wm(on):
    return '<div class="wm">PREVIEW</div>' if on else ""


def page(sub, tag, tagsub, inner, n, watermark=False, name=False, ak=False, cls=""):
    top = (f'<div class="top"><div class="t">{escape(TITLE)}<small>{sub}</small></div>'
           f'<div class="tag"><div class="lv{" ak" if ak else ""}">{tag}</div><div class="ln">{tagsub}</div></div></div>')
    nm = '<div class="name">Name<span></span>Date<span class="d"></span></div>' if name else ""
    return (f'<section class="page {cls}">{wm(watermark)}{top}{nm}<div class="main">{inner}</div>{foot(f"Page {n}")}</section>')


def checkbox(size=12):
    return (f'<svg width="{size}pt" height="{size}pt" viewBox="0 0 {size} {size}" style="flex:0 0 auto;margin-top:1pt">'
            f'<rect x=".7" y=".7" width="{size - 1.4}" height="{size - 1.4}" rx="2" fill="#fff" stroke="#000" stroke-width="1.2"/></svg>')


def svg(w, h, body, style="display:block"):
    return f'<svg xmlns="http://www.w3.org/2000/svg" width="{w}pt" height="{h}pt" viewBox="0 0 {w} {h}" style="{style}">{body}</svg>'


def t_(x, y, s, size, weight=700, family="Oswald", anchor="middle", fill="#000"):
    fam = f"'{family}'" if " " in family else family
    return f'<text x="{x:.1f}" y="{y:.1f}" font-family="{fam}" font-weight="{weight}" font-size="{size}" text-anchor="{anchor}" fill="{fill}">{escape(s)}</text>'


# ---------------------------------------------------------------- drawings
def grid_svg(rows, width=511, cell_h=17, label_w=70, blank=False, fill="#000", head=True, label_rows=None):
    """A 16-step drum grid. rows = [(name, 'x...')]; blank=True draws empty cells (worksheet). Beat columns are shaded."""
    names = label_rows or [r[0] for r in rows]
    cw = (width - label_w) / 16
    hh = 13 if head else 0
    out = []
    if head:
        for i, lab in enumerate(STEP_LABELS):
            out.append(t_(label_w + (i + .5) * cw, 9.5, lab, 7.4 if "." in lab else 9, 500 if "." in lab else 700))
    for r, name in enumerate(names):
        y = hh + r * cell_h
        out.append(t_(label_w - 6, y + cell_h * .68, name, 9.2, 500, anchor="end"))
        steps = dict(rows).get(name, "." * 16) if not blank else "." * 16
        for i in range(16):
            x = label_w + i * cw
            bg = "#e6e6e6" if i % 4 == 0 else "#fff"
            out.append(f'<rect x="{x:.2f}" y="{y:.2f}" width="{cw:.2f}" height="{cell_h}" fill="{bg}" stroke="#000" stroke-width="{1.4 if i % 4 == 0 else .6}"/>')
            if steps[i] == "x":
                out.append(f'<rect class="hit" x="{x + 2.6:.2f}" y="{y + 2.6:.2f}" width="{cw - 5.2:.2f}" height="{cell_h - 5.2:.2f}" rx="2" fill="{fill}"/>')
    h = hh + len(names) * cell_h + 1
    out.append(f'<rect x="{label_w:.2f}" y="{hh}" width="{width - label_w - .5:.2f}" height="{len(names) * cell_h}" fill="none" stroke="#000" stroke-width="1.6"/>')
    return svg(width, h, "".join(out))


def cyl(cx, top, w, h, fill="#fff", band=None):
    """A drum shell seen slightly from above: ellipse top, straight sides, curved bottom."""
    ry = w * 0.16
    body = (f'<path d="M{cx - w / 2:.1f} {top:.1f} L{cx - w / 2:.1f} {top + h:.1f} A{w / 2:.1f} {ry:.1f} 0 0 0 {cx + w / 2:.1f} {top + h:.1f} '
            f'L{cx + w / 2:.1f} {top:.1f}" fill="{fill}" stroke="#000" stroke-width="2.4" stroke-linejoin="round"/>')
    if band:
        body += (f'<path d="M{cx - w / 2:.1f} {top + h * .5:.1f} A{w / 2:.1f} {ry:.1f} 0 0 0 {cx + w / 2:.1f} {top + h * .5:.1f}" fill="none" '
                 f'stroke="{band}" stroke-width="3"/>')
    body += f'<ellipse cx="{cx:.1f}" cy="{top:.1f}" rx="{w / 2:.1f}" ry="{ry:.1f}" fill="#fff" stroke="#000" stroke-width="2.4"/>'
    return body


def stand(x, y0, y1, legs=True):
    s = f'<line x1="{x}" y1="{y0}" x2="{x}" y2="{y1}" stroke="#000" stroke-width="2.4" stroke-linecap="round"/>'
    if legs:
        s += (f'<path d="M{x} {y1 - 22} L{x - 22} {y1} M{x} {y1 - 22} L{x + 22} {y1}" stroke="#000" stroke-width="2.4" '
              f'stroke-linecap="round" fill="none"/>')
    return s


def cymbal(cx, cy, w, tilt=0, fill="#fff"):
    return (f'<g transform="rotate({tilt} {cx} {cy})"><ellipse cx="{cx}" cy="{cy}" rx="{w / 2}" ry="{w * .09:.1f}" fill="{fill}" stroke="#000" stroke-width="2.4"/>'
            f'<ellipse cx="{cx}" cy="{cy - w * .03:.1f}" rx="{w * .09:.1f}" ry="{w * .035:.1f}" fill="#000"/></g>')


KIT_POS = {  # callout anchor points on the 640 x 420 drawing
    1: (150, 92), 2: (52, 168), 3: (262, 158), 4: (378, 158), 5: (520, 102), 6: (156, 236), 7: (320, 300), 8: (520, 250)}


def drumkit_svg(mode="plain", accent=ACCENT, width="100%"):
    """Original line drawing of a 5-piece kit. mode: plain | numbers (worksheet) | answers (key)."""
    b = []
    floor = 400
    # stands first (behind)
    b.append(stand(150, 100, floor))          # crash
    b.append(stand(70, 178, floor))           # hi-hat
    b.append(stand(525, 112, floor))          # ride
    b.append(stand(160, 262, floor))          # snare
    # cymbals
    b.append(cymbal(150, 96, 150, -8))
    b.append(cymbal(70, 170, 96, 0))
    b.append(cymbal(70, 182, 96, 0))
    b.append(cymbal(525, 108, 176, 7))
    # kick (front view circle)
    b.append(f'<circle cx="320" cy="300" r="98" fill="#fff" stroke="#000" stroke-width="3"/>')
    b.append(f'<circle cx="320" cy="300" r="82" fill="none" stroke="#000" stroke-width="1.6"/>')
    b.append(f'<circle cx="320" cy="300" r="30" fill="{accent}" stroke="#000" stroke-width="2.4"/>')
    for ang in range(0, 360, 45):   # tension rods
        import math
        x, y = 320 + 90 * math.cos(math.radians(ang)), 300 + 90 * math.sin(math.radians(ang))
        b.append(f'<circle cx="{x:.1f}" cy="{y:.1f}" r="3.2" fill="#000"/>')
    b.append('<path d="M250 372 L236 400 M390 372 L404 400" stroke="#000" stroke-width="2.4" stroke-linecap="round"/>')
    # rack toms on the kick
    b.append('<line x1="300" y1="205" x2="282" y2="190" stroke="#000" stroke-width="2.4"/><line x1="340" y1="205" x2="360" y2="190" stroke="#000" stroke-width="2.4"/>')
    b.append(cyl(262, 150, 86, 46))
    b.append(cyl(378, 150, 98, 52))
    # snare
    b.append(cyl(160, 226, 112, 36, band=accent))
    # floor tom
    b.append(cyl(520, 226, 120, 120))
    b.append('<path d="M470 346 L462 400 M570 346 L578 400" stroke="#000" stroke-width="2.4" stroke-linecap="round"/>')
    b.append(f'<line x1="20" y1="{floor + 1}" x2="620" y2="{floor + 1}" stroke="#000" stroke-width="1.4"/>')
    if mode in ("numbers", "answers"):
        for n, name in KIT:
            x, y = KIT_POS[n]
            r = 13 if mode == "numbers" else 17
            b.append(f'<circle cx="{x}" cy="{y}" r="{r}" fill="{"#000" if mode == "numbers" else accent}" stroke="#fff" stroke-width="2"/>')
            b.append(t_(x, y + r * .43, str(n), r * 1.15, 700, fill="#fff"))
    return f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 640 420" width="{width}" style="display:block">{"".join(b)}</svg>'


def machine_svg(width="100%"):
    """Cover hero: an original step-sequencer drum machine (16 steps x 4 rows) drawn in black + accent."""
    rows = [("HAT", "x.x.x.x.x.x.x.x."), ("CLAP", "............x..."), ("SNARE", "....x.......x..."), ("KICK", "x.....x.x.......")]
    b = ['<rect x="10" y="10" width="780" height="380" rx="26" fill="#fff" stroke="#000" stroke-width="5"/>']
    # top strip: screen + knobs
    b.append('<rect x="44" y="40" width="250" height="66" rx="8" fill="#000"/>')
    b.append(t_(169, 84, "116 BPM", 30, 700, fill="#fff"))
    for i, x in enumerate((360, 440, 520, 600, 680)):
        b.append(f'<circle cx="{x}" cy="73" r="25" fill="#fff" stroke="#000" stroke-width="4"/>')
        a = [-60, 20, -20, 60, 0][i]
        import math
        b.append(f'<line x1="{x}" y1="73" x2="{x + 18 * math.sin(math.radians(a)):.1f}" y2="{73 - 18 * math.cos(math.radians(a)):.1f}" '
                 f'stroke="{ACCENT if i == 1 else "#000"}" stroke-width="5" stroke-linecap="round"/>')
    # grid
    x0, y0, cw, ch = 140, 140, 39.5, 52
    for r, (name, steps) in enumerate(rows):
        y = y0 + r * ch
        b.append(t_(x0 - 14, y + 31, name, 17, 700, anchor="end"))
        for i in range(16):
            x = x0 + i * cw
            on = steps[i] == "x"
            col = ACCENT if (on and name == "KICK") else ("#000" if on else "#fff")
            b.append(f'<rect x="{x + 3:.1f}" y="{y + 6}" width="{cw - 6:.1f}" height="{ch - 12}" rx="6" fill="{col}" stroke="#000" '
                     f'stroke-width="{3.2 if i % 4 == 0 else 2}"/>')
    return f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 800 400" width="{width}" style="display:block">{"".join(b)}</svg>'


# ---------------------------------------------------------------- page plan
PLAN = ["cover", "howto", "overview", "setup", "pacing", "day12", "day34", "tnotes", "tnotes2",
        "post", "kit", "listen", "famous", "proj1", "proj2", "vocab", "checklist", "rubric",
        "key", "example", "terms"]
IX = {k: i + 1 for i, k in enumerate(PLAN)}
STUDENT = ["post", "kit", "listen", "famous", "proj1", "proj2", "vocab", "checklist", "rubric"]
PREVIEW_PICKS = ["pacing", "proj1", "listen", "rubric"]
DECK_SLIDES = len(SLIDES)


def fill(s):
    return s.replace("KEYPAGE", str(IX["key"]))


def cover_page():
    return (f'<section class="page cover"><div class="store">{BRAND}</div>'
            f'<div style="width:6.6in">{machine_svg()}</div>'
            f'<div class="ctitle"><div class="a">Drum Programming in Logic Pro</div><div class="b">Beat From Nothing</div>'
            f'<div class="c">Digital Music Production 1 · Unit 1 · teacher notes, student pages, /25 rubric, intro deck</div></div>'
            f'<div class="cband">Grades 9–12 <em>·</em> 4 days <em>·</em> {len(PLAN)} pages + {DECK_SLIDES}-slide deck</div></section>')


def howto_page(n):
    parts = [("Teacher pages", "howto", "tnotes2", "How to use, overview, setup, pacing, day plans, teacher notes"),
             ("Student pages", "post", "rubric", "Day 1 post, worksheets, patterns, project sheet, vocabulary, checklist, rubric"),
             ("Answer key + example notes", "key", "example", "Kit labels, listening answers, a worked example beat and mix"),
             ("Intro deck (separate file)", None, None, f"{DECK_SLIDES} slides, .pptx + PDF, speaker notes on every slide")]
    trs = "".join(f"<tr><td><b>{a}</b></td><td style='white-space:nowrap'>{('pages ' + str(IX[f]) + '–' + str(IX[l])) if f else 'deck'}</td><td>{u}</td></tr>"
                  for a, f, l, u in parts)
    inner = (f'<div class="cols"><div>'
             f'<h2 class="s">What this is<span class="acc"> ·</span></h2>'
             f'<p>A complete 4-day drum-programming unit for high school music technology. Students build an original drum beat from nothing in '
             f'Logic Pro’s Drum Machine Designer: no loops, no pre-made grooves. They name the kit, hear how loops work, copy one famous pattern, then '
             f'build and mix their own 16-measure beat in AABA form and present it.</p>'
             f'<p>This is the first unit of my Digital Music Production 1 course, the one I teach every September. It works just as well on its own.</p>'
             f'<h2 class="s">Quick start<span class="acc"> ·</span></h2><ol>'
             f'<li>Read the overview and pacing (pages {IX["overview"]}–{IX["pacing"]}) and the teacher notes (pages {IX["tnotes"]}–{IX["tnotes2"]}).</li>'
             f'<li>Check the setup list (page {IX["setup"]}) on one student computer.</li>'
             f'<li>Program the three listening examples into Logic from the key (page {IX["key"]}): about 10 minutes.</li>'
             f'<li>Open the intro deck and cue up Billie Jean (Michael Jackson) on YouTube or Spotify. No audio is included.</li>'
             f'<li>Print or post the student pages (pages {IX["post"]}–{IX["rubric"]}).</li></ol>'
             f'<h2 class="s">The intro deck<span class="acc"> ·</span></h2>'
             f'<p>You present it: slides 1–8 on Day 1 (hook, what a loop is, beats and measures, the kit, the grid) and slides 9–15 to launch the project '
             f'on Day 2 (worked example, Logic steps, the project, deliverables, rubric). Every slide has speaker notes.</p>'
             f'<p class="small">Fonts: the deck uses Oswald and Source Sans 3 (free on Google Fonts, built into Google Slides). If PowerPoint shows a '
             f'different font, install them or use the PDF copy, which shows the intended look.</p>'
             f'</div><div>'
             f'<h2 class="s">What’s inside<span class="acc"> ·</span></h2><table class="t"><tr><th>Part</th><th>Pages</th><th>Use</th></tr>{trs}</table>'
             f'<h2 class="s">Digital or print<span class="acc"> ·</span></h2>'
             f'<p>Every student page prints in black and white. To go digital, post the PDF pages to your class page; students can fill in the '
             f'listening grid and checklist with any PDF annotation tool.</p>'
             f'<h2 class="s">Make it fit your class<span class="acc"> ·</span></h2>'
             f'<ul><li><b>Short periods or a slower group?</b> Split Day 3 into two build days (A section; then B, AABA and mix). That is my original 5-day pacing.</li>'
             f'<li><b>Need it simpler or harder?</b> See the variations on page {IX["tnotes2"]}.</li>'
             f'<li><b>Feedback board:</b> I use Padlet; any discussion board, LMS forum or paper gallery walk works.</li></ul>'
             f'<div class="box" style="margin-top:6pt"><span class="k">No copyrighted media</span>'
             f'<p style="margin:0">Songs are named by title and artist only. Play them from your own YouTube or Spotify access. The drum grids are simple '
             f'counting charts for teaching; no audio, video, lyrics or album art are included.</p></div>'
             f'</div></div>')
    return page("How to use this resource", "START HERE", "Teacher", inner, n)


def overview_page(n):
    inner = (f'<div class="cols"><div>'
             f'<h2 class="s">What students make<span class="acc"> ·</span></h2>'
             f'<p>An original drum beat built entirely from scratch in Drum Machine Designer: no loops, no pre-made grooves. Students copy one famous '
             f'pattern, then build an original one inspired by it: 16 measures in AABA form, mixed with volume and panning.</p>'
             f'<h2 class="s">What students learn<span class="acc"> ·</span></h2>'
             f'<p>Kick, snare, and hi-hat by name, sound, and function. Subdivision: eighth notes, sixteenth notes, and why they feel different. Why '
             f'repetition is a musical tool, not laziness. The satisfaction of making something from nothing.</p>'
             f'<h2 class="s">Why it comes first<span class="acc"> ·</span></h2>'
             f'<p>This is the hook. But instead of ‘drag a loop and feel like a producer,’ the hook is authorship: ‘I made that sound myself and I understand '
             f'why it works.’ Students who build a beat understand what is inside every loop they will meet for the rest of the year. Constraints force musical decisions.</p>'
             f'<h2 class="s">Mixing concept: volume &amp; panning<span class="acc"> ·</span></h2>'
             f'<p>The first mixing lesson is simple: not everything should be the same volume. The kick should feel grounded, the hi-hat should sit above it, '
             f'the snare should cut through. Turning something down to make something else feel louder is the most transferable mixing lesson there is.</p>'
             f'</div><div>'
             f'<h2 class="s">Goals<span class="acc"> ·</span></h2><ul>' + "".join(f"<li>{g}</li>" for g in GOALS) + '</ul>'
             f'<h2 class="s">At a glance<span class="acc"> ·</span></h2><table class="t">'
             f'<tr><td><b>Grades</b></td><td>9–12, beginners (no music reading needed)</td></tr>'
             f'<tr><td><b>Length</b></td><td>4 class periods of about 53 minutes: 2 learning, 1 work, 1 presenting</td></tr>'
             f'<tr><td><b>Type</b></td><td>Structured project</td></tr>'
             f'<tr><td><b>Software</b></td><td>Logic Pro: Drum Machine Designer, pattern regions, Step Sequencer (16 and 64 steps), volume and pan</td></tr>'
             f'<tr><td><b>Product</b></td><td>16-measure AABA drum beat, MP3 + compressed project, presented with three feedback questions</td></tr>'
             f'<tr><td><b>Assessment</b></td><td>5-category rubric, /25 (page {IX["rubric"]})</td></tr></table>'
             f'<h2 class="s">Standards<span class="acc"> ·</span></h2>'
             f'<p>Supports the National Core Arts Standards for Music (Music Technology, high school): generating and organizing musical ideas '
             f'(MU:Cr1, MU:Cr2), refining work through feedback (MU:Cr3), presenting (MU:Pr6) and analyzing what is heard (MU:Re7).</p>'
             f'<div class="box" style="margin-top:6pt"><span class="k">A note from me</span>'
             f'<p style="margin:0">Some students will be frustrated that they can’t just drag a loop. Hold the line. By the end of this unit they will feel '
             f'more pride in four bars they built than in four minutes they assembled. That pride is the hook, and it lasts longer.</p></div>'
             f'</div></div>')
    return page("Unit overview", "OVERVIEW", "DMP1 · Unit 1", inner, n)


def setup_page(n):
    logic = [("New drum track", "Create a Software Instrument track and choose any kit under Electronic Drum Kit in the Library (Y). This loads Drum Machine Designer."),
             ("Pattern region", "Control-click (right-click) in the track at bar 1 and choose Create Pattern Region. Double-click it to open the Step Sequencer."),
             ("Place hits", "Click squares to turn hits on. Each row is one sound; add rows for more sounds. Pattern length goes up to 64 steps."),
             ("Cycle", "Press C, then set the cycle area over bar 1 so it repeats while you work."),
             ("Tempo", "Double-click the tempo in the control bar and type the BPM."),
             ("Tracks", "Duplicate a track with Track › New Track With Duplicate Setting (Command-D). Change a sound from the Library."),
             ("Repeat", "Option-drag a region to copy it, or select it and press Command-R to repeat it."),
             ("Volume and pan", "Open the Mixer (X). Click the disclosure triangle on the Drum Machine Designer track to show each drum’s own fader and pan knob."),
             ("Export", "File › Bounce › Project or Section (Command-B), MP3. In Finder, Control-click the .logicx project and choose Compress.")]
    rows = "".join(f"<tr><td style='width:22%'><b>{a}</b></td><td>{b}</td></tr>" for a, b in logic)
    inner = (f'<div class="cols"><div>'
             f'<h2 class="s">You need<span class="acc"> ·</span></h2><ul>'
             f'<li><b>A Mac with Logic Pro</b> for each student (Logic Pro X 10.5 or later, for the Step Sequencer and pattern regions).</li>'
             f'<li><b>Headphones</b> for every student.</li>'
             f'<li><b>Projector and speakers</b> for the intro deck and the listening examples.</li>'
             f'<li><b>Billie Jean (Michael Jackson)</b> from your own YouTube or Spotify access. Optional for comparison: the other two songs on the '
             f'Famous Drum Patterns page.</li>'
             f'<li><b>A place to submit</b> an MP3 and a .zip (your LMS or class page).</li>'
             f'<li><b>A feedback board</b> for Day 4 (Padlet, an LMS discussion, or paper).</li></ul>'
             f'<h2 class="s">Optional<span class="acc"> ·</span></h2><ul>'
             f'<li><b>A small USB MIDI keyboard or pad controller</b> (an MPK Mini or similar) to audition sounds and play hits in. The project is built in the '
             f'Step Sequencer, so a keyboard is never required.</li>'
             f'<li>A short “parts of a drum set” video for Day 1 (search YouTube; any clear 3–5 minute video works).</li>'
             f'<li>The free Learning Music beat-making lessons at learningmusic.ableton.com for students who want more.</li></ul>'
             f'<h2 class="s">Before Day 1<span class="acc"> ·</span></h2>'
             + "".join(f'<div class="chk">{checkbox()}<span>{c}</span></div>' for c in [
                 "Open Logic on one student computer and make a pattern region (checks the version and the kits).",
                 f"Program the three listening examples from page {IX['key']} as three pattern regions (or three projects).",
                 "Make a quick example beat of your own to play on Day 2.",
                 "Test the projector sound and cue up Billie Jean.",
                 "Print or post the student pages; set up the submission spot and the feedback board."]) +
             f'</div><div>'
             f'<h2 class="s">Logic quick reference<span class="acc"> ·</span></h2>'
             f'<table class="t" style="font-size:9.4pt">{rows}</table>'
             f'<p class="small">Menu names are from current Logic Pro versions; older versions may differ slightly. Logic Pro is a trademark of Apple Inc.; '
             f'it is named here only to describe what the unit needs. This resource is not affiliated with Apple.</p>'
             f'</div></div>')
    return page("Setup and requirements", "SETUP", "Teacher", inner, n)


def pacing_page(n, watermark=False):
    rows = "".join(
        f"<tr><td style='white-space:nowrap'><b class='acc' style='font-family:Oswald;font-size:13pt'>DAY {d['n']}</b><br><span class='small'>{d['kind']}</span></td>"
        f"<td><b>{d['title']}</b><br><span class='small'>{d['goal']}</span></td>"
        f"<td style='font-size:10pt;line-height:1.45'>{'<br>'.join(f'{a} min · {h}' for a, h, _ in d['blocks'])}</td>"
        f"<td style='font-size:10pt'>{d['slides']}</td></tr>" for d in DAYS)
    inner = (f'<p>Four class periods of about 53 minutes: two learning days, one work day, one presenting day. Each day’s full plan, minute by minute, '
             f'is on pages {IX["day12"]}–{IX["day34"]}.</p>'
             f'<table class="t"><tr><th style="width:13%">Day</th><th style="width:37%">Focus and goal</th><th style="width:34%">Minute blocks</th><th>Show</th></tr>{rows}</table>'
             f'<div class="cols" style="margin-top:8pt"><div class="box"><span class="k">Need a fifth day?</span>'
             f'<p style="margin:0">Split Day 3: build the A section on one day; the B section, AABA and mix on the next. That is the 5-day version I '
             f'have taught. Presentations then move to Day 5.</p></div>'
             f'<div class="box"><span class="k">AABA at a glance</span>' + aaba_svg(230) + '</div></div>'
             f'<h2 class="s">Student pages by day<span class="acc"> ·</span></h2><table class="t">'
             f'<tr><td><b>Day 1</b></td><td>Day 1 post (p. {IX["post"]}), Parts of the Drum Set (p. {IX["kit"]}), Drumbeat Listening (p. {IX["listen"]}), Vocabulary (p. {IX["vocab"]})</td></tr>'
             f'<tr><td><b>Day 2</b></td><td>Famous Drum Patterns (p. {IX["famous"]}), Project Sheet (pp. {IX["proj1"]}–{IX["proj2"]}), Rubric (p. {IX["rubric"]})</td></tr>'
             f'<tr><td><b>Day 3</b></td><td>Project Sheet steps 4–7, Submission Checklist (p. {IX["checklist"]})</td></tr>'
             f'<tr><td><b>Day 4</b></td><td>Rubric for self-score and teacher score (p. {IX["rubric"]})</td></tr></table>')
    return page("Pacing guide · 4 days (2 learning / 1 work / 1 present)", "PACING", "Teacher", inner, n, watermark)


def aaba_svg(w):
    cw = w / 16
    b = []
    for sec, m0 in AABA:
        x = (m0 - 1) * cw
        col = ACCENT if sec == "B" else "#000"
        b.append(f'<rect x="{x + 1:.1f}" y="2" width="{4 * cw - 2:.1f}" height="24" rx="4" fill="{col}"/>')
        b.append(t_(x + 2 * cw, 19.5, sec, 15, 700, fill="#fff"))
        b.append(t_(x + 2, 40, str(m0), 8.5, 500, anchor="start"))
    b.append(t_(w - 2, 40, "16", 8.5, 500, anchor="end"))
    return svg(w, 44, "".join(b)) + '<p class="small" style="margin:2pt 0 0">4 measures each · 16 measures · B = A with one change</p>'


def day_box(d):
    blocks = "".join(f'<div class="blk"><div class="tm">{a}</div><div><b class="h">{h}</b> · {fill(t)}</div></div>' for a, h, t in d["blocks"])
    return (f'<div class="day"><div class="dh"><b>Day {d["n"]} · {d["title"]}</b><span>{d["kind"]} · {d["slides"]}</span></div>'
            f'<p class="goal"><b>Goal:</b> {d["goal"]}</p>{blocks}</div>')


def day12_page(n):
    return page("Day-by-day plans · Days 1 and 2", "DAYS 1–2", "Learning", day_box(DAYS[0]) + day_box(DAYS[1]), n)


def day34_page(n):
    extra = ('<div class="box"><span class="k">Circulating prompts</span><ul style="margin:0">'
             '<li>‘Can you nod your head to it? If not, what one thing would you change?’</li>'
             '<li>‘Play me your A and then your B. Can I hear the difference?’</li>'
             '<li>‘Can you clearly hear each drum sound? What could you turn down?’</li></ul></div>')
    return page("Day-by-day plans · Days 3 and 4", "DAYS 3–4", "Work + present", day_box(DAYS[2]) + day_box(DAYS[3]) + extra, n)


def tnotes_page(n):
    st = "".join(f'<p><b>{a}:</b> {b}</p>' for a, b in STRUGGLES)
    inner = (f'<div class="tn"><h2 class="s" style="margin-top:2pt">What this unit reinforces<span class="acc"> ·</span></h2>'
             f'<p>This is the entry point for the course. It sets the tone by establishing authorship over assembly as the hook: students who build a beat '
             f'from scratch understand what is inside every loop they will encounter for the rest of the year. The first two days are deliberately slow: '
             f'students name the kit, learn what a loop actually is, and count beats and measures by ear before they are asked to build anything. The AABA form '
             f'requirement plants the first formal seed: they meet the term on the project sheet and hear it in their own work before they can define it. The '
             f'mixing concept (volume and panning) is introduced here in the simplest possible context: three sounds competing for space is the most accessible '
             f'mixing problem you can give a beginner.</p>'
             f'<p>The Famous Drum Patterns activity on Day 2 is critical. Copying Billie Jean is not busywork: it trains the ear to recognize what ‘correct’ '
             f'feels like at 116 BPM with a specific kick-snare-hi-hat relationship. Students who complete it properly have a reference point for everything '
             f'they build afterward.</p>'
             f'<h2 class="s">Common student struggles<span class="acc"> ·</span></h2>{st}</div>')
    return page("Teacher notes · what to expect", "TEACHER NOTES", "1 of 2", inner, n)


def tnotes2_page(n):
    inner = (f'<div class="tn"><h2 class="s" style="margin-top:2pt">Suggested variations<span class="acc"> ·</span></h2>'
             f'<p><b>More accessible:</b> Remove the AABA structure requirement. Have students focus on building one strong 4-measure A section they are proud '
             f'of. Introduce AABA in the presentation discussion as a concept to explore next time, not a requirement they failed.</p>'
             f'<p><b>More advanced:</b> Require the B section to be built in 64-step mode (rather than 16). Challenge these students to use note repeat on the '
             f'hi-hat to create a triplet or 32nd-note accent. Ask them to include in their feedback questions: ‘What specific choice makes my B section '
             f'different, and why does it work?’</p>'
             f'<h2 class="s">Grading tips<span class="acc"> ·</span></h2><ul>'
             f'<li>Each category is scored 2–5, for a total out of 25. Students self-score first; you score in the second column (page {IX["rubric"]}).</li>'
             f'<li>Listen once for groove (Composition &amp; Structure), once for contrast between A and B (Creativity), and once with the mixer open (Volume &amp; Balance).</li>'
             f'<li>Technical Skill: open the project. Pattern regions on a clean grid, or 64-step / note repeat used with control, earn the 5.</li></ul>'
             f'<h2 class="s">Resources<span class="acc"> ·</span></h2><ul>'
             f'<li>The intro deck (.pptx and PDF), included: Day 1 loops lesson feat. Billie Jean, and the Day 2 launch.</li>'
             f'<li>A short “how to identify parts of a drum set” video, and a Logic Step Sequencer tutorial (YouTube; choose ones you like).</li>'
             f'<li>learningmusic.ableton.com/make-beats: a useful conceptual supplement for advanced students.</li></ul>'
             f'<h2 class="s">Observation notes<span class="acc"> ·</span></h2>'
             f'<p class="small">What worked, what to adjust, and anything worth remembering for next year.</p></div><div class="lines"></div>')
    return page("Teacher notes · variations and resources", "TEACHER NOTES", "2 of 2", inner, n)


def post_page(n):
    steps = [("Watch", "a short video on the parts of a drum set (your teacher shares the link)", "Learn the names of the main drums before you touch anything."),
             ("Label", f"the Parts of the Drum Set worksheet (page {IX['kit']}) using the word bank at the top", "You should be able to name every piece of the kit without looking by the end."),
             ("Loops", "the “Beat From Nothing” intro slides, feat. Billie Jean", "What is a loop? Which instruments loop in Billie Jean? Clap the pulse: each clap is a beat, and 4 beats make a measure."),
             ("Listen", f"to the Drumbeat Listening Examples (page {IX['listen']}) and follow along as a class", "Count the kicks. Count the snares. What is the hi-hat doing? What makes each one feel different?")]
    lis = "".join(f'<li style="margin-bottom:13pt"><div class="n">{i + 1}</div><div><b class="hd" style="display:inline;font-size:14pt">{a}</b> · {b}<br>'
                  f'<span class="why">{c}</span></div></li>' for i, (a, b, c) in enumerate(steps))
    words = [("Parts of the kit", "kick (bass drum), snare, hi-hat, toms, crash, ride"),
             ("Loop", "a repeated musical idea that builds on what came before"),
             ("Beat", "one clap of the pulse"), ("Measure (bar)", "4 beats"), ("Bass", "the lowest sound in the mix")]
    inner = (f'<h2 class="s" style="font-size:20pt;margin:4pt 0 2pt">Beat From Nothing · Day 1</h2>'
             f'<p class="l" style="font-size:12pt">Digital Music Production 1 · Unit 1</p>'
             f'<p style="font-size:11.5pt;margin:6pt 0 10pt">Follow this sequence in order. Do not skip ahead. Steps 3 and 4 we do together as a class.</p>'
             f'<ol class="num" style="font-size:12.5pt">{lis}</ol>'
             f'<div class="box acc-b" style="margin-top:8pt"><span class="k" style="font-size:12pt">Words you should own by the end of today</span>'
             + "".join(f'<p style="font-size:12.5pt;margin:0 0 5pt"><b>{a}:</b> {b}</p>' for a, b in words) +
             f'<p style="font-size:11.5pt;margin:6pt 0 0"><b>If you can define all of these, you’re ready for Logic next class.</b></p></div>'
             f'<h2 class="s">Helpful links<span class="acc"> ·</span></h2><ul style="font-size:11pt">'
             f'<li>A Logic Step Sequencer tutorial (your teacher shares the link)</li>'
             f'<li>Learning Music beat maker, learningmusic.ableton.com/make-beats: great for understanding beats conceptually</li></ul>'
             f'<div style="flex:1"></div><div class="box"><span class="k">My questions for next class</span><div class="lines" style="height:66pt;flex:none"></div></div>')
    return page("Day 1 classroom post", "STUDENT", "Day 1", inner, n)


def kit_page(n):
    bank = "".join(f'<span style="display:inline-block;border:1.2pt solid #000;border-radius:4pt;padding:2pt 8pt;margin:0 6pt 6pt 0;font-family:Oswald;'
                   f'font-weight:500;font-size:11pt;text-transform:uppercase;letter-spacing:.03em">{w}</span>' for w in WORD_BANK)
    blanks = "".join(f'<div class="fl"><b>{k}.</b><span></span></div>' for k, _ in KIT)
    inner = (f'<p style="font-size:11.5pt">Label the parts of the drum set. Write the name from the word bank next to each number.</p>'
             f'<div class="box" style="margin-bottom:8pt"><span class="k">Word bank</span>{bank}</div>'
             f'<div style="width:84%;margin:0 auto 6pt">{drumkit_svg("numbers")}</div>'
             f'<div style="display:grid;grid-template-columns:1fr 1fr;gap:2pt 22pt">{blanks}</div>'
             f'<div class="box" style="margin-top:4pt"><span class="k">Think about it</span>'
             f'<p style="margin:0 0 2pt">Which three pieces make the basic beat? Which is the lowest sound?</p><div class="ln2"></div></div>')
    return page("Parts of the drum set", "STUDENT", "Day 1 · answers p. " + str(IX["key"]), inner, n, name=True)


def listen_page(n, watermark=False):
    ex = "".join(f'<div style="margin:0 0 7pt"><div class="k" style="font-size:11pt">Example {i + 1}</div>'
                 f'{grid_svg([], blank=True, label_rows=LISTEN_ROWS, cell_h=21)}</div>' for i in range(3))
    inner = (f'<p style="font-size:11pt">Your teacher plays three drum beats. Each grid is one measure: 16 squares, four beats (the shaded squares are '
             f'the beats: 1, 2, 3, 4). Put an X where you hear each sound. Count. Listen. Discuss.</p>{ex}'
             f'<div class="fl"><b>Where does the snare land in all three?</b><span></span></div>'
             f'<div class="fl"><b>Which example feels busiest? Why?</b><span></span></div>'
             f'<div class="fl"><b>What is the hi-hat doing in Example 3?</b><span></span></div>')
    return page("Drumbeat listening examples", "STUDENT", "Day 1 · answers p. " + str(IX["key"]), inner, n, watermark, name=True)


def famous_page(n):
    blocks = "".join(f'<div style="margin:0 0 9pt"><div style="display:flex;justify-content:space-between;align-items:baseline">'
                     f'<span class="k" style="font-size:12pt;color:#000">{i + 1}. {s} · <span class="l" style="text-transform:none">{a}</span></span>'
                     f'<span style="font-family:Oswald;font-weight:700;font-size:12pt" class="acc">Tempo = {bpm} BPM</span></div>'
                     f'{grid_svg(rows, cell_h=21)}</div>' for i, (s, a, bpm, rows) in enumerate(FAMOUS))
    inner = (f'<p style="font-size:11pt">Choose ONE pattern and recreate it exactly in Drum Machine Designer. The tempo must match. Each grid is one '
             f'measure (16 steps). The numbers are Logic’s positions: 1 is beat 1, 1.3 is the “and” of beat 1. Call over your teacher when you’re done.</p>'
             f'{blocks}'
             f'<div style="margin:0 0 9pt"><span class="k" style="font-size:12pt;color:#000">My copy · pattern: ____________ · tempo: _____ BPM</span>'
             f'{grid_svg([], blank=True, label_rows=["Hi-hat", "Snare", "Kick", "Clap"], cell_h=21)}</div>'
             f'<div class="box"><span class="k">Check yourself</span>'
             + "".join(f'<div class="chk" style="margin:0 0 4pt;font-size:10.5pt">{checkbox()}<span>{c}</span></div>' for c in [
                 "Every X is on the right square (count from 1)", "The tempo matches", "It loops smoothly with Cycle on", "My teacher checked it: ________"]) +
             f'</div><p class="small" style="margin-top:4pt">Songs are named for reference only; your teacher can play them from YouTube or Spotify. '
             f'The grids are simplified teaching charts, not transcriptions.</p>')
    return page("Famous drum patterns", "STUDENT", "Day 2", inner, n)


def proj1_page(n, watermark=False):
    inner = (f'<h2 class="s" style="font-size:20pt;margin:2pt 0 0">Project sheet</h2>'
             f'<p class="l" style="font-size:12pt;margin-bottom:6pt">Digital Music Production 1 · Unit 1 · Structured Project · Drum Machine Designer</p>'
             f'<h2 class="s">Overview<span class="acc"> ·</span></h2>'
             f'<p style="font-size:12pt">You’re going to build an original drum beat from nothing: no loops, no pre-made grooves. Starting with the three building '
             f'blocks of a drum kit (kick, snare, hi-hat), you’ll explore rhythm, pattern, and what makes a beat feel right. You’ll also analyze a famous drum '
             f'beat to understand how real producers think about groove. By the end, you’ll have made something entirely your own, and you’ll understand '
             f'exactly why it works.</p>'
             f'<h2 class="s">Goals<span class="acc"> ·</span></h2><ul>' + "".join(f"<li style='font-size:12pt'>{g}</li>" for g in GOALS) + '</ul>'
             f'<h2 class="s">Requirements<span class="acc"> ·</span></h2>'
             + "".join(f'<div class="chk" style="font-size:12pt;margin-bottom:9pt">{checkbox(13)}<span>{r}</span></div>' for r in REQUIREMENTS) +
             f'<div class="cols" style="margin-top:6pt"><div class="box"><span class="k">Your beat’s shape</span>{aaba_svg(225)}</div>'
             f'<div class="box"><span class="k">Key vocabulary</span><p style="margin:0">14 terms on page {IX["vocab"]}: kick, snare, hi-hat, loop, beat, '
             f'measure, step sequencer, subdivision, tempo, pattern region, cycle, AABA, volume, panning.</p></div></div>')
    return page("Project sheet · overview and requirements", "STUDENT", "Project · 1 of 2", inner, n, watermark)


def proj2_page(n):
    lis = "".join(f'<li><div class="n">{i + 1}</div><div><b class="hd">Day {d} · {h}</b>{t}</div></li>' for i, (d, h, t) in enumerate(STEPS))
    inner = (f'<h2 class="s" style="margin-top:2pt">Steps / process<span class="acc"> ·</span></h2><ol class="num" style="font-size:11.5pt">{lis}</ol>'
             f'<h2 class="s" style="margin-top:2pt">Tips for success<span class="acc"> ·</span></h2><ul>' + "".join(f"<li style='font-size:11.5pt'>{t}</li>" for t in TIPS) + '</ul>')
    return page("Project sheet · steps and tips", "STUDENT", "Project · 2 of 2", inner, n)


def vocab_page(n):
    rows = "".join(f"<tr><td style='width:24%'><b>{a}</b></td><td>{b}</td><td style='width:22%'></td></tr>" for a, b in VOCAB)
    inner = (f'<p style="font-size:11pt">Read each word and its meaning. In the last column, write an example from your own beat or a song you know.</p>'
             f'<table class="t" style="font-size:10.4pt"><tr><th>Word</th><th>Meaning</th><th>My example</th></tr>{rows}</table>'
             f'<div class="box" style="margin-top:6pt"><span class="k">Quick check</span>'
             f'<div class="fl"><b>A measure has ___ beats. In the Step Sequencer, one measure has ___ squares.</b><span></span></div>'
             f'<div class="fl" style="margin:0"><b>Eighth notes: every ______ square. Sixteenth notes: every ______ square.</b><span></span></div></div>')
    return page("Key vocabulary", "STUDENT", "14 words", inner, n, name=True)


def checklist_page(n):
    inner = (f'<p style="font-size:11.5pt">Tick each box before you submit and present. If a box isn’t ticked, you aren’t done yet.</p>'
             + "".join(f'<div class="chk" style="font-size:12pt;margin-bottom:11pt">{checkbox(15)}<span>{c}</span></div>' for c in CHECKLIST) +
             f'<h2 class="s">My three feedback questions<span class="acc"> ·</span></h2>'
             f'<p class="small">Ask about something you really want to know. Example: “Does my B section feel different enough from my A?”</p>'
             f'<div class="fl"><b>1.</b><span></span></div><div class="fl"><b>2.</b><span></span></div><div class="fl"><b>3.</b><span></span></div>'
             f'<h2 class="s">Notes<span class="acc"> ·</span></h2><div class="lines"></div>')
    return page("Submission checklist", "STUDENT", "Before you present", inner, n, name=True)


def rubric_page(n, watermark=False):
    head = "<tr><th>Category</th>" + "".join(f"<th>{lv}</th>" for lv in LEVELS) + "<th>Self</th><th>Teacher</th></tr>"
    rows = "".join(f"<tr><td class='c'>{c}</td>" + "".join(f"<td>{d}</td>" for d in ds) + "<td class='sc'></td><td class='sc'></td></tr>"
                   for c, ds in RUBRIC)
    total = (f"<tr><td class='c' colspan='5' style='text-align:right;font-size:13pt'>Total</td>"
             f"<td style='font-family:Oswald;font-weight:700;font-size:13pt;text-align:right;vertical-align:bottom'>/ {5 * len(RUBRIC)}</td>"
             f"<td style='font-family:Oswald;font-weight:700;font-size:13pt;text-align:right;vertical-align:bottom'>/ {5 * len(RUBRIC)}</td></tr>")
    inner = (f'<p style="font-size:11pt">Score each category from 2 to 5. Score yourself first, then your teacher scores.</p>'
             f'<table class="rub">{head}{rows}{total}</table>'
             f'<div class="cols" style="margin-top:10pt"><div><h2 class="s" style="margin-top:0">One thing I’m proud of</h2><div class="ln2"></div><div class="ln2"></div></div>'
             f'<div><h2 class="s" style="margin-top:0">Teacher comments</h2><div class="ln2"></div><div class="ln2"></div></div></div>')
    return page("Rubric", "RUBRIC", f"/ {5 * len(RUBRIC)}", inner, n, watermark, name=True)


def key_page(n):
    lst = "".join(f'<div style="margin:0 0 6pt"><span class="k">Example {i + 1}</span>{grid_svg(rows, cell_h=13, fill=ACCENT, label_rows=LISTEN_ROWS)}</div>'
                  for i, rows in enumerate(LISTEN))
    names = "".join(f'<div style="font-size:11pt;margin:0 0 3pt"><b>{k}.</b> <span class="acc" style="font-weight:700">{nm}</span></div>' for k, nm in KIT)
    inner = (f'<div class="cols" style="gap:16pt;align-items:flex-start"><div style="flex:0 0 3.7in">'
             f'<h2 class="s" style="margin-top:0">Parts of the drum set</h2>{drumkit_svg("answers")}</div>'
             f'<div><h2 class="s" style="margin-top:0">Labels</h2>{names}'
             f'<h2 class="s">Think about it</h2><p style="font-size:10pt">Basic beat: kick, snare, hi-hat. Lowest sound: the kick (bass drum).</p></div></div>'
             f'<h2 class="s" style="margin-top:4pt">Drumbeat listening examples<span class="acc"> ·</span></h2>'
             f'<p style="font-size:10pt">Program these in Drum Machine Designer before Day 1, one pattern region each (any tempo around 90–110 BPM works). '
             f'Example 1 is the same as the Billie Jean pattern.</p>{lst}'
             f'<div class="cols" style="gap:16pt"><div><span class="k">Listening questions</span><p style="font-size:10pt;margin:0">Snare: beats 2 and 4 in all three. '
             f'Busiest: Example 3 (sixteenth-note hi-hats, a shaker, a clap that doubles the snare, a syncopated kick). Example 3’s hi-hat plays every square: '
             f'sixteenth notes.</p></div>'
             f'<div style="flex:0 0 2.3in"><span class="k">Vocabulary quick check</span><p style="font-size:10pt;margin:0">4 beats; 16 squares. Eighths: every other '
             f'square. Sixteenths: every square.</p></div></div>')
    return page("Answer key · kit labels and listening examples", "ANSWER KEY", "Teacher", inner, n, ak=True)


def example_page(n):
    mix = "".join(f"<tr><td><b>{a}</b></td><td>{b}</td><td>{c}</td><td>{d}</td></tr>" for a, b, c, d in MIX)
    row = lambda k, cap, g: (f'<div style="margin:0 0 5pt"><span class="k">{k}</span>{g}<p class="small" style="margin:2pt 0 0">{cap}</p></div>')
    inner = (f'<p>Use this to model the build on Days 2–3 or as a reference while you circulate. It is one possible answer, not the answer: every '
             f'student’s beat should sound different.</p>'
             + row("1 · Starting point", "Kick on 1 and 3, snare on 2 and 4, hi-hat in eighths. Every hit lands on the grid, so it locks in. "
                   "Now change one thing at a time and listen.", grid_svg(START, label_w=86, cell_h=14))
             + row("2 · My A section", "Extra kick on the “and” of 2 (2.3); clap with the snare on 4; open hat on 4.3, where the closed hat drops out.",
                   grid_svg(EX_A, label_w=86, cell_h=14))
             + row("3 · My B section (one change)", "Same beat, closed hi-hat in sixteenths: more energy, still clearly related to A.",
                   grid_svg(EX_B, label_w=86, cell_h=14)) +
             f'<span class="k">4 · Arrange</span>{aaba_svg(511)}'
             f'<span class="k" style="margin-top:5pt">5 · A sample mix</span>'
             f'<table class="t" style="font-size:9.3pt"><tr><th>Track</th><th>Volume</th><th>Pan</th><th>Why</th></tr>{mix}</table>'
             f'<p class="small">Numbers are starting points for a demo, not targets. Trust your ears: if you can’t hear a sound clearly, turn something else down first.</p>'
             f'<span class="k" style="margin-top:3pt">What a 5 sounds like</span>'
             f'<p style="font-size:9.8pt;margin:0">A groove you nod your head to within a bar, a B section you notice without being told, every drum sound clear, '
             f'width from panning that doesn’t distract, and a student who can say why they made each choice.</p>')
    return page("Example notes · a worked beat, step by step", "EXAMPLE", "Teacher", inner, n, ak=True)


def terms_page(n):
    return (f'<section class="page"><div class="top"><div class="t">Terms of Use<small>Thank you for your purchase!</small></div></div>'
            f'<div class="main"><p>© {YEAR} {STORE}. {BRAND} by {STORE}. All rights reserved. This purchase gives <b>one teacher</b> a license to use this '
            f'resource (the PDF and the intro deck) with their own students.</p>'
            f'<h2 class="s">You may</h2><ul><li>Print and copy pages for your own classroom and students, year after year.</li>'
            f'<li>Present the intro deck to your own classes and edit it for your own teaching.</li>'
            f'<li>Share pages with your own students through a password-protected class site or learning platform.</li>'
            f'<li>Leave the pages for a substitute teacher who teaches your own students.</li></ul>'
            f'<h2 class="s">You may not</h2><ul><li>Share, email or copy this resource for other teachers, a whole school or a district. '
            f'Additional licenses are available on Teachers Pay Teachers.</li>'
            f'<li>Post any part of it on a public website, shared drive or social media.</li>'
            f'<li>Sell, give away or claim any part of it as your own, or edit it into a new product.</li></ul>'
            f'<h2 class="s">Questions?</h2><p>Please use the Q&amp;A tab on the {BRAND} by {STORE} store on Teachers Pay Teachers. '
            f'If something looks wrong, let me know there and I will fix it.</p>'
            f'<h2 class="s">Credits</h2><ul>'
            f'<li>Curriculum, project sheet, rubric, teacher notes and intro deck: original work by {STORE}, {BRAND}.</li>'
            f'<li>Drum kit and drum machine drawings, grids and layouts: original work, {BRAND}.</li>'
            f'<li>Text fonts: Oswald (The Oswald Project Authors) and Source Sans 3 (Adobe), SIL Open Font License 1.1.</li>'
            f'<li>Songs are named for reference only (title and artist). No audio, video, lyrics or album art are included. Logic Pro, Drum Machine Designer '
            f'and Apple Loops are trademarks of Apple Inc., named only to describe compatibility; this resource is not affiliated with or endorsed by Apple. '
            f'Padlet, YouTube and Spotify are named only as examples of tools you may already use.</li></ul>'
            f'<p class="small">{DISCLOSURE}</p><div style="flex:1"></div></div>{foot(f"Page {n}")}</section>')


def render_page(key, watermark=False):
    n = IX[key]
    f = {"cover": cover_page, "howto": lambda: howto_page(n), "overview": lambda: overview_page(n), "setup": lambda: setup_page(n),
         "pacing": lambda: pacing_page(n, watermark), "day12": lambda: day12_page(n), "day34": lambda: day34_page(n),
         "tnotes": lambda: tnotes_page(n), "tnotes2": lambda: tnotes2_page(n), "post": lambda: post_page(n), "kit": lambda: kit_page(n),
         "listen": lambda: listen_page(n, watermark), "famous": lambda: famous_page(n), "proj1": lambda: proj1_page(n, watermark),
         "proj2": lambda: proj2_page(n), "vocab": lambda: vocab_page(n), "checklist": lambda: checklist_page(n),
         "rubric": lambda: rubric_page(n, watermark), "key": lambda: key_page(n), "example": lambda: example_page(n), "terms": lambda: terms_page(n)}
    return f[key]()


# ---------------------------------------------------------------- intro deck (.pptx, python-pptx) + PDF copy (LibreOffice)
def build_deck(kit_png):
    from pptx import Presentation
    from pptx.util import Inches, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.oxml.ns import qn

    BLACK, WHITE, RED, GREY = RGBColor(0, 0, 0), RGBColor(255, 255, 255), RGBColor.from_string(ACCENT[1:]), RGBColor(0xE6, 0xE6, 0xE6)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.author = f"{STORE} · {BRAND}"
    prs.core_properties.title = f"{TITLE} · Intro deck"
    prs.core_properties.comments = DISCLOSURE
    layout = prs.slide_layouts[6]

    def box(sl, x, y, w, h, fillc=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=1.5):
        s = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        if fillc is None:
            s.fill.background()
        else:
            s.fill.solid()
            s.fill.fore_color.rgb = fillc
        if line is None:
            s.line.fill.background()
        else:
            s.line.color.rgb = line
            s.line.width = Pt(lw)
        s.shadow.inherit = False
        return s

    def text(sl, x, y, w, h, runs, size=24, font="Source Sans 3", bold=False, color=BLACK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             spacing=None, caps=False):
        tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.04)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        paras = runs if isinstance(runs, list) else [runs]
        for i, ptxt in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = ptxt.upper() if caps else ptxt
            r.font.size = Pt(size)
            r.font.name = font
            r.font.bold = bold
            r.font.color.rgb = color
            if spacing:
                r._r.get_or_add_rPr().set("spc", str(spacing))
        return tb

    def chrome(sl, i, title):
        text(sl, 0.6, 0.32, 6, 0.4, BRAND, 13, "Oswald", False, BLACK, spacing=600, caps=True)
        text(sl, 0.6, 0.72, 12.1, 0.95, title, 40, "Oswald", True, BLACK, caps=True, anchor=MSO_ANCHOR.BOTTOM)
        box(sl, 0.6, 1.72, 12.13, 0.04, BLACK)
        box(sl, 0.6, 7.0, 12.13, 0.012, BLACK)
        text(sl, 0.6, 7.04, 9, 0.3, f"{TITLE} · Intro deck · © {YEAR} {STORE} · {BRAND}", 10, "Source Sans 3", False, RGBColor(0x33, 0x33, 0x33))
        text(sl, 11.2, 7.04, 1.53, 0.3, f"{i}", 10, "Oswald", True, BLACK, align=PP_ALIGN.RIGHT)

    def bullets(sl, items, x=0.6, y=2.05, w=12.1, size=28, dot=True):
        for k, it in enumerate(items):
            yy = y + k * (size / 72 * 1.75)
            if dot:
                box(sl, x, yy + size / 72 * 0.36, 0.16, 0.16, RED, shape=MSO_SHAPE.OVAL)
            text(sl, x + (0.38 if dot else 0), yy, w - 0.4, size / 72 * 1.6, it, size, "Source Sans 3", False, BLACK)

    def grid(sl, rows, x, y, w, rh, label_w=1.5, names=None):
        names = names or [r[0] for r in rows]
        cw = (w - label_w) / 16
        for i, lab in enumerate(STEP_LABELS):
            text(sl, x + label_w + i * cw - 0.1, y - 0.36, cw + 0.2, 0.32, lab, 11 if "." in lab else 14, "Oswald", "." not in lab, BLACK,
                 align=PP_ALIGN.CENTER)
        d = dict(rows)
        for r, name in enumerate(names):
            yy = y + r * rh
            text(sl, x, yy, label_w - 0.12, rh, name, 18, "Oswald", True, BLACK, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
            for i in range(16):
                xx = x + label_w + i * cw
                box(sl, xx, yy, cw, rh, GREY if i % 4 == 0 else WHITE, BLACK, lw=2 if i % 4 == 0 else 0.75)
                if d.get(name, "." * 16)[i] == "x":
                    box(sl, xx + 0.06, yy + 0.06, cw - 0.12, rh - 0.12, RED if name == "Kick" else BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)

    for i, (kind, title, body, notes) in enumerate(SLIDES, start=1):
        sl = prs.slides.add_slide(layout)
        items = [b for b in body.split("|") if b]
        if kind == "title":
            text(sl, 0, 0.55, 13.333, 0.5, BRAND, 16, "Oswald", False, BLACK, align=PP_ALIGN.CENTER, spacing=800, caps=True)
            grid(sl, START + [("Clap", "............x...")], 3.0, 1.75, 7.33, 0.42, label_w=1.1, names=["Hi-hat", "Clap", "Snare", "Kick"])
            text(sl, 0, 3.55, 13.333, 1.3, title, 88, "Oswald", True, BLACK, align=PP_ALIGN.CENTER, caps=True)
            text(sl, 0, 5.15, 13.333, 0.6, body, 30, "Oswald", False, BLACK, align=PP_ALIGN.CENTER)
            box(sl, 0, 6.55, 13.333, 0.95, BLACK)
            text(sl, 0, 6.68, 13.333, 0.6, "Digital Music Production · Unit 1 · Drum programming in Logic Pro", 20, "Oswald", False, WHITE,
                 align=PP_ALIGN.CENTER, spacing=200, caps=True)
        else:
            chrome(sl, i, title)
        if kind == "listen":
            box(sl, 0.6, 2.1, 12.13, 1.55, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            text(sl, 0.9, 2.18, 11.5, 0.8, items[0], 40, "Oswald", True, WHITE, caps=True)
            text(sl, 0.9, 2.95, 11.5, 0.6, items[1], 20, "Source Sans 3", False, WHITE)
            bullets(sl, items[2:], y=4.05, size=32)
        elif kind in ("bullets", "closer"):
            bullets(sl, items, size=32 if kind == "bullets" else 40, y=2.2 if kind == "bullets" else 2.6)
        elif kind == "columns":
            for k, it in enumerate(items):
                x = 0.6 + k * 3.08
                box(sl, x, 2.3, 2.85, 3.9, WHITE, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=2.25)
                head, _, rest = it.partition(":")
                text(sl, x + 0.15, 2.5, 2.55, 0.8, head, 30, "Oswald", True, RED if k == 0 else BLACK, align=PP_ALIGN.CENTER, caps=True)
                if rest:
                    text(sl, x + 0.15, 3.35, 2.55, 2.5, [s.strip() for s in rest.split(",")], 24, "Source Sans 3", False, BLACK, align=PP_ALIGN.CENTER)
        elif kind == "intro":
            layers = [("Drums", 0, 16, BLACK), ("Bass", 4, 16, BLACK), ("Shaker", 8, 16, BLACK), ("Synth", 12, 16, RED)]
            x0, w = 2.6, 10.1
            for m in range(0, 17, 4):
                text(sl, x0 + m / 16 * w - 0.3, 2.0, 0.6, 0.35, str(m + 1) if m < 16 else "", 14, "Oswald", True, BLACK, align=PP_ALIGN.CENTER)
            for r, (nm, a, b, col) in enumerate(layers):
                y = 2.45 + r * 0.95
                text(sl, 0.6, y, 1.85, 0.75, nm, 24, "Oswald", True, BLACK, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE, caps=True)
                for m in range(a, b, 4):
                    box(sl, x0 + m / 16 * w + 0.04, y + 0.05, 4 / 16 * w - 0.08, 0.65, col, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                    text(sl, x0 + m / 16 * w + 0.1, y + 0.1, 2.2, 0.55, "4 measures", 14, "Source Sans 3", True, WHITE, anchor=MSO_ANCHOR.MIDDLE)
            text(sl, 0.6, 6.35, 12.1, 0.5, "Each layer enters at the start of a 4-measure block and keeps looping. In Logic, every block is a region.",
                 18, "Source Sans 3", False, BLACK)
        elif kind == "kit":
            sl.shapes.add_picture(kit_png, Inches(0.6), Inches(2.0), width=Inches(6.4))
            for k, it in enumerate(items):
                head, _, rest = it.partition(":")
                y = 2.15 + k * 1.55
                box(sl, 7.3, y + 0.12, 0.2, 0.2, RED, shape=MSO_SHAPE.OVAL)
                text(sl, 7.65, y, 5.1, 0.55, head, 30, "Oswald", True, BLACK, caps=True)
                text(sl, 7.65, y + 0.58, 5.1, 0.9, rest.strip(), 19, "Source Sans 3", False, BLACK)
        elif kind == "grid":
            grid(sl, [("Eighths", "x.x.x.x.x.x.x.x."), ("Sixteenths", "x" * 16)], 0.6, 2.55, 12.13, 0.62, label_w=1.9)
            bullets(sl, items, y=4.25, size=26)
        elif kind == "pattern":
            grid(sl, FAMOUS[0][3], 0.6, 2.55, 12.13, 0.78, label_w=1.6)
            text(sl, 0.6, 5.05, 12.1, 0.6, "Billie Jean · Michael Jackson · 116 BPM", 30, "Oswald", True, BLACK, caps=True)
            text(sl, 0.6, 5.7, 12.1, 1.0, ["Hi-hat in eighth notes · snare on 2 and 4 · kick on 1 and 3",
                                           "Teacher: play the song from YouTube or Spotify and compare."], 20, "Source Sans 3", False, BLACK)
        elif kind == "steps":
            for k, it in enumerate(items):
                y = 2.05 + k * 0.95
                box(sl, 0.6, y + 0.05, 0.62, 0.62, BLACK, shape=MSO_SHAPE.OVAL)
                text(sl, 0.6, y + 0.05, 0.62, 0.62, str(k + 1), 24, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                text(sl, 1.45, y, 11.2, 0.8, it, 24, "Source Sans 3", False, BLACK, anchor=MSO_ANCHOR.MIDDLE)
        elif kind == "aaba":
            days = [("Day 1", "Name the kit · loops · listening"), ("Day 2", "Copy a famous pattern"),
                    ("Day 3", "Build A · build B · arrange AABA · mix · submit"), ("Day 4", "Present + feedback")]
            for k, (d, t) in enumerate(days):
                y = 2.05 + k * 0.72
                box(sl, 0.6, y, 1.35, 0.55, RED if k == 2 else BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                text(sl, 0.6, y, 1.35, 0.55, d, 20, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, caps=True)
                text(sl, 2.15, y, 10.5, 0.55, t, 24, "Source Sans 3", False, BLACK, anchor=MSO_ANCHOR.MIDDLE)
            x0, w = 0.6, 12.13
            for sec, m0 in AABA:
                x = x0 + (m0 - 1) / 16 * w
                box(sl, x + 0.04, 5.25, 4 / 16 * w - 0.08, 0.85, RED if sec == "B" else BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                text(sl, x, 5.25, 4 / 16 * w, 0.85, sec, 36, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                text(sl, x, 6.12, 3, 0.35, f"measures {m0}–{m0 + 3}", 14, "Source Sans 3", False, BLACK)
            text(sl, 0.6, 6.45, 12.1, 0.45, "B = A with one change: remove one element, add one sound, or change the hi-hat rhythm.", 16,
                 "Source Sans 3", True, BLACK)
        elif kind == "check":
            for k, it in enumerate(items):
                col, row = k % 2, k // 2
                x, y = 0.6 + col * 6.15, 2.15 + row * 1.5
                box(sl, x, y + 0.08, 0.42, 0.42, WHITE, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=2.25)
                text(sl, x + 0.6, y, 5.35, 1.3, it, 22, "Source Sans 3", False, BLACK)
        elif kind == "rubric":
            for k, (c, ds) in enumerate(RUBRIC):
                y = 2.05 + k * 0.92
                box(sl, 0.6, y, 4.3, 0.78, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                text(sl, 0.75, y, 4.0, 0.78, c, 20, "Oswald", True, WHITE, anchor=MSO_ANCHOR.MIDDLE, caps=True)
                text(sl, 5.1, y, 6.6, 0.78, "5 = " + ds[-1], 15, "Source Sans 3", False, BLACK, anchor=MSO_ANCHOR.MIDDLE)
                text(sl, 11.7, y, 1.03, 0.78, "/ 5", 22, "Oswald", True, RED, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
            text(sl, 9.0, 6.6, 3.73, 0.4, f"Total / {5 * len(RUBRIC)}", 20, "Oswald", True, BLACK, align=PP_ALIGN.RIGHT, caps=True)
        if kind == "closer":
            grid(sl, EX_A, 2.6, 4.7, 8.13, 0.36, label_w=1.6)
        sl.notes_slide.notes_text_frame.text = notes
    prs.save(DECK)


def ttf_fonts():
    """woff2 -> ttf for LibreOffice (deck PDF), plus a fontconfig file that adds them."""
    from fontTools.ttLib import TTFont
    d = os.path.join(BUILD, "ttf")
    os.makedirs(d, exist_ok=True)
    # The web subsets carry odd family names ("Source Sans 3 ExtraLight"); give LibreOffice a clean 4-style family each.
    styles = {"oswald-latin-500-normal": ("Oswald", "Regular", 400), "oswald-latin-700-normal": ("Oswald", "Bold", 700),
              "source-sans-3-latin-400-normal": ("Source Sans 3", "Regular", 400), "source-sans-3-latin-700-normal": ("Source Sans 3", "Bold", 700),
              "source-sans-3-latin-400-italic": ("Source Sans 3", "Italic", 400)}
    for stem, (fam, sub, wt) in styles.items():
        out = os.path.join(d, stem + ".ttf")
        if os.path.exists(out):
            continue
        t = TTFont(os.path.join(FONTS, stem + ".woff2"))
        t.flavor = None
        name = t["name"]
        ps = (fam.replace(" ", "") + "-" + sub)
        for rec in list(name.names):
            if rec.nameID in (16, 17, 21, 22):
                name.removeNames(nameID=rec.nameID)
        for nid, val in ((1, fam), (2, sub), (4, f"{fam} {sub}" if sub != "Regular" else fam), (6, ps)):
            name.setName(val, nid, 3, 1, 0x409)
            name.setName(val, nid, 1, 0, 0)
        t["OS/2"].usWeightClass = wt
        sel = 0
        if sub == "Bold":
            sel |= 1 << 5
        if sub == "Italic":
            sel |= 1
        if sub == "Regular":
            sel |= 1 << 6
        t["OS/2"].fsSelection = (t["OS/2"].fsSelection & ~0b1100001) | sel
        t["head"].macStyle = (1 if sub == "Bold" else 0) | (2 if sub == "Italic" else 0)
        t.save(out)
    conf = os.path.join(BUILD, "fonts.conf")
    with open(conf, "w") as fh:
        fh.write(f'<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig><include ignore_missing="yes">/etc/fonts/fonts.conf</include>'
                 f'<dir>{d}</dir><cachedir>{os.path.join(BUILD, "fccache")}</cachedir></fontconfig>')
    return conf


def deck_pdf():
    conf = ttf_fonts()
    env = dict(os.environ, FONTCONFIG_FILE=conf)
    prof = os.path.join(BUILD, "lo-profile")
    subprocess.run(["soffice", f"-env:UserInstallation=file://{prof}", "--headless", "--convert-to", "pdf", "--outdir", BUILD, DECK],
                   check=True, env=env, capture_output=True)
    shutil.move(os.path.join(BUILD, os.path.basename(DECK).replace(".pptx", ".pdf")), DECK_PDF)


# ---------------------------------------------------------------- square images (2000 x 2000)
SQ_CSS = """
.sq { width: 1000px; height: 1000px; position: relative; overflow: hidden; background: linear-gradient(180deg,#fff 0%,#fff 60%,#eeeeee 100%);
      display: flex; flex-direction: column; align-items: center; }
.sq .store { font-family: 'Oswald'; font-weight: 500; letter-spacing: .32em; text-transform: uppercase; font-size: 22px; margin-top: 34px; }
.sq .hd { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; text-align: center; line-height: .95; white-space: nowrap; }
.sq .sub { font-family: 'Oswald'; font-weight: 300; text-align: center; }
.sq .band { position: absolute; left: 0; right: 0; bottom: 0; background: #000; color: #fff; text-align: center; font-family: 'Oswald';
            font-weight: 500; text-transform: uppercase; letter-spacing: .1em; font-size: 26px; padding: 16px 0 18px; }
.sq .band em { font-style: normal; color: ACCENT; }
.shot { background: #fff; border: 2px solid #000; box-shadow: 8px 8px 0 #000; position: relative; overflow: hidden; }
.shot img { display: block; width: 100%; }
.lbl { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 22px; text-align: center; margin-top: 16px; }
.lbl span { font-weight: 300; text-transform: none; display: block; font-size: 18px; }
.swm { position: absolute; left: 50%; top: 50%; transform: translate(-50%,-50%) rotate(-38deg); font-family: 'Oswald'; font-weight: 700;
       color: rgba(0,0,0,.14); white-space: nowrap; letter-spacing: .06em; }
""".replace("ACCENT", ACCENT)
FONT_PRIME = ('<div style="position:absolute;left:-9999px;font-family:\'Source Sans 3\'">'
              '<span style="font-weight:400">a</span><span style="font-weight:600">a</span><i style="font-weight:400">a</i></div>')


def sq_doc(body):
    return (f'<!doctype html><html><head><meta charset="utf-8"><style>{font_css()}'
            f'html,body{{margin:0;background:#fff}}{SQ_CSS}</style></head><body>{FONT_PRIME}{body}</body></html>')


def cover_square():
    return sq_doc(
        f'<div class="sq"><div class="store">{BRAND}</div>'
        f'<div style="width:760px;margin-top:30px">{machine_svg()}</div>'
        f'<div class="hd" style="font-size:58px;margin-top:34px">Drum Programming in Logic Pro</div>'
        f'<div class="hd" style="font-size:106px">Beat From Nothing</div>'
        f'<div class="sub" style="font-size:29px;margin-top:14px">4-day unit · teacher notes · project sheet · /25 rubric · intro slides</div>'
        f'<div class="band">Grades 9–12 <em>·</em> 4 days <em>·</em> {len(PLAN)} pages + {DECK_SLIDES} slides</div></div>')


def preview_square_1(img):
    items = [(img[0], "Pacing + plans", "4 days, minute by minute"), (img[1], "Project sheet", "goals, requirements, steps"),
             (img[2], "Rubric", "5 categories, / 25")]
    shots = "".join(f'<div style="width:290px"><div class="shot"><img src="{s}"><div class="swm" style="font-size:64px">PREVIEW</div></div>'
                    f'<div class="lbl">{a}<span>{b}</span></div></div>' for s, a, b in items)
    return sq_doc(
        f'<div class="sq"><div class="store">{BRAND}</div>'
        f'<div class="hd" style="font-size:74px;margin-top:20px">Ready to teach Monday</div>'
        f'<div class="sub" style="font-size:27px;margin-top:8px">Teacher notes, day plans, student pages and the rubric in one PDF</div>'
        f'<div style="display:flex;gap:28px;margin-top:56px;align-items:flex-start">{shots}</div>'
        f'<div class="band">No loops <em>·</em> no pre-made grooves <em>·</em> built from nothing</div></div>')


def preview_square_2(slides):
    cells = "".join(f'<div class="shot" style="box-shadow:6px 6px 0 #000"><img src="{s}"><div class="swm" style="font-size:54px">PREVIEW</div></div>'
                    for s in slides)
    return sq_doc(
        f'<div class="sq"><div class="store">{BRAND}</div>'
        f'<div class="hd" style="font-size:74px;margin-top:20px">Intro deck included</div>'
        f'<div class="sub" style="font-size:27px;margin-top:8px">{DECK_SLIDES} editable slides (.pptx + PDF) with speaker notes on every slide</div>'
        f'<div style="display:grid;grid-template-columns:420px 420px;gap:30px 34px;margin-top:44px">{cells}</div>'
        f'<div class="band">Hook <em>·</em> listen <em>·</em> worked example <em>·</em> steps <em>·</em> rubric</div></div>')


def preview_square_3(imgs):
    cells = "".join(f'<div style="width:196px"><div class="shot" style="box-shadow:5px 5px 0 #000"><img src="{s}">'
                    f'<div class="swm" style="font-size:44px">PREVIEW</div></div><div class="lbl" style="font-size:17px;margin-top:9px">{lab}</div></div>'
                    for s, lab in imgs)
    return sq_doc(
        f'<div class="sq"><div class="store">{BRAND}</div>'
        f'<div class="hd" style="font-size:70px;margin-top:16px">Everything inside</div>'
        f'<div class="sub" style="font-size:27px;margin-top:8px">Student pages, answer key and a worked example beat</div>'
        f'<div style="display:grid;grid-template-columns:repeat(4,196px);gap:20px 26px;margin-top:34px">{cells}</div>'
        f'<div class="band">{len(PLAN)} pages <em>·</em> US Letter PDF <em>·</em> prints in black &amp; white</div></div>')


SHOTS1 = ["pacing", "proj1", "rubric"]
SHOTS2 = [1, 6, 9, 14]   # deck slides
SHOTS3 = [("post", "Day 1 post"), ("kit", "Drum kit labels"), ("listen", "Listening grid"), ("famous", "Famous patterns"),
          ("vocab", "Vocabulary"), ("checklist", "Checklist"), ("key", "Answer key"), ("example", "Worked example")]


# ---------------------------------------------------------------- render
def node_env():
    env = dict(os.environ)
    if not env.get("NODE_PATH"):
        env["NODE_PATH"] = subprocess.run(["npm", "root", "-g"], capture_output=True, text=True).stdout.strip()
    return env


def render(jobs):
    path = os.path.join(BUILD, "jobs.json")
    with open(path, "w") as f:
        json.dump(jobs, f)
    subprocess.run(["node", RENDER_JS, path], check=True, env=node_env())


def write(name, txt):
    p = os.path.join(BUILD, name)
    with open(p, "w") as f:
        f.write(txt)
    return p


def pdf_png(src, page_no, out, res=110):
    subprocess.run(["pdftoppm", "-png", "-r", str(res), "-f", str(page_no), "-l", str(page_no), "-singlefile", src,
                    os.path.join(BUILD, out)], check=True)


def tpt_json():
    return {
        "price": PRICE, "license_price": PRICE,
        "tax_code": "Other Digital Goods - No Physical Media",
        "grades": ["9th-grade", "10th-grade", "11th-grade", "12th-grade"],
        "subjects": ["Music"],
        "tags": ["Autumn"],
        "formats": ["PDF"],
        "pages": len(PLAN) + DECK_SLIDES,
        "answer_key": "Included",
        "zip_contents": [os.path.basename(PDF), os.path.basename(DECK), os.path.basename(DECK_PDF)],
        "files": {"product": os.path.basename(ZIP), "preview": os.path.basename(PREVIEW_PDF), "thumb1": "cover.png",
                  "thumb2": "preview-1.png", "thumb3": "preview-2.png", "thumb4": "preview-3.png"},
    }


def make_zip():
    if os.path.exists(ZIP):
        os.remove(ZIP)
    with zipfile.ZipFile(ZIP, "w", zipfile.ZIP_DEFLATED) as z:
        for p in (PDF, DECK, DECK_PDF):
            zi = zipfile.ZipInfo(f"Beat From Nothing - Hudson Beat/{os.path.basename(p)}", date_time=(2026, 10, 11, 12, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            with open(p, "rb") as fh:
                z.writestr(zi, fh.read())


def build():
    os.makedirs(BUILD, exist_ok=True)
    problems = verify_data()
    if problems:
        print("\n".join(problems))
        sys.exit("data check failed; not rendering")
    full = write("resource.html", html_doc([render_page(p) for p in PLAN], TITLE))
    prev = write("preview.html", html_doc([render_page(p, watermark=True) for p in PREVIEW_PICKS], "Preview"))
    kit = write("kit.html", f'<!doctype html><html><head><meta charset="utf-8"><style>html,body{{margin:0;background:#fff}}</style></head>'
                            f'<body><div style="width:1280px;height:840px">{drumkit_svg("plain", width="1280")}</div></body></html>')
    kit_png = os.path.join(BUILD, "kit.png")
    render([{"html": full, "pdf": PDF}, {"html": prev, "pdf": PREVIEW_PDF}, {"html": kit, "png": kit_png, "width": 1280, "height": 840, "scale": 1}])
    build_deck(kit_png)
    deck_pdf()
    make_zip()
    res = 45 if LOW_RES else 110
    for key in SHOTS1:
        pdf_png(PDF, IX[key], key, res)
    for key, _ in SHOTS3:
        pdf_png(PDF, IX[key], key, 40 if LOW_RES else 80)
    for s in SHOTS2:
        pdf_png(DECK_PDF, s, f"slide{s}", 30 if LOW_RES else 70)
    sq = [write("cover.html", cover_square()), write("p1.html", preview_square_1([f"{k}.png" for k in SHOTS1])),
          write("p2.html", preview_square_2([f"slide{s}.png" for s in SHOTS2])),
          write("p3.html", preview_square_3([(f"{k}.png", lab) for k, lab in SHOTS3]))]
    render([{"html": h, "png": o, "width": 1000, "height": 1000, "scale": 1 if LOW_RES else 2} for h, o in zip(sq, [COVER_PNG] + PREVIEW_PNGS)])
    with open(TPT_JSON, "w") as f:
        json.dump(tpt_json(), f, indent=2, ensure_ascii=False)
        f.write("\n")
    print(f"built {PDF} ({len(PLAN)} pages), {DECK} + PDF ({DECK_SLIDES} slides), {ZIP}, preview, cover + 3 previews, tpt.json"
          + (" [LOW RES]" if LOW_RES else ""))


# ---------------------------------------------------------------- checks
def steps_of(rows, name):
    return [i for i, c in enumerate(dict(rows)[name]) if c == "x"]


def verify_data():
    errs = []
    for label, rows in [(f"famous {s}", r) for s, _, _, r in FAMOUS] + [(f"listen {i + 1}", r) for i, r in enumerate(LISTEN)] + \
                       [("start", START), ("example A", EX_A), ("example B", EX_B)]:
        for name, st in rows:
            if len(st) != 16 or set(st) - set("x."):
                errs.append(f"{label} {name}: not 16 steps of x/.")
    bj = FAMOUS[0]
    if (bj[0], bj[1], bj[2]) != ("Billie Jean", "Michael Jackson", 116):
        errs.append("Billie Jean title/artist/tempo changed")
    if steps_of(bj[3], "Kick") != [0, 8] or steps_of(bj[3], "Snare") != [4, 12] or steps_of(bj[3], "Hi-hat") != list(range(0, 16, 2)):
        errs.append("Billie Jean grid is not kick 1+3, snare 2+4, hats in eighths (as the slides and notes say)")
    if [r for r in LISTEN[0]] != bj[3]:
        errs.append("listening example 1 should equal the Billie Jean pattern (key says so)")
    if START != bj[3]:
        errs.append("starting point should be kick 1+3, snare 2+4, hats eighths")
    for i, ex in enumerate(LISTEN):
        if steps_of(ex, "Snare") != [4, 12]:
            errs.append(f"listen {i + 1}: key says snare on 2 and 4")
        if any(n not in LISTEN_ROWS for n, _ in ex):
            errs.append(f"listen {i + 1}: row not on the student grid")
    if steps_of(LISTEN[2], "Hi-hat") != list(range(16)):
        errs.append("key says Example 3 hi-hat is sixteenths")
    diff = [a for (a, x), (_, y) in zip(EX_A, EX_B) if x != y]
    if diff != ["Closed hi-hat"]:
        errs.append(f"example B should differ from A in exactly the closed hi-hat row, differs in {diff}")
    if steps_of(EX_A, "Kick") != [0, 6, 8] or 14 in steps_of(EX_A, "Closed hi-hat") or steps_of(EX_A, "Open hi-hat") != [14]:
        errs.append("example A text (kick 2.3, open hat 4.3, closed hat drops out) no longer matches the grid")
    if [m for _, m in AABA] != [1, 5, 9, 13] or "".join(s for s, _ in AABA) != "AABA":
        errs.append("AABA map is not 4 x 4 measures")
    if len(VOCAB) != 14 or len(GOALS) != 4 or len(TIPS) != 6 or len(CHECKLIST) != 9 or len(STEPS) != 8:
        errs.append("owner's sheet counts changed (14 vocab, 4 goals, 6 tips, 9 checklist, 8 steps)")
    if len(RUBRIC) != 5 or any(len(d) != len(LEVELS) for _, d in RUBRIC) or 5 * len(RUBRIC) != 25:
        errs.append("rubric must be 5 categories x levels 2-5 = /25")
    if sorted(n for _, n in KIT) != sorted(WORD_BANK) or [k for k, _ in KIT] != list(range(1, 9)):
        errs.append("kit callouts must use every word-bank name once, numbered 1-8")
    if len(DAYS) != 4 or [d["kind"].split()[0] for d in DAYS] != ["Learning", "Learning", "Work", "Presenting"]:
        errs.append("pacing must be 4 days: 2 learning, 1 work, 1 present")
    for d in DAYS:
        blocks = [tuple(map(int, a.split("–"))) for a, _, _ in d["blocks"]]
        if blocks[0][0] != 0 or blocks[-1][1] != 53 or any(blocks[i][1] != blocks[i + 1][0] for i in range(len(blocks) - 1)):
            errs.append(f"Day {d['n']} minute blocks do not run 0-53 without gaps")
    if len(SLIDES) < 10 or any(not s[3].strip() for s in SLIDES):
        errs.append("deck needs 10+ slides, each with notes")
    if len(PLAN) != len(set(PLAN)):
        errs.append("plan has a repeated page")
    return errs


def pdf_fonts(path):
    out = subprocess.run(["pdffonts", path], capture_output=True, text=True, check=True).stdout.splitlines()[2:]
    return [(ln.split()[0], ln.split()[-5]) for ln in out]


def pdf_text(path):
    return subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True, check=True).stdout


def deck_texts():
    from pptx import Presentation
    prs = Presentation(DECK)
    slides, notes = [], []
    for s in prs.slides:
        slides.append(" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame))
        notes.append(s.notes_slide.notes_text_frame.text if s.has_notes_slide else "")
    media = [p for p in zipfile.ZipFile(DECK).namelist() if p.startswith("ppt/media/")]
    return prs, slides, notes, media


def verify_files():
    errs = []
    from pypdf import PdfReader
    from PIL import Image
    for p in [PDF, DECK, DECK_PDF, ZIP, PREVIEW_PDF, COVER_PNG, *PREVIEW_PNGS, UPLOAD, TPT_JSON]:
        if not os.path.exists(p):
            errs.append(f"missing {os.path.basename(p)}")
    if errs:
        return errs
    # page counts and sizes
    for path, n_exp in ((PDF, len(PLAN)), (PREVIEW_PDF, len(PREVIEW_PICKS))):
        rd = PdfReader(path)
        if len(rd.pages) != n_exp:
            errs.append(f"{os.path.basename(path)}: {len(rd.pages)} pages, expected {n_exp}")
        for i, pg_ in enumerate(rd.pages):
            w, h = float(pg_.mediabox.width), float(pg_.mediabox.height)
            if abs(w - 612) > 1 or abs(h - 792) > 1:
                errs.append(f"{os.path.basename(path)} p{i + 1}: {w:.0f}x{h:.0f}pt, not US Letter")
    if len(PdfReader(PREVIEW_PDF).pages) != 4:
        errs.append("preview PDF must be 4 pages")
    drd = PdfReader(DECK_PDF)
    if len(drd.pages) != DECK_SLIDES:
        errs.append(f"deck PDF: {len(drd.pages)} pages, expected {DECK_SLIDES}")
    for i, pg_ in enumerate(drd.pages):
        if abs(float(pg_.mediabox.width) / float(pg_.mediabox.height) - 16 / 9) > 0.01:
            errs.append(f"deck PDF p{i + 1}: not 16:9")
    # fonts: only Oswald / Source Sans 3, all embedded (no silent fallback)
    for path in (PDF, PREVIEW_PDF, DECK_PDF):
        fonts = pdf_fonts(path)
        for name, emb in fonts:
            if emb != "yes":
                errs.append(f"{os.path.basename(path)}: font {name} not embedded")
            if not re.search(r"oswald|sourcesans", name.lower().replace(" ", "")):
                errs.append(f"{os.path.basename(path)}: unexpected font {name} (fallback?)")
        names = " ".join(n for n, _ in fonts).lower()
        for fam in ("oswald", "sourcesans3"):
            if fam not in names:
                errs.append(f"{os.path.basename(path)}: {fam} missing")
    # deck: slide count, speaker notes on every slide, no embedded audio/video, only our kit drawing as media
    prs, stext, notes, media = deck_texts()
    if len(prs.slides) != DECK_SLIDES:
        errs.append(f"deck has {len(prs.slides)} slides, expected {DECK_SLIDES}")
    for i, nt in enumerate(notes):
        if len(nt.strip()) < 80:
            errs.append(f"deck slide {i + 1}: speaker notes missing or too short")
    if any(not m.lower().endswith(".png") for m in media) or len(media) > 1:
        errs.append(f"deck media should be only the drawn kit PNG, found {media}")
    if abs(prs.slide_width / prs.slide_height - 16 / 9) > 0.01:
        errs.append("deck is not 16:9")
    alldeck = " ".join(stext + notes)
    if "Billie Jean" not in alldeck or "Michael Jackson" not in alldeck or not re.search(r"YouTube or Spotify", alldeck):
        errs.append("deck must name Billie Jean / Michael Jackson and tell the teacher to use YouTube or Spotify")
    if not re.search(r"YouTube or Spotify", stext[1]) or not re.search(r"YouTube or Spotify", notes[0]):
        errs.append("deck: the listen slide and slide 1 notes must say YouTube or Spotify")
    # zip contents
    names = zipfile.ZipFile(ZIP).namelist()
    for p in (PDF, DECK, DECK_PDF):
        if not any(n.endswith("/" + os.path.basename(p)) for n in names):
            errs.append(f"zip is missing {os.path.basename(p)}")
        else:
            data = zipfile.ZipFile(ZIP).read(next(n for n in names if n.endswith("/" + os.path.basename(p))))
            if data != open(p, "rb").read():
                errs.append(f"zip copy of {os.path.basename(p)} is stale; rebuild")
    # resource text: footers, page numbers, cross-references
    rd = PdfReader(PDF)
    pages = [subprocess.run(["pdftotext", "-layout", "-f", str(i + 1), "-l", str(i + 1), PDF, "-"], capture_output=True, text=True,
                            check=True).stdout for i in range(len(rd.pages))]
    flat = [re.sub(r"\s+", " ", t) for t in pages]
    for i, t in enumerate(flat):
        if i > 0 and "Brian Lotze · Hudson Beat" not in t:
            errs.append(f"page {i + 1}: footer missing")
        if i > 0 and not re.search(rf"Page {i + 1}\s*$", t.strip()):
            errs.append(f"page {i + 1}: footer page number wrong")
    for m in re.finditer(r"\b(?:pages?|pp?\.) (\d+)(?:–(\d+))?", " ".join(flat[1:])):
        for g in m.groups():
            if g and not 1 <= int(g) <= len(PLAN):
                errs.append(f"cross-reference to page {g} does not exist")
    if f"/ {5 * len(RUBRIC)}" not in flat[IX["rubric"] - 1] or any(c.split()[0].upper() not in flat[IX["rubric"] - 1].upper() for c, _ in RUBRIC):
        errs.append("rubric page: categories or /25 total missing")
    for c, ds in RUBRIC:   # the owner's descriptor text is printed verbatim
        for d in ds:
            if re.sub(r"\s+", " ", d) not in re.sub(r"[\s]+", " ", " ".join(flat[IX["rubric"] - 1].split())).replace("- ", "-"):
                # pdftotext may wrap inside table cells; fall back to a word check
                words = re.findall(r"[A-Za-z]{5,}", d)
                if not all(w in flat[IX["rubric"] - 1] for w in words):
                    errs.append(f"rubric text missing: {d[:50]}")
    for w, _ in VOCAB:
        if w not in flat[IX["vocab"] - 1]:
            errs.append(f"vocab page missing {w}")
    for nm in WORD_BANK:
        if nm.lower() not in flat[IX["kit"] - 1].lower() or nm not in flat[IX["key"] - 1]:
            errs.append(f"kit worksheet/key missing {nm}")
    for d in DAYS:
        if f"Day {d['n']}" not in flat[IX["pacing"] - 1].replace("DAY", "Day"):
            errs.append(f"pacing page missing Day {d['n']}")
    ptxt = [pg.extract_text() for pg in PdfReader(PREVIEW_PDF).pages]
    if sum("PREVIEW" in t for t in ptxt) != len(PREVIEW_PICKS):
        errs.append("preview PDF: watermark missing on some pages")
    # banned strings in EVERY output
    blobs = {"resource PDF": pdf_text(PDF), "preview PDF": pdf_text(PREVIEW_PDF), "deck PDF": pdf_text(DECK_PDF), "deck pptx": alldeck,
             "UPLOAD.md": open(UPLOAD).read(), "tpt.json": open(TPT_JSON).read(), "zip names": " ".join(names)}
    with zipfile.ZipFile(DECK) as z:   # raw XML too (titles, properties, notes)
        blobs["deck XML"] = " ".join(z.read(n).decode("utf8", "ignore") for n in z.namelist() if n.endswith(".xml"))
    for path in (PDF, PREVIEW_PDF, DECK_PDF):
        meta = PdfReader(path).metadata or {}
        blobs[f"{os.path.basename(path)} metadata"] = " ".join(str(v) for v in meta.values())
    for label, txt in blobs.items():
        low = txt.lower()
        for b in BANNED:
            if b.lower() in low:
                errs.append(f"{label}: contains banned string '{b}'")
    for p in [COVER_PNG, *PREVIEW_PNGS]:
        if Image.open(p).size != (2000, 2000):
            errs.append(f"{os.path.basename(p)}: not 2000x2000 (low-res build?)")
    # tpt.json
    tj = json.load(open(TPT_JSON))
    if tj["price"] != PRICE or tj["license_price"] != PRICE or tj.get("tax_code") != "Other Digital Goods - No Physical Media":
        errs.append("tpt.json price/license_price/tax_code")
    if tj["grades"] != ["9th-grade", "10th-grade", "11th-grade", "12th-grade"] or tj["subjects"] != ["Music"]:
        errs.append("tpt.json grades 9-12 / subject Music")
    if not (1 <= len(tj["tags"]) <= 6 and len(tj["formats"]) <= 3):
        errs.append("tpt.json tags/formats counts")
    if set(tj["files"]) - {"product", "preview", "thumb1", "thumb2", "thumb3", "thumb4"}:
        errs.append("tpt.json: unknown file slot (tpt.js rejects it)")
    for k, f in tj["files"].items():
        if not os.path.exists(os.path.join(HERE, f)):
            errs.append(f"tpt.json file {k} missing: {f}")
    if tj["pages"] != len(rd.pages) + len(prs.slides):
        errs.append("tpt.json pages != PDF pages + slides")
    # UPLOAD.md
    up = open(UPLOAD).read()
    m = re.search(r"^## 1\. Title\s*\n+```\n(.+?)\n```\n\((\d+) characters", up, re.M | re.S)
    if not m:
        errs.append("UPLOAD.md: title block or its character count not found")
    else:
        n = len(m.group(1).strip())
        if n > 80:
            errs.append(f"UPLOAD.md: title is {n} chars (max 80)")
        if int(m.group(2)) != n:
            errs.append(f"UPLOAD.md: says {m.group(2)} characters, title has {n}")
    m2 = re.search(r"^## 2\. Description\s*\n+```\n(.+?)\n```", up, re.M | re.S)
    if not m2 or DISCLOSURE not in m2.group(1):
        errs.append("UPLOAD.md: description block missing the disclosure line")
    else:
        desc = m2.group(1)
        for need in (f"{len(PLAN)}-page", f"{DECK_SLIDES}-slide", "Logic Pro", "/25"):
            if need not in desc:
                errs.append(f"UPLOAD.md description: missing '{need}'")
        if re.search(r"  +\S|\s[,.]\s|\bteh\b|\s{2,}$", desc, re.M):
            errs.append("UPLOAD.md description: double space or stray punctuation")
    for need in (f"${PRICE}", f"{len(PLAN)} pages", os.path.basename(ZIP), os.path.basename(DECK), os.path.basename(DECK_PDF)):
        if need not in up:
            errs.append(f"UPLOAD.md: missing '{need}'")
    for f in ("release", "uploaded"):
        if os.path.exists(os.path.join(HERE, f)):
            errs.append(f"{f} exists: only the Manager / uploader writes it")
    return errs


def check():
    errs = verify_data() + verify_files()
    if errs:
        print("CHECK FAILED")
        print("\n".join(" - " + e for e in errs))
        sys.exit(1)
    print(f"CHECK OK: drum grids are 16 steps and match the notes (Billie Jean kick 1+3 / snare 2+4 / eighth hats at 116, listening key, "
          f"example A/B one change, AABA 4x4); rubric 5 x (2-5) = /25 printed verbatim; 4-day pacing blocks run 0-53; "
          f"{len(PLAN)} Letter pages with footers and valid cross-references; preview 4 watermarked pages; deck {DECK_SLIDES} slides 16:9 with speaker "
          f"notes on every slide, no audio/video (only the drawn kit PNG), PDF copy {DECK_SLIDES} pages; only Oswald/Source Sans 3, all embedded; "
          f"no {', '.join(BANNED[:3])} strings in any output; zip holds current PDF + pptx + deck PDF; images 2000x2000; tpt.json and UPLOAD.md ok.")


if __name__ == "__main__":
    if "--check" in sys.argv:
        check()
    else:
        build()
