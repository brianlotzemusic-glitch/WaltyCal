#!/usr/bin/env python3
"""DMP1 Unit 6: Spooky Tunes (TpT 009): build the teacher resource, the intro deck and the listing files.

    python3 gen.py            build everything into this folder
    python3 gen.py --lowres   same, but square images at scale 1 and small page shots (self-check only; --check rejects them)
    python3 gen.py --check    verify data, files, page/slide counts, speaker notes, fonts, banned strings, tpt.json and UPLOAD.md

The owner's own curriculum (Brian Lotze, Digital Music Production 1, "Spooky Tunes"), reformatted in the Hudson Beat
style: a 6-day scoring-to-picture project in Logic Pro. Students score a one-minute night-woods video (32 measures at
128 BPM) with reverb and delay on shared buses. No key, chord progression or song form is required (the owner's May project
sheet; owner decision 12 Oct: the Restructure's minor-key rule is dropped). Source documents (owner's Drive, read-only):
Project Sheet & Rubric, Teacher Notes, Day 1 Classroom Post, How to Bus - Reverb and Delay, the Chords and Scales
worksheet and the DMP1 Restructure (Unit 6: 6 days 1/3/2, 32 measures, reflection, concrete descriptors). Minor keys appear only
as an optional tool (tip, Minor Key Map, reference card); the rubric keeps the owner's wording.
School names and logos are removed; no student work, names or grades; no audio, video or album art. The owner's
classroom video is NOT included: the teacher supplies a clip (own footage or a royalty-free stock clip), and the
product says so. Songs are named by title and artist only, with a find-it-on-YouTube/Spotify note.

Outputs: <slug>.pdf (teacher + student resource), <slug>-intro-deck.pptx and its PDF copy, a ZIP of those three
(the TpT product file), a 4-page watermarked preview PDF, cover.png and preview-1..3.png (2000 x 2000).
Renders with Playwright + Chromium (../tools/render.js) and LibreOffice (deck PDF), fonts from ../fonts. $0 AI.
"""
import json
import math
import os
import random
import re
import shutil
import subprocess
import sys
import zipfile
from html import escape

HERE = os.path.dirname(os.path.abspath(__file__))
TPT = os.path.dirname(HERE)
FONTS = os.path.join(TPT, "fonts")
BUILD = os.path.join(HERE, "build")
RENDER_JS = os.path.join(TPT, "tools", "render.js")

SLUG = "spooky-tunes"
PDF = os.path.join(HERE, f"{SLUG}.pdf")
DECK = os.path.join(HERE, f"{SLUG}-intro-deck.pptx")
DECK_PDF = os.path.join(HERE, f"{SLUG}-intro-deck.pdf")
ZIP = os.path.join(HERE, f"{SLUG}.zip")
PREVIEW_PDF = os.path.join(HERE, f"{SLUG}-PREVIEW.pdf")
COVER_PNG = os.path.join(HERE, "cover.png")
PREVIEW_PNGS = [os.path.join(HERE, f"preview-{i}.png") for i in (1, 2, 3)]
UPLOAD = os.path.join(HERE, "UPLOAD.md")
TPT_JSON = os.path.join(HERE, "tpt.json")
ZIP_DIR = "Spooky Tunes - Hudson Beat"

STORE = "Brian Lotze"      # copyright holder and TpT seller
BRAND = "Hudson Beat"      # store line on covers and previews (owner, 4 Oct)
YEAR = 2026
ACCENT = "#E8731A"         # this product's single accent: Halloween orange (tpt/STYLE.md)
TITLE = "Spooky Tunes"
PRICE = "12.00"
LOW_RES = "--lowres" in sys.argv
DISCLOSURE = "Designed with the help of digital and AI tools, and checked by hand."

# Strings that must never appear in any output or in any file in this folder (school branding, LMS, the owner's
# third-party classroom video). Split so that this file itself never contains them.
BANNED = ["Blind" + " Brook", "BB" + "HS", "blind" + "brook", "Google" + " Classroom", "Walking Through" + " The Woods", "Lotze" + ".mov"]

# ---------------------------------------------------------------- timing: 32 measures of 4/4 at 128 BPM = 60 seconds
BPM = 128
MEASURES = 32
SPM = 4 * 60 / BPM          # seconds per measure = 1.875
KEY_CHANGE = 17             # worked example only: its optional move from A minor to D minor lands on the visual cut at m. 17


def tfmt(sec):
    m, s = int(sec // 60), round(sec - 60 * int(sec // 60), 1)
    return f"{m}:{int(s):02d}" + (f".{int(round((s - int(s)) * 10))}" if s != int(s) else "")


def m_start(m):
    return (m - 1) * SPM


def measure_at(sec):
    return int(sec // SPM) + 1


def bpm_for(seconds):
    return MEASURES * 4 * 60 / seconds


CUE_PRACTICE = [(30, 17), (45, 25), (12, 7)]                    # (seconds, measure) on the cue sheet; key has answers
BPM_PRACTICE = [(60, 128), (64, 120), (75, 102)]                # (video seconds, BPM rounded) on the cue sheet

# ---------------------------------------------------------------- minor keys (natural minor), spelled with ASCII # and b
MINOR_STEPS = [2, 1, 2, 2, 1, 2, 2]                             # W H W W H W W
KEYS = {
    "A minor": ["A", "B", "C", "D", "E", "F", "G"],
    "E minor": ["E", "F#", "G", "A", "B", "C", "D"],
    "B minor": ["B", "C#", "D", "E", "F#", "G", "A"],
    "D minor": ["D", "E", "F", "G", "A", "Bb", "C"],
    "G minor": ["G", "A", "Bb", "C", "D", "Eb", "F"],
    "C minor": ["C", "D", "Eb", "F", "G", "Ab", "Bb"],
    "F minor": ["F", "G", "Ab", "Bb", "C", "Db", "Eb"],
    "Bb minor": ["Bb", "C", "Db", "Eb", "F", "Gb", "Ab"],
}
KEY_NOTES = {"A minor": "white keys only", "E minor": "one sharp (F#)", "B minor": "two sharps (F#, C#)",
             "D minor": "one flat (Bb)", "G minor": "two flats (Bb, Eb)", "C minor": "three flats (Bb, Eb, Ab)",
             "F minor": "four flats (Bb, Eb, Ab, Db)", "Bb minor": "five flats; a half step above A minor"}
KEY1, KEY2 = "A minor", "D minor"                               # the worked example's (optional) keys
PAIRS = [("A minor", "D minor", "Smooth: they share six of seven notes, so the change feels like the path turning."),
         ("A minor", "Bb minor", "A lurch: everything moves up a half step. Use it at the scariest moment."),
         ("E minor", "A minor", "Darker and lower: a good choice when the video slows down or goes still.")]
PC = {"C": 0, "D": 2, "E": 4, "F": 5, "G": 7, "A": 9, "B": 11}


def pc(name):
    return (PC[name[0]] + name.count("#") - name.count("b")) % 12


def chord(key):
    s = KEYS[key]
    return [s[0], s[2], s[4]]


# ---------------------------------------------------------------- worked example: a 32-measure score for a night-woods clip
# (name, source kind, source, notes, reverb send dB or None, delay send dB or None, why, regions [(first, last) measures])
TRACKS = [
    ("Low drone", "library", "Logic library synth pad", "A, then D at m. 17", -12, None,
     "The floor of the piece. It moves to D on the cut at measure 17, so the change of key is heard.", [(1, 15), (17, 32)]),
    ("Wind", "download", "Free sound-effects download", "no key", -6, None,
     "Far away and very wet: it sounds like the whole forest.", [(1, 15), (21, 32)]),
    ("Footsteps", "recording", "My recording: crumpled paper close to the mic", "no key", -30, None,
     "Close and almost dry: right next to the listener.", [(5, 15), (17, 24)]),
    ("Music box", "library", "Logic library (celesta or music box)", "A minor scale, then D minor", -14, -18,
     "The only melody, and the main delay send: it echoes through the trees.", [(9, 15), (17, 28)]),
    ("Low hum", "recording", "My own voice, pitched down", "E and F, a half step that rubs", -8, None,
     "Unsettling texture. The pitch-bend wheel slides it between the two notes.", [(25, 32)]),
    ("Impact", "library", "Logic library hit", "D (one hit)", -10, -22,
     "One hit at the scariest moment, with a long tail that rings to the end.", [(29, 32)]),
]
SILENCE = [16]
MOMENTS = [  # (measure, what happens) on the worked example's score map
    (1, "Drone and wind fade in"),
    (5, "Footsteps enter, close and dry"),
    (9, "Music box melody in A minor, with delay"),
    (16, "Everything stops: one measure of silence"),
    (17, "Visual cut: the drone moves to D (a change of key)"),
    (25, "Low hum enters; the reverb send swells up"),
    (29, "Impact hit; its reverb tail rings to the end"),
]

# ---------------------------------------------------------------- the owner's words (lightly edited: school, LMS, gear, Restructure additions; no key required)
GOALS = ["Understand what a bus is and why it exists",
         "Set up a reverb bus and a delay bus in Logic using aux channels and sends",
         "Use reverb and delay to create a sense of space, distance, and atmosphere",
         "Find, record, and import sounds from multiple sources: the Logic library, your own microphone recordings, and a free sound-effects site",
         "Compose music that serves a visual, making deliberate decisions about what happens when in the video",
         "Score a one-minute video with no structural restrictions: let the video drive the decisions"]
VOCAB = [("Bus", "a shared channel that multiple tracks send their audio to; effects on the bus apply to everything that sends to it"),
         ("Aux channel", "the channel in Logic that acts as the bus destination; this is where you put the reverb or delay plugin"),
         ("Send", "a control on a track that routes some of its signal to a bus; turning it up sends more signal to the effect"),
         ("Dry signal", "the original, unaffected sound: no reverb, no delay"),
         ("Wet signal", "the processed sound, with reverb or delay applied"),
         ("Reverb", "a simulation of acoustic space; makes sounds feel like they exist in a room, hall, cave, or open air"),
         ("Decay", "how long the reverb tail lasts after the original sound ends; longer decay = larger-feeling space"),
         ("Delay", "an effect that repeats a sound after a set amount of time, like an echo"),
         ("Feedback", "in a delay, how many times the echo repeats before fading out; high feedback = many repeats"),
         ("SFX", "sound effects: non-musical sounds used to create atmosphere, texture, or narrative"),
         ("Atonal", "music that does not follow a key or scale; notes chosen for texture and feeling rather than harmonic function"),
         ("Minor key", "the set of notes in a minor scale; it usually sounds darker than a major key"),
         ("Half step", "the smallest distance on the keyboard, from one key to the very next; two notes a half step apart rub against each other"),
         ("Scoring to picture", "writing music for a video, with sounds placed at specific moments on screen")]
REQUIREMENTS = ["Reverb bus set up correctly in Logic: aux channel, Bus 1, Space Designer at 100% wet",
                "Delay bus set up correctly in Logic: aux channel, Bus 2, Tape Delay or Echo at 100% wet",
                "At least 3 tracks sending to the reverb bus",
                "At least 1 track sending to the delay bus",
                "Minimum 5 distinct sound sources (a combination of the Logic library, your own recording, and a free download)",
                "At least 1 original recording made with a microphone",
                "At least 1 sound downloaded from a free sound-effects site (such as pixabay.com) and imported into Logic",
                "Music composed to the full length of the video: 32 measures, about 1 minute",
                "No empty silence longer than 5 seconds unless it is a deliberate compositional choice",
                "Project named: Your Name – Spooky Tunes",
                "Exported as MP3 and .zip (compressed .logicx file) and submitted to your class page before presenting",
                "A short written reflection: two specific techniques you used to create atmosphere, and why you chose them"]
SOURCES = [("Logic Pro library", "Open the Loop Browser (O) or the Library. Search by keyword: try “dark”, “tension”, “drone”, “cinematic”, “wind”, “impact”. "
            "You are not limited to loops: individual one-shots and samples work well for SFX layers."),
           ("Your own recording", "Record your own sounds in the room. Scratch a surface. Close a door slowly. Tap something hollow. Crumple paper. Hum. "
            "Any recorded sound can become something unsettling with enough processing. Drag the recording into Logic and apply effects."),
           ("A free download", "Go to a free sound-effects site your teacher approves (for example pixabay.com › Sound Effects) and search for atmospheric "
            "or horror sounds. Download as MP3, then drag the file from your Downloads folder into the Tracks area to create a new audio track. "
            "Check the site’s license.")]
PICTURE = [("Watch before you build", "Play the video twice without touching Logic. What moments stand out? Where does the energy change? What would a sound do "
            "to that moment? Write notes on your cue sheet before you open a single plugin."),
           ("Mark your moments", "Add markers in Logic at the timestamps where you want something to happen. A footstep. A sudden silence. A swell. "
            "Having the map in advance stops you from improvising past the moments that matter."),
           ("Silence is a tool", "If nothing is playing for a moment and then a sound enters, the entry feels enormous. Silence is not dead air: it is tension. "
            "Use it deliberately."),
           ("Let the video breathe", "You do not need to fill every second with sound. Sparse music over a quiet visual can be more unsettling than a wall "
            "of noise. Less is often more here."),
           ("Match the energy", "If the video gets more intense, your music should respond. Build toward moments of tension and release from them. "
            "The reverb and delay buses will help: a swell of wet reverb signals that something is coming.")]
STEPS = [  # (days, heading, text)
    ("Day 1", "Watch the video", "Watch it twice without doing anything. What moments stand out? Where does the energy shift? Make notes on your cue sheet "
                                 "before touching Logic."),
    ("Days 1–2", "Set up your buses", "Follow the How to Bus pages to set up your reverb bus (Bus 1, Space Designer) and delay bus (Bus 2, Tape Delay). "
                                     "Get both working before you start collecting sounds."),
    ("Day 2", "Find and import sounds", "Browse the Logic library, record at least one original sound, and download at least one free sound. "
                                        "Drag everything into Logic. Do not worry about arrangement yet."),
    ("Day 3", "Experiment with effects", "Before you score the video, spend time just playing. Send sounds to the reverb and delay buses at different levels. "
                                         "What does a heavy reverb do to a footstep? What does a long delay do to a single note? Let yourself be surprised."),
    ("Days 3–4", "Score the video", "Import the video and add markers at your key moments. Place sounds against the video. Adjust send levels so the "
                                    "wet/dry balance serves the mood. Build toward tension, release from it."),
    ("Day 5", "Mix, reflect and submit", "Play the full piece with the video. Does the music serve what is on screen? Adjust volumes. Make sure nothing is "
                                         "clipping. Write your reflection. Export MP3 and .zip and submit before presenting."),
]
TIPS = ["Set up the buses first. Do not start collecting sounds until both buses are working. If you build 10 tracks and then realize your bus routing is "
        "wrong, you have to fix everything.",
        "Record something weird. The microphone is there: use it. A recording of your own voice processed beyond recognition is more interesting than any preset.",
        "Want it darker? Any key, or no key at all, is fine, but a minor key sounds darker than major: try one (the Minor Key Map shows how). Half steps "
        "are spooky too: try two notes a half step apart, or slide between them with the pitch-bend wheel.",
        "Atonal does not mean random. Even without a key, your sounds should feel like they belong together. Choose notes for texture and tension, not "
        "just because they sound unpleasant.",
        "Watch the video with your music at least once per working day. It is easy to lose track of whether you are actually serving the picture or just "
        "making a piece of music that happens to be playing at the same time.",
        "Keep the How to Bus pages open. If your routing stops working, the troubleshooting table will fix almost every common problem."]
CHECKLIST = ["Reverb bus set up correctly (Aux, Bus 1, Space Designer at 100% wet)",
             "Delay bus set up correctly (Aux, Bus 2, Tape Delay or Echo at 100% wet)",
             "At least 3 tracks sending to the reverb bus",
             "At least 1 track sending to the delay bus",
             "Minimum 5 distinct sound sources used",
             "At least 1 original microphone recording included",
             "At least 1 free downloaded sound imported",
             "Music composed to the full length of the video (32 measures, about 1 minute)",
             "Project named: Your Name – Spooky Tunes",
             "Exported as MP3 and .zip (compressed .logicx) and submitted before presenting",
             "Reflection written: two techniques and why",
             "Three feedback questions posted to the class feedback board before presenting"]
RUBRIC = [  # owner's rubric (May project sheet), levels 2-5, 5 categories = /25; no key criteria (owner decision, 12 Oct)
    ("Sound Design & Atmosphere",
     ["Fewer than 3 sound sources used. Sounds feel generic or unrelated to the visual. No sense of atmosphere.",
      "At least 3 sound sources used but choices feel arbitrary. Some atmospheric quality but inconsistent.",
      "At least 5 sound sources including an original recording and a free download. Sounds contribute to a consistent spooky atmosphere.",
      "Sound choices are intentional and varied. The combination of sources creates a genuinely unsettling atmosphere that serves the video at every moment."]),
    ("Composition & Pacing",
     ["Music does not respond to the video. Sounds placed arbitrarily. Long stretches of silence or constant unvarying noise.",
      "Some response to the video but key moments are missed or poorly timed. Music feels like it is playing alongside the video rather than for it.",
      "Music responds to the video at most key moments. Energy builds and releases appropriately. No unintentional dead air.",
      "Music serves the picture precisely. Key moments are acknowledged. Silence is used as a compositional tool. The piece feels like it was written "
      "for this specific video."]),
    ("Technical Skill / DAW Use",
     ["Bus routing not set up or set up incorrectly. Space Designer and Tape Delay not at 100% wet. Sends not used.",
      "Bus routing partially correct. At least one bus working but setup has errors. Some sends used but inconsistently.",
      "Both buses set up correctly. Sends used on multiple tracks. Wet/dry balance adjusted by ear.",
      "Both buses set up correctly and used with intention. Send levels vary between tracks to create spatial depth. Workflow is organized and confident."]),
    ("Reverb & Delay as Mix",
     ["Effects not used to create space or depth. Everything either fully dry or drenched indiscriminately.",
      "Effects present but applied uniformly. No sense of near/far or spatial relationship between elements.",
      "Reverb and delay used to place sounds at different perceived distances. Mix has some sense of depth and space.",
      "Effects used to create a convincing sense of physical space. Near elements are dry, distant elements are wet. Reverb and delay contribute to the "
      "atmosphere rather than just processing the sound."]),
    ("Documentation & Presentation",
     ["Project not submitted or not presented.",
      "Project submitted late or named incorrectly. Presentation made with minimal engagement.",
      "Project named correctly and submitted. Presentation made. Feedback questions posted. Reflection written.",
      "Project correctly named, submitted on time. Student presents with confidence and can explain their sound choices and how the effects serve the "
      "atmosphere. Three thoughtful feedback questions posted."]),
]
LEVELS = [2, 3, 4, 5]

DAYS = [  # 6 days, 1/3/2 (owner's DMP1 Restructure); minute blocks for ~53-minute periods, from the owner's Teacher Notes pacing guide
    {"n": 1, "kind": "Learning", "title": "Space is a sound: the bus demo", "slides": "Intro deck slides 1–9",
     "goal": "Every student has a working reverb bus by the end of class. Delay bus setup begins on Day 2.",
     "blocks": [("0–5", "The video", "Play the woods video for the class. No comment. Then ask: what would you expect to hear if this were a film?"),
                ("5–15", "Intro deck", "The hook (slides 1–3: play two spooky themes from YouTube or Spotify), what makes music spooky, dry vs. wet, "
                                       "and what a bus is (slides 4–6)."),
                ("15–35", "Bus demo", "Teacher-led, students watch with computers closed for building (the 8-step script on page DEMOPAGE, slides 7–9). "
                                      "End by having one student set up the reverb bus from scratch while the class watches."),
                ("35–50", "Reverb bus", "Students open the How to Bus pages and set up their own reverb bus, replicating the demo. Check each setup "
                                        "before anyone moves on."),
                ("50–53", "Exit check", "The one thing to remember: your track still plays dry; the bus adds the wet; the send controls how much.")]},
    {"n": 2, "kind": "Work day", "title": "Delay bus, sounds, moments", "slides": "Slides 10–11, cue sheet, key map (optional)",
     "goal": "Both buses working, at least 5 sounds imported, the video’s moments on the cue sheet.",
     "blocks": [("0–10", "Watch again", "Play the video again as a class. What moments stood out? Where does the energy shift? Students note 2–3 moments "
                                        "on the cue sheet."),
                ("10–22", "Delay bus", "Students complete the delay bus. Walk around and check every student’s Space Designer settings: confirm 100% wet "
                                       "before moving on."),
                ("22–30", "Tempo + keys (optional)", "Set the project tempo (128 BPM for a 60-second clip). Slide 10: no key is required; students may "
                                                     "work in any key or none. Anyone who wants a darker harmonic home can pick a minor key on the Minor Key Map."),
                ("30–53", "Sound exploration", "Students browse the Logic library, record something with the microphone, and download a free sound. "
                                               "Goal: at least 5 sounds in the project before any composition decisions.")]},
    {"n": 3, "kind": "Work day", "title": "Experiment, then start scoring", "slides": "Slide 11, worked example",
     "goal": "First draft started: the video imported, markers placed, at least half the video (measures 1–16) scored.",
     "blocks": [("0–5", "Near or far?", "Mini-lesson: one sound, send low then high. Which feels closer? Some sounds close and dry, some distant and wet."),
                ("5–15", "Video + markers", "Import the video, set the tempo, add markers at the moments from the cue sheet."),
                ("15–50", "Experiment + compose", "Students experiment with the effects and begin placing sounds against the video. "
                                                  "Stuck? ‘Send one sound to the reverb bus at maximum. Now pull it back slowly. Where does it start to feel "
                                                  "right? That is your send level.’"),
                ("50–53", "Check-in", "Every student plays you measures 1–16 with the video.")]},
    {"n": 4, "kind": "Work day", "title": "Composition: make the moments land", "slides": "Project sheet",
     "goal": "Full video scored, rough mix in place.",
     "blocks": [("0–5", "Make a moment land", "Mini-lesson: pick the biggest moment in the video and make the music change right there: a new sound, "
                                             "a drop to silence, a reverb swell. Students working in a key can move the drone or bass to a new root."),
                ("5–48", "Compose", "Students continue composing to the video. Circulate and ask every student to play their piece with the video. "
                                    "Does the music respond to what is on screen? If it feels like background noise, push them to acknowledge specific moments."),
                ("48–53", "Rough mix", "Pull every fader back a little; nothing in the red. Save a copy.")]},
    {"n": 5, "kind": "Finish + present", "title": "Mix, reflect, submit", "slides": "Checklist + rubric",
     "goal": "All projects submitted with a reflection. Early presentations if time allows.",
     "blocks": [("0–30", "Final mix", "Students finalize their mix. Reverb and delay levels adjusted. Nothing clipping."),
                ("30–37", "Reflection", "Two specific techniques used to create atmosphere, and why (page REFLPAGE)."),
                ("37–43", "Export + submit", "Bounce an MP3, compress the .logicx, submit both. Optional: export the audio into a copy of the movie."),
                ("43–53", "Early presentations", "Volunteers play their piece with the video.")]},
    {"n": 6, "kind": "Presenting", "title": "Watch and listen together", "slides": "Rubric",
     "goal": "Every student plays their score with the video and explains one sound.",
     "blocks": [("0–5", "Feedback questions", "Students post three feedback questions to the class board before class begins."),
                ("5–50", "Presentations", "Each student plays their piece with the video. Classmates watch both together. Ask each presenter: pick one "
                                          "sound in your piece and explain why you chose it and what you did to it. About 2–3 minutes each."),
                ("50–53", "Close", "‘Scoring to picture is one of the most in-demand skills in music production. Film, TV, games, ads: all of it is this.’")]},
]

STRUGGLES = [
    ("Space Designer not set to 100% wet", "The single most common error. If Space Designer is outputting any dry signal, students will hear double: the dry "
     "signal from the track AND the dry signal from the bus. The bus should be 100% wet. Walk around on Day 2 and check every student’s Space Designer "
     "settings before anything else."),
    ("Wrong bus number on the Aux input", "Students create an Aux channel by hand but forget to set its input to a bus, or set it to the wrong bus number. "
     "Show them explicitly: the Aux input must match the bus number on the send. Bus 1 sends to Bus 1. If they do not match, nothing happens. Creating "
     "the bus from a track’s Send slot (as the How to Bus pages do) avoids most of this."),
    ("Not hearing any reverb at all", "Usually means the Aux channel fader is all the way down, or the send level is at its lowest. Check both. The send "
     "knob needs to be turned up and the Aux fader needs to be at a reasonable level."),
    ("Blank page paralysis on the creative side", "Some students freeze without harmonic or structural constraints. The redirect: download three sounds "
     "you find unsettling, import them into Logic, and just play. You do not need a plan yet. The plan emerges from the sounds."),
    ("Everything drenched in reverb equally", "Students discover reverb and apply maximum amounts to every track. The result is an undifferentiated wash. "
     "Prompt them: some sounds should feel close and dry, some should feel distant and wet. The contrast is what creates space."),
    ("The big moment nobody hears", "Students change something at a big moment in the video (a new melody, a new key) but leave the drone or bass "
     "untouched, so nothing seems to change. Have them change the lowest sound at the same moment (move it, drop it out, or let silence in), and place "
     "that moment on a cut or a shift in the video."),
]
DEMO = [  # owner's Day 1 bus demo, 8 steps (lightly edited)
    "Open a new Logic project. Add 3 or 4 audio or instrument tracks with different sounds playing.",
    "Play the project. Ask: what if I want all of these to sound like they are in the same room? I could add reverb to each track individually, but then "
    "I have three separate reverbs to manage. Show this briefly: add Space Designer to one track, then copy it to another. Ask: what is the problem here? "
    "(Any change to one reverb does not affect the other.)",
    "Delete the individual reverbs. Open the Mixer (X). On the first track, click an empty Send slot and choose Bus 1: Logic creates an Aux channel. "
    "Rename it Reverb Bus.",
    "Check that the Aux input reads Bus 1. Add Space Designer to the Aux insert. Set it to 100% wet (Dry 0%).",
    "Turn that track’s send up. Play the project. Ask: what do you hear? What changed?",
    "Add sends to the other tracks at different levels. Ask: what happens when one track has a high send and another has a low send? (Different apparent "
    "distance.) This is the moment the concept lands.",
    "Ask: what would happen if I wanted a delay too? Walk through setting up Bus 2 with Tape Delay. The second setup goes faster: students now understand the pattern.",
    "Open the project on a student computer and ask a student to set up the reverb bus from scratch while the class watches. This is the most effective "
    "check for understanding.",
]
BUS_PART1 = [  # owner's How to Bus, Part 1 (send levels changed to realistic dB values)
    ("Open the channel strip for your track", "In the Mixer (X) or the inspector on the left, find the channel strip for the track you want to add reverb to. "
     "Look for the Sends section: it sits below the insert plugin slots.", "If you cannot see Sends, widen the inspector or the Mixer so the Sends slots show."),
    ("Click an empty Send slot and choose a bus", "In the Sends section, click one of the empty slots. A menu appears. Go to Bus and choose Bus 1 (any bus "
     "that shows no name next to it is empty).", "Watch what happens: Logic automatically creates an Aux channel strip routed to that bus. You do not need "
     "to create it manually."),
    ("Raise the Send level", "A small knob now appears in the send slot. At its lowest, nothing is going to the bus. Click and drag it upward. Start around –15 dB.",
     "This is the most critical step. A send knob at its lowest sends nothing. Nothing will happen until you raise it."),
    ("Add Space Designer to the Aux channel", "Find the Aux channel Logic just created in the Mixer. Click its first empty Insert slot and choose "
     "Reverb › Space Designer. For Spooky Tunes, try a large hall, cathedral or other large-space preset.", "The Aux channel may appear at the far right "
     "of the Mixer. Scroll if you cannot find it."),
    ("Set Space Designer to 100% wet", "Inside Space Designer, set the Dry signal to 0% and the Wet signal to 100%. The bus outputs only the reverb. The dry "
     "signal always comes from the original track.", "This is the most common mistake. If the plugin is partially dry, you hear the sound doubled. "
     "Always 100% wet on a bus plugin."),
    ("Name the Aux channel", "Double-click the name at the bottom of the Aux channel strip and type REVERB BUS.",
     "You only set this up once. Every other track that needs reverb just needs a send pointing to the same bus number."),
    ("Send more tracks to the same bus", "Select another track. Click an empty Send slot. Choose the same bus number as before. Raise the send level. "
     "That track now shares the reverb.", "Do not pick a new bus number for each track. All reverb sends point to the same bus."),
]
BUS_PART2 = [
    ("Click a Send slot and choose a new bus", "On any track you want delay on, click an empty Send slot. Choose Bus › Bus 2 (or whichever bus is still "
     "empty, not the one used for reverb). Logic creates a second Aux channel.", "You can send a track to both the reverb bus and the delay bus at the same "
     "time by adding two separate sends."),
    ("Raise the Send level", "Drag the send knob up. Start lower than your reverb send: delay is a more noticeable effect. Around –20 to –25 dB is a good "
     "starting point.", "Delay is easy to overdo. Less is almost always better."),
    ("Add Tape Delay to the Aux channel", "Find the new Aux channel in the Mixer. Click its first Insert slot and choose Delay › Tape Delay. Set the plugin "
     "to 100% wet (Dry at 0%): same rule as the reverb bus.", "Try Echo for a cleaner sound or Stereo Delay for a wider effect."),
    ("Name the Aux channel", "Double-click the channel name and type DELAY BUS.", "You now have two clearly named effect buses."),
    ("Adjust the delay time", "In Tape Delay, the delay time controls the gap between the original and its echo. For atmosphere, try longer times "
     "(600–900 ms) with higher feedback. For subtle movement, try shorter times (100–250 ms) with low feedback.",
     "Turn off tempo sync if you want the delay to float freely. Free-running delay often feels more unsettling."),
    ("Balance reverb and delay together", "Play the full project and listen to both effects. Reverb creates space. Delay creates movement. Use the fader "
     "on each Aux to balance the total effect amount. You can even send the delay bus to the reverb bus by adding a send on the delay’s Aux channel.",
     "A common approach: reverb on most things, delay on just one or two specific elements. Selective delay draws the ear."),
]
TROUBLE = [("I hear reverb, but the whole track sounds wet, with no clear dry sound", "Space Designer is not set to 100% wet. Open it and set Dry to 0%, Wet to 100%."),
           ("I turned up the send but I can’t hear any reverb", "Check that the Aux input is set to the same bus number as the send (Bus 1 for reverb, "
            "Bus 2 for delay). Also check that the Aux fader is turned up."),
           ("Everything sounds like it’s drowning in reverb", "Your send levels are too high, or your Space Designer decay is too long. Pull the send "
            "knobs down (try –20 dB) and shorten the reverb tail. Or… did you put reverb on the Stereo Out by accident?"),
           ("The delay is syncing to tempo in a weird way", "Tape Delay has a tempo sync option. Turn it off and set the time manually in milliseconds "
            "if you want free-running delay."),
           ("I put reverb directly on a track instead of using the bus", "Delete the plugin from the track’s insert, set up the bus as above, and use a send instead.")]
ATMOS = [("Heavy reverb on a quiet sound", "A faint, distant sound buried in reverb feels enormous, like something very far away in a large, empty space. "
          "Try a single low note with a long reverb tail."),
         ("Delay on a single hit", "One sharp sound with a long delay feels like an echo in an empty corridor. High feedback makes it cascade and fade. "
          "Low feedback gives one clean echo."),
         ("Reverb on a recording", "Record something yourself (your voice, a tap on a surface, a door creak) and send it to the reverb bus. Real "
          "recordings with heavy processing sound genuinely unsettling."),
         ("Different send levels", "Not everything needs the same amount of reverb. A close, dry sound next to a distant, wet sound creates depth and a "
          "sense of physical space."),
         ("Automate the send level", "Turn the send up gradually over 8 measures and the reverb swells in. Pull it out suddenly and the space collapses. "
          "Both are effective in different moments of the video.")]
BUS_QUIZ = [("On a bus, Space Designer is set to ______ wet.", "100% (Dry at 0%)"),
            ("The Aux input must match ______.", "the bus number chosen on the send (Bus 1 for reverb, Bus 2 for delay)"),
            ("A send knob at its lowest sends ______ to the bus.", "nothing"),
            ("To make a sound feel far away, turn its reverb send ______. To make it feel close, ______.", "up; turn it down (keep it mostly dry)")]
VOCAB_QUIZ = [("A bus plugin at 100% wet outputs only the ______ signal.", "wet"),
              ("Longer reverb decay = a ______-feeling space.", "larger"),
              ("High delay feedback = ______ repeats.", "many"),
              ("In A minor, the half steps are B–C and ______.", "E–F")]

# ---------------------------------------------------------------- intro deck content (owner's "Looping in Logic" pattern)
LISTEN1 = ("Halloween (Main Title)", "John Carpenter")
LISTEN2 = ("Stranger Things (Main Theme)", "Kyle Dixon & Michael Stein")
SLIDES = [  # (kind, title, body, notes)
    ("title", "Spooky Tunes", "Feat. the Halloween theme",
     "Day 1, after you have played the woods video once with no comment. Tell students this is their first open-ended project: no chord progression, no song "
     "form, no right answer. Their brief is one minute of video and the job is to score it so the music makes the video feel the way it looks. Have the two "
     "listening examples cued up on YouTube or Spotify (your own account; no audio or video is included in this deck)."),
    ("listen", "Listen first", f"{LISTEN1[0]} · {LISTEN1[1]}|Teacher: play about 60 seconds from YouTube or Spotify.|What makes it scary?|"
     "What keeps repeating?|Does it feel bright or dark?",
     "Play roughly the first minute of the Halloween theme by John Carpenter. Ask the three questions and collect answers on the board, but do not confirm "
     "anything yet: feel it before you name it. Students usually notice the piano figure that keeps circling, that it never resolves, and that it is dark. "
     "Some will notice it is hard to count in 4: the piano pattern is famously counted in 5. Keep their words for slide 4."),
    ("listen", "Listen for the space", f"{LISTEN2[0]} · {LISTEN2[1]}|Teacher: play about 60 seconds from YouTube or Spotify.|"
     "Does it sound close, or far away?|Can you hear echoes or a ‘room’ around the sound?|What would it sound like totally dry?",
     "Play the opening of the Stranger Things theme by Kyle Dixon and Michael Stein. This time the question is about space rather than notes: does the sound "
     "feel close or far away, is there a room or echo around it? Ask a student to describe what it would sound like with no space at all. That is the "
     "difference between dry and wet, which they will name on slide 5."),
    ("bullets", "What makes music spooky?", "Minor keys: darker than major|Half steps: two notes that rub|"
     "Space: reverb and delay make it feel big and far away|Silence: the sound after it feels enormous",
     "Connect their answers from the listening to four tools. Minor keys sound darker (they met major and minor earlier in the course; if not, the "
     "Minor Key Map covers it). Half steps are the smallest distance on the keyboard: play E and F together on the projector so they hear the rub. "
     "Space comes from reverb and delay. Silence is a sound too: a sound after silence feels enormous. These are tools, not rules: no key is required, "
     "and a score with no key at all is fine."),
    ("columns", "Space is a sound", "Reverb: a room, hall or cave around the sound|Delay: an echo that repeats|Dry: the original sound|Wet: the effect only",
     "Plain-English definitions. Reverb simulates a space; decay is how long the tail lasts. Delay repeats the sound after a set time; feedback is how many "
     "repeats. Dry is the untouched sound; wet is only the effect. If you can, play one short sound in Logic dry, then with heavy reverb, then with a long "
     "delay. The unit's mixing concept is reverb and delay via a bus, and its big idea is the line on this slide: space is a sound."),
    ("bus", "What the heck is a bus?", "",
     "A bus is a shared channel. Instead of a reverb plugin on every track, you build one reverb on an Aux channel and each track sends some of its sound "
     "to it. Walk the diagram: each track still plays dry straight to the Stereo Out; the send knob decides how much of it also goes to Bus 1; the Aux "
     "adds only the wet reverb. Three reasons: consistency (one shared space), control (one fader for all the reverb) and efficiency (one plugin, less CPU)."),
    ("bullets", "Near or far?", "Same footstep: reverb send at –30 dB, then at –6 dB|Which one sounds closer?|Which one sounds farther away?|"
     "Why not just put the reverb on every track?",
     "Guided listening in Logic: play one footstep or tap with a low reverb send, then a high one. Students will say the low send sounds close and the high "
     "send sounds far away. That is the whole idea of depth in a mix: near things are dry, far things are wet. The last question sets up the demo: separate "
     "reverbs on every track can’t be controlled together and don’t sound like one shared space."),
    ("steps", "Watch me: the reverb bus", "Mixer (X): on a track, click an empty Send slot, choose Bus 1|Logic creates an Aux channel. Name it REVERB BUS|"
     "On the Aux insert: Reverb › Space Designer, a large-space preset|In Space Designer: Dry 0%, Wet 100%|Raise the track’s send (start around –15 dB), "
     "then send 2 more tracks",
     "This is the worked example in Logic; do it live while students watch with laptops closed for building (the full 8-step demo script is in the teacher "
     "notes). After step 5 play the project and ask what changed. Then give two other tracks different send levels and ask which sounds nearer. Finish by "
     "having one student rebuild the reverb bus from scratch on their own computer while the class watches."),
    ("steps", "Then the delay bus", "On a track, click an empty Send slot, choose Bus 2|On the new Aux: Delay › Tape Delay. Name it DELAY BUS|"
     "Tape Delay: Dry 0%, Wet 100%|Send lower than the reverb: try –20 to –25 dB|Long time (600–900 ms) + more feedback = eerie",
     "The second setup goes faster because students now understand the pattern. Point out that a track can have two sends, one to each bus. Suggest "
     "delay on only one or two sounds: selective delay draws the ear. Turning off tempo sync lets the echoes float freely, which often feels more "
     "unsettling. Students set up the delay bus at the start of Day 2."),
    ("keys", "Optional: try a minor key", "",
     "Day 2. No key is required: students may work in any key, or in none, and atonal or sound-effect scores are fine. For anyone who wants a darker "
     "harmonic home, a minor key helps. The examples are A minor (white keys only) and D minor (one flat, B-flat). The orange keys are the first chord: root, "
     "third and fifth (A C E, D F A). The dots are the scale. The Minor Key Map and the reference card (eight minor keys) are optional reference pages. "
     "A student who changes key at a big moment should move the drone or bass too, or nobody hears the change."),
    ("score", "Worked example: scoring the woods", "",
     "A score map for a 60-second clip at 128 BPM: 32 measures of 4/4 is exactly one minute, so every 4 measures is 7.5 seconds. Walk the numbered "
     "moments: the drone and wind fade in, close dry footsteps, a music-box melody with delay, a full measure of deliberate silence, then a visual cut at "
     "measure 17 where the drone moves (this example happens to change key there; that is a choice, not a rule), a swelling reverb send and one impact "
     "with a long tail. Your own clip’s moments will differ: that is what the cue sheet is for."),
    ("bullets", "Your project: Spooky Tunes", "Score a one-minute woods video|32 measures: any key, or none|A reverb bus and a delay bus|"
     "5+ sound sources: the Logic library, your own recording, a free download",
     "Hand out the project sheet and rubric. This is the brief: make it sound the way the video looks. Stress the three sound sources (at least one "
     "original microphone recording and one free download) and that both buses come before any composing. There is no chord progression, song form or "
     "key required: atonal and sound-effect scores are welcome."),
    ("days", "The steps", "",
     "Walk the six days: Day 1 the bus demo and a working reverb bus; Day 2 the delay bus and collecting at least 5 sounds; Days 3 and 4 importing the "
     "video, placing markers and composing to the whole video; Day 5 mix, reflection and submit; Day 6 presentations with the video. "
     "Students who need more time on routing get it on Day 2."),
    ("check", "What you turn in", "Reverb bus (Bus 1) and delay bus (Bus 2), both 100% wet|3+ tracks to reverb, 1+ to delay|"
     "Music for the whole video (32 measures)|5+ sources: 1+ recording, 1+ free download|MP3 + compressed .logicx, named Your Name – Spooky Tunes|"
     "Reflection + three feedback questions",
     "These are the deliverables from the submission checklist. Show where students submit (your class page) and where the feedback questions go (any "
     "class feedback board). Export: File › Bounce › Project or Section, MP3. Compress: in Finder, Control-click the .logicx project › Compress. "
     "The reflection names two specific techniques and why."),
    ("rubric", "How it’s scored", "",
     "Five categories, each scored 2 to 5, for a total out of 25. Read the 5 column aloud so students know what great looks like: an atmosphere that "
     "serves the video at every moment, a score where key moments are acknowledged and silence is a tool, buses used with intention, near sounds dry and far sounds wet, "
     "and a confident presentation that explains the sound choices."),
    ("closer", "One last thing…", "Space is a sound.|Silence is a tool.|Celebrate the strange ones.",
     "Close the launch with the unit’s big idea. This project gives students permission to be weird. The student who uses nothing but silence and a "
     "single pitch-bent note may be doing something more sophisticated than the student who layered eight spooky loops. Tell them that out loud."),
]


# ---------------------------------------------------------------- html scaffolding (Hudson Beat, from 006-008)
def font_css(base=BUILD):
    f = os.path.relpath(FONTS, base)
    faces = [("Oswald", 300, "normal", "oswald-latin-300-normal"), ("Oswald", 500, "normal", "oswald-latin-500-normal"),
             ("Oswald", 700, "normal", "oswald-latin-700-normal"),
             ("Source Sans 3", 400, "normal", "source-sans-3-latin-400-normal"),
             ("Source Sans 3", 400, "italic", "source-sans-3-latin-400-italic"),
             ("Source Sans 3", 600, "normal", "source-sans-3-latin-600-normal"),
             ("Source Sans 3", 700, "normal", "source-sans-3-latin-700-normal")]
    return "".join(f"@font-face{{font-family:'{n}';font-weight:{w};font-style:{s};src:url('{f}/{file}.woff2') format('woff2');}}"
                   for n, w, s, file in faces)


CSS = """
@page { size: 8.5in 11in; margin: 0; }
* { box-sizing: border-box; }
html, body { margin: 0; padding: 0; background: #fff; color: #000; }
body { font-family: 'Source Sans 3', sans-serif; -webkit-print-color-adjust: exact; print-color-adjust: exact; }
.page { width: 8.5in; height: 11in; position: relative; overflow: hidden; page-break-after: always; break-after: page;
        padding: 0.42in 0.5in 0.38in; display: flex; flex-direction: column; }
.page:last-child { page-break-after: auto; break-after: auto; }
.l { font-family: 'Oswald'; font-weight: 300; }
.top { display: flex; justify-content: space-between; align-items: flex-end; border-bottom: 2.2pt solid #000; padding-bottom: 6pt; }
.top .t { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; letter-spacing: .01em; line-height: 1; font-size: 25pt; }
.top .t small { display: block; font-family: 'Oswald'; font-weight: 300; text-transform: none; font-size: 12pt; letter-spacing: 0; margin-top: 4pt; }
.tag { text-align: right; }
.tag .lv { display: inline-block; background: #000; color: #fff; font-family: 'Oswald'; font-weight: 700; font-size: 14pt; padding: 3pt 9pt 4pt; border-radius: 4pt; letter-spacing: .04em; }
.tag .lv.ak { background: ACCENT; }
.tag .ln { font-family: 'Oswald'; font-weight: 300; font-size: 11pt; margin-top: 3pt; }
.name { display: flex; gap: 18pt; font-size: 12pt; margin: 9pt 0 4pt; }
.name span { flex: 1; border-bottom: 1pt solid #000; height: 15pt; }
.name span.d { flex: 0 0 1.6in; }
.main { flex: 1; display: flex; flex-direction: column; min-height: 0; padding-top: 6pt; }
.foot { display: flex; justify-content: space-between; font-size: 8.5pt; color: #333; border-top: .8pt solid #000; padding-top: 3pt; }
.wm { position: absolute; left: 50%; top: 50%; transform: translate(-50%, -50%) rotate(-38deg); font-family: 'Oswald'; font-weight: 700;
      font-size: 150pt; color: rgba(0,0,0,.13); letter-spacing: .06em; white-space: nowrap; pointer-events: none; z-index: 9; }
h2.s { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 13pt; margin: 9pt 0 3pt; letter-spacing: .02em; }
h2.s .acc, .acc { color: ACCENT; }
p, li { font-size: 10.5pt; line-height: 1.32; margin: 0 0 4pt; }
ul, ol { margin: 0 0 4pt; padding-left: 15pt; }
.tn p, .tn li { font-size: 11pt; line-height: 1.42; margin-bottom: 6pt; }
.tn.big p, .tn.big li { font-size: 12.2pt; line-height: 1.5; margin-bottom: 8pt; }
.tn.big h2.s { font-size: 14pt; margin-top: 12pt; }
.lines { flex: 1; min-height: 44pt; background: repeating-linear-gradient(to bottom, transparent 0, transparent 21pt, #000 21pt, #000 22pt); }
table.t { border-collapse: collapse; width: 100%; font-size: 9.8pt; margin: 2pt 0 4pt; }
table.t td, table.t th { border: 1pt solid #000; padding: 3pt 5pt; text-align: left; vertical-align: top; line-height: 1.25; }
table.t th { font-family: 'Oswald'; font-weight: 500; text-transform: uppercase; font-size: 9.5pt; background: #000; color: #fff; letter-spacing: .04em; vertical-align: middle; }
.cols { display: flex; gap: 20pt; }
.cols > div { flex: 1; min-width: 0; }
.small { font-size: 9pt !important; color: #333; }
.box { border: 1.4pt solid #000; border-radius: 8pt; padding: 7pt 10pt; }
.box.acc-b { border-color: ACCENT; border-width: 2pt; }
.box.inv { background: #000; color: #fff; border-color: #000; }
.k { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 10pt; letter-spacing: .05em; color: ACCENT; display: block; margin-bottom: 2pt; }
.ln2 { border-bottom: 1pt solid #000; height: 22pt; }
.fl { display: flex; align-items: flex-end; gap: 6pt; font-size: 10.5pt; margin: 0 0 8pt; }
.fl b { font-weight: 600; white-space: nowrap; }
.fl span { flex: 1; border-bottom: 1pt solid #000; height: 14pt; }
.num { list-style: none; padding: 0; margin: 2pt 0 0; }
.num li { display: flex; gap: 10pt; margin: 0 0 9pt; }
.num .n { flex: 0 0 24pt; height: 24pt; border-radius: 12pt; background: #000; color: #fff; font-family: 'Oswald'; font-weight: 700; font-size: 13pt;
          display: flex; align-items: center; justify-content: center; }
.num .hd { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 12pt; letter-spacing: .02em; display: block; }
.why { font-style: italic; }
.blk { display: flex; gap: 9pt; margin: 0 0 6pt; font-size: 10.6pt; line-height: 1.32; }
.blk .tm { flex: 0 0 40pt; text-align: center; background: #000; color: #fff; font-family: 'Oswald'; font-weight: 700; font-size: 10pt;
           border-radius: 4pt; padding: 2pt 0 3pt; align-self: flex-start; }
.blk b.h { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 10.5pt; letter-spacing: .02em; }
.day { border: 1.4pt solid #000; border-radius: 8pt; padding: 7pt 10pt 3pt; margin-bottom: 8pt; }
.day .dh { display: flex; justify-content: space-between; align-items: baseline; border-bottom: 1pt solid #000; padding-bottom: 3pt; margin-bottom: 6pt; }
.day .dh b { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 14pt; }
.day .dh span { font-family: 'Oswald'; font-weight: 300; font-size: 11pt; }
.goal { font-size: 10.6pt; margin: 0 0 7pt; }
.chk { display: flex; gap: 8pt; align-items: flex-start; font-size: 11pt; margin: 0 0 8pt; line-height: 1.3; }
table.rub { border-collapse: collapse; width: 100%; }
table.rub td, table.rub th { border: 1.1pt solid #000; padding: 3pt 4pt; vertical-align: top; font-size: 8.6pt; line-height: 1.22; }
table.rub th { background: #000; color: #fff; font-family: 'Oswald'; font-weight: 500; font-size: 11pt; text-align: center; vertical-align: middle; }
table.rub td.c { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 10pt; width: 12%; vertical-align: middle; }
table.rub td.sc { width: 6.5%; }
.cover { align-items: center; justify-content: space-between; padding: 0.55in 0.6in 0.5in;
         background: linear-gradient(180deg, #fff 0%, #fff 55%, #ededed 100%); }
.store { font-family: 'Oswald'; font-weight: 500; font-size: 13pt; letter-spacing: .32em; text-transform: uppercase; }
.ctitle { text-align: center; }
.ctitle .a { white-space: nowrap; font-family: 'Oswald'; font-weight: 700; font-size: 30pt; line-height: 1; letter-spacing: .02em; text-transform: uppercase; }
.ctitle .b { white-space: nowrap; font-family: 'Oswald'; font-weight: 700; font-size: 84pt; line-height: .98; letter-spacing: .01em; text-transform: uppercase; }
.ctitle .c { font-family: 'Oswald'; font-weight: 300; font-size: 15pt; margin-top: 8pt; }
.cband { width: 100%; background: #000; color: #fff; text-align: center; font-family: 'Oswald'; font-weight: 500; font-size: 14pt;
         letter-spacing: .08em; text-transform: uppercase; padding: 8pt 0 9pt; border-radius: 4pt; }
.cband em { font-style: normal; color: ACCENT; }
.slot { display: inline-block; border: 1.2pt solid #000; border-radius: 4pt; height: 26pt; min-width: 44pt; margin-right: 5pt; vertical-align: middle; }
.lab { font-family: 'Oswald'; font-weight: 500; font-size: 9pt; text-transform: uppercase; letter-spacing: .05em; text-align: center; }
""".replace("ACCENT", ACCENT)

COPY = f"© {YEAR} {STORE} · {BRAND} · For single-classroom use"


def html_doc(pages, title="doc"):
    return (f'<!doctype html><html><head><meta charset="utf-8"><title>{escape(title)}</title>'
            f'<style>{font_css()}{CSS}</style></head><body>{"".join(pages)}</body></html>')


def foot(right):
    return f'<div class="foot"><span>{COPY}</span><span>{right}</span></div>'


def wm(on):
    return '<div class="wm">PREVIEW</div>' if on else ""


def page(sub, tag, tagsub, inner, n, watermark=False, name=False, ak=False, cls=""):
    top = (f'<div class="top"><div class="t">{escape(TITLE)}<small>{sub}</small></div>'
           f'<div class="tag"><div class="lv{" ak" if ak else ""}">{tag}</div><div class="ln">{tagsub}</div></div></div>')
    nm = '<div class="name">Name<span></span>Date<span class="d"></span></div>' if name else ""
    return (f'<section class="page {cls}">{wm(watermark)}{top}{nm}<div class="main">{inner}</div>{foot(f"Page {n}")}</section>')


def checkbox(size=12):
    return (f'<svg width="{size}pt" height="{size}pt" viewBox="0 0 {size} {size}" style="flex:0 0 auto;margin-top:1pt">'
            f'<rect x=".7" y=".7" width="{size - 1.4}" height="{size - 1.4}" rx="2" fill="#fff" stroke="#000" stroke-width="1.2"/></svg>')


def t_(x, y, s, size, weight=700, family="Oswald", anchor="middle", fill="#000", extra=""):
    fam = f"'{family}'" if " " in family else family
    return (f'<text x="{x:.1f}" y="{y:.1f}" font-family="{fam}" font-weight="{weight}" font-size="{size}" text-anchor="{anchor}" '
            f'fill="{fill}" {extra}>{escape(s)}</text>')


def svgwrap(vw, vh, body, width="100%"):
    return f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {vw} {vh}" width="{width}" style="display:block">{body}</svg>'


# ---------------------------------------------------------------- drawings (original, code-drawn: black + accent only)
def _branch(out, rng, x, y, ang, ln, wd, depth):
    x1, y1 = x + ln * math.sin(ang), y - ln * math.cos(ang)
    out.append(f'<line x1="{x:.1f}" y1="{y:.1f}" x2="{x1:.1f}" y2="{y1:.1f}" stroke="#000" stroke-width="{wd:.2f}" stroke-linecap="round"/>')
    if depth <= 0:
        return
    kids = 3 if (depth > 3 and rng.random() < 0.25) else 2
    for k in range(kids):
        spread = rng.uniform(0.28, 0.6) * (-1 if k == 0 else 1) if k < 2 else rng.uniform(-0.12, 0.12)
        _branch(out, rng, x1, y1, ang + spread + rng.uniform(-0.08, 0.08), ln * rng.uniform(0.66, 0.8), max(wd * 0.66, 0.9), depth - 1)


def woods_svg(vw=800, vh=420, width="100%", seed=11, trees=None, moon=(0.6, 0.34, 0.2)):
    """Hero art: bare trees under an orange moon; the ground line is a sound wave with fading delay echoes."""
    rng = random.Random(seed)
    base = vh - 58
    out = [f'<circle cx="{moon[0] * vw:.1f}" cy="{moon[1] * vh:.1f}" r="{moon[2] * vh:.1f}" fill="{ACCENT}"/>']
    trees = trees or [(0.07, 0.86, 0.1), (0.2, 0.6, -0.05), (0.38, 0.42, 0.06), (0.8, 0.72, -0.08), (0.94, 0.95, 0.05)]
    for fx, fh, lean in trees:
        h = fh * (base - 8)
        _branch(out, rng, fx * vw, base, lean, h * 0.34, max(3.0, h * 0.05), 7)
    # ground: a sound hit (accent) and its delay echoes, each quieter
    pts, bursts = [], [(0.04, 46, ACCENT), (0.27, 28, "#000"), (0.47, 17, "#000"), (0.64, 10, "#000"), (0.79, 6, "#000")]
    bw = 0.12 * vw
    out.append(f'<line x1="0" y1="{base}" x2="{vw}" y2="{base}" stroke="#000" stroke-width="3" stroke-linecap="round"/>')
    for fx, amp, col in bursts:
        x0 = fx * vw
        pts = []
        for i in range(161):
            t = i / 160
            env = math.exp(-3.2 * t) * min(1, t * 14)
            pts.append(f"{x0 + t * bw:.1f},{base - amp * env * math.sin(2 * math.pi * 11 * t):.1f}")
        out.append(f'<polyline points="{" ".join(pts)}" fill="none" stroke="{col}" stroke-width="3.2" stroke-linejoin="round" stroke-linecap="round"/>')
    out.append(f'<rect x="0" y="{base + 2}" width="{vw}" height="{vh - base}" fill="#fff"/>')
    out.append(f'<line x1="0" y1="{base}" x2="{vw}" y2="{base}" stroke="#000" stroke-width="3"/>')
    for fx, amp, col in bursts:   # redraw the lower half of each wave over the white ground
        x0 = fx * vw
        pts = [f"{x0 + i / 160 * bw:.1f},{base - amp * math.exp(-3.2 * i / 160) * min(1, i / 160 * 14) * math.sin(2 * math.pi * 11 * i / 160):.1f}"
               for i in range(161)]
        out.append(f'<polyline points="{" ".join(pts)}" fill="none" stroke="{col}" stroke-width="3.2" stroke-linejoin="round" stroke-linecap="round"/>')
    return svgwrap(vw, vh, "".join(out), width)


def bus_svg(width="100%"):
    """Signal-flow diagram: three tracks play dry to the Stereo Out and send part of their signal to a reverb bus (Aux, 100% wet)."""
    b = ['<defs><marker id="ar" viewBox="0 0 10 10" refX="8" refY="5" markerWidth="7" markerHeight="7" orient="auto-start-reverse">'
         '<path d="M0 0 L10 5 L0 10 z" fill="#000"/></marker>'
         f'<marker id="ara" viewBox="0 0 10 10" refX="8" refY="5" markerWidth="7" markerHeight="7" orient="auto-start-reverse">'
         f'<path d="M0 0 L10 5 L0 10 z" fill="{ACCENT}"/></marker></defs>']
    rail = 30
    b.append(t_(470, rail - 10, "DRY: every track still plays straight to the output", 17, 500, "Oswald"))
    b.append(f'<path d="M286 {rail} H792 V148" fill="none" stroke="#000" stroke-width="3.5" marker-end="url(#ar)"/>')
    for i, (name, lvl, angle) in enumerate([("Drone", "send –12 dB", 20), ("Footsteps", "send –30 dB", -70), ("Wind", "send –6 dB", 55)]):
        y = 62 + i * 108
        b.append(f'<path d="M250 {y + 22} H286 V{rail}" fill="none" stroke="#000" stroke-width="3.5"/>')
        b.append(f'<rect x="20" y="{y}" width="230" height="88" rx="10" fill="#fff" stroke="#000" stroke-width="3"/>')
        b.append(t_(38, y + 38, name.upper(), 25, 700, anchor="start"))
        b.append(t_(38, y + 68, lvl, 16, 600, "Source Sans 3", anchor="start"))
        cx, cy = 212, y + 50
        b.append(f'<circle cx="{cx}" cy="{cy}" r="19" fill="#fff" stroke="#000" stroke-width="3"/>')
        b.append(f'<line x1="{cx}" y1="{cy}" x2="{cx + 15 * math.sin(math.radians(angle)):.1f}" y2="{cy - 15 * math.cos(math.radians(angle)):.1f}" '
                 f'stroke="{ACCENT}" stroke-width="4.5" stroke-linecap="round"/>')
        b.append(f'<path d="M250 {y + 62} C 300 {y + 62}, 310 222, 356 222" fill="none" stroke="{ACCENT}" stroke-width="3.5" '
                 f'stroke-dasharray="9 6" marker-end="url(#ara)"/>')
    b.append(t_(312, 386, "SENDS to Bus 1", 16, 700, anchor="middle", fill=ACCENT))
    b.append('<rect x="360" y="142" width="262" height="160" rx="12" fill="#000"/>')
    b.append(t_(491, 182, "AUX · INPUT BUS 1", 22, 700, fill="#fff"))
    b.append(t_(491, 222, "REVERB BUS", 32, 700, fill="#fff"))
    b.append(t_(491, 256, "Space Designer", 19, 600, "Source Sans 3", fill="#fff"))
    b.append(t_(491, 284, "100% WET · 0% DRY", 19, 700, fill=ACCENT))
    b.append('<path d="M622 222 H694" fill="none" stroke="#000" stroke-width="3.5" marker-end="url(#ar)"/>')
    b.append(t_(658, 210, "WET", 16, 700))
    b.append('<rect x="700" y="150" width="185" height="118" rx="10" fill="#fff" stroke="#000" stroke-width="4"/>')
    b.append(t_(792, 202, "STEREO OUT", 26, 700))
    b.append(t_(792, 234, "what you hear:", 16, 400, "Source Sans 3"))
    b.append(t_(792, 254, "dry + wet together", 16, 600, "Source Sans 3"))
    return svgwrap(900, 400, "".join(b), width)


WHITE_PCS = [0, 2, 4, 5, 7, 9, 11]
BLACK_PCS = {1: 0.68, 3: 1.72, 6: 3.65, 8: 4.7, 10: 5.74}   # black key centre, in white-key widths from C


def keyboard_svg(key=None, width="100%", show_chord=True, show_scale=True, labels=True, octaves=2):
    """Two-octave keyboard from C. key=None draws a blank keyboard (worksheet). Chord tones: accent + name; other scale notes: dots."""
    ww, wh, bw, bh = 40, 150, 25, 94
    nwhite = 7 * octaves + 1
    scale = KEYS.get(key, [])
    names = {pc(n): n for n in scale}
    ch = {pc(n) for n in chord(key)} if key else set()
    out = []
    for i in range(nwhite):
        p = WHITE_PCS[i % 7]
        x = 2 + i * ww
        on_ch = show_chord and p in ch
        out.append(f'<rect x="{x}" y="2" width="{ww}" height="{wh}" rx="3" fill="{ACCENT if on_ch else "#fff"}" stroke="#000" stroke-width="2"/>')
        if p in names and (on_ch or show_scale):
            if on_ch and labels:
                out.append(t_(x + ww / 2, wh - 12, names[p], 19, 700, fill="#fff"))
            elif show_scale:
                out.append(f'<circle cx="{x + ww / 2}" cy="{wh - 22}" r="7" fill="#000"/>')
                if labels:
                    out.append(t_(x + ww / 2, wh - 38, names[p], 15, 500))
    for o in range(octaves + (1 if octaves else 0)):
        for p, cx in BLACK_PCS.items():
            x = 2 + (o * 7 + cx) * ww - bw / 2
            if x + bw > 2 + nwhite * ww:
                continue
            on_ch = show_chord and p in ch
            out.append(f'<rect x="{x:.1f}" y="2" width="{bw}" height="{bh}" rx="2" fill="{ACCENT if on_ch else "#000"}" stroke="#000" stroke-width="2"/>')
            if p in names and (on_ch or show_scale):
                if on_ch and labels:
                    out.append(t_(x + bw / 2, bh - 10, names[p], 13, 700, fill="#fff"))
                elif show_scale:
                    out.append(f'<circle cx="{x + bw / 2:.1f}" cy="{bh - 16}" r="6" fill="#fff"/>')
                    if labels:
                        out.append(t_(x + bw / 2, bh - 28, names[p], 11, 600, fill="#fff"))
    return svgwrap(nwhite * ww + 4, wh + 4, "".join(out), width)


def score_svg(blank=False, width="100%"):
    """32-measure score map at 128 BPM: time ruler, key band (example: A minor then D minor; blank: one optional bar), moments, track lanes."""
    lw, W = 128, 1000
    mw = (W - lw - 6) / MEASURES
    X = lambda m: lw + (m - 1) * mw
    lane_h = 34
    lanes = [("Track " + str(i + 1), []) for i in range(6)] if blank else [(t[0], t[7]) for t in TRACKS]
    y0 = 128
    H = y0 + len(lanes) * lane_h + 8
    b = []
    if not blank:
        for m in SILENCE:
            b.append(f'<rect x="{X(m):.1f}" y="40" width="{mw:.1f}" height="{H - 40}" fill="#e8e8e8"/>')
    for m in range(1, MEASURES + 2, 4):
        x = X(m)
        heavy = (1, MEASURES + 1) if blank else (1, KEY_CHANGE, MEASURES + 1)
        b.append(f'<line x1="{x:.1f}" y1="22" x2="{x:.1f}" y2="{H}" stroke="#000" stroke-width="{1.6 if m in heavy else .7}"/>')
        if m <= MEASURES:
            b.append(t_(x + 3, 14, f"m. {m}", 13, 700, anchor="start"))
            b.append(t_(x + 3, 32, tfmt(m_start(m)), 12, 400, "Source Sans 3", anchor="start"))
    b.append(t_(X(MEASURES + 1) - 2, 32, tfmt(m_start(MEASURES + 1)), 12, 400, "Source Sans 3", anchor="end"))
    for m in range(1, MEASURES + 1):
        b.append(f'<line x1="{X(m):.1f}" y1="{y0 - 6}" x2="{X(m):.1f}" y2="{y0}" stroke="#000" stroke-width=".7"/>')
    if blank:   # no key is required: one open bar, written in only if the student uses a key
        b.append(f'<rect x="{X(1) + 1:.1f}" y="42" width="{MEASURES * mw - 2:.1f}" height="30" rx="5" fill="#fff" stroke="#000" stroke-width="1.6"/>')
        b.append(t_(X(1) + 8, 63, "OPTIONAL · ANY KEY, OR NONE:", 15, 700, anchor="start"))
    else:       # the worked example happens to change key on the visual cut at measure 17
        b.append(f'<rect x="{X(1) + 1:.1f}" y="42" width="{(KEY_CHANGE - 1) * mw - 2:.1f}" height="30" rx="5" fill="#000"/>')
        b.append(f'<rect x="{X(KEY_CHANGE) + 1:.1f}" y="42" width="{(MEASURES - KEY_CHANGE + 1) * mw - 2:.1f}" height="30" rx="5" fill="{ACCENT}"/>')
        b.append(t_((X(1) + X(KEY_CHANGE)) / 2, 63, f"{KEY1.upper()}  ·  m. 1–{KEY_CHANGE - 1}", 16, 700, fill="#fff"))
        b.append(t_((X(KEY_CHANGE) + X(MEASURES + 1)) / 2, 63, f"{KEY2.upper()}  ·  m. {KEY_CHANGE}–{MEASURES}", 16, 700, fill="#fff"))
    b.append(t_(lw - 8, 63, "KEY", 14, 700, anchor="end"))
    b.append(t_(lw - 8, 100, "MOMENTS", 14, 700, anchor="end"))
    if not blank:
        for i, (m, _) in enumerate(MOMENTS):
            x = X(m) + mw / 2
            b.append(f'<line x1="{x:.1f}" y1="104" x2="{x:.1f}" y2="{H}" stroke="{ACCENT}" stroke-width="1.6" stroke-dasharray="4 3"/>')
            b.append(f'<circle cx="{x:.1f}" cy="94" r="12" fill="{ACCENT}"/>')
            b.append(t_(x, 99.5, str(i + 1), 15, 700, fill="#fff"))
    for r, (name, regions) in enumerate(lanes):
        y = y0 + r * lane_h
        b.append(f'<line x1="{lw}" y1="{y + lane_h:.1f}" x2="{X(MEASURES + 1):.1f}" y2="{y + lane_h:.1f}" stroke="#000" stroke-width=".5"/>')
        b.append(t_(lw - 8, y + 23, name, 14, 500, anchor="end"))
        for a, z in regions:
            b.append(f'<rect x="{X(a) + 1.5:.1f}" y="{y + 5}" width="{(z - a + 1) * mw - 3:.1f}" height="{lane_h - 10}" rx="5" fill="#000"/>')
    if not blank:
        b.append(t_(X(16) + mw / 2 + 4, H - 8, "silence", 12, 600, "Source Sans 3", anchor="start", extra=f'transform="rotate(-90 {X(16) + mw / 2 + 4:.1f} {H - 8})"'))
    b.append(f'<rect x="{lw}" y="{y0}" width="{X(MEASURES + 1) - lw:.1f}" height="{len(lanes) * lane_h}" fill="none" stroke="#000" stroke-width="1.6"/>')
    return svgwrap(W, H, "".join(b), width)


# ---------------------------------------------------------------- page plan
PLAN = ["cover", "howto", "overview", "setup", "logicref", "pacing", "day12", "day34", "day56",
        "tn1", "tn2", "tn3", "tn4",
        "post", "proj1", "proj2", "proj3", "proj4", "bus1", "bus2", "bus3", "bus4", "keys", "keyref", "cue", "vocab", "checklist", "rubric",
        "key", "example", "terms"]
IX = {k: i + 1 for i, k in enumerate(PLAN)}
PREVIEW_PICKS = ["pacing", "proj2", "bus2", "rubric"]
DECK_SLIDES = len(SLIDES)


def fill(s):
    return s.replace("DEMOPAGE", str(IX["tn3"])).replace("REFLPAGE", str(IX["proj4"]))


def H2(s, first=False):
    return f'<h2 class="s"{" style=\"margin-top:2pt\"" if first else ""}>{s}<span class="acc"> ·</span></h2>'


def cover_page():
    return (f'<section class="page cover"><div class="store">{BRAND}</div>'
            f'<div style="width:7.2in">{woods_svg()}</div>'
            f'<div class="ctitle"><div class="a">Scoring to Picture in Logic Pro</div><div class="b">Spooky Tunes</div>'
            f'<div class="c">Digital Music Production 1 · Unit 6 · reverb &amp; delay buses · sound design · intro deck</div></div>'
            f'<div class="cband">Grades 9–12 <em>·</em> 6 days <em>·</em> {len(PLAN)} pages + {DECK_SLIDES}-slide deck</div></section>')


VIDEO_BOX = ('<div class="box acc-b"><span class="k">No video is included: you supply one</span>'
             '<p style="margin:0 0 3pt">Students score a video of about one minute: a walk through dark woods at night, or any eerie scene. '
             'This resource includes no video or audio. Choose one of these:</p><ul style="margin:0">'
             '<li><b>Film your own.</b> A 60-second phone video walking a dark trail, a foggy field or an empty hallway with a flashlight. No faces needed.</li>'
             '<li><b>A free stock clip.</b> Royalty-free video sites (for example Pixabay or Pexels) have night-forest clips: search “forest night”, '
             '“dark woods” or “fog trees”. Read the site’s license and download one clip for the class.</li>'
             '<li><b>A short menu.</b> Approve two or three clips and let students choose.</li></ul>'
             '<p class="small" style="margin:3pt 0 0">Trim the clip to about 60 seconds first (QuickTime Player › Edit › Trim). At 128 BPM, '
             '32 measures is exactly 60 seconds.</p></div>')


def howto_page(n):
    parts = [("Teacher pages", "howto", "tn4", "How to use, overview, setup, Logic reference, pacing, day plans, teacher notes, bus demo script"),
             ("Student pages", "post", "rubric", "Day 1 post, project sheet, How to Bus guide, minor key map + reference, cue sheet, vocabulary, checklist, rubric"),
             ("Answer key + example", "key", "example", "Worksheet answers and a worked 32-measure score with send levels"),
             ("Intro deck (separate file)", None, None, f"{DECK_SLIDES} slides, .pptx + PDF, speaker notes on every slide")]
    trs = "".join(f"<tr><td><b>{a}</b></td><td style='white-space:nowrap'>{('pages ' + str(IX[f]) + '–' + str(IX[l])) if f else 'deck'}</td><td>{u}</td></tr>"
                  for a, f, l, u in parts)
    inner = (f'<div class="cols"><div>'
             f'{H2("What this is", True)}'
             f'<p>A complete 6-day scoring-to-picture project for high school music technology. Students score a one-minute video of a walk through the '
             f'woods at night in Logic Pro: 32 measures in any key or none, with reverb and delay on shared buses, built from the Logic library, their own '
             f'microphone recordings and free sound effects. They present their score with the video.</p>'
             f'<p>This is Unit 6 of my Digital Music Production 1 course, the first open-ended project of the year. It works just as well on its own, '
             f'and it fits October as well as November.</p>'
             f'{H2("Quick start")}<ol>'
             f'<li>Read the overview and pacing (pages {IX["overview"]}–{IX["pacing"]}) and the teacher notes (pages {IX["tn1"]}–{IX["tn4"]}).</li>'
             f'<li>Find your video (box on the right) and check the setup list (page {IX["setup"]}) on one student computer.</li>'
             f'<li>Practice the bus demo once (script on page {IX["tn3"]}, Logic steps on pages {IX["bus2"]}–{IX["bus3"]}).</li>'
             f'<li>Cue up the two listening examples on YouTube or Spotify: the Halloween theme (John Carpenter) and the Stranger Things theme '
             f'(Kyle Dixon &amp; Michael Stein). No audio is included.</li>'
             f'<li>Print or post the student pages (pages {IX["post"]}–{IX["rubric"]}).</li></ol>'
             f'{H2("The intro deck")}'
             f'<p>You present it on Day 1 (slides 1–9: hook, what makes music spooky, dry and wet, what a bus is, the reverb and delay bus in Logic) and '
             f'Day 2 (slides 10–16: minor keys as an option, a worked score, the project, steps, deliverables, rubric). Every slide has speaker notes.</p>'
             f'<p class="small">Fonts: the deck uses Oswald and Source Sans 3 (free on Google Fonts, built into Google Slides). If PowerPoint or Keynote '
             f'shows a different font, install the two fonts or present the PDF copy, which shows the intended look.</p>'
             f'</div><div>'
             f'{H2("What’s inside", True)}<table class="t"><tr><th>Part</th><th>Pages</th><th>Use</th></tr>{trs}</table>'
             f'{VIDEO_BOX}'
             f'{H2("Make it fit your class")}'
             f'<ul><li><b>Less time or less experience?</b> Pre-build the two buses in a template (page {IX["tn4"]}).</li>'
             f'<li><b>Minor keys are optional.</b> No key is required; for students who want one, the Minor Key Map and reference card (pages {IX["keys"]}–{IX["keyref"]}) teach it.</li>'
             f'<li><b>Feedback board:</b> I use Padlet; any discussion board, LMS forum or paper gallery walk works.</li></ul>'
             f'</div></div>')
    return page("How to use this resource", "START HERE", "Teacher", inner, n)


def overview_page(n):
    inner = (f'<div class="cols"><div>'
             f'{H2("What students make", True)}'
             f'<p>32 measures of atmospheric music scored to a one-minute video of a walk through the woods at night, with reverb and delay routed through '
             f'shared buses. No key, chord progression or song form is required: any key, or none.</p>'
             f'{H2("What students learn")}'
             f'<p>Minor keys feel different from major keys, and now they know why. Half steps create dissonance. Reverb creates space and distance. Pitch '
             f'bending creates unease. Delay creates echo and mystery. Music can tell a story without words. Less can be more: silence and space are creative choices.</p>'
             f'{H2("Why it lands here")}'
             f'<p>This is the first truly open-ended creative project, and it lands after students have chord vocabulary, melodic instincts, form awareness and '
             f'the beginning of mixing understanding. The video anchor removes the ‘what should I make?’ anxiety while still demanding genuine creative decisions.</p>'
             f'{H2("Mixing concept: reverb &amp; delay via bus")}'
             f'<p>Students need reverb to solve a creative problem: how do I make this sound scary? The bus/send workflow is taught as the right way to use '
             f'time-based effects: ‘Put the reverb on a bus, not directly on the track. That way you control how much of it you add.’</p>'
             f'<div class="box" style="margin-top:12pt"><span class="k">A note from me</span>'
             f'<p style="margin:0">This project gives students permission to be weird. Some of your most musically undeveloped students will surprise you, '
             f'because there is no wrong answer, only more or less effective choices. Celebrate the strange ones.</p></div>'
             f'</div><div>'
             f'{H2("Goals", True)}<ul>' + "".join(f"<li>{g}</li>" for g in GOALS) + '</ul>'
             f'{H2("At a glance")}<table class="t">'
             f'<tr><td><b>Grades</b></td><td>9–12 (no music reading needed)</td></tr>'
             f'<tr><td><b>Length</b></td><td>6 periods of about 53 minutes: 1 learning, 3 work, 2 finish and present</td></tr>'
             f'<tr><td><b>Type</b></td><td>Open-ended project, scoring to picture</td></tr>'
             f'<tr><td><b>Software</b></td><td>Logic Pro: sends, aux channels, Space Designer, Tape Delay, movie import, markers, automation</td></tr>'
             f'<tr><td><b>Product</b></td><td>A 1-minute, 32-measure score (any key, or none), MP3 + compressed project, a written reflection</td></tr>'
             f'<tr><td><b>Assessment</b></td><td>5-category rubric, /25 (page {IX["rubric"]})</td></tr></table>'
             f'{H2("Standards")}'
             f'<p>Supports the National Core Arts Standards for Music (Music Technology, high school): generating and developing musical ideas (MU:Cr1, '
             f'MU:Cr2), refining work through feedback (MU:Cr3), presenting (MU:Pr6) and connecting music to other media (MU:Cn11).</p>'
             f'</div></div>')
    return page("Unit overview", "OVERVIEW", "DMP1 · Unit 6", inner, n)


def setup_page(n):
    inner = (f'<div class="cols"><div>'
             f'{H2("You need", True)}<ul>'
             f'<li><b>A Mac with Logic Pro</b> for each student. Space Designer, Tape Delay and movie import are built in.</li>'
             f'<li><b>Headphones</b> for every student.</li>'
             f'<li><b>A way to record</b>: an audio interface and microphone (I use a Focusrite Scarlett Solo with a condenser mic), a USB microphone, '
             f'or in a pinch the Mac’s built-in mic. One or two shared stations are enough.</li>'
             f'<li><b>A one-minute video</b> (box on the right).</li>'
             f'<li><b>Internet access</b> to a free sound-effects site, or a folder of sounds you download in advance.</li>'
             f'<li><b>Projector and speakers</b> for the deck, the listening examples and presentations.</li>'
             f'<li><b>The two listening examples</b> from your own YouTube or Spotify access.</li>'
             f'<li><b>A place to submit</b> an MP3 and a .zip, and a <b>feedback board</b> for Day 6.</li></ul>'
             f'{H2("Optional")}<ul>'
             f'<li><b>A small MIDI keyboard with a pitch-bend wheel</b> (an MPK Mini or similar). Pitch bending is one of the spookiest tools in the unit.</li>'
             f'<li>A Logic template with the two buses already built, for students who need it.</li></ul>'
             f'</div><div>{VIDEO_BOX}'
             f'{H2("Before Day 1")}'
             + "".join(f'<div class="chk">{checkbox()}<span>{c}</span></div>' for c in [
                 "Choose, trim and test the video: open it in Logic on one student computer (File › Movie › Open Movie).",
                 "Practice the bus demo once (page " + str(IX["tn3"]) + "). Note 3–4 moments in the video worth pointing out.",
                 "Test the recording station: interface, mic, input level, headphones.",
                 "Check that students can reach your chosen sound-effects site, or download a folder of sounds.",
                 "Cue up the two listening examples; test the projector sound.",
                 "Print or post the student pages; set up the submission spot and the feedback board."]) +
             f'</div></div>')
    return page("Setup and requirements", "SETUP", "Teacher", inner, n)


LOGIC_REF = [("Import the video", "File › Movie › Open Movie. The movie plays in a floating window and a Movie track appears. Make sure the movie starts at bar 1. "
              "You can skip importing the movie’s own audio."),
             ("Tempo", "Double-click the tempo in the control bar and type 128 (32 measures = 60 seconds). Other lengths: BPM = 7680 ÷ seconds."),
             ("Markers", "Show the global tracks (G) to see the Marker track. Move the playhead to a moment and create a marker with Option-' "
              "(apostrophe) or the + in the Marker track. Double-click a marker to rename it."),
             ("Record a sound", "Create an audio track, choose your interface’s input, record-enable the track and press R. Keep the level out of the red."),
             ("Import a download", "Drag the MP3 from Finder into the Tracks area. Logic makes a new audio track."),
             ("Reverb bus", "On a track’s channel strip, click an empty Send slot › Bus › Bus 1. Logic creates an Aux. Put Space Designer on the Aux; Dry 0%, Wet 100%."),
             ("Delay bus", "Same steps with Bus 2 and Tape Delay (Dry 0%, Wet 100%)."),
             ("Send level", "Drag the small knob in the Send slot. Reverb: start around –15 dB. Delay: start around –20 to –25 dB."),
             ("Pitch bend", "Use the pitch-bend wheel or joystick on a MIDI keyboard while you record a software instrument part."),
             ("Automation", "Press A to show automation. Choose the send level in the track’s automation menu and draw points to swell the reverb in or out."),
             ("Mixer", "Press X. Aux channels sit to the right of the track channels."),
             ("Export", "File › Bounce › Project or Section (Command-B), MP3. In Finder, Control-click the .logicx project and choose Compress. "
              "Optional for presenting: File › Movie › Export Audio to Movie.")]


def logicref_page(n):
    rows = "".join(f"<tr><td style='width:20%'><b>{a}</b></td><td>{b}</td></tr>" for a, b in LOGIC_REF)
    inner = (f'<p>Everything students need in Logic for this unit, in one table. Students get the step-by-step version of the bus setup on the How to Bus '
             f'pages ({IX["bus1"]}–{IX["bus4"]}).</p>'
             f'<table class="t" style="font-size:11.2pt;line-height:1.38">{rows}</table>'
             f'<div class="cols" style="margin-top:6pt"><div class="box"><span class="k">Tempo math</span>'
             f'<p style="margin:0">32 measures × 4 beats = 128 beats. At 128 BPM that is exactly 60 seconds: 1 measure = 1.875 seconds and every 4 measures '
             f'= 7.5 seconds. For a 64-second clip use 120 BPM; for 75 seconds, about 102 BPM.</p></div>'
             f'<div class="box"><span class="k">Version note</span>'
             f'<p style="margin:0" class="small">Menu names are from current Logic Pro versions; older versions may differ slightly. Logic Pro, Space Designer and '
             f'Tape Delay are trademarks of Apple Inc., named only to describe what the unit needs. This resource is not affiliated with Apple.</p></div></div>')
    return page("Logic quick reference", "SETUP", "Teacher", inner, n)


def pacing_page(n, watermark=False):
    rows = "".join(
        f"<tr><td style='white-space:nowrap'><b class='acc' style='font-family:Oswald;font-size:13pt'>DAY {d['n']}</b><br><span class='small'>{d['kind']}</span></td>"
        f"<td><b>{d['title']}</b><br><span class='small'>{d['goal']}</span></td>"
        f"<td style='font-size:9.6pt;line-height:1.4'>{'<br>'.join(f'{a} · {h}' for a, h, _ in d['blocks'])}</td>"
        f"<td style='font-size:9.6pt'>{d['slides']}</td></tr>" for d in DAYS)
    inner = (f'<p>Six class periods of about 53 minutes: one learning day, three work days, then finishing and presenting. Each day’s full plan is on pages '
             f'{IX["day12"]}–{IX["day56"]}.</p>'
             f'<table class="t"><tr><th style="width:14%">Day</th><th style="width:38%">Focus and goal</th><th style="width:30%">Minute blocks</th><th>Show / use</th></tr>{rows}</table>'
             f'{H2("Student pages by day")}<table class="t">'
             f'<tr><td><b>Day 1</b></td><td>Day 1 post (p. {IX["post"]}), How to Bus (pp. {IX["bus1"]}–{IX["bus2"]})</td></tr>'
             f'<tr><td><b>Day 2</b></td><td>How to Bus Part 2 (p. {IX["bus3"]}), Minor Key Map and reference card (optional, pp. {IX["keys"]}–{IX["keyref"]}), '
             f'cue sheet (p. {IX["cue"]}), project sheet (pp. {IX["proj1"]}–{IX["proj3"]})</td></tr>'
             f'<tr><td><b>Days 3–4</b></td><td>Project sheet steps 4–5, troubleshooting (p. {IX["bus4"]}), vocabulary (p. {IX["vocab"]})</td></tr>'
             f'<tr><td><b>Day 5</b></td><td>Tips and reflection (p. {IX["proj4"]}), submission checklist (p. {IX["checklist"]})</td></tr>'
             f'<tr><td><b>Day 6</b></td><td>Rubric for self-score and teacher score (p. {IX["rubric"]})</td></tr></table>')
    return page("Pacing guide · 6 days (1 learning / 3 work / 2 finish + present)", "PACING", "Teacher", inner, n, watermark)


def day_box(d):
    blocks = "".join(f'<div class="blk"><div class="tm">{a}</div><div><b class="h">{h}</b> · {fill(t)}</div></div>' for a, h, t in d["blocks"])
    return (f'<div class="day"><div class="dh"><b>Day {d["n"]} · {d["title"]}</b><span>{d["kind"]} · {d["slides"]}</span></div>'
            f'<p class="goal"><b>Goal:</b> {d["goal"]}</p>{blocks}</div>')


def day12_page(n):
    return page("Day-by-day plans · Days 1 and 2", "DAYS 1–2", "Learn + work", day_box(DAYS[0]) + day_box(DAYS[1]), n)


def day34_page(n):
    extra = ('<div class="box"><span class="k">Circulating prompts</span><ul style="margin:0">'
             '<li>‘Play it for me with the video. Where is the music doing something because of what’s on screen?’</li>'
             '<li>‘Which sound is closest to us? Which is farthest away? How did you make that happen?’</li>'
             '<li>‘What is the biggest moment in the video, and what does your music do right there?’</li></ul></div>')
    return page("Day-by-day plans · Days 3 and 4", "DAYS 3–4", "Work", day_box(DAYS[2]) + day_box(DAYS[3]) + extra, n)


def day56_page(n):
    return page("Day-by-day plans · Days 5 and 6", "DAYS 5–6", "Finish + present", day_box(DAYS[4]) + day_box(DAYS[5]), n)


def tn1_page(n):
    inner = (f'<div class="tn big">{H2("What this unit does", True)}'
             f'<p>This is the first open-ended project in DMP1. After four units of structured constraint, students get their first taste of genuine creative '
             f'freedom. The spooky theme and the fixed video brief give just enough structure to prevent paralysis (students know what they are making and '
             f'roughly how long it needs to be), but there are no chord-progression or song-form requirements. The brief is: make something that sounds like '
             f'this video looks.</p>'
             f'<p>The primary technical goal is bussing. This is the first formal introduction to routing in Logic and students consistently find it '
             f'confusing. The unit gives you 6 days partly because the effects processing is worth exploring, but also because some students will need '
             f'multiple attempts to get the bus routing right before they can focus on the creative work.</p>'
             f'<p>The combination of bussing, an open brief and video scoring makes this unit pedagogically rich. Students who have been cautious and '
             f'mechanical in previous units often surprise themselves here. The absence of rules removes the fear of being wrong.</p>'
             f'{H2("Keys are optional")}'
             f'<p>No key is required. A student may work in any key, or in none: atonal and sound-effect scores are welcome. Minor keys are offered as a '
             f'tool, not a rule: they sound darker, and some students like a harmonic home or a change of key at a big moment in the video. For them, the '
             f'Minor Key Map (page {IX["keys"]}) teaches the minor scale and its first chord, and the reference card (page {IX["keyref"]}) gives eight keys '
             f'ready to use.</p>'
             f'<div class="box"><span class="k">How this kit differs from my classroom version</span><ul style="margin:0">'
             f'<li>No video is included: you supply a one-minute clip (page {IX["setup"]}).</li>'
             f'<li>The 32-measure length and the written reflection come from my course restructure; the rubric is my original, with no key criteria.</li>'
             f'<li>Equipment and websites are named as examples; any interface, microphone and free sound-effects site works.</li></ul></div></div>')
    return page("Teacher notes · what this unit does", "TEACHER NOTES", "1 of 4", inner, n)


def tn2_page(n):
    st = "".join(f'<p><b>{a}:</b> {b}</p>' for a, b in STRUGGLES)
    return page("Teacher notes · common student struggles", "TEACHER NOTES", "2 of 4",
                f'<div class="tn big">{H2("Common student struggles", True)}{st}</div>', n)


def tn3_page(n):
    demo = "".join(f'<li><div class="n">{i + 1}</div><div>{s}</div></li>' for i, s in enumerate(DEMO))
    inner = (f'<div class="tn big">{H2("Day 1: the bus demo (20 minutes)", True)}'
             f'<p>This is teacher-led. Students watch first, then replicate. Here is a suggested sequence.</p>'
             f'<ol class="num" style="font-size:11pt">{demo}</ol>'
             f'{H2("About the video")}'
             f'<p>Before playing it for students, watch your clip yourself and identify 3–4 moments worth noting: a visual cut, a shift in lighting, a moment of '
             f'stillness. These are the compositional opportunities you will point students toward during working days.</p>'
             f'<p>Play the video for the full class at the start of Day 1 and again at the start of Day 2. Watch it once without comment, then watch it again '
             f'and ask: what moments stood out? What did you expect to hear? This primes students to think cinematically before they open Logic.</p></div>')
    return page("Teacher notes · the bus demo and the video", "TEACHER NOTES", "3 of 4", inner, n)


def tn4_page(n):
    inner = (f'<div class="tn">{H2("Suggested variations", True)}'
             f'<p><b>More accessible:</b> Pre-build the reverb and delay buses in a Logic template and distribute it. Students open the template and focus on '
             f'sound collection and composition rather than routing. The How to Bus pages become a reference for understanding what is already built rather '
             f'than a setup guide.</p>'
             f'<p><b>More advanced:</b> Challenge students to automate the send level on at least one track so the reverb swells in or out at a specific '
             f'moment in the video. Introduce the idea of a wet/dry blend that changes over time as a compositional tool. Or, as an extra challenge, try a '
             f'half-step key change (A minor to B-flat minor) at the scariest moment.</p>'
             f'{H2("Grading tips")}<ul>'
             f'<li>Each category is scored 2–5, for a total out of 25. Students self-score first; you score in the second column (page {IX["rubric"]}).</li>'
             f'<li>Watch once with the video for Composition &amp; Pacing (are key moments acknowledged? is silence used on purpose?), once eyes-closed for Sound Design and '
             f'Reverb &amp; Delay (near and far), then open the Mixer to check both buses, 100% wet, and the sends.</li>'
             f'<li>The reflection is the quickest evidence of intention: two named techniques with a reason each.</li></ul>'
             f'{H2("Resources")}<ul>'
             f'<li>The intro deck (.pptx and PDF), included.</li>'
             f'<li>Listening: the Halloween theme (John Carpenter) and the Stranger Things theme (Kyle Dixon &amp; Michael Stein), from your own YouTube or '
             f'Spotify access. A search for spooky hip-hop instrumentals or horror film scores gives more.</li>'
             f'<li>A free sound-effects site you approve, and a short Logic “bus and send” tutorial video if you like one.</li></ul>'
             f'{H2("Observation notes")}'
             f'<p class="small">What worked, what to adjust, and anything worth remembering for next year.</p></div><div class="lines"></div>')
    return page("Teacher notes · variations, grading, resources", "TEACHER NOTES", "4 of 4", inner, n)


def post_page(n):
    steps = [("Watch the video", "Your teacher will play the woods video. Watch it without doing anything.",
              "What would you expect to hear? What moments stand out? You will be scoring this video for the next 6 days."),
             ("Watch the demo", "Your teacher will build a reverb and delay bus live in Logic. Watch carefully. Do not open your own computer for building yet.",
              "The demo shows you exactly what you are about to do. Pay attention to the order of steps."),
             ("Open the How to Bus pages", f"Find How to Bus (pages {IX['bus1']}–{IX['bus4']}) and keep it visible next to Logic.",
              "Every step you need is there. If something is not working, check the troubleshooting table before asking for help."),
             ("Set up your reverb bus", "Follow Part 1 step by step. Click an empty Send slot, choose Bus 1, add Space Designer at 100% wet on the new Aux, "
              "and raise the send.", "Call your teacher over to check your setup before moving on to the delay bus."),
             ("Set up your delay bus", "Follow Part 2. Click a Send slot, choose Bus 2, add Tape Delay at 100% wet.",
              "You do not need to add more delay sends today. Just get it set up and working.")]
    lis = "".join(f'<li style="margin-bottom:11pt"><div class="n">{i + 1}</div><div><b class="hd" style="display:inline;font-size:13.5pt">{a}</b> · {b}<br>'
                  f'<span class="why">{c}</span></div></li>' for i, (a, b, c) in enumerate(steps))
    inner = (f'<h2 class="s" style="font-size:20pt;margin:4pt 0 2pt">Spooky Tunes · Day 1</h2>'
             f'<p class="l" style="font-size:12pt">Digital Music Production 1 · Unit 6</p>'
             f'<p style="font-size:11.5pt;margin:6pt 0 4pt">Today you are learning how to bus. Watch the demo first. Then replicate it yourself using the How to Bus pages.</p>'
             f'<p style="font-size:11.5pt;margin:0 0 10pt"><b>Do not start building your Spooky Tunes project until your reverb bus is working.</b></p>'
             f'<ol class="num" style="font-size:12pt">{lis}</ol>'
             f'<div class="box acc-b" style="margin-top:4pt"><span class="k" style="font-size:12pt">The one thing to remember about buses</span>'
             f'<p style="font-size:12.5pt;margin:0 0 3pt">Your original track still plays dry. The bus adds the wet (effected) signal on top. '
             f'The send knob on your track controls how much signal goes to the bus.</p>'
             f'<p style="font-size:11.5pt;margin:0" class="why">Space Designer on the bus must be set to 100% wet. If it is outputting any dry signal, you will '
             f'hear everything doubled.</p></div>'
             f'<div style="flex:1"></div><div class="box"><span class="k">Moments I noticed in the video</span><div class="lines" style="height:66pt;flex:none"></div></div>')
    return page("Day 1 classroom post", "STUDENT", "Day 1", inner, n)


def proj1_page(n):
    inner = (f'<h2 class="s" style="font-size:20pt;margin:2pt 0 0">Project sheet</h2>'
             f'<p class="l" style="font-size:12pt;margin-bottom:6pt">Digital Music Production 1 · Unit 6 · Open-Ended · 6 Days · Reverb &amp; Delay via Bus</p>'
             f'{H2("Overview")}'
             f'<p style="font-size:12pt">This is your first open-ended project. There is no chord progression required. No song form required. No right answer. '
             f'Your brief is a one-minute video of a walk through the woods at night, and your job is to score it. The music should make the video feel the way '
             f'it looks. Use any sounds you can find or record. Experiment. Make something unsettling.</p>'
             f'<p style="font-size:12pt">The new technical skill in this unit is bussing: routing reverb and delay through shared aux channels so multiple tracks '
             f'share the same effect. Your teacher will demo this on Day 1. The How to Bus pages walk you through the setup step by step.</p>'
             f'{H2("Goals")}<ul>' + "".join(f"<li style='font-size:11.5pt'>{g}</li>" for g in GOALS) + '</ul>'
             f'<div class="box inv" style="margin-top:8pt"><span class="k">The brief</span>'
             f'<div style="display:flex;justify-content:space-between;font-family:Oswald;font-weight:700;font-size:15pt;text-transform:uppercase;letter-spacing:.03em">'
             f'<span>1 video</span><span>about 60 seconds</span><span>32 measures</span><span>128 BPM</span><span>any key or none</span></div></div>'
             f'<div style="margin-top:8pt">{score_svg(blank=True)}</div>'
             f'<p class="small" style="margin-top:3pt">Your score map: 32 measures, about one minute. Vocabulary: page {IX["vocab"]}. Rubric: page {IX["rubric"]}.</p>')
    return page("Project sheet · overview and goals", "STUDENT", "Project · 1 of 4", inner, n)


def proj2_page(n, watermark=False):
    src = "".join(f"<tr><td style='width:22%'><b>{a}</b></td><td>{b}</td></tr>" for a, b in SOURCES)
    inner = (f'{H2("Requirements", True)}'
             + "".join(f'<div class="chk" style="font-size:11pt;margin-bottom:5pt">{checkbox(12)}<span>{r}</span></div>' for r in REQUIREMENTS) +
             f'{H2("Sound sources")}<p>You have three places to find sounds for this project. Use all three.</p>'
             f'<table class="t" style="font-size:10.2pt">{src}</table>')
    return page("Project sheet · requirements and sound sources", "STUDENT", "Project · 2 of 4", inner, n, watermark)


def proj3_page(n):
    pic = "".join(f"<tr><td style='width:22%'><b>{a}</b></td><td>{b}</td></tr>" for a, b in PICTURE)
    lis = "".join(f'<li style="margin-bottom:7pt"><div class="n">{i + 1}</div><div><b class="hd">{d} · {h}</b>{t}</div></li>' for i, (d, h, t) in enumerate(STEPS))
    inner = (f'{H2("Composing to picture", True)}<p>Scoring to a video is different from writing music for its own sake. The video tells you what to do '
             f'and when. Before you build anything, watch the video twice and take notes.</p>'
             f'<table class="t" style="font-size:10.2pt">{pic}</table>'
             f'{H2("Steps / process")}<ol class="num" style="font-size:10.8pt">{lis}</ol>')
    return page("Project sheet · composing to picture and steps", "STUDENT", "Project · 3 of 4", inner, n)


def proj4_page(n):
    tech = lambda k: (f'<div class="box" style="margin-bottom:8pt"><span class="k">Technique {k}</span>'
                      f'<div class="fl"><b>What I did (name the sound, the effect or the moment):</b><span></span></div><div class="ln2"></div>'
                      f'<div class="fl" style="margin-top:6pt"><b>Why I chose it (what it does to the video):</b><span></span></div><div class="ln2"></div></div>')
    inner = (f'{H2("Tips for success", True)}<ul>' + "".join(f"<li style='font-size:11pt'>{t}</li>" for t in TIPS) + '</ul>'
             f'{H2("My reflection")}'
             f'<p style="font-size:11pt">Name two specific techniques you used to create atmosphere and explain why you chose them. Not “I used reverb”, but '
             f'“I sent the wind to the reverb bus at –6 dB so it sounds like the whole forest, and kept the footsteps dry so they feel close.”</p>'
             f'{tech(1)}{tech(2)}')
    return page("Project sheet · tips and reflection", "STUDENT", "Project · 4 of 4", inner, n, name=True)


def bus1_page(n):
    inner = (f'<p style="font-size:11pt"><b>Keep these pages open next to Logic while you work.</b></p>'
             f'{H2("What is a bus, and why does it exist?")}'
             f'<p style="font-size:11.5pt">A bus is a shared channel that multiple tracks can send their audio to. Instead of putting a reverb plugin on every single '
             f'track individually, you create one reverb bus and send as many tracks as you want to it. The reverb lives in one place. Everything that sends '
             f'to it gets the same reverb. This matters for three reasons:</p><ul>'
             f'<li style="font-size:11.5pt"><b>Consistency:</b> every track that sends to the bus sounds like it exists in the same space.</li>'
             f'<li style="font-size:11.5pt"><b>Control:</b> one fader on the bus controls how much reverb everything gets at once.</li>'
             f'<li style="font-size:11.5pt"><b>Efficiency:</b> one plugin instead of ten. Less CPU. Less mess.</li></ul>'
             f'<div style="margin:10pt 0 8pt">{bus_svg()}</div>'
             f'<div class="box acc-b"><span class="k" style="font-size:12pt">The send / return relationship</span>'
             f'<p style="font-size:11.5pt;margin:0">Your original track still plays dry: you always hear the unaffected signal. The send knob controls how much '
             f'of that signal gets routed to the bus. The bus adds wet signal (the effect) on top. Turn the send up: more reverb. Turn it down: less. '
             f'The original track is unchanged.</p></div>')
    return page("How to Bus · reverb and delay", "HOW TO BUS", "1 of 4 · What a bus is", inner, n)


def bus_steps(rows, fs=10.8):
    return '<ol class="num">' + "".join(
        f'<li style="margin-bottom:8pt"><div class="n">{i + 1}</div><div style="font-size:{fs}pt;line-height:1.32"><b class="hd">{a}</b>{b}<br>'
        f'<span class="why">{c}</span></div></li>' for i, (a, b, c) in enumerate(rows)) + '</ol>'


def bus2_page(n, watermark=False):
    inner = (f'{H2("Part 1 · Set up your reverb bus", True)}'
             f'<p style="font-size:11pt">Start from the track you want to add reverb to, not from a blank aux channel. Logic creates the bus automatically when '
             f'you make the send.</p>{bus_steps(BUS_PART1)}')
    return page("How to Bus · Part 1: the reverb bus", "HOW TO BUS", "2 of 4 · Reverb", inner, n, watermark)


def bus3_page(n):
    inner = (f'{H2("Part 2 · Set up your delay bus", True)}'
             f'<p style="font-size:11pt">Same process as the reverb bus: start from the track, click a Send slot, choose an empty bus, and Logic creates the Aux. '
             f'Do this after your reverb bus is confirmed working.</p>{bus_steps(BUS_PART2)}')
    return page("How to Bus · Part 2: the delay bus", "HOW TO BUS", "3 of 4 · Delay", inner, n)


def bus4_page(n):
    tr = "".join(f"<tr><td style='width:38%'><b>{a}</b></td><td>{b}</td></tr>" for a, b in TROUBLE)
    at = "".join(f"<tr><td style='width:26%'><b>{a}</b></td><td>{b}</td></tr>" for a, b in ATMOS)
    qz = "".join(f'<div class="fl" style="margin-bottom:5pt"><b>{i + 1}. {q}</b><span></span></div>' for i, (q, _) in enumerate(BUS_QUIZ))
    inner = (f'{H2("Troubleshooting", True)}<table class="t" style="font-size:9.8pt"><tr><th>Problem</th><th>Fix</th></tr>{tr}</table>'
             f'{H2("Using reverb &amp; delay for atmosphere")}<table class="t" style="font-size:9.8pt">{at}</table>'
             f'<div class="box" style="margin-top:4pt"><span class="k">Check yourself · answers on page {IX["key"]}</span>{qz}</div>')
    return page("How to Bus · troubleshooting and atmosphere", "HOW TO BUS", "4 of 4 · Fixes", inner, n)


def key_slot(label, n_slots):
    return (f'<div style="display:flex;align-items:center;gap:8pt;margin:0 0 6pt"><b style="font-family:Oswald;font-size:11pt;text-transform:uppercase;'
            f'flex:0 0 1.25in">{label}</b>' + "".join('<span class="slot"></span>' for _ in range(n_slots)) + '</div>')


def keys_page(n):
    def sec(label):
        return (f'<div class="box" style="margin-bottom:6pt;padding:5pt 10pt"><div class="fl" style="font-size:12pt;margin-bottom:5pt"><b>{label}:</b><span></span>'
                f'<b>minor</b></div>'
                f'<div style="display:flex;gap:6pt;margin-bottom:2pt"><div style="flex:0 0 1.25in"></div>'
                + "".join(f'<div class="lab" style="width:44pt;margin-right:5pt">{x}</div>' for x in ("Root", "Third", "Fifth")) + '</div>'
                + key_slot("First minor chord", 3) + key_slot("Minor scale", 7) +
                f'<div style="width:60%;margin:2pt auto 0">{keyboard_svg(None)}</div>'
                f'<p class="small" style="text-align:center;margin:2pt 0 0">Colour or circle the chord. Put a dot on every scale note.</p></div>')
    inner = (f'<p style="font-size:10.5pt;margin:0 0 5pt"><b>This page is optional.</b> Your project does not need a key: any key, or none, is fine, and '
             f'atonal or sound-effect scores are welcome. Want a darker sound? Try a minor key. Some composers also change key at a big moment in the video.</p>'
             f'<div class="box acc-b" style="margin-bottom:8pt"><span class="k">The recipe</span>'
             f'<p style="margin:0 0 2pt"><b>Minor scale:</b> start on any key, then go up Whole, Half, Whole, Whole, Half, Whole, Whole (a half step is the very '
             f'next key, black or white; a whole step skips one). <b>First minor chord:</b> scale notes 1, 3 and 5 (root, third, fifth).</p>'
             f'<p style="margin:0"><b>Spooky tip:</b> the half steps in a minor scale sit between notes 2–3 and 5–6. Play them together and they rub.</p></div>'
             f'{sec("My key")}{sec("A second key (only if I change key)")}'
             f'<div class="box"><span class="k">If you change key: pairs to try</span>'
             + "".join(f'<p style="margin:0 0 2pt"><b>{a} › {b}:</b> {why}</p>' for a, b, why in PAIRS) + '</div>')
    return page("Minor key map · optional", "STUDENT", f"Optional · reference p. {IX['keyref']}", inner, n, name=True)


def keyref_page(n):
    cards = "".join(
        f'<div class="box" style="padding:6pt 8pt"><div style="display:flex;justify-content:space-between;align-items:baseline">'
        f'<b style="font-family:Oswald;font-size:15pt">{k}</b><span class="small">{KEY_NOTES[k]}</span></div>'
        f'<div style="margin:3pt 0">{keyboard_svg(k)}</div>'
        f'<p style="margin:0;font-size:9.8pt"><b>Chord:</b> {" ".join(chord(k))} &nbsp; <b>Scale:</b> {" ".join(s)}</p></div>'
        for k, s in KEYS.items())
    inner = (f'<p style="font-size:10.5pt">Optional: no key is required, and atonal or sound-effect scores are fine. Want a darker sound? Pick one of these. '
             f'Orange keys: the first minor chord (root, third, fifth). Dots: the rest of the scale. “b” means flat and “#” means sharp.</p>'
             f'<div style="display:grid;grid-template-columns:1fr 1fr;gap:8pt 12pt">{cards}</div>')
    return page("Minor key reference card · eight keys", "STUDENT", "Reference", inner, n)


def cue_page(n):
    rows = "".join("<tr><td style='height:27pt'></td><td></td><td></td><td></td><td></td></tr>" for _ in range(8))
    inner = (f'<div class="box" style="margin-bottom:6pt"><span class="k">Tempo math</span>'
             f'<div class="fl" style="margin-bottom:4pt"><b>My video is</b><span style="flex:0 0 50pt"></span><b>seconds. My tempo = 7680 ÷</b>'
             f'<span style="flex:0 0 50pt"></span><b>=</b><span style="flex:0 0 50pt"></span><b>BPM</b></div>'
             f'<p class="small" style="margin:0">A 60-second video = 128 BPM: 1 measure = 1.875 seconds, 4 measures = 7.5 seconds. '
             f'Measure number = (seconds ÷ 1.875), round down, then + 1.</p>'
             f'<div style="display:flex;gap:16pt;margin-top:5pt">' + "".join(f'<div class="fl" style="flex:1;margin:0"><b>{tfmt(s)} is measure</b><span></span></div>'
                                                                              for s, _ in CUE_PRACTICE) + '</div></div>'
             f'<p style="font-size:10.5pt;margin-top:4pt">Watch the video twice. Pause at every moment that matters and fill in a row.</p>'
             f'<table class="t"><tr><th style="width:12%">Time</th><th style="width:11%">Measure</th><th>What I see</th><th>What I’ll do with sound</th>'
             f'<th style="width:11%">Key? (optional)</th></tr>{rows}</table>'
             f'<span class="k" style="margin-top:6pt">My score map</span>{score_svg(blank=True)}'
             f'<p class="small" style="margin-top:3pt">Name your tracks and sketch where each sound plays. Using a key? Write it in the bar (optional).</p>')
    return page("Cue sheet · map your moments", "STUDENT", "Days 2–3", inner, n, name=True)


def vocab_page(n):
    rows = "".join(f"<tr><td style='width:22%'><b>{a}</b></td><td>{b}</td><td style='width:20%'></td></tr>" for a, b in VOCAB)
    qz = "".join(f'<div class="fl" style="margin-bottom:5pt"><b>{q}</b><span></span></div>' for q, _ in VOCAB_QUIZ)
    inner = (f'<p style="font-size:11pt">Read each word and its meaning. In the last column, write where you used it in your own project.</p>'
             f'<table class="t" style="font-size:10pt"><tr><th>Word</th><th>Meaning</th><th>In my project</th></tr>{rows}</table>'
             f'<div class="box" style="margin-top:4pt"><span class="k">Quick check</span>{qz}</div>')
    return page("Key vocabulary", "STUDENT", f"{len(VOCAB)} words", inner, n, name=True)


def checklist_page(n):
    inner = (f'<p style="font-size:11.5pt">Tick each box before you submit and present. If a box isn’t ticked, you aren’t done yet.</p>'
             + "".join(f'<div class="chk" style="font-size:11.5pt;margin-bottom:7pt">{checkbox(14)}<span>{c}</span></div>' for c in CHECKLIST) +
             f'{H2("My three feedback questions")}'
             f'<p class="small">Ask about something you really want to know. Example: “Does the silence before the cut make the next sound feel bigger?”</p>'
             f'<div class="fl"><b>1.</b><span></span></div><div class="fl"><b>2.</b><span></span></div><div class="fl"><b>3.</b><span></span></div>'
             f'{H2("Notes")}<div class="lines"></div>')
    return page("Submission checklist", "STUDENT", "Before you present", inner, n, name=True)


def rubric_page(n, watermark=False):
    head = "<tr><th>Category</th>" + "".join(f"<th>{lv}</th>" for lv in LEVELS) + "<th>Self</th><th>Teacher</th></tr>"
    rows = "".join(f"<tr><td class='c'>{c}</td>" + "".join(f"<td>{d}</td>" for d in ds) + "<td class='sc'></td><td class='sc'></td></tr>"
                   for c, ds in RUBRIC)
    total = (f"<tr><td class='c' colspan='5' style='text-align:right;font-size:13pt'>Total</td>"
             f"<td style='font-family:Oswald;font-weight:700;font-size:13pt;text-align:right;vertical-align:bottom'>/ {5 * len(RUBRIC)}</td>"
             f"<td style='font-family:Oswald;font-weight:700;font-size:13pt;text-align:right;vertical-align:bottom'>/ {5 * len(RUBRIC)}</td></tr>")
    inner = (f'<p style="font-size:11pt">Score each category from 2 to 5. Score yourself first, then your teacher scores.</p>'
             f'<table class="rub">{head}{rows}{total}</table>'
             f'<div class="cols" style="margin-top:8pt"><div><h2 class="s" style="margin-top:0">One sound I’m proud of</h2><div class="ln2"></div></div>'
             f'<div><h2 class="s" style="margin-top:0">Teacher comments</h2><div class="ln2"></div></div></div>')
    return page("Rubric", "RUBRIC", f"/ {5 * len(RUBRIC)}", inner, n, watermark, name=True)


def key_page(n):
    def kk(k, ms):
        return (f'<div class="box" style="padding:6pt 8pt"><span class="k">{k} · worked example, m. {ms}</span>'
                f'<p style="margin:0 0 2pt;font-size:10pt"><b>Chord (root, third, fifth):</b> {" · ".join(chord(k))} &nbsp; <b>Scale:</b> {" ".join(KEYS[k])}</p>'
                f'{keyboard_svg(k)}</div>')
    bq = "".join(f'<p style="margin:0 0 2pt;font-size:10pt"><b>{i + 1}.</b> {a}</p>' for i, (_, a) in enumerate(BUS_QUIZ))
    vq = "".join(f'<p style="margin:0 0 2pt;font-size:10pt"><b>{i + 1}.</b> {a}</p>' for i, (_, a) in enumerate(VOCAB_QUIZ))
    cq = " &nbsp;·&nbsp; ".join(f"{tfmt(s)} = measure {m}" for s, m in CUE_PRACTICE)
    bp = " &nbsp;·&nbsp; ".join(f"{s} s = {b} BPM" for s, b in BPM_PRACTICE)
    inner = (f'{H2("Minor key map · example answer", True)}'
             f'<p style="font-size:10pt">Optional page: only for students who chose a key. Check answers against the reference card (page {IX["keyref"]}). '
             f'These are the worked example’s keys.</p>'
             f'<div class="cols" style="gap:12pt">{kk(KEY1, "1–16")}{kk(KEY2, "17–32")}</div>'
             f'<div class="cols" style="gap:14pt;margin-top:8pt"><div>{H2("How to Bus · check yourself", True)}{bq}</div>'
             f'<div>{H2("Vocabulary · quick check", True)}{vq}</div></div>'
             f'{H2("Cue sheet · tempo math")}'
             f'<p style="font-size:10pt">At 128 BPM: {cq}. Tempo for other lengths (7680 ÷ seconds): {bp} (75 s gives 102.4; round to 102).</p>'
             f'<div class="box" style="margin-top:4pt"><span class="k">Checking the buses in a student project</span>'
             f'<p style="margin:0;font-size:10pt">Open the Mixer (X). Two Aux channels named REVERB BUS (input Bus 1, Space Designer) and DELAY BUS (input Bus 2, '
             f'Tape Delay or Echo), each plugin at 100% wet. At least 3 tracks show a Bus 1 send above its lowest setting and at least 1 shows a Bus 2 send. '
             f'No reverb or delay plugin sits directly on a track’s inserts.</p></div>')
    return page("Answer key · key map, bus check, vocabulary, tempo", "ANSWER KEY", "Teacher", inner, n, ak=True)


def example_page(n):
    rows = "".join(f"<tr><td><b>{a}</b><br><span class='small'>{c}</span></td><td>{d}</td>"
                   f"<td style='white-space:nowrap'>{'–' if r is None else f'{r} dB'}</td><td style='white-space:nowrap'>{'off' if dl is None else f'{dl} dB'}</td>"
                   f"<td>{w}</td></tr>" for a, _, c, d, r, dl, w, _ in TRACKS)
    mom = "".join(f'<p style="margin:0 0 2pt;font-size:9.8pt"><b class="acc">{i + 1}</b> &nbsp;<b>m. {m} · {tfmt(m_start(m))}</b> &nbsp;{t}</p>'
                  for i, (m, t) in enumerate(MOMENTS))
    inner = (f'<p style="font-size:10.2pt">Use this to model the score on Days 3–4 or as a reference while you circulate. It is one possible answer for a 60-second '
             f'clip at 128 BPM, not the answer: your video’s moments will be different, and every student’s score should sound different.</p>'
             f'{score_svg()}'
             f'<div class="cols" style="gap:14pt;margin-top:6pt"><div style="flex:0 0 3.35in">{H2("The moments", True)}{mom}</div>'
             f'<div><span class="k" style="margin-top:2pt">What a 5 sounds like</span><p style="font-size:9.8pt;margin:0">The music changes because the picture '
             f'changes; the biggest change lands on a cut; near sounds are dry and far sounds are wet; one moment of real silence; and a student who can name '
             f'two techniques and say why. Requirement check: 6 sources (3 library, 2 recordings, 1 download), 6 tracks to the reverb bus, 2 to the delay bus.</p></div></div>'
             f'<table class="t" style="font-size:9pt;margin-top:4pt"><tr><th>Track · source</th><th>Notes</th><th>Reverb send</th><th>Delay send</th><th>Why</th></tr>{rows}</table>'
             f'<p class="small">Send levels are starting points for a demo, not targets. Trust your ears.</p>')
    return page("Example notes · a worked 32-measure score", "EXAMPLE", "Teacher", inner, n, ak=True)


def terms_page(n):
    return (f'<section class="page"><div class="top"><div class="t">Terms of Use<small>Thank you for your purchase!</small></div></div>'
            f'<div class="main"><p>© {YEAR} {STORE}. {BRAND} by {STORE}. All rights reserved. This purchase gives <b>one teacher</b> a license to use this '
            f'resource (the PDF and the intro deck) with their own students.</p>'
            f'<h2 class="s">You may</h2><ul><li>Print and copy pages for your own classroom and students, year after year.</li>'
            f'<li>Present the intro deck to your own classes and edit it for your own teaching.</li>'
            f'<li>Share pages with your own students through a password-protected class site or learning platform.</li>'
            f'<li>Leave the pages for a substitute teacher who teaches your own students.</li></ul>'
            f'<h2 class="s">You may not</h2><ul><li>Share, email or copy this resource for other teachers, a whole school or a district. '
            f'Additional licenses are available on Teachers Pay Teachers.</li>'
            f'<li>Post any part of it on a public website, shared drive or social media.</li>'
            f'<li>Sell, give away or claim any part of it as your own, or edit it into a new product.</li></ul>'
            f'<h2 class="s">Questions?</h2><p>Please use the Q&amp;A tab on the {BRAND} by {STORE} store on Teachers Pay Teachers. '
            f'If something looks wrong, let me know there and I will fix it.</p>'
            f'<h2 class="s">Credits</h2><ul>'
            f'<li>Curriculum, project sheet, rubric, teacher notes, How to Bus guide and intro deck: original work by {STORE}, {BRAND}.</li>'
            f'<li>Woods, keyboard, signal-flow and score-map drawings and layouts: original work, {BRAND}.</li>'
            f'<li>Text fonts: Oswald (The Oswald Project Authors) and Source Sans 3 (Adobe), SIL Open Font License 1.1.</li>'
            f'<li>No video, audio, lyrics or album art are included. Songs are named for reference only (title and artist); play them from your own '
            f'YouTube or Spotify access. You supply the video students score; follow the license of any clip you download.</li>'
            f'<li>Logic Pro, Space Designer and Tape Delay are trademarks of Apple Inc.; Focusrite and Scarlett are trademarks of Focusrite Audio Engineering; '
            f'they are named only to describe compatible equipment, and this resource is not affiliated with or endorsed by either company. Padlet, YouTube, '
            f'Spotify, Pixabay and Pexels are named only as examples of tools you may already use.</li></ul>'
            f'<p class="small">{DISCLOSURE}</p><div style="flex:1"></div></div>{foot(f"Page {n}")}</section>')


def render_page(key, watermark=False):
    n = IX[key]
    f = {"cover": cover_page, "howto": lambda: howto_page(n), "overview": lambda: overview_page(n), "setup": lambda: setup_page(n),
         "logicref": lambda: logicref_page(n), "pacing": lambda: pacing_page(n, watermark), "day12": lambda: day12_page(n),
         "day34": lambda: day34_page(n), "day56": lambda: day56_page(n), "tn1": lambda: tn1_page(n), "tn2": lambda: tn2_page(n),
         "tn3": lambda: tn3_page(n), "tn4": lambda: tn4_page(n), "post": lambda: post_page(n), "proj1": lambda: proj1_page(n),
         "proj2": lambda: proj2_page(n, watermark), "proj3": lambda: proj3_page(n), "proj4": lambda: proj4_page(n),
         "bus1": lambda: bus1_page(n), "bus2": lambda: bus2_page(n, watermark), "bus3": lambda: bus3_page(n), "bus4": lambda: bus4_page(n),
         "keys": lambda: keys_page(n), "keyref": lambda: keyref_page(n), "cue": lambda: cue_page(n), "vocab": lambda: vocab_page(n),
         "checklist": lambda: checklist_page(n), "rubric": lambda: rubric_page(n, watermark), "key": lambda: key_page(n),
         "example": lambda: example_page(n), "terms": lambda: terms_page(n)}
    return f[key]()


# ---------------------------------------------------------------- intro deck (.pptx, python-pptx) + PDF copy (LibreOffice)
RUBRIC_SHORT = ["Sounds from 5+ sources create an unsettling atmosphere that serves the video at every moment.",
                "The music serves the picture: key moments are acknowledged and silence is a tool.",
                "Both buses set up correctly and used with intention; send levels vary to create depth.",
                "Near sounds dry, far sounds wet: a convincing sense of physical space.",
                "On time, named correctly, and a confident presentation that explains the sound choices."]
DECK_DAYS = [("Day 1", "Bus demo · your reverb bus working"), ("Day 2", "Delay bus · collect 5+ sounds · mark your moments"),
             ("Day 3", "Import the video · markers · start scoring"), ("Day 4", "Score the whole video · rough mix"),
             ("Day 5", "Final mix · reflection · submit"), ("Day 6", "Present with the video")]
WOODS_WIDE = dict(vw=1200, vh=380, seed=5, moon=(0.5, 0.42, 0.25),
                  trees=[(0.05, 0.92, 0.08), (0.16, 0.62, -0.05), (0.3, 0.44, 0.05), (0.7, 0.5, -0.06), (0.84, 0.8, 0.04), (0.96, 0.96, -0.03)])


def build_deck(img):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

    BLACK, WHITE, ACC = RGBColor(0, 0, 0), RGBColor(255, 255, 255), RGBColor.from_string(ACCENT[1:])
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.author = f"{STORE} · {BRAND}"
    prs.core_properties.title = f"{TITLE} · Intro deck"
    prs.core_properties.comments = DISCLOSURE
    layout = prs.slide_layouts[6]

    def box(sl, x, y, w, h, fillc=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=1.5):
        s = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        if fillc is None:
            s.fill.background()
        else:
            s.fill.solid()
            s.fill.fore_color.rgb = fillc
        if line is None:
            s.line.fill.background()
        else:
            s.line.color.rgb = line
            s.line.width = Pt(lw)
        s.shadow.inherit = False
        return s

    def text(sl, x, y, w, h, runs, size=24, font="Source Sans 3", bold=False, color=BLACK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             spacing=None, caps=False):
        tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.04)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        paras = runs if isinstance(runs, list) else [runs]
        for i, ptxt in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = ptxt.upper() if caps else ptxt
            r.font.size = Pt(size)
            r.font.name = font
            r.font.bold = bold
            r.font.color.rgb = color
            if spacing:
                r._r.get_or_add_rPr().set("spc", str(spacing))
        return tb

    def chrome(sl, i, title):
        text(sl, 0.6, 0.32, 6, 0.4, BRAND, 13, "Oswald", False, BLACK, spacing=600, caps=True)
        text(sl, 0.6, 0.72, 12.1, 0.95, title, 40, "Oswald", True, BLACK, caps=True, anchor=MSO_ANCHOR.BOTTOM)
        box(sl, 0.6, 1.72, 12.13, 0.04, BLACK)
        box(sl, 0.6, 7.0, 12.13, 0.012, BLACK)
        text(sl, 0.6, 7.04, 9, 0.3, f"{TITLE} · Intro deck · © {YEAR} {STORE} · {BRAND}", 10, "Source Sans 3", False, RGBColor(0x33, 0x33, 0x33))
        text(sl, 11.2, 7.04, 1.53, 0.3, f"{i}", 10, "Oswald", True, BLACK, align=PP_ALIGN.RIGHT)

    def bullets(sl, items, x=0.6, y=2.05, w=12.1, size=28, dot=True):
        for k, it in enumerate(items):
            yy = y + k * (size / 72 * 1.75)
            if dot:
                box(sl, x, yy + size / 72 * 0.36, 0.16, 0.16, ACC, shape=MSO_SHAPE.OVAL)
            text(sl, x + (0.38 if dot else 0), yy, w - 0.4, size / 72 * 1.6, it, size, "Source Sans 3", False, BLACK)

    def pic(sl, key, x, y, w):
        path, (pw, ph) = img[key]
        sl.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))
        return y + w * ph / pw

    for i, (kind, title, body, notes) in enumerate(SLIDES, start=1):
        sl = prs.slides.add_slide(layout)
        items = [b for b in body.split("|") if b]
        if kind == "title":
            text(sl, 0, 0.42, 13.333, 0.5, BRAND, 16, "Oswald", False, BLACK, align=PP_ALIGN.CENTER, spacing=800, caps=True)
            pic(sl, "woods", 2.0, 1.0, 9.33)
            text(sl, 0, 4.0, 13.333, 1.3, title, 80, "Oswald", True, BLACK, align=PP_ALIGN.CENTER, caps=True)
            text(sl, 0, 5.4, 13.333, 0.6, body, 28, "Oswald", False, BLACK, align=PP_ALIGN.CENTER)
            box(sl, 0, 6.55, 13.333, 0.95, BLACK)
            text(sl, 0, 6.68, 13.333, 0.6, "Digital Music Production · Unit 6 · Scoring to picture in Logic Pro", 20, "Oswald", False, WHITE,
                 align=PP_ALIGN.CENTER, spacing=200, caps=True)
        else:
            chrome(sl, i, title)
        if kind == "listen":
            song, _, artist = items[0].partition(" · ")
            box(sl, 0.6, 2.0, 12.13, 2.0, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            text(sl, 0.9, 2.1, 11.5, 0.7, song, 36, "Oswald", True, WHITE, caps=True)
            text(sl, 0.9, 2.82, 11.5, 0.5, artist, 24, "Oswald", False, ACC)
            text(sl, 0.9, 3.38, 11.5, 0.5, items[1], 18, "Source Sans 3", False, WHITE)
            bullets(sl, items[2:], y=4.3, size=30)
        elif kind in ("bullets", "closer"):
            bullets(sl, items, size=30 if kind == "bullets" else 40, y=2.2 if kind == "bullets" else 2.3)
            if kind == "closer":
                pic(sl, "woods", 7.6, 5.0, 5.13)
        elif kind == "columns":
            for k, it in enumerate(items):
                x = 0.6 + k * 3.08
                box(sl, x, 2.3, 2.85, 3.9, WHITE, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=2.25)
                head, _, rest = it.partition(":")
                text(sl, x + 0.15, 2.5, 2.55, 0.8, head, 30, "Oswald", True, ACC if k < 2 else BLACK, align=PP_ALIGN.CENTER, caps=True)
                text(sl, x + 0.15, 3.4, 2.55, 2.6, rest.strip(), 24, "Source Sans 3", False, BLACK, align=PP_ALIGN.CENTER)
        elif kind == "bus":
            pic(sl, "bus", 1.37, 1.95, 10.6)
        elif kind == "steps":
            for k, it in enumerate(items):
                y = 2.05 + k * 0.95
                box(sl, 0.6, y + 0.05, 0.62, 0.62, BLACK, shape=MSO_SHAPE.OVAL)
                text(sl, 0.6, y + 0.05, 0.62, 0.62, str(k + 1), 24, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                text(sl, 1.45, y - 0.05, 11.2, 0.9, it, 22, "Source Sans 3", False, BLACK, anchor=MSO_ANCHOR.MIDDLE)
        elif kind == "keys":
            pic(sl, "keys", 0.6, 2.0, 12.13)
        elif kind == "score":
            yb = pic(sl, "score", 0.6, 2.0, 12.13)
            text(sl, 0.6, yb + 0.15, 12.1, 0.5, "32 measures at 128 BPM = 60 seconds. Every 4 measures = 7.5 seconds. Numbers = moments in the video.",
                 17, "Source Sans 3", False, BLACK)
        elif kind == "days":
            for k, (d, t) in enumerate(DECK_DAYS):
                y = 2.0 + k * 0.78
                box(sl, 0.6, y, 1.45, 0.6, ACC if k in (2, 3) else BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                text(sl, 0.6, y, 1.45, 0.6, d, 20, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, caps=True)
                text(sl, 2.25, y, 10.4, 0.6, t, 24, "Source Sans 3", False, BLACK, anchor=MSO_ANCHOR.MIDDLE)
        elif kind == "check":
            for k, it in enumerate(items):
                col, row = k % 2, k // 2
                x, y = 0.6 + col * 6.15, 2.15 + row * 1.5
                box(sl, x, y + 0.08, 0.42, 0.42, WHITE, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=2.25)
                text(sl, x + 0.6, y, 5.35, 1.35, it, 22, "Source Sans 3", False, BLACK)
        elif kind == "rubric":
            for k, ((c, _), short) in enumerate(zip(RUBRIC, RUBRIC_SHORT)):
                y = 2.05 + k * 0.92
                box(sl, 0.6, y, 4.3, 0.78, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                text(sl, 0.75, y, 4.0, 0.78, c, 18, "Oswald", True, WHITE, anchor=MSO_ANCHOR.MIDDLE, caps=True)
                text(sl, 5.1, y, 6.6, 0.78, "5 = " + short, 16, "Source Sans 3", False, BLACK, anchor=MSO_ANCHOR.MIDDLE)
                text(sl, 11.7, y, 1.03, 0.78, "/ 5", 22, "Oswald", True, ACC, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
            text(sl, 9.0, 6.6, 3.73, 0.4, f"Total / {5 * len(RUBRIC)}", 20, "Oswald", True, BLACK, align=PP_ALIGN.RIGHT, caps=True)
        sl.notes_slide.notes_text_frame.text = notes
    prs.save(DECK)


def ttf_fonts():
    """woff2 -> ttf for LibreOffice (deck PDF), plus a fontconfig file that adds them."""
    from fontTools.ttLib import TTFont
    d = os.path.join(BUILD, "ttf")
    os.makedirs(d, exist_ok=True)
    styles = {"oswald-latin-500-normal": ("Oswald", "Regular", 400), "oswald-latin-700-normal": ("Oswald", "Bold", 700),
              "source-sans-3-latin-400-normal": ("Source Sans 3", "Regular", 400), "source-sans-3-latin-700-normal": ("Source Sans 3", "Bold", 700),
              "source-sans-3-latin-400-italic": ("Source Sans 3", "Italic", 400)}
    for stem, (fam, sub, wt) in styles.items():
        out = os.path.join(d, stem + ".ttf")
        if os.path.exists(out):
            continue
        t = TTFont(os.path.join(FONTS, stem + ".woff2"))
        t.flavor = None
        name = t["name"]
        ps = (fam.replace(" ", "") + "-" + sub)
        for rec in list(name.names):
            if rec.nameID in (16, 17, 21, 22):
                name.removeNames(nameID=rec.nameID)
        for nid, val in ((1, fam), (2, sub), (4, f"{fam} {sub}" if sub != "Regular" else fam), (6, ps)):
            name.setName(val, nid, 3, 1, 0x409)
            name.setName(val, nid, 1, 0, 0)
        t["OS/2"].usWeightClass = wt
        sel = 0
        if sub == "Bold":
            sel |= 1 << 5
        if sub == "Italic":
            sel |= 1
        if sub == "Regular":
            sel |= 1 << 6
        t["OS/2"].fsSelection = (t["OS/2"].fsSelection & ~0b1100001) | sel
        t["head"].macStyle = (1 if sub == "Bold" else 0) | (2 if sub == "Italic" else 0)
        t.save(out)
    conf = os.path.join(BUILD, "fonts.conf")
    with open(conf, "w") as fh:
        fh.write(f'<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig><include ignore_missing="yes">/etc/fonts/fonts.conf</include>'
                 f'<dir>{d}</dir><cachedir>{os.path.join(BUILD, "fccache")}</cachedir></fontconfig>')
    return conf


def deck_pdf(src=DECK, out=DECK_PDF, fonts=True):
    """fonts=False renders with system fonts only: the font-fallback view a teacher without Oswald would see (self-check)."""
    env = dict(os.environ, FONTCONFIG_FILE=ttf_fonts()) if fonts else dict(os.environ)
    prof = os.path.join(BUILD, "lo-profile" if fonts else "lo-profile-fallback")
    tmp = os.path.join(BUILD, "lo-out")
    os.makedirs(tmp, exist_ok=True)
    subprocess.run(["soffice", f"-env:UserInstallation=file://{prof}", "--headless", "--convert-to", "pdf", "--outdir", tmp, src],
                   check=True, env=env, capture_output=True)
    shutil.move(os.path.join(tmp, os.path.basename(src).replace(".pptx", ".pdf")), out)


# ---------------------------------------------------------------- square images (2000 x 2000)
SQ_CSS = """
.sq { width: 1000px; height: 1000px; position: relative; overflow: hidden; background: linear-gradient(180deg,#fff 0%,#fff 60%,#eeeeee 100%); }
.sq .store { position: absolute; left: 0; right: 0; top: 30px; text-align: center; font-family: 'Oswald'; font-weight: 500; letter-spacing: .32em;
             text-transform: uppercase; font-size: 22px; }
.sq .fillw { position: absolute; left: 0; right: 0; top: 66px; bottom: 66px; display: flex; flex-direction: column; align-items: center; justify-content: space-evenly; }
.sq .hd { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; text-align: center; line-height: .95; white-space: nowrap; }
.sq .sub { font-family: 'Oswald'; font-weight: 300; text-align: center; }
.sq .band { position: absolute; left: 0; right: 0; bottom: 0; background: #000; color: #fff; text-align: center; font-family: 'Oswald';
            font-weight: 500; text-transform: uppercase; letter-spacing: .1em; font-size: 26px; padding: 16px 0 18px; }
.sq .band em { font-style: normal; color: ACCENT; }
.shot { background: #fff; border: 2px solid #000; box-shadow: 8px 8px 0 #000; position: relative; overflow: hidden; }
.shot img { display: block; width: 100%; }
.lbl { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 22px; text-align: center; margin-top: 14px; }
.lbl span { font-weight: 300; text-transform: none; display: block; font-size: 18px; }
.swm { position: absolute; left: 50%; top: 50%; transform: translate(-50%,-50%) rotate(-38deg); font-family: 'Oswald'; font-weight: 700;
       color: rgba(0,0,0,.14); white-space: nowrap; letter-spacing: .06em; }
""".replace("ACCENT", ACCENT)
FONT_PRIME = ('<div style="position:absolute;left:-9999px;font-family:\'Source Sans 3\'">'
              '<span style="font-weight:400">a</span><span style="font-weight:600">a</span><i style="font-weight:400">a</i></div>')


def sq_doc(body):
    return (f'<!doctype html><html><head><meta charset="utf-8"><style>{font_css()}'
            f'html,body{{margin:0;background:#fff}}{SQ_CSS}</style></head><body>{FONT_PRIME}{body}</body></html>')


def sq(inner, band):
    return sq_doc(f'<div class="sq"><div class="store">{BRAND}</div><div class="fillw">{inner}</div><div class="band">{band}</div></div>')


def cover_square():
    return sq(f'<div style="width:960px">{woods_svg()}</div>'
              f'<div><div class="hd" style="font-size:56px">Scoring to Picture in Logic Pro</div>'
              f'<div class="hd" style="font-size:150px;margin-top:6px">Spooky Tunes</div>'
              f'<div class="sub" style="font-size:30px;margin-top:12px">6-day unit · reverb &amp; delay buses · sound design · intro slides</div></div>',
              f'Grades 9–12 <em>·</em> 6 days <em>·</em> {len(PLAN)} pages + {DECK_SLIDES} slides')


def preview_square_1(img):
    rot = [-4, 0, 4]
    shots = "".join(f'<div style="position:absolute;left:{x}px;top:{y}px;width:372px;transform:rotate({r}deg);z-index:{z}">'
                    f'<div class="shot"><img src="{s}"><div class="swm" style="font-size:80px">PREVIEW</div></div></div>'
                    for (s, _, _), x, y, r, z in zip(img, (40, 314, 588), (24, 0, 24), rot, (1, 3, 2)))
    labels = "".join(f'<div class="lbl" style="width:300px">{a}<span>{b}</span></div>' for _, a, b in img)
    return sq(f'<div><div class="hd" style="font-size:76px">Ready to teach Monday</div>'
              f'<div class="sub" style="font-size:28px;margin-top:8px">Teacher notes, day plans, student pages and the rubric in one PDF</div></div>'
              f'<div style="position:relative;width:1000px;height:520px">{shots}</div>'
              f'<div style="display:flex;gap:30px">{labels}</div>',
              'You supply the video <em>·</em> we supply everything else')


def preview_square_2(slides):
    cells = "".join(f'<div class="shot" style="box-shadow:6px 6px 0 #000"><img src="{s}"><div class="swm" style="font-size:52px">PREVIEW</div></div>'
                    for s in slides)
    return sq(f'<div><div class="hd" style="font-size:76px">Intro deck included</div>'
              f'<div class="sub" style="font-size:28px;margin-top:8px">{DECK_SLIDES} editable slides (.pptx + PDF) with speaker notes on every slide</div></div>'
              f'<div style="display:grid;grid-template-columns:420px 420px;gap:26px 34px">{cells}</div>',
              'Hook <em>·</em> listen <em>·</em> worked example <em>·</em> steps <em>·</em> rubric')


def preview_square_3(imgs):
    cells = "".join(f'<div style="width:214px"><div class="shot" style="box-shadow:5px 5px 0 #000"><img src="{s}">'
                    f'<div class="swm" style="font-size:46px">PREVIEW</div></div><div class="lbl" style="font-size:18px;margin-top:9px">{lab}</div></div>'
                    for s, lab in imgs)
    return sq(f'<div><div class="hd" style="font-size:72px">Everything inside</div>'
              f'<div class="sub" style="font-size:28px;margin-top:8px">How to Bus guide, minor key map, cue sheet, answer key and a worked score</div></div>'
              f'<div style="display:grid;grid-template-columns:repeat(4,214px);gap:22px 24px">{cells}</div>',
              f'{len(PLAN)} pages <em>·</em> US Letter PDF <em>·</em> prints in black &amp; white')


SHOTS1 = [("pacing", "Pacing + plans", "6 days, minute by minute"), ("proj2", "Project sheet", "requirements, sound sources"),
          ("rubric", "Rubric", "5 categories, / 25")]
SHOTS2 = [1, 6, 8, 10, 11, 15]   # deck slides
SHOTS3 = [("post", "Day 1 post"), ("bus1", "How to Bus"), ("keys", "Minor key map"), ("keyref", "Key reference"),
          ("cue", "Cue sheet"), ("checklist", "Checklist"), ("key", "Answer key"), ("example", "Worked score")]


# ---------------------------------------------------------------- render
def node_env():
    env = dict(os.environ)
    if not env.get("NODE_PATH"):
        env["NODE_PATH"] = subprocess.run(["npm", "root", "-g"], capture_output=True, text=True).stdout.strip()
    return env


def render(jobs):
    path = os.path.join(BUILD, "jobs.json")
    with open(path, "w") as f:
        json.dump(jobs, f)
    subprocess.run(["node", RENDER_JS, path], check=True, env=node_env())


def write(name, txt):
    p = os.path.join(BUILD, name)
    with open(p, "w") as f:
        f.write(txt)
    return p


def pdf_png(src, page_no, out, res=110):
    subprocess.run(["pdftoppm", "-png", "-r", str(res), "-f", str(page_no), "-l", str(page_no), "-singlefile", src,
                    os.path.join(BUILD, out)], check=True)


def art_doc(w, h, inner):
    return (f'<!doctype html><html><head><meta charset="utf-8"><style>{font_css()}html,body{{margin:0;background:#fff}}</style></head>'
            f'<body><div style="width:{w}px;height:{h}px;overflow:hidden">{inner}</div></body></html>')


def deck_art():
    """PNG art for the deck (original drawings, rendered with the brand fonts)."""
    kb = lambda k, ms, col: (f'<div style="width:700px"><div style="font-family:Oswald;font-weight:700;font-size:48px;text-transform:uppercase;'
                             f'color:{col}">{k}</div><div style="font-family:Oswald;font-weight:300;font-size:28px;margin:2px 0 12px">'
                             f'{ms} · chord {" ".join(chord(k))}</div>{keyboard_svg(k, octaves=1)}'
                             f'<div style="font-family:\'Source Sans 3\';font-size:26px;margin-top:10px">Scale: {" ".join(KEYS[k])}</div></div>')
    keys_html = (f'<div style="display:flex;gap:150px;padding:10px 0 0 25px">{kb(KEY1, "white keys only", "#000")}{kb(KEY2, "one flat", ACCENT)}</div>'
                 f'<div style="font-family:\'Source Sans 3\';font-weight:600;font-size:28px;margin-top:22px">Optional: any key, or none. Orange = the first '
                 f'chord. Dots = the rest of the scale. Half steps rub: B–C and E–F.</div>')
    W = WOODS_WIDE
    arts = {"woods": (1500, round(1500 * W["vh"] / W["vw"]), woods_svg(**W)),
            "bus": (1800, 800, bus_svg()),
            "keys": (1600, 640, keys_html),
            "score": (2000, 680, score_svg())}
    jobs, out = [], {}
    for k, (w, h, inner) in arts.items():
        p = os.path.join(BUILD, f"deck-{k}.png")
        jobs.append({"html": write(f"deck-{k}.html", art_doc(w, h, inner)), "png": p, "width": w, "height": h, "scale": 1})
        out[k] = (p, (w, h))
    return jobs, out


def tpt_json():
    return {
        "price": PRICE, "license_price": PRICE,
        "tax_code": "Other Digital Goods - No Physical Media",
        "grades": ["9th-grade", "10th-grade", "11th-grade", "12th-grade"],
        "subjects": ["Music"],
        "tags": ["Halloween", "Autumn"],
        "formats": ["PDF"],
        "pages": len(PLAN) + DECK_SLIDES,
        "answer_key": "Included",
        "zip_contents": [os.path.basename(PDF), os.path.basename(DECK), os.path.basename(DECK_PDF)],
        "files": {"product": os.path.basename(ZIP), "preview": os.path.basename(PREVIEW_PDF), "thumb1": "cover.png",
                  "thumb2": "preview-1.png", "thumb3": "preview-2.png", "thumb4": "preview-3.png"},
    }


def make_zip():
    if os.path.exists(ZIP):
        os.remove(ZIP)
    with zipfile.ZipFile(ZIP, "w", zipfile.ZIP_DEFLATED) as z:
        for p in (PDF, DECK, DECK_PDF):
            zi = zipfile.ZipInfo(f"{ZIP_DIR}/{os.path.basename(p)}", date_time=(2026, 10, 12, 12, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            with open(p, "rb") as fh:
                z.writestr(zi, fh.read())


def build():
    os.makedirs(BUILD, exist_ok=True)
    problems = verify_data()
    if problems:
        print("\n".join(problems))
        sys.exit("data check failed; not rendering")
    full = write("resource.html", html_doc([render_page(p) for p in PLAN], TITLE))
    prev = write("preview.html", html_doc([render_page(p, watermark=True) for p in PREVIEW_PICKS], "Preview"))
    art_jobs, art = deck_art()
    render([{"html": full, "pdf": PDF}, {"html": prev, "pdf": PREVIEW_PDF}] + art_jobs)
    build_deck(art)
    deck_pdf()
    make_zip()
    res = 45 if LOW_RES else 120
    for key, _, _ in SHOTS1:
        pdf_png(PDF, IX[key], key, res)
    for key, _ in SHOTS3:
        pdf_png(PDF, IX[key], key, 40 if LOW_RES else 80)
    for s in SHOTS2:
        pdf_png(DECK_PDF, s, f"slide{s}", 30 if LOW_RES else 70)
    sqs = [write("cover.html", cover_square()), write("p1.html", preview_square_1([(f"{k}.png", a, b) for k, a, b in SHOTS1])),
           write("p2.html", preview_square_2([f"slide{s}.png" for s in SHOTS2])),
           write("p3.html", preview_square_3([(f"{k}.png", lab) for k, lab in SHOTS3]))]
    render([{"html": h, "png": o, "width": 1000, "height": 1000, "scale": 1 if LOW_RES else 2} for h, o in zip(sqs, [COVER_PNG] + PREVIEW_PNGS)])
    with open(TPT_JSON, "w") as f:
        json.dump(tpt_json(), f, indent=2, ensure_ascii=False)
        f.write("\n")
    print(f"built {PDF} ({len(PLAN)} pages), {DECK} + PDF ({DECK_SLIDES} slides), {ZIP}, preview, cover + 3 previews, tpt.json"
          + (" [LOW RES]" if LOW_RES else ""))


# ---------------------------------------------------------------- checks
def covered(m):
    return any(a <= m <= z for t in TRACKS for a, z in t[7])


def verify_data():
    errs = []
    words = {"white keys only": 0, "one": 1, "two": 2, "three": 3, "four": 4, "five": 5}
    for k, s in KEYS.items():
        root = pc(s[0])
        if [(pc(n) - root) % 12 for n in s] != [sum(MINOR_STEPS[:i]) for i in range(7)]:
            errs.append(f"{k}: spelled notes are not a natural minor scale")
        if [n[0] for n in s] != [("ABCDEFG" * 2)[("ABCDEFG".index(s[0][0]) + i)] for i in range(7)]:
            errs.append(f"{k}: letters do not run in order (spelling)")
        if [(pc(n) - root) % 12 for n in chord(k)] != [0, 3, 7]:
            errs.append(f"{k}: first chord is not a minor triad")
        acc = sum(1 for n in s if len(n) > 1)
        lead = KEY_NOTES[k].split(" (")[0].split(";")[0]
        if words.get(lead if lead in words else lead.split()[0], -1) != acc:
            errs.append(f"{k}: '{KEY_NOTES[k]}' does not match {acc} accidentals")
        if k != f"{s[0]} minor":
            errs.append(f"{k}: name does not match its root")
    if MINOR_STEPS[1] != 1 or MINOR_STEPS[4] != 1 or sum(MINOR_STEPS) != 12:
        errs.append("minor steps: half steps must sit between notes 2-3 and 5-6")
    if (pc("C") - pc("B")) % 12 != 1 or (pc("F") - pc("E")) % 12 != 1:
        errs.append("A minor half steps are not B-C and E-F")
    a, d, bb = ({pc(n) for n in KEYS[x]} for x in ("A minor", "D minor", "Bb minor"))
    if len(a & d) != 6:
        errs.append("pair A/D minor: should share six of seven notes")
    if {(p + 1) % 12 for p in a} != bb:
        errs.append("pair A/Bb minor: Bb minor should be A minor a half step up")
    if abs(SPM - 1.875) > 1e-9 or abs(MEASURES * SPM - 60) > 1e-9 or bpm_for(60) != BPM:
        errs.append("timing: 32 measures at 128 BPM must be 60 s, 1.875 s per measure")
    for s, m in CUE_PRACTICE:
        if measure_at(s) != m:
            errs.append(f"cue practice: {s}s is measure {measure_at(s)}, key says {m}")
    for s, b in BPM_PRACTICE:
        if round(bpm_for(s)) != b:
            errs.append(f"bpm practice: {s}s -> {bpm_for(s):.1f}, key says {b}")
    if tfmt(28.125) != "0:28.1" or tfmt(7.5) != "0:07.5" or tfmt(60) != "1:00":
        errs.append("time formatting")
    if not all(1 <= m <= MEASURES for m, _ in MOMENTS) or KEY_CHANGE not in [m for m, _ in MOMENTS] or SILENCE[0] not in [m for m, _ in MOMENTS]:
        errs.append("moments: out of range, or missing the key change / the silence")
    for m in range(1, MEASURES + 1):
        if covered(m) == (m in SILENCE):
            errs.append(f"example score: measure {m} {'should be silent' if m in SILENCE else 'is empty'}")
    if len(SILENCE) * SPM >= 5:
        errs.append("example silence is 5 s or longer")
    srcs = [t[1] for t in TRACKS]
    rev = sum(t[4] is not None for t in TRACKS)
    dly = sum(t[5] is not None for t in TRACKS)
    if (srcs.count("library"), srcs.count("recording"), srcs.count("download"), rev, dly) != (3, 2, 1, 6, 2) or len(TRACKS) < 5:
        errs.append("example: counts changed; update the 'requirement check' line (3 library, 2 recordings, 1 download, 6 reverb, 2 delay)")
    for t in TRACKS:
        if t[0] == "Low drone" and not (KEY1[0] in t[3] and KEY2[0] in t[3] and t[7][1][0] == KEY_CHANGE):
            errs.append("example: the drone must move from A to D on the visual cut at measure 17 (the example's optional key change)")
        if t[0] == "Low hum" and (pc("F") - pc("E")) % 12 != 1:
            errs.append("example: low hum half step")
        for v in (t[4], t[5]):
            if v is not None and not -40 <= v <= 0:
                errs.append(f"example: send level {v} dB out of range")
    sends = {t[0]: t[4] for t in TRACKS}
    for nm, v in (("Low drone", -12), ("Footsteps", -30), ("Wind", -6)):
        if sends[nm] != v:
            errs.append(f"bus diagram shows {nm} at {v} dB but the example says {sends[nm]}")
    counts = dict(vocab=(VOCAB, 14), goals=(GOALS, 6), req=(REQUIREMENTS, 12), check=(CHECKLIST, 12), steps=(STEPS, 6), tips=(TIPS, 6),
                  demo=(DEMO, 8), part1=(BUS_PART1, 7), part2=(BUS_PART2, 6), trouble=(TROUBLE, 5), atmos=(ATMOS, 5), sources=(SOURCES, 3),
                  picture=(PICTURE, 5), struggles=(STRUGGLES, 6))
    for nm, (lst, nexp) in counts.items():
        if len(lst) != nexp:
            errs.append(f"{nm}: {len(lst)} items, expected {nexp}")
    if len(RUBRIC) != 5 or any(len(dd) != len(LEVELS) for _, dd in RUBRIC) or 5 * len(RUBRIC) != 25 or len(RUBRIC_SHORT) != 5:
        errs.append("rubric must be 5 categories x levels 2-5 = /25")
    if len(DAYS) != 6 or [dd["n"] for dd in DAYS] != list(range(1, 7)) or len(DECK_DAYS) != 6:
        errs.append("pacing must be 6 days (and 6 on the deck)")
    for dd in DAYS:
        blocks = [tuple(map(int, x.split("–"))) for x, _, _ in dd["blocks"]]
        if blocks[0][0] != 0 or blocks[-1][1] != 53 or any(blocks[i][1] != blocks[i + 1][0] for i in range(len(blocks) - 1)):
            errs.append(f"Day {dd['n']} minute blocks do not run 0-53 without gaps")
    if len(SLIDES) < 12 or any(len(s[3].strip()) < 80 for s in SLIDES):
        errs.append("deck needs 12+ slides, each with real speaker notes")
    if len(PLAN) != len(set(PLAN)):
        errs.append("plan has a repeated page")
    # owner decision (12 Oct): no key is required. Requirements, checklist, goals and rubric must not ask for one.
    for nm, lst in (("requirements", REQUIREMENTS), ("checklist", CHECKLIST), ("goals", GOALS), ("deck checklist", [SLIDES[13][2]]),
                    ("rubric", [d for _, ds in RUBRIC for d in ds] + RUBRIC_SHORT)):
        for item in lst:
            if re.search(r"\bminor\b|\bkeys\b|key change|\bkey [12]\b|measures? 17", item, re.I):
                errs.append(f"{nm}: asks for a key ('{item[:60]}'); no key is required (owner, 12 Oct)")
    if RUBRIC[1][1][3] != ("Music serves the picture precisely. Key moments are acknowledged. Silence is used as a compositional tool. "
                           "The piece feels like it was written for this specific video."):
        errs.append("rubric: Composition & Pacing level 5 must keep the owner's wording")
    return errs


def pdf_fonts(path):
    out = subprocess.run(["pdffonts", path], capture_output=True, text=True, check=True).stdout.splitlines()[2:]
    return [(ln.split()[0], ln.split()[-5]) for ln in out]


def pdf_text(path):
    return subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True, check=True).stdout


def deck_texts():
    from pptx import Presentation
    prs = Presentation(DECK)
    slides, notes = [], []
    for s in prs.slides:
        slides.append(" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame))
        notes.append(s.notes_slide.notes_text_frame.text if s.has_notes_slide else "")
    media = [p for p in zipfile.ZipFile(DECK).namelist() if p.startswith("ppt/media/")]
    return prs, slides, notes, media


def verify_files():
    errs = []
    from pypdf import PdfReader
    from PIL import Image
    for p in [PDF, DECK, DECK_PDF, ZIP, PREVIEW_PDF, COVER_PNG, *PREVIEW_PNGS, UPLOAD, TPT_JSON]:
        if not os.path.exists(p):
            errs.append(f"missing {os.path.basename(p)}")
    if errs:
        return errs
    for path, n_exp in ((PDF, len(PLAN)), (PREVIEW_PDF, len(PREVIEW_PICKS))):
        rd = PdfReader(path)
        if len(rd.pages) != n_exp:
            errs.append(f"{os.path.basename(path)}: {len(rd.pages)} pages, expected {n_exp}")
        for i, pg_ in enumerate(rd.pages):
            w, h = float(pg_.mediabox.width), float(pg_.mediabox.height)
            if abs(w - 612) > 1 or abs(h - 792) > 1:
                errs.append(f"{os.path.basename(path)} p{i + 1}: {w:.0f}x{h:.0f}pt, not US Letter")
    if len(PdfReader(PREVIEW_PDF).pages) != 4:
        errs.append("preview PDF must be 4 pages")
    drd = PdfReader(DECK_PDF)
    if len(drd.pages) != DECK_SLIDES:
        errs.append(f"deck PDF: {len(drd.pages)} pages, expected {DECK_SLIDES}")
    for i, pg_ in enumerate(drd.pages):
        if abs(float(pg_.mediabox.width) / float(pg_.mediabox.height) - 16 / 9) > 0.01:
            errs.append(f"deck PDF p{i + 1}: not 16:9")
    for path in (PDF, PREVIEW_PDF, DECK_PDF):
        fonts = pdf_fonts(path)
        for name, emb in fonts:
            if emb != "yes":
                errs.append(f"{os.path.basename(path)}: font {name} not embedded")
            if not re.search(r"oswald|sourcesans", name.lower().replace(" ", "")):
                errs.append(f"{os.path.basename(path)}: unexpected font {name} (fallback?)")
        names = " ".join(n for n, _ in fonts).lower()
        for fam in ("oswald", "sourcesans3"):
            if fam not in names:
                errs.append(f"{os.path.basename(path)}: {fam} missing")
    # deck: slide count, speaker notes on every slide, only our 4 drawn PNGs as media
    prs, stext, notes, media = deck_texts()
    if len(prs.slides) != DECK_SLIDES:
        errs.append(f"deck has {len(prs.slides)} slides, expected {DECK_SLIDES}")
    for i, nt in enumerate(notes):
        if len(nt.strip()) < 80:
            errs.append(f"deck slide {i + 1}: speaker notes missing or too short")
    if any(not m.lower().endswith(".png") for m in media) or len(media) != 4:
        errs.append(f"deck media should be exactly the 4 drawn PNGs (woods, bus, keys, score), found {media}")
    if abs(prs.slide_width / prs.slide_height - 16 / 9) > 0.01:
        errs.append("deck is not 16:9")
    alldeck = " ".join(stext + notes)
    for s in (LISTEN1, LISTEN2):
        for part in s:
            if part.lower() not in alldeck.lower():
                errs.append(f"deck must name {part}")
    for i in (1, 2):
        if "YouTube or Spotify" not in stext[i]:
            errs.append(f"deck slide {i + 1}: listen slide must say YouTube or Spotify")
    if "YouTube or Spotify" not in notes[0]:
        errs.append("deck slide 1 notes must say YouTube or Spotify")
    names = zipfile.ZipFile(ZIP).namelist()
    for p in (PDF, DECK, DECK_PDF):
        n = f"{ZIP_DIR}/{os.path.basename(p)}"
        if n not in names:
            errs.append(f"zip is missing {os.path.basename(p)}")
        elif zipfile.ZipFile(ZIP).read(n) != open(p, "rb").read():
            errs.append(f"zip copy of {os.path.basename(p)} is stale; rebuild")
    if len(names) != 3:
        errs.append(f"zip should hold exactly 3 files, has {names}")
    rd = PdfReader(PDF)
    pages = [subprocess.run(["pdftotext", "-layout", "-f", str(i + 1), "-l", str(i + 1), PDF, "-"], capture_output=True, text=True,
                            check=True).stdout for i in range(len(rd.pages))]
    flat = [re.sub(r"\s+", " ", t) for t in pages]
    if not any(re.fullmatch(r"\s*SCORING TO PICTURE IN LOGIC PRO\s*", ln) for ln in pages[0].splitlines()):
        errs.append("PDF cover: 'SCORING TO PICTURE IN LOGIC PRO' is not on one line (orphaned word?)")
    for i, t in enumerate(flat):
        if i > 0 and "Brian Lotze · Hudson Beat" not in t:
            errs.append(f"page {i + 1}: footer missing")
        if i > 0 and not re.search(rf"Page {i + 1}\s*$", t.strip()):
            errs.append(f"page {i + 1}: footer page number wrong (overflow?)")
    for m in re.finditer(r"\b(?:pages?|pp?\.) (\d+)(?:–(\d+))?", " ".join(flat[1:])):
        for g in m.groups():
            if g and not 1 <= int(g) <= len(PLAN):
                errs.append(f"cross-reference to page {g} does not exist")
    rub = flat[IX["rubric"] - 1]
    if f"/ {5 * len(RUBRIC)}" not in rub:
        errs.append("rubric page: /25 total missing")
    for c, ds in RUBRIC:
        for d in ds:
            if not all(w in rub for w in re.findall(r"[A-Za-z]{6,}", d)):
                errs.append(f"rubric text missing: {d[:50]}")
    for w, _ in VOCAB:
        if w not in flat[IX["vocab"] - 1]:
            errs.append(f"vocab page missing {w}")
    for k, s in KEYS.items():
        if " ".join(s) not in flat[IX["keyref"] - 1]:
            errs.append(f"key reference missing scale for {k}")
    for dd in DAYS:
        if f"DAY {dd['n']}" not in flat[IX["pacing"] - 1]:
            errs.append(f"pacing page missing Day {dd['n']}")
    # no output may still say two keys are required (owner, 12 Oct). Phrases split so this file never trips its own scan.
    no_key_rule = [r"two (different )?minor " + "keys?", r"\b2 minor " + "keys", r"\bkey " + r"[12]\b", r"two " + "keys", r"one " + "musical rule",
                   r"choose your two", r"pick two minor", r"second minor " + "key"]
    for label, txt in (("resource PDF", " ".join(flat)), ("preview PDF", pdf_text(PREVIEW_PDF)), ("deck PDF", pdf_text(DECK_PDF)),
                       ("deck pptx", alldeck), ("UPLOAD.md", open(UPLOAD).read()), ("tpt.json", open(TPT_JSON).read())):
        flat_txt = re.sub(r"\s+", " ", txt)
        for pat in no_key_rule:
            for mm in re.finditer(pat, flat_txt, re.I):
                errs.append(f"{label}: still states a key rule: '...{flat_txt[max(0, mm.start() - 30):mm.end() + 30]}...'")
    ptxt = [pg.extract_text() for pg in PdfReader(PREVIEW_PDF).pages]
    if sum("PREVIEW" in t for t in ptxt) != len(PREVIEW_PICKS):
        errs.append("preview PDF: watermark missing on some pages")
    # banned strings: every output, the deck XML, PDF metadata, and every text file in this folder (gen.py included)
    blobs = {"resource PDF": pdf_text(PDF), "preview PDF": pdf_text(PREVIEW_PDF), "deck PDF": pdf_text(DECK_PDF), "deck pptx": alldeck,
             "zip names": " ".join(names)}
    with zipfile.ZipFile(DECK) as z:
        blobs["deck XML"] = " ".join(z.read(n).decode("utf8", "ignore") for n in z.namelist() if n.endswith(".xml") or n.endswith(".rels"))
    for path in (PDF, PREVIEW_PDF, DECK_PDF):
        meta = PdfReader(path).metadata or {}
        blobs[f"{os.path.basename(path)} metadata"] = " ".join(str(v) for v in meta.values())
    for fn in os.listdir(HERE):
        if fn == "QA.md":   # QA's own review record (not shipped) names the strings it searched for
            continue
        if fn.endswith((".md", ".json", ".py", ".txt")) or fn in ("release", "uploaded"):
            blobs[fn] = open(os.path.join(HERE, fn), encoding="utf8", errors="ignore").read()
    for label, txt in blobs.items():
        low = txt.lower()
        for b in BANNED:
            if b.lower() in low:
                errs.append(f"{label}: contains banned string '{b}'")
    for p in [COVER_PNG, *PREVIEW_PNGS]:
        if Image.open(p).size != (2000, 2000):
            errs.append(f"{os.path.basename(p)}: not 2000x2000 (low-res build?)")
    tj = json.load(open(TPT_JSON))
    if tj != tpt_json():
        errs.append("tpt.json is stale; rebuild")
    if tj["price"] != PRICE or tj["license_price"] != PRICE or tj.get("tax_code") != "Other Digital Goods - No Physical Media":
        errs.append("tpt.json price/license_price/tax_code")
    if tj["grades"] != ["9th-grade", "10th-grade", "11th-grade", "12th-grade"] or tj["subjects"] != ["Music"]:
        errs.append("tpt.json grades 9-12 / subject Music")
    if not (1 <= len(tj["tags"]) <= 6 and len(tj["formats"]) <= 3):
        errs.append("tpt.json tags/formats counts")
    for k, f in tj["files"].items():
        if not os.path.exists(os.path.join(HERE, f)):
            errs.append(f"tpt.json file {k} missing: {f}")
    if tj["pages"] != len(rd.pages) + len(prs.slides):
        errs.append("tpt.json pages != PDF pages + slides")
    up = open(UPLOAD).read()
    m = re.search(r"^## 1\. Title\s*\n+```\n(.+?)\n```\n\((\d+) characters", up, re.M | re.S)
    if not m:
        errs.append("UPLOAD.md: title block or its character count not found")
    else:
        n = len(m.group(1).strip())
        if n > 80:
            errs.append(f"UPLOAD.md: title is {n} chars (max 80)")
        if int(m.group(2)) != n:
            errs.append(f"UPLOAD.md: says {m.group(2)} characters, title has {n}")
    m2 = re.search(r"^## 2\. Description\s*\n+```\n(.+?)\n```", up, re.M | re.S)
    if not m2 or not m2.group(1).rstrip().endswith(DISCLOSURE):
        errs.append("UPLOAD.md: description must end with the disclosure line")
    else:
        desc = m2.group(1)
        for need in (f"{len(PLAN)}-page", f"{DECK_SLIDES}-slide", "Logic Pro", "/25", "No video", "YouTube or Spotify"):
            if need not in desc:
                errs.append(f"UPLOAD.md description: missing '{need}'")
        if re.search(r"  +\S|\s[,.]\s|\bteh\b|\s{2,}$", desc, re.M):
            errs.append("UPLOAD.md description: double space or stray punctuation")
    for need in (f"${PRICE}", f"{len(PLAN)} pages", os.path.basename(ZIP), os.path.basename(DECK), os.path.basename(DECK_PDF)):
        if need not in up:
            errs.append(f"UPLOAD.md: missing '{need}'")
    for f in ("release", "uploaded"):
        if os.path.exists(os.path.join(HERE, f)):
            errs.append(f"{f} exists: only the Manager / uploader writes it")
    return errs


def check():
    errs = verify_data() + verify_files()
    if errs:
        print("CHECK FAILED")
        print("\n".join(" - " + e for e in errs))
        sys.exit(1)
    print(f"CHECK OK: 8 minor keys spelled and checked (scale W-H-W-W-H-W-W, chord 0-3-7, accidental counts); 32 measures x 1.875 s = 60 s at 128 BPM, "
          f"cue/BPM practice answers; worked score covers every measure but the 1 deliberate silent one, sends and source counts match the text and the bus "
          f"diagram; no key required anywhere (requirements, checklist, rubric, PDF, deck, UPLOAD.md); rubric 5 x (2-5) = /25 printed; 6-day blocks run 0-53; {len(PLAN)} Letter pages with footers and valid cross-references, cover "
          f"title on one line; preview 4 watermarked pages; deck {DECK_SLIDES} slides 16:9 with speaker notes on every slide, only 4 drawn PNGs, PDF copy "
          f"{DECK_SLIDES} pages; only Oswald/Source Sans 3, all embedded; no banned strings (school name, LMS, the classroom video's title) in any output "
          f"or file here; zip holds current PDF + pptx + deck PDF; images 2000x2000; tpt.json and UPLOAD.md ok.")


if __name__ == "__main__":
    if "--fallback" in sys.argv:   # self-check: the deck rendered without Oswald / Source Sans 3
        deck_pdf(DECK, os.path.join(BUILD, "deck-fallback.pdf"), fonts=False)
        print("wrote build/deck-fallback.pdf")
    elif "--check" in sys.argv:
        check()
    else:
        build()
