#!/usr/bin/env python3
"""DMP1 Unit 2: Stems & Grooves (TpT 013): build the teacher resource, the intro deck and the listing files.

    python3 gen.py            build everything into this folder
    python3 gen.py --lowres   same, but square images at scale 1 and small page shots (self-check only; --check rejects them)
    python3 gen.py --check    verify the drum patterns, transitions, section maps, stem recipes, data, files, page/slide counts,
                              day slots, speaker notes, fonts, label fit and overlap, page overflow, banned strings, tpt.json and UPLOAD.md

The owner's own curriculum (Brian Lotze, Digital Music Production 1), reformatted in the Hudson Beat style: a 6-day drum unit in
Logic Pro where students write drums over a stem (a recording with the drums taken out). Source: the owner's Drive folder
"2 - Stems & Grooves": Project Sheet & Rubric (2 Oct 2026: Part 1 Foundation, Part 2 Full Song, rubric 5 categories scored 2-5
= /25), Teacher Notes (6-day pacing, struggles, variations), Day 1 Classroom Post; and the DMP1 Restructure, Unit 2 (6 days,
1/3/2, mixing concept: gain staging). The kit follows the newest project sheet. The Restructure's older description (Group 1:
AABA over a 16-bar stem; Group 2: full-song drums with reverb) is kept as: the Foundation stem is 16 measures of 4-measure
sections (the Day 1 post), shown in AABA, and the reverb bus is an optional, ungraded extension (Day 4). Flagged to the owner.

No stems, audio or Logic sessions are included (the owner's stems are third-party). The kit says so, lists where a teacher can
find stems they may use, and gives a recipe for making two stems in Logic (verified below: chords, bass roots, section lengths,
durations).

Owner text is kept word for word except: the LMS name ("your class page"), the Day 1 post's "Check the example!" link (the owner's
Logic session, not included: "Your teacher will play an example of each"), a closing parenthesis in Part 2 step 1, and the day
plans' 5-minute save blocks (Days 1 and 3 stop at minute 48 in the owner's notes; Day 4 has no minute blocks). All disclosed on
teacher notes page 1. Owner inconsistencies kept as written and flagged: the checklist names both projects "Your Name -
Foundations" (the requirements say Foundation and Full Song); the Day 1 post's "A, B, A, B, etc."; "a real professional
recording" when the teacher supplies the stem.

Every drum pattern, transition, section map, stem chord and duration in the kit is computed or checked by verify_music().

Outputs: <slug>.pdf (teacher + student resource), <slug>-intro-deck.pptx and its PDF copy, a ZIP of those three
(the TpT product file), a 4-page watermarked preview PDF, cover.png and preview-1..3.png (2000 x 2000).
Renders with Playwright + Chromium (../tools/render.js) and LibreOffice (deck PDF), fonts from ../fonts. $0 AI.
"""
import json
import math
import os
import re
import shutil
import subprocess
import sys
import zipfile
from html import escape

HERE = os.path.dirname(os.path.abspath(__file__))
TPT = os.path.dirname(HERE)
FONTS = os.path.join(TPT, "fonts")
BUILD = os.path.join(HERE, "build")
RENDER_JS = os.path.join(TPT, "tools", "render.js")

SLUG = "stems-and-grooves"
PDF = os.path.join(HERE, f"{SLUG}.pdf")
DECK = os.path.join(HERE, f"{SLUG}-intro-deck.pptx")
DECK_PDF = os.path.join(HERE, f"{SLUG}-intro-deck.pdf")
ZIP = os.path.join(HERE, f"{SLUG}.zip")
PREVIEW_PDF = os.path.join(HERE, f"{SLUG}-PREVIEW.pdf")
COVER_PNG = os.path.join(HERE, "cover.png")
PREVIEW_PNGS = [os.path.join(HERE, f"preview-{i}.png") for i in (1, 2, 3)]
UPLOAD = os.path.join(HERE, "UPLOAD.md")
TPT_JSON = os.path.join(HERE, "tpt.json")
OVERFLOW_JSON = os.path.join(BUILD, "overflow.json")
ZIP_DIR = "Stems and Grooves - Hudson Beat"

STORE = "Brian Lotze"      # copyright holder and TpT seller
BRAND = "Hudson Beat"      # store line on covers and previews (owner, 4 Oct)
YEAR = 2026
ACCENT = "#1E8A50"         # this product's single accent: green (008 red, 009 orange, 010 violet, 011 teal, 012 blue)
ACC_LIGHT = "#BFE8D0"      # light tint of the accent, for text on black in the deck
CLIP_RED = "#D7262B"       # the meter's clip zone only (colour is the content there; the word RED is printed beside it)
TITLE = "Stems & Grooves"
TITLE_H = escape(TITLE)
PRICE = "10.00"
UNIT = 2
DAYS_N = 6
MINUTES = 53
LOW_RES = "--lowres" in sys.argv
DISCLOSURE = "Designed with the help of digital and AI tools, and checked by hand."
LISTING_TITLE = "Drum Fills in Logic Pro | Stems & Transitions | High School Music Technology"
NAME1 = "“Your Name - Foundation”"            # the owner's requirement wording
NAME2 = "“Your Name - Full Song”"
NAME_CHECK = "“Your Name - Foundations”"      # the owner's checklist wording, both projects (kept; flagged)

# Strings that must never appear in any output or in any file in this folder (school branding, LMS, the owner's third-party stems).
# Split so this file never contains them.
BANNED = ["Blind" + " Brook", "BB" + "HS", "blind" + "brook", "Google" + " Classroom", "Lan" + "tieri", "Chard" + "onnay", "Koi" + "shii",
          "Gal" + "ias", "Saud" + "ade", "Virus" + " En Mi", "Ben" + " Flowers", "Baum" + "Xmedia"]

# ================================================================ drum patterns (16 steps = one 4/4 measure in sixteenths)
LANES = ["Crash", "Open hi-hat", "Closed hi-hat", "Snare", "High tom", "Low tom", "Kick"]   # top to bottom, as on the drawings
STEPS = 16


def pat(**rows):
    """A one-measure pattern: lane name (underscores for spaces) -> 'x...' string of 16 steps."""
    out = {}
    for k, v in rows.items():
        out[{"open": "Open hi-hat", "hat": "Closed hi-hat", "snare": "Snare", "hitom": "High tom", "lotom": "Low tom", "kick": "Kick",
             "crash": "Crash"}[k]] = v
    return out


# Worked example (Hudson Beat, original). A: laid back, closed hi-hat eighths. B: four-on-the-floor kick, open hi-hats off the beat.
PATTERN_A = pat(hat="x.x.x.x.x.x.x.x.", snare="....x.......x...", kick="x.......x.x.....")
PATTERN_B = pat(open="..x...x...x...x.", hat="x...x...x...x...", snare="....x.......x...", kick="x...x...x...x...")
# The four transitions, each the measure right before a section change (the measure before it is PATTERN_A).
TRANS = {
    "Drum fill": pat(hat="x.x.x.x.........", snare="....x...xxxx....", hitom="............xx..", lotom="..............xx",
                     kick="x.......x......."),
    "Breakdown": pat(kick="x...x...x...x..."),
    "Cut": pat(hat="x.x.x.xx........", snare="....x...........", kick="x.....x........."),
    "Build": pat(hat="........xxxxxxxx", snare="....x...x.x.xxxx", hitom="............x.x.", kick="x...x...x.x.xxxx"),
}
TYPES = list(TRANS)        # fill, breakdown, cut, build (the owner's order)
# "Name that transition" (student quick check): new examples, shuffled; the answers are computed by classify().
QUIZ = [pat(snare="x.x.x.x.........", hitom="........x.x.....", lotom="............x.x.", kick="x..............x"),
        pat(kick="x...x...x...x.x.", snare="........x.x.xxxx", hat="............xxxx"),
        pat(hat="x.x.x.x.x.xx....", snare="....x...........", kick="x.......x.x....."),
        pat(hat="xxxxxxxxxxxxxxxx")]
# The owner's four transition types (Day 1 post box and project sheet), word for word.
TYPE_DEF = {"Drum fill": "a rhythmic variation in the last 1–2 measures that signals a change is coming",
            "Breakdown": "strip to one element (kick only, hi-hat only) before the new section hits",
            "Cut": "complete silence for 1–2 beats before the next section drops",
            "Build": "add elements progressively in the last measure to increase energy into the change"}
TYPE_WHEN = {"Drum fill": "the classic: works almost anywhere, especially into a bigger section",
             "Breakdown": "before a section that should hit harder, or as the music thins out",
             "Cut": "the most powerful and the easiest to do badly; a cut before a quiet verse",
             "Build": "a build before a big chorus"}


def lanes_of(*measures):
    used = {ln for m in measures for ln, s in m.items() if "x" in s}
    return [ln for ln in LANES if ln in used]


def beat_lanes(m):
    """Number of lanes that play in each of the 4 beats."""
    return [sum(1 for s in m.values() if "x" in s[b * 4:b * 4 + 4]) for b in range(4)]


def trailing_silence(m):
    n = 0
    for i in range(STEPS - 1, -1, -1):
        if any(s[i] == "x" for s in m.values()):
            break
        n += 1
    return n


def same(a, b):
    return all(a.get(ln, "." * STEPS) == b.get(ln, "." * STEPS) for ln in LANES)


def classify(m, groove=PATTERN_A):
    """Which of the owner's four transition types a measure is (checked against the definitions):
    breakdown = one element only; cut = silence for the last 1-2 beats (4-8 sixteenths) after something played;
    build = lanes enter beat by beat (never fewer, more at the end, starting from at most one); fill = any other change to the groove."""
    act = lanes_of(m)
    if len(act) == 1:
        return "Breakdown"
    t = trailing_silence(m)
    if 4 <= t <= 8 and t < STEPS:
        return "Cut"
    c = beat_lanes(m)
    if all(c[i] <= c[i + 1] for i in range(3)) and c[3] > c[0] and c[0] <= 1:
        return "Build"
    if not same(m, groove):
        return "Drum fill"
    return None


# ================================================================ the stems: two recipes the teacher can make in Logic (Hudson Beat, original)
LETTERS = "CDEFGAB"
NAT = dict(C=0, D=2, E=4, F=5, G=7, A=9, B=11)
C_MAJOR = ["C", "D", "E", "F", "G", "A", "B"]
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII"]


def triad(rn):
    """Diatonic triad in C major from a Roman numeral: (notes, quality); the case must match the quality."""
    deg = ROMAN.index(rn.upper().rstrip("°"))
    notes = [C_MAJOR[deg], C_MAJOR[(deg + 2) % 7], C_MAJOR[(deg + 4) % 7]]
    i1 = (NAT[notes[1]] - NAT[notes[0]]) % 12
    i2 = (NAT[notes[2]] - NAT[notes[1]]) % 12
    q = {(4, 3): "major", (3, 4): "minor", (3, 3): "diminished"}[(i1, i2)]
    return notes, q


def chord_name(rn):
    n, q = triad(rn)
    return n[0] + {"major": "", "minor": "m", "diminished": "dim"}[q]


TEMPO = 96   # BPM for both recipes
PROGS = {"A": ["I", "vi", "IV", "V"], "B": ["IV", "V", "iii", "vi"], "Verse": ["I", "vi", "IV", "V"], "Chorus": ["IV", "I", "V", "vi"],
         "Bridge": ["vi", "IV", "vi", "V"], "Intro": ["I", "vi", "IV", "V"], "Outro": ["IV", "V", "I", "I"]}
PARTS = {"A": ["Piano chords", "Bass"], "B": ["Piano chords", "Bass", "Pad", "Synth plays the top note of each chord"],
         "Intro": ["Pad"], "Verse": ["Piano chords", "Bass"], "Chorus": ["Piano chords", "Bass", "Pad", "Synth plays the top note of each chord"],
         "Bridge": ["Pad", "Bass"], "Outro": ["Pad"]}
FOUNDATION = [("A", 4), ("A", 4), ("B", 4), ("A", 4)]             # the Day 1 post: 16 measures, sections of 4 (AABA, as in Unit 1)
ROMAN_ORDER = ["I", "ii", "iii", "IV", "V", "vi"]
CHORDS_USED = [r for r in ROMAN_ORDER if any(r in p for p in PROGS.values())]   # every chord the two stems use, in scale order
CHORDS_USED_TEXT = ", ".join(CHORDS_USED[:-1]) + " and " + CHORDS_USED[-1]
FULL_SONG = [("Intro", 4), ("Verse", 8), ("Chorus", 8), ("Verse", 8), ("Chorus", 8), ("Bridge", 8), ("Chorus", 8), ("Outro", 4)]
LETTER = {"Intro": "i", "Verse": "a", "Chorus": "b", "Bridge": "c", "Outro": "o"}
# Worked Full Song map: one transition at every section change (the measure before the change), with the reason.
SONG_TRANS = [("Build", "the music grows from the pad into the verse"),
              ("Drum fill", "announce the first chorus"),
              ("Breakdown", "back down to the smaller verse"),
              ("Drum fill", "the same move as before: listeners expect it now"),
              ("Cut", "the bridge is quieter: silence sets it up"),
              ("Build", "a build before the big last chorus"),
              ("Breakdown", "thin out into the outro")]


def energy(sec):
    return len(PARTS[sec])


def starts(sections):
    out, m = [], 1
    for _, n in sections:
        out.append(m)
        m += n
    return out


def seconds(measures, bpm=TEMPO):
    return measures * 4 * 60 / bpm


def mmss(s):
    return f"{int(s // 60)}:{int(round(s % 60)):02d}"


# ================================================================ songs (title and artist only; find them on YouTube or Spotify)
ITAT = ("In the Air Tonight", "Phil Collins")
SONG_NOTE = "Songs are named by title and artist only. Play them from your own YouTube or Spotify access; no audio, video or lyrics are included."
STEM_NOTE = ("No stems, audio or Logic sessions are included in this kit. You supply the stems (recordings with the drums taken out): "
             "make your own in Logic with the recipe on page @MAKE, or use audio you have the right to use in class (page @STEMS).")

# ================================================================ the owner's words (Project Sheet & Rubric, Teacher Notes, Day 1 post)
OVERVIEW = ("This project has two parts. In Foundation, you will write two contrasting drum patterns and place them correctly as sections "
            "change in a real professional recording. In Full Song, you will take those patterns and focus on the moments between sections — "
            "the transitions. How do you signal to a listener that something is about to change? By the end, you will have a complete drum "
            "track that feels like it belongs to the music, not just on top of it.")
GOALS = ["Listen to a piece of music and identify where its sections change",
         "Write two contrasting drum patterns that serve the music they accompany",
         "Use transitions (fills, breakdowns, cuts, builds) to signal section changes",
         "Apply gain staging so the drum track sits correctly in the overall mix",
         "Experience what it means to respond musically to something that already exists"]
VOCAB = [("Stem track", "a single isolated instrument or group from a real recording - in this project, everything except the drums"),
         ("Section", "a distinct part of a song - verse, chorus, bridge, intro, outro - that has its own feel or energy"),
         ("A section / B section", "the two main drum patterns in this project: A for one type of section, B for another"),
         ("Transition", "the moment between two sections where the drum pattern signals a change is coming"),
         ("Drum fill", "a rhythmic variation - usually one or two measures - that breaks the pattern and announces a new section"),
         ("Breakdown", "stripping the beat down to one element (just kick, just hi-hat) before a new section hits"),
         ("Cut", "complete silence for a beat or two before the next section drops"),
         ("Build", "adding elements progressively in the last measure before a section change to increase energy"),
         ("Gain staging", "setting all track volumes so nothing clips (goes into the red) and everything is audible - the foundation of a "
          "clean mix"),
         ("Contrast", "a meaningful musical difference between two patterns - not just slightly different, but noticeably different in feel")]
REQ1 = ["Choose one stem track from the Foundation folder on your class page",
        "Listen to the full track and identify at least two distinct sections (A and B) before building anything",
        "Write an A pattern: a drum groove that fits the feel of one section type",
        "Write a B pattern: a drum groove that clearly contrasts with A in feel, energy, or rhythm",
        "Place patterns correctly so they follow the section changes in the track",
        "Minimum 3 layers per pattern (kick, snare, hi-hat at minimum)",
        "All track volumes adjusted so nothing clips and every element is audible",
        f"Project named correctly: {NAME1}",
        "Exported as MP3 and .zip (compressed .logicx file) and submitted via your class page before presenting"]
REQ2 = ["Using a new Full Song stem from your class page, identify every section change in the track",
        "At each section change, add a deliberate transition using one of the four options: drum fill, breakdown, cut, or build",
        "Every section change must have a transition… no pattern can simply loop through a boundary without acknowledgment",
        "At least two different transition types must be used across the full track",
        "Transitions should feel appropriate to the music… a cut before a quiet verse, a build before a big chorus",
        "All track volumes and panning adjusted… mix serves the stem, not the other way around",
        f"Project named correctly: {NAME2}",
        "Exported as MP3 and .zip (compressed .logicx file) and submitted via your class page before presenting"]
STEPS1 = [("Listen before you build", "Download your stem track and listen to the whole thing twice before you open Drum Machine Designer. "
           "Count measures. Notice where the energy changes. Label sections in the arrangement markers: A, B, A, B, etc. You cannot write "
           "drums for music you have not listened to."),
          ("Build your A pattern", "Create a pattern region at an A section. Build a drum groove that fits the feel of that section. At "
           "minimum: kick, snare, closed hi-hat. Aim for 3–5 total layers. Use 64-step mode if you want more control over individual measures."),
          ("Build your B pattern", "Create a new, contrasting pattern region at a B section. It should feel different from A. If A is busy, "
           "B might be stripped back. If A is laid back, B might add energy. The listener should clearly feel the difference."),
          ("Place and check", "Copy your patterns across all the correct sections. Play the full track back from beginning to end. Ask "
           "yourself: do my drums follow the music, or does the music follow my drums? It should always be the former.")]
STEPS2 = [("Map every section change (Label the sections with letters (a,b,c…)", "Go through the full track and mark every section "
           "boundary in your arrangement markers. You should know exactly how many transitions exist before you write a single one."),
          ("Choose a transition type for each boundary", "For every section change, decide which of the four transition types fits best:"),
          ("Build each transition", "Work through the section changes one at a time. Use the 64-step sequencer to control exactly what "
           "happens in the transition measure. Listen after each one before moving to the next."),
          ("Play it all the way through", "Listen to the complete track from start to finish. Every section change should feel intentional. "
           "If a transition does not add anything, it is not doing its job. Adjust until every seam in the music is earned."),
          ("Mix and submit", "Adjust all volume levels so the drums sit correctly in the mix. Nothing should overpower the stems. Name the "
           "project correctly, export MP3 and .zip, and submit via your class page before your presentation.")]
# Part 2 step 2's four bullets (project sheet wording; the Day 1 post box words them as TYPE_DEF)
STEP2_TYPES = [("Drum fill", "a rhythmic variation in the last 1–2 measures that signals the change"),
               ("Breakdown", "strip to one element (kick only, hi-hat only) before the new section"),
               ("Cut", "complete silence for 1–2 beats before the next section drops"),
               ("Build", "add elements progressively in the last measure to increase energy")]
TIPS = ["Listen first. Always. You cannot write drum parts for music you have not deeply listened to. Play it twice before you touch anything.",
        "Your A and B patterns should be obviously different. If you have to look at the screen to tell them apart, they are not different "
        "enough.",
        "A transition that works feels inevitable in hindsight. If it surprises you in a good way, it will surprise the listener too.",
        "The cut is the most powerful transition and the easiest to execute badly. Silence only works if the music earns it. Listen carefully "
        "before and after.",
        "When mixing, the stem is the star. Your drums should support it, not compete with it. If anything in your drum track is louder than "
        "the loudest element in the stem, pull it back.",
        "You do not need to fill every measure with activity. Space is musical. A simple groove that serves the track is worth more than a "
        "complicated one that fights it."]
CHECK = [("Foundation", "A and B patterns written and placed correctly across all sections"),
         ("Foundation", "Minimum 3 layers per pattern"),
         ("Foundation", "All volumes adjusted — nothing clipping, everything audible"),
         ("Foundation", f"Project named: {NAME_CHECK}"),
         ("Foundation", "Exported as MP3 and .zip (compressed .logicx file) and submitted via your class page before presenting"),
         ("Full Song", "Every section change is labeled and has a deliberate transition"),
         ("Full Song", "At least two different transition types used"),
         ("Full Song", "Transitions feel appropriate to the music at each moment"),
         ("Full Song", "Mix adjusted so drums serve the stem"),
         ("Full Song", f"Project named: {NAME_CHECK}"),
         ("Full Song", "Exported as MP3 and .zip (compressed .logicx file) and submitted via your class page before presenting"),
         (None, "Three feedback questions posted to Padlet before presenting")]
RUBRIC = [  # the owner's rubric, word for word: levels 2-5, 5 categories = /25
    ("Foundations",
     ["A and B patterns are nearly identical. Drums do not respond to the stem or follow its sections.",
      "A/B distinction exists but contrast is minimal. Patterns placed but feel independent of the music.",
      "A and B patterns clearly contrast in feel or energy. Patterns follow section changes in the stem.",
      "A and B patterns have distinct personalities and serve the music at every moment. Placement is precise and intentional."]),
    ("Full Song",
     ["No deliberate transitions. Patterns loop through section boundaries without acknowledgment.",
      "Some section changes acknowledged but transitions feel accidental or mismatched to the music.",
      "All section changes have a deliberate transition. At least two different transition types used.",
      "Every transition feels inevitable - the right type chosen for each moment. The music earns each section change."]),
    ("Technical Skill / DAW Use",
     ["Significant workflow errors. Pattern regions placed incorrectly or step sequencer used with major confusion.",
      "Patterns placed but some section changes missed or misaligned. Some workflow confusion.",
      "Patterns correctly placed across all section changes. Step sequencer used with confidence.",
      "64-step mode used for transition detail. Pattern regions managed cleanly. Workflow is confident and organized."]),
    ("Gain Staging & Mix",
     ["No meaningful volume adjustments. Drum track overwhelms or is buried under the stem.",
      "Some volume adjustment but mix is unbalanced. Drums compete with the stem rather than supporting it.",
      "Volumes adjusted so drums and stem coexist. Nothing is clipping. Mix is mostly appropriate.",
      "Drum track sits perfectly in relation to the stem. Every element is audible. Nothing clips. Mix feels like one cohesive piece of "
      "music."]),
    ("Documentation & Presentation",
     ["Project not submitted or not presented.",
      "Project submitted late or named incorrectly. Presentation made with minimal engagement.",
      "Project named correctly and submitted. Presentation made. Padlet questions posted.",
      "Project correctly named, submitted on time. Confident presentation. Three thoughtful Padlet questions posted."]),
]
LEVELS = [2, 3, 4, 5]
RUBRIC_TOTAL = 25   # checked in Drive: TOTAL / 25 (5 categories x 5)
RUBRIC_SHORT = ["A and B have distinct personalities; placement precise and intentional.",
                "Every transition feels inevitable: the right type for each moment.",
                "64-step mode for transition detail; clean, confident workflow.",
                "Drums sit perfectly with the stem; nothing clips; one cohesive piece.",
                "Named, on time, confident, three thoughtful Padlet questions."]
# ---- Teacher Notes
WHAT = ["This unit introduces the most important listening skill in the curriculum: responding to music that already exists. Students are "
        "not building in a vacuum — they have to hear what a professional recording is doing and fit into it. That constraint forces musical "
        "decisions that free-form building never does.",
        "Foundation plants AABA thinking without naming it yet. Students are writing two contrasting patterns and placing them based on what "
        "they hear. This is form comprehension through action. Full Song introduces the concept of musical transitions — that the seam "
        "between sections is as compositionally significant as the sections themselves. This idea will echo through the EDM project in Unit 8 "
        "and every project after it.",
        "The mixing concept introduced here is gain staging — the foundation of every mix decision students will make for the rest of the "
        "year. Frame it simply: nothing in the red, everything audible, the drums serve the stem. Students encounter this as a practical "
        "problem immediately because the stem tracks have varying levels and their drums often overwhelm them."]
STRUGGLES = [
    ("Skipping the listening step", "The most common failure is students who open Drum Machine Designer before they have listened to their "
     "stem. Enforce the pre-listening requirement. Walk around Day 1 and ask every student to tell you where their sections are before they "
     "create a single pattern region. If they cannot answer, they are not ready to build."),
    ("A and B patterns that are barely different", "Students will often change one hi-hat hit and call it a B section. Push back during "
     "working days with a simple test: play A, then immediately play B. Can you feel the difference without looking? If not, it needs more "
     "contrast. A useful prompt: if A is busy and driving, what would the opposite feel like?"),
    ("Transitions that are just louder versions of the pattern", "Students who are not yet confident with drum fills will often just double "
     "the hi-hat and call it a transition. Show examples of all four transition types on Day 3 before they start Full Song. The cut and "
     "breakdown are the most underused and often the most effective."),
    ("Drums that drown out the stem", "Students push faders up instinctively. Address this with a direct demonstration: play a student "
     "project with drums too loud, then pull the drum bus down until the stem is clearly the lead. Ask: which version sounds more like a "
     "real song? The answer is usually obvious. Introduce the rule here: if you can clearly hear every drum hit and the stem is getting "
     "lost, your mix is backwards."),
    ("Not knowing where sections are", "Some stems have less obvious section changes. For students who are struggling, add arrangement "
     "markers to their Logic file before they download it so the sections are pre-labeled. This is appropriate scaffolding at this stage "
     "of the year.")]
VAR_EASY = ("Pre-label all section changes in the stem file before distributing. Allow students to complete Foundation only and skip Full "
            "Song. Accept one transition type (fill only) rather than requiring two different types.")
VAR_HARD = ("Require all four transition types to be used at least once. Challenge students to write a third pattern (a C section) for any "
            "bridge or breakdown in the stem. Ask them to write a one-sentence rationale for every transition choice on their Padlet post.")
# ---- DMP1 Restructure, Unit 2
PLAN_LEARN = ("AABA form through doing — not naming. How drums change when a section changes. Why musical continuity matters. How to listen "
              "to a piece of music and respond to it. That mixing decisions (volume, panning) are musical decisions, not technical ones.")
PLAN_WHY = ("The stems project is the most musically sophisticated thing in the early curriculum. It removes the blank-page problem while "
            "requiring genuine musical thinking. Students hear what a complete track sounds like and have to fit into it — that is a listening "
            "and responding skill that transfers to every project for the rest of the year.")
PLAN_MIX = ("Gain staging lands here because students encounter it naturally — the stem tracks have varying levels and students immediately "
            "notice when their drums are too loud or too quiet relative to the music. Introduce the concept directly: 'Never push anything "
            "into the red. It is almost always better to turn other things down than to turn one thing up.' This principle will echo through "
            "every project for the rest of the year.")
RED_RULE = "Never push anything into the red. It is almost always better to turn other things down."
DRUM_LINE = "Drums in songs tend to remain relatively consistent. That is what ties everything together and makes it sound like one song."
CLOSE = ("Your drums were written in response to music that already existed. That is what session musicians do. That is a professional "
         "skill.")
# ---- Day 1 Classroom Post
POST_INTRO = "Today you are starting Part 1 of this project: Foundation. Follow this sequence in order."
POST_BOLD = "Do not open Drum Machine Designer until you have completed Steps 1 and 2."
POST_STEPS = [
    ("Download", "Open your class page and download one stem track from the Foundation folder. Open it in Logic Pro.",
     "Pick something that catches your ear. You are going to spend a lot of time with this track."),
    ("Listen — twice", "Play the full 16 measures from beginning to end. Do not touch anything. Then play it again.",
     "What is this music doing? Where does the energy change? Where does a new section feel like it starts?"),
    ("Map your sections", "Add arrangement markers in Logic to label every section change you hear: A, B, A, B, etc. Each section should be 4 "
     "measures long.", "You cannot write drums for music you have not listened to. This step is not optional."),
    ("Build your A pattern", "Create a pattern region at an A section. Build a groove that fits the feel of that section. Start with kick, "
     "snare, and hi-hat — then make it yours.", "Does this groove fit the music? Play it alongside the stem and listen."),
    ("Start your B pattern", "Once your A pattern feels solid, create a new pattern region at a B section and start building something that "
     "contrasts with A.", "A and B should feel different. If you have to look at the screen to tell them apart, they are not different enough yet.")]
POST_EXAMPLE = "Your teacher will play an example of each."     # replaces the owner's "Check the example!" link (a Logic session)

# ================================================================ pacing: 6 days of about 53 minutes (Teacher Notes; Restructure 1/3/2)
# block = (minutes, heading, text, deck slides (first, last) or None). Every deck slide belongs to exactly one block (checked).
DAYS = [
    {"n": 1, "kind": "Learning", "title": "Learning + Foundation start",
     "goal": "Every student has a stem downloaded, sections labeled, and an A pattern started.",
     "blocks": [("0–5", "The stem, no drums", "Slides 1–2. Play 60 seconds of one of the stem tracks without the drums. Ask: what is this "
                 "music doing? Where does the energy change? Do not give answers yet.", (1, 2)),
                ("5–15", "Introduce the project", f"Slides 3–6. Explain Foundation and Full Song. Show the four transition types with audio "
                 f"examples ({ITAT[0]}, then the four you programmed from page @KEY2). Keep this brief — students will understand "
                 "transitions better once they have built patterns.", (3, 6)),
                ("15–25", "Listen twice", "Slide 7. Students download their stem track and listen twice before touching Logic. Circulate "
                 "and ask every student: where are your sections? They should be able to tell you (section map, page @MAP).", (7, 7)),
                ("25–48", "Markers + A pattern", "Slides 8–9: the worked example, then a short demo. Students add arrangement markers and "
                 "begin building their A pattern. Prompt: does your groove fit the feel of the section it lives in?", (8, 9)),
                ("48–53", "Save", f"Students save, named {NAME1}. Quick check of the day’s goal as they pack up.", None)]},
    {"n": 2, "kind": "Work day", "title": "Working: Foundation",
     "goal": "Foundation complete. Both patterns placed correctly. Mix adjusted.",
     "blocks": [("0–5", "Quick check", "Slide 10. Ask two students to play their A pattern alongside the stem. Does it fit? Whole class "
                 "listens and responds briefly.", (10, 10)),
                ("5–45", "A, B, place", "Slide 11. Students complete A pattern, build B pattern, and place both across the full stem. "
                 "Teacher circulates. Key prompt for B section: play A, then play B. Can you feel the difference without looking at the "
                 "screen?", (11, 11)),
                ("45–53", "Gain staging", "Slide 12. Students adjust volumes. Nothing should be clipping. Introduce the gain staging rule "
                 "explicitly: the stem is the lead, the drums support it (student page @GAIN).", (12, 12))]},
    {"n": 3, "kind": "Work day", "title": "Transition workshop + Full Song start",
     "goal": "At least two transitions built. Students understand the difference between the four types.",
     "blocks": [("0–15", "Transition workshop", "Slide 13. Show all four types with a real example: drum fill, breakdown, cut, build. Play "
                 "each one twice and ask students to describe what it does to the energy. This is the most important instructional moment "
                 "of the unit.", (13, 13)),
                ("15–48", "Map + build", "Slides 14–15: a mapped full song and a 64-step demo. Students map every section change in their "
                 f"stem and begin adding transitions. Teacher circulates. Key question at every transition: why did you choose this type "
                 f"for this moment?", (14, 15)),
                ("48–53", "Save", f"Students save the new project, named {NAME2}. Exit question: which transition type have you not tried "
                 "yet?", None)]},
    {"n": 4, "kind": "Work day", "title": "Working: Full Song",
     "goal": "All section changes have deliberate transitions. Mix adjusted. Project ready to submit.",
     "blocks": [("0–53", "Full working day", "Slides 16–17. Students complete all transitions and finalize mix. Common prompt: play the full "
                 "track from beginning to end without stopping. Every section change should feel like it was meant to happen. If anything "
                 "feels accidental, it needs work. Early finishers may try the optional reverb bus.", (16, 17))]},
    {"n": 5, "kind": "Submit", "title": "Submission + presentation prep",
     "goal": "All projects submitted before Day 6.",
     "blocks": [("0–20", "Finish + submit", "Slides 18–20: both checklists and the rubric. Students who are not yet done finish and submit. "
                 "Students who are done write their three Padlet feedback questions.", (18, 20)),
                ("20–53", "Early presentations", "Early presentations if time allows. Otherwise this is a full working and finalization day.",
                 None)]},
    {"n": 6, "kind": "Present", "title": "Presenting",
     "goal": "Every student plays their full drum track alongside the stem; every student hears the professional skill they just used.",
     "blocks": [("0–5", "Last submissions", "Any remaining submissions. Padlet questions posted.", None),
                ("5–53", "Presentations", "Slides 21–22. Presentations. Each student plays their full drum track alongside the stem. "
                 "Classmates respond to at least 2 projects. Close with the observation from page @TN1.", (21, 22))]},
]
DECK_DAY_LINES = ["Listen to the stem; the four transitions; map the sections; start the A pattern",
                  "Finish A, build a contrasting B, place both; gain staging",
                  "Transition workshop; map every section change; build transitions in 64 steps",
                  "Finish every transition; play it through; optional reverb bus",
                  "Checklists and rubric; submit; write three feedback questions",
                  "Play it with the stem; why that transition? You mixed a song"]

# ================================================================ intro deck (the owner's "Looping in Logic" pattern); (day, kind, title, body, notes)
SLIDES = [
    (1, "title", "Stems & Grooves", f"Feat. {ITAT[0]}",
     f"Day 1, the first minute. Have your Foundation stem open in Logic on the class speakers, and {ITAT[0]} ({ITAT[1]}) cued up on YouTube "
     "or Spotify from your own account; no audio or video is included in this deck. Stems are not included either: see the stems pages of "
     "the kit to find or make them. Open with: today you write drums for music that already exists."),
    (1, "listen", "Listen: music with no drums", "Your Foundation stem · teacher plays 60 seconds|No drums yet: that part is your job.|"
     "What is this music doing?|Where does the energy change?|Where does a new section feel like it starts?",
     "From my teacher notes, minutes 0 to 5: play 60 seconds of one of the stem tracks without the drums. Ask what the music is doing and "
     "where the energy changes. Do not give answers yet. The point is to hear before naming anything; students answer again after they "
     "have mapped their own stem."),
    (1, "bullets", "What the heck is a stem?", "One instrument or group, pulled out of a recording|Your stem: everything except the drums|"
     "Your job: write drums that belong|Listen first, then respond",
     "Plain-English definition, from the project sheet vocabulary: a stem track is a single isolated instrument or group from a real "
     "recording; in this project, everything except the drums. The unit's principle from my course plan is listen first, then respond. "
     "Drums that belong to the music, not just on top of it."),
    (1, "days", "Six days, two parts", "",
     "The shape of the unit. Part 1, Foundation: two contrasting patterns, A and B, placed where the sections change in a 16-measure stem. "
     "Part 2, Full Song: a new, longer stem, and a deliberate transition at every section change. Day 5 is for finishing and submitting, "
     "Day 6 for presenting. Hand out the project sheet now."),
    (1, "listen", "Listen: a famous drum entrance", f"{ITAT[0]} · {ITAT[1]}|Teacher: play from about 3:00 to 4:00 from YouTube or Spotify.|"
     "What keeps time before the big drums arrive?|What happens right before they come in?|Why does it hit so hard?",
     f"About three minutes. For most of {ITAT[0]} a quiet drum machine keeps time and the big drums are held back; they finally crash in "
     "with a famous tom fill near 3:40. That fill is a transition: it announces the new section. It hits hard because of the space before "
     "it. Keep it brief, as my notes say: students understand transitions better once they have built patterns."),
    (1, "types", "Four ways to announce a change", "",
     "The four transition types, word for word from the Day 1 post. Play each of the four examples you programmed from the worked "
     "transitions page, once each today; Day 3 is the full workshop. Students use these in Part 2, Full Song, so today is just a first "
     "hearing."),
    (1, "steps", "Listen before you build", "Download one stem from the Foundation folder|Play all 16 measures twice: touch nothing|"
     "Where does the energy change? Count measures|Arrangement markers: A or B, 4 measures each|Only then open Drum Machine Designer",
     "Minutes 15 to 25. Students download their stem and listen twice before touching Logic. Circulate and ask every student: where are "
     "your sections? They should be able to tell you. Students who struggle can use the section map worksheet, or you can pre-label the "
     "markers for them."),
    (1, "pic", "Worked example: A and B over the stem", "patterns",
     "One possible answer over the Foundation stem from the recipe: 16 measures, four sections of 4. Pattern A is laid back with closed "
     "hi-hat eighths; pattern B pushes energy with a kick on every beat and open hi-hats off the beat. Placed A, A, B, A to follow the stem. "
     "You can tell them apart without looking at the screen."),
    (1, "steps", "Watch me: build your A pattern", "New track: Software Instrument, an Electronic Drum Kit|"
     "Control-click at the first A: Create Pattern Region|Kick, snare, closed hi-hat first|Play it with the stem: does it fit?|"
     "Save it: Your Name - Foundation",
     "A three-minute demo. An Electronic Drum Kit loads Drum Machine Designer. Control-click in the track at the first A marker, Create "
     "Pattern Region, double-click to open the Step Sequencer. Start with kick, snare and closed hi-hat, aim for 3 to 5 layers, and keep "
     "playing it with the stem."),
    (2, "bullets", "Does it fit?", "Two volunteers: A pattern with the stem|Does the groove fit the feel of the section?|"
     "Do the drums follow the music?|Or does the music follow the drums?",
     "Day 2, minutes 0 to 5, from my notes: ask two students to play their A pattern alongside the stem. Does it fit? The whole class "
     "listens and responds briefly. The project sheet's question: do my drums follow the music, or does the music follow my drums? It "
     "should always be the former."),
    (2, "bullets", "Play A, then play B", "Can you feel the difference without looking?|If A is busy, B might be stripped back|"
     "If A is laid back, B might add energy|Then copy both across every section",
     "Minutes 5 to 45: students finish A, build B and place both across the full stem. The key prompt from my notes: play A, then "
     "immediately play B. If they have to look at the screen to tell them apart, they are not different enough. Push back on the student "
     "who changed one hi-hat hit and called it B."),
    (2, "gain", "Mixing moment: gain staging", "",
     "The mixing concept for this unit, minutes 45 to 53. Students adjust volumes so nothing clips. Say the rule from my tutorial: never "
     "push anything into the red; it is almost always better to turn other things down. Play a project with the drums too loud, pull the "
     "drums down until the stem leads, and ask which sounds more like a real song."),
    (3, "pic", "Transition workshop", "trans",
     "Day 3, the most important instructional moment of the unit. Play each of the four examples twice, programmed in Logic from the "
     "worked transitions page, and ask what it does to the energy. Fill: the groove breaks into snare and toms. Breakdown: only the kick "
     "is left. Cut: two beats of silence. Build: a sound joins on every beat."),
    (3, "pic", "Worked example: map every change", "songmap",
     "One possible map of the Full Song stem from the recipe: 56 measures, eight sections, so seven section changes and seven transitions. "
     "Four types are used, each chosen for its moment: a cut before the quieter bridge, a build before the big last chorus, as the project "
     "sheet says. Students map their own stem before they write a single transition."),
    (3, "steps", "Watch me: a transition in 64 steps", "Copy the groove into the measures before the change|"
     "Step Sequencer: set the pattern length to 64 steps|Change only the last measure (or two)|Play into the new section: does it announce it?|"
     "Listen, then move to the next change",
     "A short demo. A 64-step pattern covers four measures, so the transition sits in the same region as the groove. Change only the last "
     "one or two measures. Listen after each transition before moving to the next, and ask the key question at every one: why this type "
     "for this moment?"),
    (4, "bullets", "Play it all the way through", "Every section change should feel intentional|If a transition adds nothing, rework it|"
     "A cut only works if the music earns it|The stem is the star: drums support it",
     "Day 4 is a full working day. The common prompt from my notes: play the full track from beginning to end without stopping. Every "
     "section change should feel like it was meant to happen. If anything feels accidental, it needs work. Remind students of the cut tip: "
     "silence only works if the music earns it."),
    (4, "steps", "Optional: a reverb bus", "Mixer (X): on the drums, click a Send slot, choose a bus|Logic adds an Aux channel strip|"
     "On the Aux: an Audio FX slot, then ChromaVerb|Turn the Send knob up slowly|Not graded: the full lesson is in Unit 6",
     "Optional, for early finishers, and not graded: my course plan mentions reverb for the full song, and the bus and send workflow is "
     "taught properly in Unit 6, Spooky Tunes. A little shared reverb can glue the drums to the stem. Less is more: the drums should not "
     "sound far away."),
    (5, "check", "Foundation: what you turn in", "A and B written and placed across all sections|At least 3 layers per pattern|"
     "Nothing clipping, everything audible|Named as on your project sheet|MP3 + .zip, submitted before presenting",
     "Day 5. The Foundation part of the submission checklist, condensed. Show where students submit (your class page). Export: File, "
     "Bounce, Project or Section, MP3. Compress: in Finder, Control-click the .logicx project and choose Compress. Check the project name "
     "you want on the project sheet."),
    (5, "check", "Full Song: what you turn in", "Every section change labeled, with a transition|At least two transition types|"
     "Each transition fits the music|Mix: the drums serve the stem|Named; MP3 + .zip submitted|Three Padlet feedback questions posted",
     "The Full Song part of the checklist, plus the three feedback questions every student posts before presenting. Students who are done "
     "write their questions now; students who are not yet done finish and submit. All projects are submitted before Day 6."),
    (5, "rubric", "How it’s scored", "",
     f"Five categories, each scored 2 to 5, for a total out of {RUBRIC_TOTAL}. Read the 5 column aloud: distinct A and B patterns placed "
     "precisely; every transition feels inevitable; 64-step mode used for transition detail; drums that sit perfectly with the stem with "
     "nothing clipping; on time, confident, three thoughtful Padlet questions."),
    (6, "bullets", "Presentation day", "Play your full drum track with the stem|Point to each transition: why that type?|"
     "Respond to at least 2 projects|You responded to music that already existed",
     "Day 6. Each student plays their full drum track alongside the stem. Classmates respond to at least 2 projects. Ask about one "
     "transition: why did you choose this type for this moment? Name what students just did: they adjusted volumes and panned "
     "instruments around a real song."),
    (6, "closer", "One last thing…", "You listened first.|You wrote drums that belong.|You announced every change.|You mixed a song.",
     "Close with the observation from my teacher notes: your drums were written in response to music that already existed. That is what "
     "session musicians do. That is a professional skill. And from my course plan, name it: you mixed a song. Next unit, Chords by Ear."),
]
DECK_SLIDES = len(SLIDES)
DECK_PICS = ["hero", "patterns", "trans", "songmap", "meter"]
SIDE = {"patterns": ["A: laid back, closed hi-hat eighths", "B: kick on every beat, open hi-hats", "Placed A, A, B, A: it follows the stem",
                     "Tell them apart without looking"],
        "trans": ["Fill: the groove breaks into snare, then toms", "Breakdown: one element left (the kick)",
                  "Cut: two beats of complete silence", "Build: a sound joins on every beat"],
        "songmap": ["8 sections, 7 changes, 7 transitions", "Four types, each for its moment", "A cut before the quieter bridge",
                    "A build before the big last chorus"]}

# ================================================================ html scaffolding (Hudson Beat, from 006-012)


def font_css(base=BUILD):
    f = os.path.relpath(FONTS, base)
    faces = [("Oswald", 300, "normal", "oswald-latin-300-normal"), ("Oswald", 500, "normal", "oswald-latin-500-normal"),
             ("Oswald", 700, "normal", "oswald-latin-700-normal"),
             ("Source Sans 3", 400, "normal", "source-sans-3-latin-400-normal"),
             ("Source Sans 3", 400, "italic", "source-sans-3-latin-400-italic"),
             ("Source Sans 3", 600, "normal", "source-sans-3-latin-600-normal"),
             ("Source Sans 3", 700, "normal", "source-sans-3-latin-700-normal")]
    return "".join(f"@font-face{{font-family:'{n}';font-weight:{w};font-style:{s};src:url('{f}/{file}.woff2') format('woff2');}}"
                   for n, w, s, file in faces)


CSS = """
@page { size: 8.5in 11in; margin: 0; }
* { box-sizing: border-box; }
html, body { margin: 0; padding: 0; background: #fff; color: #000; }
body { font-family: 'Source Sans 3', sans-serif; -webkit-print-color-adjust: exact; print-color-adjust: exact; }
.page { width: 8.5in; height: 11in; position: relative; overflow: hidden; page-break-after: always; break-after: page;
        padding: 0.42in 0.5in 0.38in; display: flex; flex-direction: column; }
.page:last-child { page-break-after: auto; break-after: auto; }
.l { font-family: 'Oswald'; font-weight: 300; }
.top { display: flex; justify-content: space-between; align-items: flex-end; border-bottom: 2.2pt solid #000; padding-bottom: 6pt; }
.top .t { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; letter-spacing: .01em; line-height: 1; font-size: 25pt; }
.top .t small { display: block; font-family: 'Oswald'; font-weight: 300; text-transform: none; font-size: 12pt; letter-spacing: 0; margin-top: 4pt; }
.tag { text-align: right; }
.tag .lv { display: inline-block; background: #000; color: #fff; font-family: 'Oswald'; font-weight: 700; font-size: 14pt; padding: 3pt 9pt 4pt; border-radius: 4pt; letter-spacing: .04em; }
.tag .lv.ak { background: ACCENT; }
.tag .ln { font-family: 'Oswald'; font-weight: 300; font-size: 11pt; margin-top: 3pt; }
.name { display: flex; gap: 18pt; font-size: 12pt; margin: 9pt 0 4pt; }
.name span { flex: 1; border-bottom: 1pt solid #000; height: 15pt; }
.name span.d { flex: 0 0 1.6in; }
.main { flex: 1; display: flex; flex-direction: column; min-height: 0; padding-top: 6pt; overflow: hidden; }
.foot { display: flex; justify-content: space-between; font-size: 8.5pt; color: #333; border-top: .8pt solid #000; padding-top: 3pt; }
.wm { position: absolute; left: 50%; top: 50%; transform: translate(-50%, -50%) rotate(-38deg); font-family: 'Oswald'; font-weight: 700;
      font-size: 150pt; color: rgba(0,0,0,.13); letter-spacing: .06em; white-space: nowrap; pointer-events: none; z-index: 9; }
h2.s { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 13pt; margin: 9pt 0 3pt; letter-spacing: .02em; }
h2.s .acc, .acc { color: ACCENT; }
p, li { font-size: 10.5pt; line-height: 1.32; margin: 0 0 4pt; }
ul, ol { margin: 0 0 4pt; padding-left: 15pt; }
.tn p, .tn li { font-size: 11.4pt; line-height: 1.45; margin-bottom: 6pt; }
.tn.big p, .tn.big li { font-size: 12.2pt; line-height: 1.46; margin-bottom: 8pt; }
.tn.huge p { font-size: 14.2pt; line-height: 1.55; margin-bottom: 14pt; }
.lines { flex: 1; min-height: 44pt; background: repeating-linear-gradient(to bottom, transparent 0, transparent 21pt, #000 21pt, #000 22pt); }
table.t { border-collapse: collapse; width: 100%; font-size: 9.8pt; margin: 2pt 0 4pt; }
table.t td, table.t th { border: 1pt solid #000; padding: 3pt 5pt; text-align: left; vertical-align: top; line-height: 1.25; }
table.t th { font-family: 'Oswald'; font-weight: 500; text-transform: uppercase; font-size: 9.5pt; background: #000; color: #fff; letter-spacing: .04em; vertical-align: middle; }
.cols { display: flex; gap: 20pt; }
.cols > div { flex: 1; min-width: 0; }
.small { font-size: 9pt !important; color: #333; }
.box { border: 1.4pt solid #000; border-radius: 8pt; padding: 7pt 10pt; }
.box.acc-b { border-color: ACCENT; border-width: 2pt; }
.box.inv { background: #000; color: #fff; border-color: #000; }
.k { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 10pt; letter-spacing: .05em; color: ACCENT; display: block; margin-bottom: 2pt; }
.box.inv .k { color: ACC_LIGHT; }
.ln2 { border-bottom: 1pt solid #000; height: 22pt; }
.fl { display: flex; align-items: flex-end; gap: 6pt; font-size: 10.5pt; margin: 0 0 8pt; }
.fl b { font-weight: 600; white-space: nowrap; }
.fl span { flex: 1; border-bottom: 1pt solid #000; height: 14pt; }
.num { list-style: none; padding: 0; margin: 2pt 0 0; }
.num li { display: flex; gap: 10pt; margin: 0 0 9pt; }
.num .n { flex: 0 0 24pt; height: 24pt; border-radius: 12pt; background: #000; color: #fff; font-family: 'Oswald'; font-weight: 700; font-size: 13pt;
          display: flex; align-items: center; justify-content: center; }
.num .hd { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 12pt; letter-spacing: .02em; }
.why { font-style: italic; }
.blk { display: flex; gap: 9pt; margin: 0 0 5pt; font-size: 10.4pt; line-height: 1.3; }
.blk .tm { flex: 0 0 40pt; text-align: center; background: #000; color: #fff; font-family: 'Oswald'; font-weight: 700; font-size: 10pt;
           border-radius: 4pt; padding: 2pt 0 3pt; align-self: flex-start; }
.blk b.h { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 10.5pt; letter-spacing: .02em; }
.day { border: 1.4pt solid #000; border-radius: 8pt; padding: 7pt 10pt 3pt; margin-bottom: 8pt; }
.day .dh { display: flex; justify-content: space-between; align-items: baseline; border-bottom: 1pt solid #000; padding-bottom: 3pt; margin-bottom: 6pt; }
.day .dh b { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 14pt; }
.day .dh span { font-family: 'Oswald'; font-weight: 300; font-size: 11pt; }
.goal { font-size: 10.6pt; margin: 0 0 7pt; }
.chk { display: flex; gap: 8pt; align-items: flex-start; font-size: 11pt; margin: 0 0 8pt; line-height: 1.3; }
table.rub { border-collapse: collapse; width: 100%; }
table.rub td, table.rub th { border: 1.1pt solid #000; padding: 5pt 5pt; vertical-align: top; font-size: 9.9pt; line-height: 1.3; }
table.rub th { background: #000; color: #fff; font-family: 'Oswald'; font-weight: 500; font-size: 12pt; text-align: center; vertical-align: middle; }
table.rub td.c { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 10.5pt; width: 14%; vertical-align: middle; }
table.rub td.sc { width: 7%; }
.cover { align-items: center; justify-content: space-between; padding: 0.55in 0.6in 0.5in;
         background: linear-gradient(180deg, #fff 0%, #fff 55%, #ededed 100%); }
.store { font-family: 'Oswald'; font-weight: 500; font-size: 13pt; letter-spacing: .32em; text-transform: uppercase; }
.ctitle { text-align: center; }
.ctitle .a, .ctitle .b, .ctitle .c { white-space: nowrap; }
.ctitle .a { font-family: 'Oswald'; font-weight: 700; line-height: 1; letter-spacing: .02em; text-transform: uppercase; }
.ctitle .b { font-family: 'Oswald'; font-weight: 700; line-height: .98; letter-spacing: .01em; text-transform: uppercase; }
.ctitle .c { font-family: 'Oswald'; font-weight: 300; margin-top: 8pt; }
.cband { width: 100%; background: #000; color: #fff; text-align: center; font-family: 'Oswald'; font-weight: 500; font-size: 14pt;
         letter-spacing: .08em; text-transform: uppercase; padding: 8pt 0 9pt; border-radius: 4pt; white-space: nowrap; }
.cband em { font-style: normal; color: ACCENT; }
.rule { background: #000; color: #fff; border-radius: 8pt; padding: 9pt 12pt; }
.rule .k { color: ACC_LIGHT; }
.rule b { font-size: 14pt; line-height: 1.3; display: block; }
.stp { font-family: 'Oswald'; font-weight: 500; letter-spacing: .12em; }
.dp .blk { font-size: 11.4pt; line-height: 1.4; margin-bottom: 8pt; }
.dp .goal { font-size: 11.6pt; }
.dp .day { padding-bottom: 5pt; margin-bottom: 12pt; }
.roomy p, .roomy li { font-size: 11.6pt; line-height: 1.42; margin-bottom: 6pt; }
.roomy .chk { font-size: 11.6pt; margin-bottom: 10pt; }
table.map { border-collapse: collapse; width: 100%; }
table.map td, table.map th { border: 1.2pt solid #000; padding: 5pt 6pt; font-size: 10.5pt; height: 25pt; }
table.map th { background: #000; color: #fff; font-family: 'Oswald'; font-weight: 500; text-transform: uppercase; font-size: 10pt; letter-spacing: .04em; height: auto; }
""".replace("ACCENT", ACCENT).replace("ACC_LIGHT", ACC_LIGHT)

COPY = f"© {YEAR} {STORE} · {BRAND} · For single-classroom use"


def html_doc(pages, title="doc"):
    return (f'<!doctype html><html><head><meta charset="utf-8"><title>{escape(title)}</title>'
            f'<style>{font_css()}{CSS}</style></head><body>{"".join(pages)}</body></html>')


def foot(right):
    return f'<div class="foot"><span>{COPY}</span><span>{right}</span></div>'


def wm(on):
    return '<div class="wm">PREVIEW</div>' if on else ""


def page(sub, tag, tagsub, inner, n, watermark=False, name=False, ak=False, cls="", key=""):
    top = (f'<div class="top"><div class="t">{TITLE_H}<small>{sub}</small></div>'
           f'<div class="tag"><div class="lv{" ak" if ak else ""}">{tag}</div><div class="ln">{tagsub}</div></div></div>')
    nm = '<div class="name">Name<span></span>Date<span class="d"></span></div>' if name else ""
    return (f'<section class="page {cls}" data-key="{key}">{wm(watermark)}{top}{nm}<div class="main">{inner}</div>{foot(f"Page {n}")}</section>')


def checkbox(size=12):
    return (f'<svg width="{size}pt" height="{size}pt" viewBox="0 0 {size} {size}" style="flex:0 0 auto;margin-top:1pt">'
            f'<rect x=".7" y=".7" width="{size - 1.4}" height="{size - 1.4}" rx="2" fill="#fff" stroke="#000" stroke-width="1.2"/></svg>')


def t_(x, y, s, size, weight=700, family="Oswald", anchor="middle", fill="#000", extra=""):
    fam = f"'{family}'" if " " in family else family
    return (f'<text x="{x:.1f}" y="{y:.1f}" font-family="{fam}" font-weight="{weight}" font-size="{size}" text-anchor="{anchor}" '
            f'fill="{fill}" {extra}>{escape(s)}</text>')


def svgwrap(vw, vh, body, width="100%"):
    return f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {vw:.0f} {vh:.0f}" width="{width}" style="display:block">{body}</svg>'


_ADV = {}


def text_w(s, size, font="oswald"):
    """Advance width of s at size (viewBox units / px / pt) from the font's own metrics: Oswald 700, or Source Sans 3 400 ("ss")."""
    if font not in _ADV:
        from fontTools.ttLib import TTFont
        f = TTFont(os.path.join(FONTS, "oswald-latin-700-normal.woff2" if font == "oswald" else "source-sans-3-latin-400-normal.woff2"))
        cmap, hm = f.getBestCmap(), f["hmtx"].metrics
        d = {chr(c): hm[g][0] for c, g in cmap.items() if c < 0x2200}
        d["upm"] = f["head"].unitsPerEm
        _ADV[font] = d
    a = _ADV[font]
    return sum(a.get(ch, a["upm"] * 0.6) for ch in s) * size / a["upm"]


# ================================================================ drawings (original, code-drawn: black + accent only); every label is logged
LABELS = []     # (diagram, text, size, room): text_w(text, size) must be <= room
SPANS = []      # (diagram, row, x0, x1): labels on one row must not overlap (checked)


def lab(diagram, text, size, room):
    LABELS.append((diagram, text, size, room))


def span(diagram, row, text, size, cx, anchor="middle", lo=0.0, hi=1e9):
    w = text_w(text, size)
    x0 = cx - w / 2 if anchor == "middle" else (cx if anchor == "start" else cx - w)
    SPANS.append((diagram, row, x0, x0 + w, lo, hi))


def grid_body(name, measures, captions, cw=20, ch=22, gut=150, ls=15, accent_from=1, lanes=None):
    """A drum grid, one 16-step block per measure. Grey columns are beats. Hits in measures >= accent_from are drawn in the accent.
    Returns (svg body, width, height)."""
    lanes = lanes or lanes_of(*measures)
    out = []
    head = 22
    W = gut + STEPS * len(measures) * cw + 4
    for mi in range(len(measures)):
        for b in range(4):
            x = gut + (mi * STEPS + b * 4) * cw
            out.append(t_(x + cw / 2, 16, str(b + 1), 14, 700))
            lab(f"{name} beats", str(b + 1), 14, cw * 2)
    for r, ln in enumerate(lanes):
        y = head + r * ch
        out.append(t_(gut - 8, y + ch * 0.7, ln, ls, 500, anchor="end"))
        lab(f"{name} lanes", ln, ls, gut - 14)
        for mi, m in enumerate(measures):
            s = m.get(ln, "." * STEPS)
            for i in range(STEPS):
                x = gut + (mi * STEPS + i) * cw
                bg = "#e4e4e4" if i % 4 == 0 else "#fff"
                out.append(f'<rect x="{x:.1f}" y="{y:.1f}" width="{cw:.1f}" height="{ch}" fill="{bg}" stroke="#000" '
                           f'stroke-width="{1.3 if i % 4 == 0 else .55}"/>')
                if s[i] == "x":
                    out.append(f'<rect x="{x + 3:.1f}" y="{y + 3:.1f}" width="{cw - 6:.1f}" height="{ch - 6}" rx="2.5" '
                               f'fill="{ACCENT if mi >= accent_from else "#000"}"/>')
    gh = len(lanes) * ch
    for mi in range(len(measures) + 1):
        x = gut + mi * STEPS * cw
        out.append(f'<line x1="{x:.1f}" y1="{head - 4}" x2="{x:.1f}" y2="{head + gh + 4}" stroke="#000" stroke-width="3"/>')
    out.append(f'<rect x="{gut}" y="{head}" width="{STEPS * len(measures) * cw:.1f}" height="{gh}" fill="none" stroke="#000" stroke-width="1.8"/>')
    for mi, cap in enumerate(captions):
        cx = gut + (mi + .5) * STEPS * cw
        out.append(t_(cx, head + gh + 21, cap, 16, 700, fill=ACCENT if mi >= accent_from else "#000"))
        lab(f"{name} captions", cap, 16, STEPS * cw - 12)
    return "".join(out), W, head + gh + 28


def grid_svg(name, measures, captions, width="100%", **kw):
    b, W, H = grid_body(name, measures, captions, **kw)
    return svgwrap(W, H, b, width)


def wave(x, y, w, h, level, seed, fill="#fff"):
    """A stylised stem waveform (deterministic), symmetric around the middle of the box."""
    out, n = [], max(4, int(w // 7))
    for i in range(n):
        a = level * (0.45 + 0.55 * abs(math.sin(seed * 1.7 + i * 0.83) * math.cos(seed + i * 0.37)))
        bh = max(2.0, a * h)
        out.append(f'<rect x="{x + i * w / n + 1:.1f}" y="{y + (h - bh) / 2:.1f}" width="{max(1.5, w / n - 2.5):.1f}" height="{bh:.1f}" '
                   f'rx="1" fill="{fill}" opacity=".55"/>')
    return "".join(out)


def strip_body(name, sections, trans=None, W=1600, h=118, show_letters=True):
    """A stem's section map: one block per section (width = measures), a waveform for its energy, the start measure under it,
    and (Full Song) the transition chosen for each section change. Returns (body, width, height)."""
    total = sum(n for _, n in sections)
    mw = (W - 4) / total
    out, x = [], 2.0
    emax = max(energy(s) for s, _ in sections)
    fills = {"B": ACCENT, "Chorus": ACCENT, "Bridge": "#555", "Intro": "#6e6e6e", "Outro": "#6e6e6e"}
    st = starts(sections)
    for k, (sec, n) in enumerate(sections):
        w = n * mw
        out.append(f'<rect x="{x + 2:.1f}" y="2" width="{w - 4:.1f}" height="{h}" rx="8" fill="{fills.get(sec, "#000")}"/>')
        out.append(wave(x + 10, 50, w - 20, h - 56, energy(sec) / emax, k + 1))
        head = sec if not show_letters else (sec if sec in ("A", "B") else sec)
        sz = 26 if len(head) > 1 else 32
        while sz > 18 and text_w(head.upper(), sz) > w - 16:
            sz -= 1
        out.append(t_(x + w / 2, 36, head.upper(), sz, 700, fill="#fff"))
        lab(name, head.upper(), sz, w - 16)
        out.append(t_(x + 4, h + 24, f"m. {st[k]}", 15, 500, anchor="start"))
        lab(f"{name} measures", f"m. {st[k]}", 15, w - 8)
        x += w
    out.append(t_(W - 2, h + 24, f"{total}", 15, 500, anchor="end"))
    H = h + 32
    if trans:
        x = 2.0
        y = h + 64
        for k, (sec, n) in enumerate(sections[:-1]):
            x += n * mw
            tname = trans[k][0]
            out.append(f'<line x1="{x:.1f}" y1="{h + 30}" x2="{x:.1f}" y2="{y - 20}" stroke="{ACCENT}" stroke-width="3"/>')
            out.append(f'<circle cx="{x:.1f}" cy="{h + 30}" r="5" fill="{ACCENT}"/>')
            label = tname.replace("Drum fill", "Fill").upper()
            out.append(t_(x, y, label, 19, 700, fill=ACCENT))
            span(name, "transitions", label, 19, x, hi=W)
            lab(name, label, 19, W)
        H = y + 10
    return "".join(out), W, H


def strip_svg(name, sections, trans=None, width="100%", **kw):
    b, W, H = strip_body(name, sections, trans, **kw)
    return svgwrap(W, H, b, width)


def stack(parts, gap=26):
    """Stack (body, w, h) drawings vertically, centred."""
    W = max(w for _, w, _ in parts)
    out, y = [], 0
    for b, w, h in parts:
        out.append(f'<g transform="translate({(W - w) / 2:.1f},{y:.1f})">{b}</g>')
        y += h + gap
    return "".join(out), W, y - gap


def side(parts, gap=40):
    H = max(h for _, _, h in parts)
    out, x = [], 0
    for b, w, h in parts:
        out.append(f'<g transform="translate({x:.1f},{(H - h) / 2:.1f})">{b}</g>')
        x += w + gap
    return "".join(out), x - gap, H


def patterns_body():
    strip = strip_body("patterns strip", FOUNDATION, W=1500, h=96)
    grid = grid_body("patterns grid", [PATTERN_A, PATTERN_B], ["A · measures 1–8 and 13–16", "B · measures 9–12"], cw=40, ch=32, gut=170, ls=20,
                     lanes=lanes_of(PATTERN_A, PATTERN_B))
    return stack([strip, grid], 30)


def trans_cell(t, name, cw=22, ch=27, caption=None):
    return grid_body(name, [PATTERN_A, TRANS[t]], ["Groove (A)", caption or f"{t}"], cw=cw, ch=ch, gut=150, ls=18,
                     lanes=lanes_of(PATTERN_A, TRANS[t]))


def trans_body():
    cells = [trans_cell(t, f"trans {t}") for t in TYPES]
    return stack([side(cells[:2], 60), side(cells[2:], 60)], 34)


def songmap_body(W=1700, h=104):
    return strip_body(f"songmap {W}", FULL_SONG, SONG_TRANS, W=W, h=h)


def hero_body():
    strip = strip_body("hero strip", FULL_SONG, SONG_TRANS, W=1100, h=175)
    grid = grid_body("hero grid", [PATTERN_A, TRANS["Drum fill"]], ["Verse groove", "Fill into the chorus"], cw=28, ch=40, gut=150, ls=19,
                     lanes=lanes_of(PATTERN_A, TRANS["Drum fill"]))
    return stack([strip, grid], 34)


# meters: (panel title, levels as a fraction of full scale where 1.0 = 0 dB, the top of the clean zone)
METERS = [("Backwards: drums too loud", {"Stem": 0.55, "Drums": 0.93, "Master": 1.1}),
          ("Right: the stem leads", {"Stem": 0.8, "Drums": 0.62, "Master": 0.88})]


def meter_body():
    out, mh, top, bw = [], 330, 40, 70
    W = 0
    for p, (title, lv) in enumerate(METERS):
        x0 = p * 470
        out.append(t_(x0 + 200, 24, title.upper(), 21, 700, fill=ACCENT if p else "#000"))
        lab("meter", title.upper(), 21, 420)
        y0 = top + 38
        full = 1.18
        zero = y0 + mh * (1 - 1.0 / full)
        for k, (nm, v) in enumerate(lv.items()):
            x = x0 + 40 + k * 120
            out.append(f'<rect x="{x}" y="{y0}" width="{bw}" height="{mh}" fill="#fff" stroke="#000" stroke-width="2"/>')
            out.append(f'<rect x="{x}" y="{y0}" width="{bw}" height="{zero - y0:.1f}" fill="{CLIP_RED}" opacity=".18"/>')
            hgt = mh * min(v, full) / full
            yb = y0 + mh - hgt
            col = {"Stem": "#000", "Drums": ACCENT, "Master": "#777"}[nm]
            clean_top = max(yb, zero)
            out.append(f'<rect x="{x + 6}" y="{clean_top:.1f}" width="{bw - 12}" height="{y0 + mh - clean_top:.1f}" fill="{col}"/>')
            if yb < zero:
                out.append(f'<rect x="{x + 6}" y="{yb:.1f}" width="{bw - 12}" height="{zero - yb:.1f}" fill="{CLIP_RED}"/>')
            out.append(t_(x + bw / 2, y0 + mh + 26, nm.upper(), 18, 700))
            lab("meter", nm.upper(), 18, 112)
        out.append(f'<line x1="{x0 + 30}" y1="{zero:.1f}" x2="{x0 + 400}" y2="{zero:.1f}" stroke="#000" stroke-width="2" stroke-dasharray="7 5"/>')
        out.append(t_(x0 + 200, y0 - 10, "ABOVE THE LINE (0 dB) = RED = CLIPPING", 14, 700, fill=CLIP_RED))
        lab("meter", "ABOVE THE LINE (0 dB) = RED = CLIPPING", 14, 400)
        W = x0 + 420
    return "".join(out), W, top + 38 + 330 + 36


def body_svg(bwh, width="100%"):
    b, W, H = bwh
    return svgwrap(W, H, b, width)


def verify_labels():
    """Every drawn label fits its cell, block or column; labels on one row do not overlap; nothing leaves the drawing."""
    LABELS.clear()
    SPANS.clear()
    _, Wp, _ = patterns_body()
    _, Wt, _ = trans_body()
    _, Ws, _ = songmap_body()
    songmap_body(W=1250, h=96)
    _, Wh, _ = hero_body()
    meter_body()
    for i, q in enumerate(QUIZ):
        grid_body(f"quiz {i}", [PATTERN_A, q], ["Groove (A)", "?"], cw=16, ch=20, gut=140, lanes=lanes_of(PATTERN_A, q))
    errs = [f"{d}: label '{t}' at {s} is {text_w(t, s):.0f} wide, room {r:.0f}" for d, t, s, r in LABELS if text_w(t, s) > r]
    if not LABELS:
        errs.append("label check found no labels")
    rows, widths = {}, {}
    for d, row, x0, x1, lo, hi in SPANS:
        rows.setdefault((d, row), []).append((x0, x1))
        widths[d] = hi
    for (d, row), iv in rows.items():
        iv.sort()
        for (a0, a1), (b0, b1) in zip(iv, iv[1:]):
            if b0 < a1 + 10:
                errs.append(f"{d} {row}: labels overlap ({a0:.0f}-{a1:.0f} and {b0:.0f}-{b1:.0f})")
        if iv[0][0] < 0 or iv[-1][1] > widths.get(d, 1e9):
            errs.append(f"{d} {row}: a label runs off the drawing")
    return errs


# ================================================================ page plan
PLAN = ["cover", "howto", "overview", "setup", "stems", "makestem", "logicref", "pacing", "day12", "day34", "day56", "tn1", "tn2", "tn3",
        "post1", "transitions", "gain", "map", "proj1", "proj2", "proj3", "proj4", "rubric", "transcheck", "key1", "key2", "key3", "terms"]
IX = {k: i + 1 for i, k in enumerate(PLAN)}
PREVIEW_PICKS = ["pacing", "transitions", "proj2", "rubric"]
STUDENT = ("post1", "rubric")


def rng(a, b):
    return str(a) if a == b else f"{a}–{b}"


def day_slides(d):
    sl = [s for blk in d["blocks"] if blk[3] for s in range(blk[3][0], blk[3][1] + 1)]
    return (min(sl), max(sl)) if sl else None


def slides_label(d):
    r = day_slides(d)
    return f"Intro deck slides {rng(*r)}" if r and r[0] != r[1] else (f"Intro deck slide {r[0]}" if r else "No slides")


def fill(s):
    """Page placeholders are written @NAME so they never collide with ordinary words."""
    pages = {"MAP": IX["map"], "GAIN": IX["gain"], "KEY1": IX["key1"], "KEY2": IX["key2"], "KEY3": IX["key3"], "TN1": IX["tn1"],
             "STEMS": IX["stems"], "MAKE": IX["makestem"], "TRANS": IX["transitions"], "RUB": IX["rubric"]}
    out = re.sub(r"@([A-Z0-9]+)", lambda m: str(pages[m.group(1)]), s)
    if "@" in out:
        raise ValueError(f"unfilled placeholder in: {s}")
    return out


def H2(s, first=False):
    st = ' style="margin-top:2pt"' if first else ""
    return f'<h2 class="s"{st}>{s}<span class="acc"> ·</span></h2>'


COVER_A = "Drums over stems in Logic Pro"
COVER_B = "Stems & Grooves"
COVER_C = f"DMP1 · Unit {UNIT} · listen first · fills, breakdowns, cuts, builds · gain staging"
COVER_BAND = f"Grades 9–12 · {DAYS_N} days · {len(PLAN)} pages + {DECK_SLIDES}-slide deck"
PDF_COVER = dict(a=30, b=68, c=14.5, band=14, room=7.3 * 72)
SQ_COVER = dict(a=50, b=134, c=28, band=26, room=940)
SQ_B_LINES = ["Stems &", "Grooves"]   # the square stacks the big title in two equal lines (one line is too wide at 10% height)


def cover_page():
    s = PDF_COVER
    band = COVER_BAND.replace(" · ", " <em>·</em> ")
    return (f'<section class="page cover" data-key="cover"><div class="store">{BRAND}</div>'
            f'<div style="width:7.3in">{body_svg(hero_body())}</div>'
            f'<div class="ctitle"><div class="a" style="font-size:{s["a"]}pt">{COVER_A}</div><div class="b" style="font-size:{s["b"]}pt">{escape(COVER_B)}</div>'
            f'<div class="c" style="font-size:{s["c"]}pt">{COVER_C}</div></div>'
            f'<div class="cband">{band}</div></section>')


def deck_day_line():
    parts = []
    for d in DAYS:
        r = day_slides(d)
        parts.append(f"Day {d['n']} (slide{'s' if r[0] != r[1] else ''} {rng(*r)})")
    return ", ".join(parts[:-1]) + " and " + parts[-1]


def howto_page(n):
    parts = [("Teacher pages", "howto", "tn3", "How to use, overview, setup, the stems (find or make them), Logic reference, pacing, day plans, "
              "teacher notes"),
             ("Student pages", "post1", "transcheck", "Day 1 post, the four transitions, gain staging, section map, the project sheet (4 pages), "
              "rubric, name that transition"),
             ("Answer keys + examples", "key1", "key3", "Worked A and B patterns, the four transitions and quiz answers, a mapped full song"),
             ("Intro deck (separate file)", None, None, f"{DECK_SLIDES} slides, .pptx + PDF, speaker notes on every slide")]
    trs = "".join(f"<tr><td><b>{a}</b></td><td style='white-space:nowrap'>{('pages ' + str(IX[f]) + '–' + str(IX[l])) if f else 'deck'}</td><td>{u}</td></tr>"
                  for a, f, l, u in parts)
    inner = (f'<div class="cols"><div>'
             f'{H2("What this is", True)}'
             f'<p>A complete {DAYS_N}-day drum unit for high school music technology. Students write drums for music that already exists: a '
             f'stem, a recording with the drums taken out. In Part 1, Foundation, they write two contrasting patterns, A and B, and place them '
             f'where the sections change in a 16-measure stem. In Part 2, Full Song, they map every section change in a longer stem and announce '
             f'each one with a fill, a breakdown, a cut or a build. The mixing concept is gain staging: nothing in the red.</p>'
             f'<p>This is Unit {UNIT} of my Digital Music Production 1 course, after Unit 1, Beat From Nothing (a separate resource), where '
             f'students learned Drum Machine Designer and the Step Sequencer.</p>'
             f'{H2("Quick start")}<ol>'
             f'<li>Read the overview and pacing (pages {IX["overview"]} and {IX["pacing"]}) and the teacher notes (pages {IX["tn1"]}–{IX["tn3"]}).</li>'
             f'<li><b>Get your stems</b> (pages {IX["stems"]}–{IX["makestem"]}): a few 16-measure Foundation stems and a few Full Song stems. '
             f'None are included.</li>'
             f'<li>Program the four example transitions (page {IX["key2"]}) and cue up {ITAT[0]} ({ITAT[1]}) on YouTube or Spotify.</li>'
             f'<li>Print the student pages (pages {IX["post1"]}–{IX["transcheck"]}), one set per student.</li></ol>'
             f'{H2("The intro deck")}'
             f'<p>Every slide has a day: {deck_day_line()}. The day plans name the slide for each block. Every slide has speaker notes.</p>'
             f'<p class="small">Fonts: the deck uses Oswald and Source Sans 3 (free on Google Fonts, built into Google Slides). If PowerPoint or Keynote '
             f'shows a different font, install the two fonts or present the PDF copy, which shows the intended look.</p>'
             f'</div><div>'
             f'{H2("What’s inside", True)}<table class="t"><tr><th>Part</th><th>Pages</th><th>Use</th></tr>{trs}</table>'
             f'<div class="box acc-b" style="margin-top:6pt"><span class="k">No stems included</span><p style="margin:0">{fill(STEM_NOTE)}</p></div>'
             f'<div class="box" style="margin-top:6pt"><span class="k">About the songs</span><p style="margin:0">{SONG_NOTE}</p></div>'
             f'{H2("Make it fit your class")}'
             f'<ul><li><b>File names:</b> my requirements say {NAME1} and {NAME2}; my checklist says {NAME_CHECK} for both. Tell students '
             f'which names you want.</li>'
             f'<li><b>Grading:</b> my rubric, five categories scored 2–5, /{RUBRIC_TOTAL} (page {IX["rubric"]}).</li>'
             f'<li><b>More accessible / more advanced:</b> variations on page {IX["tn3"]}.</li>'
             f'<li><b>Feedback board:</b> I use Padlet for the three feedback questions; any class board works.</li></ul>'
             f'</div></div>')
    return page("How to use this resource", "START HERE", "Teacher", inner, n, key="howto")


def overview_page(n):
    inner = (f'<div class="cols"><div>'
             f'{H2("What students make", True)}'
             f'<p>Two drum tracks written to fit stems the teacher supplies. <b>Foundation:</b> an A pattern and a contrasting B pattern, '
             f'at least 3 layers each, placed where the sections change in a 16-measure stem. <b>Full Song:</b> drums for a full-length stem '
             f'with a deliberate transition at every section change, at least two types, mixed so the drums serve the stem.</p>'
             f'{H2("What students learn")}<p>{PLAN_LEARN}</p>'
             f'{H2("Why this unit lands here")}<p>{PLAN_WHY}</p>'
             f'{H2("Mixing concept: gain staging")}'
             f'<p>The stem tracks have varying levels, and students notice at once when their drums are too loud or too quiet. The rule: '
             f'“{RED_RULE}”</p>'
             f'</div><div>'
             f'{H2("Goals", True)}<ul>' + "".join(f"<li>{g}</li>" for g in GOALS) + '</ul>'
             f'{H2("At a glance")}<table class="t">'
             f'<tr><td><b>Grades</b></td><td>9–12 (no music reading needed)</td></tr>'
             f'<tr><td><b>Length</b></td><td>{DAYS_N} periods of about {MINUTES} minutes: learning day, three work days, submit, present</td></tr>'
             f'<tr><td><b>Software</b></td><td>Logic Pro: Drum Machine Designer, pattern regions, Step Sequencer (16 and 64 steps), '
             f'arrangement markers, volume and pan, export</td></tr>'
             f'<tr><td><b>You supply</b></td><td>The stems (none included): find or make them, pages {IX["stems"]}–{IX["makestem"]}</td></tr>'
             f'<tr><td><b>Product</b></td><td>Two projects, Foundation and Full Song, each an MP3 + compressed project; three feedback questions</td></tr>'
             f'<tr><td><b>Assessment</b></td><td>My 5-category rubric, scored 2–5, /{RUBRIC_TOTAL} (page {IX["rubric"]})</td></tr></table>'
             f'{H2("Standards")}'
             f'<p>Supports the National Core Arts Standards for Music (Music Technology, high school): developing musical ideas (MU:Cr2.1), '
             f'refining work (MU:Cr3.1) and presenting work (MU:Pr6.1).</p>'
             f'<div class="box" style="margin-top:4pt"><span class="k">From my tutorial</span><p style="margin:0"><i>“{DRUM_LINE}”</i></p></div>'
             f'</div></div>')
    return page("Unit overview", "OVERVIEW", f"DMP1 · Unit {UNIT}", inner, n, key="overview")


def setup_page(n):
    inner = (f'<div class="cols roomy"><div>'
             f'{H2("You need", True)}<ul>'
             f'<li><b>A Mac with Logic Pro</b> for each student (Logic Pro X 10.5 or later, for the Step Sequencer and pattern regions), and '
             f'<b>headphones</b>.</li>'
             f'<li><b>Stems</b>: a Foundation folder (several 16-measure stems, sections of 4 measures) and a Full Song folder (several '
             f'full-length stems with clear sections), on your class page. None are included: pages {IX["stems"]}–{IX["makestem"]}.</li>'
             f'<li><b>Projector and speakers</b> for the deck, the stems and {ITAT[0]} ({ITAT[1]}), from your own YouTube or Spotify access.</li>'
             f'<li><b>The four example transitions</b>, programmed once in a Logic project from page {IX["key2"]} (10 minutes), for Days 1 and 3.</li>'
             f'<li><b>Printed student pages</b>, one set each: the Day 1 post, the four transitions, gain staging, the section map, the project '
             f'sheet and rubric, name that transition.</li>'
             f'<li><b>A place to submit</b> an MP3 and a .zip, and a class feedback board for the three questions (I use Padlet).</li></ul>'
             f'{H2("A MIDI keyboard?")}<p>Not needed. Every step in this unit is clicked into the Step Sequencer.</p>'
             f'<div class="box" style="margin-top:6pt"><span class="k">Skills from Unit 1 this unit uses</span><p style="margin:0">Drum Machine '
             f'Designer, pattern regions, the Step Sequencer (16 and 64 steps), volume and pan in the Mixer, and export. Students new to them '
             f'need about 10 extra minutes on Day 1 with the Logic reference (page {IX["logicref"]}).</p></div>'
             f'</div><div>'
             f'{H2("Before Day 1", True)}'
             + "".join(f'<div class="chk">{checkbox()}<span>{c}</span></div>' for c in [
                 "Put the Foundation stems on your class page, each named with its tempo (for example “Stem 1 - 96 BPM”).",
                 "Open one stem in Logic yourself: set the tempo, drop it at bar 1, check the click lines up.",
                 "Program the four example transitions (page " + str(IX["key2"]) + ").",
                 "Test the class speakers; cue up the song.",
                 "Print the student pages; set up the submission spot and the feedback board."]) +
             f'{H2("Before Day 3")}'
             + "".join(f'<div class="chk">{checkbox()}<span>{c}</span></div>' for c in [
                 "Put the Full Song stems on your class page.",
                 "Optional scaffolding: pre-label the sections with arrangement markers for students who need it."]) +
             f'<div class="box acc-b" style="margin-top:8pt"><span class="k">How students open a stem</span>'
             f'<p style="margin:0">New empty project. Set the project tempo to the stem’s BPM. Drag the audio file into the Tracks area so it '
             f'starts at bar 1. Turn on the metronome (K): the click should line up with the music. Then add the drum track.</p></div>'
             f'</div></div>')
    return page("Setup and requirements", "SETUP", "Teacher", inner, n, key="setup")


def stems_page(n):
    good = ["<b>No drums.</b> Everything else can stay: bass, chords, melody, vocals.",
            "<b>A steady tempo you know.</b> Put the BPM in the file name so students can set Logic to match.",
            "<b>Clear sections.</b> Energy that changes where the sections change: instruments in or out, louder or softer.",
            "<b>Starts on the downbeat of bar 1</b>, with no pickup or silence before it, so the grid lines up.",
            "<b>Foundation:</b> 16 measures, sections of 4 measures (A or B), as in my Day 1 post.",
            "<b>Full Song:</b> 2 to 4 minutes with an intro, verses, choruses and ideally a bridge.",
            "<b>WAV or MP3</b>, one file per stem."]
    inner = (f'<div class="box inv"><span class="k">What is (and is not) in this kit</span><p style="margin:0;font-size:11.5pt">'
             f'{fill(STEM_NOTE)} My own classroom stems come from other artists, so I cannot include them.</p></div>'
             f'<div class="cols roomy" style="margin-top:6pt"><div>'
             f'{H2("Where to get stems")}'
             f'<p><b>1. Make your own (safest).</b> Build a short backing track in Logic with no drums and bounce it to an audio file. The recipe '
             f'on page {IX["makestem"]} makes one Foundation stem and one Full Song stem in about 30 minutes.</p>'
             f'<p><b>2. A multitrack library (check its terms first).</b> I source new tracks from the multitrack library at '
             f'cambridge-mt.com/ms/mtk. Before you use it, read the library’s own terms and each track’s: some libraries allow personal '
             f'mixing practice but not handing files to students. If the terms allow it, leave out (or mute) the drum tracks and bounce the rest into one stem.</p>'
             f'<p><b>3. Your own recordings</b>: your school band, a colleague, or you, playing without a drummer.</p>'
             f'<p><b>4. Royalty-free music</b> whose license allows you to edit it and share it with students, if you can get a version without drums.</p>'
             f'<p class="small">Only use audio you have the right to use and share in class. Do not hand out stems pulled from commercial songs.</p>'
             f'</div><div>'
             f'{H2("What makes a good stem")}<ul>' + "".join(f"<li>{g}</li>" for g in good) + '</ul>'
             f'{H2("Two folders on your class page")}'
             f'<table class="t"><tr><th>Folder</th><th>Used</th><th>What</th></tr>'
             f'<tr><td><b>Foundation</b></td><td>Days 1–2</td><td>Several 16-measure stems; students pick one that catches their ear</td></tr>'
             f'<tr><td><b>Full Song</b></td><td>Days 3–6</td><td>Several full-length stems; each student uses a new one</td></tr></table>'
             f'<p class="small">My project sheet calls the stem “a real professional recording”. With stems you made, that sentence means your stem.</p>'
             f'</div></div>')
    return page("The stems: what you supply", "STEMS", "Teacher", inner, n, key="stems")


def recipe_table(sections, total_label):
    rows, done = [], set()
    st = starts(sections)
    for (sec, nm), s in zip(sections, st):
        prog = PROGS[sec]
        rows.append(f"<tr><td style='padding:2pt 5pt'><b>{sec}</b></td><td style='white-space:nowrap;padding:2pt 5pt'>{s}–{s + nm - 1}</td>"
                    f"<td style='white-space:nowrap;padding:2pt 5pt'>{' – '.join(prog)} &nbsp;({' '.join(chord_name(r) for r in prog)})</td>"
                    f"<td style='padding:2pt 5pt'>{', '.join(PARTS[sec])}</td></tr>")
    total = sum(nm for _, nm in sections)
    return (f"<table class='t' style='font-size:9.6pt'><tr><th>Section</th><th>Measures</th><th>Chords (one per measure)</th>"
            f"<th>Instruments</th></tr>{''.join(rows)}</table>"
            f"<p class='small' style='margin:0 0 4pt'>{total} measures at {TEMPO} BPM = {mmss(seconds(total))} ({total_label}).</p>")


def makestem_page(n):
    steps = [f"New project, tempo {TEMPO} BPM, 4/4. Software Instrument tracks: a piano, a bass, a pad and a simple synth. No drum track.",
             "Turn on Cycle over 4 measures. With the Pencil in the piano roll, enter one chord per measure: three notes stacked, all white keys "
             "(the chord names below). The bass plays each chord’s root, low, as a whole note.",
             "Copy the 4-measure progression (Option-drag) to build each section; mute or delete instruments so each section has the "
             "instruments in the table. More instruments = more energy.",
             "Add arrangement markers for your own answer key. Then File › Bounce › Project or Section, WAV or MP3, starting at bar 1.",
             f"Name the file with its tempo, for example “Foundation 1 - {TEMPO} BPM”, and post it in the right folder."]
    inner = (f'<p style="margin:0 0 3pt">Two original stems, C major, white keys only. Every chord is a white-key chord of C major: '
             f'{CHORDS_USED_TEXT}. Change the chords, sounds or tempo to make more; keep the section lengths.</p>'
             f'<ol style="margin:0 0 4pt">' + "".join(f"<li>{s}</li>" for s in steps) + '</ol>'
             f'{H2("Foundation stem · 16 measures · A A B A")}{recipe_table(FOUNDATION, "Foundation")}'
             f'<div style="width:5.6in;margin:0 auto 2pt">{strip_svg("recipe foundation", FOUNDATION, W=1500, h=70)}</div>'
             f'{H2("Full Song stem · 56 measures")}{recipe_table(FULL_SONG, "Full Song")}'
             f'<p class="small" style="margin:0">Chord spellings: I = C E G · ii = D F A · iii = E G B · IV = F A C · V = G B D · vi = A C E. '
             f'Worked drum answers for both stems: pages {IX["key1"]}–{IX["key3"]}.</p>')
    return page("Make your own stems in Logic", "STEMS", "Recipe · 30 min", inner, n, key="makestem")


LOGIC_REF = [("Open a stem", "New empty project. Set the tempo (double-click it in the control bar) to the stem’s BPM. Drag the audio file into "
              "the Tracks area at bar 1."),
             ("New drum track", "Track › New Software Instrument Track; in the Library (Y) choose any kit under Electronic Drum Kit. This loads "
              "Drum Machine Designer."),
             ("Arrangement markers", "Turn on Global Tracks (G). In the Arrangement track, click + to add a marker, drag its edges to fit the "
              "section, and rename it (A, B, Verse…)."),
             ("Pattern region", "Control-click (right-click) in the drum track at the start of a section and choose Create Pattern Region. "
              "Double-click it to open the Step Sequencer."),
             ("Place hits", "Click steps to turn hits on. Each row is one sound; add rows for more sounds."),
             ("64-step mode", "In the Step Sequencer, set the pattern length to 64 steps: four measures in one pattern, so a transition can sit "
              "in its last measure."),
             ("Copy across", "Option-drag a region to copy it, or select it and press Command-R to repeat it."),
             ("Cycle", "Press C and drag the cycle area over a section to loop it while you work."),
             ("Volume and pan", "Open the Mixer (X). Click the disclosure triangle on the Drum Machine Designer track to show each drum’s own "
              "fader and pan knob."),
             ("Clipping", "The peak display above each fader turns red when that track clips. Turn something down, then click the display to reset it."),
             ("Export", "File › Bounce › Project or Section (Command-B), MP3. In Finder, Control-click the .logicx project and choose Compress.")]


def logicref_page(n):
    rows = "".join(f"<tr><td style='width:22%;padding:5pt 6pt'><b>{a}</b></td><td style='padding:5pt 6pt'>{b}</td></tr>" for a, b in LOGIC_REF)
    inner = (f'<p>Everything students need in Logic for this unit, in one table. Demonstrate once (slides 9 and 15), then let students work.</p>'
             f'<table class="t" style="font-size:12.2pt;line-height:1.46">{rows}</table>'
             f'<div class="cols" style="margin-top:8pt"><div class="box"><span class="k">The order that works</span>'
             f'<p style="margin:0">Listen twice, markers, then drums. Groove first (kick, snare, closed hi-hat), then layers, then transitions, '
             f'then the mix.</p></div>'
             f'<div class="box"><span class="k">Version note</span>'
             f'<p style="margin:0" class="small">Menu names are from current Logic Pro versions; older versions may differ slightly. Logic Pro is a '
             f'trademark of Apple Inc., named only to describe what the unit needs. This resource is not affiliated with Apple.</p></div></div>')
    return page("Logic quick reference", "SETUP", "Teacher", inner, n, key="logicref")


def pacing_page(n, watermark=False):
    rows = "".join(
        f"<tr><td style='white-space:nowrap'><b class='acc' style='font-family:Oswald;font-size:13pt'>DAY {d['n']}</b><br><span class='small'>{d['kind']}</span></td>"
        f"<td><b>{d['title']}</b><br><span class='small'>{d['goal']}</span></td>"
        f"<td style='font-size:9.3pt;line-height:1.36'>{'<br>'.join(f'{a} · {h}' for a, h, _, _ in d['blocks'])}</td>"
        f"<td style='font-size:9.3pt'>{slides_label(d)}</td></tr>" for d in DAYS)
    inner = (f'<p>Six class periods of about {MINUTES} minutes, as in my course plan: a learning day, three work days, a submission day and a '
             f'presenting day. Foundation first, then Full Song. Each day’s full plan is on pages {IX["day12"]}–{IX["day56"]}.</p>'
             f'<table class="t"><tr><th style="width:12%">Day</th><th style="width:38%">Focus and goal</th><th style="width:31%">Minute blocks</th>'
             f'<th>Deck</th></tr>{rows}</table>'
             f'{H2("Student pages by day")}<table class="t">'
             f'<tr><td style="width:12%"><b>Day 1</b></td><td>Day 1 post (p. {IX["post1"]}), project sheet (pp. {IX["proj1"]}–{IX["proj4"]}), '
             f'section map Part 1 (p. {IX["map"]})</td></tr>'
             f'<tr><td><b>Day 2</b></td><td>Gain staging (p. {IX["gain"]})</td></tr>'
             f'<tr><td><b>Day 3</b></td><td>The four transitions (p. {IX["transitions"]}), section map Part 2 (p. {IX["map"]}), name that '
             f'transition (p. {IX["transcheck"]})</td></tr>'
             f'<tr><td><b>Days 4–6</b></td><td>Checklist (p. {IX["proj4"]}) and rubric (p. {IX["rubric"]})</td></tr></table>'
             f'<div class="box" style="margin-top:6pt"><span class="k">My line for the whole unit</span><p style="margin:0;font-size:12pt">'
             f'<b>“{RED_RULE}”</b> And before anything: listen first, then respond.</p></div>')
    return page(f"Pacing guide · {DAYS_N} days", "PACING", "Teacher", inner, n, watermark, key="pacing")


def day_box(d):
    def blk(a, h, t, s):
        tag = f' <span class="small">[{"slide" if s[0] == s[1] else "slides"} {rng(*s)}]</span>' if s else ""
        return f'<div class="blk"><div class="tm">{a}</div><div><b class="h">{h}</b>{tag} · {fill(t)}</div></div>'
    return (f'<div class="day"><div class="dh"><b>Day {d["n"]} · {d["title"]}</b><span>{d["kind"]} · {slides_label(d)}</span></div>'
            f'<p class="goal"><b>Goal:</b> {d["goal"]}</p>{"".join(blk(*b) for b in d["blocks"])}</div>')


def day12_page(n):
    return page("Day-by-day plans · Days 1 and 2", "DAYS 1–2", "Foundation", '<div class="dp">' + day_box(DAYS[0]) + day_box(DAYS[1]) + '</div>', n, key="day12")


def day34_page(n):
    prompts = ('<div class="box"><span class="k">Circulating prompts (Days 1–4)</span><ul style="margin:0">'
               '<li>‘Where are your sections? Show me before you build.’</li>'
               '<li>‘Play A, then play B. Can you feel the difference without looking?’</li>'
               '<li>‘Why did you choose this type for this moment?’</li></ul></div>')
    listen = (f'<div class="box acc-b" style="margin-top:8pt"><span class="k">What to listen for at every transition</span><ul style="margin:0">'
              f'<li>{TIPS[2]}</li><li>{TIPS[3]}</li><li>Only the measure before the change moves; the groove around it stays. {DRUM_LINE}</li></ul></div>')
    prompts = prompts + listen
    return page("Day-by-day plans · Days 3 and 4", "DAYS 3–4", "Full Song", '<div class="dp">' + day_box(DAYS[2]) + day_box(DAYS[3]) + prompts + '</div>', n, key="day34")


def day56_page(n):
    close = (f'<div class="box acc-b"><span class="k">The close (Day 6)</span><p style="margin:0;font-size:11.5pt">“{CLOSE}”</p>'
             f'<p style="margin:3pt 0 0">From my course plan: name what students just did. <b>“You mixed a song.”</b> They adjusted volumes and '
             f'panned instruments without knowing that was mixing; naming it gives them ownership of a skill they already have.</p></div>')
    qs = ('<div class="box" style="margin-top:8pt"><span class="k">Three feedback questions: examples to get students started (Day 5)</span>'
          '<ul style="margin:0"><li>“Does my B pattern feel different enough from my A pattern?”</li>'
          '<li>“Which of my transitions worked best, and did any feel accidental?”</li>'
          '<li>“Do my drums sit under the stem, or on top of it?”</li></ul>'
          '<p class="small" style="margin:3pt 0 0">Students write their own; these only model the kind of question that gets useful answers.</p></div>')
    close = close + qs
    return page("Day-by-day plans · Days 5 and 6", "DAYS 5–6", "Submit + present", '<div class="dp">' + day_box(DAYS[4]) + day_box(DAYS[5]) + close + '</div>', n, key="day56")


TN1_DIFF = [
    "My course plan describes an older version of this unit (Group 1: drums in AABA over a 16-bar stem; Group 2: full-song drums with "
    "volume, panning and reverb). This kit follows my newest project sheet, Foundation and Full Song. AABA lives in the 16-measure "
    f"Foundation stem (sections of 4), and the reverb bus is an optional, ungraded extension (slide 17, page {IX['gain']}).",
    "My stems and my example Logic session are not included (third-party audio). Pages "
    f"{IX['stems']}–{IX['makestem']} replace them. The Day 1 post’s “Check the example!” link becomes “{POST_EXAMPLE}”",
    "My LMS is written as “your class page”. Padlet is kept: it is the board I use; any feedback board works.",
    "Day plans: Days 1 and 3 end with a 5-minute save block (my notes stop at minute 48); Day 4 is one 53-minute block. Slide numbers are "
    "added to every block.",
    "Part 2 step 1: one closing parenthesis added. Otherwise the wording of my project sheet, rubric (/25), Day 1 post and teacher notes is "
    "kept as I wrote it, including the checklist’s “Your Name - Foundations” for both projects (see page 2).",
    "My Day 1 post and project sheet say to label sections “A, B, A, B, etc.”: that is an example of labelling, kept as written. The "
    f"recipe stem and worked example in this kit are A, A, B, A (pages {IX['makestem']} and {IX['key1']}); your stems may put A and B in any order.",
    "Added to my day plans: the Day 3 exit question (“which transition type have you not tried yet?”) and the Day 4 note that early "
    "finishers may try the optional reverb bus.",
    f"New: the stems pages, the Logic reference, the four transitions page, gain staging, the section map, name that transition, the worked "
    f"examples and answer keys, and the deck."]


def tn1_page(n):
    inner = (f'<div class="tn">{H2("What this unit reinforces", True)}'
             + "".join(f"<p>{p}</p>" for p in WHAT) +
             f'<div class="box" style="margin-top:4pt"><span class="k">How this kit differs from my classroom version</span><ul style="margin:0">'
             + "".join(f"<li style='font-size:10.4pt;line-height:1.35;margin-bottom:3pt'>{x}</li>" for x in TN1_DIFF) + '</ul></div>'
             f'<div class="box acc-b" style="margin-top:6pt"><span class="k">The close (Day 6)</span><p style="margin:0">“{CLOSE}”</p></div></div>')
    return page("Teacher notes · the big idea", "TEACHER NOTES", "1 of 3", inner, n, key="tn1")


def tn2_page(n):
    st = "".join(f'<p><b>{a}:</b> {b}</p>' for a, b in STRUGGLES)
    inner = f'<div class="tn huge">{H2("Common student struggles", True)}{st}</div>'
    return page("Teacher notes · common struggles", "TEACHER NOTES", "2 of 3", inner, n, key="tn2")


def tn3_page(n):
    inner = (f'<div class="tn">{H2("Suggested variations", True)}'
             f'<p><b>More accessible:</b> {VAR_EASY}</p><p><b>More advanced:</b> {VAR_HARD}</p>'
             f'{H2("Mixing concept: gain staging (Day 2)")}<p>{PLAN_MIX}</p>'
             f'{H2("Grading tips")}<ul>'
             f'<li>Five categories, each scored 2–5, /{RUBRIC_TOTAL} (page {IX["rubric"]}). Foundations and Technical Skill come from the '
             f'Foundation project; Full Song from the second; Gain Staging and Documentation from both.</li>'
             f'<li>Play A, then B, without looking: is the contrast obvious? Then play every section change of the full song: is there a '
             f'deliberate transition at each, two types or more?</li>'
             f'<li>Watch the meters on the loudest section: any red? Is the stem still the lead?</li>'
             f'<li>Worked examples and answer keys: pages {IX["key1"]}–{IX["key3"]}.</li></ul>'
             f'{H2("Resources")}<ul>'
             f'<li>Stem tracks: the Foundation folder and the Full Song folder on your class page (update each year as needed); pages '
             f'{IX["stems"]}–{IX["makestem"]}.</li>'
             f'<li>cambridge-mt.com/ms/mtk: a multitrack library I use for sourcing new tracks (check each track’s terms).</li>'
             f'<li>The intro deck (.pptx and PDF), included, in place of my older video tutorials.</li></ul>'
             f'{H2("Observation notes")}'
             f'<p class="small">What worked, what to adjust, and anything worth remembering for next year.</p></div><div class="lines"></div>')
    return page("Teacher notes · variations, mixing, grading", "TEACHER NOTES", "3 of 3", inner, n, key="tn3")


def post1_page(n):
    lis = "".join(f'<li style="margin-bottom:11pt;font-size:12pt;line-height:1.36"><div class="n">{i + 1}</div><div><b class="hd" style="font-size:13pt">{a}</b> · {bb}<br>'
                  f'<span class="why">{cc}</span></div></li>' for i, (a, bb, cc) in enumerate(POST_STEPS))
    box = ("".join(f'<p style="margin:0 0 2pt;font-size:11pt"><b>{t}</b> — {TYPE_DEF[t]}</p>' for t in TYPES)
           + f'<p style="margin:3pt 0 0;font-size:11pt" class="why">{POST_EXAMPLE}</p>')
    inner = (f'<h2 class="s" style="font-size:20pt;margin:4pt 0 2pt">Stems &amp; Grooves — Day 1</h2>'
             f'<p class="l" style="font-size:12pt">Digital Music Production 1 · Unit {UNIT}</p>'
             f'<p style="font-size:12pt;margin:8pt 0 6pt">{POST_INTRO}</p>'
             f'<p style="font-size:12pt;margin:0 0 10pt"><b>{POST_BOLD}</b></p>'
             f'<ol class="num">{lis}</ol>'
             f'<div class="box acc-b" style="margin-top:2pt"><span class="k" style="font-size:11.5pt">The four transition types — you will use '
             f'these in Part 2:</span>{box}</div>')
    return page("Classroom post", "STUDENT", "Day 1", inner, n, key="post1")


def transitions_page(n, watermark=False):
    blocks = "".join(
        f'<div style="margin:0 0 5pt"><div style="display:flex;gap:8pt;align-items:baseline"><b style="font-family:Oswald;font-size:15pt;'
        f'text-transform:uppercase;color:{ACCENT}">{i + 1} · {t}</b><span style="font-size:10.8pt">{TYPE_DEF[t]}.</span></div>'
        f'<p class="why" style="margin:0 0 2pt;font-size:10pt">Use it: {TYPE_WHEN[t]}.</p>'
        f'<div style="width:6.2in;margin:0 auto">{grid_svg(f"page trans {t}", [PATTERN_A, TRANS[t]], ["Groove (A)", "Transition measure"], cw=16, ch=19, gut=140, lanes=lanes_of(PATTERN_A, TRANS[t]))}</div></div>'
        for i, t in enumerate(TYPES))
    inner = (f'<p style="margin:0 0 5pt;font-size:10.8pt">A transition is the moment between two sections where the drums signal that a change '
             f'is coming. Each example shows one measure of the groove, then the transition measure (green) right before the new section. '
             f'Each grid column is a sixteenth note; grey columns are the beats.</p>{blocks}')
    return page("The four transitions", "STUDENT", "Reference", inner, n, watermark, key="transitions")


def gain_page(n):
    steps = ["Start with the stem: it is the lead. Leave it where it sits.",
             "Bring the drums up until they support the stem: you can hear every drum, and the stem is still the star.",
             "Play the loudest section and watch the meters, including the master (Stereo Out).",
             "Anything in the red? Turn other things down. Do not turn one thing up.",
             "Balance the drums against each other: kick and snare first, hi-hats just above them."]
    inner = (f'<div class="rule"><span class="k">The rule</span><b>{RED_RULE}</b></div>'
             f'<p style="margin:6pt 0 4pt;font-size:11pt"><b>Gain staging</b> — {dict(VOCAB)["Gain staging"]}.</p>'
             f'<div style="width:6.3in;margin:2pt auto 2pt">{body_svg(meter_body())}</div>'
             f'<div class="cols" style="gap:14pt"><div>{H2("Gain staging in 5 steps", True)}<ol>'
             + "".join(f"<li style='font-size:10.8pt'>{s}</li>" for s in steps) + '</ol></div>'
             f'<div>{H2("Check your mix", True)}'
             + "".join(f'<div class="chk" style="font-size:10.8pt;margin-bottom:5pt">{checkbox(12)}<span>{c}</span></div>' for c in [
                 "Nothing clips: no red anywhere, in the loudest section too",
                 "Every drum is audible",
                 "Nothing in my drums is louder than the loudest part of the stem",
                 "It sounds like one song, not drums on top of a song"]) +
             f'</div></div>'
             f'<div class="box" style="margin-top:4pt"><span class="k">Going further (optional, not graded): a reverb bus</span>'
             f'<p style="margin:0;font-size:10.4pt">In the Mixer (X), click a Send slot on the drums and choose a bus. Logic adds an Aux channel '
             f'strip: put ChromaVerb in its Audio FX slot, then turn the Send knob up slowly. A little shared reverb glues the drums to the '
             f'stem. Unit 6 teaches buses properly.</p></div>')
    return page("Gain staging: nothing in the red", "STUDENT", "Day 2", inner, n, name=True, key="gain")


def map_page(n):
    boxes = "".join(f'<td style="width:25%;height:46pt;vertical-align:top"><span class="small">Measures {s}–{s + 3}</span><br>'
                    f'<span style="font-size:10pt">Section letter:</span> <span style="display:inline-block;width:46pt;border-bottom:1pt solid #000"></span></td>'
                    for s in (1, 5, 9, 13))
    rows = "".join(f"<tr><td style='text-align:center'>{i}</td><td></td><td></td><td></td><td></td></tr>" for i in range(1, 10))
    inner = (f'{H2("Part 1 · Foundation stem (16 measures)", True)}'
             f'<p style="margin:0 0 4pt">Listen twice before you touch anything. Then label each 4-measure section A or B.</p>'
             f'<table class="map"><tr>{boxes}</tr></table>'
             f'<div class="fl" style="margin-top:7pt"><b>My stem:</b><span></span><b>Tempo (BPM):</b><span style="flex:0 0 0.9in"></span></div>'
             f'<div class="fl"><b>Where does the energy change? What changes?</b><span></span></div><div class="ln2" style="margin-bottom:6pt"></div>'
             f'<div class="fl"><b>My A pattern will feel:</b><span></span><b>My B pattern will feel:</b><span></span></div>'
             f'{H2("Part 2 · Full Song stem: map every section change")}'
             f'<p style="margin:0 0 4pt">Label the sections with letters (a, b, c…). Count the changes before you write a single transition: '
             f'every change needs one.</p>'
             f'<table class="map"><tr><th style="width:7%">#</th><th style="width:14%">Starts at measure</th><th style="width:17%">Section + letter</th>'
             f'<th style="width:22%">Transition into the next (fill, breakdown, cut, build)</th><th>Why this type for this moment?</th></tr>{rows}</table>'
             f'<div class="fl" style="margin-top:7pt"><b>Number of section changes:</b><span style="flex:0 0 0.8in"></span>'
             f'<b>Transition types I used:</b><span></span></div>')
    return page("Section map", "STUDENT", "Days 1 + 3", inner, n, name=True, key="map")


def proj_head(sub):
    return (f'<h2 class="s" style="font-size:20pt;margin:2pt 0 0">Stems &amp; Grooves</h2>'
            f'<p class="l" style="font-size:12pt;margin-bottom:4pt">Digital Music Production 1 · Unit {UNIT} · Structured Project · '
            f'{DAYS_N} Days · Drum Machine Designer + Stems{sub}</p>')


def proj1_page(n):
    rows = "".join(f"<tr><td style='width:24%;padding:4pt 6pt'><b>{a}</b></td><td style='padding:4pt 6pt'>{b}</td></tr>" for a, b in VOCAB)
    inner = (proj_head("") + f'{H2("Overview")}<p style="font-size:11pt">{OVERVIEW}</p>'
             f'{H2("Goals")}<ul>' + "".join(f"<li style='font-size:11pt'>{g}</li>" for g in GOALS) + '</ul>'
             f'{H2("Key Vocabulary")}<table class="t" style="font-size:10.6pt">{rows}</table>')
    return page("Project sheet", "STUDENT", "1 of 4", inner, n, name=True, key="proj1")


def part_band(t):
    return f'<div class="box inv" style="padding:4pt 10pt;margin:6pt 0 4pt"><b style="font-family:Oswald;font-size:13pt;letter-spacing:.04em">{t}</b></div>'


def proj2_page(n, watermark=False):
    req = lambda lst: "".join(f'<div class="chk" style="font-size:12pt;margin-bottom:8pt;line-height:1.35">{checkbox(13)}<span>{r}</span></div>' for r in lst)
    inner = (f'{H2("Requirements", True)}{part_band("PART 1 — FOUNDATION")}{req(REQ1)}'
             f'{part_band("PART 2 — FULL SONG")}{req(REQ2)}')
    return page("Project sheet · requirements", "STUDENT", "2 of 4", inner, n, watermark, name=True, key="proj2")


def steplist(steps, extra=None):
    out = []
    for i, (h, t) in enumerate(steps):
        sub = ""
        if extra and i == 1:
            sub = "<ul style='margin:3pt 0 0'>" + "".join(f"<li style='font-size:10pt;margin-bottom:1pt'><b>{a}</b> — {b}</li>" for a, b in extra) + "</ul>"
        out.append(f'<li style="margin-bottom:5pt"><div class="n">{i + 1}</div><div><b class="hd" style="font-size:11pt">{h}</b>'
                   f'<p style="margin:1pt 0 0;font-size:10.2pt;line-height:1.28">{t}</p>{sub}</div></li>')
    return f'<ol class="num">{"".join(out)}</ol>'


def proj3_page(n):
    s2 = [(STEPS2[0][0] + ")", STEPS2[0][1])] + STEPS2[1:]   # one closing parenthesis added (disclosed)
    inner = (f'{H2("Steps / Process", True)}{part_band("PART 1 — FOUNDATION")}{steplist(STEPS1)}'
             f'{part_band("PART 2 — FULL SONG")}{steplist(s2, STEP2_TYPES)}')
    return page("Project sheet · steps", "STUDENT", "3 of 4", inner, n, name=True, key="proj3")


def proj4_page(n):
    inner = (f'{H2("Tips for Success", True)}<ul>' + "".join(f"<li style='font-size:10.6pt;margin-bottom:3pt'>{t}</li>" for t in TIPS) + '</ul>'
             f'{H2("Submission Checklist")}'
             + "".join(f'<div class="chk" style="font-size:10.6pt;margin-bottom:4pt">{checkbox(12)}<span>'
                       f'{("<b>" + p + ":</b> ") if p else ""}{c}</span></div>' for p, c in CHECK) +
             f'{H2("Notes")}<div class="lines"></div>')
    return page("Project sheet · tips and checklist", "STUDENT", "4 of 4", inner, n, name=True, key="proj4")


def rubric_page(n, watermark=False):
    head = "<tr><th>Category</th>" + "".join(f"<th>{lv}</th>" for lv in LEVELS) + "<th>Score</th></tr>"
    rows = "".join(f"<tr><td class='c'>{escape(c)}</td>" + "".join(f"<td>{d}</td>" for d in ds) + "<td class='sc'></td></tr>" for c, ds in RUBRIC)
    total = (f"<tr><td class='c' colspan='5' style='text-align:right;font-size:14pt'>Total</td>"
             f"<td style='font-family:Oswald;font-weight:700;font-size:14pt;text-align:right;vertical-align:bottom'>/ {RUBRIC_TOTAL}</td></tr>")
    inner = (f'<h2 class="s" style="font-size:18pt;margin:2pt 0 0">Stems &amp; Grooves — Rubric</h2>'
             f'<p class="l" style="font-size:12pt;margin-bottom:4pt">Digital Music Production 1 · Unit {UNIT} · /{RUBRIC_TOTAL}</p>'
             f'<p style="font-size:10.5pt;margin-bottom:4pt">Each category is scored from 2 to 5.</p>'
             f'<table class="rub">{head}{rows}{total}</table>'
             f'<div class="cols" style="margin-top:6pt"><div><h2 class="s" style="margin-top:0">My favorite transition, and why…</h2>'
             f'<div class="ln2"></div><div class="ln2"></div></div>'
             f'<div><h2 class="s" style="margin-top:0">Teacher comments</h2><div class="ln2"></div><div class="ln2"></div></div></div>')
    return page("Rubric", "RUBRIC", f"/ {RUBRIC_TOTAL}", inner, n, watermark, key="rubric")


def transcheck_page(n):
    cells = "".join(
        f'<div style="margin:0 0 6pt"><div style="display:flex;gap:10pt;align-items:baseline"><b style="font-family:Oswald;font-size:15pt">{i + 1}</b>'
        f'<div class="fl" style="flex:1;margin:0"><b>Transition type:</b><span style="flex:0 0 1.5in"></span><b>What does it do to the energy?</b><span></span></div></div>'
        f'<div style="width:6.9in;margin:3pt auto 0">{grid_svg(f"page quiz {i}", [PATTERN_A, q], ["Groove (A)", "?"], cw=16, ch=21, gut=140, lanes=lanes_of(PATTERN_A, q))}</div></div>'
        for i, q in enumerate(QUIZ))
    inner = (f'<p style="margin:0 0 6pt;font-size:11pt">Each grid shows one measure of the groove, then the measure right before a new section. '
             f'Name the transition: drum fill, breakdown, cut or build. Use the definitions on page {IX["transitions"]}.</p>{cells}')
    return page("Name that transition", "STUDENT", "Day 3 · quick check", inner, n, name=True, key="transcheck")


def contrast(a, b):
    return sum(1 for ln in LANES for i in range(STEPS) if a.get(ln, "." * STEPS)[i] != b.get(ln, "." * STEPS)[i])


def key1_page(n):
    lay = lambda p: ", ".join(lanes_of(p))
    inner = (f'<p style="margin:0 0 4pt">One possible answer for the Foundation project, over the Foundation stem from the recipe (page '
             f'{IX["makestem"]}): {sum(m for _, m in FOUNDATION)} measures at {TEMPO} BPM, sections of 4 measures in the order '
             f'{", ".join(s for s, _ in FOUNDATION)}.</p>'
             f'<div style="width:7.2in;margin:2pt auto 4pt">{body_svg(patterns_body())}</div>'
             f'<table class="t" style="font-size:10pt"><tr><th style="width:13%">Pattern</th><th style="width:28%">Layers</th><th>Feel, and why it fits</th></tr>'
             f'<tr><td><b>A</b> (m. 1–8, 13–16)</td><td>{len(lanes_of(PATTERN_A))}: {lay(PATTERN_A)}</td><td>Laid back: closed hi-hat eighths, '
             f'snare on 2 and 4, a kick that skips ahead on the “and” of 3. It sits under the piano-and-bass sections.</td></tr>'
             f'<tr><td><b>B</b> (m. 9–12)</td><td>{len(lanes_of(PATTERN_B))}: {lay(PATTERN_B)}</td><td>More energy: a kick on every beat and '
             f'open hi-hats off the beat. The stem adds a pad and a synth here, so the drums lift with it.</td></tr></table>'
             f'<div class="cols" style="gap:14pt;margin-top:4pt"><div class="box"><span class="k">Contrast check</span><p style="margin:0;font-size:10.2pt">'
             f'A and B differ in {contrast(PATTERN_A, PATTERN_B)} of their 16-step cells: a new sound (open hi-hat), a new kick pattern and a '
             f'different hi-hat rhythm. Play A, then B: you can hear it without looking.</p></div>'
             f'<div class="box"><span class="k">Gain staging check</span><p style="margin:0;font-size:10.2pt">Loudest section (B): no red on any '
             f'meter. Kick and snare clear, hi-hats just under them; the stem still leads.</p></div></div>'
             f'{H2("Program it in Logic (x = a hit, . = a rest, 16 steps)")}'
             f'<table class="t" style="font-size:10.4pt"><tr><th style="width:22%">Sound</th><th>A pattern</th><th>B pattern</th></tr>'
             + "".join(f"<tr><td><b>{ln}</b></td><td class='stp'>{PATTERN_A.get(ln, '—')}</td><td class='stp'>{PATTERN_B.get(ln, '—')}</td></tr>"
                       for ln in lanes_of(PATTERN_A, PATTERN_B)) + '</table>'
             f'{H2("Placement")}<table class="t" style="font-size:10.4pt"><tr><th>Section</th><th>Measures</th><th>Pattern region</th></tr>'
             + "".join(f"<tr><td><b>{sec}</b></td><td>{st_}–{st_ + nm - 1}</td><td>{sec} pattern, 4 measures (one region, repeated)</td></tr>"
                       for (sec, nm), st_ in zip(FOUNDATION, starts(FOUNDATION))) + '</table>'
             f'<p class="small" style="margin-top:5pt">Students’ answers will differ. Grade against the rubric: is the contrast obvious, does '
             f'the placement follow the stem’s sections, at least 3 layers, nothing clipping?</p>')
    return page("Example · Foundation patterns", "EXAMPLE", "Teacher", inner, n, ak=True, key="key1")


def describe(t):
    m = TRANS[t]
    c = beat_lanes(m)
    return {"Drum fill": "Beats 1–2 keep the groove; beats 3–4 break into snare sixteenths, then the high and low toms.",
            "Breakdown": "Only the kick is left, on every beat: one element.",
            "Cut": f"The groove plays to the end of beat 2, then {trailing_silence(m) // 4} beats of complete silence.",
            "Build": f"Sounds join beat by beat: {', '.join(str(x) for x in c)} sounds on beats 1–4."}[t]


def key2_page(n):
    rows = "".join(f"<tr><td><b>{t}</b></td><td>{describe(t)}</td><td style='font-size:9.4pt'>{'<br>'.join(f'{ln}: <span class=\"stp\">{TRANS[t][ln]}</span>' for ln in lanes_of(TRANS[t]))}</td></tr>"
                   for t in TYPES)
    quiz = "".join(f"<tr><td style='text-align:center'><b>{i + 1}</b></td><td><b>{classify(q)}</b></td><td>{qwhy(q)}</td></tr>"
                   for i, q in enumerate(QUIZ))
    inner = (f'<p style="margin:0 0 4pt">The four transitions on page {IX["transitions"]} and slide 13, step by step, so you can program them in '
             f'Logic before Day 1 (each comes after one measure of the A groove). x = a hit, . = a rest; 16 steps per measure.</p>'
             f'<table class="t" style="font-size:10.4pt"><tr><th style="width:12%">Type</th><th style="width:36%">What happens</th>'
             f'<th>Steps (transition measure)</th></tr>{rows}</table>'
             f'{H2("Name that transition · answers (page " + str(IX["transcheck"]) + ")")}'
             f'<table class="t" style="font-size:10.8pt"><tr><th style="width:7%">#</th><th style="width:16%">Answer</th><th>How to tell</th></tr>{quiz}</table>'
             f'{H2("Listening answers")}<table class="t" style="font-size:10.8pt">'
             f'<tr><td style="width:24%"><b>The stem, no drums</b> (slide 2)</td><td>Answers depend on your stem. Listen for instruments '
             f'entering or leaving and the music getting louder or softer: that is where the sections change.</td></tr>'
             f'<tr><td><b>{ITAT[0]}</b> ({ITAT[1]})</td><td>A quiet drum machine and keyboards keep time; the big drums are held back. Right '
             f'before they arrive there is a famous tom fill. It hits hard because of the space before it: the fill is a transition that '
             f'announces a new section.</td></tr></table>')
    return page("Answer key · transitions", "ANSWER KEY", "Teacher", inner, n, ak=True, key="key2")


def qwhy(q):
    t = classify(q)
    if t == "Breakdown":
        return f"Only one sound is left ({lanes_of(q)[0].lower()})."
    if t == "Cut":
        return f"The groove stops with {trailing_silence(q) // 4} beat{'s' if trailing_silence(q) > 4 else ''} of silence before the new section."
    if t == "Build":
        return f"Sounds join beat by beat ({', '.join(str(x) for x in beat_lanes(q))} on beats 1–4)."
    return "The groove breaks into a rhythmic variation across the drums (snare, then toms)."


def key3_page(n):
    st = starts(FULL_SONG)
    rows = []
    for k, ((sec, nm), s) in enumerate(zip(FULL_SONG, st)):
        tr = SONG_TRANS[k] if k < len(SONG_TRANS) else None
        groove = {"Intro": "none, or hi-hat only", "Verse": "A", "Chorus": "B", "Bridge": "A, half the kicks", "Outro": "none"}[sec]
        rows.append(f"<tr><td><b>{sec}</b> ({LETTER[sec]})</td><td>{s}–{s + nm - 1}</td><td>{energy(sec)}</td><td>{groove}</td>"
                    f"<td>{(f'<b>{tr[0]}</b> in m. {s + nm - 1}: {tr[1]}') if tr else 'End: let the last chord ring'}</td></tr>")
    total = sum(nm for _, nm in FULL_SONG)
    used = sorted({t for t, _ in SONG_TRANS}, key=TYPES.index)
    inner = (f'<p style="margin:0 0 4pt">One possible map of the Full Song stem from the recipe (page {IX["makestem"]}): {total} measures at '
             f'{TEMPO} BPM ({mmss(seconds(total))}), {len(FULL_SONG)} sections, so {len(FULL_SONG) - 1} section changes and '
             f'{len(SONG_TRANS)} transitions.</p>'
             f'<div style="width:7.4in;margin:8pt auto 12pt">{body_svg(songmap_body(W=1250, h=96))}</div>'
             f'<table class="t" style="font-size:11pt;line-height:1.55"><tr><th style="width:15%">Section</th><th style="width:10%">Measures</th>'
             f'<th style="width:9%">Energy</th><th style="width:15%">Groove</th><th>Transition into the next section, and why</th></tr>{"".join(rows)}</table>'
             f'<p style="font-size:10pt;margin:2pt 0 4pt">Energy = the number of instruments in the stem in that section. Types used: '
             f'{", ".join(used)} ({len(used)}; the project needs at least two).</p>'
             f'<div class="box" style="margin-bottom:6pt"><span class="k">Check the math</span><p style="margin:0;font-size:10.4pt">'
             f'{" + ".join(str(nm) for _, nm in FULL_SONG)} = {total} measures. At {TEMPO} BPM one measure of 4/4 lasts {seconds(1):.1f} seconds, '
             f'so the stem runs {total} × {seconds(1):.1f} = {seconds(total):.0f} seconds ({mmss(seconds(total))}). Each transition sits in the '
             f'last measure of its section: measures {", ".join(str(s + nm - 1) for (_, nm), s in list(zip(FULL_SONG, st))[:-1])}.</p></div>'
             f'<div class="cols" style="gap:14pt"><div class="box"><span class="k">My rule of thumb</span><p style="margin:0;font-size:10.2pt">'
             f'“A cut before a quiet verse, a build before a big chorus.” Here every build leads into more energy and the cut leads into less.</p></div>'
             f'<div class="box"><span class="k">Grading the Full Song</span><p style="margin:0;font-size:10.2pt">Count the section changes in '
             f'the student’s stem, then play each one: a deliberate transition at every change, two types or more, each fitting its moment.</p></div></div>')
    return page("Example · a mapped full song", "EXAMPLE", "Teacher", inner, n, ak=True, key="key3")


def terms_page(n):
    return (f'<section class="page" data-key="terms"><div class="top"><div class="t">Terms of Use<small>Thank you for your purchase!</small></div></div>'
            f'<div class="main"><p>© {YEAR} {STORE}. {BRAND} by {STORE}. All rights reserved. This purchase gives <b>one teacher</b> a license to use this '
            f'resource (the PDF and the intro deck) with their own students.</p>'
            f'<h2 class="s">You may</h2><ul><li>Print and copy pages for your own classroom and students, year after year.</li>'
            f'<li>Present the intro deck to your own classes and edit it for your own teaching.</li>'
            f'<li>Share pages with your own students through a password-protected class site or learning platform.</li>'
            f'<li>Leave the pages for a substitute teacher who teaches your own students.</li></ul>'
            f'<h2 class="s">You may not</h2><ul><li>Share, email or copy this resource for other teachers, a whole school or a district. '
            f'Additional licenses are available on Teachers Pay Teachers.</li>'
            f'<li>Post any part of it on a public website, shared drive or social media.</li>'
            f'<li>Sell, give away or claim any part of it as your own, or edit it into a new product.</li></ul>'
            f'<h2 class="s">Questions?</h2><p>Please use the Q&amp;A tab on the {BRAND} by {STORE} store on Teachers Pay Teachers. '
            f'If something looks wrong, let me know there and I will fix it.</p>'
            f'<h2 class="s">Credits</h2><ul>'
            f'<li>Curriculum, project sheet, rubric, teacher notes, classroom post and intro deck: original work by {STORE}, {BRAND}.</li>'
            f'<li>Drum grids, section maps, meters, stem recipes, worked examples and layouts: original work, {BRAND}.</li>'
            f'<li>Text fonts: Oswald (The Oswald Project Authors) and Source Sans 3 (Adobe), SIL Open Font License 1.1.</li>'
            f'<li>No stems, audio, video, lyrics or album art are included. Songs are named for reference only (title and artist); play them from '
            f'your own YouTube or Spotify access. Song titles belong to their owners; this resource is not affiliated with or endorsed by the '
            f'artists or their labels. Any stem you use must be one you have the right to use in class.</li>'
            f'<li>Logic Pro and ChromaVerb are trademarks of Apple Inc., named only to describe compatibility; this resource is not affiliated with '
            f'or endorsed by Apple. Padlet, YouTube, Spotify and the Cambridge-MT multitrack library are named only as examples of tools you may '
            f'already use.</li></ul>'
            f'<p class="small">{DISCLOSURE}</p><div style="flex:1"></div></div>{foot(f"Page {n}")}</section>')


def render_page(key, watermark=False):
    n = IX[key]
    f = {"cover": cover_page, "howto": lambda: howto_page(n), "overview": lambda: overview_page(n), "setup": lambda: setup_page(n),
         "stems": lambda: stems_page(n), "makestem": lambda: makestem_page(n), "logicref": lambda: logicref_page(n),
         "pacing": lambda: pacing_page(n, watermark), "day12": lambda: day12_page(n), "day34": lambda: day34_page(n),
         "day56": lambda: day56_page(n), "tn1": lambda: tn1_page(n), "tn2": lambda: tn2_page(n), "tn3": lambda: tn3_page(n),
         "post1": lambda: post1_page(n), "transitions": lambda: transitions_page(n, watermark), "gain": lambda: gain_page(n),
         "map": lambda: map_page(n), "proj1": lambda: proj1_page(n), "proj2": lambda: proj2_page(n, watermark),
         "proj3": lambda: proj3_page(n), "proj4": lambda: proj4_page(n), "rubric": lambda: rubric_page(n, watermark),
         "transcheck": lambda: transcheck_page(n), "key1": lambda: key1_page(n), "key2": lambda: key2_page(n),
         "key3": lambda: key3_page(n), "terms": lambda: terms_page(n)}
    return f[key]()


# ================================================================ intro deck (.pptx, python-pptx) + PDF copy (LibreOffice)
def slide_day(i):
    return SLIDES[i - 1][0]


def build_deck(img):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

    BLACK, WHITE, ACC = RGBColor(0, 0, 0), RGBColor(255, 255, 255), RGBColor.from_string(ACCENT[1:])
    LIGHT = RGBColor.from_string(ACC_LIGHT[1:])
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.author = f"{STORE} · {BRAND}"
    prs.core_properties.title = f"{TITLE} · Intro deck"
    prs.core_properties.comments = DISCLOSURE
    layout = prs.slide_layouts[6]

    def box(sl, x, y, w, h, fillc=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=1.5):
        s = sl.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        if fillc is None:
            s.fill.background()
        else:
            s.fill.solid()
            s.fill.fore_color.rgb = fillc
        if line is None:
            s.line.fill.background()
        else:
            s.line.color.rgb = line
            s.line.width = Pt(lw)
        s.shadow.inherit = False
        return s

    def text(sl, x, y, w, h, runs, size=24, font="Source Sans 3", bold=False, color=BLACK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             spacing=None, caps=False):
        tb = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.04)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        paras = runs if isinstance(runs, list) else [runs]
        for k, ptxt in enumerate(paras):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.alignment = align
            r = p.add_run()
            r.text = ptxt.upper() if caps else ptxt
            r.font.size = Pt(size)
            r.font.name = font
            r.font.bold = bold
            r.font.color.rgb = color
            if spacing:
                r._r.get_or_add_rPr().set("spc", str(spacing))
        return tb

    def chrome(sl, i, title, day):
        text(sl, 0.6, 0.32, 6, 0.4, BRAND, 13, "Oswald", False, BLACK, spacing=600, caps=True)
        box(sl, 11.33, 0.3, 1.4, 0.42, ACC, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        text(sl, 11.33, 0.3, 1.4, 0.42, f"Day {day}", 15, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, caps=True)
        text(sl, 0.6, 0.72, 12.1, 0.95, title, 40, "Oswald", True, BLACK, caps=True, anchor=MSO_ANCHOR.BOTTOM)
        box(sl, 0.6, 1.72, 12.13, 0.04, BLACK)
        box(sl, 0.6, 7.0, 12.13, 0.012, BLACK)
        text(sl, 0.6, 7.04, 9, 0.3, f"{TITLE} · Intro deck · © {YEAR} {STORE} · {BRAND}", 10, "Source Sans 3", False, RGBColor(0x33, 0x33, 0x33))
        text(sl, 11.2, 7.04, 1.53, 0.3, f"{i}", 10, "Oswald", True, BLACK, align=PP_ALIGN.RIGHT)

    def bullets(sl, items, x=0.6, y=2.05, w=12.1, size=28, gap=1.75):
        for k, it in enumerate(items):
            yy = y + k * (size / 72 * gap)
            box(sl, x, yy + size / 72 * 0.36, 0.16, 0.16, ACC, shape=MSO_SHAPE.OVAL)
            text(sl, x + 0.38, yy, w - 0.4, size / 72 * 1.6, it, size, "Source Sans 3", False, BLACK)

    def pic(sl, key, x, y, w=None, h=None):
        path, (pw, ph) = img[key]
        if w is None:
            w = h * pw / ph
        sl.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w))
        return y + w * ph / pw

    def fit(key, maxw, maxh):
        pw, ph = img[key][1]
        return min(maxw, maxh * pw / ph)

    def tile(sl, x, y, w, h, head, sub, home=False, hs=40, ss=20):
        box(sl, x, y, w, h, ACC if home else BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        text(sl, x, y + 0.08, w, h * 0.42, head, hs, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        text(sl, x + 0.15, y + h * 0.45, w - 0.3, h * 0.52, sub, ss, "Source Sans 3", False, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP)

    for i, (day, kind, title, body, notes) in enumerate(SLIDES, start=1):
        sl = prs.slides.add_slide(layout)
        items = [b for b in body.split("|") if b]
        if kind == "title":
            text(sl, 0, 0.42, 13.333, 0.5, BRAND, 16, "Oswald", False, BLACK, align=PP_ALIGN.CENTER, spacing=800, caps=True)
            w = fit("hero", 9.6, 3.25)
            pic(sl, "hero", (13.333 - w) / 2, 0.98, w)
            text(sl, 0, 4.3, 13.333, 1.1, title, 72, "Oswald", True, BLACK, align=PP_ALIGN.CENTER, caps=True)
            text(sl, 0, 5.5, 13.333, 0.6, body, 28, "Oswald", False, BLACK, align=PP_ALIGN.CENTER)
            box(sl, 0, 6.55, 13.333, 0.95, BLACK)
            text(sl, 0, 6.68, 13.333, 0.6, f"Digital Music Production · Unit {UNIT} · Day {day} · Drums over stems in Logic Pro", 20, "Oswald",
                 False, WHITE, align=PP_ALIGN.CENTER, spacing=200, caps=True)
        else:
            chrome(sl, i, title, day)
        if kind == "listen":
            head, _, sub = items[0].partition(" · ")
            box(sl, 0.6, 2.0, 12.13, 2.0, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            text(sl, 0.9, 2.1, 11.5, 0.7, head, 36, "Oswald", True, WHITE, caps=True)
            text(sl, 0.9, 2.82, 11.5, 0.5, sub, 24, "Oswald", False, LIGHT)
            text(sl, 0.9, 3.38, 11.5, 0.5, items[1], 18, "Source Sans 3", False, WHITE)
            bullets(sl, items[2:], y=4.3, size=30, gap=1.6)
        elif kind == "bullets":
            n_ = len(items)
            size = 34 if n_ <= 4 else 31
            bullets(sl, items, size=size, y=2.2, gap=2.2 if n_ <= 4 else 1.95)
        elif kind == "closer":
            bullets(sl, items, w=6.4, size=30, y=2.1, gap=1.9)
            w = fit("hero", 5.4, 4.6)
            pic(sl, "hero", 12.73 - w, 2.15, w)
        elif kind == "days":
            for k, d in enumerate(DAYS):
                y = 1.98 + k * 0.83
                box(sl, 0.6, y, 1.5, 0.7, ACC if d["n"] in (1, 2) else BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                text(sl, 0.6, y, 1.5, 0.7, f"Day {d['n']}", 24, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, caps=True)
                part = "Foundation" if d["n"] <= 2 else ("Full Song" if d["n"] <= 4 else ("Submit" if d["n"] == 5 else "Present"))
                text(sl, 2.35, y - 0.04, 10.4, 0.42, d["title"] + "  ·  " + part, 21, "Oswald", True, BLACK, caps=True)
                text(sl, 2.35, y + 0.37, 10.4, 0.36, DECK_DAY_LINES[k], 16, "Source Sans 3", False, BLACK)
        elif kind == "types":
            for k, t in enumerate(TYPES):
                col, row = k % 2, k // 2
                x, y = 0.6 + col * 6.17, 1.98 + row * 2.3
                tile(sl, x, y, 5.96, 2.1, f"{k + 1}  ·  {t}", TYPE_DEF[t], home=(k == 0), hs=30, ss=20)
            text(sl, 0.6, 6.5, 12.1, 0.4, "You will use these in Part 2. Every section change gets one.", 18, "Source Sans 3", False, BLACK,
                 align=PP_ALIGN.CENTER)
        elif kind == "pic":
            pw_, ph_ = img[body][1]
            side_ = SIDE.get(body)
            if side_:   # a wide drawing across the slide, the points in two columns under it
                w = fit(body, 12.13, 3.25)
                yb = pic(sl, body, 0.6 + (12.13 - w) / 2, 1.98, w)
                for k, s in enumerate(side_):
                    col, row = k % 2, k // 2
                    x, yy = 0.6 + col * 6.17, max(yb + 0.3, 5.3) + row * 0.75
                    box(sl, x, yy + 0.12, 0.16, 0.16, ACC, shape=MSO_SHAPE.OVAL)
                    text(sl, x + 0.32, yy, 5.7, 0.6, s, 22, "Source Sans 3", False, BLACK)
            else:
                w = fit(body, 12.13, 4.85)
                pic(sl, body, 0.6 + (12.13 - w) / 2, 1.98, w)
        elif kind == "gain":
            w = fit("meter", 7.4, 4.85)
            pic(sl, "meter", 0.6, 2.0, w)
            bullets(sl, ["Never into the red", "Turn other things down", "The stem leads", "Drums support it", "Every element audible"],
                    x=8.4, y=2.15, w=4.35, size=24, gap=2.0)
        elif kind == "check":
            for k, it in enumerate(items):
                col, row = k % 2, k // 2
                x, y = 0.6 + col * 6.15, 2.15 + row * 1.5
                box(sl, x, y + 0.08, 0.42, 0.42, WHITE, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=2.25)
                text(sl, x + 0.6, y, 5.35, 1.4, it, 24, "Source Sans 3", False, BLACK)
        elif kind == "rubric":
            for k, ((c, _), short) in enumerate(zip(RUBRIC, RUBRIC_SHORT)):
                y = 1.98 + k * 0.9
                box(sl, 0.6, y, 4.2, 0.78, BLACK, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
                text(sl, 0.75, y, 3.9, 0.78, c, 19, "Oswald", True, WHITE, anchor=MSO_ANCHOR.MIDDLE, caps=True)
                text(sl, 5.0, y, 6.7, 0.78, "5 = " + short, 18, "Source Sans 3", False, BLACK, anchor=MSO_ANCHOR.MIDDLE)
                text(sl, 11.75, y, 0.98, 0.78, "/ 5", 24, "Oswald", True, ACC, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
            text(sl, 6.0, 6.5, 6.73, 0.45, f"Scored 2 to 5 · Total / {RUBRIC_TOTAL}", 22, "Oswald", True, BLACK, align=PP_ALIGN.RIGHT, caps=True)
        elif kind == "steps":
            for k, it in enumerate(items):
                y = 2.05 + k * 0.95
                box(sl, 0.6, y + 0.05, 0.62, 0.62, BLACK, shape=MSO_SHAPE.OVAL)
                text(sl, 0.6, y + 0.05, 0.62, 0.62, str(k + 1), 24, "Oswald", True, WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
                text(sl, 1.45, y - 0.05, 11.2, 0.9, it, 28, "Source Sans 3", False, BLACK, anchor=MSO_ANCHOR.MIDDLE)
        sl.notes_slide.notes_text_frame.text = notes
    prs.save(DECK)


def ttf_fonts():
    """woff2 -> ttf for LibreOffice (deck PDF), plus a fontconfig file that adds them."""
    from fontTools.ttLib import TTFont
    d = os.path.join(BUILD, "ttf")
    os.makedirs(d, exist_ok=True)
    styles = {"oswald-latin-500-normal": ("Oswald", "Regular", 400), "oswald-latin-700-normal": ("Oswald", "Bold", 700),
              "source-sans-3-latin-400-normal": ("Source Sans 3", "Regular", 400), "source-sans-3-latin-700-normal": ("Source Sans 3", "Bold", 700),
              "source-sans-3-latin-400-italic": ("Source Sans 3", "Italic", 400)}
    for stem, (fam, sub, wt) in styles.items():
        out = os.path.join(d, stem + ".ttf")
        if os.path.exists(out):
            continue
        t = TTFont(os.path.join(FONTS, stem + ".woff2"))
        t.flavor = None
        name = t["name"]
        ps = (fam.replace(" ", "") + "-" + sub)
        for rec in list(name.names):
            if rec.nameID in (16, 17, 21, 22):
                name.removeNames(nameID=rec.nameID)
        for nid, val in ((1, fam), (2, sub), (4, f"{fam} {sub}" if sub != "Regular" else fam), (6, ps)):
            name.setName(val, nid, 3, 1, 0x409)
            name.setName(val, nid, 1, 0, 0)
        t["OS/2"].usWeightClass = wt
        sel = 0
        if sub == "Bold":
            sel |= 1 << 5
        if sub == "Italic":
            sel |= 1
        if sub == "Regular":
            sel |= 1 << 6
        t["OS/2"].fsSelection = (t["OS/2"].fsSelection & ~0b1100001) | sel
        t["head"].macStyle = (1 if sub == "Bold" else 0) | (2 if sub == "Italic" else 0)
        t.save(out)
    conf = os.path.join(BUILD, "fonts.conf")
    with open(conf, "w") as fh:
        fh.write(f'<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig><include ignore_missing="yes">/etc/fonts/fonts.conf</include>'
                 f'<dir>{d}</dir><cachedir>{os.path.join(BUILD, "fccache")}</cachedir></fontconfig>')
    return conf


def deck_pdf(src=DECK, out=DECK_PDF, fonts=True):
    env = dict(os.environ, FONTCONFIG_FILE=ttf_fonts()) if fonts else dict(os.environ)
    prof = os.path.join(BUILD, "lo-profile" if fonts else "lo-profile-fallback")
    tmp = os.path.join(BUILD, "lo-out")
    os.makedirs(tmp, exist_ok=True)
    subprocess.run(["soffice", f"-env:UserInstallation=file://{prof}", "--headless", "--convert-to", "pdf", "--outdir", tmp, src],
                   check=True, env=env, capture_output=True)
    shutil.move(os.path.join(tmp, os.path.basename(src).replace(".pptx", ".pdf")), out)


# ================================================================ square images (2000 x 2000)
SQ_CSS = """
.sq { width: 1000px; height: 1000px; position: relative; overflow: hidden; background: linear-gradient(180deg,#fff 0%,#fff 60%,#eeeeee 100%); }
.sq .store { position: absolute; left: 0; right: 0; top: 30px; text-align: center; font-family: 'Oswald'; font-weight: 500; letter-spacing: .32em;
             text-transform: uppercase; font-size: 22px; }
.sq .fillw { position: absolute; left: 0; right: 0; top: 66px; bottom: 70px; display: flex; flex-direction: column; align-items: center; justify-content: space-evenly; }
.sq .hd { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; text-align: center; line-height: .95; white-space: nowrap; }
.sq .sub { font-family: 'Oswald'; font-weight: 300; text-align: center; white-space: nowrap; }
.sq .band { position: absolute; left: 0; right: 0; bottom: 0; background: #000; color: #fff; text-align: center; font-family: 'Oswald';
            font-weight: 500; text-transform: uppercase; letter-spacing: .1em; font-size: 26px; padding: 16px 0 18px; white-space: nowrap; }
.sq .band em { font-style: normal; color: ACCENT; }
.shot { background: #fff; border: 2px solid #000; box-shadow: 8px 8px 0 #000; position: relative; overflow: hidden; }
.shot img { display: block; width: 100%; }
.lbl { font-family: 'Oswald'; font-weight: 700; text-transform: uppercase; font-size: 22px; text-align: center; margin-top: 14px; }
.lbl span { font-weight: 300; text-transform: none; display: block; font-size: 18px; }
.swm { position: absolute; left: 50%; top: 50%; transform: translate(-50%,-50%) rotate(-38deg); font-family: 'Oswald'; font-weight: 700;
       color: rgba(0,0,0,.14); white-space: nowrap; letter-spacing: .06em; }
""".replace("ACCENT", ACCENT)
FONT_PRIME = ('<div style="position:absolute;left:-9999px;font-family:\'Source Sans 3\'">'
              '<span style="font-weight:400">a</span><span style="font-weight:600">a</span><i style="font-weight:400">a</i></div>')


def sq_doc(body):
    return (f'<!doctype html><html><head><meta charset="utf-8"><style>{font_css()}'
            f'html,body{{margin:0;background:#fff}}{SQ_CSS}</style></head><body>{FONT_PRIME}{body}</body></html>')


def sq(inner, band):
    return sq_doc(f'<div class="sq"><div class="store">{BRAND}</div><div class="fillw">{inner}</div><div class="band">{band}</div></div>')


SQ_SUB = f"{DAYS_N}-day unit · drums over stems · fills, cuts, builds · intro slides"
SQ_BAND = f"Grades 9–12 · {DAYS_N} days · {len(PLAN)} pages + {DECK_SLIDES} slides"


def cover_square():
    s = SQ_COVER
    return sq(f'<div style="width:900px">{body_svg(hero_body())}</div>'
              f'<div><div class="hd" style="font-size:{s["a"]}px">{COVER_A}</div>'
              + "".join(f'<div class="hd" style="font-size:{s["b"]}px;margin-top:4px">{escape(x)}</div>' for x in SQ_B_LINES) +
              f'<div class="sub" style="font-size:{s["c"]}px;margin-top:10px">{SQ_SUB}</div></div>',
              SQ_BAND.replace(" · ", " <em>·</em> "))


SQ_HEADS = {1: ("Ready to teach Monday", 76, "Teacher notes, six day plans, student pages and the /25 rubric in one PDF", 28),
            2: ("Intro deck included", 76, f"{DECK_SLIDES} editable slides (.pptx + PDF) with speaker notes on every slide", 28),
            3: ("Everything inside", 72, "Post, transitions, gain staging, map, project sheet, examples", 28)}
SQ_BANDS = {1: "Listen <em>·</em> A + B <em>·</em> transitions <em>·</em> gain staging", 2: "Hook <em>·</em> listen <em>·</em> worked example "
            "<em>·</em> steps <em>·</em> rubric", 3: f"{len(PLAN)} pages <em>·</em> US Letter PDF <em>·</em> stems: make your own"}


def sq_head(i):
    h, hs, s, ss = SQ_HEADS[i]
    return f'<div><div class="hd" style="font-size:{hs}px">{h}</div><div class="sub" style="font-size:{ss}px;margin-top:8px">{s}</div></div>'


def preview_square_1(img):
    rot = [-4, 0, 4]
    shots = "".join(f'<div style="position:absolute;left:{x}px;top:{y}px;width:372px;transform:rotate({r}deg);z-index:{z}">'
                    f'<div class="shot"><img src="{s}"><div class="swm" style="font-size:80px">PREVIEW</div></div></div>'
                    for (s, _, _), x, y, r, z in zip(img, (40, 314, 588), (24, 0, 24), rot, (1, 3, 2)))
    labels = "".join(f'<div class="lbl" style="width:300px">{a}<span>{b}</span></div>' for _, a, b in img)
    return sq(f'{sq_head(1)}<div style="position:relative;width:1000px;height:520px">{shots}</div>'
              f'<div style="display:flex;gap:30px">{labels}</div>', SQ_BANDS[1])


def preview_square_2(slides):
    cells = "".join(f'<div class="shot" style="box-shadow:6px 6px 0 #000"><img src="{s}"><div class="swm" style="font-size:52px">PREVIEW</div></div>'
                    for s in slides)
    return sq(f'{sq_head(2)}<div style="display:grid;grid-template-columns:388px 388px;gap:20px 34px">{cells}</div>', SQ_BANDS[2])


def preview_square_3(imgs):
    cells = "".join(f'<div style="width:214px"><div class="shot" style="box-shadow:5px 5px 0 #000"><img src="{s}">'
                    f'<div class="swm" style="font-size:46px">PREVIEW</div></div><div class="lbl" style="font-size:18px;margin-top:9px">{lab_}</div></div>'
                    for s, lab_ in imgs)
    return sq(f'{sq_head(3)}<div style="display:grid;grid-template-columns:repeat(4,214px);gap:22px 24px">{cells}</div>', SQ_BANDS[3])


SHOTS1 = [("pacing", "Pacing + plans", "6 days, minute by minute"), ("transitions", "Four transitions", "fill, breakdown, cut, build"),
          ("key3", "Worked examples", "patterns and a mapped song")]
SHOTS2 = [1, 6, 8, 12, 13, 14]   # deck slides
SHOTS3 = [("post1", "Day 1 post"), ("gain", "Gain staging"), ("map", "Section map"), ("proj1", "Project sheet"),
          ("proj3", "Steps"), ("rubric", "Rubric / 25"), ("transcheck", "Quick check"), ("makestem", "Make a stem")]


# ================================================================ render
def node_env():
    env = dict(os.environ)
    if not env.get("NODE_PATH"):
        env["NODE_PATH"] = subprocess.run(["npm", "root", "-g"], capture_output=True, text=True).stdout.strip()
    return env


def render(jobs):
    path = os.path.join(BUILD, "jobs.json")
    with open(path, "w") as f:
        json.dump(jobs, f)
    subprocess.run(["node", RENDER_JS, path], check=True, env=node_env())


OVERFLOW_JS = r"""
const { chromium } = require('playwright');
const url = require('url');
(async () => {
  const browser = await chromium.launch();
  const page = await browser.newPage();
  await page.goto(url.pathToFileURL(process.argv[2]).href, { waitUntil: 'load' });
  await page.evaluate(() => document.fonts.ready);
  await page.emulateMedia({ media: 'print' });
  const res = await page.evaluate(() => [...document.querySelectorAll('section.page')].map((s, i) => {
    const m = s.querySelector('.main');
    const out = { page: i + 1, key: s.dataset.key, main_over: m ? Math.max(0, m.scrollHeight - m.clientHeight) : 0, cells: [] };
    if (m) {
      const mb = m.getBoundingClientRect();
      m.querySelectorAll('*').forEach(el => {
        if (el.closest('svg') && el.tagName.toLowerCase() !== 'svg') return;
        const r = el.getBoundingClientRect();
        if (r.width && (r.right > mb.right + 1 || r.bottom > mb.bottom + 1 || r.left < mb.left - 1))
          out.cells.push((el.tagName + ': ' + (el.textContent || '').trim()).slice(0, 50));
      });
      let maxb = mb.top;
      m.querySelectorAll('*').forEach(el => { const r = el.getBoundingClientRect(); if (r.height && !el.classList.contains('lines')) maxb = Math.max(maxb, r.bottom); });
      out.free = Math.round(mb.bottom - maxb);
      if (m.querySelector('.lines')) out.free = 0;
    }
    return out;
  }));
  require('fs').writeFileSync(process.argv[3], JSON.stringify(res, null, 1));
  await browser.close();
})().catch(e => { console.error(e); process.exit(1); });
"""


def overflow_scan(html):
    js = os.path.join(BUILD, "overflow.js")
    with open(js, "w") as f:
        f.write(OVERFLOW_JS)
    subprocess.run(["node", js, html, OVERFLOW_JSON], check=True, env=node_env())


def write(name, txt):
    p = os.path.join(BUILD, name)
    with open(p, "w") as f:
        f.write(txt)
    return p


def pdf_png(src, page_no, out, res=110):
    subprocess.run(["pdftoppm", "-png", "-r", str(res), "-f", str(page_no), "-l", str(page_no), "-singlefile", src,
                    os.path.join(BUILD, out)], check=True)


def art_doc(w, h, inner):
    return (f'<!doctype html><html><head><meta charset="utf-8"><style>{font_css()}{CSS}'
            f'html,body{{margin:0;background:#fff}}</style></head>'
            f'<body><div style="width:{w}px;height:{h}px;overflow:hidden;background:#fff">{inner}</div></body></html>')


def deck_art():
    """PNG art for the deck (original drawings rendered with the brand fonts)."""
    svgs = {"hero": (1900, body_svg(hero_body())), "patterns": (2000, body_svg(patterns_body())), "trans": (2000, body_svg(trans_body())),
            "songmap": (2000, body_svg(songmap_body(h=190))), "meter": (1300, body_svg(meter_body()))}
    arts = {}
    for k, (w, svg) in svgs.items():   # exact height from the drawing's own viewBox, so nothing is cropped and nothing is padded
        vw, vh = map(float, re.search(r'viewBox="0 0 ([\d.]+) ([\d.]+)"', svg).groups())
        arts[k] = (w, int(round(w * vh / vw)) + 2, f'<div style="width:{w}px">{svg}</div>')
    jobs, out = [], {}
    for k, (w, h, inner) in arts.items():
        p = os.path.join(BUILD, f"deck-{k}.png")
        jobs.append({"html": write(f"deck-{k}.html", art_doc(w, h, inner)), "png": p, "width": w, "height": h, "scale": 1})
        out[k] = (p, (w, h))
    return jobs, out


def tpt_json():
    return {
        "price": PRICE, "license_price": PRICE,
        "tax_code": "Other Digital Goods - No Physical Media",
        "grades": ["9th-grade", "10th-grade", "11th-grade", "12th-grade"],
        "subjects": ["Music"],
        "tags": ["Autumn"],
        "formats": ["PDF"],
        "pages": len(PLAN) + DECK_SLIDES,
        "answer_key": "Included",
        "zip_contents": [os.path.basename(PDF), os.path.basename(DECK), os.path.basename(DECK_PDF)],
        "files": {"product": os.path.basename(ZIP), "preview": os.path.basename(PREVIEW_PDF), "thumb1": "cover.png",
                  "thumb2": "preview-1.png", "thumb3": "preview-2.png", "thumb4": "preview-3.png"},
    }


def make_zip():
    if os.path.exists(ZIP):
        os.remove(ZIP)
    with zipfile.ZipFile(ZIP, "w", zipfile.ZIP_DEFLATED) as z:
        for p in (PDF, DECK, DECK_PDF):
            zi = zipfile.ZipInfo(f"{ZIP_DIR}/{os.path.basename(p)}", date_time=(2026, 10, 14, 12, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            with open(p, "rb") as fh:
                z.writestr(zi, fh.read())


def build():
    os.makedirs(BUILD, exist_ok=True)
    problems = verify_music() + verify_data() + verify_labels()
    if problems:
        print("\n".join(problems))
        sys.exit("data check failed; not rendering")
    full = write("resource.html", html_doc([render_page(p) for p in PLAN], TITLE))
    prev = write("preview.html", html_doc([render_page(p, watermark=True) for p in PREVIEW_PICKS], "Preview"))
    art_jobs, art = deck_art()
    render([{"html": full, "pdf": PDF}, {"html": prev, "pdf": PREVIEW_PDF}] + art_jobs)
    overflow_scan(full)
    build_deck(art)
    deck_pdf()
    make_zip()
    res = 45 if LOW_RES else 120
    for key, _, _ in SHOTS1:
        pdf_png(PDF, IX[key], key, res)
    for key, _ in SHOTS3:
        pdf_png(PDF, IX[key], key, 40 if LOW_RES else 80)
    for s in SHOTS2:
        pdf_png(DECK_PDF, s, f"slide{s}", 30 if LOW_RES else 70)
    sqs = [write("cover.html", cover_square()), write("p1.html", preview_square_1([(f"{k}.png", a, b) for k, a, b in SHOTS1])),
           write("p2.html", preview_square_2([f"slide{s}.png" for s in SHOTS2])),
           write("p3.html", preview_square_3([(f"{k}.png", lab_) for k, lab_ in SHOTS3]))]
    render([{"html": h, "png": o, "width": 1000, "height": 1000, "scale": 1 if LOW_RES else 2} for h, o in zip(sqs, [COVER_PNG] + PREVIEW_PNGS)])
    with open(TPT_JSON, "w") as f:
        json.dump(tpt_json(), f, indent=2, ensure_ascii=False)
        f.write("\n")
    print(f"built {PDF} ({len(PLAN)} pages), {DECK} + PDF ({DECK_SLIDES} slides), {ZIP}, preview, cover + 3 previews, tpt.json"
          + (" [LOW RES]" if LOW_RES else ""))


# ================================================================ checks
def verify_music():
    """Drum patterns, transitions, quiz answers, section maps, stem chords and durations, computed."""
    errs = []
    allp = [("A", PATTERN_A), ("B", PATTERN_B)] + list(TRANS.items()) + [(f"quiz {i + 1}", q) for i, q in enumerate(QUIZ)]
    for nm, p in allp:
        for ln, s in p.items():
            if ln not in LANES or len(s) != STEPS or set(s) - {"x", "."}:
                errs.append(f"pattern {nm}: lane {ln} must be 16 steps of x/. ({s})")
    # Foundation patterns: at least 3 layers (kick, snare, hi-hat at minimum); snare on 2 and 4 in both; clear contrast
    for nm, p in (("A", PATTERN_A), ("B", PATTERN_B)):
        layers = lanes_of(p)
        if len(layers) < 3 or "Kick" not in layers or "Snare" not in layers or not ({"Closed hi-hat", "Open hi-hat"} & set(layers)):
            errs.append(f"pattern {nm}: needs kick, snare and hi-hat (3 layers minimum), has {layers}")
        if p["Snare"][4] != "x" or p["Snare"][12] != "x":
            errs.append(f"pattern {nm}: snare on beats 2 and 4")
        if p["Kick"][0] != "x":
            errs.append(f"pattern {nm}: kick on beat 1")
    if contrast(PATTERN_A, PATTERN_B) < 8 or lanes_of(PATTERN_A) == lanes_of(PATTERN_B):
        errs.append("A and B must contrast clearly (8+ cells different and a different set of sounds)")
    # the four transitions match the owner's definitions (classify), each is the measure before a change
    for t, m in TRANS.items():
        got = classify(m)
        if got != t:
            errs.append(f"transition '{t}' classifies as {got}")
        if same(m, PATTERN_A):
            errs.append(f"transition '{t}' is the same as the groove")
    if trailing_silence(TRANS["Cut"]) not in (4, 8):
        errs.append("cut: the silence must be exactly 1 or 2 beats")
    if lanes_of(TRANS["Breakdown"]) not in (["Kick"], ["Closed hi-hat"]):
        errs.append("breakdown: kick only or hi-hat only (the owner's examples)")
    fill_ = TRANS["Drum fill"]
    if not any(fill_.get(ln, "." * STEPS)[8:] != PATTERN_A.get(ln, "." * STEPS)[8:] for ln in LANES) or \
            not ({"High tom", "Low tom"} & set(lanes_of(fill_))):
        errs.append("drum fill: a rhythmic variation (toms) in the transition measure")
    # quiz: every type once, answers computed; quiz grids differ from the reference grids
    qa = [classify(q) for q in QUIZ]
    if sorted(qa) != sorted(TYPES):
        errs.append(f"name-that-transition answers must be the four types once each, got {qa}")
    if qa == TYPES:
        errs.append("name-that-transition is in the same order as the reference page")
    for q in QUIZ:
        if any(same(q, m) for m in TRANS.values()):
            errs.append("a quiz grid repeats a reference transition")
    # section maps
    if [n for _, n in FOUNDATION] != [4, 4, 4, 4] or "".join(s for s, _ in FOUNDATION) != "AABA":
        errs.append("Foundation stem: 16 measures, 4 sections of 4, AABA")
    if sum(n for _, n in FULL_SONG) != 56 or len(SONG_TRANS) != len(FULL_SONG) - 1:
        errs.append("Full Song: 56 measures and one transition per section change")
    if len({t for t, _ in SONG_TRANS}) < 2 or any(t not in TYPES for t, _ in SONG_TRANS):
        errs.append("Full Song: at least two transition types, all from the four")
    for k, (t, why) in enumerate(SONG_TRANS):
        a, b = FULL_SONG[k][0], FULL_SONG[k + 1][0]
        if t == "Build" and energy(b) <= energy(a):
            errs.append(f"Full Song change {k + 1}: a build should lead into more energy ({a} -> {b})")
        if t in ("Cut", "Breakdown") and energy(b) >= energy(a) and t == "Cut":
            errs.append(f"Full Song change {k + 1}: the cut leads into a quieter section in this example ({a} -> {b})")
        if t == "Breakdown" and energy(b) >= energy(a):
            errs.append(f"Full Song change {k + 1}: the breakdown leads into less energy in this example ({a} -> {b})")
    if energy("B") <= energy("A") or energy("Chorus") <= energy("Verse") or energy("Bridge") >= energy("Chorus"):
        errs.append("stem recipes: B/chorus must carry more instruments than A/verse; the bridge fewer than the chorus")
    # stem chords: diatonic in C major, case matches quality, one per measure, sections are whole repeats of the progression
    for sec, prog in PROGS.items():
        for rn in prog:
            notes, q = triad(rn)
            if (q == "major") != rn.isupper() or q == "diminished" or any(x not in C_MAJOR for x in notes):
                errs.append(f"stem chords {sec}: {rn} = {notes} {q} (case must match quality; no diminished chords)")
    for sec, nm in FOUNDATION + FULL_SONG:
        if nm % len(PROGS[sec]):
            errs.append(f"{sec}: {nm} measures is not a whole number of {len(PROGS[sec])}-chord progressions")
    want = {"I": "C E G", "ii": "D F A", "iii": "E G B", "IV": "F A C", "V": "G B D", "vi": "A C E"}
    for rn, w in want.items():
        if " ".join(triad(rn)[0]) != w:
            errs.append(f"chord spelling {rn}: {triad(rn)[0]}")
    if abs(seconds(16) - 40) > 1e-9 or mmss(seconds(56)) != "2:20":
        errs.append("durations at 96 BPM: 16 measures = 0:40, 56 measures = 2:20")
    for k, (t, why) in enumerate(SONG_TRANS):
        if t == "Build" and FULL_SONG[k + 1][0] == "Chorus" and "build before" not in why and k != 0:
            errs.append("the bridge-to-chorus build should quote the owner's 'a build before a big chorus'")
    return errs


def verify_data():
    errs = []
    counts = dict(goals=(GOALS, 5), vocab=(VOCAB, 10), req1=(REQ1, 9), req2=(REQ2, 8), steps1=(STEPS1, 4), steps2=(STEPS2, 5),
                  step2types=(STEP2_TYPES, 4), tips=(TIPS, 6), check=(CHECK, 12), rubric=(RUBRIC, 5), struggles=(STRUGGLES, 5),
                  post=(POST_STEPS, 5), what=(WHAT, 3), short=(RUBRIC_SHORT, 5))
    for nm, (lst, nexp) in counts.items():
        if len(lst) != nexp:
            errs.append(f"{nm}: {len(lst)} items, expected {nexp}")
    if any(len(dd) != len(LEVELS) for _, dd in RUBRIC) or 5 * len(RUBRIC) != RUBRIC_TOTAL or RUBRIC_TOTAL != 25:
        errs.append("rubric must be the owner's 5 categories x levels 2-5 = /25")
    if [c for c, _ in RUBRIC] != ["Foundations", "Full Song", "Technical Skill / DAW Use", "Gain Staging & Mix", "Documentation & Presentation"]:
        errs.append("rubric categories must be the owner's, in order")
    # no LMS name or new requirement: requirement and checklist items keep the owner's wording (only the LMS name changes)
    for lst in (REQ1, REQ2, [c for _, c in CHECK]):
        for it in lst:
            if "Google" in it:
                errs.append("an LMS name was left in a list")
    if sum(NAME_CHECK in c for _, c in CHECK) != 2 or NAME1 not in REQ1[7] or NAME2 not in REQ2[6]:
        errs.append("owner's project names must stay as written (requirements Foundation / Full Song; checklist Foundations twice)")
    if [t for t, _ in STEP2_TYPES] != TYPES or list(TYPE_DEF) != TYPES:
        errs.append("transition types must be the owner's four, in order")
    if "A, B, A, B, etc." not in POST_STEPS[2][1] or "16 measures" not in POST_STEPS[1][1]:
        errs.append("Day 1 post wording changed (keep 'A, B, A, B, etc.' and '16 measures')")
    # pacing: 6 days, blocks 0-53, every deck slide in exactly one block of the day the slide says
    if len(DAYS) != DAYS_N or [d["n"] for d in DAYS] != list(range(1, DAYS_N + 1)) or len(DECK_DAY_LINES) != DAYS_N:
        errs.append(f"pacing must be {DAYS_N} days (and {DAYS_N} on the deck)")
    seen = {}
    for d in DAYS:
        blocks = [tuple(map(int, x.split("–"))) for x, _, _, _ in d["blocks"]]
        if blocks[0][0] != 0 or blocks[-1][1] != MINUTES or any(blocks[i][1] != blocks[i + 1][0] for i in range(len(blocks) - 1)):
            errs.append(f"Day {d['n']} minute blocks do not run 0-{MINUTES} without gaps")
        if not day_slides(d):
            errs.append(f"Day {d['n']} has no deck slide")
        for _, h, txt, sl in d["blocks"]:
            nameds = set()
            for a_, b_ in re.findall(r"[Ss]lides? (\d+)(?:–(\d+))?", txt):
                nameds |= set(range(int(a_), int(b_ or a_) + 1))
            if not sl:
                if nameds:
                    errs.append(f"Day {d['n']} '{h}' names a slide but has no slide slot")
                continue
            if nameds != set(range(sl[0], sl[1] + 1)):
                errs.append(f"Day {d['n']} '{h}': text names slides {sorted(nameds)} but the slot is {sl}")
            for s in range(sl[0], sl[1] + 1):
                if s in seen:
                    errs.append(f"slide {s} is slotted twice (Day {seen[s]} and Day {d['n']})")
                seen[s] = d["n"]
                if not 1 <= s <= DECK_SLIDES:
                    errs.append(f"Day {d['n']} names slide {s}, which does not exist")
                elif slide_day(s) != d["n"]:
                    errs.append(f"slide {s} says Day {slide_day(s)} but the Day {d['n']} plan uses it")
    for i in range(1, DECK_SLIDES + 1):
        if i not in seen:
            errs.append(f"deck slide {i} has no day slot in the day plans")
    if [s[0] for s in SLIDES] != sorted(s[0] for s in SLIDES):
        errs.append("deck slides must run in day order")
    if len(SLIDES) < 12 or any(len(s[4].strip()) < 80 for s in SLIDES):
        errs.append("deck needs 12+ slides, each with real speaker notes")
    kinds = [k for _, k, *_ in SLIDES]
    for need in ("title", "bullets", "listen", "pic", "steps", "check", "rubric", "closer"):
        if need not in kinds:
            errs.append(f"deck pattern: no '{need}' slide")
    if sorted([b for _, k, _, b, _ in SLIDES if k == "pic"] + ["hero", "meter"]) != sorted(DECK_PICS):
        errs.append("deck pictures must be hero + patterns + trans + songmap + meter")
    # the Day 1 post, the deck's Day 1 demo and the day plan agree: listen twice before Drum Machine Designer; 16 measures; A pattern
    if "twice" not in POST_STEPS[1][0] or "16 measures" not in SLIDES[6][3] or "Drum Machine Designer" not in SLIDES[6][3]:
        errs.append("Day 1: post, deck slide 7 and plan must agree (listen twice, 16 measures, then Drum Machine Designer)")
    if "A pattern" not in DAYS[0]["goal"] or "A pattern" not in SLIDES[8][2]:
        errs.append("Day 1 goal and demo: the A pattern")
    if len(PLAN) != len(set(PLAN)):
        errs.append("plan has a repeated page")
    if len(LISTING_TITLE) > 80 or "Logic Pro" not in LISTING_TITLE[:30]:
        errs.append(f"listing title is {len(LISTING_TITLE)} chars, or 'Logic Pro' is not in the first 30")
    for label, txt, sz in (("cover A", COVER_A, PDF_COVER["a"]), ("cover B", COVER_B, PDF_COVER["b"])):
        if text_w(txt.upper(), sz) * 1.01 > PDF_COVER["room"]:
            errs.append(f"PDF {label} '{txt}' is {text_w(txt.upper(), sz):.0f}pt wide, room {PDF_COVER['room']:.0f}")
    if text_w(COVER_C, PDF_COVER["c"]) > PDF_COVER["room"]:
        errs.append("PDF cover subtitle too wide")
    if text_w(COVER_BAND.upper(), PDF_COVER["band"]) * 1.12 > PDF_COVER["room"]:
        errs.append("PDF cover band too wide")
    if " ".join(SQ_B_LINES) != COVER_B:
        errs.append("square cover lines must spell the title")
    for txt, sz in [(COVER_A.upper(), SQ_COVER["a"]), (SQ_SUB, SQ_COVER["c"])] + [(x.upper(), SQ_COVER["b"]) for x in SQ_B_LINES]:
        if text_w(txt, sz) * 1.01 > SQ_COVER["room"]:
            errs.append(f"square cover line '{txt}' is {text_w(txt, sz):.0f}px wide, room {SQ_COVER['room']}")
    if text_w(SQ_BAND.upper(), 26) * 1.12 > 960:
        errs.append("square cover band too wide")
    if SQ_COVER["b"] * 0.75 < 100:
        errs.append("square cover big line under 10% of the image height")
    for i, (h, hs, s, ss) in SQ_HEADS.items():
        if text_w(h.upper(), hs) * 1.01 > 960 or text_w(s, ss) > 960:
            errs.append(f"preview {i} headline/subtitle too wide")
    for s in TN1_DIFF + [STEM_NOTE]:
        if "Blind" in s:
            errs.append("school name in the disclosure")
    for need in ("Check the example", "your class page", "minute 48", "parenthesis", "reverb", "not included"):
        if need not in " ".join(TN1_DIFF):
            errs.append(f"teacher notes page 1 disclosure: missing '{need}'")
    return errs


def pdf_fonts(path):
    out = subprocess.run(["pdffonts", path], capture_output=True, text=True, check=True).stdout.splitlines()[2:]
    return [(ln.split()[0], ln.split()[-5]) for ln in out]


def pdf_text(path):
    return subprocess.run(["pdftotext", "-layout", path, "-"], capture_output=True, text=True, check=True).stdout


def deck_texts():
    from pptx import Presentation
    prs = Presentation(DECK)
    slides, notes = [], []
    for s in prs.slides:
        slides.append(" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame))
        notes.append(s.notes_slide.notes_text_frame.text if s.has_notes_slide else "")
    media = [p for p in zipfile.ZipFile(DECK).namelist() if p.startswith("ppt/media/")]
    return prs, slides, notes, media


def deck_overflow(prs):
    """Every shape stays inside the slide; Oswald text fits its box (lines available x width); no two text boxes overlap;
    free text keeps 0.1 in clear of filled boxes; no text over a picture; pictures stay above the footer rule."""
    errs = []
    W, H = prs.slide_width, prs.slide_height
    for i, s in enumerate(prs.slides, start=1):
        boxes = []
        for sh in s.shapes:
            if sh.left < 0 or sh.top < 0 or sh.left + sh.width > W + 10 or sh.top + sh.height > H + 10:
                errs.append(f"deck slide {i}: a shape runs off the slide")
            if sh.has_text_frame and sh.text_frame.text:
                r = sh.text_frame.paragraphs[0].runs
                if r and r[0].font.name == "Oswald" and r[0].font.size:
                    widest = max((text_w(p.text, r[0].font.size.pt) for p in sh.text_frame.paragraphs), default=0)
                    room = sh.width / 12700 - 6
                    lines_avail = int((sh.height / 12700) // (r[0].font.size.pt * 1.15)) or 1
                    if widest > room * lines_avail:
                        errs.append(f"deck slide {i}: '{sh.text_frame.text[:30]}' ({widest:.0f}pt) does not fit {room:.0f}pt x {lines_avail} line(s)")
                if r and r[0].font.name == "Source Sans 3" and r[0].font.size and sh.shape_type == 17:
                    sz = r[0].font.size.pt
                    room = sh.width / 12700 - 6
                    need = sum(max(1, math.ceil(text_w(p.text, sz, "ss") / room)) for p in sh.text_frame.paragraphs)
                    if need * sz * 1.2 > sh.height / 12700 + sz * 0.45:
                        errs.append(f"deck slide {i}: '{sh.text_frame.text[:30]}' wraps to {need} line(s), box holds less (overlap risk)")
                if sh.shape_type == 17:
                    boxes.append((sh.left, sh.top, sh.left + sh.width, sh.top + sh.height, sh.text_frame.text[:20]))
        pics = [(sh.left, sh.top, sh.left + sh.width, sh.top + sh.height) for sh in s.shapes if sh.shape_type == 13]
        for P in pics:
            if P[3] > 914400 * 6.95:
                errs.append(f"deck slide {i}: a picture runs into the footer")
        pad = 914400 // 10
        fills = [(sh.left, sh.top, sh.left + sh.width, sh.top + sh.height) for sh in s.shapes
                 if sh.shape_type == 1 and sh.width > 914400 // 2 and sh.height > 914400 // 2]

        def inside(A, F):
            return A[0] >= F[0] - 12700 and A[2] <= F[2] + 12700 and A[1] >= F[1] - 12700 and A[3] <= F[3] + 12700
        for A in boxes:
            if any(inside(A, F) for F in fills):
                continue
            for F in fills:
                ox = min(A[2], F[2]) - max(A[0], F[0])
                near = min(A[3], F[3] + pad) - max(A[1], F[1] - pad) > 0
                if ox > 12700 and near:
                    errs.append(f"deck slide {i}: text '{A[4]}' touches a box (needs 0.1 in clearance)")
        for P in pics:
            for A in boxes:
                if min(A[2], P[2]) - max(A[0], P[0]) > 12700 * 4 and min(A[3], P[3]) - max(A[1], P[1]) > 12700 * 4:
                    errs.append(f"deck slide {i}: text '{A[4]}' overlaps a picture")
        for a in range(len(boxes)):
            for b in range(a + 1, len(boxes)):
                A, B = boxes[a], boxes[b]
                ox = min(A[2], B[2]) - max(A[0], B[0])
                oy = min(A[3], B[3]) - max(A[1], B[1])
                if ox > 12700 * 4 and oy > 12700 * 4:
                    inside_ = (A[0] >= B[0] and A[2] <= B[2] and A[1] >= B[1] and A[3] <= B[3]) or \
                              (B[0] >= A[0] and B[2] <= A[2] and B[1] >= A[1] and B[3] <= A[3])
                    if not inside_:
                        errs.append(f"deck slide {i}: text boxes '{A[4]}' and '{B[4]}' overlap")
    return errs


def verify_files():
    errs = []
    from pypdf import PdfReader
    from PIL import Image
    for p in [PDF, DECK, DECK_PDF, ZIP, PREVIEW_PDF, COVER_PNG, *PREVIEW_PNGS, UPLOAD, TPT_JSON, OVERFLOW_JSON]:
        if not os.path.exists(p):
            errs.append(f"missing {os.path.basename(p)}")
    if errs:
        return errs
    for path, n_exp in ((PDF, len(PLAN)), (PREVIEW_PDF, len(PREVIEW_PICKS))):
        rd = PdfReader(path)
        if len(rd.pages) != n_exp:
            errs.append(f"{os.path.basename(path)}: {len(rd.pages)} pages, expected {n_exp}")
        for i, pg_ in enumerate(rd.pages):
            w, h = float(pg_.mediabox.width), float(pg_.mediabox.height)
            if abs(w - 612) > 1 or abs(h - 792) > 1:
                errs.append(f"{os.path.basename(path)} p{i + 1}: {w:.0f}x{h:.0f}pt, not US Letter")
    if len(PdfReader(PREVIEW_PDF).pages) != 4:
        errs.append("preview PDF must be 4 pages")
    for o in json.load(open(OVERFLOW_JSON)):
        if o["main_over"] > 1:
            errs.append(f"page {o['page']} ({o['key']}): content overflows by {o['main_over']}px (clipped)")
        for c in o["cells"]:
            errs.append(f"page {o['page']} ({o['key']}): text wider than its box: '{c}'")
        if o["key"] not in ("cover", "terms") and o.get("free", 0) > 200:
            errs.append(f"page {o['page']} ({o['key']}): {o['free']}px empty at the bottom (dead space)")
    drd = PdfReader(DECK_PDF)
    if len(drd.pages) != DECK_SLIDES:
        errs.append(f"deck PDF: {len(drd.pages)} pages, expected {DECK_SLIDES}")
    for i, pg_ in enumerate(drd.pages):
        if abs(float(pg_.mediabox.width) / float(pg_.mediabox.height) - 16 / 9) > 0.01:
            errs.append(f"deck PDF p{i + 1}: not 16:9")
    for path in (PDF, PREVIEW_PDF, DECK_PDF):
        fonts = pdf_fonts(path)
        for name, emb in fonts:
            if emb != "yes":
                errs.append(f"{os.path.basename(path)}: font {name} not embedded")
            if not re.search(r"oswald|sourcesans", name.lower().replace(" ", "")):
                errs.append(f"{os.path.basename(path)}: unexpected font {name} (fallback?)")
        names = " ".join(n for n, _ in fonts).lower()
        for fam in ("oswald", "sourcesans3"):
            if fam not in names:
                errs.append(f"{os.path.basename(path)}: {fam} missing")
    prs, stext, notes, media = deck_texts()
    if len(prs.slides) != DECK_SLIDES:
        errs.append(f"deck has {len(prs.slides)} slides, expected {DECK_SLIDES}")
    for i, nt in enumerate(notes):
        if len(nt.strip()) < 80:
            errs.append(f"deck slide {i + 1}: speaker notes missing or too short")
        if not re.search(rf"\bDay {slide_day(i + 1)}\b", stext[i], re.I):
            errs.append(f"deck slide {i + 1}: no 'Day {slide_day(i + 1)}' slot shown on the slide")
    errs += deck_overflow(prs)
    if any(not m.lower().endswith(".png") for m in media) or len(media) != len(DECK_PICS):
        errs.append(f"deck media should be exactly the {len(DECK_PICS)} drawn PNGs, found {media}")
    if abs(prs.slide_width / prs.slide_height - 16 / 9) > 0.01:
        errs.append("deck is not 16:9")
    alldeck = " ".join(stext + notes)
    if ITAT[0].lower() not in alldeck.lower() or ITAT[1].lower() not in alldeck.lower():
        errs.append(f"deck must name {ITAT[0]} by {ITAT[1]}")
    for i, (_, k, _, body, _) in enumerate(SLIDES):
        if k == "listen" and ITAT[0] in body and "YouTube or Spotify" not in stext[i]:
            errs.append(f"deck slide {i + 1}: a song slide must say YouTube or Spotify")
    if "YouTube or Spotify" not in notes[0] or "not included" not in notes[0]:
        errs.append("deck slide 1 notes must say YouTube or Spotify and that stems are not included")
    names = zipfile.ZipFile(ZIP).namelist()
    for p in (PDF, DECK, DECK_PDF):
        n = f"{ZIP_DIR}/{os.path.basename(p)}"
        if n not in names:
            errs.append(f"zip is missing {os.path.basename(p)}")
        elif zipfile.ZipFile(ZIP).read(n) != open(p, "rb").read():
            errs.append(f"zip copy of {os.path.basename(p)} is stale; rebuild")
    if len(names) != 3:
        errs.append(f"zip should hold exactly 3 files, has {names}")
    if any(re.search(r"\.(wav|mp3|aif|aiff|m4a|logicx|mov|mp4)$", nm, re.I) for nm in names):
        errs.append("zip holds audio, video or a Logic session")
    rd = PdfReader(PDF)
    pages = [subprocess.run(["pdftotext", "-layout", "-f", str(i + 1), "-l", str(i + 1), PDF, "-"], capture_output=True, text=True,
                            check=True).stdout for i in range(len(rd.pages))]
    flat_ = [re.sub(r"\s+", " ", t) for t in pages]
    raw0 = subprocess.run(["pdftotext", "-f", "1", "-l", "1", PDF, "-"], capture_output=True, text=True, check=True).stdout
    lines0 = [re.sub(r"\s+", " ", ln.strip()) for ln in raw0.splitlines() if ln.strip()]
    for need in (COVER_A.upper(), COVER_B.upper(), COVER_C):
        if need not in lines0:
            errs.append(f"PDF cover: '{need}' is not on a line of its own (orphaned word?)")
    for i, t in enumerate(flat_):
        if i > 0 and "Brian Lotze · Hudson Beat" not in t:
            errs.append(f"page {i + 1}: footer missing")
        if i > 0 and not re.search(rf"Page {i + 1}\s*$", t.strip()):
            errs.append(f"page {i + 1}: footer page number wrong (overflow?)")
    for m in re.finditer(r"\b(?:pages?|pp?\.) (\d+)(?:–(\d+))?", " ".join(flat_[1:])):
        for g in m.groups():
            if g and not 1 <= int(g) <= len(PLAN):
                errs.append(f"cross-reference to page {g} does not exist")
    for m in re.finditer(r"\bslides? (\d+)(?:–(\d+))?", " ".join(flat_[1:])):
        for g in m.groups():
            if g and not 1 <= int(g) <= DECK_SLIDES:
                errs.append(f"PDF names deck slide {g}, which does not exist")
    t = flat_[IX["rubric"] - 1]
    if f"/ {RUBRIC_TOTAL}" not in t:
        errs.append(f"rubric: /{RUBRIC_TOTAL} total missing")
    for c, ds in RUBRIC:
        for d in ds:
            if not all(w in t for w in re.findall(r"[A-Za-z]{6,}", d)):
                errs.append(f"rubric text missing: {d[:50]}")
    for pk, lst in (("proj2", REQ1 + REQ2), ("proj4", [c for _, c in CHECK] + TIPS), ("proj1", GOALS + [b for _, b in VOCAB]),
                    ("proj3", [b for _, b in STEPS1 + STEPS2] + [b for _, b in STEP2_TYPES]), ("tn2", [b for _, b in STRUGGLES]),
                    ("tn1", WHAT), ("post1", [b for _, b, _ in POST_STEPS] + list(TYPE_DEF.values()))):
        tt = flat_[IX[pk] - 1].replace("- ", "")
        for it in lst:
            if not all(w in tt for w in re.findall(r"[A-Za-z]{6,}", it)):
                errs.append(f"{pk}: owner item missing: {it[:50]}")
    k2 = flat_[IX["key2"] - 1]
    for i, q in enumerate(QUIZ):
        if classify(q) not in k2:
            errs.append(f"answer key: quiz {i + 1} answer missing")
    for t_ in TYPES:
        if any(TRANS[t_][ln] not in k2.replace(" ", "") for ln in lanes_of(TRANS[t_])):
            errs.append(f"answer key: the {t_} steps are not printed")
    for d in DAYS:
        if f"DAY {d['n']}" not in flat_[IX["pacing"] - 1]:
            errs.append(f"pacing page missing Day {d['n']}")
    st_ = flat_[IX["stems"] - 1] + flat_[IX["howto"] - 1]
    if "No stems" not in st_ or "not included" not in " ".join(flat_):
        errs.append("the kit must say clearly that no stems are included")
    ptxt = [pg.extract_text() for pg in PdfReader(PREVIEW_PDF).pages]
    if sum("PREVIEW" in t for t in ptxt) != len(PREVIEW_PICKS):
        errs.append("preview PDF: watermark missing on some pages")
    blobs = {"resource PDF": pdf_text(PDF), "preview PDF": pdf_text(PREVIEW_PDF), "deck PDF": pdf_text(DECK_PDF), "deck pptx": alldeck,
             "zip names": " ".join(names)}
    with zipfile.ZipFile(DECK) as z:
        blobs["deck XML"] = " ".join(z.read(n).decode("utf8", "ignore") for n in z.namelist() if n.endswith(".xml") or n.endswith(".rels"))
    for path in (PDF, PREVIEW_PDF, DECK_PDF):
        meta = PdfReader(path).metadata or {}
        blobs[f"{os.path.basename(path)} metadata"] = " ".join(str(v) for v in meta.values())
    for fn in os.listdir(HERE):
        if fn == "QA.md":
            continue
        if fn.endswith((".md", ".json", ".py", ".txt")) or fn in ("release", "uploaded"):
            blobs[fn] = open(os.path.join(HERE, fn), encoding="utf8", errors="ignore").read()
    flat = " ".join(blobs["resource PDF"].split())
    for gone in ("allows classroom use", "never tell", "I, IV, V and vi", "are word for word"):
        if gone in flat:
            errs.append(f"resource PDF: QA round 1 wording is back ('{gone}')")
    for need in (f"white-key chord of C major: {CHORDS_USED_TEXT}", "read the library’s own terms", "“A, B, A, B, etc.”: that is an example",
                 "the Day 3 exit question", "A, A, B, A (pages"):
        if " ".join(need.split()) not in flat and need.replace("’", "'") not in flat:
            errs.append(f"resource PDF: missing disclosure/fix text '{need}'")
    if CHORDS_USED != ["I", "iii", "IV", "V", "vi"]:
        errs.append(f"stem chords changed ({CHORDS_USED}); recheck the page {IX['makestem']} sentence")
    for label, txt in blobs.items():
        low = txt.lower()
        for b in BANNED:
            if b.lower() in low:
                errs.append(f"{label}: contains banned string '{b}'")
    for p in [COVER_PNG, *PREVIEW_PNGS]:
        if Image.open(p).size != (2000, 2000):
            errs.append(f"{os.path.basename(p)}: not 2000x2000 (low-res build?)")
    tj = json.load(open(TPT_JSON))
    if tj != tpt_json():
        errs.append("tpt.json is stale; rebuild")
    if tj["price"] != PRICE or tj["license_price"] != PRICE or tj.get("tax_code") != "Other Digital Goods - No Physical Media":
        errs.append("tpt.json price/license_price/tax_code")
    if tj["grades"] != ["9th-grade", "10th-grade", "11th-grade", "12th-grade"] or tj["subjects"] != ["Music"]:
        errs.append("tpt.json grades 9-12 / subject Music")
    proven_tags = {"Halloween", "Autumn", "Winter", "Thanksgiving"}
    if not (1 <= len(tj["tags"]) <= 6 and len(tj["formats"]) <= 3) or not set(tj["tags"]) <= proven_tags or tj["formats"] != ["PDF"]:
        errs.append("tpt.json tags/formats: only proven values (tags from 001-007, format PDF)")
    for k, f in tj["files"].items():
        if not os.path.exists(os.path.join(HERE, f)):
            errs.append(f"tpt.json file {k} missing: {f}")
    if tj["pages"] != len(rd.pages) + len(prs.slides):
        errs.append("tpt.json pages != PDF pages + slides")
    up_ = open(UPLOAD).read()
    m = re.search(r"^## 1\. Title\s*\n+```\n(.+?)\n```\n\((\d+) characters", up_, re.M | re.S)
    if not m:
        errs.append("UPLOAD.md: title block or its character count not found")
    else:
        n = len(m.group(1).strip())
        if n > 80:
            errs.append(f"UPLOAD.md: title is {n} chars (max 80)")
        if int(m.group(2)) != n:
            errs.append(f"UPLOAD.md: says {m.group(2)} characters, title has {n}")
        if m.group(1).strip() != LISTING_TITLE:
            errs.append("UPLOAD.md title differs from LISTING_TITLE in gen.py")
    m2 = re.search(r"^## 2\. Description\s*\n+```\n(.+?)\n```", up_, re.M | re.S)
    if not m2 or not m2.group(1).rstrip().endswith(DISCLOSURE):
        errs.append("UPLOAD.md: description must end with the disclosure line")
    else:
        desc = m2.group(1)
        if not desc.startswith("Drum fills and transitions for high school music technology:"):
            errs.append("UPLOAD.md description: first line must lead with the search phrase (METADATA-GUIDE)")
        for need in (f"{len(PLAN)}-page", f"{DECK_SLIDES}-slide", "Logic Pro", f"/{RUBRIC_TOTAL}", "YouTube or Spotify", "GarageBand",
                     f"{DAYS_N}-day", "gain staging", "stems are not included", ITAT[0]):
            if need not in desc:
                errs.append(f"UPLOAD.md description: missing '{need}'")
        if re.search(r"  +\S|\s[,.]\s|\bteh\b|\s{2,}$", desc, re.M):
            errs.append("UPLOAD.md description: double space or stray punctuation")
        if "/15" in desc or "checked by computer" in desc or "any DAW" in desc:
            errs.append("UPLOAD.md description: wrong rubric total or an unverifiable claim")
    for need in (f"${PRICE}", f"{len(PLAN)} pages", os.path.basename(ZIP), os.path.basename(DECK), os.path.basename(DECK_PDF)):
        if need not in up_:
            errs.append(f"UPLOAD.md: missing '{need}'")
    for f in ("release", "uploaded"):
        if os.path.exists(os.path.join(HERE, f)):
            errs.append(f"{f} exists: only the Manager / uploader writes it")
    return errs


def check():
    errs = verify_music() + verify_data() + verify_labels() + verify_files()
    if errs:
        print("CHECK FAILED")
        print("\n".join(" - " + e for e in errs))
        sys.exit(1)
    print(f"CHECK OK: A and B patterns (3+ layers, kick on 1, snare on 2 and 4, {contrast(PATTERN_A, PATTERN_B)} cells of contrast); the "
          f"four transitions match the owner's definitions by classifier (fill, breakdown one element, cut {trailing_silence(TRANS['Cut']) // 4} "
          f"beats of silence, build lanes {beat_lanes(TRANS['Build'])}); quiz answers computed {[classify(q) for q in QUIZ]}; Foundation "
          f"16 measures AABA, Full Song 56 measures, 7 changes, 7 transitions, {len({t for t, _ in SONG_TRANS})} types, builds into more "
          f"energy and the cut into less; stem chords diatonic in C (case = quality), whole progressions per section, 0:40 and 2:20 at "
          f"{TEMPO} BPM; {len(LABELS)} diagram labels fit, transition labels do not overlap; owner's lists and rubric 5 x (2-5) = "
          f"/{RUBRIC_TOTAL} printed; {DAYS_N} days run 0-{MINUTES}; every one of {DECK_SLIDES} slides has exactly one day slot matching its "
          f"Day badge; {len(PLAN)} Letter pages, no overflow or dead space, footers and valid page/slide references; cover lines whole; "
          f"preview 4 watermarked pages; deck 16:9, notes on every slide, text in bounds, no overlapping text, {len(DECK_PICS)} drawn PNGs, "
          f"PDF copy {DECK_SLIDES} pages; only Oswald/Source Sans 3, embedded; no banned strings; zip current, no audio; images 2000x2000; "
          f"tpt.json proven values; UPLOAD.md title {len(LISTING_TITLE)} chars, SEO opening and disclosure.")


if __name__ == "__main__":
    if "--fallback" in sys.argv:
        deck_pdf(DECK, os.path.join(BUILD, "deck-fallback.pdf"), fonts=False)
        print("wrote build/deck-fallback.pdf")
    elif "--check" in sys.argv:
        check()
    else:
        build()
